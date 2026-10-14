#!/usr/bin/env python3
"""large_test.py — 大量の入力で、仕事が1つのエージェントの文脈に収まる大きさに割られるか・途中から続けられるかを確かめる。

  python3 evals/large_test.py -w <fixture を描画済みの作業フォルダ>            # 既定は 43 倍（約 600 ページ）
  python3 evals/large_test.py -w <作業> --copies 100

架空データ（evals/make_fixture.py → inventory.py → render_pages.py）の機械抽出と期待値のカードを、資料 ID を変えて何倍にも複製し、
600 ページ規模の作業フォルダを機械で作る（LLM は使わない。画像は複製しない）。確かめること:
  1 読取の束が、どれも予算（見込みトークン）に収まる。束の一覧の表示が長くならない
  2 材料（カードの文字数）が上限を超える分類を、点検が FAIL にする。plan-flows が節ごとの大きさを出す
  3 同じ種類の FAIL が何百件あっても、点検の画面に出る行数は増えない（全件は --report のファイル）
  4 範囲を決めないページ1行の要約（digest）は、章ごとの集計に切り替わる
  5 途中でカードが欠けても、status.py が残りを出し、束が残りのページだけになる。読み残しは半分の予算で束ね直される
  6 フローが百本を超えても、業務マトリクス一覧を書き出せる（1本が、作業の行と処理時間の行の 2 行。作業のノードがすべてセルになる）
"""
import argparse
import json
import pathlib
import shutil
import subprocess
import sys
import tempfile

SKILL = pathlib.Path(__file__).resolve().parent.parent
S = SKILL / "scripts"


def run(script, *args):
    r = subprocess.run([sys.executable, str(S / script), *map(str, args)], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def build(base, out, copies):
    """架空データを copies 倍にした作業フォルダを作る"""
    exp = SKILL / "evals" / "fixture_expected"
    m = json.loads((base / "manifest.json").read_text(encoding="utf-8"))
    docs, pages = [], 0
    for k in range(copies):
        for d in m["docs"]:
            new = f"M{k * len(m['docs']) + int(d['doc_id'][1:]):03d}"
            nd = dict(d, doc_id=new, rel_path=f"{k:03d}/{d['rel_path']}", folder=[f"{k:03d}"] + list(d["folder"]), page_list=[])
            for p in d["page_list"]:
                pj = json.loads((base / p["json"]).read_text(encoding="utf-8"))
                pj["doc_id"] = new
                dst = out / "pages" / new / pathlib.Path(p["json"]).name
                dst.parent.mkdir(parents=True, exist_ok=True)
                dst.write_text(json.dumps(pj, ensure_ascii=False), encoding="utf-8")
                nd["page_list"].append(dict(p, json=f"pages/{new}/{dst.name}", png=f"pages/{new}/{dst.stem}.png" if p.get("png") else None))
                card = exp / "cards" / d["doc_id"] / dst.name
                if card.exists():
                    c = json.loads(card.read_text(encoding="utf-8"))
                    c["doc_id"] = new
                    cd = out / "cards" / new / dst.name
                    cd.parent.mkdir(parents=True, exist_ok=True)
                    cd.write_text(json.dumps(c, ensure_ascii=False, indent=1), encoding="utf-8")
                pages += 1
            docs.append(nd)
    (out / "manifest.json").write_text(json.dumps(dict(m, docs=docs), ensure_ascii=False), encoding="utf-8")
    allp = [{"doc_id": d["doc_id"], "pages": [p["page"] for p in d["page_list"]]} for d in docs]
    (out / "structure").mkdir(exist_ok=True)
    (out / "structure" / "hierarchy.json").write_text(json.dumps({
        "is_hierarchical": "no", "basis": ["inferred"], "rationale": "合成データ", "docs": [{"doc_id": d["doc_id"], "role": "other"} for d in docs],
        "tree": [{"id": "H1", "level": 1, "title": "全部", "kind": "process", "source": "inferred", "children": [],
                  "refs": [dict(x, rel="describes") for x in allp]}],
        "doc_relations": [], "sequence": [], "toc_flow_list": {"found": False, "source": [], "entries": []}, "questions": []}, ensure_ascii=False), encoding="utf-8")
    shutil.copy(exp / "glossary.json", out / "glossary.json")
    (out / "taxonomy.json").write_text(json.dumps({
        "basis": "llm", "rationale": "合成データ。全ページを1つの分類に入れて、材料の上限を確かめる",
        "categories": [{"id": "F01", "number": "", "name": "全部", "parent": None, "hierarchy_id": "H1", "level": 1, "basis": "llm",
                        "pages": allp, "make_flow": True},
                       {"id": "F02", "number": "", "name": "小さい分類", "parent": None, "hierarchy_id": None, "level": 1, "basis": "llm",
                        "pages": [allp[0]], "make_flow": True}],
        "unassigned_pages": [], "questions": [{"id": "q1", "severity": "high", "category": "分類", "effect": "フローの範囲", "question": "この分け方でよいか。",
                                               "assumption": "", "sources": [], "ask": ""}]}, ensure_ascii=False), encoding="utf-8")
    return docs, pages


def many_flows(out, docs, per):
    """期待値のフロー3本を、資料の組（per 件ずつ）ごとに複製する。分類は、組ごとの見出しの下に3つ。戻り値は (フローの数, 作業のノードの数)"""
    exp = SKILL / "evals" / "fixture_expected"
    src = [json.loads(f.read_text(encoding="utf-8")) for f in sorted((exp / "flows").glob("*.flow.json"))]
    cats, n_steps = [], 0
    shutil.rmtree(out / "flows", ignore_errors=True)
    (out / "flows").mkdir()
    for k in range(len(docs) // per):
        top = f"F{k + 1:03d}"
        new_doc = lambda d: f"M{k * per + int(d[1:]):03d}"
        new_flow = lambda i: f"{top}-{int(i[1:])}"
        cats.append({"id": top, "number": "", "name": f"第 {k + 1} 組", "parent": None, "hierarchy_id": None, "level": 1, "basis": "llm", "pages": [],
                     "make_flow": False})
        for fl in src:
            fl = json.loads(json.dumps(fl))
            fl.update(flow_id=new_flow(fl["flow_id"]), parent=top, level=2, basis="llm")
            for n in fl["nodes"]:
                for x in n.get("sources") or []:
                    x["doc_id"] = new_doc(x["doc_id"])
                if n.get("link_flow"):
                    n["link_flow"] = new_flow(n["link_flow"])
            for key in ("unused_pages", "excluded_items"):
                for x in fl.get(key) or []:
                    x["doc_id"] = new_doc(x["doc_id"])
            for q in fl.get("questions") or []:
                for x in q.get("sources") or []:
                    x["doc_id"] = new_doc(x["doc_id"])
            (out / "flows" / f"{fl['flow_id']}.flow.json").write_text(json.dumps(fl, ensure_ascii=False), encoding="utf-8")
            n_steps += sum(n["type"] not in ("start", "end", "file") for n in fl["nodes"])
            cats.append({"id": fl["flow_id"], "number": fl.get("number", ""), "name": fl["name"], "parent": top, "hierarchy_id": None, "level": 2,
                         "basis": "llm", "pages": [], "make_flow": True})
    tax = json.loads((out / "taxonomy.json").read_text(encoding="utf-8"))
    (out / "taxonomy.json").write_text(json.dumps(dict(tax, categories=cats), ensure_ascii=False), encoding="utf-8")
    return len(cats) - len(docs) // per, n_steps


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-w", "--work", required=True)
    ap.add_argument("--copies", type=int, default=43)
    ap.add_argument("--budget", type=int, default=100000)
    a = ap.parse_args()
    base = pathlib.Path(a.work).resolve()
    tmp = pathlib.Path(tempfile.mkdtemp(prefix="asisflow_large_"))
    res = []

    def check(name, ok, detail=""):
        res.append(ok)
        print(f"{'OK  ' if ok else 'NG  '} {name}" + (f" — {detail}" if detail else ""))
    try:
        W = tmp / "work"
        W.mkdir()
        docs, n_pages = build(base, W, a.copies)
        print(f"合成した入力: 資料 {len(docs)} 件・{n_pages} ページ・カード {len(list(W.glob('cards/*/p*.json')))} 枚\n")

        code, out = run("pages.py", "batches", "-w", W, "--all", "--budget", a.budget)
        bs = json.loads((W / "batches.json").read_text(encoding="utf-8"))
        multi = [b for b in bs if len(b["pages"]) > 1]
        check("読取の束が、どれも予算に収まる", bool(bs) and all(b["est_tokens"] <= a.budget for b in multi) and sum(len(b["pages"]) for b in bs) == n_pages,
              f"{len(bs)} 束・最大 {max(b['est_tokens'] for b in bs):,} トークン・1束 最大 {max(len(b['pages']) for b in bs)} ページ")
        check("束の一覧の表示が長くならない", len(out.splitlines()) <= 50, f"{len(out.splitlines())} 行（全件は batches.json）")

        code, out = run("check_flow.py", "-w", W, "--stage", "structure")
        check("材料が大きすぎる分類を FAIL にする", code == 1 and "F01「全部」の材料が大きすぎる" in out and "F02「小さい分類」の材料" not in out)
        code, out = run("pages.py", "plan-flows", "-w", W)
        check("plan-flows が、大きすぎる分類と節ごとの大きさを出す（行数は増えない）", "★大きすぎる" in out and "そのまま1エージェントで" in out
              and len(out.splitlines()) <= 50, f"{len(out.splitlines())} 行")

        for f in sorted(W.glob("cards/*/p*.json")):                  # 同じ種類の誤りを全カードに仕込む
            c = json.loads(f.read_text(encoding="utf-8"))
            c["header"] = "柱"
            f.write_text(json.dumps(c, ensure_ascii=False), encoding="utf-8")
        code, out = run("check_extract.py", "-w", W, "--report", W / "out" / "チェック_読取.txt")
        full = (W / "out" / "チェック_読取.txt").read_text(encoding="utf-8").splitlines()
        check("同じ種類の FAIL が何百件あっても、画面の行数は増えない", code == 1 and len(out.splitlines()) <= 40 and len(full) > n_pages,
              f"画面 {len(out.splitlines())} 行／ファイル {len(full)} 行")
        for f in sorted(W.glob("cards/*/p*.json")):
            c = json.loads(f.read_text(encoding="utf-8"))
            c.pop("header")
            f.write_text(json.dumps(c, ensure_ascii=False), encoding="utf-8")

        code, out = run("pages.py", "digest", "-w", W)
        check("範囲を決めない digest は、集計に切り替わり、行数が増えない", "範囲を決める" in out and len(out.splitlines()) <= 50, f"{len(out.splitlines())} 行")
        code, out = run("pages.py", "outline", "-w", W, "--only", ",".join(d["doc_id"] for d in docs[:4]))
        check("資料を選べば、章ごとの集計が出る", "第1章 申請と承認" in out and len(out.splitlines()) <= 40, f"{len(out.splitlines())} 行")
        code, out = run("pages.py", "digest", "-w", W, "--only", docs[0]["doc_id"])
        check("範囲を決めた digest は、ページ1行で出る", len(out.splitlines()) == len(docs[0]["page_list"]))

        gone = sorted(W.glob("cards/*/p*.json"))[100:140]            # 途中でカードが欠けた（束が途中で止まった）
        for f in gone:
            f.unlink()
        left_doc = gone[0].parent.name
        left = [int(f.stem[1:]) for f in gone if f.parent.name == left_doc]
        (W / "reports").mkdir(exist_ok=True)
        (W / "reports" / f"read_{left_doc}_{left[0]:04d}.json").write_text(json.dumps({
            "doc_id": left_doc, "pages_done": [], "pages_left": left, "stopped_because": "文脈の残りが少ない", "unreadable": 0, "uncertain": 0,
            "hints": []}, ensure_ascii=False), encoding="utf-8")
        code, out = run("status.py", "-w", W, "--fast")
        check("status.py が、残りのページと次のコマンドを出す", f"残り {len(gone)}" in out and "pages.py batches" in out, out.splitlines()[4][:70])
        code, out = run("pages.py", "batches", "-w", W)
        bs = json.loads((W / "batches.json").read_text(encoding="utf-8"))
        check("束が、残りのページだけになる", sum(len(b["pages"]) for b in bs) == len(gone), f"{len(bs)} 束・{sum(len(b['pages']) for b in bs)} ページ")
        retry = [b for b in bs if b["retry"]]
        check("読み残しは、半分の予算で束ね直される", bool(retry) and all(b["est_tokens"] <= a.budget // 2 or len(b["pages"]) == 1 for b in retry),
              f"{len(retry)} 束・最大 {max([b['est_tokens'] for b in retry] or [0]):,} トークン")

        n_flows, n_steps = many_flows(W, docs, len(docs) // a.copies)
        code, out = run("build_matrix.py", "-w", W)
        rows = cells = 0
        if code == 0:
            from openpyxl import load_workbook
            ws = load_workbook(W / "out" / "業務マトリクス一覧.xlsx", read_only=True)["業務マトリクス"]
            for k, r in enumerate(ws.iter_rows(min_row=2, min_col=10, values_only=True)):
                if k % 2 == 0:                                       # 作業の行（そのすぐ下は、処理時間の行）。作業の列は 10 列目から（左に、件数・処理時間の合計）
                    rows, cells = rows + 1, cells + sum(v is not None for v in r)
        check("フローが百本を超えても、業務マトリクス一覧を書き出せる", code == 0 and rows == n_flows and cells == n_steps,
              f"{rows} 行・作業 {cells} セル" if code == 0 else out[-200:])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("\nlarge_test " + ("OK" if all(res) else "NG"))
    sys.exit(0 if all(res) else 1)


if __name__ == "__main__":
    main()
