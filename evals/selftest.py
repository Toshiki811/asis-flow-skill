#!/usr/bin/env python3
"""selftest.py — チェックのスクリプトが、仕込んだ誤りをちゃんと FAIL にするかを確かめる。

  python3 evals/make_fixture.py <入力>                                   # 架空の入力一式
  python3 scripts/inventory.py <入力> -w <作業> && python3 scripts/render_pages.py -w <作業>
  python3 evals/selftest.py -w <作業>

<作業> は、上の手順で fixture を描画済みの作業フォルダ。ここに evals/fixture_expected（カード・階層・用語・分類・フロー）を写し、
(1) 正しいデータで FAIL 0 になること、(2) 納品物の表記が決まりどおりであること（刷られたページ番号・シート名とセル）、
(3) 誤りを1つずつ仕込んだ写しで、狙った FAIL が出ること、(4) 部品が単体で正しく動くこと、を確かめる。
スクリプトを直したら、これを通してから使う。
"""
import argparse
import json
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

SKILL = pathlib.Path(__file__).resolve().parent.parent
S = SKILL / "scripts"
sys.path.insert(0, str(S))
from common import q_targets  # noqa: E402


def run(script, work, *args):
    r = subprocess.run([sys.executable, str(S / script), "-w", str(work), *args], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def edit(path, fn):
    d = json.loads(path.read_text(encoding="utf-8"))
    fn(d)
    path.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")


def node(fl, nid):
    return next(n for n in fl["nodes"] if n["id"] == nid)


def drop_error_branch(fl):
    gone = {"D-002", "P-009", "P-010", "P-011", "P-012"}
    fl["nodes"] = [n for n in fl["nodes"] if n["id"] not in gone]
    fl["edges"] = [e for e in fl["edges"] if e["from"] not in gone and e["to"] not in gone] + [{"from": "P-008", "to": "S-001"}]
    fl["questions"] = [q for q in fl["questions"] if not set(q_targets(q)) & gone and not (q.get("edge") and set(q["edge"]) & gone)]


def exclude_one_item(fl):
    """M02 p.2 を出典にしたノードを全部外し、そのページの手順4つのうち1つだけを「対象外」と書く（残りの3つは、どこにも無い）"""
    drop_error_branch(fl)
    fl.setdefault("excluded_items", []).append({"doc_id": "M02", "page": 2, "item": "s1", "reason": "エラー時の対応は別フローで描く"})


def patchwork(c):
    """原典の語だけで作った継ぎはぎ（「課長が承認する」の主語を、次の文の「部長」にすり替える）"""
    c["steps"][1].update(actor="部長", quote="申請金額が5万円未満の場合は部長が承認する。")


def swap_cards(w, doc, a, b):
    """カード2枚の中身を入れ替える（ページ画像を1つずれて読んだ状態）。doc_id・page はファイルの場所に合わせる"""
    pa, pb = w / f"cards/{doc}/p{a:04d}.json", w / f"cards/{doc}/p{b:04d}.json"
    ca, cb = json.loads(pa.read_text(encoding="utf-8")), json.loads(pb.read_text(encoding="utf-8"))
    ca["page"], cb["page"] = b, a
    pa.write_text(json.dumps(cb, ensure_ascii=False), encoding="utf-8")
    pb.write_text(json.dumps(ca, ensure_ascii=False), encoding="utf-8")


def shift_sources(w, doc, delta):
    """全フローで、ある資料を出典にしたページを delta だけずらす（刷られた番号や目次の番号で出典を書いた状態）"""
    for f in (w / "flows").glob("*.flow.json"):
        def fn(fl):
            for n in fl["nodes"]:
                for s_ in n.get("sources") or []:
                    if s_["doc_id"] == doc:
                        s_["page"] += delta
        edit(f, fn)


def drop_refs(h, doc):
    def walk(ns):
        for n in ns:
            n["refs"] = [r for r in n["refs"] if r["doc_id"] != doc]
            walk(n["children"])
    walk(h["tree"])


def tree_ref(h, node, doc):
    def walk(ns):
        for n in ns:
            if n["id"] == node:
                return next(r for r in n["refs"] if r["doc_id"] == doc)
            r = walk(n["children"])
            if r:
                return r
    return walk(h["tree"])


def relabel_sheet(m):
    p = next(p for d in m["docs"] if d["doc_id"] == "M03" for p in d["page_list"] if p["page"] == 1)
    p["label"] = "シート不明(1)"
    p.pop("sheet", None)


def make_book(w, shift=0):
    """表紙を数えない冊子の作業フォルダを作る（表紙・目次・4つの節が2ページずつ。刷られた番号はファイルの番号 − 1）。
    shift は、分類のページの割当を何ページずらすか。−1 は、目次の番号（刷られた番号）をファイルの番号としてそのまま使った状態。
    節が2ページ以上あると、ずれていても目次の指すページは割当に入ったままなので、見出しの位置でしか見つからない"""
    names, plist = ["受付", "審査", "支払", "保管"], []

    def page(i, ptype, title, text):
        for sub, body in (("pages", {"doc_id": "M01", "page": i, "text": text,
                                     "printed_candidates": [{"n": str(i - 1), "x": 0.5, "y": 0.96}] if i > 1 else []}),
                          ("cards", {"doc_id": "M01", "page": i, "printed_page": str(i - 1) if i > 1 else None, "page_type": ptype,
                                     "title": title, "section_path": [title], "steps": [], "decisions": []})):
            f = w / sub / "M01" / f"p{i:04d}.json"
            f.parent.mkdir(parents=True, exist_ok=True)
            f.write_text(json.dumps(body, ensure_ascii=False), encoding="utf-8")
        plist.append({"page": i, "label": f"p.{i}", "kind": "page", "png": None, "json": f"pages/M01/p{i:04d}.json"})
    page(1, "cover", "事務の手引", ["事務の手引"])
    page(2, "toc", "目次", ["目次"] + [f"{k} {n} …… {2 * k}" for k, n in enumerate(names, 1)])
    for k, n in enumerate(names, 1):
        page(2 * k + 1, "procedure", f"{k} {n}", [f"{k} {n}", "書類を確かめる。"])
        page(2 * k + 2, "procedure", "記録", ["記録を残す。"])
    put = lambda rel, body: (w / rel).parent.mkdir(parents=True, exist_ok=True) or (w / rel).write_text(json.dumps(body, ensure_ascii=False), encoding="utf-8")
    put("manifest.json", {"version": 1, "inputs": [], "docs": [{"doc_id": "M01", "name": "手引.pdf", "rel_path": "手引.pdf", "folder": [], "kind": "pdf",
                                                              "pages": len(plist), "page_list": plist, "duplicate_of": None}]})
    own = lambda k: [2 * k + 1 + shift, 2 * k + 2 + shift]
    put("structure/hierarchy.json", {
        "is_hierarchical": "yes", "basis": ["toc"], "rationale": "目次の4つの節", "docs": [{"doc_id": "M01", "role": "manual"}],
        "tree": [{"id": f"H{k}", "level": 1, "title": n, "kind": "process", "source": "toc", "children": [],
                  "refs": [{"doc_id": "M01", "pages": own(k), "rel": "describes"}]} for k, n in enumerate(names, 1)],
        "doc_relations": [], "sequence": [], "questions": [],
        "toc_flow_list": {"found": True, "source": [{"doc_id": "M01", "pages": [2]}],
                          "entries": [{"number": str(k), "title": n, "level": 1, "page": 2 * k} for k, n in enumerate(names, 1)]}})
    put("glossary.json", {"terms": [], "questions": []})
    put("taxonomy.json", {"basis": "toc", "rationale": "目次のとおり", "unassigned_pages": [], "questions": [], "categories": [
        {"id": f"F{k:02d}", "number": str(k), "name": n, "parent": None, "hierarchy_id": f"H{k}", "level": 1, "basis": "toc",
         "toc_ref": {"doc_id": "M01", "page": 2, "text": f"{k} {n}"}, "pages": [{"doc_id": "M01", "pages": own(k)}], "make_flow": True}
        for k, n in enumerate(names, 1)]})


def put(w, rel, body):
    f = w / rel
    f.parent.mkdir(parents=True, exist_ok=True)
    f.write_text(json.dumps(body, ensure_ascii=False), encoding="utf-8")


def add_doc(w, doc, name, folder, pages):
    """合成の資料を1つ足す（機械抽出・カード・manifest）。pages は [(種別, 見出し, [行…])]。番号の印字は無い"""
    plist = []
    for i, (ptype, title, text) in enumerate(pages, 1):
        put(w, f"pages/{doc}/p{i:04d}.json", {"doc_id": doc, "page": i, "text": text, "printed_candidates": []})
        put(w, f"cards/{doc}/p{i:04d}.json", {"doc_id": doc, "page": i, "printed_page": None, "page_type": ptype, "title": title,
                                              "section_path": [title], "steps": [], "decisions": []})
        plist.append({"page": i, "label": f"p.{i}", "kind": "page", "png": None, "json": f"pages/{doc}/p{i:04d}.json"})
    mf = w / "manifest.json"
    m = json.loads(mf.read_text(encoding="utf-8")) if mf.exists() else {"version": 1, "inputs": [], "docs": []}
    m["docs"].append({"doc_id": doc, "name": name, "rel_path": "/".join(folder + [name]), "folder": folder, "kind": "pdf", "pages": len(plist),
                      "page_list": plist, "duplicate_of": None})
    put(w, "manifest.json", m)


def tiny_flow(w, cid, name, doc, quote, lane="担当者", **extra):
    """作業1つだけのフロー（縦軸の列の出方を見るためのもの）"""
    put(w, f"flows/{cid}.flow.json", dict({
        "flow_id": cid, "name": name, "number": "", "parent": None, "level": 1, "basis": "llm", "summary": "", "start_condition": "申請があったとき",
        "end_condition": "記録したとき", "lanes": [{"id": "L1", "name": lane, "kind": "role"}],
        "nodes": [{"id": "start", "type": "start", "lane": "L1", "label": "開始"},
                  {"id": "P-001", "type": "process", "lane": "L1", "label": quote.rstrip("。"),
                   "sources": [{"doc_id": doc, "page": 1, "quote": quote, "from_image": False}]},
                  {"id": "end", "type": "end", "lane": "L1", "label": "終了"}],
        "edges": [{"from": "start", "to": "P-001"}, {"from": "P-001", "to": "end"}], "unused_pages": [], "excluded_items": [], "questions": []}, **extra))


def make_mixed(w):
    """目次のある冊子（make_book）に、業務でない目次項目（用語集）と、目次に無い業務の手順書（M02）を足した作業フォルダ（混在）"""
    make_book(w)
    for i, (ptype, title, text) in ((11, ("reference", "5 用語集", ["5 用語集", "受付とは、書類を受け取ることをいう。"])),):
        put(w, f"pages/M01/p{i:04d}.json", {"doc_id": "M01", "page": i, "text": text, "printed_candidates": [{"n": str(i - 1), "x": 0.5, "y": 0.96}]})
        put(w, f"cards/M01/p{i:04d}.json", {"doc_id": "M01", "page": i, "printed_page": str(i - 1), "page_type": ptype, "title": title,
                                            "section_path": [title], "steps": [], "decisions": []})
    edit(w / "manifest.json", lambda m: (m["docs"][0]["page_list"].append({"page": 11, "label": "p.11", "kind": "page", "png": None,
                                                                           "json": "pages/M01/p0011.json"}), m["docs"][0].update(pages=11)))
    edit(w / "pages/M01/p0002.json", lambda pj: pj["text"].append("5 用語集 …… 10"))
    add_doc(w, "M02", "備品貸出の手順.pdf", ["別冊"], [("procedure", "備品の貸出", ["備品の貸出", "総務担当は、貸出簿に記入する。"]),
                                                    ("procedure", "返却", ["返却を受けたら、貸出簿に返却日を記入する。"])])

    def hier(h):
        h["docs"].append({"doc_id": "M02", "role": "procedure"})
        h["tree"] += [{"id": "H5", "level": 1, "title": "用語集", "kind": "process", "source": "toc", "children": [],
                       "refs": [{"doc_id": "M01", "pages": [11], "rel": "describes"}]},
                      {"id": "H6", "level": 1, "title": "備品の貸出", "kind": "process", "source": "heading", "children": [],
                       "refs": [{"doc_id": "M02", "pages": [1, 2], "rel": "describes"}]}]
        h["basis"] = ["toc", "filename"]

    def tax(t):
        t.update(basis="mixed", rationale="目次の4つの節をそのまま分類にし、目次に無い別冊の手順書を1つの分類として足した。")
        t["categories"] += [
            {"id": "F05", "number": "5", "name": "用語集", "parent": None, "hierarchy_id": "H5", "level": 1, "basis": "toc",
             "toc_ref": {"doc_id": "M01", "page": 2, "text": "5 用語集"}, "pages": [], "make_flow": False, "out_of_scope": "用語の説明で、業務の手順ではない"},
            {"id": "F06", "number": "", "name": "備品の貸出", "parent": None, "hierarchy_id": "H6", "level": 1, "basis": "llm",
             "name_ref": {"kind": "heading", "doc_id": "M02", "page": 1, "text": "備品の貸出"}, "pages": [{"doc_id": "M02", "pages": [1, 2]}], "make_flow": True}]
        t["unassigned_pages"] = [{"doc_id": "M01", "page": 1, "reason": "表紙"}, {"doc_id": "M01", "page": 2, "reason": "目次"},
                                 {"doc_id": "M01", "page": 11, "reason": "用語の説明で、業務の手順ではない"}]
        t["questions"] = [{"id": "q1", "severity": "high", "category": "分類", "effect": "フローの範囲", "flow": "F06", "question": "別冊の「備品の貸出」を、1つの業務分類として足してよいですか。目次のどれかに含めますか。",
                           "assumption": "目次に無い業務として、末尾に足した", "sources": [{"doc_id": "M02", "page": 1}], "ask": ""}]
    edit(w / "structure/hierarchy.json", hier)
    edit(w / "taxonomy.json", tax)
    tiny_flow(w, "F06", "備品の貸出", "M02", "総務担当は、貸出簿に記入する。")


def make_loose(w):
    """目次の無い、ばらばらの手順書2つの作業フォルダ。フォルダ名を上の段にした分類と、原典に名前の無い分類"""
    add_doc(w, "M01", "備品貸出の手順.pdf", ["総務"], [("procedure", "備品の貸出", ["備品の貸出", "総務担当は、貸出簿に記入する。"])])
    add_doc(w, "M02", "メモ.pdf", [], [("procedure", "", ["立て替えた経費は、月末までに精算書を出す。"])])
    put(w, "structure/hierarchy.json", {
        "is_hierarchical": "partial", "basis": ["folder"], "rationale": "フォルダ「総務」の下に手順書が1つ。もう1つはフォルダに入っていない",
        "docs": [{"doc_id": "M01", "role": "procedure"}, {"doc_id": "M02", "role": "other"}],
        "tree": [{"id": "H1", "level": 1, "title": "総務", "kind": "domain", "source": "folder", "refs": [], "children": [
                     {"id": "H1.1", "level": 2, "title": "備品の貸出", "kind": "process", "source": "heading", "children": [],
                      "refs": [{"doc_id": "M01", "pages": [1], "rel": "describes"}]}]},
                 {"id": "H2", "level": 1, "title": "経費の立替の精算", "kind": "process", "source": "inferred", "children": [],
                  "refs": [{"doc_id": "M02", "pages": [1], "rel": "describes"}]}],
        "doc_relations": [], "sequence": [], "questions": [], "toc_flow_list": {"found": False, "source": [], "entries": []}})
    put(w, "glossary.json", {"terms": [], "questions": []})
    put(w, "taxonomy.json", {"basis": "llm", "rationale": "手順書ごとに、きっかけから成果までを1つの分類にした。", "unassigned_pages": [], "categories": [
        {"id": "F01", "number": "", "name": "総務", "parent": None, "hierarchy_id": "H1", "level": 1, "basis": "llm",
         "name_ref": {"kind": "title", "doc_id": "M01", "text": "総務"}, "pages": [], "make_flow": False},
        {"id": "F01-1", "number": "", "name": "備品の貸出", "parent": "F01", "hierarchy_id": "H1.1", "level": 2, "basis": "llm",
         "name_ref": {"kind": "heading", "doc_id": "M01", "page": 1, "text": "備品の貸出"}, "pages": [{"doc_id": "M01", "pages": [1]}], "make_flow": True},
        {"id": "F02", "number": "", "name": "経費の立替の精算", "parent": None, "hierarchy_id": "H2", "level": 1, "basis": "llm",
         "pages": [{"doc_id": "M02", "pages": [1]}], "make_flow": True}],
        "questions": [{"id": "q1", "severity": "high", "category": "分類", "effect": "フローの範囲", "question": "手順書ごとに分類しました。この分け方でよいですか。部門別にまとめる案もあります。",
                       "assumption": "手順書ごとの2分類", "sources": [], "ask": ""}]})
    tiny_flow(w, "F01-1", "備品の貸出", "M01", "総務担当は、貸出簿に記入する。", parent="F01", level=2)
    tiny_flow(w, "F02", "経費の立替の精算", "M02", "立て替えた経費は、月末までに精算書を出す。")


def cat(t, cid):
    return next(c for c in t["categories"] if c["id"] == cid)


def move_before(t, cid, other):
    """分類 cid を、分類 other の前に動かす"""
    c = cat(t, cid)
    t["categories"].remove(c)
    t["categories"].insert(t["categories"].index(cat(t, other)), c)


AXIS_CASES = [
    # (名前, もとにする作業フォルダ, 仕込み(work), 出るはずの文言)。どれも check_flow.py --stage structure
    ("業務でない目次項目に、対象外の理由が無い", "mixed", lambda w: edit(w / "taxonomy.json", lambda t: cat(t, "F05").pop("out_of_scope")), "対象外の理由（out_of_scope）を書く"),
    ("対象外の分類に、ページを割り当てたまま", "mixed", lambda w: edit(w / "taxonomy.json", lambda t: (
        cat(t, "F05").update(pages=[{"doc_id": "M01", "pages": [11]}]), t["unassigned_pages"].pop())), "対象外なのに、ページが割り当てられている"),
    ("目次の項目の扱いが決まっていない（分類が無い）", "mixed", lambda w: edit(w / "taxonomy.json", lambda t: t["categories"].remove(cat(t, "F05"))), "の扱いが決まっていない"),
    ("足した分類を、目次の項目の間に置く", "mixed", lambda w: edit(w / "taxonomy.json", lambda t: move_before(t, "F06", "F03")), "より前にある"),
    ("名前の根拠の文字が、そのページに無い", "mixed", lambda w: edit(w / "taxonomy.json", lambda t: cat(t, "F06")["name_ref"].update(text="備品の管理")), "見当たらない"),
    ("名前が、根拠の文字に含まれていない", "mixed", lambda w: edit(w / "taxonomy.json", lambda t: cat(t, "F06").update(name="備品管理業務")), "に含まれていない"),
    ("目次の文字（toc_ref）が、目次のページに無い", "mixed", lambda w: edit(w / "taxonomy.json", lambda t: cat(t, "F02")["toc_ref"].update(text="2 審査と承認")), "toc_ref の文字"),
    ("目次の資料でない資料を、根拠なしで目次の項目に結び付ける", "mixed", lambda w: edit(w / "structure/hierarchy.json", lambda h: h["tree"][2]["refs"].append(
        {"doc_id": "M02", "pages": [2], "rel": "describes"})), "目次の資料でない M02 を結び付けたのに"),
    ("原典の手がかりの無い上の段を作る", "loose", lambda w: edit(w / "taxonomy.json", lambda t: cat(t, "F01").pop("name_ref")), "原典に無い上の段は作らない"),
    ("資料名に無い文字を、名前の根拠にする", "loose", lambda w: edit(w / "taxonomy.json", lambda t: (cat(t, "F01").update(name="管理"),
                                                                                    cat(t, "F01")["name_ref"].update(text="管理"))), "見当たらない"),
]


def axis_tests(tmp):
    """目次あり（対象外の項目つき）・混在・目次なしの合成データで、縦軸の検査と、業務マトリクスの左の列を確かめる"""
    from openpyxl import load_workbook
    out, base = [], {"mixed": tmp / "mixed", "loose": tmp / "loose"}
    make_mixed(base["mixed"])
    make_loose(base["loose"])
    sheets = {}
    for name, w in base.items():
        code, text = run("check_flow.py", w, "--stage", "structure")
        out.append((f"縦軸: 正しい{'混在' if name == 'mixed' else '目次なし'}のデータは FAIL 0", code == 0 and "0 FAIL" in text))
        if code != 0:
            print(text[-1200:])
        c1, _t = run("build_matrix.py", w)
        c2, t2 = run("check_flow.py", w, "--stage", "matrix")
        out.append((f"縦軸: {'混在' if name == 'mixed' else '目次なし'}のデータでも、業務マトリクス一覧を書き出せて、点検が通る", c1 == 0 and c2 == 0 and "0 FAIL" in t2))
        wb = load_workbook(w / "out" / "業務マトリクス一覧.xlsx", rich_text=True)
        sheets[name] = {ws.title: [["" if c.value is None else str(c.value) for c in r] for r in ws.iter_rows(min_row=2)] for ws in wb}
    mx, ax = sheets["mixed"]["業務マトリクス"], {r[1].strip(): r for r in sheets["mixed"]["縦軸の根拠"]}
    out.append(("縦軸（混在）: 足した分類は、根拠が「本文の見出し」で場所つき。その分類を指す確認事項が行に出る",
                len(mx) == 2 and mx[0][:2] == ["備品の貸出", "—"] and mx[1][:2] == ["備品の貸出｜処理時間（分）", "—"] and mx[0][2].startswith("本文の見出し\n") and " p.1「備品の貸出」" in mx[0][2] and "M02" not in mx[0][2] and "Q-001" in mx[0][6]))
    out.append(("縦軸（混在）: 目次の項目は、フローを描く・対象外（理由つき）、足した分類は「目次に無く足した」",
                ax["受付"][3] == "フローを描く" and ax["用語集"][3] == "対象外（業務でない）" and "業務の手順ではない" in ax["用語集"][5]
                and ax["備品の貸出"][3] == "目次に無く足した"))
    lo = {(r[1] if r[1] != "—" else r[0]): r for r in sheets["loose"]["業務マトリクス"]}
    out.append(("縦軸（目次なし）: フォルダ名が大分類になり、原典に名前の無い分類は「LLM が命名」。1 段の分類は、大分類に名前（上から詰める。中分類は「—」）",
                lo["備品の貸出"][0] == "総務" and lo["備品の貸出"][2].startswith("本文の見出し") and lo["経費の立替の精算"][:2] == ["経費の立替の精算", "—"]
                and lo["経費の立替の精算"][2] == "LLM が命名"))
    lax = sheets["loose"]["縦軸の根拠"]
    out.append(("縦軸の根拠（目次なし）: 業務の木の項目に結び付いた分類を、「目次に無く足した」の行に重ねて出さない。名前の根拠は「理由」に",
                not any(r[3] == "目次に無く足した" for r in lax) and len(lax) == 3 and any(r[1].strip() == "備品の貸出" and r[5].startswith("名前: 本文の見出し") for r in lax)))
    code, text = run("pages.py", base["mixed"], "orphans")
    out.append(("目次に無いページの一覧: 割り当て済みなら 0 件", code == 0 and "まとまり 0 件" in text))
    w = tmp / "orph"
    shutil.copytree(base["mixed"], w)
    edit(w / "taxonomy.json", lambda t: t["categories"].remove(cat(t, "F06")))
    code, text = run("pages.py", w, "orphans")
    out.append(("目次に無いページの一覧: 分類に入っていない手順書を、まとまりで出す（候補なし → 新しい分類の候補）",
                code == 0 and "M02" in text and "「備品の貸出」" in text and "候補なし" in text))
    w = tmp / "drafted"                                              # 分類の下書き: 目次と階層から、分類・目次の場所・ページの割当を機械で組む
    make_book(w)
    booked = json.loads((w / "taxonomy.json").read_text(encoding="utf-8"))
    (w / "taxonomy.json").unlink()
    edit(w / "structure/hierarchy.json", lambda h: [n.update(refs=[dict(r, pages=[]) for r in n["refs"]]) for n in h["tree"]])
    put(w, "cards/M01/p0002.json", dict(json.loads((w / "cards/M01/p0002.json").read_text(encoding="utf-8")),
                                        toc_entries=[{"level": 1, "number": str(k), "title": n, "page": 2 * k, "is_flow": True}
                                                     for k, n in enumerate(["受付", "審査", "支払", "保管"], 1)]))
    c1, _t = run("pages.py", w, "taxonomy", "--draft")
    c2, t2 = run("check_flow.py", w, "--stage", "structure")
    drafted = json.loads((w / "taxonomy.json").read_text(encoding="utf-8")) if c1 == 0 else {"categories": []}
    key = lambda t: [(c["id"], c["number"], c["name"], c["make_flow"], c["pages"]) for c in t["categories"]]
    out.append(("分類の下書き: 目次の番号（刷られた番号）からページを引いて、手で書いた分類と同じ割当になり、点検が通る",
                c1 == 0 and c2 == 0 and key(drafted) == key(booked)))
    if key(drafted) != key(booked):
        print(key(drafted), key(booked), t2[-600:])
    for name, which, mutate, expect in AXIS_CASES:
        w = tmp / "axis_case"
        if w.exists():
            shutil.rmtree(w)
        shutil.copytree(base[which], w)
        mutate(w)
        code, text = run("check_flow.py", w, "--all-lines", "--stage", "structure")
        hit = [ln for ln in text.splitlines() if ln.startswith("FAIL") and expect in ln]
        out.append((f"{name}: " + (hit[0][:100] if hit else f"FAIL「{expect}」が出ない（終了コード {code}）"), code == 1 and bool(hit)))
        if not hit:
            print("\n".join(ln for ln in text.splitlines() if ln.startswith(("FAIL", "WARN")))[:1000])
    return out


def edit_book(w, fn):
    """書き出した業務マトリクス一覧（xlsx）を、手で直した状態にする"""
    from openpyxl import load_workbook
    book = w / "out" / "業務マトリクス一覧.xlsx"
    wb = load_workbook(book, rich_text=True)
    fn(wb)
    wb.save(book)


def blank_step(wb):
    wb["業務マトリクス"].cell(row=6, column=12).value = None          # F03 の3つ目の作業（1 フロー＝2 行。作業の列は 10 列目から）


def drop_exit(wb):
    ws = wb["業務フローのノード一覧"]
    col = [c.value for c in ws[1]].index("行き先")
    c = next(c for r in ws.iter_rows(min_row=2) for c in [r[col]] if r[2].value == "F03" and "\n" in str(c.value or ""))
    c.value = str(c.value).split("\n")[0]


def drop_question(wb):
    ws = wb["業務マトリクス"]
    c = next(c for c in ws[6] if re.search(r"Q-\d+", str(c.value or "")))
    c.value = re.sub(r"Q-\d+", "", str(c.value), count=1)


def wrong_lookup(wb):
    """仮の値の式を、別の動作の種類を引く式に書き換える（F01 の 1 つ目の作業）"""
    c = wb["業務マトリクス"].cell(row=3, column=10)
    c.value = str(c.value).replace('"入力する"', '"調べる"')


def break_total(wb):
    """合計の式を、1 つ手前までの足し算にする（F01）"""
    c = wb["業務マトリクス"].cell(row=3, column=9)
    c.value = "=SUM(J3:K3)"


def redraw(w, fn):
    """書き出した業務フローの draw.io を、手で直した状態にする（fn: XML の文字 → 文字）"""
    for p in (w / "out").glob("業務フロー*.drawio"):
        p.write_text(fn(p.read_text(encoding="utf-8")), encoding="utf-8")


def drop_page(x, pid):
    return re.sub(rf'\s*<diagram [^>]*id="{pid}">.*?</diagram>', "", x, count=1, flags=re.S)


def dup_cell(x):
    """F01 のページで、2 つの出典の字のセルの id を同じにする"""
    page = re.search(r'<diagram [^>]*id="flow-F01">.*?</diagram>', x, flags=re.S).group(0)
    a, b = re.findall(r'<mxCell id="(src-P-\d+)"', page)[:2]
    return x.replace(page, page.replace(f'<mxCell id="{b}"', f'<mxCell id="{a}"', 1))


def cut_edge(x):
    page = re.search(r'<diagram [^>]*id="flow-F01">.*?</diagram>', x, flags=re.S).group(0)
    return x.replace(page, re.sub(r'(edge="1"[^>]*?target=")[^"]+"', r'\1nowhere"', page, count=1))


def head_col(ws, name, row=None):
    """見出しの名前から列の番号（1 から）"""
    for r in ws.iter_rows(min_row=1, max_row=row or 3):
        for c in r:
            if c.value == name:
                return c.column, c.row
    raise KeyError(name)


def rename_major(wb):
    ws = wb["業務マトリクス"]
    col, row = head_col(ws, "大分類")
    ws.cell(row=row + 1, column=col).value = "別の大分類"


def unlink_time(wb):
    """業務フローのノード一覧の処理時間（分）を、業務マトリクスの時間のセルを指さない式にする"""
    ws = wb["業務フローのノード一覧"]
    col, row = head_col(ws, "処理時間（分）")
    c = next(ws.cell(row=r, column=col) for r in range(row + 1, ws.max_row + 1) if str(ws.cell(row=r, column=col).value or "").startswith("="))
    c.value = "=5"


def step(c, sid):
    return next(x for x in c["steps"] if x["id"] == sid)


def bare_source(fl, nid):
    """カードの項目を引いていない出典（概要の1行など）"""
    return next(x for x in node(fl, nid)["sources"] if not x.get("item"))


M02P1 = "cards/M02/p0001.json"


def reverse_steps(fl):
    """支払データの作成（M02 の s2）と、振込の実行（s3）の順を入れ替える"""
    pay, check, send = (next(n["id"] for n in fl["nodes"] if n["label"] == x) for x in ("支払データを作成する", "支払データを確認する", "ファームバンキングで振込を実行する"))
    for e in fl["edges"]:
        if e["to"] == pay:
            e["to"] = send
        elif (e["from"], e["to"]) == (check, send):
            e["to"] = next(x["to"] for x in fl["edges"] if x["from"] == send)
        elif e["from"] == send:
            e["to"] = pay
POINTER_CASES = [
    # 引用を行の番号で指す方式（lines）。期待データを番号に引き直した写しに、1つずつ仕込む: (名前, スクリプト, 引数, 仕込み, 出るはずの文字, FAIL か WARN か)
    ("番号: 行の番号が範囲の外", "check_extract.py", [], lambda w: edit(w / M02P1, lambda c: step(c, "s1").update(lines=[99, 99])), "が範囲の外", "FAIL"),
    ("番号: 引用を手で書き換える", "check_extract.py", [],
     lambda w: edit(w / M02P1, lambda c: step(c, "s2").update(quote="2. 経理部は、支払データを作成する。")), "の文字と違う", "FAIL"),
    ("番号: 画像の中の文字に、行の番号を書く", "check_extract.py", [],
     lambda w: edit(w / M02P1, lambda c: step(c, "s1").update(from_image=True)), "行の番号（lines）がある", "FAIL"),
    ("番号: 別の行を指す（指し違い）", "check_extract.py", [],
     lambda w: (edit(w / M02P1, lambda c: step(c, "s1").update(lines=[6, 6])), run("quotes.py", w, "fill", "--only", "M02")), "合っていない疑い", "WARN"),
    ("番号: 手順を、紙面と逆の順に並べる", "check_extract.py", [], lambda w: edit(w / M02P1, lambda c: c["steps"].reverse()), "紙面の行の順", "WARN"),
    ("番号: 1つの文を、同じ担当の手順2つに分ける", "check_extract.py", [],
     lambda w: (edit(w / M02P1, lambda c: step(c, "s3").update(lines=[3, 3])), run("quotes.py", w, "fill", "--only", "M02")), "同じ1文", "WARN"),
    ("番号: 作業の文を、どの項目にも指さない", "check_extract.py", [], lambda w: edit(w / M02P1, lambda c: c["steps"].pop(2)), "どれにも指されていない", "WARN"),
    ("順序: カードの続く手順を、フローが逆向きにつなぐ", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F03.flow.json", reverse_steps), "と逆につないでいる", "WARN"),
    ("番号: フローの出典の行の番号が範囲の外", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: bare_source(f, "P-003").update(lines=[99, 99])), "が範囲の外", "FAIL"),
    ("番号: フローの出典の引用を手で書き換える", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: bare_source(f, "P-003").update(lines=[5, 5], quote="申請の締め日は毎月末。")), "の文字と違う", "FAIL"),
    ("番号: カードの行を直したのに、フローに写していない", "check_flow.py", ["--stage", "flow"],
     lambda w: (edit(w / M02P1, lambda c: step(c, "s1").update(lines=[3, 3])), run("quotes.py", w, "fill", "--only", "M02")), "の引用と違う", "FAIL"),
]


def heading_tests(tmp):
    """本文の見出しの候補（detect_structure.py）: 番号で始まる短い行を拾い、文・数量・柱の繰り返し・表計算の行は拾わない"""
    w = tmp / "heads"
    head = ["第1章 受付", "事務の手引"]                               # 柱（章名と資料名が、どのページの上にも出る）
    add_doc(w, "M01", "手引.pdf", [], [
        ("procedure", "", head + ["第1章 受付", "1 申請の受付", "（1） 書類の確認", "① 申請書を受け取ります。", "3 月 31 日までに提出する", "2020 年度の様式"]),
        ("procedure", "", head + ["（2） 記録", "1.2 補足", "① 台帳", "担当者は、台帳に記録する。"]),
        ("procedure", "", ["第2章 支払", "事務の手引", "1 支払の準備", "（1） 書類の確認"])])
    add_doc(w, "M02", "一覧.xlsx", [], [("table", "", ["1 総務部", "2 経理部"])])
    edit(w / "manifest.json", lambda m: m["docs"][1].update(kind="xlsx"))
    code, text = run("detect_structure.py", w)
    sig = json.loads((w / "structure/signals.json").read_text(encoding="utf-8")) if code == 0 else {"docs": [{}, {}]}
    got = [(h["page"], h["text"], h["level"]) for h in sig["docs"][0].get("body_headings") or []]
    want = [(1, "第1章 受付", 1), (1, "1 申請の受付", 2), (1, "（1） 書類の確認", 4), (2, "（2） 記録", 4), (2, "1.2 補足", 3), (2, "① 台帳", 5),
            (3, "第2章 支払", 1), (3, "1 支払の準備", 2), (3, "（1） 書類の確認", 4)]
    return [("本文の見出し: 番号で始まる短い行を、段つきで拾う（文・「3 月 31 日」・「2020 年度」は拾わない。別の章の同じ見出しは拾う）", got == want),
            ("本文の見出し: 前のページの上にもある行（柱）は、初めて出たページだけ", [g[:2] for g in got].count((2, "第1章 受付")) == 0 and (1, "第1章 受付", 1) in got),
            ("本文の見出し: 表計算の番号つきの行は拾わない", not sig["docs"][1].get("body_headings")),
            ("本文の見出し: 目次が拾えない資料は、候補の一覧を画面に出す", "■ M01 の本文の見出しの候補（9 件" in text and "p0002 " + "  " * 4 + "① 台帳" in text)]


def pointer_tests(tmp, good):
    """引用を行の番号で指す方式: 期待データの引用を番号に引き直しても点検が通り、番号から引用が入り、仕込みが見つかること"""
    out, ptr = [], tmp / "ptr"
    shutil.copytree(good, ptr)
    cards = [json.loads(f.read_text(encoding="utf-8")) for f in sorted((ptr / "cards").glob("*/*.json"))]
    typed = sum(1 for c in cards for place in ("steps", "decisions", "notes") for x in c.get(place) or [] if not x.get("from_image"))
    _c, located = run("quotes.py", ptr, "locate", "--rewrite")
    c_fill, filled = run("quotes.py", ptr, "fill")
    out.append((f"番号: 期待データの引用 {typed} 件（画像の中の文字を除く全部）の行が引ける", f"引けた引用 {typed} 件" in located and "行が見つからない" not in located and c_fill == 0))
    c1, o1 = run("check_extract.py", ptr)
    c2, o2 = run("check_flow.py", ptr, "--stage", "flow")
    c3, _o = run("validate.py", ptr)
    out.append(("番号: 引き直したデータで、読取・フローの点検が FAIL 0・WARN 0",
                c1 == 0 and "0 FAIL / 0 WARN" in o1 and f"行の番号で指した引用 {typed} 件" in o1 and c2 == 0 and "0 FAIL / 0 WARN" in o2 and c3 == 0))
    _c, again = run("quotes.py", ptr, "fill")
    out.append(("番号: 引用を入れ直しても、変わらない", "カード 0 枚" in again and "フロー 0 本" in again))
    # 行を直すと、カードの引用も、その項目を引くフローの出典も、その行の文字になる。手で書き換えた引用は、行の文字に戻る
    w = tmp / "ptr_fill"
    shutil.copytree(ptr, w)
    edit(w / M02P1, lambda c: (step(c, "s1").update(lines=[3, 3]), step(c, "s4").update(quote="手で書いた引用")))
    run("quotes.py", w, "fill")
    card = json.loads((w / M02P1).read_text(encoding="utf-8"))
    line3, line5 = "2. 経理部は、支払データを作成し、経理課長の確認を受ける。", "4. 振込結果は会計システムに自動で連携される。"
    src = next(x for x in node(json.loads((w / "flows/F03.flow.json").read_text(encoding="utf-8")), "P-001")["sources"] if x.get("item") == "s1")
    out.append(("番号: 行を直して入れ直すと、カードの引用と、その項目を引くフローの出典が、その行の文字になる",
                step(card, "s1")["quote"] == line3 and step(card, "s4")["quote"] == line5 and src["quote"] == line3))
    edit(w / M02P1, lambda c: step(c, "s2").update(lines=[99, 99]))
    c_bad, o_bad = run("quotes.py", w, "fill", "--only", "M02")
    out.append(("番号: 範囲の外の番号は、入れずに知らせる", c_bad == 1 and "が範囲の外" in o_bad))
    _c, shown = run("pages.py", ptr, "lines", "--only", "M02", "--pages", "1-2")
    out.append(("番号: 行の番号つきの表示（文字層の無いページは、画像から読むと出す）", "   2| 1. 経理部は、毎月末に" in shown and "文字層なし" in shown))
    r = subprocess.run([sys.executable, str(S / "diff_runs.py"), "cards", str(ptr), str(tmp / "ptr_fill"), "--min", "1.0"], capture_output=True, text=True)
    r0 = subprocess.run([sys.executable, str(S / "diff_runs.py"), "cards", str(ptr), str(ptr), "--min", "1.0"], capture_output=True, text=True)
    out.append(("番号: 揺れの測定は、指した行で手順を対応させる（同じものは 100%、行の違う手順は差に出る）", r0.returncode == 0 and r.returncode == 1))
    for name, script, args, mutate, expect, level in POINTER_CASES:
        w = tmp / "ptr_case"
        if w.exists():
            shutil.rmtree(w)
        shutil.copytree(ptr, w)
        mutate(w)
        code, text = run(script, w, "--all-lines", *args)
        hit = [ln for ln in text.splitlines() if ln.startswith(level) and expect in ln]
        passed = bool(hit) and (code == 1 if level == "FAIL" else True)
        out.append((f"{name}: " + (hit[0][:100] if hit else f"{level}「{expect}」が出ない（終了コード {code}）\n" + "\n".join(
            ln for ln in text.splitlines() if ln.startswith(("FAIL", "WARN")))[:900]), passed))
    return out


def make_outline(w):
    """概要のページ（図＋番号つきの手順）と、詳しいページ（見出しの下に、箇条書きと条件つきの手順）の合成データ"""
    add_doc(w, "D1", "月次の締め.pdf", [], [
        ("flow_diagram", "月次の締め", ["業務の流れ", "1 残高の確認", "担当者は、残高を確認します。", "2 報告", "担当者は、課長に報告します。"]),
        ("procedure", "月次の締め（手順）", ["（1） 残高の確認", "\uf06c", "現金を数えます。", "\uf06c", "預金通帳と照合します。", "差異があれば、原因を調べます。",
                                           "（2） 報告", "結果を課長に報告します。"]),
        ("reference", "参考", ["勘定科目の一覧"])])

    def step(i, action, line, **kw):
        return dict({"id": f"s{i}", "actor": "", "action": action, "system": "", "inputs": [], "outputs": [], "condition": "", "section": "", "lines": [line, line],
                     "quote": "", "from_image": False, "origin": "body"}, **kw)
    edit(w / "cards/D1/p0001.json", lambda c: c.update(
        steps=[step(1, "残高を確認する", 3, no=1, actor="担当者"), step(2, "課長に報告する", 5, no=2, actor="担当者")],
        diagram={"nodes": [{"id": "g1", "label": "1 残高の確認"}, {"id": "g2", "label": "2 報告"}], "edges": [{"from": "g1", "to": "g2"}]}))
    edit(w / "cards/D1/p0002.json", lambda c: c.update(
        steps=[step(1, "現金を数える", 2), step(2, "預金通帳と照合する", 3, inputs=["預金通帳"]), step(3, "原因を調べる", 4, condition="差異があれば"), step(4, "結果を課長に報告する", 6)],
        decisions=[{"id": "d1", "question": "差異があるか", "branches": [{"label": "差異があれば", "to": "原因を調べる"}], "section": "", "lines": [4, 4], "quote": "",
                    "from_image": False, "origin": "body"}]))
    put(w, "glossary.json", {"terms": [{"kind": "actor", "name": "担当者", "aliases": []}, {"kind": "actor", "name": "課長", "aliases": []},
                                       {"kind": "document", "name": "預金通帳", "aliases": []}], "questions": []})
    put(w, "taxonomy.json", {"categories": [{"id": "F01", "number": "", "name": "月次の締め", "parent": None, "level": 1, "basis": "llm", "make_flow": True,
                                             "pages": [{"doc_id": "D1", "pages": [1, 2, 3]}]}], "unassigned_pages": [], "questions": []})
    run("quotes.py", w, "fill")


def bad_edits(base, change, expect):
    """直しの一覧を1か所壊して当てる: 終了コード 1・狙った文字が出る・flow.json に触らない"""
    e = json.loads((SKILL / "evals/fixture_edits/F03.edits.json").read_text(encoding="utf-8"))
    change(e)
    (base / "flows/F03.edits.json").write_text(json.dumps(e, ensure_ascii=False), encoding="utf-8")
    before = (base / "flows/F03.flow.json").read_bytes()
    code, out = run("apply_edits.py", base, "--flow", "F03")
    return code == 1 and expect in out and (base / "flows/F03.flow.json").read_bytes() == before, out


def draft_tests(tmp):
    """フローの下書き方式: 下書きの決め方（まとめる候補・箇条書きの推定・分岐）と、直しの一覧の当て方"""
    from apply_edits import make
    from draft_flow import Draft
    out, w = [], tmp / "outline"
    w.mkdir()
    make_outline(w)
    d = Draft(w, "F01")
    out.append(("下書き: 概要の手順を、同じ番号の見出しの下の最初の手順にまとめる（番号と見出しが同じ）",
                {a: b for a, (b, _why) in d.merged.items()} == {"D1.1.s1": "D1.2.s1", "D1.1.s2": "D1.2.s4"}
                and [k for k in d.seq if k not in d.merged] == ["D1.2.s1", "D1.2.s2", "D1.2.d1", "D1.2.s3", "D1.2.s4"]))
    fl, err, _d = make(w, "F01", {})
    by = {n["label"]: n for n in (fl or {}).get("nodes", [])}
    lane = {x["id"]: x["name"] for x in (fl or {}).get("lanes", [])}
    eg = {(e["from"], e["to"]): e for e in (fl or {}).get("edges", [])}
    nid = lambda label: by[label]["id"]
    ok = fl is not None and not err and len(by) == 8 and lane[by["現金を数える"]["lane"]] == "担当者" \
        and [(s["page"], s["item"]) for s in by["現金を数える"]["sources"]] == [(2, "s1"), (1, "s1")] and by["預金通帳"]["direction"] == "in"
    out.append(("下書き: ラベルはカードの文、担当は概要の手順から引き継ぎ、帳票は参照・作るから付く", ok))
    ok = ok and eg[(nid("現金を数える"), nid("預金通帳と照合する"))].get("inferred") is True and "inferred" not in eg[(nid("預金通帳と照合する"), nid("差異があるか"))] \
        and eg[(nid("差異があるか"), nid("原因を調べる"))]["label"] == "差異があれば" and eg[(nid("差異があるか"), nid("結果を課長に報告する"))].get("inferred") is True \
        and "inferred" not in eg[(nid("原因を調べる"), nid("結果を課長に報告する"))]
    out.append(("下書き: 番号の無い箇条書きの項から項への線と、分岐の書かれていない側だけが推定になる", ok))
    qs = {q["id"]: q for q in (fl or {}).get("questions", [])}
    out.append(("下書き: 決まった形の確認事項を機械が作る（担当の無い作業は1件・箇条書きの順・分岐の書かれていない側・開始条件）。"
                "終了条件は、終了に入る最後の作業から機械が入れる（確認事項にしない）",
                sorted(qs) == ["m-else-2", "m-lane", "m-list-1", "m-start"] and qs["m-lane"]["nodes"] == [nid("預金通帳と照合する"), nid("原因を調べる")]
                and (fl or {}).get("end_condition") == "結果を課長に報告したとき"
                and (fl or {}).get("unused_pages") == [{"doc_id": "D1", "page": 3, "reason": "手順・分岐の記載が無いページ（参考）"}]))
    fl2, _e, _d = make(w, "F01", {})
    fl3, err3, _d = make(w, "F01", {"default_lane": "課長", "ops": [
        {"op": "unmerge", "item": "D1.1.s2"}, {"op": "exclude", "item": "D1.1.s2", "reason": "概要の1行。詳しい手順に同じ作業がある"},
        {"op": "decide", "item": "D1.2.s3", "question": "差異があるか", "yes": "ある", "no": "ない"}, {"op": "exclude", "item": "D1.2.d1", "reason": "手順の条件を判断にした"},
        {"op": "split", "item": "D1.2.s4", "parts": [{"label": "結果を報告する", "lane": "担当者"}, {"label": "報告を受ける", "lane": "課長", "inferred": True}]}]})
    lb = {n["label"]: n for n in (fl3 or {}).get("nodes", [])}
    out.append(("直しの一覧: 同じ入力からは同じフロー。まとめを外す・対象外・条件を判断にする・担当で分ける、が当たる",
                fl == fl2 and not err3 and [x["item"] for x in fl3["excluded_items"]] == ["s2", "d1"] and "報告を受ける" in lb and lb["差異があるか"]["type"] == "decision"
                and [(s["page"], s["item"]) for s in lb["差異があるか"]["sources"]] == [(2, "s3")] and not lb["差異があるか"].get("inferred")
                and {x["name"] for x in fl3["lanes"]} == {"担当者", "課長"}))
    # 架空データ（資料4つ）: 直しの一覧から作ったフローが、点検を通る
    fx = tmp / "ptr_edits"
    shutil.copytree(tmp / "ptr", fx)
    shutil.copy2(SKILL / "evals/fixture_edits/F03.edits.json", fx / "flows/F03.edits.json")
    c0, o0 = run("apply_edits.py", fx, "--flow", "F03")
    first = (fx / "flows/F03.flow.json").read_bytes()
    c1, o1 = run("check_flow.py", fx, "--stage", "flow")
    c2, _o = run("validate.py", fx)
    run("apply_edits.py", fx, "--flow", "F03")
    f3 = json.loads(first.decode("utf-8"))
    kinds = [n["type"] for n in f3["nodes"]]
    ok = c0 == 0 and c1 == 0 and "0 FAIL / 0 WARN" in o1 and c2 == 0 and (fx / "flows/F03.flow.json").read_bytes() == first
    out.append(("直しの一覧: 架空データの「支払処理」を直しの一覧から作ると、点検が FAIL 0・WARN 0。やり直しても同じ" + ("" if ok else "\n" + o0[-600:] + o1[-900:]), ok))
    acts = {n["label"]: (n.get("action"), n.get("volume"), n.get("volume_reason")) for n in f3["nodes"] if n["type"] == "process"}
    out.append(("直しの一覧: 動作の種類（op time・足す箱・分ける箱）が、人の作業の箱すべてに付く。量は、理由つきで変えた箱だけ。判断・システム処理には付かない",
                all(a for a, _v, _r in acts.values()) and acts["支払データを作成する"][0] == "入力する" and acts["支払データを確認する"][0] == "確認する"
                and acts["再振込の記録を支払台帳に記入する"] == ("入力する", "小", "台帳に 1 行、書くだけ")
                and sum(1 for _a, v, _r in acts.values() if v) == 1 and not any(n.get("action") for n in f3["nodes"] if n["type"] != "process")))
    lab = {n["id"]: n["label"] for n in f3["nodes"]}
    back = [e for e in f3["edges"] if lab.get(e["from"]) == "申請者に差し戻す"]
    other = [e.get("label") for e in f3["edges"] if lab.get(e["from"]) == "不備があるか" and e.get("inferred")]
    out.append(("直しの一覧: 足したノードだけが推定になり、システムのレーンの作業はシステム処理になる。分岐の先の差し戻しは、推定の終了と確認事項になる"
                "（受け手の箱は足さない）。書かれていない側のラベルは、書かれた側の反対の語（あり → なし）。エラーのあとの戻りがある"
                + ("" if back and other else f"\n{back}{other}"),
                sorted(n["label"] for n in f3["nodes"] if n.get("inferred")) == ["振込エラーがあるか"] and kinds.count("system") == 1
                and [(e["to"], e.get("inferred")) for e in back] == [("end2", True)] and other == ["なし"]
                and any(q["id"].startswith("m-dead") and q.get("edge") == [back[0]["from"], "end2"] for q in f3["questions"])
                and kinds.count("decision") == 2 and [x["item"] for x in f3["excluded_items"]] == ["s1"]
                and any(q["id"] == "q1" and q.get("node") for q in f3["questions"]) and sum(1 for e in f3["edges"] if e.get("inferred")) == 4))
    c3, o3 = run("draft_flow.py", fx, "--flow", "F03")
    out.append(("下書きの表示: キー・担当・ラベル・原文が出る", c3 == 0 and "M02.1.s2  手順  [経理部]  支払データを作成し、経理課長の確認を受ける" in o3 and "原文（画像）" in o3))
    for name, change, expect in [
            ("決まっていない種類の直し", lambda e: e["ops"].append({"op": "rewrite", "item": "M02.1.s1"}), "形の違反"),
            ("下書きに無い項目を指す", lambda e: e["ops"].append({"op": "label", "item": "M02.1.s9", "label": "x"}), "下書きに無い"),
            ("同じ担当の作業に分ける", lambda e: e["ops"].append({"op": "split", "item": "M02.1.s1", "parts": [{"label": "出力する", "lane": "経理部"}, {"label": "確認する", "lane": "経理部"}]}),
             "同じレーン"),
            ("理由なしで対象外にする", lambda e: e["ops"][1].update(reason=" "), "reason"),
            ("条件の無い手順を判断にする", lambda e: e["ops"].append({"op": "decide", "item": "M02.2.s1", "question": "エラーか", "yes": "はい", "no": "いいえ"}), "条件"),
            ("線を外して、たどり着けない項目を作る", lambda e: e["ops"].append({"op": "disconnect", "from": "M02.2.s1", "to": "M02.2.s2"}), "たどり着けない"),
            ("推定でない線に、根拠の場所が無い", lambda e: e["ops"].append({"op": "connect", "from": "M02.2.s4", "to": "M02.1.s4", "inferred": False}), "evidence"),
            ("用語の正本に無いレーン", lambda e: e["ops"].append({"op": "lane", "item": "M02.1.s1", "lane": "財務部"}), "用語の正本"),
            ("機械の確認事項の名前を使う", lambda e: e["questions"][0].update(id="m-lane"), "m- で始まる"),
            ("差し戻しを受ける側の箱を足す", lambda e: e["ops"].append({"op": "add", "key": "x9", "type": "process", "after": "M03.1.s3", "label": "差し戻しを受け取る", "lane": "申請者",
                                                                 "source": {"doc_id": "M03", "page": 2, "quote": "申請者に差し戻す"}}), "差し戻しを受ける側の箱は足さない"),
            ("決まっていない動作の種類を書く", lambda e: e["ops"].append({"op": "time", "item": "M02.1.s1", "action": "打ち込む"}), "形の違反"),
            ("理由なしで、量を変える", lambda e: e["ops"].append({"op": "time", "item": "M02.1.s1", "volume": "大"}), "reason"),
            ("判断に、動作の種類を付ける", lambda e: e["ops"].append({"op": "time", "item": "M03.2.d1", "action": "確認する"}), "手順でない"),
            ("補足の直しに、補足も参考も無い", lambda e: e["ops"].append({"op": "detail", "item": "M02.1.s1"}), "detail か reference を書く"),
            ("決まっていないレーンの種類", lambda e: e.update(lane_kinds={"経理部": "部門"}), "形の違反")]:
        passed, text = bad_edits(fx, change, expect)
        out.append((f"直しの一覧の誤り: {name} → 何も書かずに知らせる" + ("" if passed else f"（「{expect}」が出ない）\n" + text[-500:]), passed))
    return out


def no_start(w):
    """F02 の開始条件を空にし、それを聞く機械の問い（m-start）も外す"""
    def f(fl):
        fl["start_condition"] = ""
        fl["questions"] = [q for q in fl.get("questions") or [] if q.get("id") != "m-start"]
    edit(w / "flows/F02.flow.json", f)


def glossary_doc_id(w):
    """用語の確認事項の文に資料 ID を書いて、報告書を書き出し直す"""
    edit(w / "glossary.json", lambda g: g["questions"].append({"id": "q9", "category": "用語", "effect": "システム・帳票", "question": "M02 p.1 の「伝票」と M01 の「申請書」は同じものですか、別のものですか",
                                                             "assumption": "同じものとした", "sources": [{"doc_id": "M02", "page": 1}], "ask": "経理部"}))
    run("build_reports.py", w)


CASES = [
    # (名前, スクリプト, 引数, 仕込み(work), 出るはずの文言)
    # ---- 業務マトリクス一覧 ----
    ("マトリクス: セルから作業を1つ消す", "check_flow.py", ["--stage", "matrix"], lambda w: edit_book(w, blank_step), "作業のセルが、フローの作業のノードと合わない"),
    ("マトリクス: 行き先の文を1つ消す", "check_flow.py", ["--stage", "matrix"], lambda w: edit_book(w, drop_exit), "行き先から組み直した線が"),
    ("マトリクス: 確認事項の番号を1つ消す", "check_flow.py", ["--stage", "matrix"], lambda w: edit_book(w, drop_question), "確認事項の番号が"),
    ("マトリクス: 仮の処理時間の式が、別の動作の種類を引いている", "check_flow.py", ["--stage", "matrix"], lambda w: edit_book(w, wrong_lookup), "仮の値の式が"),
    ("マトリクス: 処理時間の合計の式を、書き換える", "check_flow.py", ["--stage", "matrix"], lambda w: edit_book(w, break_total), "合計の式が"),
    ("マトリクス: フローを直したあと、書き出し直していない", "check_flow.py", ["--stage", "matrix"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-001").update(label="経費精算申請書を起票する")), "いまのデータと違う"),
    ("マトリクス: ブックが無い", "check_flow.py", ["--stage", "matrix"], lambda w: (w / "out" / "業務マトリクス一覧.xlsx").unlink(), "が無い"),
    ("原典に無いノードに、推定の印が無い", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F03.flow.json", lambda f: node(f, "D-002").pop("inferred")), "推定（inferred）の印が無い"),
    # ---- 形（スキーマ）----
    ("カードに決めていない項目", "check_extract.py", [],
     lambda w: edit(w / "cards/M01/p0003.json", lambda c: c.update(header="経費精算業務マニュアル")), "スキーマに無い項目"),
    ("カードの型違い（番号を数で書く）", "check_extract.py", [],
     lambda w: edit(w / "cards/M04/p0002.json", lambda c: c.update(printed_page=1)), "型が違う"),
    ("カードの決まっていない語", "check_extract.py", [],
     lambda w: edit(w / "cards/M01/p0003.json", lambda c: c["steps"][0].update(origin="memo")), "決まった語でない"),
    ("フローに決めていない項目", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-001").update(memo="あとで確認")), "スキーマに無い項目"),
    ("確認事項の決まっていない種別", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: f["questions"][0].update(category="読取")), "決まった語でない"),
    # ---- 階層・業務分類・フローの見出し・余りもの ----
    ("階層: 根拠の文があるのに、出典のページが無い", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h["tree"][0]["children"][0]["refs"][1].pop("sources")), "出典のページ（sources）が無い"),
    ("階層: 「続く」の根拠の出典が、無いページ", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h["sequence"][0].update(sources=[{"doc_id": "M01", "page": 99}])), "根拠の出典が、存在しないページ"),
    ("階層: 確認事項の出典が、無いページ", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h["questions"][0].update(sources=[{"doc_id": "M02", "page": 99}])), "の出典が、存在しないページ"),
    ("業務分類: 段が親と合わない", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "taxonomy.json", lambda t: t["categories"][2].update(level=3)), "親から数えた段"),
    ("フローの名前が業務分類と違う", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: f.update(name="別の名前")), "業務分類と違う"),
    ("業務分類に無いフローの直しの一覧", "check_flow.py", ["--stage", "flow"],
     lambda w: (w / "flows/F09.edits.json").write_text("{}", encoding="utf-8"), "業務分類に無いフローの直しの一覧"),
    ("図: out/ に、いまのデータから書き出す図の名前に無い図がある", "check_flow.py", ["--stage", "drawio"],
     lambda w: shutil.copy2(w / "out/業務フロー_as-is.drawio", w / "out/業務フロー_as-is_手直し.drawio"), "書き出す図の名前に無い"),
    ("フロー: 開始条件が空なのに、それを聞く確認事項が無い", "check_flow.py", ["--stage", "flow"], no_start, "それを聞く確認事項が無い"),
    ("納品物: 用語の確認事項の文に資料 ID が出る", "check_flow.py", ["--stage", "deliverables"], glossary_doc_id, "納品物に「M0"),
    # ---- 補足の中の作業・項目の網羅 ----
    ("補足の中の作業の読み落とし", "check_extract.py", [],
     lambda w: edit(w / "cards/M01/p0005.json", lambda c: c.update(steps=[s_ for s_ in c["steps"] if s_["id"] != "s5"])), "補足の中に作業らしい文がある"),
    ("作業でない理由の書き落とし", "check_extract.py", [],
     lambda w: edit(w / "cards/M02/p0001.json", lambda c: c["notes"][0].pop("why_not_action")), "補足の中に作業らしい文がある"),
    ("カードの手順がフローに無い", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-001")["sources"][0].pop("item")), "がフローに無い"),
    ("補足の中の作業がフローに無い", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: [s_.pop("item") for n in f["nodes"] for s_ in n.get("sources") or [] if s_.get("item") == "s5"]),
     "補足の中の作業"),
    ("出典が引く項目が無い", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-001")["sources"][0].update(item="s9")), "カードの項目 s9 が"),
    ("出典の引用が項目と合わない", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-001")["sources"][0].update(item="s3")), "引いた項目 s3"),
    # ---- ページの取り違え ----
    ("カード2枚の入れ替え", "check_extract.py", [], lambda w: swap_cards(w, "M04", 2, 3), "ページの取り違え"),
    ("カードの番号が刷られた番号と違う", "check_extract.py", [],
     lambda w: edit(w / "cards/M04/p0003.json", lambda c: c.update(printed_page="3")), "ページの端に刷られた数字"),
    ("出典のページが1つずれる", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F03.flow.json", lambda f: [s_.update(page=2) for n in f["nodes"] for s_ in n.get("sources") or []
                                                         if (s_["doc_id"], s_["page"]) == ("M02", 1)]), "ページずれている"),
    ("全ノードの出典が1ページずれる", "check_flow.py", ["--stage", "flow"], lambda w: shift_sources(w, "M01", -1), "系統的なずれ"),
    ("目次が指すページが割当に無い", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "taxonomy.json", lambda t: t["categories"][3].update(
         pages=[{"doc_id": "M01", "pages": [5]}, {"doc_id": "M02", "pages": [1, 2]}, {"doc_id": "M03", "pages": [1, 2]}, {"doc_id": "M04", "pages": [3]}])),
     "目次の番号"),
    ("ページ画像と機械抽出がずれている", "check_extract.py", [],
     lambda w: edit(w / "pages/M01/p0005.json", lambda pj: pj.update(misaligned=True, align=0.12)), "ページ画像の文字と機械抽出の文字が合わない"),
    ("シートとページ画像の対応が取れない", "check_extract.py", [], lambda w: edit(w / "manifest.json", relabel_sheet), "どのシートのものか対応が取れていない"),
    # ---- 資料の紐付け・用語・大きさ ----
    ("根拠の無い紐付け", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: tree_ref(h, "H1.3", "M02").update(evidence="")), "根拠（evidence）が無い"),
    ("どの業務にも結び付かない資料", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: drop_refs(h, "M04")), "どの業務にも結び付いていない"),
    ("上位・下位が業務の木に落ちていない", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: tree_ref(h, "H1.3", "M02").update(rel="describes")), "どの節を詳しくしているか"),
    ("役割の無い資料", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h.update(docs=h["docs"][:3])), "役割（docs[].role）が書かれていない"),
    ("分類の親子の循環", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "taxonomy.json", lambda t: t["categories"][0].update(parent="F01")), "循環"),
    ("材料が大きすぎる分類", "check_flow.py", ["--stage", "structure", "--max-material", "3000"], lambda w: None, "材料が大きすぎる"),
    ("用語の正本に無いレーン名", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: f["lanes"][1].update(name="課長代理")), "用語の正本"),
    ("根拠なしに寄せた用語", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "glossary.json", lambda g: g["terms"][1].update(aliases=["課長"])), "同じものとした根拠"),
    ("続く業務へのリンクが無い", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h["sequence"].append({"from": "H1.1", "to": "H1.3", "evidence": "（仕込み）"})),
     "へのリンク（subprocess・link_flow）が無い"),
    # ---- 以前からの点検 ----
    ("出典の無いノード", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-001").update(sources=[])), "出典が無い"),
    ("作り話の引用", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-001")["sources"][0].update(quote="申請者は部長の事前承認を得てから申請する。")),
     "引用が原典"),
    ("存在しないページへの出典", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-001")["sources"][0].update(page=99)), "存在しないページを出典"),
    ("ラベルの無い分岐", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: [e.pop("label") for e in f["edges"] if e["from"] == "D-002" and e["to"] == "P-004"]),
     "ラベルの無い分岐"),
    ("出口が1本の分岐", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: f.update(edges=[e for e in f["edges"] if not (e["from"] == "D-001" and e["to"] == "D-002")])),
     "分岐の出口が 1 本"),
    ("行き止まり", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: f.update(edges=[e for e in f["edges"] if e["from"] != "P-003"])), "終了にたどり着けない"),
    ("推定なのに確認事項なし", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: f.update(questions=[q for q in f["questions"] if "P-003" not in q_targets(q)])),
     "推定（inferred）で置いたのに"),
    ("札（使うシステム）を1つ消す", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-002").pop("systems")), "札（systems）が、出典に引いた手順の system と合わない"),
    ("札に、手順に無いシステムを書く", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: node(f, "P-001").update(systems=["経費精算システム"])), "札（systems）が、出典に引いた手順の system と合わない"),
    ("原典にある作業（カードの手順を引くノード）に推定の印を付ける", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: node(f, "P-001").update(inferred=True)), "を出典に引いているのに、推定"),
    ("まとめた確認事項（nodes）が、存在しないノードを指す", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: f["questions"][0].update(nodes=["P-001", "P-002", "P-999"])), "存在しないノード P-999"),
    ("担当の書かれていない手順を、確認事項なしでレーンに置く", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F02.flow.json", lambda f: (node(f, "P-005").pop("inferred", None),
                                                            f.update(questions=[q for q in f["questions"] if "P-005" not in q_targets(q)]))),
     "もとの手順に担当が書かれていない"),
    ("使われていないページ", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F03.flow.json", drop_error_branch), "どのノードにも使われていない内容ページ"),
    ("対象外の項目を1つだけ書いて、残りを落とす", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F03.flow.json", exclude_one_item), "の手順 s2「"),
    ("ほかの分類のページを unused_pages に書く", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: f.update(unused_pages=[{"doc_id": "M02", "page": 1, "reason": "この分類では使わない"}])),
     "この分類に割り当てたページでない"),
    ("確認事項の id の重複", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: f["questions"].append(dict(f["questions"][0], question="別の問い"))), "確認事項の id が重複"),
    ("帳票への線", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: f["edges"].append({"from": "P-002", "to": "F-001"})), "の端が帳票"),
    ("壊れたフロー（JSON として読めない）", "check_flow.py", ["--stage", "flow"],
     lambda w: (w / "flows/F01.flow.json").write_text('{"flow_id": "F01", "nodes": [', encoding="utf-8"), "JSON として読めない"),
    ("同じ名前の用語が2件", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "glossary.json", lambda g: g["terms"].append({"kind": "system", "name": g["terms"][0]["name"], "aliases": []})), "2つの用語"),
    ("継ぎはぎの引用（主語のすり替え）", "check_extract.py", [],
     lambda w: edit(w / "cards/M01/p0005.json", patchwork), "引用が機械抽出の文字に見当たらない"),
    ("手順のあるページを白紙にする", "check_extract.py", [],
     lambda w: edit(w / "cards/M02/p0002.json", lambda c: c.update(page_type="blank")), "種別が blank なのに"),
    ("壊れたカード（JSON として読めない）", "check_extract.py", [],
     lambda w: (w / "cards/M01/p0003.json").write_text('{"doc_id": "M01",', encoding="utf-8"), "カードが JSON として読めない"),
    ("仮置きの語", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-002").update(label="担当不明の処理")), "仮置きの語"),
    ("目次と違う分類名", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "taxonomy.json", lambda t: t["categories"][2].update(name="承認ワークフロー")), "目次"),
    ("目次があるのに LLM 分類", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "taxonomy.json", lambda t: t.update(basis="llm")), "目次に業務フローの一覧があるのに"),
    ("割り当てていないページ", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "taxonomy.json", lambda t: t["categories"][3].update(pages=[{"doc_id": "M01", "pages": [6]}])),
     "どの業務分類にも割り当てていない内容ページ"),
    ("カードの作り話の引用", "check_extract.py", [],
     lambda w: edit(w / "cards/M01/p0003.json", lambda c: c["steps"][0].update(quote="申請者は上長の押印をもらう。")), "機械抽出の文字に見当たらない"),
    ("カードが無いページ", "check_extract.py", [], lambda w: (w / "cards/M02/p0001.json").unlink(), "カードが無いページ"),
    ("手順のページなのに空", "check_extract.py", [],
     lambda w: edit(w / "cards/M02/p0001.json", lambda c: c.update(steps=[])), "手順・分岐・図がどれも空"),
    ("図を手で直した（ラベル）", "check_flow.py", ["--stage", "drawio"],
     lambda w: [p.write_text(p.read_text(encoding="utf-8").replace("経費精算システムにログインする", "システムへログイン"), encoding="utf-8")
                for p in (w / "out").glob("*.drawio")], "ラベルが flow.json と違う"),
    ("図を手で直した（札を消す）", "check_flow.py", ["--stage", "drawio"],
     lambda w: [p.write_text(p.read_text(encoding="utf-8").replace("［経費精算システム］", "", 1), encoding="utf-8")
                for p in (w / "out").glob("*.drawio")], "の札「［経費精算システム］」が図に無い"),
    ("テンプレートに無い色", "check_flow.py", ["--stage", "drawio"],
     lambda w: [p.write_text(p.read_text(encoding="utf-8").replace("fillColor=#d5e8d4", "fillColor=#ff0000", 1), encoding="utf-8")
                for p in (w / "out").glob("*.drawio")], "テンプレートに無い色"),
    ("ノードの重なり", "check_layout.py", ["--no-png"],
     lambda w: [edit(p, lambda l: l["pages"][1]["nodes"].update({"P-002": l["pages"][1]["nodes"]["P-001"]})) for p in (w / "out").glob("*.layout.json")],
     "が重なる"),
    ("出典の字（枠の外）に、ノードが重なる", "check_layout.py", ["--no-png"],
     lambda w: [edit(p, lambda l: l["pages"][1]["nodes"].update({"P-002": l["pages"][1]["captions"]["P-001"]})) for p in (w / "out").glob("*.layout.json")],
     "の出典の字が、ノード"),
    ("線が、出典の字（枠の外）を通る", "check_layout.py", ["--no-png"],
     lambda w: [edit(p, lambda l: l["pages"][1]["edges"].append({"id": "x", "from": "start", "to": "end", "points": (
         lambda b: [[b[0] - 5, b[1] + b[3] / 2], [b[0] + b[2] + 5, b[1] + b[3] / 2]])(l["pages"][1]["captions"]["P-001"])})) for p in (w / "out").glob("*.layout.json")],
     "の出典の字を通る"),
    ("箱の名前が、動詞で終わらない（「〜へ」）", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "SUB-001").update(label="承認へ")), "動詞で終わっていない"),
    ("判断の名前が、問いの形でない", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F03.flow.json", lambda f: node(f, "D-001").update(label="不備あり？")), "「〜か」で終わっていない"),
    ("times.json に、ブックから取り込んでいない値（LLM が書いた数）", "check_flow.py", ["--stage", "matrix"],
     lambda w: (w / "times.json").write_text(json.dumps({"version": 1, "cells": [{"flow": "F01", "node": "P-002", "value": 42}]}), encoding="utf-8"), "ブックから取り込んでいない値"),
    ("確認事項に、答えで変わる所（effect）が無い", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: f["questions"][0].pop("effect")), "effect"),
    # ---- 図（draw.io）: ページ・セル・線・形・リンク ----
    ("図: フローのページを消す", "check_flow.py", ["--stage", "drawio"], lambda w: redraw(w, lambda x: drop_page(x, "flow-F02")), "フロー F02 のページが draw.io に無い"),
    ("図: 全体構成のページを消す", "check_flow.py", ["--stage", "drawio"], lambda w: redraw(w, lambda x: drop_page(x, "overview")), "全体構成のページが無い"),
    ("図: XML として読めない", "check_flow.py", ["--stage", "drawio"], lambda w: redraw(w, lambda x: x[:-40]), "XML として読めない"),
    ("図: セルの id が重複", "check_flow.py", ["--stage", "drawio"], lambda w: redraw(w, dup_cell), "セルの id が重複"),
    ("図: 線の行き先のセルが無い", "check_flow.py", ["--stage", "drawio"], lambda w: redraw(w, cut_edge), "の両端が無い"),
    ("図: テンプレートに無い形", "check_flow.py", ["--stage", "drawio"], lambda w: redraw(w, lambda x: x.replace("rhombus", "hexagon", 1)), "はテンプレートに無い"),
    ("図: リンク先のページが無い", "check_flow.py", ["--stage", "drawio"],
     lambda w: redraw(w, lambda x: x.replace('link="data:page/id,flow-F02"', 'link="data:page/id,flow-F09"', 1)), "のリンク先のページ"),
    # ---- 業務マトリクス一覧: 行・分類・シート・時間のセル ----
    ("マトリクス: 業務フローのノード一覧の行を1つ消す", "check_flow.py", ["--stage", "matrix"],
     lambda w: edit_book(w, lambda wb: wb["業務フローのノード一覧"].delete_rows(3)), "行の数が、いまのデータと違う"),
    ("マトリクス: 大分類を書き換える", "check_flow.py", ["--stage", "matrix"], lambda w: edit_book(w, rename_major), "大分類・中分類が"),
    ("マトリクス: 一覧のシートを消す", "check_flow.py", ["--stage", "matrix"], lambda w: edit_book(w, lambda wb: wb.remove(wb["業務フローのノード一覧"])), "が、本に無い"),
    ("マトリクス: ノード一覧の処理時間が、業務マトリクスの時間のセルを指さない", "check_flow.py", ["--stage", "matrix"],
     lambda w: edit_book(w, unlink_time), "時間のセル"),
    ("マトリクス: ブックが壊れている", "check_flow.py", ["--stage", "matrix"], lambda w: (w / "out" / "業務マトリクス一覧.xlsx").write_bytes(b"PK\x03\x04broken"), "を読めない"),
    # ---- 配置 ----
    ("線が、ノードを貫く", "check_layout.py", ["--no-png"],
     lambda w: [edit(p, lambda l: l["pages"][1]["edges"].append({"id": "x", "from": "start", "to": "end", "points": (
         lambda b: [[b[0] - 20, b[1] + b[3] / 2], [b[0] + b[2] + 20, b[1] + b[3] / 2]])(l["pages"][1]["nodes"]["P-002"])})) for p in (w / "out").glob("*.layout.json")],
     "を貫く"),
    # ---- 階層 ----
    ("階層: 業務の木の id が重複", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h["tree"][0]["children"][1].update(id=h["tree"][0]["children"][0]["id"])), "階層ノードの id が重複"),
    ("階層: 業務の木が、無い資料を指す", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: tree_ref(h, "H1.3", "M04").update(doc_id="M09")), "が存在しない資料を指す: M09"),
    ("階層: 業務の木が、無いページを指す", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: tree_ref(h, "H1.3", "M04").update(pages=[99])), "が存在しないページを指す"),
    ("階層: 資料の役割が、無い資料を指す", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h["docs"].append({"doc_id": "M09", "role": "manual"})), "docs が存在しない資料を指す"),
    ("階層: 資料どうしの関係が、無い資料を指す", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h["doc_relations"][0].update(to="M09")), "資料間の関係が存在しない資料を指す"),
    ("階層: 「続く」が、業務の木に無い id を指す", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h["sequence"][0].update(to="H9")), "「続く」が業務の木に無い id"),
    ("階層: 目次に一覧があるとしたのに、項目が空", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h["toc_flow_list"].update(entries=[])), "目次の資料（source）か項目（entries）が空"),
    ("階層: 階層ありなのに、業務の木が空", "check_flow.py", ["--stage", "structure"],
     lambda w: edit(w / "structure/hierarchy.json", lambda h: h.update(tree=[])), "tree が空"),
    # ---- フロー: 形の要の所 ----
    ("フロー: レーンが無いノード", "check_flow.py", ["--stage", "flow"], lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-001").update(lane="L9")), "レーン「L9」が無い"),
    ("フロー: ラベルが空のノード", "check_flow.py", ["--stage", "flow"], lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "P-002").update(label=" ")), "ラベルが空"),
    ("フロー: ノードの id が重複", "check_flow.py", ["--stage", "flow"], lambda w: edit(w / "flows/F01.flow.json", lambda f: f["nodes"].append(dict(node(f, "P-001")))), "ノードの id が重複"),
    ("フロー: レーンの id が重複", "check_flow.py", ["--stage", "flow"], lambda w: edit(w / "flows/F01.flow.json", lambda f: f["lanes"].append(dict(f["lanes"][0], name="経理部"))), "レーンの id が重複"),
    ("フロー: 開始ノードが無い", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: f.update(nodes=[n for n in f["nodes"] if n["type"] != "start"], edges=[e for e in f["edges"] if e["from"] != "start"])), "開始ノードが無い"),
    ("フロー: 開始ノードに入る線", "check_flow.py", ["--stage", "flow"], lambda w: edit(w / "flows/F01.flow.json", lambda f: f["edges"].append({"from": "P-001", "to": "start"})), "に入る線がある"),
    ("フロー: 終了ノードから出る線", "check_flow.py", ["--stage", "flow"], lambda w: edit(w / "flows/F01.flow.json", lambda f: f["edges"].append({"from": "end", "to": "P-001"})), "から出る線がある"),
    ("フロー: 下位フローへのリンク先が、業務分類に無い", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: node(f, "SUB-001").update(link_flow="F09")), "link_flow「F09」が無い"),
    ("フロー: 帳票名が用語の正本に無い", "check_flow.py", ["--stage", "flow"],
     lambda w: edit(w / "flows/F01.flow.json", lambda f: next(n for n in f["nodes"] if n["type"] == "file").update(label="謎の帳票")), "が用語の正本（glossary.json）に無い"),
]


def matrix_looks(work):
    """業務マトリクス一覧（xlsx）を読み戻して、行・セル・行き先・確認事項が、もとのフローのとおりかを見る"""
    import zipfile
    from openpyxl import load_workbook
    from build_matrix import BOOK, EXIT_RE, LOOK, MATRIX_HEAD, NODE_HEAD, NODES, PROGRAM
    book = work / "out" / BOOK
    wb = load_workbook(book, rich_text=True)
    text = lambda c: "" if c.value is None else str(c.value)
    flows = {f.name.split(".")[0]: json.loads(f.read_text(encoding="utf-8")) for f in sorted((work / "flows").glob("*.flow.json"))}
    qs = json.loads((work / "questions.json").read_text(encoding="utf-8"))["questions"]
    both = [[text(c) for c in r] for r in wb["業務マトリクス"].iter_rows(min_row=2)]
    rows, under = both[::2], both[1::2]                              # 作業の行・そのすぐ下の、処理時間の行
    col = {h: i for i, h in enumerate(NODE_HEAD)}
    nodes = {}
    for r in wb[NODES].iter_rows(min_row=2):
        v = [text(c) for c in r]
        nodes.setdefault(v[col["フローID"]], []).append(v)
    left = len(MATRIX_HEAD)                                          # 作業の列の左の欄（…件数・処理時間の合計のあとに、作業の列）
    cells = lambda r: [c for c in r[left:] if c]
    steps = lambda fl: [n for n in fl["nodes"] if n["type"] not in ("start", "end", "file")]
    edges = lambda F: {(v[col["ノード番号"]], mt["to"], mt["label"] or "", "（推定）" in mt["marks"])
                       for v in nodes[F] for ln in v[col["行き先"]].split("\n") if ln for mt in [EXIT_RE.match(ln)] if mt}
    f3 = rows[list(flows).index("F03")]
    xml = "".join(zipfile.ZipFile(book).read(n).decode("utf-8") for n in zipfile.ZipFile(book).namelist() if n.startswith("xl/worksheets/sheet"))
    return [
        ("業務マトリクス一覧: シートが決まった並び（「処理時間の基準」は「読み方」の前。起動点一覧は無い。プログラムが無ければ、プログラム用のシートも無い）であり、"
         "フロー1本が、作業の行と処理時間の行の 2 行",
         wb.sheetnames == [s for s in LOOK["order"] if s not in PROGRAM] and "起動点一覧" not in LOOK["order"] and len(both) == 2 * len(flows) and list(nodes) == list(flows)
         and all(any("｜処理時間（分）" in x for x in r[:2]) for r in under)),
        ("業務マトリクス一覧: 処理時間の行は、作業のセルの真下に、基準の表を引く式。左に合計の式、その左に、空の件数。見出しは「開始条件」「終了条件」",
         all(len(cells(t)) == len(steps(fl)) and all(c.startswith('=VLOOKUP("') for c in cells(t)) and t[left - 1].startswith("=SUM(") and t[left - 2] == "" and r[left - 2] == ""
             for r, t, fl in zip(rows, under, flows.values()))
         and [text(c) for c in wb["業務マトリクス"][1]][4:9] == ["開始条件", "終了条件", "フロー全体の確認事項", "件数（人が入れる）", "処理時間の合計（分・全作業）"]),
        ("業務マトリクス一覧: 開始・終了・帳票を除くノードが、どれも1つのセルになる",
         all(len(cells(r)) == len(steps(fl)) for r, fl in zip(rows, flows.values()))),
        ("業務マトリクス一覧: 業務フローのノード一覧の「行き先」から線を組み直すと、フローの線と同じ",
         all(edges(F) == {(e["from"], e["to"], e.get("label") or "", bool(e.get("inferred"))) for e in fl["edges"]} for F, fl in flows.items())),
        ("業務マトリクス一覧: そのフローの確認事項の番号が、すべて行に出る",
         all(set(re.findall(r"Q-\d+", " ".join(r))) == {q["id"] for q in qs if q.get("flow") == F} for r, F in zip(rows, flows))),
        ("業務マトリクス一覧: 資料が複数のフローは、原典が資料ごとに、題名で出る（手順書・別紙・全体マニュアル）",
         all(d in f3[3] for d in ("支払処理手順 p.1〜2", "経費精算チェックシート「手順」シート", "経費精算業務マニュアル スライド 6", "「支払処理手順」は「経費精算業務マニュアル」の下位"))
         and "M0" not in f3[3]),
        ("業務マトリクス一覧: 戻りの線は「へ戻る」、推定の線は（推定）と書く", any("へ戻る" in c for c in cells(f3)) and any("→（推定）" in c for c in cells(f3))),
        ("業務マトリクス一覧: 空白だけの区切りを書かない（あると Excel がブックを開かない）", not re.search(r"<t(?: [^>]*)?>\s*</t>", xml)),
    ]


def time_tests(tmp, good):
    """業務マトリクス一覧の処理時間: 仮の値は、基準の表（動作の種類 × 量）を引く式。人が入れた値・件数・表の直しは、書き出し直しても残る。
    番号を振り直しても・フローを作り直しても付いていき、箱が無くなっても、黙って消えない。基準の表の無いテンプレートでは、時間の行を出さない"""
    from openpyxl import load_workbook
    import build_matrix as BM
    import office_to_pdf
    import table_template as TT
    out, w = [], tmp / "times"
    shutil.copytree(good, w)
    book, M = w / "out" / BM.BOOK, "業務マトリクス"

    def tagged(wb):
        """({(フロー, ノード): 時間のセル}, {フロー: 件数のセル})"""
        t, n = {}, {}
        for row in wb[M].iter_rows():
            for c in row:
                text = c.comment.text if c.comment is not None else ""
                m, k = BM.TAG_TIME.match(text), BM.TAG_COUNT.match(text)
                if m:
                    t[(m.group(1), m.group(2))] = c
                elif k:
                    n[k.group(1)] = c
        return t, n

    def base(wb, action, volume):
        tb = TT.time_table(wb, BM.SPEC["time"])[0]
        return wb[tb["title"]].cell(tb["rows"][action], tb["col"] + tb["cols"][volume] - 1)

    def human(fn, path=None):
        """人が、ブックに値を入れたことにする"""
        wb = load_workbook(path or book)
        fn(wb, *tagged(wb))
        wb.save(path or book)

    def vals():
        d = json.loads((w / "times.json").read_text(encoding="utf-8")) if (w / "times.json").exists() else {}
        return ({(x["flow"], x["node"]): x["value"] for x in d.get("cells") or []}, {x["flow"]: x["value"] for x in d.get("counts") or []},
                {(x["action"], x["volume"]): (x["value"], x.get("from")) for x in d.get("table") or []})

    def now():
        wb = load_workbook(book)
        return (wb,) + tagged(wb)
    soft = lambda c: isinstance(c.value, str) and bool(BM.OURS.match(c.value))
    # 書き出したまま: 式と、見分けの書式
    wb, t, n = now()
    ws = wb[M]
    rules = [f for rng in ws.conditional_formatting for r in rng.rules for f in r.formula]
    looked = lambda key: re.fullmatch(r'=VLOOKUP\("([^"]+)",.+,(\d),FALSE\)', t[key].value).groups()
    out.append(("処理時間: 仮の値は、基準の表を「動作の種類 × 量」で引く式（量は、理由つきの指定・見る帳票 2 つ以上で変わる。出典が 3 つあるだけでは変わらない。判断・システム処理・別のフローへは、機械が決める）",
                all(soft(c) for c in t.values()) and looked(("F01", "P-001")) == ("入力する", "2") and looked(("F01", "P-002")) == ("入力する", "3")
                and looked(("F02", "P-006")) == ("承認する", "3") and looked(("F03", "P-003")) == ("確認する", "4") and looked(("F03", "D-001")) == ("判断する", "3")
                and looked(("F03", "S-001")) == ("システムが処理する", "3") and looked(("F01", "SUB-001")) == ("別のフローへ進む", "3")))
    out.append(("処理時間: 式のセル＝仮（条件付き書式 ISFORMULA）。基準の表に名前が付いている。人の値の置き場（times.json）は、値が入るまで作らない",
                any("ISFORMULA" in f for f in rules) and BM.TIME_NAME in wb.defined_names and BM.TIME_SHEET in wb.sheetnames and not (w / "times.json").exists()
                and all(c.value is None for c in n.values()) and sorted(n) == ["F01", "F02", "F03"]))
    # LibreOffice で計算し直す: 仮の値が表のとおり。表の 1 行を直すと、その種類の仮の値だけが変わる
    office = office_to_pdf.find_soffice()
    if office:
        def shown(path, name):
            d = tmp / name
            d.mkdir()
            subprocess.run([office, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(d), str(path)], capture_output=True, timeout=300)
            got = load_workbook(d / path.name, data_only=True)
            ws = got[M]
            return {k: ws[c.coordinate].value for k, c in t.items()}, [ws.cell(i, 9).value for i in (3, 5, 7)], got[BM.NODES]
        v0, tot0, nl = shown(book, "calc0")
        col = {c.value: k for k, c in enumerate(nl[1])}
        listed = {(r[col["フローID"]].value, r[col["ノード番号"]].value): (r[col["処理時間（分）"]].value, r[col["時間の根拠"]].value) for r in nl.iter_rows(min_row=2)}
        edited = tmp / "edited.xlsx"
        shutil.copy(book, edited)
        human(lambda wb, _t, _n: setattr(base(wb, "入力する", "中"), "value", 30), edited)
        v1, tot1, _nl = shown(edited, "calc1")
        out.append(("処理時間: 計算し直すと、仮の値が基準の表のとおり（入力する×小 2・入力する×中 4・確認する×大 15・判断 2）。合計は、その行の足し算。ノード一覧にも同じ値と「仮」",
                    (v0[("F01", "P-001")], v0[("F01", "P-002")], v0[("F03", "P-003")], v0[("F03", "D-001")]) == (2, 4, 15, 2) and tot0 == [7, 24, 53]
                    and listed[("F01", "P-002")] == (4, "仮") and listed[("F03", "S-001")] == (0, "仮")))
        out.append(("処理時間: 基準の表の「入力する × 中」を 30 に直すと、その種類・量の仮の値だけが変わる（入力する × 小は変わらない。合計も）",
                    {k for k in v0 if v0[k] != v1[k]} == {("F01", "P-002"), ("F03", "P-005"), ("F03", "P-008"), ("F03", "P-011")} and v1[("F01", "P-002")] == 30
                    and v1[("F01", "P-001")] == 2 and tot1 == [33, 24, 131]))
    else:
        out.append(("処理時間: 計算し直しての確かめは、飛ばした（LibreOffice が無い）", True))

    def put(wb, t, n):
        t[("F01", "P-002")].value, t[("F02", "P-001")].value, t[("F02", "P-005")].value, t[("F03", "P-007")].value = 20, "8分", 9, 10
        n["F01"].value, base(wb, "確認する", "中").value = 120, 6
    human(put)
    c0, o0 = run("check_flow.py", w, "--stage", "matrix")
    c1, o1 = run("build_matrix.py", w)
    c2, o2 = run("check_flow.py", w, "--stage", "matrix")
    wb, t, n = now()
    kept = ({("F01", "P-002"): 20, ("F02", "P-001"): 8, ("F02", "P-005"): 9, ("F03", "P-007"): 10}, {"F01": 120}, {("確認する", "中"): (6, "book")})
    out.append(("処理時間: 人が入れた値（数・「8分」）・件数・基準の表の直しは、点検では FAIL にならず（まだ取り込んでいない、と知らせる）、書き出し直すと times.json に入り、ブックに戻る"
                + ("" if c0 == 0 and c1 == 0 and c2 == 0 else o0[-400:] + o1[-400:] + o2[-400:]),
                c0 == 0 and "値の違うセルが 5 個" in o0 and c1 == 0 and vals() == kept and c2 == 0 and "0 FAIL / 0 WARN" in o2
                and [t[k].value for k in sorted(kept[0])] == [20, 8, 9, 10] and sum(soft(c) for c in t.values()) == len(t) - 4 and n["F01"].value == 120 and n["F02"].value is None
                and base(wb, "確認する", "中").value == 6 and base(wb, "確認する", "大").value == 15))
    from validate import validate
    side, snap = json.loads(BM.side_path(book).read_text(encoding="utf-8")), BM.baseline(w)
    cells = {(f["flow"], x["node"]): x for f in snap["flows"] for x in f["cells"]}
    tb = TT.time_table(wb, BM.SPEC["time"])[0]
    note = wb[tb["title"]].cell(min(tb["rows"].values()) - 1, tb["col"]).comment
    moved = re.search(r"基準の表の版 (T-[0-9a-f]{6})（前の書き出し (T-[0-9a-f]{6}) から変わった", o1)
    out.append(("処理時間: 基準の表を直して書き出すと、版（表の値の印）が変わったと知らせ、表の見出しのコメントにも出る。"
                "現状の時間の写しは、形が正しく、合計がセルの足し算・人の値に「人」・版がブックと同じ",
                moved is not None and moved.group(1) == side["table_version"] != moved.group(2) and note is not None and note.text.count("基準の表の版") == 1
                and side["table_version"] in note.text and validate(snap, "baseline") == [] and snap["table_version"] == side["table_version"]
                and snap["total"] == sum(x["value"] or 0 for x in cells.values()) > 0 and snap["human"] == 4 and snap["soft"] == len(t) - 4
                and (cells[("F01", "P-002")]["value"], cells[("F01", "P-002")]["by"]) == (20, "人") and cells[("F01", "P-001")]["by"] == "仮"
                and [f["count"] for f in snap["flows"]] == [120, None, None]))
    c1, o1 = run("check_flow.py", w, "--stage", "matrix")
    c2, o2 = run("build_matrix.py", w, "--calibrate")
    wb, t, n = now()
    out.append(("処理時間: 人の値が 3 つ以上ある「動作の種類 × 量」は、点検が表の値と中央値を並べ、--calibrate で、表を中央値に引き直す",
                "基準の表「確認する × 中」は 6 分・人が入れた 3 セルの中央値は 9 分" in o1 and c2 == 0 and vals()[2] == {("確認する", "中"): (9, "calibrate")} and base(wb, "確認する", "中").value == 9))
    human(lambda wb, t, n: setattr(t[("F01", "P-002")], "value", None))
    run("build_matrix.py", w)
    wb, t, n = now()
    out.append(("処理時間: 人が値を消すと、仮の値（式）に戻る", ("F01", "P-002") not in vals()[0] and soft(t[("F01", "P-002")]) and t[("F02", "P-001")].value == 8))

    def bare(wb, t, n):
        t[("F01", "P-003")].comment, t[("F01", "P-003")].value, n["F03"].comment, n["F03"].value = None, 3, None, 30
    human(bare)
    c1, o1 = run("check_flow.py", w, "--stage", "matrix")
    c2, o2 = run("build_matrix.py", w)
    out.append(("処理時間: コメント（だれの時間かの札）を消したセルの値も、書き出したときの番地で突き合わせて取り込む。点検は WARN",
                c1 == 0 and "コメントが消えている" in o1 and c2 == 0 and "番地で突き合わせた" in o2 and vals()[0].get(("F01", "P-003")) == 3 and vals()[1] == {"F01": 120, "F03": 30}
                and sorted(now()[2]) == ["F01", "F02", "F03"]))
    # 番号の振り直し: 値の番号も、一緒に直る
    for f in (w / "flows/F02.flow.json", w / "times.json"):
        f.write_text(f.read_text(encoding="utf-8").replace("P-005", "P-015"), encoding="utf-8")
    c1, o1 = run("renumber.py", w, "--only", "F02", "--dry-run")
    mid = vals()[0]
    c2, o2 = run("renumber.py", w, "--only", "F02")
    c3, o3 = run("build_matrix.py", w)
    out.append(("処理時間: 番号を振り直すと（renumber.py）、人が入れた値の番号も、一緒に振り直す（--dry-run では変えない）",
                c1 == 0 and mid.get(("F02", "P-015")) == 9 and c2 == 0 and "times.json" in o2 and vals()[0].get(("F02", "P-005")) == 9 and ("F02", "P-015") not in vals()[0]
                and c3 == 0 and "WARN" not in o3 and now()[1][("F02", "P-005")].value == 9))
    # 作業名が変わる・箱が無くなる
    edit(w / "flows/F02.flow.json", lambda f: node(f, "P-001").update(label="申請内容を領収書と照合する"))
    edit(w / "times.json", lambda d: d["cells"].append({"flow": "F01", "node": "P-099", "value": 4, "label": "昔の作業を行う", "book": "業務マトリクス一覧.xlsx", "date": "2026-10-01"}))
    c1, o1 = run("build_matrix.py", w)
    c2, o2 = run("check_flow.py", w, "--stage", "matrix")
    c3, o3 = run("build_matrix.py", w, "--accept", "--forget")
    c4, o4 = run("build_matrix.py", w)
    out.append(("処理時間: 作業名が変わった箱の値は、当てたまま WARN（--accept で認める）。無くなった箱の値は、times.json に残して WARN（--forget で消す）",
                c1 == 0 and "作業名が変わった: F02 P-001" in o1 and "箱が、もう無い: F01 P-099" in o1 and c2 == 0 and "作業名が変わった" in o2 and "もう無い" in o2
                and c3 == 0 and c4 == 0 and "WARN" not in o4 and vals()[0].get(("F02", "P-001")) == 8 and ("F01", "P-099") not in vals()[0]))
    # フローを作り直して、番号がずれる（直しの一覧に、箱を 2 つ足す）: 値は、作業名で付いていく
    e = json.loads((SKILL / "evals/fixture_edits/F03.edits.json").read_text(encoding="utf-8"))
    e["ops"].append({"op": "add", "key": "x9", "type": "process", "after": "M02.1.s1", "label": "一覧を経理課長に渡す", "lane": "経理部", "action": "提出・回付する",
                     "source": {"doc_id": "M02", "page": 1, "quote": "経理部は、毎月末に承認済みの経費申請を一覧で出力する。"}})
    e["ops"].append({"op": "add", "key": "x8", "type": "process", "after": "x9", "label": "一覧を控えに綴じる", "lane": "経理部", "action": "保管する",
                     "source": {"doc_id": "M02", "page": 1, "quote": "経理部は、毎月末に承認済みの経費申請を一覧で出力する。"}})
    (w / "flows/F03.edits.json").write_text(json.dumps(e, ensure_ascii=False), encoding="utf-8")
    c1, o1 = run("apply_edits.py", w, "--flow", "F03")
    c2, o2 = run("build_matrix.py", w)
    c3, o3 = run("check_flow.py", w, "--stage", "matrix")
    f3 = {x["label"]: x["id"] for x in json.loads((w / "flows/F03.flow.json").read_text(encoding="utf-8"))["nodes"]}
    out.append(("処理時間: フローを作り直して番号がずれても、人が入れた値は、同じ作業名の箱に付いていく" + ("" if c1 == 0 and c2 == 0 and c3 == 0 else o1[-300:] + o2[-300:] + o3[-500:]),
                c1 == 0 and c2 == 0 and f3["支払データを確認する"] == "P-008" and "F03 P-007 → P-008" in o2 and vals()[0].get(("F03", "P-008")) == 10 and ("F03", "P-007") not in vals()[0]
                and now()[1][("F03", "P-008")].value == 10 and c3 == 0 and "0 FAIL" in o3))
    # 読めないブック: 人の値を取り込めないので、上書きしない
    whole, before = book.read_bytes(), vals()
    book.write_bytes(b"not a workbook")
    c1, o1 = run("build_matrix.py", w)
    same = book.read_bytes() == b"not a workbook"
    c2, o2 = run("build_matrix.py", w, "--fresh")
    out.append(("処理時間: いまあるブックを読めないときは、上書きせずに止まる。--fresh で、times.json の値から作り直す",
                c1 != 0 and "上書きしない" in o1 and same and c2 == 0 and vals() == before and now()[1][("F02", "P-001")].value == 8 and len(book.read_bytes()) > len(whole) / 2))
    # 基準の表の無いテンプレートで、補うもとのテンプレートも無いとき（前の版のスキルで書き出した、処理時間の行の無いブックと同じ形）: 処理時間の行を出さない
    w3, old = tmp / "times_old", tmp / "matrix_old.xlsx"
    shutil.copytree(good, w3)
    wb = load_workbook(SKILL / "assets" / BM.TEMPLATE)
    del wb.defined_names[BM.TIME_NAME]
    del wb[BM.TIME_SHEET]
    ws = wb[M]
    for c in list(ws[2]):
        if c.value in ("{処理時間の合計}", "{件数}"):
            ws.delete_cols(c.column)
    wb.save(old)
    bare = subprocess.run([sys.executable, "-c", "import build_matrix as BM; BM.SPEC['time'].pop('default'); BM.main()", "-w", str(w3), "--template", str(old), "--calibrate"],
                          capture_output=True, text=True, cwd=str(S))
    c1, o1 = bare.returncode, bare.stdout + bare.stderr
    c2, o2 = run("check_flow.py", w3, "--stage", "matrix")
    got = load_workbook(w3 / "out" / BM.BOOK)
    out.append(("処理時間: 基準の表を補えないとき・前の版のスキルで書き出したブック（処理時間の行が無い。1 フロー＝1 行）も、点検が通る" + ("" if c1 == 0 and c2 == 0 else o1[-400:] + o2[-400:]),
                c1 == 0 and "処理時間は扱わない" in o1 and c2 == 0 and "0 FAIL / 0 WARN" in o2 and got[M].max_row == 4 and not any(tagged(got)) and not (w3 / "times.json").exists()
                and "①" in [c.value for c in got[M][1]]))
    # 動作の種類の無い箱（古い作業フォルダ）: 止めない
    edit(w3 / "flows/F01.flow.json", lambda f: node(f, "P-001").pop("action"))
    c1, o1 = run("check_flow.py", w3, "--stage", "flow")
    c2, _o = run("build_matrix.py", w3)
    out.append(("処理時間: 動作の種類（action）の無い作業は、点検が WARN（止めない）。仮の値は「その他」の行を引く",
                c1 == 0 and "動作の種類（action）の無い作業が 1 個（P-001）" in o1 and "0 FAIL / 1 WARN" in o1 and c2 == 0
                and 'VLOOKUP("その他"' in str(tagged(load_workbook(w3 / "out" / BM.BOOK))[0][("F01", "P-001")].value)))
    # システムの画面で作る箱（使うシステムの札あり）は「入力する」（決まり 3.22）。札が無ければ「作成する」のままでよい
    edit(w3 / "flows/F01.flow.json", lambda f: (node(f, "P-001").update(action="入力する"), node(f, "P-002").update(action="作成する")))
    edit(w3 / "flows/F02.flow.json", lambda f: node(f, "P-005").update(action="作成する"))
    c1, o1 = run("check_flow.py", w3, "--stage", "flow")
    out.append(("処理時間: 使うシステムの札が付いた作業の動作の種類が「作成する」なら、点検が WARN で「入力する」を案内する（札の無い箱は出さない）",
                c1 == 0 and "「作成する」の箱が 1 個（P-002）" in o1 and "F01:" in o1 and "F02: 使うシステムの札" not in o1 and "0 FAIL / 1 WARN" in o1
                and node(json.loads((w3 / "flows/F02.flow.json").read_text(encoding="utf-8")), "P-005").get("systems") in (None, [])))
    return out


CODE_SAMPLE = """Attribute VB_Name = "集計"
Option Explicit
Private Const 置き場 As String = "\\\\srv\\share\\"

' ボタンから動く
Public Sub Main()
    Dim ws As Worksheet, 月 As String, i As Long
    On Error GoTo ErrHandler
    Application.ScreenUpdating = False
    月 = Worksheets("メニュー").Range("B2").Value
    If 月 = "" Then
        MsgBox "対象月を入れてください"
        Exit Sub
    End If
    If MsgBox("始めますか", vbYesNo) = vbNo Then Exit Sub
    Set ws = Worksheets("一覧")
    For i = 2 To 10
        ws.Cells(i, 3).Value = 税額(ws.Cells(i, 2).Value)
    Next i
    Call 保存する(月)
    Exit Sub
ErrHandler:
    MsgBox "エラー: " & Err.Description
End Sub

Private Sub 保存する(月 As String)
    Worksheets("一覧").ExportAsFixedFormat Type:=0, Filename:=置き場 & "一覧_" & 月 & ".pdf"
End Sub

Public Sub 取込()
    Dim cn As Object
    Set cn = CreateObject("ADODB.Connection")
    cn.Open "Provider=X;Data Source=db01;Initial Catalog=hanbai;User ID=u;Password=himitsu99;"
    cn.Execute "UPDATE 入金 SET 済 = 1 WHERE 額 > " & 税額(100)
End Sub

Public Function 税額(額 As Currency) As Currency
    If 額 < 0 Then
        税額 = 0
    Else
        税額 = Int(額 * 0.1)
    End If
End Function
"""


LANG_SAMPLES = {
    'PY': '"""入金データを取り込む（架空）"""\nimport csv\nimport sqlite3\nimport sys\n\nDB = "keiri.db"\n\n\ndef load_rows(path):\n    rows = []\n    with open(path, encoding="cp932") as f:\n        for row in csv.DictReader(f):\n            if not row["請求番号"]:\n                continue\n            rows.append(row)\n    return rows\n\n\ndef save(rows):\n    con = sqlite3.connect(DB)\n    for r in rows:\n        # 金額がマイナスなら、エラーの表に入れる\n        if int(r["入金額"]) < 0:\n            con.execute("INSERT INTO 入金エラー (請求番号) VALUES (?)", (r["請求番号"],))\n        else:\n            con.execute("UPDATE 請求 SET 入金額 = ? WHERE 請求番号 = ?", (r["入金額"], r["請求番号"]))\n    con.commit()\n    con.close()\n\n\nif __name__ == "__main__":\n    month = input("対象月（yyyymm）を入れてください: ")\n    try:\n        rows = load_rows(f"\\\\\\\\srv\\\\share\\\\入金_{month}.csv")\n    except FileNotFoundError:\n        print("ファイルがありません")\n        sys.exit(1)\n    save(rows)\n    print(f"{len(rows)} 件を取り込みました")\n',
    'SQL': "-- 月次の締め（架空）\nUSE keiri;\n-- 当月分を締めの表に移す\nINSERT INTO 月次締め (年月, 取引先, 金額)\nSELECT 年月, 取引先, SUM(金額) FROM 売上明細 WHERE 年月 = '202609' GROUP BY 年月, 取引先;\nUPDATE 売上明細 SET 締め済 = 1 WHERE 年月 = '202609';\nDELETE FROM 作業用 WHERE 年月 < '202601';\n",
    'BAT': '@echo off\nrem 夜間のバックアップ（架空）\nset SRC=\\\\srv\\keiri\\月次\nset DST=D:\\backup\\%date:~0,4%\nif not exist %DST% mkdir %DST%\nxcopy /s /y %SRC%\\*.xlsx %DST%\nif errorlevel 1 goto :error\ncall :notify\nsqlcmd -S db01 -i shime.sql\nexit /b 0\n\n:error\necho コピーに失敗しました >> D:\\backup\\error.log\nexit /b 1\n\n:notify\necho 終わりました\nexit /b\n',
    'PS1': '# 請求データを書き出す（架空）\nparam($Month)\n$rows = Import-Csv "\\\\srv\\share\\請求_$Month.csv"\nforeach ($r in $rows) {\n    if ($r.金額 -ge 1000000) {\n        Add-Content -Path "D:\\out\\要確認.txt" -Value $r.請求番号\n    } else {\n        Write-Host $r.請求番号\n    }\n}\ntry {\n    Send-MailMessage -To "kacho@example.com" -Subject "書き出し完了"\n} catch {\n    Out-File -FilePath "D:\\out\\error.log" -InputObject $_\n}\n',
    'SH': '#!/bin/sh\n# 日次の取込（架空）\nDIR=/data/in\nif [ ! -f "$DIR/uriage.csv" ]; then\n    echo "ファイルが無い"\n    exit 1\nfi\nfor f in $DIR/*.csv; do\n    cp "$f" /data/work/\ndone\npsql -f load.sql\nmail -s "取込完了" kacho@example.com < /data/work/result.txt\n',
}


def lang_tests():
    """VBA 以外の言語の構造の取り出し（Python は構文から、SQL・バッチ・PowerShell・シェルは行の形から）"""
    import code_struct as cs
    out = []

    def build(lang, name, key):
        st = cs.build_struct("M09", lang, [{"name": name, "kind": "file", "lines": LANG_SAMPLES[key].splitlines()}])
        procs = {p["name"]: p for p in st["procedures"]}
        io = {(x["kind"], x["name"], x["way"]) for p in st["procedures"] for x in p["io"]}
        return st, procs, io
    st, procs, io = build("python", "nyukin.py", "PY")
    kinds = [n["kind"] for n in procs["（本体）"]["body"]]
    out.append(("Python: スクリプトの本体が起動点。入力の求め・関数の呼び出し・例外の分岐・途中の終了を拾う",
                st["tier"] == "parsed" and [e["trigger"] for e in st["entries"]] == ["script"] and kinds == ["interact", "call", "if", "call", "seq"]
                and procs["（本体）"]["body"][2]["cond"].startswith("例外が起きたとき") and st["inline"] == ["nyukin.py.load_rows", "nyukin.py.save"]))
    out.append(("Python: ファイル・データベース・SQL のテーブル（読む・書く）と、繰り返しの中の continue", {("テーブル", "入金エラー", "out"), ("テーブル", "請求", "out"),
                ("データベース", "keiri.db", "in")} <= io and any(n["kind"] == "exit" and n["what"] == "continue" for n in __import__("code_lang").walk(procs["load_rows"]["body"]))
                and cs.tree_errors(st) == []))
    out.append(("Python: 読めない書き方のファイルは、機械で読まずに LLM の読取に回す", cs.build_struct("M09", "python", [{"name": "x.py", "lines": ["def f(:", "  pass"]}])["tier"] == "llm"))
    st, procs, io = build("sql", "shime.sql", "SQL")
    out.append(("SQL: 文ごとに、読むテーブルと書くテーブル", st["tier"] == "heuristic" and {("テーブル", "月次締め", "out"), ("テーブル", "売上明細", "in"), ("テーブル", "売上明細", "out"),
                ("テーブル", "作業用", "out")} <= io and ("テーブル", "作業用", "in") not in io and cs.tree_errors(st) == []))
    st, procs, io = build("batch", "backup.bat", "BAT")
    body = procs["（本体）"]["body"]
    out.append(("バッチ: 1行の if・goto とラベル・call :ラベル（別の手続き）・ファイルの操作・外部の道具",
                [n["kind"] for n in body] == ["seq", "if", "seq", "if", "call", "seq", "exit", "label", "seq", "exit"] and body[3]["branches"][0]["body"][0]["what"] == "goto"
                and "notify" in procs and ("ファイル", "error.log", "out") in io and ("データベース", "sqlcmd", "out") in io and cs.tree_errors(st) == []))
    st, procs, io = build("powershell", "export.ps1", "PS1")
    body = procs["（本体）"]["body"]
    out.append(("PowerShell: foreach・if/else・try/catch、変数の入ったファイル名", [n["kind"] for n in body] == ["seq", "seq", "loop", "seq", "if"]
                and ("ファイル", "請求_＜Month＞.csv", "in") in io and ("ファイル", "要確認.txt", "out") in io and ("メール", "Send-MailMessage", "out") in io and cs.tree_errors(st) == []))
    st, procs, io = build("shell", "daily.sh", "SH")
    body = procs["（本体）"]["body"]
    out.append(("シェル: if〜fi・for〜done・途中の終了、cp と外部の道具", [n["kind"] for n in body] == ["seq", "if", "loop", "seq", "seq"] and body[1]["branches"][0]["body"][-1]["kind"] == "exit"
                and ("データベース", "psql", "out") in io and ("メール", "mail", "out") in io and cs.tree_errors(st) == []))
    return out


TREE_JAVA = r"""package example.batch;

import java.io.*;
import java.sql.*;
import java.util.*;

/** 受注データを取り込むバッチ（架空） */
public class OrderBatch {

    private static final String IN_FILE = "orders.csv";

    public static void main(String[] args) throws Exception {
        List<String[]> rows = readRows(IN_FILE);
        int errors = 0;
        if (rows.isEmpty()) {
            System.out.println("取り込むデータがありません");
            return;
        }
        // 1件ずつ登録する
        for (String[] row : rows) {
            if (row.length < 3) {
                errors++;
                continue;
            }
            register(row);
        }
        try {
            writeReport(rows.size(), errors);
        } catch (IOException e) {
            System.err.println("結果を書けませんでした");
            throw e;
        }
        System.out.println("done");
    }

    static List<String[]> readRows(String path) throws IOException {
        List<String[]> out = new ArrayList<>();
        BufferedReader br = new BufferedReader(new FileReader(path));
        String line;
        while ((line = br.readLine()) != null) {
            out.add(line.split(","));
        }
        br.close();
        return out;
    }

    static void register(String[] row) throws Exception {
        Connection cn = DriverManager.getConnection("jdbc:sqlserver://db01;databaseName=hanbai;user=batch;password=himitsu99");
        PreparedStatement ps = cn.prepareStatement("INSERT INTO 受注 (受注番号, 顧客コード, 金額) VALUES (?, ?, ?)");
        ps.setString(1, row[0]);
        ps.executeUpdate();
        switch (row[2]) {
            case "0":
                cn.createStatement().executeUpdate("UPDATE 顧客 SET 最終受注日 = GETDATE()");
                break;
            default:
                break;
        }
        cn.close();
    }

    static void writeReport(int total, int errors) throws IOException {
        PrintWriter pw = new PrintWriter(new FileWriter("取込結果.txt"));
        pw.println("件数: " + total + " エラー: " + errors);
        pw.close();
    }
}
"""

TREE_COBOL = r"""000100 IDENTIFICATION DIVISION.
000200 PROGRAM-ID. URIAGE01.
000300*    売上の月次集計（架空）
000400 PROCEDURE DIVISION.
000500 MAIN-RTN.
000600     OPEN INPUT URIAGE-FILE
000700     PERFORM UNTIL WS-EOF = 'Y'
000800         READ URIAGE-FILE
000900             AT END MOVE 'Y' TO WS-EOF
001000         END-READ
001100         IF WS-EOF NOT = 'Y'
001200             PERFORM KEISAN-RTN
001300         END-IF
001400     END-PERFORM
001500     CLOSE URIAGE-FILE
001600     STOP RUN.
001700 KEISAN-RTN.
001800     EVALUATE UR-KUBUN
001900         WHEN '1'
002000             ADD UR-KINGAKU TO WS-GOKEI
002100         WHEN OTHER
002200             DISPLAY 'KUBUN ERROR'
002300     END-EVALUATE.
"""


def tree_samples():
    """架空の Java・COBOL と、そのブロックの木（LLM が書くもの。行の番号は、ソースの文字から引く）"""
    J = TREE_JAVA.splitlines()
    at = lambda text: next(k for k, ln in enumerate(J, 1) if text in ln)
    M, C = "OrderBatch.java", "URIAGE01.cbl"
    one = lambda kind, text, **kw: dict({"kind": kind, "lines": [at(text), at(text)]}, **kw)
    main, rd, reg, rep = at("void main"), at("readRows(String path)"), at("void register"), at("void writeReport")
    java = {"project": "M01", "procedures": [
        {"module": M, "name": "main", "lines": [main, rd - 2], "body": [
            {"kind": "call", "lines": [main + 1, main + 1], "to": f"{M}.readRows"},
            {"kind": "seq", "lines": [main + 2, main + 2], "tech": "変数の準備"},
            {"kind": "if", "lines": [main + 3, main + 6], "branches": [{"label": "then", "body": [one("interact", "取り込むデータがありません", how="msgbox"), one("exit", "return;", what="proc")]}]},
            {"kind": "loop", "lines": [at("for (String[] row"), at("register(row);") + 1], "body": [
                {"kind": "if", "lines": [at("row.length < 3"), at("row.length < 3") + 3], "branches": [{"label": "then", "body": [one("seq", "errors++"), one("exit", "continue;", what="continue")]}]},
                one("call", "register(row);", to=f"{M}.register")]},
            one("call", "writeReport(rows", to=f"{M}.writeReport"),
            {"kind": "if", "lines": [at("catch (IOException"), at("throw e;") + 1], "branches": [{"label": "then", "body": [one("interact", "System.err", how="msgbox"), one("exit", "throw e;", what="proc")]}]},
            one("seq", '"done"', tech="ログ")]},
        {"module": M, "name": "readRows", "lines": [rd, reg - 2], "body": [
            {"kind": "seq", "lines": [rd + 1, rd + 3], "io": [{"kind": "ファイル", "name": "orders.csv", "way": "in", "line": at("IN_FILE =")}]},
            {"kind": "loop", "lines": [at("while ((line"), at("while ((line") + 2], "body": [one("seq", "out.add")]},
            one("seq", "br.close", tech="後片付け"), one("exit", "return out;", what="proc")]},
        {"module": M, "name": "register", "lines": [reg, rep - 2], "body": [
            {"kind": "seq", "lines": [reg + 1, reg + 4], "io": [{"kind": "データベース", "name": "hanbai", "way": "out", "line": reg + 1}, {"kind": "テーブル", "name": "受注", "way": "out", "line": reg + 2}]},
            {"kind": "select", "lines": [at("switch (row[2])"), at("default:") + 2], "branches": [
                {"label": "case", "head": [at('case "0":'), at('case "0":')], "body": [
                    one("seq", "UPDATE 顧客", io=[{"kind": "テーブル", "name": "顧客", "way": "out", "line": at("UPDATE 顧客")}]),
                    {"kind": "exit", "lines": [at("UPDATE 顧客") + 1, at("UPDATE 顧客") + 1], "what": "select"}]},
                {"label": "else", "head": [at("default:"), at("default:")], "body": [{"kind": "exit", "lines": [at("default:") + 1, at("default:") + 1], "what": "select"}]}]},
            one("seq", "cn.close", tech="後片付け")]},
        {"module": M, "name": "writeReport", "lines": [rep, rep + 4], "body": [
            {"kind": "seq", "lines": [rep + 1, rep + 3], "io": [{"kind": "ファイル", "name": "取込結果.txt", "way": "out", "line": rep + 1}]}]}],
        "entries": [{"procedure": f"{M}.main", "trigger": "entry", "label": "バッチ「OrderBatch」を実行", "line": main}]}
    cobol = {"project": "M02", "procedures": [
        {"module": C, "name": "MAIN-RTN", "lines": [5, 16], "body": [
            {"kind": "seq", "lines": [6, 6], "io": [{"kind": "ファイル", "name": "URIAGE-FILE", "way": "in", "line": 6}]},
            {"kind": "loop", "lines": [7, 14], "body": [{"kind": "seq", "lines": [8, 10]},
                                                       {"kind": "if", "lines": [11, 13], "branches": [{"label": "then", "body": [{"kind": "call", "lines": [12, 12], "to": f"{C}.KEISAN-RTN"}]}]}]},
            {"kind": "seq", "lines": [15, 15], "tech": "後片付け"}, {"kind": "exit", "lines": [16, 16], "what": "proc"}]},
        {"module": C, "name": "KEISAN-RTN", "lines": [17, 23], "body": [
            {"kind": "select", "lines": [18, 23], "branches": [{"label": "case", "head": [19, 19], "body": [{"kind": "seq", "lines": [20, 20]}]},
                                                              {"label": "else", "head": [21, 21], "body": [{"kind": "interact", "lines": [22, 22], "how": "msgbox"}]}]}]}],
        "entries": [{"procedure": f"{C}.MAIN-RTN", "trigger": "entry", "label": "プログラム「URIAGE01」を実行", "line": 2}]}
    return (M, J, java), (C, TREE_COBOL.splitlines(), cobol)


def tree_tests(tmp):
    """機械で構文を読めない言語: LLM が書いたブロックの木を、機械が確かめて構造にする。誤りを1つずつ仕込む"""
    import copy

    import code_lang
    import code_struct as cs
    out = []
    (M, J, java), (C, K, cobol) = tree_samples()
    jm, km = [{"name": M, "kind": "file", "lines": J}], [{"name": C, "kind": "file", "lines": K}]
    st = cs.build_struct("M01", "java", jm, tree=java)
    procs = {p["name"]: p for p in st["procedures"]}
    io = {(x["kind"], x["name"], x["way"]) for p in st["procedures"] for x in p["io"]}
    out.append(("木が無い: 手続きは空で、木が無いことが構造に残る（点検が FAIL にする）", cs.build_struct("M01", "java", jm)["read"] == {"tree": False, "total_lines": len(code_lang.plain_lines("java", J)[0])}))
    out.append(("Java の木 → 構造: 手続き 4・起動点 1・条件の文字と実行文の行は機械がソースから入れる・呼ぶ側の図に入れる手続き 3",
                st["tier"] == "llm" and code_lang.tree_problems(java, jm, "java") == [] and len(procs) == 4 and [e["label"] for e in st["entries"]] == ["バッチ「OrderBatch」を実行"]
                and procs["main"]["body"][2]["cond"].startswith("if (rows.isEmpty())") and len(st["inline"]) == 3 and cs.tree_errors(st) == []
                and {("ファイル", "orders.csv", "in"), ("テーブル", "受注", "out"), ("テーブル", "顧客", "out"), ("ファイル", "取込結果.txt", "out")} <= io
                and st["read"]["tree"] and st["read"]["lines"] < st["read"]["total_lines"]))
    st2 = cs.build_struct("M02", "cobol", km, tree=cobol)
    p2 = {p["name"]: p for p in st2["procedures"]}
    out.append(("COBOL の木 → 構造: 行の番号の欄と注釈行（7 桁目の *）を除く・PERFORM の繰り返し・EVALUATE の振り分け・STOP RUN",
                code_lang.tree_problems(cobol, km, "cobol") == [] and [n["kind"] for n in p2["MAIN-RTN"]["body"]] == ["seq", "loop", "seq", "exit"]
                and p2["MAIN-RTN"]["body"][1]["cond"] == "PERFORM UNTIL WS-EOF = 'Y'" and p2["KEISAN-RTN"]["body"][0]["branches"][1]["body"][0]["interact"]["text"] == "KUBUN ERROR"
                and 3 not in code_lang.plain_lines("cobol", K)[0] and cs.tree_errors(st2) == []))

    def case(name, change, expect, tree=java, mods=jm, lang="java"):
        t = copy.deepcopy(tree)
        change(t)
        errs = code_lang.tree_problems(t, mods, lang)
        hit = next((e for e in errs if expect in e), None)
        out.append((f"木の誤り: {name}: {(hit or '')[:100]}" if hit else f"木の誤り: {name}: 「{expect}」が出ない\n{errs[:4]}", bool(hit)))
    main = lambda t: t["procedures"][0]["body"]
    case("まとまりを1つ消す（行の網羅）", lambda t: main(t).pop(), "どのノードにも入っていない行")
    case("ノードの範囲を広げて、次のノードと重ねる", lambda t: main(t)[1]["lines"].__setitem__(1, main(t)[1]["lines"][1] + 1), "重なっているか、順が逆")
    case("ふつうの文を分岐と書く", lambda t: main(t)[1].update(kind="if", branches=[{"label": "then", "body": []}]), "分岐の語が無い")
    case("繰り返しでない行を繰り返しと書く", lambda t: main(t)[2].update(kind="loop", body=main(t)[2].pop("branches")[0]["body"]), "繰り返しの語が無い")
    case("木に無い手続きを呼ぶ", lambda t: main(t)[0].update(to="OrderBatch.java.load"), "木の手続きに無い")
    case("呼び先の名前が、その行に無い", lambda t: main(t)[0].update(to="OrderBatch.java.register"), "呼び先の名前")
    case("入出力の名前が、その行に無い", lambda t: t["procedures"][3]["body"][0]["io"][0].update(name="kekka.txt"), "行の文字に無い")
    case("抜ける語の無い行を、抜けると書く", lambda t: main(t)[1].update(kind="exit", what="proc") or main(t)[1].pop("tech"), "抜ける・戻る語")
    case("手続きの範囲が、ほかの手続きと重なる", lambda t: t["procedures"][1]["lines"].__setitem__(0, t["procedures"][0]["lines"][1]), "重なる")
    case("起動点の名前が、根拠の行にもファイル名にも無い", lambda t: t["entries"][0].update(label="バッチ「受注取込」を実行"), "かぎかっこの中の名前")
    case("COBOL: 注釈でない行を落とす", lambda t: t["procedures"][0]["body"].pop(2), "どのノードにも入っていない行", cobol, km, "cobol")
    other = copy.deepcopy(java)
    other["procedures"][0]["body"][1]["lines"] = [other["procedures"][0]["body"][1]["lines"][0]] * 2
    other["procedures"][3]["body"] = [{"kind": "seq", "lines": [x, x]} for x in range(other["procedures"][3]["lines"][0] + 1, other["procedures"][3]["lines"][0] + 4)]
    ag = code_lang.tree_agreement(java, other)
    kind = copy.deepcopy(java)
    kind["procedures"][3]["body"][0]["kind"] = "interact"
    ag3 = code_lang.tree_agreement(java, kind)
    out.append(("2回読んだ木の一致: 手続きは全部一致。まとまりの区切り方だけの違いは、つないで比べるので一致（書いたままでは 4 つ違う）。種類の違いは、食い違い",
                ag["procedures"] == [4, 4] and ag["rate"] == 1.0 and ag["raw"][0] == ag["raw"][1] - 4 and 0.8 < ag3["rate"] < 1 and ag3["flow"][0] < ag3["flow"][1]
                and code_lang.tree_agreement(java, java)["rate"] == 1.0))
    # 通しで: ページ化 → 木 → 構造 → 台帳 → アクティビティ図（名付けなし）→ 点検（コードの段）
    src, w = tmp / "tree_in", tmp / "tree_work"
    (src / "受注バッチ").mkdir(parents=True)
    (src / "売上集計").mkdir(parents=True)
    (src / "受注バッチ" / M).write_text(TREE_JAVA, encoding="utf-8")
    (src / "売上集計" / C).write_text(TREE_COBOL, encoding="utf-8")
    subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(w)], capture_output=True, text=True)
    run("render_pages.py", w)
    run("code_struct.py", w)
    code0, text0 = run("check_flow.py", w, "--stage", "code")
    (w / "code").mkdir(exist_ok=True)
    (w / "code/M01.tree.json").write_text(json.dumps(java, ensure_ascii=False), encoding="utf-8")
    (w / "code/M02.tree.json").write_text(json.dumps({"project": "M02", "skip": "利用者と決めた範囲の外", "procedures": [], "entries": []}, ensure_ascii=False), encoding="utf-8")
    run("code_struct.py", w)
    code1, text1 = run("check_flow.py", w, "--stage", "code")
    run("systems.py", w, "draft")
    edit(w / "structure/systems.json", lambda j: [s.update(kind="core") for s in j["systems"]])
    code2, text2 = run("activity.py", w, "apply", "--all", "--no-edits")
    run("build_activity.py", w)
    code3, text3 = run("check_layout.py", w, "--no-png")
    act = json.loads((w / "activities/A-01-1.activity.json").read_text(encoding="utf-8")) if (w / "activities/A-01-1.activity.json").exists() else {"nodes": [], "lines": []}
    page = (w / "pages/M01/p0001.json").read_text(encoding="utf-8")
    out.append(("通し（Java の木あり・COBOL は読まないと決めた）: 木が無いうちは FAIL、書くと FAIL 0。図は「LLM の読取」と出て、実行文が箱か対象外に1回ずつ入る。認証情報は伏せる",
                code0 != 0 and text0.count("tree.json が無い") == 2 and code1 == 0 and "0 FAIL" in text1 and code2 == 0 and "0 FAIL" in text3 and act.get("tier") == "llm"
                and len({ln for o in act["lines"] for ln in o["lines"]}) == sum(len(o["lines"]) for o in act["lines"]) == 36 and "himitsu99" not in page))
    from systems import machine_questions
    local = {k for k, _q in machine_questions(w)}
    out.append(("機械が作る確認事項: 読まないと決めたまとまり・読んだ範囲が一部だけのまとまり", {"skipped:M02"} <= local and "partial:M01" not in local))
    code4, text4 = (edit(w / "code/M01.tree.json", lambda t: t["procedures"][0]["body"].pop()), run("check_flow.py", w, "--stage", "code"))[1]
    out.append(("木を直したあと作り直していない: 取り出し直した構造と違う", code4 != 0 and "LLM が書いた木から取り出し直した構造と違う" in text4))
    return out


TREE_WEB = r"""package example.web;

public class OrderController {

    @PostMapping("/orders/new")
    public String create(Order order, BindingResult result, RedirectAttributes ra) {
        if (order == null) {
            abort("empty");
        }
        if (isDuplicate(order)) {
            result.rejectValue("code", "duplicate", "同じ番号の注文があります");
            return "orders/form";
        }
        // 保存する
        this.orders.save(order);
        int n = 0;
        log.info("saved");
        ra.addFlashAttribute("message", "注文を登録しました");
        return finish(order);
    }

    private boolean isDuplicate(Order order) {
        return this.orders.findByCode(order.getCode()) != null;
    }

    private String finish(Order order) {
        order.setStatus("done");
        // ---------------------------------
        // 担当者に知らせる
        // ---------------------------------
        this.mailer.send(order);
        if (order.isUrgent()) {
            abort("urgent");
        }
        if (order.isClosed()) {
            stop();
        }
        return "redirect:/orders";
    }

    private void stop() {
        System.exit(1);
    }

    private void abort(String why) {
        log.error(why);
        System.exit(9);
    }
}
"""

VBA_END = """Sub Main()
    Call 確認
    中止
    MsgBox "続き"
End Sub
Sub 確認()
    If Range("A1").Value = "" Then
        MsgBox "未入力"
        End
    End If
End Sub
Sub 中止()
    MsgBox "中止します"
    End
End Sub
"""


TREE_HOOK = r"""@RequestMapping("/shop")
public class ItemController {

    @ModelAttribute("shop")
    public Shop findShop(int shopId) {
        return this.shops.findById(shopId);
    }

    @PostMapping("/items/new")
    public String create(Item item) {
        this.items.save(item);
        return "redirect:/items";
    }

    @PostMapping("/items/edit")
    public String update(Item item) {
        this.items.save(item);
        return "redirect:/items";
    }
}
"""

VBA_LOOP = """Sub 探す()
    Do
        r = r + 1
        If Cells(r, 1).Value = "" Then
            Exit Do
        End If
        Cells(r, 2).Value = 1
    Loop
    MsgBox "終わり"
End Sub
"""


def schema_tests(tmp, good):
    """決まった語: コードに残した定数・呼び名の表が、スキーマの決まった語と同じか。
    形の確かめ: 人・LLM が書くもの、機械が読み戻すものは、崩れていれば止まる（黙って捨てない・振り直さない）"""
    import activity
    import apply_edits
    import build_drawio
    import build_matrix as BM
    import common as C
    import extract_access
    import systems
    import table_template as TT
    import terms
    from validate import validate

    def enum(kind, *path):
        return C.schema_enum(kind, *path)
    out = [
        ("決まった語: 処理時間の動作の種類（コードとスキーマ）", BM.TIME_ACTIONS == enum("times", "$defs", "row", "properties", "action")),
        ("決まった語: アクティビティ図の、図にしない理由・対象外の理由", list(dict.fromkeys(list(activity.FOLD_REASONS) + list(activity.EXCLUDE_REASONS)))
         == enum("activity_edits", "$defs", "op", "properties", "reason")),
        ("決まった語: 用語の対応の根拠", list(terms.BASIS) == enum("code_terms", "$defs", "term", "properties", "basis")),
        ("決まった語: Access の書き出しの種類", list(extract_access.LANGUAGE) == enum("manifest", "$defs", "doc", "properties", "access", "properties", "object")),
        ("決まった語: 呼び名の表（重要度・ノード・資料の役割・資料の関係・階層・システムの種類・結び付きの根拠・階層の根拠）",
         list(C.SEV_LABEL) == enum("common", "$defs", "question", "properties", "severity") and set(C.NODE_LABEL) == set(C.NODE_TYPES)
         and list(C.ROLE_LABEL) == enum("hierarchy", "$defs", "doc", "properties", "role")
         and list(C.REL_LABEL) == enum("hierarchy", "$defs", "relation", "properties", "type")
         and list(C.HIERARCHY_LABEL) == enum("hierarchy", "properties", "is_hierarchical")
         and list(C.SYSTEM_KIND) == enum("systems", "$defs", "system", "properties", "kind")
         and list(systems.BASIS) == enum("systems", "$defs", "link", "properties", "basis")
         and set(build_drawio.BASIS) == set(C.schema_enum("hierarchy", "properties", "basis", "items"))),
        ("決まった語: スキーマの中の写し（直しの一覧の動作の種類・量・直しの種類、確認事項の下書きの種別・答えで変わる所）",
         enum("edits", "$defs", "op", "properties", "action") == enum("edits", "$defs", "op", "properties", "parts", "items", "properties", "action")
         == enum("flow", "$defs", "node", "properties", "action") == list(C.ACTIONS)
         and enum("edits", "$defs", "op", "properties", "volume") == enum("edits", "$defs", "op", "properties", "parts", "items", "properties", "volume")
         == enum("flow", "$defs", "node", "properties", "volume")
         and enum("common", "$defs", "question_draft", "properties", "category") == enum("common", "$defs", "question", "properties", "category")
         and enum("common", "$defs", "question_draft", "properties", "effect") == enum("common", "$defs", "question", "properties", "effect")
         and list(apply_edits.OPS) == enum("edits", "$defs", "op", "properties", "op") and list(activity.OPS) == enum("activity_edits", "$defs", "op", "properties", "op")),
        ("決まった語: ID の書式（資料・フロー・図・ノード・業務の木・確認事項の番号）は common の 1 か所。ほかのスキーマは写さずに引く",
         not [(f.name, p) for f in sorted((C.ASSETS / "schemas").glob("*.schema.json")) if f.name != "common.schema.json"
              for p in re.findall(r'"pattern":\s*"((?:[^"\\]|\\.)*)"', f.read_text(encoding="utf-8"))
              if json.loads(f'"{p}"') in {d.get("pattern") for d in json.loads((C.ASSETS / "schemas/common.schema.json").read_text(encoding="utf-8"))["$defs"].values()}]),
        ("決まった語: 入出力の種類（機械の取り出しの語を、業務の語にそろえる。ブロックの木も同じ語を引く）", set(C.IO_KIND_WAS.values()) <= set(C.IO_KINDS)
         and json.loads((C.ASSETS / "schemas/code_tree.schema.json").read_text(encoding="utf-8"))["$defs"]["io"]["properties"]["kind"]["$ref"] == "common#/$defs/io_kind"),
        ("決まった語: 照合の報告の種類は、採否表に書ける", set(enum("report_review", "$defs", "finding", "properties", "kind"))
         | set(enum("report_code_review", "$defs", "finding", "properties", "kind")) <= set(enum("review", "$defs", "finding", "properties", "kind"))),
    ]
    m = json.loads((good / "manifest.json").read_text(encoding="utf-8"))
    pj = json.loads((good / "pages" / "M01" / "p0001.json").read_text(encoding="utf-8"))
    out.append(("形: 棚卸し・機械抽出に、宣言していない項目があれば違反（正しいものは通る）",
                not validate(m, "manifest") and not validate(pj, "page") and validate(dict(m, extra=1), "manifest")
                and validate(dict(m, docs=[dict(m["docs"][0], extra=1)]), "manifest") and validate(dict(pj, extra=1), "page")))
    rv = {"verdicts": [{"node": "P-001", "verdict": "支持", "evidence": "「申請書を作成する」"}], "findings": []}
    out.append(("形: 照合の報告（作り方を知らない担当が書く）を確かめられる",
                not validate(rv, "report_review") and validate(dict(rv, verdicts=[{"node": "P-1", "verdict": "OK"}]), "report_review")
                and not validate({"activities": ["A-01-1"], "checked": {"A-01-1": ["P-001"]}, "findings": [{"activity": "A-01-1", "node": "P-001", "kind": "名前が違う", "what": "x"}]}, "report_code_review")
                and validate({"activities": ["A-01-1"], "findings": []}, "report_code_review")
                and validate({"activities": ["A-01-1"], "findings": [{"activity": "A-01-1", "kind": "違う", "what": "x"}]}, "report_code_review")))
    w = tmp / "shape"
    shutil.copytree(good, w)
    bad_times = json.dumps({"version": 1, "cells": [{"flow": "F01", "node": "P-001", "value": "十"}], "counts": [], "table": []}, ensure_ascii=False)
    (w / "times.json").write_text(bad_times, encoding="utf-8")
    code, text = run("build_matrix.py", w)
    out.append(("形: 人が入れた処理時間（times.json）が崩れていれば、書き出さずに止まる（黙って捨てて上書きしない）",
                code != 0 and "times.json の形が崩れている" in text and (w / "times.json").read_text(encoding="utf-8") == bad_times))
    (w / "times.json").unlink()
    qs = (w / "questions.json").read_text(encoding="utf-8")
    broken = json.dumps({"next": 5, "questions": [{"origin": "flow:F01"}], "assigned": {}}, ensure_ascii=False)
    (w / "questions.json").write_text(broken, encoding="utf-8")
    code, text = run("build_reports.py", w)
    out.append(("形: Q 番号の台帳（questions.json）が崩れていれば、Q-001 から振り直さずに止まる",
                code != 0 and "questions.json の形が崩れている" in text and (w / "questions.json").read_text(encoding="utf-8") == broken))
    (w / "questions.json").write_text(qs, encoding="utf-8")
    (w / "reports").mkdir(exist_ok=True)
    (w / "reports" / "read_M01_0001.json").write_text(json.dumps({"doc_id": "M01", "pages_done": [1], "pages_left": "2"}), encoding="utf-8")
    code, text = run("pages.py", w, "batches")
    _c, st = run("status.py", w)
    out.append(("形: 読取の報告が崩れていれば、束を作らない（読み残したページが束に戻らなくなる）。進み具合にも出る",
                code != 0 and "読取の報告の形が崩れている" in text and "形の崩れた報告 1" in st))
    lay = next((w / "out").glob("*.layout.json"))
    edit(lay, lambda d: d["pages"][0].update(bounds=[0, 0]))
    code, text = run("check_layout.py", w, "--no-png")
    out.append(("形: 図の配置（layout.json）が崩れていれば、配置の点検が FAIL", code != 0 and "形が崩れている" in text))
    side = TT.side_path(w / "out" / BM.BOOK)
    edit(side, lambda d: next(iter(d["sheets"].values())).update(header_rows="1"))
    try:
        TT.read_back(w / "out" / BM.BOOK, BM.SPEC)
        stopped = False
    except TT.TemplateError as e:
        stopped = "形が崩れている" in str(e)
    out.append(("形: 列の対応（columns.json）が崩れていれば、読み戻さない（人の値を別の欄に取り込まない）", stopped))
    return out


def review_gate_tests(tmp, good):
    """原典照合の門（check_flow --stage review）・採否表の下書き（review.py draft）・照合のあとの番号の引き継ぎ（apply_edits）"""
    import os
    import renumber as RN
    import terms as T
    out = []
    w = tmp / "gate_review"
    shutil.copytree(good, w)
    (w / "reports").mkdir(exist_ok=True)

    def report(F, verdict="支持", finding=True, rename=None):
        fl = json.loads((w / "flows" / f"{F}.flow.json").read_text(encoding="utf-8"))
        vs = [{"node": rename.get(n["id"], n["id"]) if rename else n["id"], "verdict": verdict, "evidence": f"「{n['label']}」とある"}
              for n in fl["nodes"] if n["type"] not in ("start", "end")]
        fs = [{"kind": "表記", "target": vs[0]["node"], "text": "箱の名前に、原典の語を使うとよい"}] if finding else []
        (w / "reports" / f"review_{F}.json").write_text(json.dumps({"verdicts": vs, "findings": fs}, ensure_ascii=False), encoding="utf-8")
    for F in ("F01", "F02", "F03"):
        report(F, finding=F == "F01")
    c0, o0 = run("review.py", w, "draft")
    c1, o1 = run("check_flow.py", w, "--stage", "review")

    def decide(F, decision="不採用", action="原典の語のままで、意味が通じる"):
        edit(w / "review" / f"{F}.review.json", lambda r: [x.update(decision=decision, action=action) for x in r["findings"]])
    decide("F01")
    c2, o2 = run("check_flow.py", w, "--stage", "review")
    cb, ob = run("build_reports.py", w)
    _c, st = run("status.py", w)
    out.append(("照合の門: 報告から採否表の下書きを機械が作る。採否が空のうちは FAIL、書けば FAIL 0。採否表に支持率が出て、進み具合は「済み」",
                c0 == 0 and c1 == 1 and "採否・対応の書かれていない指摘 1 件" in o1 and c2 == 0 and "0 FAIL" in o2 and "支持率 100%" in o2
                and cb == 0 and "支持率 100%" in (w / "out" / "原典照合レビュー_採否表.md").read_text(encoding="utf-8") and "8 原典照合     済み" in st))
    c3, _o = run("review.py", w, "draft")
    kept = json.loads((w / "review/F01.review.json").read_text(encoding="utf-8"))["findings"][0]["decision"] == "不採用"
    out.append(("照合の門: 下書きを作り直しても、採否は引き継ぐ", c3 == 0 and kept))

    def gate(name, change, expect):
        x = tmp / "gate_case"
        if x.exists():
            shutil.rmtree(x)
        shutil.copytree(w, x)
        change(x)
        code, text = run("check_flow.py", x, "--stage", "review")
        hit = [ln for ln in text.splitlines() if ln.startswith("FAIL") and expect in ln]
        out.append((f"照合の門: {name} → FAIL" + ("" if hit else f"（「{expect}」が出ない）\n" + text[-600:]), code == 1 and bool(hit)))
    gate("採否表の判定を、報告と違う文にした", lambda x: edit(x / "review/F02.review.json", lambda r: r["verdicts"][0].update(evidence="言い換えた根拠")), "照合の報告（reports/）と違う")
    gate("採用したのに、図を直していない", lambda x: edit(x / "review/F01.review.json", lambda r: r["findings"][0].update(decision="採用", action="箱の名前を直した")),
         "その指摘の報告を受けたときのまま")
    gate("採用したのに、図を直していない（照合した版の写しの無い、前の版の採否表）",
         lambda x: (edit(x / "review/F01.review.json", lambda r: (r["findings"][0].update(decision="採用", action="箱の名前を直した"), r.pop("reviewed"))),
                    os.utime(x / "flows/F01.flow.json", (1, 1))), "採否表より古い")
    gate("支持率が 95% 未満", lambda x: [edit(x / f"review/{F}.review.json", lambda r: [v.update(verdict="一部支持") for v in r["verdicts"][:3]]) for F in ("F01", "F02", "F03")], "95% 未満")
    gate("図の無い採否表", lambda x: (x / "review/F09.review.json").write_text(json.dumps({"flow_id": "F09", "verdicts": [], "findings": []}), encoding="utf-8"),
         "業務フローにもアクティビティ図にも無い")
    gate("採否表の無いフロー", lambda x: (x / "review/F02.review.json").unlink(), "採否表（review/F02.review.json）が無い")
    gate("照合の報告が無い（採否表を手で書いた）", lambda x: (x / "reports/review_F02.json").unlink(), "照合の報告（reports/）が無い")
    y = tmp / "gate_case_bad"
    shutil.copytree(w, y)
    edit(y / "review/F02.review.json", lambda r: r["verdicts"][0].update(verdict="まあまあ"))
    code, text = run("build_reports.py", y)
    out.append(("採否表の形が崩れていれば（採否の未記入のほか）、報告書を書かずに止まる", code != 0 and "形が違う" in text))
    # アクティビティ図の照合: 担当が突き合わせた箱（checked）だけを判定に数える。指摘のある箱は「不支持」
    import review as RV
    a = tmp / "act_review"
    (a / "reports").mkdir(parents=True)
    (a / "reports/code_review_A-01-1.json").write_text(json.dumps({"activities": ["A-01-1"], "checked": {"A-01-1": ["P-001", "P-002"]}, "findings": [
        {"activity": "A-01-1", "kind": "条件・出口が逆", "node": "P-002", "what": "ソースでは 0 件のときに抜ける"}]}, ensure_ascii=False), encoding="utf-8")
    fig = {"nodes": [{"id": "start", "type": "start"}] + [{"id": f"P-00{i}", "type": "process"} for i in (1, 2, 3)] + [{"id": "end", "type": "end"}]}
    got = RV.from_reports(a, "A-01-1", fig)
    out.append(("アクティビティ図の照合: 突き合わせた箱だけを判定に数え（突き合わせていない箱は支持にしない）、指摘のある箱は不支持",
                got is not None and {v["node"]: v["verdict"] for v in got[0]} == {"P-001": "支持", "P-002": "不支持"} and len(got[1]) == 1))
    edit(a / "reports/code_review_A-01-1.json", lambda d: d["findings"][0].update(node="P-009"))
    try:
        RV.from_reports(a, "A-01-1", fig)
        lost = False
    except ValueError as e:
        lost = "無い箱を指す" in str(e)
    out.append(("アクティビティ図の照合: 図に無い箱を指す指摘は、下書きを作らずに止まる（黙って消えない）", lost))
    z = tmp / "start_fix"
    shutil.copytree(good, z)
    no_start(z)
    run("renumber.py", z, "--only", "F02")
    out.append(("開始条件が空のフローに、renumber.py が機械の問い（m-start）を付ける",
                any(q.get("id") == "m-start" for q in json.loads((z / "flows/F02.flow.json").read_text(encoding="utf-8"))["questions"])))
    gate("採用の無いまま、判定したノードの番号が変わった（前の版の採否表）", lambda x: (edit(x / "review/F02.review.json", lambda r: r.pop("reviewed")),
                                                                  edit(x / "flows/F02.flow.json", lambda f: RN.apply_ids(f, {f["nodes"][1]["id"]: "P-099"}))), "判定したノード")
    gate("照合のあとに、箱の番号を振り直した", lambda x: edit(x / "flows/F02.flow.json", lambda f: RN.apply_ids(f, {f["nodes"][1]["id"]: "P-099"})), "P-099 照合のあとに足した")
    out += diff_review_tests(tmp, w)
    # 照合のあとは番号を保つ（apply_edits）。直しの一覧のあるフローは、当て直したものと同じか
    k = tmp / "keepnum"
    shutil.copytree(good, k)
    run("quotes.py", k, "locate", "--rewrite")
    run("quotes.py", k, "fill")
    shutil.copy2(SKILL / "evals/fixture_edits/F03.edits.json", k / "flows/F03.edits.json")
    run("apply_edits.py", k, "--flow", "F03")
    fl = json.loads((k / "flows/F03.flow.json").read_text(encoding="utf-8"))
    lab = {n["id"]: n["label"] for n in fl["nodes"]}
    RN.apply_ids(fl, {"P-001": "P-002", "P-002": "P-001"})
    (k / "flows/F03.flow.json").write_text(json.dumps(fl, ensure_ascii=False), encoding="utf-8")
    (k / "review").mkdir(exist_ok=True)
    (k / "review/F03.review.json").write_text(json.dumps({"flow_id": "F03", "verdicts": [], "findings": []}), encoding="utf-8")
    ck, ok_ = run("apply_edits.py", k, "--flow", "F03")
    after = {n["id"]: n["label"] for n in json.loads((k / "flows/F03.flow.json").read_text(encoding="utf-8"))["nodes"]}
    (k / "review/F03.review.json").unlink()
    run("apply_edits.py", k, "--flow", "F03")
    fresh = {n["id"]: n["label"] for n in json.loads((k / "flows/F03.flow.json").read_text(encoding="utf-8"))["nodes"]}
    passed = ck == 0 and "前の番号を引き継いだ" in ok_ and after["P-001"] == lab["P-002"] and after["P-002"] == lab["P-001"] and fresh == lab
    out.append(("照合のあとは、直しの一覧を当て直しても、前の番号を引き継ぐ（採否表が無ければ、振り直す）" + ("" if passed else f"\n{ok_[-400:]}\n{after}\n{fresh}\n{lab}"), passed))
    edit(k / "flows/F03.flow.json", lambda f: f["nodes"][1].update(label="手で直した名前を付ける"))
    cf, of = run("check_flow.py", k, "--stage", "flow")
    out.append(("直しの一覧のあるフローを手で直すと、当て直したものと違うので FAIL" + ("" if cf == 1 and "当て直したものと違う" in of else "\n" + of[-800:]),
                cf == 1 and "当て直したものと違う" in of))
    tm = T.Terms(tmp, {"M09": {"modules": []}})
    tm.rows = [{"project": "M09", "name": "SIME_BI", "word": "締め日", "basis": "comment"}]
    out.append(("用語の対応: 根拠がコメント・定義なのに場所が無いと、誤り", any("場所（source の module・line）が無い" in x for x in tm.problems())))
    x = tmp / "kinds"
    shutil.copytree(good, x)
    edit(x / "flows/F01.flow.json", lambda f: f["lanes"][0].update(kind="部署"))
    _c, text = run("check_flow.py", x, "--stage", "flow", "--all-lines")
    out.append(("決まった語の誤りは、形の点検の FAIL 1 行だけ（同じ誤りを二重に出さない）", sum(1 for ln in text.splitlines() if ln.startswith("FAIL") and "部署" in ln) == 1))
    edit(x / "flows/F01.flow.json", lambda f: f.update(name="別の名前"))
    run("renumber.py", x)
    out.append(("フローの名前は、renumber.py が業務分類から写す", json.loads((x / "flows/F01.flow.json").read_text(encoding="utf-8"))["name"]
                == next(c["name"] for c in json.loads((x / "taxonomy.json").read_text(encoding="utf-8"))["categories"] if c["id"] == "F01")))
    return out


def diff_review_tests(tmp, w):
    """差分照合（2026-10-15）: 採否表に照合した版の写しを残し、照合のあとに変わった・足した箱だけを照合し直させる"""
    import review as RV
    from validate import validate
    out = []
    flow = lambda x, F="F02": json.loads((x / "flows" / f"{F}.flow.json").read_text(encoding="utf-8"))
    boxes = [n["id"] for n in flow(w)["nodes"] if n["type"] not in ("start", "end")]
    P = boxes[1]

    def case(name):
        x = tmp / f"diff_{name}"
        shutil.copytree(w, x)
        return x

    def fails(x):
        code, text = run("check_flow.py", x, "--stage", "review")
        return code, [ln for ln in text.splitlines() if ln.startswith("FAIL")], text
    rv = json.loads((w / "review/F02.review.json").read_text(encoding="utf-8"))
    out.append(("差分照合: 採否表の下書きに、照合した版の写し（箱ごとの要約・報告の印・報告を受けたときの図の印）が残り、形が正しい",
                set(rv["reviewed"]["nodes"]) == set(boxes) and list(rv["reviewed"]["reports"]) == ["review_F02.json"]
                and rv["reviewed"]["figures"] == {"review_F02.json": RV.figure_mark(RV.figures(w)["F02"])} and validate(rv, "review") == []))
    x = case("label")
    edit(x / "flows/F02.flow.json", lambda f: node(f, P).update(label="照合のあとに直した名前を付ける"))
    code, hit, text = fails(x)
    out.append(("差分照合: 照合のあとに箱の名前を変えると、その箱だけを「照合したあとに変わった」と FAIL にし、差分照合のコマンドを出す",
                code == 1 and len(hit) == 1 and f"変わった・足した箱 1 個（{P} 名前が変わった）" in hit[0] and f"--only F02 --nodes {P}" in hit[0]))
    y = case("edge")
    e0 = flow(y)["edges"][1]
    edit(y / "flows/F02.flow.json", lambda f: f["edges"][1].update(label="いいえ" if f["edges"][1].get("label") != "いいえ" else "はい"))
    code, hit, text = fails(y)
    out.append(("差分照合: 線のラベルを変えると、線の両端の箱を「線が変わった」と挙げる（開始・終了は照合しないので挙げない）",
                code == 1 and all(f"{i} 線が変わった" in hit[0] for i in (e0["from"], e0["to"]) if i in boxes) and "start" not in hit[0]))
    reports = x / "reports"
    (reports / "review_F02_差分1.json").write_text(json.dumps({"verdicts": [{"node": P, "verdict": "支持", "evidence": "「直した名前」とある"}],
                                                             "findings": [{"kind": "表記", "target": P, "text": "差分照合の指摘"}]}, ensure_ascii=False), encoding="utf-8")
    code, hit, text = fails(x)
    out.append(("差分照合: 差分照合の報告を置いたまま、下書きを作り直していなければ FAIL（報告が採否表より新しい）",
                code == 1 and any("照合の報告（reports/）と違う" in ln or "増えた・変わった" in ln for ln in hit)))
    c1, o1 = run("review.py", x, "draft", "--only", "F02")
    edit(x / "review/F02.review.json", lambda r: [f.update(decision="不採用", action="原典の語のままでよい") for f in r["findings"] if not f["decision"]])
    code, hit, text = fails(x)
    rv = json.loads((x / "review/F02.review.json").read_text(encoding="utf-8"))
    out.append(("差分照合: 差分照合の報告から下書きを作り直すと、その箱の判定を置き換え、写しも取り直して、門が通る（指摘は足される）",
                c1 == 0 and code == 0 and rv["reviewed"]["nodes"][P]["label"] == "照合のあとに直した名前を付ける"
                and next(v for v in rv["verdicts"] if v["node"] == P)["evidence"] == "「直した名前」とある" and rv["findings"][-1]["text"] == "差分照合の指摘"
                and "review_F02_差分1.json" in rv["reviewed"]["reports"]))
    edit(x / "review/F02.review.json", lambda r: r["findings"][-1].update(decision="採用", action="名前を直す"))
    code, hit, _t = fails(x)
    out.append(("差分照合: 差分照合の指摘を採用して、図を直していなければ FAIL（その報告を受けたときの図のまま）",
                code == 1 and any(f"{rv['findings'][-1]['id']}）があるのに" in ln and "報告を受けたときのまま" in ln for ln in hit)))
    edit(x / "flows/F02.flow.json", lambda f: node(f, P).update(label="差分照合の指摘で直した名前を付ける"))
    code, hit, _t = fails(x)
    out.append(("差分照合: 採用した指摘を図に反映すると、直した箱が、また照合していない箱になる（反映の FAIL は消える）",
                code == 1 and len(hit) == 1 and f"{P} 名前が変わった" in hit[0]))
    edit(reports / "review_F02_差分1.json", lambda d: d["verdicts"].append({"node": "P-999", "verdict": "支持", "evidence": "「無い」"}))
    c2, o2 = run("review.py", x, "draft", "--only", "F02")
    out.append(("差分照合: 差分照合の報告が、図に無い箱を判定していれば、下書きを作らずに止まる", c2 == 1 and "無い箱を判定している（P-999" in o2))
    z = case("legacy")
    edit(z / "review/F02.review.json", lambda r: r.pop("reviewed"))
    code, _hit, text = fails(z)
    c3, o3 = run("review.py", z, "snapshot", "--only", "F02")
    c4, o4 = run("review.py", z, "snapshot", "--only", "F02", "--assume-current")
    code2, _h, text2 = fails(z)
    rv = json.loads((z / "review/F02.review.json").read_text(encoding="utf-8"))
    out.append(("差分照合: 写しの無い前の版の採否表は WARN（FAIL にしない）。下書きを作り直しても、写しを勝手に作らない。"
                "snapshot は --assume-current を付けたときだけ、いまの図を照合済みとして写し、日を残す",
                code == 0 and "WARN  F02: 照合した版の写し（reviewed）が無い" in text and c3 != 0 and "--assume-current" in o3 and c4 == 0
                and code2 == 0 and "照合した版の写し" not in text2 and rv["reviewed"]["figures"] == {} and rv["reviewed"].get("assumed")))
    v = case("legacy_draft")
    edit(v / "review/F02.review.json", lambda r: r.pop("reviewed"))
    run("review.py", v, "draft")
    out.append(("差分照合: 前の版の採否表から下書きを作り直しても、照合した版の写しを勝手に作らない",
                "reviewed" not in json.loads((v / "review/F02.review.json").read_text(encoding="utf-8"))))
    # 重ね方: 業務フローは、差分照合の判定が、前の判定を置き換える（厳しいほうではない）。アクティビティ図は、差分照合で見た箱の判定を置き換える
    a = tmp / "diff_merge"
    (a / "reports").mkdir(parents=True)
    fig = {"nodes": [{"id": "start", "type": "start"}] + [{"id": f"P-00{i}", "type": "process"} for i in (1, 2, 3)] + [{"id": "end", "type": "end"}]}
    (a / "reports/review_F05.json").write_text(json.dumps({"verdicts": [{"node": f"P-00{i}", "verdict": "一部支持" if i == 2 else "支持", "evidence": "前"} for i in (1, 2, 3)],
                                                        "findings": []}, ensure_ascii=False), encoding="utf-8")
    (a / "reports/review_F05_差分1.json").write_text(json.dumps({"verdicts": [{"node": "P-002", "verdict": "支持", "evidence": "後"}], "findings": []}, ensure_ascii=False), encoding="utf-8")
    (a / "reports/code_review_A-01-1.json").write_text(json.dumps({"activities": ["A-01-1"], "checked": {"A-01-1": ["P-001", "P-002"]}, "findings": [
        {"activity": "A-01-1", "kind": "条件・出口が逆", "node": "P-002", "what": "前の指摘"}]}, ensure_ascii=False), encoding="utf-8")
    (a / "reports/code_review_A-01-1_差分1.json").write_text(json.dumps({"activities": ["A-01-1"], "checked": {"A-01-1": ["P-002", "P-003"]}, "findings": []},
                                                                     ensure_ascii=False), encoding="utf-8")
    fv, _f = RV.from_reports(a, "F05", fig)
    av, af = RV.from_reports(a, "A-01-1", fig)
    out.append(("差分照合: 差分照合の判定は、前の判定を置き換える（業務フロー: 一部支持 → 支持。アクティビティ図: 指摘のあった箱も、差分照合で見て指摘が無ければ支持。前の指摘は残る）",
                {x["node"]: (x["verdict"], x["evidence"]) for x in fv} == {"P-001": ("支持", "前"), "P-002": ("支持", "後"), "P-003": ("支持", "前")}
                and {x["node"]: x["verdict"] for x in av} == {"P-001": "支持", "P-002": "支持", "P-003": "支持"} and len(af) == 1))
    return out


def output_check_tests(tmp, good):
    """書き出しのあとの点検: 図・配置・報告書を作り直して比べる／計算の結果・PNG へのリンク／照合シートの過不足／人が入れた値を守る／納品物の文字の走査"""
    import zipfile
    from openpyxl import load_workbook
    import build_matrix as BM
    import table_template as TT
    out = []

    def copy(name):
        x = tmp / name
        if x.exists():
            shutil.rmtree(x)
        shutil.copytree(good, x)
        return x

    def fails(text, expect):
        return [ln for ln in text.splitlines() if ln.startswith("FAIL") and expect in ln]

    def stage(name, x, st, expect):
        code, text = run("check_flow.py", x, "--stage", st, "--all-lines")
        hit = fails(text, expect)
        out.append((f"{name} → FAIL" + ("" if code == 1 and hit else f"（「{expect}」が出ない。終了コード {code}）\n" + text[-700:]), code == 1 and bool(hit)))

    # 図と配置: いまのデータから作り直したものと比べる（日時を除いて完全一致）
    x = copy("regen")
    p = x / "out" / "業務フロー_as-is.drawio"
    t = p.read_text(encoding="utf-8")
    m = re.search(r'<mxGeometry x="([0-9.]+)"', t)
    p.write_text(t[:m.start(1)] + str(float(m.group(1)) + 7) + t[m.end(1):], encoding="utf-8")
    stage("図を手で直した（箱を 7 ずらした。ラベルと線の両端は同じ）", x, "drawio", "いまのデータから作り直した図と違う")
    p.write_text(t, encoding="utf-8")
    edit(x / "out" / "業務フロー_as-is.layout.json", lambda d: d["pages"][1]["bounds"].__setitem__(2, d["pages"][1]["bounds"][2] + 1))
    stage("配置（layout.json）を手で直した", x, "drawio", "いまのデータから作り直した配置と違う")
    # 報告書: 作り直した文と同じか。確認事項一覧の xlsx は、機械が埋める欄の全部が、いまのデータから作った行と同じか
    x = copy("reports_hand")
    code0, text0 = run("check_flow.py", x, "--stage", "reports")
    out.append(("報告書: 書き出したままなら FAIL 0", code0 == 0 and "0 FAIL" in text0))
    with open(x / "out" / "出典対応表.md", "a", encoding="utf-8") as f:
        f.write("\n手で足した行\n")
    stage("出典対応表を手で直した", x, "reports", "out/出典対応表.md が、いまのデータから作り直したものと違う")
    x = copy("reports_book")
    book = x / "out" / "確認事項一覧.xlsx"
    wb = load_workbook(book)
    ws = wb.worksheets[0]
    col = [c.value for c in ws[1]].index("重要度") + 1
    ws.cell(2, col).value = "低" if ws.cell(2, col).value != "低" else "高"
    wb.save(book)
    stage("確認事項一覧.xlsx の重要度を手で変えた", x, "reports", "いまのデータと違うセル")
    x = copy("reports_book2")
    book = x / "out" / "確認事項一覧.xlsx"
    wb = load_workbook(book)
    ws = wb.worksheets[0]
    ws.cell(2, [c.value for c in ws[1]].index("仮置きした内容") + 1).value = "手で書き換えた"
    wb.save(book)
    stage("確認事項一覧.xlsx の仮置きした内容を手で変えた", x, "reports", "「仮置きした内容」")
    # 業務マトリクス一覧の計算の結果: 保存した値が、式を計算し直した値と同じか
    x = copy("cached")
    book = x / "out" / BM.BOOK
    with zipfile.ZipFile(book) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    name = next(n for n, b in parts.items() if n.startswith("xl/worksheets/") and re.search(rb"<f>SUM\([^<]*</f><v>[0-9.]+</v>", b))
    parts[name] = re.sub(rb"(<f>SUM\([^<]*</f><v>)([0-9.]+)(</v>)", lambda m: m.group(1) + b"98765" + m.group(3), parts[name], count=1)
    with zipfile.ZipFile(book, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    stage("保存した計算の結果（合計）を、式と違う値にした", x, "matrix", "式のセルに保存した計算の結果が、式を計算し直した値と違う")
    x = copy("uncached")
    edit_book(x, lambda wb: None)                                     # openpyxl で保存し直すと、計算の結果が消える
    _c, text = run("check_flow.py", x, "--stage", "matrix")
    out.append(("計算の結果が保存されていないブック → WARN（門で FAIL にする）", "計算の結果が保存されていない式のセル" in text))
    # フロー図の PNG へのリンク
    x = copy("links")
    (x / "out" / "png").mkdir(parents=True, exist_ok=True)
    (x / "out" / "png" / "業務フロー_as-is_p02_flow-F01.png").write_bytes(b"")
    stage("PNG があるのに、リンクを付ける前に書き出した", x, "matrix", "PNG があるのに、フロー図へのリンクの無い行")
    code, text = run("build_matrix.py", x)
    linked = code == 0 and any(c.hyperlink is not None and c.hyperlink.target == "png/業務フロー_as-is_p02_flow-F01.png"
                               for row in load_workbook(x / "out" / BM.BOOK)["業務マトリクス"].iter_rows() for c in row)
    out.append(("PNG を作ってから書き出すと、行の名前に図へのリンクが付く", linked))
    (x / "out" / "png" / "業務フロー_as-is_p02_flow-F01.png").unlink()
    stage("リンクの先の PNG が無い", x, "matrix", "フロー図へのリンクの先（png/業務フロー_as-is_p02_flow-F01.png）が無い")
    # 照合シート: 作れなかったフローがあれば終了コード 1。消えたフローのシートは消す
    x = copy("sheets")
    (x / "out" / "照合シート_F09.pdf").write_bytes(b"%PDF-1.4\n")
    r = subprocess.run([sys.executable, str(S / "make_compare.py"), "-w", str(x)], capture_output=True, text=True)
    out.append(("照合シート: PNG の無いフローは作れず、終了コード 1。いまの業務フローに無いシートは消す",
                r.returncode == 1 and "照合シートを作れなかったフロー 3 本" in r.stdout + r.stderr and not (x / "out" / "照合シート_F09.pdf").exists()))

    # 人が入れた値を守る（書き出す前に止まり、ブックを上書きしない）
    def guard(name, change, expect, keep=True):
        x = copy("guard")
        change(x)
        book = x / "out" / BM.BOOK
        before = book.read_bytes()
        code, text = run("build_matrix.py", x)
        out.append((f"人が入れた値を守る: {name} → 止まる" + ("" if code != 0 and expect in text else f"（「{expect}」が出ない）\n" + text[-500:]),
                    code != 0 and expect in text and (not keep or book.read_bytes() == before)))
    guard("ブックを表計算ソフトで開いたまま（~$ がある）", lambda x: (x / "out" / ("~$" + BM.BOOK)).write_bytes(b""), "を表計算ソフトで開いている")

    def stray(x):
        def fn(wb):
            c = next(c for row in wb["業務マトリクス"].iter_rows() for c in row if c.comment is not None and c.comment.text.startswith("処理時間 F01 "))
            c.comment, c.value = None, 7
        edit_book(x, fn)
        TT.side_path(x / "out" / BM.BOOK).unlink()
    guard("札（コメント）も列の対応も無い時間の値", stray, "だれの時間か決められない数がある")

    def typed(x):
        def fn(wb):
            c = next(c for row in wb["業務フローのノード一覧"].iter_rows(min_row=2) for c in row if isinstance(c.value, str) and re.match(r"^='?業務マトリクス'?!", c.value))
            c.value = 5
        edit_book(x, fn)
    guard("業務フローのノード一覧の時間の列に、数を打った", typed, "に、数が打ってある")

    def table(x):
        def fn(wb):
            tb, _bad = TT.time_table(wb, BM.SPEC["time"])
            wb[tb["title"]].cell(tb["rows"]["入力する"], tb["col"] + tb["cols"]["小"] - 1).value = "すぐ"
        edit_book(x, fn)
    guard("処理時間の基準の表に、数でない値を入れた", table, "処理時間の基準の表を読めない")
    # 納品物の文字の走査
    x = copy("scan")
    code0, text0 = run("check_flow.py", x, "--stage", "deliverables")
    with open(x / "out" / "出典対応表.md", "a", encoding="utf-8") as f:
        f.write("\n注: SYS-03 の画面で、H1.2 の業務。M01 p.3 を見る。 flows/F01.flow.json の start_condition。【下書き】\n")
    code, text = run("check_flow.py", x, "--stage", "deliverables", "--all-lines")
    want = ["内部の番号 SYS-03", "内部の番号 H1.2", "「M01」", "作業フォルダの場所 flows/F01.flow.json", "項目名 start_condition", "「【下書き】」"]
    lost = [w for w in want if not fails(text, w)]
    out.append(("納品物の文字の走査: 正しいデータは FAIL 0。内部の番号・資料 ID・作業フォルダの場所・項目名・【下書き】を入れると、それぞれ FAIL"
                + ("" if code0 == 0 and code == 1 and not lost else f"（出ない: {lost}）\n{text0[-400:]}\n{text[-900:]}"), code0 == 0 and code == 1 and not lost))
    return out


def answer_tests(tmp, good):
    """人が埋める欄（確認事項一覧の回答・様式で足した列）を、書き出し直しても同じ行に戻す（answers.json）。崩れた・動かしたブックでは止まる"""
    import datetime
    from openpyxl import load_workbook
    import build_matrix as BM
    out = []
    w = tmp / "answers"
    shutil.copytree(good, w)
    qb = w / "out" / "確認事項一覧.xlsx"

    def q_book(fn):
        wb = load_workbook(qb)
        ws = wb.worksheets[0]
        fn(ws, [c.value for c in ws[1]])
        wb.save(qb)

    def answered(path=qb):
        ws = load_workbook(path).worksheets[0]
        head = [c.value for c in ws[1]]
        return {r[0].value: (r[head.index("回答")].value, r[head.index("回答日")].value) for r in ws.iter_rows(min_row=2) if r[head.index("回答")].value}
    first = {}

    def put(ws, head):
        row = list(ws.iter_rows(min_row=2))[1]
        first["q"] = row[0].value
        row[head.index("回答")].value = "はい"
        row[head.index("回答日")].value = datetime.datetime(2026, 10, 20)
        row[head.index("回答日")].number_format = "yyyy/m/d"
    q_book(put)
    code, text = run("build_reports.py", w)
    out.append(("回答: 書き出し直しても、同じ番号の行に戻る（日付は日付のまま）",
                code == 0 and answered() == {first["q"]: ("はい", datetime.datetime(2026, 10, 20))} and "人が埋める欄の値を取り込んだ" in text))
    code, text = run("check_flow.py", w, "--stage", "reports")
    out.append(("回答: 回答済みなのに確認事項に残っていれば、点検が WARN（FAIL にはしない）", code == 0 and "回答済みなのに、まだ確認事項に残っている 1 件" in text))
    edit(w / "answers.json", lambda d: d["books"]["確認事項一覧.xlsx"]["確認事項"][first["q"]].update(seen="前の質問の文"))
    edit(w / "answers.json", lambda d: d["books"]["確認事項一覧.xlsx"]["確認事項"].update({"Q-999": {"cells": {"回答": {"value": "はい"}}, "seen": "無くなった問い"}}))
    _c, text = run("check_flow.py", w, "--stage", "reports")
    out.append(("回答: 確認事項が無くなった回答・取り込んだあとで質問の文が変わった回答は、点検が WARN",
                "確認事項に無くなった番号の回答 1 件（Q-999）" in text and f"質問の文が変わった確認事項 1 件（{first['q']}）" in text))
    back = tmp / "戻ってきた確認事項一覧.xlsx"                             # お客様から戻ってきたブックだけ（列の対応 columns.json は無い）
    shutil.copy2(qb, back)
    wb = load_workbook(back)
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    row = list(ws.iter_rows(min_row=2))[2]
    row[head.index("回答")].value, second = "いいえ", row[0].value
    wb.save(back)
    code, text = run("build_reports.py", w, "--keep", str(back))
    out.append(("回答: 戻ってきたブックだけ（列の対応が無い）でも、テンプレートの見出しで列を探して取り込む（--keep）",
                code == 0 and answered().get(second, ("",))[0] == "いいえ" and answered().get(first["q"], ("",))[0] == "はい"))
    q_book(lambda ws, head: [setattr(r[head.index(h)], "value", None) for r in ws.iter_rows(min_row=2) for h in ("回答", "回答日")])
    run("build_reports.py", w)
    left = json.loads((w / "answers.json").read_text(encoding="utf-8"))["books"]["確認事項一覧.xlsx"]["確認事項"]
    out.append(("回答: ブックで消した回答は answers.json からも消える（ブックに無い番号の回答は残す）", answered() == {} and list(left) == ["Q-999"]))
    (w / "out" / "~$確認事項一覧.xlsx").write_bytes(b"")
    code, text = run("build_reports.py", w)
    (w / "out" / "~$確認事項一覧.xlsx").unlink()
    out.append(("回答: 確認事項一覧を開いたままなら、書き出さずに止まる", code != 0 and "表計算ソフトで開いている" in text))
    q_book(lambda ws, head: ws.insert_cols(11))
    before = qb.read_bytes()
    code, text = run("build_reports.py", w)
    kept = qb.read_bytes() == before
    code2, _t = run("build_reports.py", w, "--fresh")
    out.append(("回答: 列を足したブックは、取り込まずに止まり、上書きしない（--fresh で作り直せる）",
                code != 0 and "列を足した・動かした" in text and kept and code2 == 0))
    (w / "answers.json").write_text('{"version": 1, "books": {"確認事項一覧.xlsx": 1}}', encoding="utf-8")
    code, text = run("build_reports.py", w)
    out.append(("回答: answers.json が崩れていれば、書き出さずに止まる", code != 0 and "answers.json の形が崩れている" in text))
    # 業務マトリクス一覧: 様式で足した列（業務マトリクスはフローの行、ノード一覧はフローとノード番号の行）
    m = tmp / "answers_matrix"
    shutil.copytree(good, m)
    tpl = tmp / "tpl_plus.xlsx"
    wb = load_workbook(SKILL / "assets" / BM.TEMPLATE)
    for name in (BM.M, BM.NODES):
        wb[name].cell(1, wb[name].max_column + 1).value = "備考（お客様）"
    wb.save(tpl)
    run("build_matrix.py", m, "--template", str(tpl))
    book = m / "out" / BM.BOOK
    wb = load_workbook(book)
    ws = wb[BM.M]
    ws.cell(2, ws.max_column).value = "月末に集中"
    ws = wb[BM.NODES]
    head = [c.value for c in ws[1]]
    fid, nid, last = head.index("フローID") + 1, head.index("ノード番号") + 1, ws.max_column
    target = (ws.cell(4, fid).value, ws.cell(4, nid).value)
    ws.cell(4, last).value = "省略することがある"
    for c in range(1, last + 1):                                     # ノード一覧の 3 行目と 4 行目を入れ替える（並べ替えた）
        ws.cell(3, c).value, ws.cell(4, c).value = ws.cell(4, c).value, ws.cell(3, c).value
    wb.save(book)
    code, text = run("build_matrix.py", m, "--template", str(tpl))
    wb = load_workbook(book)
    ws = wb[BM.NODES]
    noted = [(ws.cell(r, fid).value, ws.cell(r, nid).value) for r in range(2, ws.max_row + 1) if ws.cell(r, last).value == "省略することがある"]
    out.append(("様式で足した列: 業務マトリクスはフローの行、ノード一覧はフローとノード番号の行に戻る（並べ替えたブックからも）",
                code == 0 and wb[BM.M].cell(2, wb[BM.M].max_column).value == "月末に集中" and noted == [target]))
    ws = wb[BM.M]
    for c in range(1, ws.max_column + 1):                            # 業務マトリクスの 1 つ目と 2 つ目のフローの行を入れ替える
        ws.cell(2, c).value, ws.cell(4, c).value = ws.cell(4, c).value, ws.cell(2, c).value
    wb.save(book)
    before = book.read_bytes()
    code, text = run("build_matrix.py", m, "--template", str(tpl))
    out.append(("様式で足した列: 業務マトリクスの行を動かしたブックは、取り込まずに止まり、上書きしない",
                code != 0 and "書き出したときの行と違う" in text and book.read_bytes() == before))
    # 前の書き出しの様式を覚える。--template を付けずに書き出し直しても、既定の様式に黙って戻らない
    m2 = tmp / "answers_remember"
    shutil.copytree(good, m2)
    run("build_matrix.py", m2, "--template", str(tpl))
    code, text = run("build_matrix.py", m2)
    wb = load_workbook(m2 / "out" / BM.BOOK)
    out.append(("テンプレート: --template を付けずに書き出し直しても、前の書き出しの様式（足した列）を使う",
                code == 0 and "前の書き出しと同じ" in text and wb[BM.M].cell(1, wb[BM.M].max_column).value == "備考（お客様）"))
    tpl.rename(tmp / "tpl_moved.xlsx")
    code, text = run("build_matrix.py", m2)
    (tmp / "tpl_moved.xlsx").rename(tpl)
    out.append(("テンプレート: 前の書き出しの様式のファイルが無くなっていれば、既定の様式に戻さずに止まる", code != 0 and "前の書き出しで使ったテンプレート" in text))
    # 数として読めない時間の値は、取り込めないので止まる（書き出すと消える）
    m3 = tmp / "answers_unread"
    shutil.copytree(good, m3)
    book3 = m3 / "out" / BM.BOOK
    wb = load_workbook(book3)
    next(c for row in wb[BM.M].iter_rows() for c in row if c.comment is not None and BM.TAG_TIME.match(c.comment.text)).value = "約10分"
    wb.save(book3)
    before = book3.read_bytes()
    code, text = run("build_matrix.py", m3)
    out.append(("処理時間: 数として読めない値（「約10分」）があれば、取り込まずに止まり、上書きしない", code != 0 and "数として読めない値" in text and book3.read_bytes() == before))
    return out


def gate_tests(tmp, good):
    """納品の門（deliver.py）: 照合がまだなら止まる／書き出しが古ければやり直して通り、写す／書き出しが止まれば止まる／
    出力先に前の残りがあれば止まる（消さない）／資料 ID の WARN は門で FAIL。draw.io が無ければ、通る場合は飛ばす（合格に数えない）"""
    import os
    from check_layout import drawio_cli
    out = []
    w, dest = tmp / "gate", tmp / "gate_dest"
    shutil.copytree(good, w)
    code, text = run("deliver.py", w, "--to", str(dest))
    out.append(("納品の門: 原典照合がまだなら、止まって工程 8 を案内する（写さない）", code == 1 and "原典照合（工程 8）がまだ" in text and not dest.exists()))
    if drawio_cli() is None:
        return out + [("納品の門: draw.io が無いので、PNG の要る試験（門を通る・やり直す・写す）は飛ばした", None)]
    (w / "reports").mkdir(exist_ok=True)
    for f in sorted((w / "flows").glob("*.flow.json")):
        fl = json.loads(f.read_text(encoding="utf-8"))
        vs = [{"node": n["id"], "verdict": "支持", "evidence": f"「{n['label']}」とある"} for n in fl["nodes"] if n["type"] not in ("start", "end")]
        (w / "reports" / f"review_{fl['flow_id']}.json").write_text(json.dumps({"verdicts": vs, "findings": []}, ensure_ascii=False), encoding="utf-8")
    run("review.py", w, "draft")
    code, text = run("deliver.py", w, "--to", str(dest))
    files = {p.relative_to(dest).as_posix() for p in dest.rglob("*") if p.is_file()}
    want = {"納品物/業務フロー_as-is.drawio", "納品物/業務マトリクス一覧.xlsx", "納品物/確認事項一覧.xlsx", "納品物/確認事項一覧.md", "納品物/出典対応表.md",
            "納品物/出典対応表.csv", "納品物/階層と業務分類.md", "作業の記録/チェック_フロー.txt", "作業の記録/チェック_配置.txt", "作業の記録/原典照合レビュー_採否表.md",
            "作業の記録/納品の数.md", "作業の記録/納品の数.json", "作業の記録/照合シート_F01.pdf", "作業の記録/現状の処理時間.json"}
    _c, st = run("status.py", w)
    nums = json.loads((w / "out" / "納品の数.json").read_text(encoding="utf-8")) if (w / "out" / "納品の数.json").exists() else {}
    passed = code == 0 and want <= files and any(f.startswith("納品物/png/") for f in files) and "書き出しをやり直す" in text \
        and bool(re.search(r"10 納品の門\s+済み", st)) and nums.get("review", {}).get("flow", {}).get("support") == nums.get("nodes") and nums.get("fails") == 0 \
        and re.search(r"合計 [0-9.]+ 分・基準の表の版 T-[0-9a-f]{6}", (w / "out" / "納品の数.md").read_text(encoding="utf-8")) is not None
    out.append(("納品の門: 照合が済めば、古い書き出し（PNG が無い）をやり直して通り、納品物・作業の記録（現状の処理時間の写しも）・納品の数（合計と基準の表の版）を写す。"
                "進み具合は「10 納品の門 済み」"
                + ("" if passed else f"\n{text[-1500:]}\n{sorted(want - files)}"), passed))
    import fitz
    import fnmatch
    from common import RECORDS
    f01 = json.loads((w / "flows/F01.flow.json").read_text(encoding="utf-8"))
    P = next(n for n in f01["nodes"] if n["type"] == "process" and n.get("sources"))
    cd, od = run("make_compare.py", w, "--only", "F01", "--nodes", P["id"])
    diff = w / "out" / "差分照合シート_F01.pdf"
    doc = fitz.open(str(diff)) if diff.exists() else None
    first, pages = (doc[0].get_text().replace("\u2012", "-"), doc.page_count) if doc else ("", 0)   # 取り出すと、- が ‒ になる
    cb, _ob = run("make_compare.py", w, "--only", "F01,F02", "--nodes", P["id"])
    cc, oc = run("make_compare.py", w, "--only", "F01", "--nodes", "start")
    run("make_compare.py", w, "--only", "F01")
    out.append(("差分照合のシート: 1 ページ目に照合し直す箱、続けて、その箱の出典ページだけ。フローを 2 本・開始の箱を指定すると止まる。全体を作り直すと消える。門が写す記録には入らない",
                cd == 0 and pages == 1 + len({(s["doc_id"], s["page"]) for s in P["sources"]}) and "差分照合" in first and P["id"] in first
                and cb != 0 and cc != 0 and "開始・終了は照合しない" in oc and not diff.exists()
                and not any(fnmatch.fnmatch(diff.name, pat) for pat in RECORDS)))
    os.utime(w / "flows" / "F02.flow.json")
    (w / "out" / "~$業務マトリクス一覧.xlsx").write_bytes(b"")
    code, text = run("deliver.py", w, "--to", str(dest))
    (w / "out" / "~$業務マトリクス一覧.xlsx").unlink()
    out.append(("納品の門: 書き出しが止まれば（開いたままのブック）、門も止まる", code == 1 and "書き出しが止まった" in text and "表計算ソフトで開いている" in text))
    (dest / "納品物" / "前の図.drawio").write_text("x", encoding="utf-8")
    code, text = run("deliver.py", w, "--to", str(dest))
    out.append(("納品の門: 出力先に、いまの一式に無いファイルがあれば、写さずに止まる（消さない）",
                code != 0 and "前の納品物の残り" in text and (dest / "納品物" / "前の図.drawio").exists()))
    edit(w / "flows/F01.flow.json", lambda f: f["nodes"][1].update(detail="M01 の様式を使う"))
    code, text = run("deliver.py", w, "--to", str(tmp / "gate_dest2"))
    out.append(("納品の門: 文章の中で資料を ID で呼んでいる（点検では WARN）は、門では FAIL", code == 1 and "資料を ID で呼んでいる" in text and "門では FAIL" in text
                and not (tmp / "gate_dest2").exists()))
    return out


def mech_tests(tmp, good):
    """LLM に決めさせていた所を、機械が決まりで決める（6b）: 確認事項の重要度（決まり 5.4）・終了条件（3.6）・分岐のラベルと判断の名前（3.4）"""
    from apply_edits import ask_form, opposite
    from common import collect_questions, severity_of, when_done
    out = []
    pairs = [("会計伝票を提出する", "会計伝票を提出したとき"), ("支払いを行う", "支払いを行ったとき"), ("承認を得る", "承認を得たとき"), ("結果を確かめる。", "結果を確かめたとき"),
             ("請求書を送る", "請求書を送ったとき"), ("通知を書く", "通知を書いたとき"), ("次の工程に進む", "次の工程に進んだとき"), ("月次の締め", "「月次の締め」を終えたとき")]
    bad = [(a, when_done(a)) for a, b in pairs if when_done(a) != b]
    out.append(("終了条件（決まり 3.6）: 最後の作業の名前を「〜したとき」の形にする（動詞で終わらない名前は「「…」を終えたとき」）" + (f"\n{bad}" if bad else ""), not bad))
    pairs = [("不備がある", "不備がない"), ("差異があれば", "差異がなければ"), ("あり", "なし"), ("なし", "あり"), ("はい", "いいえ"), ("10万円以上", "10万円未満"),
             ("不一致", "一致"), ("承認する", "承認しない"), ("超える", None)]
    bad = [(a, opposite(a)) for a, b in pairs if opposite(a) != b]
    bad += [(a, ask_form(a)) for a, b in [("不備があるか？", "不備があるか"), ("エラー？", "エラーか"), ("承認するか", "承認するか")] if ask_form(a) != b]
    out.append(("分岐（決まり 3.4）: 書かれていない側のラベルは、書かれた側の反対の語（表に無ければ「その他」）。判断の名前の「？」は「か」" + (f"\n{bad}" if bad else ""), not bad))
    cases = [({"effect": "担当", "nodes": ["P-001", "P-002"]}, "mid"), ({"effect": "システム・帳票", "node": "P-001"}, "low"),
             ({"effect": "フローの範囲", "node": "end"}, "low"), ({"effect": "フローの範囲", "nodes": ["start", "end2"]}, "low"),
             ({"effect": "フローの範囲", "edge": ["P-003", "end"]}, "high"), ({"effect": "フローの範囲"}, "high"), ({"effect": "箱", "node": "P-001"}, "mid"),
             ({"effect": "順序", "edge": ["P-001", "P-002"]}, "mid"), ({"effect": "順序", "nodes": ["P-001", "P-002"]}, "high"), ({"effect": "分岐"}, "mid")]
    bad = [(q, want, severity_of(q)) for q, want in cases if severity_of(q) != want]
    out.append(("重要度（決まり 5.4）: 担当 → 中／システム・帳票 → 低／フローの範囲 → 高（開始・終了だけなら低）／箱・順序・分岐 → 指す箱 2 つ以上で高、ほかは中"
                + (f"\n{bad}" if bad else ""), not bad))
    w = tmp / "mech"
    shutil.copytree(good, w)
    e = json.loads((SKILL / "evals/fixture_edits/F03.edits.json").read_text(encoding="utf-8"))
    e["questions"][0]["severity"] = "high"
    (w / "flows/F03.edits.json").write_text(json.dumps(e, ensure_ascii=False), encoding="utf-8")
    code, text = run("apply_edits.py", w, "--flow", "F03")
    qs = {q["id"]: q for q in json.loads((w / "flows/F03.flow.json").read_text(encoding="utf-8"))["questions"]}
    out.append(("重要度: 直しの一覧に書いた重要度は使わず、決まりで付け直して知らせる（担当の問い → 中）。機械の問いも同じ決まり（箇条書きの順 → 高・開始条件 → 低）",
                code == 0 and qs["q1"]["severity"] == "mid" and "q1: 重要度は、決まり 5.4 で中にした" in text
                and all(q["severity"] == severity_of(q) for q in qs.values()) and qs.get("m-start", {}).get("severity", "low") == "low"))
    dead = next(q for q in json.loads((w / "flows/F03.flow.json").read_text(encoding="utf-8"))["questions"] if q["id"].startswith("m-dead"))
    edit(w / "flows/F03.flow.json", lambda f: f["questions"].append(dict(dead, id="q9", question="差し戻したあとは、どうしますか。")))
    code, text = run("check_flow.py", w, "--stage", "flow")
    out.append(("点検: 直しの一覧の問いが、機械の問いと同じ所・同じ effect を指せば WARN（決まり 5.6）", f"q9 は、機械の確認事項 {dead['id']} と同じ所を指し" in text))
    edit(w / "flows/F03.flow.json", lambda f: f.update(questions=[q for q in f["questions"] if q["id"] != "q9"]))
    edit(w / "flows/F03.flow.json", lambda f: f["questions"][0].update(severity="low"))
    code, text = run("check_flow.py", w, "--stage", "flow")
    c2, t2 = run("renumber.py", w, "--only", "F03")
    fixed = json.loads((w / "flows/F03.flow.json").read_text(encoding="utf-8"))["questions"][0]["severity"] == "mid"
    out.append(("点検: flow.json の重要度が決まりと違えば FAIL。renumber.py が付け直す（JSON をまるごと書いたフローも）",
                code == 1 and "決まり 5.4 の mid と違う" in text and c2 == 0 and fixed and "重要度を、決まり 5.4 で中にした" in t2))
    from common import settle_flow
    fl = {"end_condition": "", "nodes": [{"id": "start", "type": "start"}, {"id": "D-001", "type": "decision", "label": "不備がある？"},
                                         {"id": "P-001", "type": "process", "label": "伝票を提出する"}, {"id": "end", "type": "end"}],
          "edges": [{"from": "start", "to": "D-001"}, {"from": "D-001", "to": "P-001", "label": "不備がある"}, {"from": "D-001", "to": "P-001", "label": "その他", "inferred": True},
                    {"from": "P-001", "to": "end"}], "questions": [{"id": "q1", "effect": "箱", "severity": "high", "node": "P-001"}]}
    notes = settle_flow(fl)
    out.append(("決まりの当て方（settle_flow。apply_edits・renumber が呼ぶ）: 判断の「？」→「か」・「その他」→ 反対の語・空の終了条件・空の開始条件を聞く問い・重要度。2 回目は何も変えない",
                fl["nodes"][1]["label"] == "不備があるか" and fl["edges"][2]["label"] == "不備がない" and fl["end_condition"] == "伝票を提出したとき"
                and fl["questions"][0]["severity"] == "mid" and any(q["id"] == "m-start" and q["severity"] == "low" for q in fl["questions"])
                and len(notes) == 5 and settle_flow(fl) == []))
    from draft_flow import Draft
    d2, d3 = Draft(w, "F02"), Draft(w, "F03")
    out.append(("下書き（決まり 3.21・3.8）: 「部長の承認を得る」だけの手順は部長のレーン。「自動で連携される」手順は、そのシステムのレーン（システム処理の箱になる）",
                d2.nodes["M04.2.s2"]["lane"] == "部長" and d3.nodes["M02.1.s4"]["lane"] == "会計システム" and "承認の相手のレーン" in d2.show()))
    from apply_edits import make
    base = {"start_condition": "", "end_condition": "", "ops": []}
    f1, e1, _d = make(w, "F02", dict(base, ops=[{"op": "exclude", "item": "M04.2.s2", "flow": "F01"}]))
    f2, e2, _d = make(w, "F02", dict(base, ops=[{"op": "exclude", "item": "M04.2.s2", "same_as": "M01.5.s1"}]))
    why = lambda fl: [x["reason"] for x in (fl or {}).get("excluded_items", []) if x["item"] == "s2" and x["page"] == 2]
    out.append(("対象外の理由（決まった形）: flow → 「〈分類名〉（F01）のフローで描く」、same_as → 「「…」と同じ作業（そちらをフローに入れた）」"
                + ("" if not (e1 or e2) else f"\n{e1}{e2}"),
                not e1 and not e2 and why(f1) == ["「経費申請」（F01）のフローで描く"] and len(why(f2)) == 1 and why(f2)[0].endswith("と同じ作業（そちらをフローに入れた）")))
    from common import labels, manifest
    f3, e3, _d = make(w, "F02", dict(base, ops=[{"op": "merge", "into": "M01.5.s1", "items": ["M04.2.s2"], "values": {"M01.5.s1": "3万円以上", "M04.2.s2": "5万円以上"}}]))
    want = f"3万円以上。{labels(manifest(w), w).title('M04')}では、5万円以上"
    out.append(("値だけが違う作業をまとめる（決まり 3.12）: 補足を「まとめ先の値。〈資料の題名〉では、〈値〉」の形で書く" + (f"\n{e3}" if e3 else ""),
                not e3 and any(n.get("detail") == want for n in (f3 or {}).get("nodes", []))))
    qa = [q for q in (f3 or {}).get("questions", []) if q["id"].startswith("m-actor")]
    out.append(("担当の違う作業をまとめたら、黙って捨てずに確認事項（資料間の矛盾・担当）を機械が作る（決まり 3.12）",
                len(qa) == 1 and qa[0]["category"] == "資料間の矛盾" and qa[0]["effect"] == "担当" and "上長・部長" in qa[0]["question"]))
    from common import added_category_questions
    tax = {"basis": "mixed", "categories": [{"id": "F01", "name": "経費申請", "basis": "toc"}, {"id": "F09", "name": "備品の貸出", "basis": "llm",
                                                                                              "name_ref": {"doc_id": "M05", "page": 1}},
                                            {"id": "F10", "name": "出張の手配", "basis": "llm", "pages": [{"doc_id": "M06", "pages": [2, 3]}]}],
           "questions": [{"id": "q1", "category": "分類", "flow": "F10", "question": "x"}]}
    got = added_category_questions(tax)
    out.append(("目次に無く足した分類ごとの確認事項は、機械が決まった形で作る（LLM がその分類を指す問いを書いていれば作らない）（決まり 2.21）",
                [(x, q["flow"], q["sources"]) for x, q in got] == [("m-added-F09", "F09", [{"doc_id": "M05", "page": 1}])] and "備品の貸出" in got[0][1]["question"]
                and added_category_questions(dict(tax, basis="llm")) == []))
    w2 = tmp / "mech_links"                                          # 紐付け（決まり 2.20）: LLM は候補の番号だけを渡し、根拠の文とページは機械が写す
    shutil.copytree(good, w2)
    edit(w2 / "structure/hierarchy.json", lambda h: h.update(doc_relations=[r for r in h["doc_relations"] if r["type"] != "appendix_of"]))
    c1, t1 = run("pages.py", w2, "links", "--accept", "R1,S1")
    h = json.loads((w2 / "structure/hierarchy.json").read_text(encoding="utf-8"))
    cand = json.loads((w2 / "structure/link_candidates.json").read_text(encoding="utf-8"))
    r1 = next(r for r in cand["relations"] if r["id"] == "R1")
    c2, t2 = run("pages.py", w2, "links", "--accept", "R1,S1")
    c3, t3 = run("pages.py", w2, "links", "--accept", "R9")
    out.append(("紐付け: 候補の番号で採ると、根拠の文とページを候補のまま hierarchy.json に写す（2 回目は足さない。無い番号は止まる）",
                c1 == 0 and {k: r1[k] for k in ("from", "to", "type", "evidence", "sources")} in h["doc_relations"]
                and any((s["from"], s["to"]) == ("H1.1", "H1.3") for s in h["sequence"]) and "写した 0 件" in t2 and c3 != 0 and "R9" in t3))
    edit(w2 / "taxonomy.json", lambda t: [p.update(pages=[x for x in p["pages"] if x != 5]) for c in t["categories"] for p in c.get("pages") or [] if p["doc_id"] == "M01"])
    edit(w2 / "structure/hierarchy.json", lambda h: [r.update(pages=[x for x in r["pages"] if x != 5]) for r in h["tree"][0]["children"][1]["refs"] if r["doc_id"] == "M01"])
    c4, t4 = run("pages.py", w2, "orphans", "--accept")
    refs = json.loads((w2 / "structure/hierarchy.json").read_text(encoding="utf-8"))["tree"][0]["children"][1]["refs"]
    out.append(("紐付け: どの分類にも入っていないページの見出しが、目次の中で一意に一致すれば、機械が業務の木に結び付ける（「2. 承認」→ 承認）" + ("" if c4 == 0 else t4[-300:]),
                c4 == 0 and any(r["doc_id"] == "M01" and 5 in r["pages"] for r in refs)
                and any(5 in p["pages"] for c in json.loads((w2 / "taxonomy.json").read_text(encoding="utf-8"))["categories"] if c["id"] == "F02"
                        for p in c["pages"] if p["doc_id"] == "M01")))
    shown = d3.show()
    out.append(("下書き: 量の案（ボタンを押す・実行する → 小）を出す（付けるかどうかは LLM）", "量の案: 小" in shown and "量の候補: 特大" not in shown))
    dr = lambda kind, a_, b_: subprocess.run([sys.executable, str(S / "diff_runs.py"), kind, str(a_), str(b_), "--min", "1.0"], capture_output=True, text=True)
    same, moved, fl = dr("structure", good, good), dr("structure", good, w2), dr("flows", good, good)
    out.append(("揺れの測り方（diff_runs.py）: 階層・分類・用語を比べる（同じものは 100%、紐付けを足した写しとは差が出る）。フローは、開始・終了条件・重要度・動作の種類も見る",
                same.returncode == 0 and moved.returncode == 1 and "業務どうしの「続く」" in moved.stdout and "ページの割当" in moved.stdout
                and fl.returncode == 0 and all(x in fl.stdout for x in ("両方にある確認事項の重要度", "終了条件", "対応したノードの動作の種類"))))
    edit(w / "structure/hierarchy.json", lambda h: h["questions"][0].update(severity="low", effect="フローの範囲"))
    got = [q for q in collect_questions(w, write=False) if q["origin"] == "hierarchy"]
    out.append(("重要度: 階層・用語・分類・台帳（LLM が書くファイル）の問いも、書いた重要度は使わず決まりで決める", bool(got) and got[0]["severity"] == "high"))
    # レーンの順（決まり 4.2）: 機械が並べる。直しの一覧の lane_order で指定できる。JSON をまるごと書いたフローは renumber.py --lanes のときだけ
    from common import order_lanes
    pairs = [("A", "C"), ("A", "C"), ("C", "B"), ("A", "B")]
    out.append(("レーンの順: 開始の担当を左端、システムを右端に置き、残りは線がまたぐレーンの数が最小の順。同じなら出てくる順",
                order_lanes(["A", "B", "C", "S"], pairs, keep=["A"], right=["S"]) == ["A", "C", "B", "S"] and order_lanes(["A", "B", "C"], [], keep=["A"]) == ["A", "B", "C"]
                and order_lanes(["A", "B", "C"], pairs, keep=["B", "A"]) == ["B", "A", "C"]))
    ln = tmp / "lanes"
    shutil.copytree(good, ln)
    run("quotes.py", ln, "locate", "--rewrite")
    run("quotes.py", ln, "fill")
    e = json.loads((SKILL / "evals/fixture_edits/F03.edits.json").read_text(encoding="utf-8"))
    (ln / "flows/F03.edits.json").write_text(json.dumps(e, ensure_ascii=False), encoding="utf-8")
    run("apply_edits.py", ln, "--flow", "F03")
    names = [x["name"] for x in json.loads((ln / "flows/F03.flow.json").read_text(encoding="utf-8"))["lanes"]]
    pick = next(x for x in reversed(names) if x != names[0] and x != "会計システム")
    (ln / "flows/F03.edits.json").write_text(json.dumps(dict(e, lane_order=[pick]), ensure_ascii=False), encoding="utf-8")
    c1, _o = run("apply_edits.py", ln, "--flow", "F03")
    after = [x["name"] for x in json.loads((ln / "flows/F03.flow.json").read_text(encoding="utf-8"))["lanes"]]
    (ln / "flows/F03.edits.json").write_text(json.dumps(dict(e, lane_order=["財務部"]), ensure_ascii=False), encoding="utf-8")
    c2, o2 = run("apply_edits.py", ln, "--flow", "F03")
    out.append(("レーンの順: 直しの一覧の lane_order のとおりに左から並ぶ（システムは右端のまま）。フローに無いレーンを書けば、何も書かずに止まる",
                c1 == 0 and after[0] == pick and after[-1] == "会計システム" and sorted(after) == sorted(names) and c2 == 1 and "lane_order の「財務部」" in o2))
    before = [x["name"] for x in json.loads((ln / "flows/F02.flow.json").read_text(encoding="utf-8"))["lanes"]]
    run("renumber.py", ln, "--only", "F02")
    kept = [x["name"] for x in json.loads((ln / "flows/F02.flow.json").read_text(encoding="utf-8"))["lanes"]]
    c3, o3 = run("renumber.py", ln, "--only", "F02", "--lanes")
    fl2 = json.loads((ln / "flows/F02.flow.json").read_text(encoding="utf-8"))
    lane_name = {x["id"]: x["name"] for x in fl2["lanes"]}
    start = lane_name[next(n["lane"] for n in fl2["nodes"] if n["id"] == "start")]
    out.append(("レーンの順: JSON をまるごと書いたフローは、renumber.py --lanes のときだけ並べ直す（開始の担当が左端）",
                kept == before and c3 == 0 and "レーンを並べ直した" in o3 and [x["name"] for x in fl2["lanes"]][0] == start))
    return out


def review_fix_tests(tmp):
    """原典との照合で見つかった所: 枠組みが先に呼ぶ手続き・前置きつきの URL・条件の無い繰り返し・線の重なり・箱の名前の記号・ファイルが2つ以上のまとまりの名前"""
    import copy

    import code_lang
    import code_struct as cs
    out = []
    L = TREE_HOOK.splitlines()
    at = lambda t: next(k for k, ln in enumerate(L, 1) if t in ln)
    M = "ItemController.java"
    T = lambda name, way, line: [{"kind": "テーブル", "name": name, "way": way, "line": line}]
    proc = lambda name, word, io: {"module": M, "name": name, "lines": [at(f" {name}("), at(f" {name}(") + 3], "body": [
        {"kind": "seq", "lines": [at(f" {name}(") + 1] * 2, "io": io}, {"kind": "exit", "lines": [at(f" {name}(") + 2] * 2, "what": "proc"}]}
    fs = at("findShop(")
    tree = {"project": "M01", "procedures": [
        {"module": M, "name": "findShop", "lines": [fs, fs + 2], "body": [{"kind": "seq", "lines": [fs + 1, fs + 1], "io": T("shops", "in", fs + 1), "leave": "proc"}]},
        proc("create", "new", T("items", "out", at(" create(") + 1)), proc("update", "edit", T("items", "out", at(" update(") + 1))],
        "entries": [{"procedure": f"{M}.create", "trigger": "entry", "label": "画面から「/shop/items/new」に送信", "line": at("/items/new"), "also": 1, "before": [f"{M}.findShop"]},
                    {"procedure": f"{M}.update", "trigger": "entry", "label": "画面から「/shop/items/edit」に送信", "line": at("/items/edit"), "also": 1, "before": [f"{M}.findShop"]}]}
    mods = [{"name": M, "kind": "file", "lines": L}]
    st = cs.build_struct("M01", "java", mods, tree=tree)
    procs = {p["name"]: p for p in st["procedures"]}
    first = procs["create"]["body"][0]
    out.append(("木の決まり: 枠組みが起動点の前に呼ぶ手続き（before）は、図の先頭の呼び出しになる（実行文は持たない）。2つの起動点から呼ばれるので、別の図",
                code_lang.tree_problems(tree, mods, "java") == [] and cs.tree_errors(st) == [] and st["unreached"] == []
                and (first["kind"], first["to"], first.get("exec") or [], first.get("framework")) == ("call", f"{M}.findShop", [], True)
                and [u["procedure"].split(".")[-1] for u in st["units"]] == ["create", "update", "findShop"]))

    def case(name, change, expect):
        t = copy.deepcopy(tree)
        change(t)
        errs = code_lang.tree_problems(t, mods, "java")
        hit = next((e for e in errs if expect in e), None)
        out.append((f"木の誤り: {name}: {(hit or '')[:100]}" if hit else f"木の誤り: {name}: 「{expect}」が出ない\n{errs[:4]}", bool(hit)))
    case("前置きつきの URL を、also なしで書く", lambda t: t["entries"][0].pop("also"), "also に書く")
    case("also の行と根拠の行をつないでも、名前にならない", lambda t: t["entries"][0].update(label="画面から「/store/items/new」に送信"), "この順につないだものでもない")
    case("before の手続きが、木に無い", lambda t: t["entries"][0].update(before=[f"{M}.nothing"]), "before")
    # 条件の無い繰り返し・線の重なり・箱の名前の記号・ファイルが2つ以上のまとまりの名前を、通しで
    src, w = tmp / "loop_in", tmp / "loop_work"
    (src / "src").mkdir(parents=True)
    (src / "src" / "探す.bas").write_text('Attribute VB_Name = "探す"\n' + VBA_LOOP, encoding="utf-8")
    (src / "src" / "補助.bas").write_text('Attribute VB_Name = "補助"\nPublic Const 上限 = 100\n', encoding="utf-8")
    subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(w)], capture_output=True, text=True)
    run("render_pages.py", w)
    run("code_struct.py", w)
    run("systems.py", w, "draft")
    led = json.loads((w / "structure/systems.json").read_text(encoding="utf-8")) if (w / "structure/systems.json").exists() else {"systems": []}
    out.append(("台帳の下書き: ファイルが2つ以上のまとまりの名前は、先頭のファイルの名前だけにしない（ありふれた名前のフォルダなら「… ほか 1 ファイル」）",
                [x["name"] for x in led["systems"]] == ["探す ほか 1 ファイル"]))
    edit(w / "structure/systems.json", lambda j: [x.update(kind="script") for x in j["systems"]])
    _c, listing = run("activity.py", w, "list")
    aid = next((x.split()[0] for x in listing.splitlines() if "探す.探す" in x), "A-01-1")
    (w / "activities").mkdir(exist_ok=True)
    (w / f"activities/{aid}.edits.json").write_text(json.dumps({"activity": aid, "name": "空の行まで印を付ける", "summary": "1列目が空になるまで、2列目に 1 を入れる。", "questions": [], "ops": [
        {"op": "exclude", "item": "探す.4", "reason": "変数の準備"},
        {"op": "name", "item": "探す.5", "label": "1列目が空か", "branches": ["空", "入っている"]},
        {"op": "name", "item": "探す.8", "label": "2列目に 1 を入れる（A & B < C）"}]}, ensure_ascii=False), encoding="utf-8")
    code1, text1 = run("activity.py", w, "apply", "--id", aid)
    act = json.loads((w / f"activities/{aid}.activity.json").read_text(encoding="utf-8")) if (w / f"activities/{aid}.activity.json").exists() else {"nodes": [], "edges": [], "branches": []}
    ids = {n["label"]: n["id"] for n in act["nodes"]}
    dec, box = ids.get("1列目が空か"), ids.get("2列目に 1 を入れる（A & B < C）")
    msg = next((n["id"] for n in act["nodes"] if "終わり" in n["label"] and n["type"] != "end"), None)
    out.append(("アクティビティ図: 条件の無い繰り返し（Do … Loop）は、判断にしない。中の最後から頭へ戻り、終わるのは中の「抜ける」だけ",
                code1 == 0 and [n["type"] for n in act["nodes"]].count("decision") == 1 and {"from": box, "to": dec} in [{"from": e["from"], "to": e["to"]} for e in act["edges"]]
                and [e["to"] for e in act["edges"] if e["from"] == dec and e.get("label") == "空"] == [msg]
                and [b["state"] for b in act["branches"] if b["key"] == "探す.3"] == ["endless"]))
    run("build_activity.py", w)
    code2, text2 = run("check_flow.py", w)
    code3, text3 = run("check_layout.py", w, "--no-png")
    out.append(("通し（条件の無い繰り返し）: 箱の名前に & や < があっても、図とデータの照合が合う。点検は FAIL 0（配置も）",
                code2 == 0 and "0 FAIL" in text2 and code3 == 0 and "0 FAIL" in text3))
    lay = next(iter((w / "out").glob("*.layout.json")), None)
    if lay is not None:                                              # 線を、わざと重ねる: 出どころも行き先も違う線が、同じ所を走る
        def overlap(j):
            pg = next(p for p in j["pages"] if len(p["edges"]) >= 3)
            a, b = next((x, y) for x in pg["edges"] for y in pg["edges"] if x["from"] != y["from"] and x["to"] != y["to"])
            b["points"] = [[a["points"][0][0], a["points"][0][1] + 2], [a["points"][0][0], a["points"][0][1] + 40]]
            a["points"] = [a["points"][0], [a["points"][0][0], a["points"][0][1] + 60]]
        edit(lay, overlap)
        code4, text4 = run("check_layout.py", w, "--no-png")
        out.append(("配置の点検: 出どころも行き先も違う線が、同じ所を重なって走ると FAIL", code4 == 1 and "同じ所を重なって走る" in text4))
    return out


def web_tree():
    """架空の画面の受け口（Java）と、そのブロックの木。1行が2つの働きを兼ねる書き方・条件の中の呼び出し・プログラムごと終える文を含む"""
    L = TREE_WEB.splitlines()
    at = lambda t: next(k for k, ln in enumerate(L, 1) if t in ln)
    M = "OrderController.java"
    one = lambda kind, where, **kw: dict({"kind": kind, "lines": [at(where), at(where)]}, **kw)
    cr, dup, fin, ab, sp = at("public String create"), at("private boolean isDuplicate"), at("private String finish"), at("private void abort"), at("private void stop")
    tree = {"project": "M01", "procedures": [
        {"module": M, "name": "create", "lines": [cr, dup - 2], "body": [
            {"kind": "if", "lines": [cr + 1, cr + 3], "branches": [{"label": "then", "body": [one("call", 'abort("empty")', to=f"{M}.abort")]}]},
            {"kind": "if", "lines": [cr + 4, cr + 7], "calls": [f"{M}.isDuplicate"], "branches": [{"label": "then", "body": [
                one("interact", "rejectValue", how="msgbox", text="同じ番号の注文があります"), one("exit", 'return "orders/form"', what="proc")]}]},
            one("seq", "orders.save", io=[{"kind": "テーブル", "name": "orders", "way": "out", "line": at("orders.save")}]),
            one("seq", "int n = 0", tech="変数の準備"), one("seq", "log.info", tech="ログ"),
            one("interact", "addFlashAttribute", how="msgbox", text="注文を登録しました"),
            one("call", "return finish(order)", to=f"{M}.finish", leave="proc")]},
        {"module": M, "name": "isDuplicate", "lines": [dup, dup + 2], "body": [
            one("seq", "findByCode", io=[{"kind": "テーブル", "name": "orders", "way": "in", "line": at("findByCode")}], leave="proc")]},
        {"module": M, "name": "finish", "lines": [fin, sp - 2], "body": [
            {"kind": "seq", "lines": [fin + 1, at("mailer.send")], "io": [{"kind": "メール", "name": "mailer", "way": "out", "line": at("mailer.send")}]},
            {"kind": "if", "lines": [at("isUrgent"), at("isUrgent") + 2], "branches": [{"label": "then", "body": [one("call", 'abort("urgent")', to=f"{M}.abort")]}]},
            {"kind": "if", "lines": [at("isClosed"), at("isClosed") + 2], "branches": [{"label": "then", "body": [one("call", "stop();", to=f"{M}.stop")]}]},
            one("exit", 'return "redirect:/orders"', what="proc")]},
        {"module": M, "name": "stop", "lines": [sp, sp + 2], "body": [one("exit", "System.exit(1)", what="program")]},
        {"module": M, "name": "abort", "lines": [ab, ab + 3], "body": [one("seq", "log.error", tech="ログ"), one("exit", "System.exit(9)", what="program")]}],
        "entries": [{"procedure": f"{M}.create", "trigger": "entry", "label": "画面から「/orders/new」に送信（POST）", "line": at("@PostMapping")}]}
    return M, L, at, tree


def tree_rule_tests(tmp):
    """木の決まりのうち、2人が読んで割れた所: 1行が2つの働きを兼ねる・条件の中の呼び出し・プログラムごと終える・まとまりの区切り・やり取りの文字"""
    import copy

    import code_lang
    import code_struct as cs
    out = []
    M, L, at, tree = web_tree()
    mods = [{"name": M, "kind": "file", "lines": L}]
    st = cs.build_struct("M01", "java", mods, tree=tree)
    procs = {p["name"]: p for p in st["procedures"]}
    kinds = lambda name: [(n["kind"], n["lines"], n.get("exec")) for n in procs[name]["body"]]
    out.append(("木の決まり: 条件の中の呼び出し（calls）で、呼び先がたどれる。呼ぶ側の図には入れない（別の図）",
                code_lang.tree_problems(tree, mods, "java") == [] and cs.tree_errors(st) == [] and st["unreached"] == []
                and {"to": f"{M}.isDuplicate", "line": at("if (isDuplicate"), "head": True} in procs["create"]["calls"]
                and [u["procedure"].split(".")[-1] for u in st["units"]] == ["create", "isDuplicate"] and sorted(x.split(".")[-1] for x in st["inline"]) == ["finish", "stop"]))
    ret = at("return finish(order)")
    out.append(("木の決まり: 1行が「する」と「戻る」を兼ねる（leave）: 行は先のノードが持ち、そのあとに「抜ける」が付く。読むものは残る",
                kinds("create")[-2:] == [("call", [ret, ret], [ret]), ("exit", [ret, ret], None)]
                and [(n["kind"], [i["name"] for i in n.get("io", [])]) for n in procs["isDuplicate"]["body"]] == [("seq", ["orders"]), ("exit", [])]))
    out.append(("木の決まり: やり取りの文字は、書いた text（項目の名前やキーでなく、見せる文言）",
                [n["interact"]["text"] for n in code_lang.walk(procs["create"]["body"]) if n["kind"] == "interact"] == ["同じ番号の注文があります", "注文を登録しました"]))
    f1, f2 = at("order.setStatus"), at("mailer.send")
    out.append(("木の決まり: まとまりの中のコメントの行で、機械が切る。手がかりは、飾りの線でなく、コメントの文字",
                [(n["kind"], n["lines"], n.get("hint", "")) for n in procs["finish"]["body"][:2]] == [("seq", [f1, f1], ""), ("seq", [f2, f2], "担当者に知らせる")]
                and [n.get("tech") for n in procs["create"]["body"][3:5]] == ["変数の準備", "ログ"]))
    out.append(("木の決まり: プログラムごと終える手続き（最後の文が、プログラムの終了）に印が付く。図にする文が無く、何か所からも呼ばれる手続きは、図にしない",
                [p["name"] for p in st["procedures"] if p.get("ends")] == ["stop", "abort"] and procs["abort"]["body"][-1]["what"] == "program"
                and st.get("blank") == [f"{M}.abort"] and procs["abort"]["boxes"] == 0))

    def case(name, change, expect):
        t = copy.deepcopy(tree)
        change(t)
        errs = code_lang.tree_problems(t, mods, "java")
        hit = next((e for e in errs if expect in e), None)
        out.append((f"木の誤り: {name}: {(hit or '')[:100]}" if hit else f"木の誤り: {name}: 「{expect}」が出ない\n{errs[:4]}", bool(hit)))
    body = lambda t, k=0: t["procedures"][k]["body"]
    case("calls の名前が、条件の行に無い", lambda t: body(t)[1].update(calls=[f"{M}.finish"]), "calls の名前")
    case("calls の手続きが、木に無い", lambda t: body(t)[1].update(calls=[f"{M}.nothing"]), "木の手続きに無い")
    case("calls を、条件でないブロックに付ける", lambda t: body(t)[2].update(calls=[f"{M}.finish"]), "calls は付けられない")
    case("leave を、戻る語の無い行に付ける", lambda t: body(t)[2].update(leave="proc"), "抜ける・戻る語")
    case("leave を、分岐に付ける", lambda t: body(t)[1].update(leave="proc"), "leave は付けられない")
    case("exit に io を付ける", lambda t: body(t, 2)[3].update(io=[{"kind": "画面", "name": "redirect", "way": "out", "line": at('return "redirect')}]), "io は付けられない")
    case("やり取りの text が、その行に無い", lambda t: body(t)[5].update(text="登録に失敗しました"), "その行の文字に無い")
    split = copy.deepcopy(tree)                                      # まとまりの区切り方だけが違う木
    split["procedures"][2]["body"][0:1] = [{"kind": "seq", "lines": [f1, f1]}, {"kind": "seq", "lines": [f2, f2], "io": tree["procedures"][2]["body"][0]["io"]}]
    ag = code_lang.tree_agreement(tree, split)
    same = [p["body"] for p in cs.build_struct("M01", "java", mods, tree=split)["procedures"]] == [p["body"] for p in st["procedures"]]
    tech = copy.deepcopy(tree)
    tech["procedures"][0]["body"][3].pop("tech")
    ag2 = code_lang.tree_agreement(tree, tech)
    out.append(("2回読んだ木の一致: まとまりの区切り方だけの違いは、食い違いに数えない（できる構造も同じ）。技術的な文の印の違いは、数える",
                ag["rate"] == 1.0 and ag["raw"][0] < ag["raw"][1] and same and ag2["rate"] < 1 and ag2["flow"][0] == ag2["flow"][1] == ag["flow"][0]))
    K = TREE_COBOL.replace("001600     STOP RUN.", "001600     CONTINUE\n001650     STOP RUN.").splitlines()
    keep = code_lang.plain_lines("cobol", K)[0]
    stop = next(k for k, ln in enumerate(K, 1) if "STOP RUN" in ln)
    out.append(("木の決まり: COBOL の何もしない文（CONTINUE）は、どのブロックにも入れなくてよい", stop - 1 not in keep and stop in keep))
    sv = cs.build_struct("M09", "vba", [{"name": "Module1", "kind": "standard", "lines": VBA_END.splitlines()}])
    pv = {p["name"]: p for p in sv["procedures"]}
    out.append(("VBA: End は、プログラムごと終える（呼んだ側に戻らない）。いつも終える手続きに印が付く",
                cs.tree_errors(sv) == [] and [n.get("what") for n in code_lang.walk(pv["確認"]["body"]) if n["kind"] == "exit"] == ["program"]
                and [p["name"] for p in sv["procedures"] if p.get("ends")] == ["中止"]))
    # 通しで: 図では、プログラムごと終える所から「終了（途中で終わる）」につながる（呼んだ側の続きに、戻らない）
    src, w = tmp / "web_in", tmp / "web_work"
    (src / "受注画面").mkdir(parents=True)
    (src / "受注画面" / M).write_text(TREE_WEB, encoding="utf-8")
    subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(w)], capture_output=True, text=True)
    run("render_pages.py", w)
    (w / "code").mkdir(exist_ok=True)
    (w / "code/M01.tree.json").write_text(json.dumps(tree, ensure_ascii=False), encoding="utf-8")
    run("code_struct.py", w)
    code1, text1 = run("check_flow.py", w, "--stage", "code")
    run("systems.py", w, "draft")
    edit(w / "structure/systems.json", lambda j: [x.update(kind="core") for x in j["systems"]])
    run("activity.py", w, "apply", "--all", "--no-edits")
    act = json.loads((w / "activities/A-01-1.activity.json").read_text(encoding="utf-8")) if (w / "activities/A-01-1.activity.json").exists() else {"nodes": [], "edges": []}
    dec = lambda word: next((n["id"] for n in act["nodes"] if n["type"] == "decision" and word in (n.get("label") or "")), None)
    yes = lambda word: [e["to"] for e in act["edges"] if e["from"] == dec(word) and e.get("label") == "はい"]
    no = lambda word: [e["to"] for e in act["edges"] if e["from"] == dec(word) and e.get("label") == "いいえ"]
    out.append(("通し（プログラムごと終える）: 異常終了の分岐は消えず、「はい」は「終了（プログラムを止める）」へ（図にする文の無い手続き・呼ぶ側に入れた手続きの、どちらでも）、「いいえ」は続きへ。点検（コードの段）は FAIL 0",
                code1 == 0 and "0 FAIL" in text1 and sorted(yes("order == null") + yes("isUrgent") + yes("isClosed")) == ["end3", "end3-2", "end3-3"]   # 終了の丸は、抜ける所ごと
                and all(n["label"] == "終了（プログラムを止める）" for n in act["nodes"] if n["id"].startswith("end3")) and sum(n["id"].startswith("end3") for n in act["nodes"]) == 3
                and no("order == null") == [dec("isDuplicate")] and no("isUrgent") == [dec("isClosed")] and no("isClosed") == ["end"]
                and not any(n["type"] == "subprocess" for n in act["nodes"])
                and any("同じ番号の注文があります" in (n.get("label") or "") for n in act["nodes"])))
    return out


def access_tests(tmp):
    """Access を書き出したテキスト: 書き出しの道具の体裁／読み取り（フォームの定義・マクロ・クエリ・テーブル）／通しで作って点検／機械が作る確認事項"""
    import code_struct as cs
    import extract_access as xa
    out = []
    tool = SKILL / "assets" / "access_export"
    core = {}
    for name in ("access_export.vbs", "access_export.bas"):
        raw = (tool / name).read_bytes()
        text = raw.decode("ascii", "replace").replace("\r\n", "\n")
        core[name] = (text[text.index("'==== core begin"):text.index("'==== core end ====")], all(b < 128 for b in raw), text)
    lines = core["access_export.bas"][2].splitlines()
    st = cs.build_struct("T", "vba", [{"name": "AsisFlowAccessExport", "kind": "standard", "lines": lines}])
    words = "\n".join(lines).lower()
    out.append(("書き出しの道具: .vbs と .bas の中身（共通の部分）が同じ・ASCII だけ・ブロックの対応が取れている・元のファイルを開く命令を使っていない",
                core["access_export.vbs"][0] == core["access_export.bas"][0] and core["access_export.vbs"][1] and core["access_export.bas"][1]
                and cs.tree_errors(st) == [] and len(st["procedures"]) == 15 and "opencurrentdatabase" not in words and ".run " not in words and "opendatabase(ssrc, false, true)" in words))
    src, good = tmp / "acc_in", tmp / "acc_good"
    subprocess.run([sys.executable, str(SKILL / "evals" / "make_fixture.py"), str(src), "--access"], capture_output=True, text=True)
    (src / "受注管理.accdb").write_bytes(b"\x00\x01\x00\x00Standard ACE DB\x00" + b"\x00" * 64)        # Access のファイルそのもの（中身は読まない）
    (src / "受注管理_export" / xa.TMP_NAME).write_bytes(b"\x00" * 16)                                   # 道具が消し残した、仮のデータベース
    r = subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(good)], capture_output=True, text=True)
    run("render_pages.py", good)
    man = json.loads((good / "manifest.json").read_text(encoding="utf-8"))
    by = {d["name"]: d for d in man["docs"]}
    out.append(("棚卸し・ページ化: 書き出したテキスト 9 件はプログラムの資料。Access のファイルは開かず、書き出しのほうを読む。仮のデータベースは数えない",
                len(man["docs"]) == 10 and all(d["kind"] == "code" for n, d in by.items() if n != "受注管理.accdb") and by["受注管理.accdb"].get("read_via") == "受注管理_export"
                and by["受注管理.accdb"]["page_list"] == [] and by["F_メニュー.form.txt"]["access"] == {"object": "form", "name": "メニュー"}
                and by["M_modCommon.bas"]["access"]["name"] == "modCommon" and "Access の書き出し" in r.stdout + json.dumps(man, ensure_ascii=False)))
    shutil.copytree(SKILL / "evals" / "fixture_access_expected", good, dirs_exist_ok=True)
    ids = {d["name"]: d["doc_id"] for d in man["docs"]}
    edit(good / "structure/systems.json", lambda j: j["systems"][0].update(docs=sorted(v for n, v in ids.items() if n != "受注管理.accdb")))
    log = []
    for script, args in [("code_struct.py", []), ("activity.py", ["apply", "--all"]), ("build_activity.py", []), ("build_matrix.py", []), ("build_reports.py", []),
                         ("check_extract.py", []), ("check_flow.py", []), ("check_layout.py", ["--no-png"])]:
        code, text = run(script, good, *args)
        log.append((script, code == 0 and "0 FAIL" in text if script.startswith("check") else code == 0, text))
    bad = [(sc, text) for sc, passed, text in log if not passed]
    out.append(("Access の書き出しだけの入力（マニュアル無し）: 通しで作って、点検が FAIL 0" + ("" if not bad else f"（{bad[0][0]}）\n{bad[0][1][-700:]}"), not bad))
    st = next(iter(cs.load_structs(good).values()))
    ent = {e["label"]: (e["trigger"], e["procedure"]) for e in st["entries"]}
    out.append(("起動点: フォームの定義でイベントに割り当てたものだけ（コード・マクロ・埋め込みマクロ）。見出しはボタンの文字・項目に付いたラベル。割り当ての無い手続きは、たどれない手続き",
                len(st["entries"]) == 9 and ent.get("フォーム「メニュー」の「月次集計」ボタンをクリック") == ("button", "月次集計.（マクロ）")
                and ent.get("フォーム「受注入力」の「数量」の値を変えたあと", ("",))[0] == "form" and ent.get("フォーム「メニュー」の「終了」ボタンをクリック", ("", ""))[1].endswith("（埋め込みマクロ）")
                and "データベースを開いたとき（AutoExec マクロ）" in ent and any("最初に開く" in k for k in ent) and ent.get("レポート「請求書」で、データが無いとき", ("",))[0] == "event"
                and st["unreached"] == ["Form_受注入力.cmd削除_Click"] and cs.tree_errors(st) == []))
    procs = {p["id"]: p for p in st["procedures"]}
    mac = procs["月次集計.（マクロ）"]
    kinds = [n["kind"] for n in mac["body"]]
    io = {(x["kind"], x["name"], x["way"]) for p in st["procedures"] for x in p["io"]}
    out.append(("マクロ: アクションの並び・条件つきの行と「...」・VBA の関数の呼び出し・クエリの定義から引いたテーブル・書き出すファイル",
                kinds == ["seq", "seq", "seq", "if", "seq", "call", "interact", "seq"] and [n["kind"] for n in mac["body"][3]["branches"][0]["body"]] == ["interact", "exit"]
                and mac["body"][5]["to"] == "modCommon.月次ログを書く" and {("クエリ", "Q_月次売上_追加", "out"), ("テーブル", "T_月次売上", "out"), ("テーブル", "T_受注", "in"),
                                                                      ("ファイル", "月次売上.xlsx", "out"), ("ファイル", "受注データ.csv", "in"), ("帳票", "請求書", "out")} <= io
                and mac["body"][2]["io"][2].get("via", {}).get("module") == "access_queries"))
    acc = st["access"]
    out.append(("Access の定義: フォーム 2・レポート 1・マクロ 2・クエリ 6・テーブル 6（うちリンク 1）・起動時に開くフォーム・書き出せなかったもの",
                (len(acc["forms"]), len(acc["reports"]), len(acc["macros"]), len(acc["queries"]), len(acc["tables"])) == (2, 1, 2, 6, 6) and acc["startup_form"] == "メニュー"
                and [x["name"] for x in acc["failed"]] == ["旧請求書"] and sum(bool(t.get("linked")) for t in acc["tables"]) == 1 and acc["source"] == "受注管理.accdb"
                and next(q for q in acc["queries"] if q["name"] == "Q_在庫_引当")["writes"] == ["T_在庫"]))
    act = json.loads((good / "activities/A-01-3.activity.json").read_text(encoding="utf-8"))
    files = [n for n in act["nodes"] if n["type"] == "file"]
    qdoc = ids["access_queries.sql"]
    out.append(("アクティビティ図（マクロ）: 構造の段は「行の形から」。クエリの定義から引いたテーブルの出典に、クエリの定義の行が入る",
                act["tier"] == "heuristic" and any(n["label"] == "テーブル T_受注" and [s["doc_id"] for s in n["sources"]] == [ids["X_月次集計.macro.txt"], qdoc] for n in files)))
    qs = json.loads((good / "questions.json").read_text(encoding="utf-8"))["questions"]
    local = {q["local"]: q for q in qs}
    out.append(("機械が作る確認事項: 書き出せなかったもの・たどれない手続き。マニュアルが無いときの紐付けの問いは、システムごとに1つ。Access のファイルを読めていない、とは出さない",
                {f"export-failed:{st['doc_id']}", f"unreached:{st['doc_id']}", "unlinked:SYS-01"} <= set(local) and not any(k.startswith("unlinked:M") for k in local)
                and not any(q["origin"] == "inventory" for q in qs)))
    lay = xa.parse_layout(["Version =20", "Begin Form", '    Caption ="長い見出しの"', '        "続き"', "    PrtMip = Begin", "        0x00", "    End", "    Begin", "        Begin Section",
                           '            Name ="詳細"', '            OnFormat ="[Event Procedure]"', "        End", "    End", "End", "CodeBehindForm", "Option Explicit"])
    out.append(("フォームの定義: 長い文字の続き・16 進の値のブロック・セクションのイベント・コードの最初の行", lay["caption"] == "長い見出しの続き" and lay["code_from"] == 16
                and [(e["owner"], e["handler"], e["kind"]) for e in lay["events"]] == [("詳細", "詳細_Format", "code")]))
    st2 = cs.build_struct("T", "accessmacro", [{"name": "新しいマクロ", "kind": "macro", "lines": ["Version =196611", "Begin", '    Comment ="_AXL:<?xml version=\\"1.0\\"?><UserInterfaceMacro/>"', "End"]}])
    out.append(("マクロ: アクションの行が無い（新しい形だけの）マクロは、読めなかったものとして出す（黙って落とさない）",
                st2["procedures"] == [] and [x["what"] for x in st2["access"]["unread"]] == ["マクロ"] and st2["tier"] == "parsed"))

    def case(name, change, expect, script="check_flow.py", args=()):
        w = tmp / f"acc_case{len(out)}"
        shutil.copytree(good, w)
        change(w)
        code, text = run(script, w, *args)
        hit = next((ln for ln in text.splitlines() if expect in ln), None)
        out.append((f"{name}: {(hit or '')[:110]}" if hit and code != 0 else f"{name}: 「{expect}」が出ない（終了コード {code}）\n{text[-600:]}", bool(hit) and code != 0))
        shutil.rmtree(w, ignore_errors=True)
    twin = tmp / "acc_twin"                                           # 同じプログラムから2回作った図の揺れ: 片方だけ、2つのまとまりを1つにまとめ、名前を変える
    shutil.copytree(good, twin)

    def vary(e):
        e["ops"] = [{"op": "merge", "items": ["（マクロ）.7", "（マクロ）.13"]}] + [o for o in e["ops"] if o.get("item") != "（マクロ）.13"]
        e["ops"][1]["label"] = "月次売上を作り直す"
    edit(twin / "activities/A-01-3.edits.json", vary)
    run("activity.py", twin, "apply", "--id", "A-01-3")
    r = subprocess.run([sys.executable, str(S / "diff_runs.py"), "activities", str(good), str(twin), "--only", "A-01-3", "--min", "0.9"], capture_output=True, text=True)
    out.append(("揺れの測定（アクティビティ図）: 箱の対応・まとめ方・ラベルの違いを数え、下限を下回ると終了コード 1",
                r.returncode == 1 and "| 箱の対応（同じ行を箱にしたか） | 12 | 13 |" in r.stdout and "片方だけにある箱 P-001" in r.stdout and "アクティビティ図 1 枚" in r.stdout))
    shutil.rmtree(twin, ignore_errors=True)
    fpage = f"pages/{ids['F_メニュー.form.txt']}/p0001.json"
    case("Access: フォームの定義のページを書き換える（ボタンの割り当てを消す）",
         lambda w: (w / fpage).write_text((w / fpage).read_text(encoding="utf-8").replace('OnClick =\\"月次集計\\"', 'Tag =\\"月次集計\\"'), encoding="utf-8"),
         "取り出し直した構造と違う", args=["--stage", "code"])
    case("Access: マクロの図の箱を1つ消す（行の網羅）", lambda w: edit(w / "activities/A-01-3.activity.json", lambda a: a["lines"].pop(1)), "A-01-3", args=["--stage", "activity"])
    return out


def direct_tests(tmp):
    """Access のファイルの直接読み（access_direct.py）: 定義・マクロ・クエリの読み取り（架空のバイト列で）／直接読みのテキストから、通しで構造と確認事項まで。
    実物のファイルでの答え合わせは evals/real_samples.py（サンプルは、リポジトリに入れていない）"""
    import struct
    import access_direct as ad
    import code_struct as cs
    import extract_access as xa
    out = []
    u = lambda t: t.encode("utf-16-le")
    rec = lambda pid, typ, val, wide=4: struct.pack("<IHIII", 7, pid, typ, wide, len(val)) + val
    FD, FE, FF = b"\xfd\x00\x00\x00", b"\xfe\x00\x00\x00", b"\xff\x00\x00\x00"
    label = struct.pack("<H", 100) + rec(20, 10, u("lbl数量"), 64) + rec(17, 12, u("数量")) + FD
    box = struct.pack("<H", 109) + rec(20, 10, u("txt数量"), 64) + rec(86, 12, u("[Event Procedure]")) + rec(999, 12, u("[Event Procedure]")) + FE + label
    btn = (struct.pack("<H", 104) + rec(20, 10, u("cmd登録"), 64) + rec(17, 12, u("登録\r\nする")) + rec(126, 12, u("[Event Procedure]"))
           + rec(7, 11, b"\x89PNG" * 3, 0xFFFF0000) + rec(998, 12, u("月次取込")) + FD)
    sec = struct.pack("<H", 152) + rec(20, 10, u("詳細"), 64) + rec(41, 3, b"\x18\x00") + FF + struct.pack("<H", 2) + box + btn
    blob = struct.pack("<HHIH", 21, 20, 0, 0x96) + rec(17, 12, u("受注入力")) + rec(156, 12, u("T_受注")) + rec(127, 12, u("[Event Procedure]")) + rec(3, 1, b"\x01", 1) \
        + FF + struct.pack("<H", 1) + sec
    code = ["Attribute VB_Name = \"Form_受注入力\"", "Option Explicit", "Private Sub Form_Load()", "    Me.Requery", "End Sub", "Private Sub txt数量_AfterUpdate()", "    Call 金額を計算する",
            "End Sub", "Private Sub txt数量_Change()", "    Me.Dirty = True", "End Sub", "Private Sub cmd登録_Click()", "    DoCmd.OpenQuery \"Q_受注_追加\"", "End Sub",
            "Private Sub 金額を計算する()", "    Me.txt金額 = Me.txt数量 * 100", "End Sub"]
    handlers = {x.split("(")[0].split()[-1].lower(): x.split("(")[0].split()[-1] for x in code if x.startswith("Private Sub")}
    tree, version = ad.layout_tree(blob)
    lines, notes = ad.layout_text(tree, version, "Form", handlers, {"月次取込"})
    lay = xa.parse_layout(lines + ["CodeBehindForm"] + code)
    ev = {(e["owner"], e["prop"], e["kind"], bool(e.get("guess"))) for e in lay["events"]}
    ctl = {c["name"]: c for c in lay["controls"]}
    out.append(("直接読み（フォームの定義）: 入れ子・付いているラベル・改行のある表示の文字・バイナリの値を越えて、最後まで読む。番号の分かるイベントはそのまま、分からないものは、手続きの名前から決めて印を付ける",
                ev == {("Form", "OnLoad", "code", False), ("txt数量", "AfterUpdate", "code", False), ("txt数量", "OnChange", "code", True), ("cmd登録", "OnClick", "code", False),
                       ("cmd登録", "OnUnknown998", "macro", False)}
                and ctl["txt数量"]["label"] == "数量" and ctl["cmd登録"]["caption"] == "登録\nする" and (lay["caption"], lay["record_source"]) == ("受注入力", "T_受注")
                and notes["guess"] == ["txt数量（1 件）"] and not notes["open"]))

    def refuse(data, word):
        try:
            ad.layout_tree(data)
        except ad.Unreadable as e:
            return word in str(e)
        return False
    out.append(("直接読み（フォームの定義）: 前の形式・途中で切れた定義・余りのある定義は、読めていないとする（半端に読まない）",
                refuse(struct.pack("<HHIH", 19, 19, 0, 0x96) + b"\x0f\x13\x1c\x30", "前の形式") and refuse(blob[:-9], "形が合わない") and refuse(blob + b"\x00\x00", "最後まで読めない")))
    xml = ('<?xml version="1.0" encoding="UTF-16" standalone="no"?>\r\n<UserInterfaceMacro MinimumClientDesignVersion="14.0.0000.0000" xmlns="http://example.invalid/macro">'
           '<Statements><Comment>月次の取り込み</Comment><ConditionalBlock><If><Condition>[Forms]![メニュー]![chk確定]=True</Condition><Statements>'
           '<Action Name="RunCode"><Argument Name="FunctionName">月次ログを書く()</Argument></Action>'
           '<Action Name="OpenQuery"><Argument Name="QueryName">Q_受注_追加</Argument><Argument Name="DataMode">Edit</Argument></Action></Statements></If>'
           '<Else><Statements><Action Name="MessageBox"><Argument Name="Message">確定のチェックを入れてください</Argument></Action><Action Name="StopMacro"/></Statements></Else>'
           '</ConditionalBlock><Action Name="ImportExportSpreadsheet"><Argument Name="TableName">T_受注</Argument><Argument Name="FileName">受注.xlsx</Argument></Action>'
           '</Statements></UserInterfaceMacro>')
    cut = len(xml) // 2
    mblob = b"\x52\x00\x01\x00" + b"\xff" * 8 + u("__*L0_") + b"\x00\x00" + u("_AXL:" + xml[:cut]) + b"\x00\x00" + b"\x10\x00\xff\xff" + u("_AXL:" + xml[cut:]) + b"\x00\x00\x00"
    got = ad.macro_xml(mblob)
    mlines = ad.macro_text(got or "<x/>")
    rows, axl = xa.macro_rows(mlines)
    joined = "".join(v[5:] for r in rows for nm, v, _k in r["props"] if nm == "Comment")
    st = cs.build_struct("T", "accessmacro", [
        {"name": "月次取込", "kind": "macro", "lines": mlines},
        {"name": "modCommon", "kind": "standard", "lines": ["Public Function 月次ログを書く()", "    Debug.Print 1", "End Function"]},
        {"name": "access_queries", "kind": "queries", "lines": ["-- query: Q_受注_追加\ttype=append", "INSERT INTO T_受注 (品番)", "SELECT W.品番", "FROM W_取込 AS W;"]}])
    mac = next((p for p in st["procedures"] if p["kind"] == "macro"), {"body": [{}]})
    top = [n["kind"] for n in mac["body"]]
    blk = next((n for n in mac["body"] if n["kind"] == "if"), {"branches": []})
    io = {(x["kind"], x["name"], x["way"]) for x in mac.get("io", [])}
    out.append(("直接読み（マクロ）: 入れ物の中の XML の定義を、順につないで取り出す。アクション・条件を1行ずつに分けても、つなげば元の XML。条件のブロック・呼び出し・クエリのテーブルを読み、"
                "ファイルを読み書きするアクションは「読めていない」に出す",
                got == xml and joined == xml.replace("\r\n", "\n") and axl and mac.get("read") == "axl" and st["tier"] == "parsed" and top == ["if", "seq"]
                and [(b["label"], [n["kind"] for n in b["body"]]) for b in blk["branches"]] == [("then", ["call", "seq"]), ("else", ["interact", "exit"])]
                and blk["cond"] == "[Forms]![メニュー]![chk確定]=True" and {("クエリ", "Q_受注_追加", "out"), ("テーブル", "T_受注", "out"), ("テーブル", "W_取込", "in")} <= io
                and [x["what"] for x in st["access"]["unread"]] == ["マクロのアクションの入出力"] and cs.tree_errors(st) == []))
    out.append(("直接読み（マクロ）: XML の定義が無い・壊れているマクロは、読めていないとする", ad.macro_xml(u("OpenForm") + b"\x00\x00") is None and ad.macro_xml(u("_AXL:<a><b></a>") + b"\x00\x00") is None))
    row = lambda attr, **kw: {"Attribute": attr, "Order": struct.pack(">I", kw.pop("n", 1)), **kw}
    sel = [row(1, Flag=1), row(3, Flag=2), row(5, Name1="T_受注", n=1), row(5, Name1="T_得意先", n=2), row(5, Name1="T 商品", Name2="商品", n=3),
           row(6, Expression="T_受注.*", n=1), row(6, Expression="Sum(T_受注.数量)", Name1="数量 計", n=2),
           row(7, Flag=2, Name1="T_受注", Name2="T_得意先", Expression="T_受注.得意先ID = T_得意先.ID", n=1),
           row(7, Flag=1, Name1="T_受注", Name2="商品", Expression="T_受注.品番 = 商品.品番", n=2), row(7, Flag=1, Name1="T_受注", Name2="商品", Expression="T_受注.色 = 商品.色", n=3),
           row(8, Expression="((T_受注.数量)>0)"), row(11, Expression="T_受注.受注日", Name1="D")]
    app = [row(1, Flag=3, Name1="T_月次 売上"), row(3, Flag=0), row(5, Name1="T_受注"), row(6, Expression="T_受注.品番", Name2="品番", n=1), row(6, Expression="T_受注.数量", Name2="数量", n=2)]
    upd = [row(1, Flag=4), row(2, Flag=8, Name1="締め日"), row(5, Name1="T_在庫"), row(6, Expression="[数量]-1", Name2="T_在庫.数量"), row(8, Expression="(T_在庫.品番)=\"A\"")]
    uni = [row(1, Flag=9), row(3, Flag=1), row(5, Expression="SELECT 品番 FROM T_受注\r\n", Name2="X7YZ_____1", n=1), row(5, Expression="SELECT 品番 FROM T_在庫", Name2="X7YZ_____2", n=2)]
    out.append(("直接読み（クエリ）: 部品（MSysQueries の行）から SQL を組み立てる（結合の入れ子・同じ2表の条件は AND・別名と空白のある名前・パラメータ・追加・更新・ユニオン・パススルー）",
                ad.query_sql(sel, "select")[0] == "SELECT DISTINCT T_受注.*, Sum(T_受注.数量) AS [数量 計]\nFROM (T_受注 LEFT JOIN T_得意先 ON T_受注.得意先ID = T_得意先.ID) INNER JOIN [T 商品] AS 商品 "
                "ON (T_受注.品番 = 商品.品番) AND (T_受注.色 = 商品.色)\nWHERE ((T_受注.数量)>0)\nORDER BY T_受注.受注日 DESC;"
                and ad.query_sql(app, "append")[0] == "INSERT INTO [T_月次 売上] (品番, 数量)\nSELECT T_受注.品番, T_受注.数量\nFROM T_受注;"
                and ad.query_sql(upd, "update")[0] == "PARAMETERS 締め日 DateTime;\nUPDATE T_在庫\nSET T_在庫.数量 = [数量]-1\nWHERE (T_在庫.品番)=\"A\";"
                and ad.query_sql(uni, "union")[0] == "SELECT 品番 FROM T_受注\nUNION ALL SELECT 品番 FROM T_在庫;"
                and ad.query_sql([row(1, Flag=8, Name1="ODBC;DSN=基幹;PWD=secret", Expression="EXEC 月次\x00")], "passthrough") == ("EXEC 月次", "ODBC;DSN=基幹;PWD=" + cs.MASK)
                and xa.sql_tables(ad.query_sql(app, "append")[0]) == (["T_月次 売上"], ["T_受注"])))
    names = b"\x00" * 4 + b"".join(b"\x04" + bytes([len(u(n)) + 4]) + u(n) + struct.pack("<I", k) for k, n in ((0, "メニュー"), (9, "受注入力")))

    def opens(data):
        f = tmp / "direct_probe.accdb"
        f.write_bytes(data)
        try:
            ad.read(f)
        except ad.Unreadable as e:
            return str(e)
        return ""
    head = b"\x00\x01\x00\x00Standard ACE DB\x00"
    out.append(("直接読み: 入れ物の目次。Access のファイルでないもの・97 以前の形式・システム テーブルを読めないもの（暗号化）は、読まない",
                ad.dir_names(names) == {0: "メニュー", 9: "受注入力"} and ad.dir_names(names[:-3]) is None and "形をしていない" in opens(b"PK\x03\x04" + b"\x00" * 64)
                and "97 以前" in opens(b"\x00\x01\x00\x00Standard Jet DB\x00\x00" + b"\x00" * 9000) and "暗号化" in opens(head + b"\x02" + b"\x5a" * 13000)))
    # 直接読みのテキスト（架空）から、通しで: 棚卸し → ページ化 → 構造 → 台帳 → 確認事項
    src, work = tmp / "direct_in", tmp / "direct_work"
    folder = src / "受注管理_direct"
    folder.mkdir(parents=True)
    (src / "受注管理.accdb").write_bytes(head + b"\x00" * 64)
    put = lambda name, ls: (folder / name).write_text("\n".join(ls) + "\n", encoding="utf-8")
    put("F_受注入力.form.txt", lines + ["CodeBehindForm"] + code)
    put("F_検索.form.txt", ad.stub_layout("Form", "前の形式で保存されている（版 19）") + ["CodeBehindForm", "Attribute VB_Name = \"Form_検索\"", "Private Sub cmd検索_Click()", "    DoCmd.OpenQuery \"Q_受注_追加\"",
                                                                               "End Sub", "Private Sub 絞り込む()", "    Me.Filter = \"\"", "End Sub"])
    put("X_月次取込.macro.txt", mlines)
    put("M_modCommon.bas", ["Attribute VB_Name = \"modCommon\"", "Public Function 月次ログを書く()", "    Debug.Print 1", "End Function"])
    put("access_queries.sql", ["-- query: Q_受注_追加\ttype=append", "INSERT INTO T_受注 (品番)", "SELECT W.品番", "FROM W_取込 AS W;"])
    put("access_tables.txt", ["[table]\tT_受注\trows=3", "  ID\tautonumber\tpk", "  品番\ttext(20)", "[table]\tW_取込\trows=0", "  品番\ttext(20)"])
    put("access_export.txt", ["tool\tasis-flow access_direct", "source\t受注管理.accdb", "via\tdirect", "access\tfile-format-3", "form\t受注入力\tF_受注入力.form.txt\tok",
                              "form\t検索\tF_検索.form.txt\tok (layout unread: 前の形式で保存されている（版 19）)", "macro\t月次取込\tX_月次取込.macro.txt\tok", "module\tmodCommon\tM_modCommon.bas\tok",
                              "macro\t旧メニュー\t\tFAILED XML の定義が入っていない（前の形のマクロ）。直接読みでは読めない", "tables\t2\taccess_tables.txt\tok", "queries\t1\taccess_queries.sql\tok",
                              "note\tフォーム「受注入力」のイベントの割り当てのうち、種類の番号が分からないものは、コードの手続きの名前から決めた: txt数量（1 件）",
                              "note\tリレーションと、起動時に開くフォームは、直接読みでは読んでいない"])
    r = subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(work)], capture_output=True, text=True)
    run("render_pages.py", work)
    man = json.loads((work / "manifest.json").read_text(encoding="utf-8"))
    by = {d["name"]: d for d in man["docs"]}
    out.append(("直接読みのテキスト: 棚卸しは「ファイルを直接読んだもの」と記し、Access のファイルは、そのテキストのほうを読む（直接読みの印が付く）。直接読みを、もう一度はしない",
                by["受注管理.accdb"].get("read_via") == "受注管理_direct" and by["受注管理.accdb"].get("read_direct") is True and "直接読んだもの" in by["M_modCommon.bas"]["note"]
                and "直接読み（" not in r.stdout and not (work / "direct").exists()))
    run("code_struct.py", work)
    st = next(iter(cs.load_structs(work).values()))
    ent = {e["label"]: e for e in st["entries"]}
    guess = ent.get("フォーム「受注入力」の「数量」の値を変えたとき", {"evidence": ""})
    out.append(("直接読みのテキスト（起動点）: 定義を読めたフォームは、割り当てのとおり（種類を名前から決めたものは、根拠にそう書く）。定義を読めていないフォームは、手続きの名前から候補を出し、未確認と記す",
                "手続きの名前から" in guess["evidence"] and "フォーム「検索」の「cmd検索」をクリック（割り当ては未確認）" in ent and "フォーム「受注入力」の「登録\nする」ボタンをクリック" in ent
                and any("種類は読めていない" in k for k in ent) and not any("絞り込む" in e["procedure"] or "金額を計算する" in e["procedure"] for e in st["entries"])
                and {x["what"] for x in st["access"]["unread"]} == {"フォームの定義", "マクロのアクションの入出力"} and st["access"].get("via") == "direct" and cs.tree_errors(st) == []))
    run("systems.py", work, "draft")
    import systems
    qs = dict(systems.machine_questions(work))
    did = st["doc_id"]
    out.append(("直接読みのテキスト（確認事項）: 直接読んだこと・読めていない点・読めなかったもの（書き出しを頼む言い方）が、自動で載る",
                "直接読みました" in qs.get(f"direct:{did}", {}).get("question", "") and "リレーション" in qs[f"direct:{did}"]["question"]
                and "直接読むやり方では、読めませんでした" in qs.get(f"export-failed:{did}", {}).get("question", "") and "直接読んだテキスト" in qs.get(f"unread:{did}", {}).get("question", "")))
    alone = tmp / "direct_alone"
    alone.mkdir()
    (alone / "暗号化.accdb").write_bytes(head + b"\x02" + b"\x5a" * 13000)
    w2 = tmp / "direct_work2"
    r = subprocess.run([sys.executable, str(S / "inventory.py"), str(alone), "-w", str(w2)], capture_output=True, text=True)
    code2, text2 = run("render_pages.py", w2)
    out.append(("直接読み: 読めない Access のファイルだけの入力は、棚卸しで理由を知らせ、ページ化で、書き出しの手順を案内する（止まらない）",
                "直接は読めない（システム テーブルを読めない" in r.stdout and "書き出しの道具" in text2 and "暗号化されている" in text2))
    return out


FIX_MODULE = """Attribute VB_Name = "Module1"
Option Explicit

Public Sub 印刷する()
Attribute 印刷する.VB_ProcData.VB_Invoke_Func = "P\\n14"
    Dim 期限 As Date: 期限 = #12/31/2026 10:30:00 AM#: If 期限 < Date Then Exit Sub
    Sheets("メニュー").更新する
    Dim 明細 As cls明細
    Set 明細 = New cls明細
    明細.読み込む
#If VBA7 Then
    Call 取込
#Else
    MsgBox "古い版では動きません": Exit Sub
#End If
    ActiveSheet.PrintOut
End Sub

Public Sub 取込()
Attribute 取込.VB_ProcData.VB_Invoke_Func = " \\n14"
    Range("A1").Value = 1
End Sub

Public Sub 月次集計()
Attribute 月次集計.VB_ProcData.VB_Invoke_Func = "m\\n14"
    Application.CommandBars("帳票").Controls("再集計").OnAction = ActiveSheet.Name & ".再計算"
End Sub

Public Sub RibbonLoad(ribbon As Object)
    取込
End Sub

Public Sub 請求書作成(control As Object)
    月次集計
End Sub

Public Sub 手で動かす()
    MsgBox "完了"
End Sub
"""
FIX_SHEET = """Option Explicit

Public Sub 更新する()
    Me.再計算
End Sub

Public Sub 再計算()
    Range("B2").Value = Range("B1").Value * 2
End Sub

Private Sub Worksheet_Deactivate()
    Me.Calculate
End Sub
"""
FIX_CLASS = """Option Explicit
Private m_件数 As Long

Private Sub Class_Initialize()
    m_件数 = 0
End Sub

Public Sub 読み込む()
    m_件数 = Range("A1").Value
End Sub

Public Property Get 件数() As Long
    件数 = m_件数
End Property

Public Property Let 件数(ByVal v As Long)
    m_件数 = v
End Property
"""


def make_macro_parts(tmp):
    """マクロ付きブックの部品（架空。zip の中の XML とバイナリを、手で組み立てる。VBA の本体は入れない）→ (xlsm のパス, xlsb のパス)"""
    import struct
    import zipfile
    M = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rels = lambda items: ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                          + "".join(f'<Relationship Id="{i}" Type="{R}/{t}" Target="{g}"/>' for i, t, g in items) + "</Relationships>")
    sp = lambda i, name, text, macro: (f'<xdr:sp macro="{macro}"><xdr:nvSpPr><xdr:cNvPr id="{i}" name="{name}"/></xdr:nvSpPr>'
                                       f"<xdr:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></xdr:txBody></xdr:sp>")
    book = tmp / "fix_book.xlsm"
    with zipfile.ZipFile(book, "w") as z:
        z.writestr("xl/workbook.xml", f'<workbook xmlns="{M}" xmlns:r="{R}"><workbookPr codeName="DieseArbeitsmappe"/><sheets><sheet name="メニュー" sheetId="1" r:id="rId1"/>'
                   '<sheet name="データ" sheetId="2" state="hidden" r:id="rId2"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", rels([("rId1", "worksheet", "worksheets/sheet1.xml"), ("rId2", "worksheet", "/xl/worksheets/sheet2.xml")]))
        z.writestr("xl/worksheets/sheet1.xml", f'<worksheet xmlns="{M}" xmlns:r="{R}"><sheetPr codeName="Sheet1"/><sheetData/><controls>'
                   '<control shapeId="1025" r:id="rId3" name="Button 1"><controlPr defaultSize="0" macro="[0]!月次集計"/></control></controls></worksheet>')
        z.writestr("xl/worksheets/sheet2.xml", f'<worksheet xmlns="{M}"><sheetPr codeName="wsData"/><sheetData/></worksheet>')
        z.writestr("xl/worksheets/_rels/sheet1.xml.rels", rels([("rId1", "drawing", "../drawings/drawing1.xml"), ("rId2", "vmlDrawing", "../drawings/vmlDrawing1.vml"),
                                                               ("rId3", "ctrlProp", "../ctrlProps/ctrlProp1.xml")]))
        z.writestr("xl/drawings/drawing1.xml", '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                   'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">' + sp(2, "角丸 1", "印刷", "&apos;帳票.xlsm&apos;!Module1.印刷する")
                   + '<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="3" name="グループ 1"/></xdr:nvGrpSpPr>' + sp(4, "矢印 1", "取込む", "[0]!取込") + sp(5, "矢印 2", "", "[0]!取込")
                   + "</xdr:grpSp>" + sp(6, "四角 9", "説明の文字", "") + "</xdr:wsDr>")
        z.writestr("xl/drawings/vmlDrawing1.vml", '<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:x="urn:schemas-microsoft-com:office:excel">'
                   '<v:shape id="_x0000_s1025" type="#_x0000_t201"><v:textbox><div><font>月次 </font><font>集計</font></div></v:textbox>'
                   '<x:ClientData ObjectType="Button"><x:FmlaMacro>[0]!月次集計</x:FmlaMacro></x:ClientData></v:shape></xml>')
        z.writestr("xl/ctrlProps/ctrlProp1.xml", '<formControlPr xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" objectType="Button" macro="[0]!月次集計"/>')
        z.writestr("customUI/customUI14.xml", '<customUI xmlns="http://schemas.microsoft.com/office/2009/07/customui" onLoad="RibbonLoad"><ribbon><tabs><tab id="t" label="帳票">'
                   '<group id="g" label="作成"><button id="b1" label="請求書を作る" onAction="請求書作成"/><button id="b2" label="ヘルプ"/></group></tab></tabs></ribbon></customUI>')

    def rec(rid, body):                                              # バイナリ ブックのレコード（番号と長さは、7 ビットずつ）
        head = bytes([rid]) if rid < 128 else bytes([(rid & 0x7F) | 0x80, rid >> 7])
        n, size = len(body), b""
        while True:
            size += bytes([(n & 0x7F) | (0x80 if n > 0x7F else 0)])
            n >>= 7
            if not n:
                return head + size + body
    ws = lambda t: struct.pack("<I", len(t)) + t.encode("utf-16-le")
    xlsb = tmp / "fix_book.xlsb"
    with zipfile.ZipFile(xlsb, "w") as z:
        z.writestr("xl/workbook.bin", rec(131, b"") + rec(143, b"") + rec(156, struct.pack("<II", 0, 1) + ws("rId1") + ws("集計"))
                   + rec(156, struct.pack("<II", 2, 2) + ws("rId2") + ws("裏")) + rec(144, b"") + rec(132, b""))
        z.writestr("xl/_rels/workbook.bin.rels", rels([("rId1", "worksheet", "worksheets/sheet1.bin"), ("rId2", "worksheet", "worksheets/sheet2.bin")]))
        z.writestr("xl/worksheets/sheet1.bin", rec(129, b"") + rec(147, b"\x00" * 19 + ws("Sheet1")) + rec(145, b"") + rec(130, b""))
        z.writestr("xl/worksheets/sheet2.bin", rec(129, b"") + rec(145, b"") + rec(147, b"\x00" * 19 + ws("NG")) + rec(130, b""))
    return book, xlsb


def sample_fix_tests(tmp):
    """実物のサンプルで見つけて直した所を、架空のデータで固定する（マクロ付きブックの部品・VBA の書き方・起動点・道具のブック・版管理アドインの書き出し）"""
    import zipfile
    import code_struct as cs
    import common
    import extract_vba as xv
    out = []
    book, xlsb = make_macro_parts(tmp)
    btn = xv.buttons_of(str(book))
    out.append(("ブックの部品: 図形の割り当て（ブック名つき・[0]!・グループの中）。名前は、書いてある文字（無ければ図形の名前）。割り当ての無い図形は拾わない",
                btn[:3] == [{"sheet": "メニュー", "shape": "印刷", "macro": "Module1.印刷する"}, {"sheet": "メニュー", "shape": "取込む", "macro": "取込"},
                            {"sheet": "メニュー", "shape": "矢印 2", "macro": "取込"}]))
    out.append(("ブックの部品: フォームのボタンは、3か所（シート・旧来の描画・部品の性質）に書いてあっても1個。名前は、ボタンに書いてある文字",
                btn[3:] == [{"sheet": "メニュー", "shape": "月次 集計", "macro": "月次集計"}]))
    out.append(("ブックの部品: シートとブックの、VBA の中での名前。リボンの定義の割り当て（読み込み時・ボタン。割り当ての無いボタンは拾わない）",
                xv.sheet_codenames(str(book)) == {"Sheet1": "メニュー", "wsData": "データ"} and xv.book_codename(str(book)) == "DieseArbeitsmappe"
                and xv.ribbon_of(str(book)) == [{"label": "customUI", "macro": "RibbonLoad", "what": "load"}, {"label": "請求書を作る", "macro": "請求書作成", "what": "button"}]))
    with zipfile.ZipFile(xlsb) as z:
        parts = xv.sheet_parts(z)
    out.append(("ブックの部品: バイナリ ブック（xlsb）のシート名・表示の状態・VBA の中での名前（セルの手前まで）",
                parts == [("集計", "xl/worksheets/sheet1.bin", "visible"), ("裏", "xl/worksheets/sheet2.bin", "veryHidden")] and xv.sheet_codenames(str(xlsb)) == {"Sheet1": "集計"}
                and xv.unread_of(str(xlsb), [1]) == [] and xv.unread_of(str(xlsb), []) == []))
    broken = tmp / "fix_none.xlsm"
    with zipfile.ZipFile(broken, "w") as z:
        z.writestr("xl/styles.xml", "<x/>")
    old = tmp / "fix_old.xls"
    old.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 600)
    out.append(("ブックの部品: 読めていないものを、黙って「無い」にしない（部品の無いブック／旧形式）。マクロが無ければ、何も出さない",
                any("ブックの部品を読めていない" in x for x in xv.unread_of(str(broken), [1])) and any("旧形式" in x for x in xv.unread_of(str(old), [1]))
                and xv.unread_of(str(old), []) == [] and xv.buttons_of(str(old)) == [] and xv.is_encrypted(str(book)) is False))

    def cmg(data, seed):                                             # 保護の状態（PROJECT の CMG）を、仕様の式で作る
        key = 0x7B
        raw, un1, enc1, enc2 = [seed, seed ^ 2, seed ^ key], key, seed ^ key, seed ^ 2
        for x in [7] * ((seed & 6) // 2) + list(len(data).to_bytes(4, "little")) + list(data):
            b = x ^ ((enc2 + un1) & 0xFF)
            raw.append(b)
            enc2, enc1, un1 = enc1, b, x
        return bytes(raw).hex().upper()
    out.append(("閲覧のロック: 保護の印を解いて見る（ロックあり・なし。種が違っても同じ）。解けなければ「分からない」",
                [xv.cmg_locked(cmg(bytes([v, 0, 0, 0]), sd)) for v, sd in ((4, 0x41), (0, 0x41), (5, 0x46), (1, 0x13))] == [True, False, True, False]
                and xv.cmg_locked("zz") is None and xv.cmg_locked(cmg(b"\x04\x00", 0x41)) is None and xv.cmg_locked("") is None and xv.is_locked(str(book)) is None))

    sheets = {"Sheet1": "メニュー"}
    mods = [{"name": "Module1", "kind": "standard", "lines": FIX_MODULE.splitlines()}, {"name": "Sheet1", "kind": "sheet", "lines": FIX_SHEET.splitlines(), "sheet": "メニュー"},
            {"name": "cls明細", "kind": "class", "lines": FIX_CLASS.splitlines()},
            {"name": "DieseArbeitsmappe", "kind": cs.module_kind("DieseArbeitsmappe", ".cls", sheets, "DieseArbeitsmappe"), "lines": ["Private Sub Workbook_Open()", "    Module1.取込", "End Sub"]}]
    st = cs.build_struct("M01", "vba", mods, sheets=sheets, buttons=btn[:3], ribbon=xv.ribbon_of(str(book)))
    procs = {x["id"]: x for x in st["procedures"]}
    ent = {e["label"]: (e["trigger"], e["procedure"]) for e in st["entries"]}
    out.append(("VBA: ブックのモジュールは、言語で名前が違っても見分ける。名前だけでは決めない",
                mods[3]["kind"] == "book" and cs.module_kind("DieseArbeitsmappe", ".cls", sheets, "") == "class" and ent.get("ブックを開いたとき") == ("auto", "DieseArbeitsmappe.Workbook_Open")))
    main = procs["Module1.印刷する"]
    kinds = [n["kind"] for n in main["body"]]
    out.append(("VBA: 1行に文が幾つもある書き方（: 区切り）と、日付の定数（#…#）。行は、先に出てくるノードが1回だけ持つ",
                cs.tree_errors(st) == [] and kinds[:2] == ["seq", "if"] and main["body"][0]["exec"] == [6] and not main["body"][1].get("exec")
                and main["body"][1]["cond"] == "期限 < Date" and main["exec"] == [6, 7, 9, 10, 11, 12, 14, 16]))
    cc = main["body"][5]
    out.append(("VBA: 条件つきコンパイル（#If … #Else … #End If）は分岐。中の Exit を、無条件のものにしない",
                kinds == ["seq", "if", "call", "call", "call", "if", "seq"] and cc["cond"] == "コンパイルの条件: VBA7" and cc["lines"] == [11, 15]
                and [(b["label"], [n["kind"] for n in b["body"]]) for b in cc["branches"]] == [("then", ["call"]), ("else", ["interact", "exit"])]))
    out.append(("VBA: 呼び出しを追う（シート名で呼ぶシートの手続き・Me.・オブジェクトの変数・New でクラスの初期化）",
                [c["to"] for c in main["calls"]] == ["Sheet1.更新する", "cls明細.Class_Initialize", "cls明細.読み込む", "Module1.取込"]
                and [c["to"] for c in procs["Sheet1.更新する"]["calls"]] == ["Sheet1.再計算"]))
    out.append(("VBA: 同じ名前の Property（Get・Let）は、別の手続き。クラスの初期化は、起動点にしない",
                {"cls明細.件数", "cls明細.件数（Let）"} <= set(procs) and st["unreached"] == ["cls明細.件数", "cls明細.件数（Let）"]
                and not any(v[1].startswith("cls明細.") for v in ent.values())))
    out.append(("起動点: 同じマクロの図形は1つにまとめる。リボン（読み込み時・ボタン）。式で組み立てた OnAction。シートのイベント",
                ent.get("シート「メニュー」の「取込む」ボタン（ほか 1 個の図形にも割り当て）") == ("button", "Module1.取込")
                and ent.get("ブックを開いたとき（リボンの読み込み）") == ("auto", "Module1.RibbonLoad") and ent.get("リボンの「請求書を作る」") == ("button", "Module1.請求書作成")
                and ent.get("コードで割り当てたボタン・メニュー「再集計」（OnAction）") == ("button", "Sheet1.再計算")
                and ent.get("シート「メニュー」: ほかのシートに移ったとき") == ("event", "Sheet1.Worksheet_Deactivate")))
    out.append(("起動点: ショートカット（大文字は Shift も。キーの所が空白なら、割り当て無し）。割り当ての無い公開の Sub は「手で実行」",
                ent.get("ショートカット Ctrl+Shift+P") == ("shortcut", "Module1.印刷する") and ent.get("ショートカット Ctrl+M") == ("shortcut", "Module1.月次集計")
                and sum(1 for v in ent.values() if v[0] == "shortcut") == 2 and ent.get("マクロ「手で動かす」を実行（ボタンの割り当ては見つからない）") == ("manual", "Module1.手で動かす")
                and len(ent) == 10))
    st2 = cs.build_struct("M01", "vba", mods[:1], unread=["ボタン・図形へのマクロの割り当て"])
    out.append(("起動点: ボタンの割り当てを読めていない形式では、「見つからない」と言い切らない",
                [e["label"] for e in st2["entries"] if e["trigger"] == "manual"] == ["マクロ「手で動かす」を実行（ボタンの割り当ては読めていない）"]))
    pg = lambda k: {"page": 1, "kind": k, "json": "x"}
    man = {"docs": [{"doc_id": "M01", "tool": True, "page_list": [dict(pg("sheet"), page=1), dict(pg("sheet"), page=2)]}, {"doc_id": "M02", "page_list": [pg("sheet")]}]}
    out.append(("道具のブック: シートはカードにしない（出典には引ける）", [d["doc_id"] for d, _p in common.all_pages(man)] == ["M02"]
                and [d["doc_id"] for d, _p in common.all_pages(man, code=True)] == ["M01", "M01", "M02"]))

    src, work = tmp / "vcs_in", tmp / "vcs_work"
    subprocess.run([sys.executable, str(SKILL / "evals" / "make_fixture.py"), str(src), "--vcs"], capture_output=True, text=True)
    r = subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(work)], capture_output=True, text=True)
    run("render_pages.py", work)
    man = json.loads((work / "manifest.json").read_text(encoding="utf-8"))
    by = {d["rel_path"].split("/", 1)[-1]: d for d in man["docs"]}
    out.append(("アドインの書き出し: 棚卸し・ページ化（下のフォルダも1つの書き出し。画像・テーマ・行の中身・設定は数えない。Access のファイルは開かず、書き出しを読む）",
                len(man["docs"]) == 13 and "読まなかったファイル: 6 件" in r.stdout and not (work / "direct").exists()
                and by["受注管理.accdb"].get("read_via") == "受注管理.accdb.src" and not by["受注管理.accdb"].get("read_direct") and by["受注管理.accdb"]["page_list"] == []
                and by["forms/メニュー.cls"]["access"] == {"object": "form", "name": "メニュー", "root": "受注管理.accdb.src", "part": "code"}
                and by["queries/Q_月次%2F売上_追加.sql"]["access"]["name"] == "Q_月次/売上_追加" and by["tbldefs/L_得意先.json"]["code"]["secrets"] == 1
                and all(d.get("rendered") for d in man["docs"] if d["name"] != "受注管理.accdb")))
    run("code_struct.py", work)
    sts = cs.load_structs(work)
    st = next(iter(sts.values()))
    procs = {x["id"]: x for x in st["procedures"]}
    ent = {e["label"]: e["procedure"] for e in st["entries"]}
    out.append(("アドインの書き出し: まとまりは1つ（フォームの定義と、隣のファイルのコードをつなぐ）。起動点は、定義で割り当てたものと AutoExec・最初に開くフォーム",
                len(sts) == 1 and st["tier"] == "parsed" and cs.tree_errors(st) == [] and len(st["procedures"]) == 10
                and ent == {"フォーム「メニュー」を開いたとき（データベースを開くと、最初に開く）": "Form_メニュー.Form_Load", "フォーム「メニュー」の「月次集計」ボタンをクリック": "Form_メニュー.cmd月次_Click",
                            "フォーム「メニュー」の「受注データ取込」ボタンをクリック": "月次取込.（マクロ）", "データベースを開いたとき（AutoExec マクロ）": "AutoExec.（マクロ）"}
                and st["unreached"] == ["Form_メニュー.cmd削除_Click", "modCommon.今月の初日", "cls明細.件数（Let）"]
                and st.get("referenced") == [{"procedure": "modCommon.今月の初日", "by": "フォーム「メニュー」の式"}]))
    mac = procs["AutoExec.（マクロ）"]
    out.append(("アドインの書き出し: マクロは、XML の定義（_AXL）のとおりに読む（条件のブロック。Access が足した行は数えない）",
                mac.get("read") == "axl" and [n["kind"] for n in mac["body"]] == ["if"]
                and [(b["label"], [n["kind"] for n in b["body"]]) for b in mac["body"][0]["branches"]] == [("then", ["call"]), ("else", ["interact"])]
                and [c["to"] for c in mac["calls"]] == ["modCommon.起動する"] and not procs["月次取込.（マクロ）"].get("read")))
    acc = st["access"]
    out.append(("アドインの書き出し: クエリ（名前の記号を戻す）・テーブルの定義・リンク テーブル（パスワードは伏せる）・元のファイルの名前",
                acc["source"] == "受注管理.accdb" and acc["startup_form"] == "メニュー" and acc["unread"] == [] and acc["failed"] == []
                and acc["queries"] == [{"name": "Q_月次/売上_追加", "type": "append", "lines": [1, 4], "writes": ["T_月次売上"], "reads": ["T_受注", "L_得意先"]}]
                and [t["name"] for t in acc["tables"]] == sorted(t["name"] for t in acc["tables"]) and len(acc["tables"]) == 3
                and next(t for t in acc["tables"] if t.get("linked"))["connect"] == "ODBC;DSN=基幹;PWD=" + cs.MASK
                and "himitsu99" not in "".join(f.read_text(encoding="utf-8") for f in (work / "pages").rglob("*.json"))))
    r1 = subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(work), "--tool", "M99"], capture_output=True, text=True)
    same = json.loads((work / "manifest.json").read_text(encoding="utf-8")) == man
    r2 = subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(work), "--tool", "M02"], capture_output=True, text=True)
    on = [d["doc_id"] for d in json.loads((work / "manifest.json").read_text(encoding="utf-8"))["docs"] if d.get("tool")]
    subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(work), "--untool", "M02"], capture_output=True, text=True)
    off = [d["doc_id"] for d in json.loads((work / "manifest.json").read_text(encoding="utf-8"))["docs"] if d.get("tool")]
    out.append(("道具のブックの印（--tool・--untool）: 無い資料 ID なら、何も書き換えずに止まる。付けた印は、取り消せる",
                r1.returncode != 0 and "見つからない" in r1.stderr + r1.stdout and same and r2.returncode == 0 and on == ["M02"] and off == []))
    return out


VBA_EXPR = """Sub Main()
    MsgBox 件数() & " 件"
    Sheets("一覧").Range("A1").Value = 合計(1)
    Application.CommandBars("x").ShowPopup 位置(1), 位置(2)
    x = 税額(100)
    y = 税額(1) + 合計(2)
    Call 保存する
    Sheets("メニュー").更新する
End Sub
Function 件数() As Long
    件数 = Range("A1").Value
End Function
Function 合計(n As Long) As Long
    合計 = Range("B1").Value + n
End Function
Function 位置(n As Long) As Long
    位置 = n * 15
End Function
Function 税額(n As Long) As Long
    税額 = n * 0.1
End Function
Sub 保存する()
    ThisWorkbook.Save
End Sub
Sub 記録()
    Range("単価").Copy Worksheets("作業").Range("a2")
    Sheets("作業").Visible = True
    Sheets("作業").Select
    Range("A2:A1000").Select
    Selection.ClearContents
    Range("B1").Value = 1
    If Range("C1").Value = 1 Then
        Range("D1").Value = 2
    End If
    Range("E1").Value = 3
    Sheets("月次").Activate
    ActiveSheet.Range("A1").Copy
    Sheets("x").Range("A1").PasteSpecial
End Sub
"""

VBA_COPY = """Private m状態 As String
Private Sub txt数量_BeforeUpdate(Cancel As Integer)
    Dim n As Long, strName As String
    n = 0
    blnFound = True
    m = n
    Cancel = True
    m状態 = "済"
    時刻 = Now
    x = n + 1
    y = Me.txt数量
    z = ""
    n = n + 1
    合計 = 合計 + n
End Sub
Function 探す(ByVal key As String, Optional ByRef found As Boolean = False) As Object
    Set 探す = Nothing
    found = False
    Set 探す = ctl
End Function
"""


def vba_stmt_tests():
    """VBA の文の読み方のうち、実物で見つけた穴: 式の中の呼び出しで文そのものが消える／写す先の向き／シートを選んだあとの「表示中のシート」／条件の中の読み書きの向き"""
    import code_struct as cs
    out = []
    at = lambda t: next(k for k, ln in enumerate(VBA_EXPR.splitlines(), 1) if t in ln)
    st = cs.build_struct("M09", "vba", [{"name": "Module1", "kind": "standard", "lines": VBA_EXPR.splitlines()},
                                        {"name": "Sheet1", "kind": "sheet", "lines": ["Public Sub 更新する()", '    Range("A1").Value = 1', "End Sub"], "sheet": "メニュー"}], sheets={"Sheet1": "メニュー"})
    procs = {p["name"]: p for p in st["procedures"]}
    body = [(n["kind"], n["lines"][0], n.get("exec") or [], (n.get("to") or "").split(".")[-1]) for n in procs["Main"]["body"]]
    a, b, c, d, e = at("MsgBox 件数"), at("= 合計(1)"), at("ShowPopup"), at("x = 税額"), at("y = 税額")
    out.append(("VBA: 式の中で関数を呼ぶ文（メッセージ・シートに書く・メニューを出す）は、文そのものが箱に残る。呼び出しは、その前に置く（行は、文のノードが持つ）",
                cs.tree_errors(st) == [] and body[:6] == [("call", a, [], "件数"), ("interact", a, [a], ""), ("call", b, [], "合計"), ("seq", b, [b], ""), ("call", c, [], "位置"), ("seq", c, [c], "")]
                and [(i["name"], i["way"]) for i in procs["Main"]["body"][3]["io"]] == [("一覧", "out")]))
    out.append(("VBA: 文そのものが呼び出しのとき（Call・変数に結果を入れる・シート名で手続きを呼ぶ）は、呼び出しのノード1つ。変数に入れる式が続く（F(1) + G(2)）ときは、最後の呼び出しが文を持つ",
                body[6:] == [("call", d, [d], "税額"), ("call", e, [], "税額"), ("call", e, [e], "合計"), ("call", e + 1, [e + 1], "保存する"), ("call", e + 2, [e + 2], "更新する")]
                and not procs["Main"]["body"][-1].get("io")))
    import code_lang
    called = {n["to"] for p in st["procedures"] for n in code_lang.walk(p["body"]) if n["kind"] == "call"}
    out.append(("VBA: 呼ぶ側の図に入れる手続きには、どれも、入れる所（呼び出しのノード）がある（図から落ちない）", st["inline"] != [] and set(st["inline"]) <= called and st["unreached"] == []))
    rec = procs["記録"]["body"]
    io = lambda n: [(i["name"], i["way"]) for i in n.get("io", [])]
    r0 = at('Range("単価").Copy')
    out.append(("VBA: 写す先を引数に書く Copy は、先が「書く」・元が「読む」。表示・非表示の切り替えは、読み書きに数えないが、別のシートを出す・選ぶ文で、まとまりを切る",
                rec[0]["lines"] == [r0, r0] and io(rec[0]) == [(cs.ACTIVE, "in"), ("作業", "out")]))
    out.append(("VBA: シートを選んだあとの Selection・シート名の無い Range は、そのシート。分岐をまたぐと、分からない（表示中のシート）に戻す",
                io(rec[1]) == [("作業", "out")] and rec[1]["lines"] == [r0 + 1, r0 + 5] and io(rec[3]) == [(cs.ACTIVE, "out")]
                and io(rec[4]) == [("月次", "in")] and io(rec[5]) == [("x", "out")]))
    cp = cs.build_struct("M09", "vba", [{"name": "Form_受注", "kind": "form", "lines": VBA_COPY.splitlines()}])
    pc = {p["name"]: [(n["lines"], n.get("tech", "")) for n in p["body"]] for p in cp["procedures"]}
    out.append(("VBA: 手続きの中だけの変数に、定数か、ほかの変数の値をそのまま入れるだけの文・今の日時を取るだけの文は「変数の準備」。"
                "引数（呼んだ側に返す値）・戻り値・モジュールの変数・関数かもしれない名前・計算・数を足す文・項目の値を入れる文は、残す",
                pc["txt数量_BeforeUpdate"] == [([4, 6], "変数の準備"), ([7, 8], ""), ([9, 9], "変数の準備"), ([10, 11], ""), ([12, 12], "変数の準備"), ([13, 14], "")]
                and pc["探す"] == [([17, 19], "")]))
    out.append(("VBA: 条件の式の = は、比べているだけ（読む）。中の文は「書く」", rec[2]["kind"] == "if" and io(rec[2]) == [(cs.ACTIVE, "in")] and io(rec[2]["branches"][0]["body"][0]) == [(cs.ACTIVE, "out")]))
    fm = cs.build_struct("M09", "vba", [{"name": "Module1", "kind": "standard", "lines": VBA_FMT.splitlines()}])
    tech, ios = {}, {}
    for n in code_lang.walk(fm["procedures"][0]["body"]):
        if n["kind"] == "seq":
            for ln in range(n["lines"][0], n["lines"][1] + 1):
                tech[ln], ios[ln] = n.get("tech", ""), [(i["name"], i["way"]) for i in n.get("io") or []]
    out.append(("VBA: 書式だけの文（字・塗り・列の幅）は、シートへの書き込みにせず「書式・表示の調整」。最後の行を取るだけ・行の位置を 1 つ進めるだけの文は「変数の準備」"
                "（機械が対象外にする。決まり 6.10）。表示する件数を数える文は、残す",
                tech.get(5) == tech.get(6) == "書式・表示の調整" and ios.get(5) == [] and tech.get(4) == "変数の準備" and tech.get(11) == "変数の準備"
                and tech.get(10) == "" and tech.get(9) == "" and ("一覧", "out") in ios.get(9, [])))
    return out


VBA_FMT = """Public Sub 整える()
    Dim ws As Worksheet, r As Long, 件数 As Long, 最終行 As Long
    Set ws = Worksheets("一覧")
    最終行 = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row
    ws.Range("A1:F1").Font.Bold = True
    ws.Columns("A:F").AutoFit
    r = 2
    Do While r <= 最終行
        ws.Cells(r, 3).Value = ws.Cells(r, 1).Value * 2
        件数 = 件数 + 1
        r = r + 1
    Loop
    MsgBox 件数 & " 件"
End Sub
"""


VBA_LINK = """Public Sub つなぎ直す()
    Dim 件数 As Long
    CurrentDb.TableDefs("取込データ").Connect = "Text;DSN=取込の定義;FMT=Delimited;HDR=NO;DATABASE=" & CurrentProject.Path
    CurrentDb.TableDefs("取込データ").RefreshLink
    Set db = OpenDatabase("", False, False, "Text;DSN=取込の定義;FMT=Delimited;DATABASE=C:\\data\\売上")
    cn.Open "ODBC;DSN=販売DB;UID=user"
    件数 = DCount("*", "[月次の集計]")
    件数 = DCount("*", "受注")
End Sub
"""


def io_fix_tests():
    """VBA の読む・書くのうち、照合で見つかった取り違え: リンクの接続先を決める文／Access のリンクの文字の DSN／保存してあるクエリを相手にした DLookup・DCount"""
    import code_struct as cs
    st = cs.build_struct("M09", "vba", [{"name": "Module1", "kind": "standard", "lines": VBA_LINK.splitlines()},
                                        {"name": "クエリ 月次の集計", "kind": "query", "lines": ["SELECT 受注.金額 FROM 受注;"]}])
    io = [(i["kind"], i["name"], i["way"]) for n in st["procedures"][0]["body"] for i in n.get("io") or []]
    return [
        ("VBA: リンクしたテーブルの接続先を決めるだけの文（.Connect ＝ 接続の文字）は、データベースを読んだことにしない", not any(nm == "取込の定義" for _k, nm, _w in io)
         and sum(k == "データベース" for k, _n, _w in io) == 2),
        ("VBA: Access のリンクの文字（Text;DSN=…;DATABASE=…）の DSN は、取り込みの定義の名前。読む先は DATABASE。ODBC の DSN は、読む先のまま",
         ("データベース", "売上", "in") in io and ("データベース", "販売DB", "in") in io),
        ("VBA: DLookup・DCount の相手が、保存してあるクエリなら、クエリと呼ぶ（そのクエリが読むテーブルも付く）。テーブルなら、テーブル",
         ("クエリ", "月次の集計", "in") in io and ("テーブル", "月次の集計", "in") not in io and io.count(("テーブル", "受注", "in")) >= 1),
    ]


def code_tests():
    """プログラムのソースの部品: 認証情報を伏せる・ページを手続きの切れ目で切る・構造を取り出す"""
    import code_struct as cs
    out = []
    lines = CODE_SAMPLE.splitlines()
    masked, hit = cs.mask_secrets(lines)
    out.append(("コード: 接続文字列のパスワードを伏せる（行の数は変えない）", len(masked) == len(lines) and len(hit) == 1 and "himitsu99" not in "\n".join(masked)
                and "Password=" + cs.MASK in masked[hit[0] - 1]))
    out.append(("コード: 伏せたあとは、もう伏せる所が無い", cs.mask_secrets(masked)[1] == []))
    out.append(("コード: 変数に入れた文字のパスワード・トークンも伏せる",
                cs.mask_secrets(['strPassword = "abc"', 'api_key: "K123456"', "x = 1", 'pw = InputBox("パスワード")'])[1] == [1, 2]))
    long = ["Option Explicit"] + [ln for k in range(6) for ln in [f"Sub P{k}()"] + ["    x = 1"] * 120 + ["End Sub", ""]]
    cuts = cs.page_cuts("vba", long, 300)
    out.append(("コード: 300 行を超えるソースを、手続きの頭で切る", len(cuts) >= 3 and all(long[c - 1].startswith("Sub ") for c in cuts[1:])))
    out.append(("コード: 1つの手続きが 300 行を超えるときだけ、途中で切る", cs.page_cuts("vba", ["Sub A()"] + ["x = 1"] * 700 + ["End Sub"], 300) == [1, 301, 601]))
    st = cs.build_struct("M09", "vba", [{"name": "集計", "kind": "standard", "lines": masked}],
                         buttons=[{"sheet": "メニュー", "shape": "集計", "macro": "Main"}])
    procs = {x["name"]: x for x in st["procedures"]}
    ent = {e["procedure"]: e for e in st["entries"]}
    out.append(("コード: 起動点は、ボタンのマクロと、引数の無い公開の Sub（呼ばれていないもの）", set(ent) == {"集計.Main", "集計.取込"}
                and ent["集計.Main"]["trigger"] == "button" and "「集計」ボタン" in ent["集計.Main"]["label"] and ent["集計.取込"]["trigger"] == "manual"))
    kinds = [n["kind"] for n in procs["Main"]["body"]]
    out.append(("コード: 木の並び（エラーの設定・まとまり・分岐・繰り返し・呼び出し・抜ける）",
                kinds == ["onerror", "seq", "seq", "if", "if", "seq", "loop", "call", "exit"]
                and [n.get("tech") for n in procs["Main"]["body"] if n["kind"] == "seq"] == ["画面の更新の制御", None, "変数の準備"]))
    out.append(("コード: エラー処理の節を、本体から分ける", [h["label"] for h in procs["Main"]["handlers"]] == ["ErrHandler"]
                and procs["Main"]["handlers"][0]["body"][0]["kind"] == "interact"))
    ask = procs["Main"]["body"][4]
    out.append(("コード: 1行の If と、条件の中の MsgBox（操作者に聞く）", ask.get("single") and ask["interact"]["how"] == "msgbox" and ask["branches"][0]["body"][0]["kind"] == "exit"))
    io = {(x["kind"], x["name"], x["way"]) for p in st["procedures"] for x in p["io"]}
    out.append(("コード: 入出力（シートの読み書き・Set で束ねた変数・組み立てたファイル名・SQL のテーブル・接続先）",
                {("シート", "メニュー", "in"), ("シート", "一覧", "out"), ("シート", "一覧", "in"), ("ファイル", "一覧_＜月＞.pdf", "out"),
                 ("テーブル", "入金", "out"), ("データベース", "hanbai", "in")} <= io))
    out.append(("コード: 入出力の種類は、決まった語（common の io_kind。レポート・印刷は帳票）", {k for k, _n, _w in io} <= set(cs.IO_KINDS)))
    out.append(("コード: 2か所から呼ばれる手続きは別の単位、1か所だけの小さな手続きは呼ぶ側に入れる",
                [u["procedure"] for u in st["units"] if u["kind"] == "shared"] == ["集計.税額"] and st["inline"] == ["集計.保存する"] and st["unreached"] == []))
    out.append(("コード: 実行文の行が、どれも木のどこかに1回だけ入る", cs.tree_errors(st) == [] and procs["Main"]["exec"] == [8, 9, 10, 11, 12, 13, 15, 16, 17, 18, 20, 21, 23]))
    bad = json.loads(json.dumps(st))
    next(x for x in bad["procedures"] if x["name"] == "税額")["body"][0]["branches"][0]["body"] = []
    out.append(("コード: 木からノードを抜くと、実行文の抜けとして出る", any("どのノードにも入っていない" in e for e in cs.tree_errors(bad))))
    return out


VBA_CHECKS = """Public Sub 点検する()
    Dim 値 As String
    値 = Worksheets("台帳").Range("A1").Value
    結果を書く "見出し", (値 = "番号")
    値 = Worksheets("台帳").Range("B1").Value
    結果を書く "日付", (値 <> "")
    結果を書く "シートの数", (Worksheets.Count > 1)
    MsgBox "点検が終わりました"
End Sub

Private Sub 結果を書く(項目 As String, 合格 As Boolean)
    Worksheets("結果").Range("A1").Value = 項目
    Worksheets("結果").Range("B1").Value = 合格
End Sub

Public Sub 写す()
    Worksheets("控え").Range("A1").Value = Worksheets("台帳").Range("A1").Value
    Worksheets("控え").Range("B1").Value = Worksheets("台帳").Range("B1").Value
End Sub
"""


def call_run_tests(tmp):
    """同じ共通の処理を続けて呼ぶ並びを、1つの箱にまとめる（直しの一覧の merge に、呼び出しを入れる）"""
    out, src, w = [], tmp / "run_in", tmp / "run_work"
    (src / "src").mkdir(parents=True)
    (src / "src" / "点検.bas").write_text('Attribute VB_Name = "点検"\n' + VBA_CHECKS, encoding="utf-8")
    subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(w)], capture_output=True, text=True)
    for sc, args in (("render_pages.py", []), ("code_struct.py", []), ("systems.py", ["draft"])):
        run(sc, w, *args)
    edit(w / "structure/systems.json", lambda j: [x.update(kind="euc_excel") for x in j["systems"]])
    _c, listing = run("activity.py", w, "list")
    ids = {name: x.split()[0] for x in listing.splitlines() for name in ("点検する", "結果を書く", "写す") if f"点検.{name}" in x}
    main, sub, copy_id = ids.get("点検する", "A-01-2"), ids.get("結果を書く", "A-01-1"), ids.get("写す", "A-01-3")
    (w / "activities").mkdir(exist_ok=True)
    put = lambda aid, name, ops: (w / f"activities/{aid}.edits.json").write_text(json.dumps(
        {"activity": aid, "name": name, "summary": "", "questions": [], "ops": ops}, ensure_ascii=False), encoding="utf-8")
    put(sub, "点検の結果を書く", [{"op": "name", "item": "結果を書く.13", "label": "「結果」に、項目と合否を書く"}])
    put(copy_id, "台帳を控えに写す", [{"op": "split", "item": "写す.18", "at": [19]}, {"op": "name", "item": "写す.18", "label": "「台帳」の「A1」を「控え」に写す"},
                                 {"op": "name", "item": "写す.19", "label": "「台帳」の「B1」を「控え」に写す"}])
    keys = ["点検する.4", "点検する.5", "点検する.6", "点検する.7", "点検する.8"]
    put(main, "台帳を点検する", [{"op": "merge", "items": keys}, {"op": "name", "item": keys[0], "label": "「台帳」を 3 項目確かめ、結果を書く"}])
    code, text = run("activity.py", w, "apply", "--all")
    act = json.loads((w / f"activities/{main}.activity.json").read_text(encoding="utf-8")) if (w / f"activities/{main}.activity.json").exists() else {"nodes": [], "lines": []}
    box = next((n for n in act["nodes"] if n.get("calls")), {})
    run("build_activity.py", w)
    code2, text2 = run("check_flow.py", w, "--stage", "activity")
    xml = "".join(p.read_text(encoding="utf-8") for p in (w / "out").glob("アクティビティ図*.drawio"))
    out.append(("直しの一覧: 同じ共通の処理を 3 回続けて呼ぶ並びを、1つの箱にまとめる（箱に、呼ぶ先の図と回数が出る。行の網羅は保たれる）" + ("" if code == 0 and code2 == 0 else text[-300:] + text2[-400:]),
                code == 0 and box.get("type") == "process" and box.get("calls") == [{"to": sub, "times": 3, "label": "点検の結果を書く"}]
                and sum(n["type"] == "subprocess" for n in act["nodes"]) == 0 and code2 == 0 and "0 FAIL" in text2 and f"別ページ {sub} 点検の結果を書く（3 回）" in xml))
    code, text = 0, ""
    act = json.loads((w / f"activities/{copy_id}.activity.json").read_text(encoding="utf-8")) if (w / f"activities/{copy_id}.activity.json").exists() else {"nodes": []}
    second = next((n["id"] for n in act["nodes"] if n.get("label") == "「台帳」の「B1」を「控え」に写す"), None)
    out.append(("直しの一覧: まとまりを分けても、同じシートをほかの行でも読む・書く部分に、読む・書くが付く（はじめの行の側だけに付けない）" + ("" if code == 0 else text[-300:]),
                code == 0 and sorted((f["label"], f.get("direction")) for f in act["nodes"] if f.get("attach_to") == second) == [("シート 台帳", "in"), ("シート 控え", "out")]))
    put(main, "台帳を点検する", [{"op": "merge", "items": keys[:2]}, {"op": "name", "item": keys[0], "label": "「台帳」の見出しを確かめ、結果を書く"}])
    code, text = run("activity.py", w, "apply", "--id", main)
    out.append(("直しの一覧: 呼び出しが 3 回に満たない並びは、まとめられない", code != 0 and "3 回以上続けて呼ぶ並びだけ" in text))
    return out


def code_pipeline_tests(tmp):
    """プログラムのソースを含む入力を、通しで: ページ化 → 構造 → 台帳 → アクティビティ図 → フロー → 図・一覧 → 点検。そのあと、誤りを1つずつ仕込む"""
    out, src, good = [], tmp / "code_in", tmp / "code_good"
    subprocess.run([sys.executable, str(SKILL / "evals" / "make_fixture.py"), str(src), "--code"], capture_output=True, text=True)
    subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(good)], capture_output=True, text=True)
    run("render_pages.py", good)
    shutil.copytree(SKILL / "evals" / "fixture_code_expected", good, dirs_exist_ok=True)
    log = []
    for script, args in [("code_struct.py", []), ("quotes.py", ["fill"]), ("activity.py", ["apply", "--all"]), ("apply_edits.py", ["--flow", "F01"]), ("build_drawio.py", []),
                         ("build_matrix.py", []), ("build_activity.py", []), ("check_extract.py", []), ("check_flow.py", []), ("check_layout.py", ["--no-png"])]:
        code, text = run(script, good, *args)
        passed = code == 0 and "0 FAIL" in text if script.startswith("check") else code == 0
        log.append((script, passed, text))
    bad = [(sc, text) for sc, passed, text in log if not passed]
    out.append(("プログラムを含む入力: 通しで作って、点検が FAIL 0" + ("" if not bad else f"（{bad[0][0]}）\n{bad[0][1][-700:]}"), not bad))
    fl = json.loads((good / "flows/F01.flow.json").read_text(encoding="utf-8"))
    boxes = [(n["label"], n["activity"], n["sources"][0]["item"]) for n in fl["nodes"] if n.get("activity")]
    out.append(("起動点の箱: 紐付けた手順のすぐあとに、アクティビティ図の名前と番号で置かれる",
                boxes == [("月次の請求を集計する", "A-01-2", "e2"), ("入金を消し込む", "A-01-3", "e3")]
                and [l["kind"] for l in fl["lanes"]][-1] == "system" and any(q["id"] == "m-sys-M01-e3" and q["category"] == "システム" for q in fl["questions"])))
    act = json.loads((good / "activities/A-01-2.activity.json").read_text(encoding="utf-8"))
    lane = {l["id"]: l["name"] for l in act["lanes"]}
    kinds = [n["type"] for n in act["nodes"] if n["type"] != "file"]
    out.append(("アクティビティ図: レーン（操作者＝紐付けた手順の担当／プログラム／外部）・判断 6・呼び出し 1・外部 1・終了 2 種類",
                list(lane.values()) == ["経理担当", "月次集計", "外部"] and kinds.count("decision") == 6 and kinds.count("subprocess") == 1 and kinds.count("system") == 1
                and sorted({n["label"] for n in act["nodes"] if n["type"] == "end"}) == ["終了", "終了（途中で終わる）"]
                and next(n for n in act["nodes"] if n["type"] == "subprocess")["link_flow"] == "A-01-5"))
    owned = sorted(ln for o in act["lines"] for ln in o["lines"])
    out.append(("アクティビティ図: 実行文 33 行が、箱か対象外に1回ずつ入る（呼ぶ側に入れた手続きの行を含む）", len(owned) == 33 and len(set(owned)) == 33
                and sum(len(o["lines"]) for o in act["lines"] if o["state"] == "excluded") == 4))
    qs = json.loads((good / "questions.json").read_text(encoding="utf-8"))["questions"]
    local = {q["local"] for q in qs if q["origin"] == "systems"}
    out.append(("機械が作る確認事項: 業務フローに出てこない起動点・ボタンの割り当てが見つからないマクロ（同じ起動点の 2 つは 1 つに）",
                {"unlinked:M01.e1", "unlinked:M01.e4", "manual:M01.e2", "manual:M01.e3"} <= local and "unlinked:M01.e2" not in local and "manual:M01.e4" not in local))
    from build_matrix import model
    mx = model(good)
    out.append(("業務マトリクス一覧: システム一覧・起動点一覧・プログラムのノード一覧に、番号と場所が出る（システムの番号は数だけ。SYS- は内部の番号）",
                any(r[0] == "01" and "F01 S-001（▶ A-01-2）" in r[7] for r in mx["システム一覧"]) and any(r[0] == "2" and r[4] == "A-01-2" and r[5] == "F01 S-001" for r in mx["起動点一覧"])
                and any(r[0] == "A-01-2" and r[3] == "対象外" for r in mx["プログラムのノード一覧"]) and any("▶ A-01-2" in c for c in mx["業務マトリクス"][1])))
    ends = [n["id"] for n in act["nodes"] if n["type"] == "end" and n["id"] != "end"]
    out.append(("アクティビティ図: 途中で終わる所は、抜ける所ごとに終了の丸（1 つの丸へ線を集めない）",
                len(ends) >= 2 and all(sum(e["to"] == i for e in act["edges"]) == 1 for i in ends) and sum(e["to"] == "end" for e in act["edges"]) >= 1))
    out.append(("用語の対応: 一覧のシートに、もとの名前・業務の言葉・根拠・定義の場所が出る。つづりから読んだ言葉は、推定として一覧に出る（名前だけの問いなので、確認事項にはしない）",
                any(r[1] == "売上フォルダ" and r[2] == "売上データの置き場所" and r[3] == "定義" and "4 行" in r[4] for r in mx["プログラムの用語対応"])
                and any(r[1] == "最終行" and "推定" in r[3] for r in mx["プログラムの用語対応"])
                and not any(str(q.get("local", "")).startswith("terms-") for q in qs) and all(q.get("effect") for q in qs)))
    run("pages.py", good, "terms")                                   # マニュアルの表記の一覧（structure/terms.json）は、別のファイル
    code, text = run("check_flow.py", good, "--stage", "activity")
    out.append(("用語の対応: マニュアルの表記の一覧（pages.py terms の structure/terms.json）と、別のファイル（structure/code_terms.json）。両方あっても、点検は FAIL 0",
                (good / "structure/terms.json").exists() and (good / "structure/code_terms.json").exists() and code == 0 and "0 FAIL" in text))
    lay = next(pg for lp in sorted((good / "out").glob("アクティビティ図*.layout.json")) for pg in json.loads(lp.read_text(encoding="utf-8"))["pages"] if pg["id"] == "act-A-01-2")
    box, eds = lay["nodes"], [(e["from"], e["to"]) for e in act["edges"]]
    side = [(u, a, b) for u in box for a in box for b in box if (u, a) in eds and (u, b) in eds and a < b and box[a][0] != box[b][0]
            and abs(box[a][1] + box[a][3] / 2 - box[b][1] - box[b][3] / 2) < 1 and {n["id"]: n["lane"] for n in act["nodes"]}[a] == {n["id"]: n["lane"] for n in act["nodes"]}[b]]
    flay = next(pg for lp in sorted((good / "out").glob("*.layout.json")) if "アクティビティ図" not in lp.name for pg in json.loads(lp.read_text(encoding="utf-8"))["pages"] if pg["id"] == "flow-F01")
    cx = {}
    for n in fl["nodes"]:
        if n["type"] != "file" and n["id"] in flay["nodes"]:
            cx.setdefault(n["lane"], set()).add(round(flay["nodes"][n["id"]][0] + flay["nodes"][n["id"]][2] / 2))
    out.append(("配置: アクティビティ図は、分岐の先が、同じレーンの同じ高さに横に並ぶ。業務フローの図は、レーンに 1 列のまま",
                bool(side) and all(len(v) == 1 for v in cx.values())))
    dec, caps = [n["id"] for n in act["nodes"] if n["type"] == "decision"], lay.get("captions") or {}
    shown = "".join(re.findall(r'id="src-[^"]*" value="([^"]*)"', next((good / "out").glob("アクティビティ図*.drawio")).read_text(encoding="utf-8")))
    out.append(("出典: アクティビティ図でも、箱・ひし形の枠の外（すぐ下・真ん中の線の右）に、切らずに出る。ひし形は、番号と条件だけの大きさ（出典で、縦に伸びない）",
                bool(dec) and all(box[i][2] == 150 and box[i][3] <= 110 for i in dec)
                and all(i in caps and caps[i][1] >= box[i][1] + box[i][3] and caps[i][0] > box[i][0] + box[i][2] / 2 for i in dec)
                and "…" not in shown and "「" in shown and "行" in shown and not re.search(r"(?<![A-Za-z0-9])M\d{2}", shown)))
    labs = [n["label"] for n in act["nodes"]]
    out.append(("名前: 機械が付ける名前（メッセージ・外部）も、動詞で終わる（見せる文字は前に、相手の名前はかっこ書き）",
                "メールを送る（Outlook）" in labs and any(x.endswith("と、メッセージを出す") for x in labs) and not any(x.startswith("メッセージを出す「") for x in labs)))
    import build_activity as BA
    acts = BA.load_activities(good)
    ctx, notes = BA.context(good, acts), []
    one = BA.ActivityDrawing(act, ctx, 1).build()
    wide = BA.activity_page(act, ctx, notes.append)
    real, BA.page_problems = BA.page_problems, lambda pg: [("fail", "重なる")]
    try:
        fallen = BA.activity_page(act, ctx, notes.append)
    finally:
        BA.page_problems = real
    out.append(("配置の安全網: 横に並べたページが配置の点検に当たれば、列を減らして描き直し、最後は縦一列に戻す",
                wide.nodes != one.nodes and fallen.nodes == one.nodes and any("描き直す" in x for x in notes)))
    page = json.loads((good / "pages/M02/p0001.json").read_text(encoding="utf-8"))
    out.append(("ソースのページ: 認証情報を伏せ、行の番号がソースの行と合う", "abc12345" not in json.dumps(page, ensure_ascii=False) and page["code"]["secrets"] == [73]
                and page["text"][6].startswith("Public Sub Main")))

    def case(name, script, args, change, expect, before=()):
        w = tmp / f"code_case{len(out)}"
        shutil.copytree(good, w)
        change(w)
        for sc, ar in before:
            run(sc, w, *ar)
        code, text = run(script, w, *args)
        hit = next((ln for ln in text.splitlines() if expect in ln), None)
        out.append((f"{name}: {(hit or '')[:110]}" if hit and code != 0 else f"{name}: 「{expect}」が出ない（終了コード {code}）\n{text[-600:]}", bool(hit) and code != 0))
        shutil.rmtree(w, ignore_errors=True)
    sysj, a2 = "structure/systems.json", "activities/A-01-2"
    case("コード: 伏せ字を外す", "check_flow.py", ["--stage", "code"],
         lambda w: (w / "pages/M02/p0001.json").write_text((w / "pages/M02/p0001.json").read_text(encoding="utf-8").replace("Password=＊＊＊", "Password=abc12345"), encoding="utf-8"),
         "認証情報らしい文字が伏せられていない")
    case("コード: 構造を手で直す", "check_flow.py", ["--stage", "code"], lambda w: edit(w / "code/M01.struct.json", lambda s: s["entries"].pop()), "取り出し直した構造と違う")
    case("台帳: 名指しの文字が、プログラムの側に無い", "check_flow.py", ["--stage", "systems"], lambda w: edit(w / sysj, lambda s: s["links"][0].update(evidence="集計ボタン")),
         "プログラムの側の名前")
    case("台帳: 名指しの文字が、マニュアルの手順に無い", "check_flow.py", ["--stage", "systems"], lambda w: edit(w / sysj, lambda s: s["links"][0].update(step="M03.1.s5")),
         "マニュアルの手順 M03.1.s5 の文・帳票に無い")
    case("台帳: 無い起動点を指す", "check_flow.py", ["--stage", "systems"], lambda w: edit(w / sysj, lambda s: s["excluded_entries"].append({"entry": "M01.e9", "reason": "中身が空"})),
         "という起動点が無い")
    case("台帳: 用語の正本のシステムが、台帳に無い", "check_flow.py", ["--stage", "systems"], lambda w: edit(w / sysj, lambda s: s["systems"].pop()), "台帳に無い")
    case("台帳: 図にしない理由が「中身が空」なのに、実行文がある（決まり 6.4）", "check_flow.py", ["--stage", "systems"],
         lambda w: edit(w / sysj, lambda s: s["excluded_entries"].append({"entry": "M01.e1", "reason": "中身が空"})), "「M01.e1」の理由は「中身が空」なのに")
    case("台帳: 図にしない理由が「ほかの起動点と同じ処理を呼ぶだけ」なのに、呼ばない", "check_flow.py", ["--stage", "systems"],
         lambda w: edit(w / sysj, lambda s: s["excluded_entries"].append({"entry": "M01.e1", "reason": "ほかの起動点と同じ処理を呼ぶだけ"})), "「ほかの起動点と同じ処理を呼ぶだけ」なのに")
    import code_struct as cs
    ui_ok = all(cs.UI_WRITE.search(x) for x in ("Me.lstResult.AddItem 名前", "lblStatus.Caption = \"完了\"", "Me.txtMonth.Value = \"\"")) \
        and not any(cs.UI_WRITE.search(x) for x in ("ws.Range(\"A1\").Value = 1", "Sheets(\"集計\").Cells(1, 1).Value = x"))
    out.append(("補助の候補: 画面の部品に書く文（リストに行を足す・ラベルを変える）のある手続きは、候補に入れない（決まり 6.32）",
                ui_ok and cs.touches({"a": {"calls": [], "io": [], "interactions": [], "dynamic": [], "ui": [5]}}, "a") == ["画面の部品への書き込み"]))
    case("台帳: プログラムの資料の一部が、システムに入っていない", "check_flow.py", ["--stage", "systems"], lambda w: edit(w / sysj, lambda s: s["systems"][0]["docs"].remove("M02")),
         "同じ1つのシステムの docs に入っていない")
    case("フロー: 起動点の箱から、アクティビティ図の番号を消す", "check_flow.py", ["--stage", "flow"],
         lambda w: edit(w / "flows/F01.flow.json", lambda f: next(n for n in f["nodes"] if n.get("activity")).pop("activity")), "プログラムのソース")
    case("フロー: 推定の紐付けの確認事項を消す", "check_flow.py", ["--stage", "flow"],
         lambda w: edit(w / "flows/F01.flow.json", lambda f: f.update(questions=[q for q in f["questions"] if not q["id"].startswith("m-sys-")])), "プログラムの起動点の箱が")
    case("フロー: 紐付けを足したのに、箱を置いていない", "check_flow.py", ["--stage", "flow"],
         lambda w: edit(w / sysj, lambda s: s["links"].append({"entry": "M01.e1", "step": "M03.1.s2", "basis": "inferred"})), "足りない A-01-1")
    case("アクティビティ図: 行の持ち主を1つ消す", "check_flow.py", ["--stage", "activity"], lambda w: edit(w / f"{a2}.activity.json", lambda a: a["lines"].pop(0)), "行の網羅")
    case("アクティビティ図: 分岐の扱いを1つ消す", "check_flow.py", ["--stage", "activity"], lambda w: edit(w / f"{a2}.activity.json", lambda a: a["branches"].pop(0)), "分岐の網羅")
    case("アクティビティ図: 名前を付けない箱がある", "check_flow.py", ["--stage", "activity"],
         lambda w: edit(w / f"{a2}.edits.json", lambda e: e.update(ops=[o for o in e["ops"] if o.get("item") != "Main.15"])), "名前が付いていない", [("activity.py", ["apply", "--all"])])
    case("アクティビティ図: かぎかっこの中に、ソースに無い名前を書く", "check_flow.py", ["--stage", "activity"],
         lambda w: edit(w / f"{a2}.edits.json", lambda e: next(o for o in e["ops"] if o.get("item") == "Main.15").update(label="対象月を「設定」から読み取る")),
         "かぎかっこの中の名前「設定」", [("activity.py", ["apply", "--all"])])
    case("アクティビティ図: 直しの一覧を直したあと、作り直していない", "check_flow.py", ["--stage", "activity"],
         lambda w: edit(w / f"{a2}.edits.json", lambda e: e.update(name="集計する")), "作り直したものと違う")
    case("アクティビティ図: 図を1枚消す", "check_flow.py", ["--stage", "activity"], lambda w: (w / "activities/A-01-5.activity.json").unlink(), "アクティビティ図がまだ無い")
    case("直しの一覧: 決まっていない理由で畳む（スキーマの決まった語に無い）", "activity.py", ["apply", "--id", "A-01-2"],
         lambda w: edit(w / f"{a2}.edits.json", lambda e: e["ops"].append({"op": "fold", "item": "Main.32", "reason": "細かいから"})), "決まった語でない")
    case("直しの一覧: 対象外の理由で畳む（畳む理由の語でない）", "activity.py", ["apply", "--id", "A-01-2"],
         lambda w: edit(w / f"{a2}.edits.json", lambda e: e["ops"].append({"op": "fold", "item": "Main.32", "reason": "ログ"})), "reason は、決まった語から選ぶ")
    case("直しの一覧: 抜ける所（Exit Function）のある分岐を畳む（決まり 6.11）", "activity.py", ["apply", "--id", "A-01-5"],
         lambda w: edit(w / "activities/A-01-5.edits.json", lambda e: e["ops"].append({"op": "fold", "item": "消費税計算.93", "reason": "型や空の確認"})), "抜ける所")
    case("直しの一覧: シートに書く文を対象外にする", "activity.py", ["apply", "--id", "A-01-2"],
         lambda w: edit(w / f"{a2}.edits.json", lambda e: e["ops"].append({"op": "exclude", "item": "Main.33", "reason": "変数の準備"})), "対象外にできない")
    case("直しの一覧: 間に分岐のあるまとまりをまとめる", "activity.py", ["apply", "--id", "A-01-2"],
         lambda w: edit(w / f"{a2}.edits.json", lambda e: e["ops"].append({"op": "merge", "items": ["Main.15", "Main.23"]})), "隣り合う「まとまり」だけ")
    case("直しの一覧: 無い項目を指す", "activity.py", ["apply", "--id", "A-01-2"],
         lambda w: edit(w / f"{a2}.edits.json", lambda e: e["ops"].append({"op": "name", "item": "Main.999", "label": "x"})), "という項目が、下書きに無い")
    case("図: アクティビティ図の箱の名前を手で直す", "check_flow.py", ["--stage", "drawio"],
         lambda w: [p.write_text(p.read_text(encoding="utf-8").replace("対象月が入っているか", "対象月の確認"), encoding="utf-8") for p in (w / "out").glob("アクティビティ図*.drawio")],
         "ラベルが flow.json と違う")
    # 名付け（読む人は、業務部門の担当）: プログラムの中の名前は箱に書かない。言い換えは、用語の対応に
    redo, act_stage, termj = [("activity.py", ["apply", "--all"])], ["--stage", "activity"], "structure/code_terms.json"
    rename = lambda item, **kw: lambda w: edit(w / f"{a2}.edits.json", lambda e: next(o for o in e["ops"] if o["op"] == "name" and o.get("item") == item).update(**kw))
    case("名付け: 箱の名前に、プログラムの中の名前を書く", "check_flow.py", act_stage, rename("Main.23", label="ClearContents のあと、売上データを「請求一覧」に読み込む"),
         "プログラムの中の名前「ClearContents」", redo)
    case("名付け: 箱の補足に、プログラムの中の名前を書く", "check_flow.py", act_stage, rename("集計表を保存する.56", detail="ThisWorkbook も保存する"),
         "プログラムの中の名前「ThisWorkbook」", redo)
    case("名付け: 図の名前に、手続きの名前を書く", "check_flow.py", act_stage, lambda w: edit(w / f"{a2}.edits.json", lambda e: e.update(name="Main の処理")),
         "図の名前・説明に、プログラムの中の名前「Main」", redo)
    case("名付け: かぎかっこの中に、プログラムの中だけの名前を書く", "check_flow.py", act_stage, rename("Main.15", label="「対象月」を「メニュー」から読み取る"),
         "かぎかっこの中の「対象月」は、プログラムの中だけの名前", redo)
    case("名付け: 箱の名前が、動詞で終わらない（体言止め）", "check_flow.py", act_stage, rename("Main.23", label="売上データの「請求一覧」への読み込み"), "動詞で終わっていない", redo)
    case("名付け: 判断の名前が、問いの形でない", "check_flow.py", act_stage, rename("Main.16", label="対象月の入力あり"), "「〜か」で終わっていない", redo)
    case("直しの一覧: 確認事項に、答えで変わる所（effect）が無い", "activity.py", ["apply", "--id", "A-01-2"],
         lambda w: edit(w / f"{a2}.edits.json", lambda e: e["questions"][0].pop("effect")), "effect")
    case("用語の対応: ソースに無い名前を載せる", "check_flow.py", act_stage, lambda w: edit(w / termj, lambda t: t["terms"][1].update(name="最終行番号")), "のソースに無い")
    case("用語の対応: 業務の言葉が空", "check_flow.py", act_stage, lambda w: edit(w / termj, lambda t: t["terms"][1].update(word="")), "業務の言葉（word）が空")
    case("用語の対応: 定義の場所に、その名前が無い", "check_flow.py", act_stage, lambda w: edit(w / termj, lambda t: t["terms"][0]["source"].update(line=5)), "定義の場所")
    case("用語の対応: 決まっていない根拠を書く", "check_flow.py", act_stage, lambda w: edit(w / termj, lambda t: t["terms"][1].update(basis="たぶん")), "terms[1].basis")
    case("用語の対応: 同じ名前を 2 回載せる", "check_flow.py", act_stage, lambda w: edit(w / termj, lambda t: t["terms"].append(dict(t["terms"][1]))), "2回載っている")
    w = tmp / "code_terms"                                           # 意味が読めない名前: 「不明」として載せれば、もとの名前のまま「」で書ける。確認事項になる
    shutil.copytree(good, w)
    edit(w / termj, lambda t: t["terms"].append({"project": "M01", "name": "対象月", "word": "", "basis": "unknown"}))
    rename("Main.15", label="「対象月」を「メニュー」から読み取る")(w)
    run("activity.py", w, "apply", "--all")
    code, text = run("check_flow.py", w, *act_stage)
    run("build_drawio.py", w)
    qs2 = json.loads((w / "questions.json").read_text(encoding="utf-8"))["questions"]
    code2, text2 = run("activity.py", w, "terms", "--id", "A-01-2")
    out.append(("用語の対応: 意味が読めない名前は「不明」で載せて、もとの名前のまま書ける。一覧に「不明」で出る（確認事項にはしない）。候補の一覧（activity.py terms）に、登録の状況が出る" +
                ("" if code == 0 else text[-400:]), code == 0 and "0 FAIL" in text and not any(str(q.get("local", "")).startswith("terms-") for q in qs2)
                and any(r[1] == "対象月" and "不明" in r[3] for r in model(w)["プログラムの用語対応"])
                and code2 == 0 and "ErrHandler" in text2 and "未登録" in text2))
    w = tmp / "code_terms_fill"                                      # 根拠は、定義の行のコメント・定数の定義から機械が入れる
    shutil.copytree(good, w)
    edit(w / termj, lambda t: (t["terms"][0].pop("basis"), t["terms"][0].pop("source"),
                               t["terms"].append({"project": "M01", "name": "Main", "word": "集計の処理", "basis": "spelling"})))
    code, text = run("activity.py", w, "apply", "--all")
    rows = json.loads((w / termj).read_text(encoding="utf-8"))["terms"]
    code2, text2 = run("check_flow.py", w, *act_stage)
    out.append(("用語の対応: 根拠を書かなければ、定数の定義から機械が「定義」と場所を入れる。定義の行にコメントがあるのに「つづり」なら WARN" + ("" if code == 0 else text[-300:]),
                code == 0 and rows[0].get("basis") == "definition" and rows[0].get("source") == {"module": "月次集計", "line": 4}
                and "「Main」: 定義の行にコメントがある" in text2))
    from systems import machine_questions
    mq = dict(machine_questions(good))
    code, text = run("activity.py", good, "list", "--group", "2")
    out.append(("台帳の機械の問い: 同じ起動点に、紐付けと実行のしかたの 2 つを聞くときは、1 つにまとめる（決まり 5.6）。組分けは、呼ばれる側から先の順で機械が出す",
                "manual:M01.e4" not in mq and "あわせて、どの操作で実行していますか" in mq.get("unlinked:M01.e4", {}).get("question", "")
                and "manual:M01.e2" in mq and code == 0 and "組 1: A-01-5" in text))
    r = subprocess.run([sys.executable, str(S / "diff_runs.py"), "systems", str(good), str(good), "--min", "1.0"], capture_output=True, text=True)
    out.append(("揺れの測り方（diff_runs.py）: 台帳・用語の対応・木（機械がつなぎ直したあと）を比べる。同じものは 100%",
                r.returncode == 0 and all(x in r.stdout for x in ("紐付け（起動点 × 手順）", "両方にある名前の言葉", "木のブロック"))))
    w = tmp / "code_jargon"
    shutil.copytree(good, w)
    rename("Main.23", label="売上データを配列に読み込む")(w)
    run("activity.py", w, "apply", "--all")
    code, text = run("check_flow.py", w, *act_stage)
    out.append(("名付け: 箱に技術の用語（配列・状態コード など）があれば WARN（「」の中は見ない）", "技術の用語「配列」" in text))
    w = tmp / "code_draft"                                           # 名前の無い箱が残る図は、下書き
    shutil.copytree(good, w)
    edit(w / f"{a2}.edits.json", lambda e: e.update(ops=[o for o in e["ops"] if o.get("item") != "Main.15"]))
    run("activity.py", w, "apply", "--all")
    code, text = run("build_activity.py", w)
    xml = "".join(p.read_text(encoding="utf-8") for p in (w / "out").glob("アクティビティ図*.drawio"))
    out.append(("下書きの印: 名前の無い箱が残る図は、ページの名前と題に【下書き】、1 ページ目の一覧にも出る", code == 0 and xml.count("【下書き】") >= 3 and "納品物ではない" in xml))
    # 畳む・分ける・図に戻す（直しの一覧の、ほかの種類）
    w = tmp / "code_ops"
    shutil.copytree(good, w)
    edit(w / f"{a2}.edits.json", lambda e: e.update(ops=[o for o in e["ops"] if o.get("item") not in ("Main.32", "Main.33")] + [
        {"op": "fold", "item": "Main.32", "reason": "型や空の確認"}, {"op": "name", "item": "Main.32", "label": "金額の大きい行に印を付ける"},
        {"op": "split", "item": "メールを送る.62", "at": [64]}, {"op": "name", "item": "メールを送る.64", "label": "宛先と件名を入れて送る"}, {"op": "include", "item": "Main.12"},
        {"op": "name", "item": "Main.12", "label": "画面の更新を止める"}], questions=[]))
    code, text = run("activity.py", w, "apply", "--id", "A-01-2")
    a = json.loads((w / f"{a2}.activity.json").read_text(encoding="utf-8")) if code == 0 else {"nodes": [], "branches": [], "lines": []}
    by = {n["label"]: n for n in a["nodes"]}
    code2, text2 = run("check_flow.py", w, "--stage", "activity")
    out.append(("直しの一覧: 畳む（分岐が1つの箱に）・分ける（行で2つに）・図に戻す（対象外の文を箱に）。行と分岐の網羅は保たれる" + ("" if code == 0 else text[-400:]),
                code == 0 and by.get("金額の大きい行に印を付ける", {}).get("folded") == "型や空の確認" and "宛先と件名を入れて送る" in by and "画面の更新を止める" in by
                and any(b["key"] == "Main.32" and b["state"] == "folded" for b in a["branches"]) and code2 == 0 and "0 FAIL" in text2))
    return out


VBA_CLOSE = """Private 直前の状態 As String

Public Sub 月末の締め()
    Dim 名前 As String
    名前 = 帳票の名前("請求")
    Worksheets("一覧").Range("A1").Value = 名前
    名前 = 帳票の名前("入金")
    Worksheets("一覧").Range("A2").Value = 名前
    印を付ける "請求"
    印を付ける "入金"
    状態を出す 1
    MsgBox "締めが終わりました"
    状態を出す 2
End Sub

Private Function 帳票の名前(種類 As String) As String
    帳票の名前 = 種類 & "_" & Format(Date, "yyyymm")
End Function

Private Sub 印を付ける(種類 As String)
    Worksheets("控え").Range("A1").Value = 種類
End Sub

Private Sub 状態を出す(番号 As Long)
    直前の状態 = "状態 " & Format(番号, "0000")
    Debug.Print 直前の状態
End Sub
"""


def narrow_tests(tmp):
    """図にする範囲を絞る: 補助の手続き（入出力・やり取りの無い、共通の処理）を、台帳に理由つきで書くと、図にしない。呼び出しは、呼ぶ側のふつうの箱になる"""
    out, src, w = [], tmp / "narrow_in", tmp / "narrow_work"
    (src / "src").mkdir(parents=True)
    (src / "src" / "締め.bas").write_text('Attribute VB_Name = "締め"\n' + VBA_CLOSE, encoding="utf-8")
    subprocess.run([sys.executable, str(S / "inventory.py"), str(src), "-w", str(w)], capture_output=True, text=True)
    for sc, args in (("render_pages.py", []), ("code_struct.py", []), ("systems.py", ["draft"])):
        run(sc, w, *args)
    sysj = w / "structure/systems.json"
    edit(sysj, lambda j: [x.update(kind="euc_excel") for x in j["systems"]])
    _c, ent = run("systems.py", w, "entries")
    star = [ln for ln in ent.splitlines() if "★補助の候補" in ln]
    out.append(("補助の候補: 入出力・やり取りの無い、共通の処理だけを、機械が候補に出す（シートに書く手続きは、候補にしない）",
                len(star) == 2 and "M01.締め.帳票の名前" in star[0] and "M01.締め.状態を出す" in star[1]))

    def listing():
        _c, text = run("activity.py", w, "list")
        return {name: ln for ln in text.splitlines() for name in ("月末の締め", "帳票の名前", "印を付ける", "状態を出す") if f"締め.{name}（" in ln}
    ids = {k: v.split()[0] for k, v in listing().items()}
    main, helper, mark, state = ids.get("月末の締め", "A-01-1"), ids.get("帳票の名前", "A-01-2"), ids.get("印を付ける", "A-01-3"), ids.get("状態を出す", "A-01-4")
    (w / "activities").mkdir(exist_ok=True)
    put = lambda aid, name, ops: (w / f"activities/{aid}.edits.json").write_text(json.dumps(
        {"activity": aid, "name": name, "summary": "", "questions": [], "ops": ops}, ensure_ascii=False), encoding="utf-8")
    put(helper, "帳票の名前を作る", [{"op": "name", "item": "帳票の名前.18", "label": "種類と年月をつないで、帳票の名前を作る"}])
    put(mark, "控えに印を付ける", [{"op": "name", "item": "印を付ける.22", "label": "「控え」に、種類を書く"}])
    put(state, "状態の番号を控える", [{"op": "name", "item": "状態を出す.26", "label": "状態の番号を、4 桁にそろえて控える"}])
    put(main, "月末の締めをする", [{"op": "name", "item": "月末の締め.7", "label": "請求の帳票の名前を「一覧」に書く"}, {"op": "name", "item": "月末の締め.9", "label": "入金の帳票の名前を「一覧」に書く"}])
    code0, text0 = run("activity.py", w, "apply", "--all")
    run("build_activity.py", w)
    xml0 = "".join(f.read_text(encoding="utf-8") for f in (w / "out").glob("アクティビティ図*.drawio"))
    # 台帳に、理由つきで書く
    edit(sysj, lambda j: j.update(folded_procedures=[{"procedure": "M01.締め.帳票の名前", "reason": "保存先・名前の組み立て"}, {"procedure": "M01.締め.状態を出す", "reason": "ログの表示"}]))
    after = listing()
    out.append(("図にしない手続き: 台帳に理由つきで書くと、一覧（activity.py list）に「図にしない」と出る。ほかの図の番号は変わらない" + ("" if code0 == 0 else text0[-300:]),
                code0 == 0 and f'id="act-{helper}"' in xml0 and "図にしない（保存先・名前の組み立て）" in after.get("帳票の名前", "")
                and all(after.get(k, "").split()[:1] == [ids.get(k)] for k in ("月末の締め", "印を付ける"))))
    _c, draft = run("activity.py", w, "draft", "--id", main)
    out.append(("下書き: 補助の手続きを呼ぶ行は、ふつうのまとまりになり、その手続きの中身が、名付けの手がかりに出る",
                "補助の手続き 締め.帳票の名前 を呼ぶ" in draft and 'Format(Date, "yyyymm")' in draft and f"別の図 {helper}" not in draft and f"別の図 {mark}" in draft))
    first, second = "請求の帳票の名前を作り、「一覧」に書く", "入金の帳票の名前を作り、「一覧」に書く"
    put(main, "月末の締めをする", [{"op": "merge", "items": ["月末の締め.6", "月末の締め.7"]}, {"op": "name", "item": "月末の締め.6", "label": first},
                               {"op": "merge", "items": ["月末の締め.8", "月末の締め.9"]}, {"op": "name", "item": "月末の締め.8", "label": second}])
    code, text = run("activity.py", w, "apply", "--all")
    act = json.loads((w / f"activities/{main}.activity.json").read_text(encoding="utf-8")) if code == 0 else {"nodes": [], "lines": []}
    labels = [n["label"] for n in act["nodes"] if n["type"] == "process"]
    logs = sorted(ln for o in act["lines"] if o["state"] == "excluded" and o.get("reason") == "ログ" for ln in o["lines"])
    code1, text1 = run("build_activity.py", w)
    xml = "".join(f.read_text(encoding="utf-8") for f in (w / "out").glob("アクティビティ図*.drawio"))
    code2, text2 = run("check_flow.py", w, "--stage", "systems")
    code3, text3 = run("check_flow.py", w, "--stage", "activity")
    code4, text4 = run("check_flow.py", w, "--stage", "drawio")
    good = code == 0 and code1 == 0 and code2 == 0 and code3 == 0 and code4 == 0
    out.append(("図にしない手続き: 呼び出しは、呼ぶ側のふつうの箱になり、となりのまとまりとまとめられる（別ページへの箱にならない）。その手続きの図のページは、作られない。点検は FAIL 0"
                + ("" if good else (text + text1)[-300:] + text2[-300:] + text3[-400:] + text4[-300:]),
                good and first in labels and second in labels and [n.get("link_flow") for n in act["nodes"] if n["type"] == "subprocess"] == [mark, mark]
                and f'id="act-{helper}"' not in xml and f'id="act-{mark}"' in xml and "図にしない手続き（補助）2（実行文 3 行）" in text2 and "0 FAIL" in text3))
    out.append(("図にしない手続き: ログの表示だけをする補助を呼ぶ行は、機械が対象外にする（理由「ログ」。名前を付けなくてよい）", logs == [12, 14] and f'id="act-{state}"' not in xml))
    out.append(("図にしない手続き: 前に作った図のデータが残っていれば、知らせる（図には出ない）", f"{helper}: 図にしない手続き" in text3 and "前に作った図のデータが残っている" in text3))
    _c, st = run("status.py", w)
    out.append(("図にしない手続き: 進み具合（status.py）は、図にしない手続きを、残りの図に数えない", "アクティビティ図 2/2 枚" in st))
    code, text = run("build_matrix.py", w)
    from build_matrix import model
    mx = model(w) if code == 0 else {}
    row = next((r for r in mx.get("プログラムの図にしない手続き", []) if r[1] == "締め.帳票の名前"), None)
    ent_row = next((r for r in mx.get("起動点一覧", []) if r[3] == "締め.帳票の名前"), None)
    code5, text5 = run("check_flow.py", w, "--stage", "matrix")
    out.append(("業務マトリクス一覧: シート「プログラムの図にしない手続き」に、理由・呼ぶ側の図・呼ぶ所が出る。起動点一覧には、番号の代わりに、図にしない理由が出る" + ("" if code == 0 and code5 == 0 else text[-300:] + text5[-400:]),
                row is not None and row[2] == "図にしない" and row[3] == "保存先・名前の組み立て" and row[4] == main and "締め.月末の締め 6 行" in row[5] and "締め.月末の締め 8 行" in row[5]
                and ent_row is not None and ent_row[4] == "—" and "図にしない: 保存先・名前の組み立て" in ent_row[8] and code5 == 0 and "0 FAIL" in text5
                and all(r[2] == "図にしない" for r in mx["プログラムの図にしない手続き"][1:])))
    if code == 0:                                                    # プログラムを読む前に書き出した本（プログラム用のシートが空で、入っていない）のまま: 点検が知らせる
        import table_template as TT
        from openpyxl import load_workbook
        book = w / "out" / "業務マトリクス一覧.xlsx"
        wb, side = load_workbook(book), json.loads(TT.side_path(book).read_text(encoding="utf-8"))
        del wb["プログラムのノード一覧"]
        wb.save(book)
        side["sheets"].pop("プログラムのノード一覧")
        side["empty"] = ["プログラムのノード一覧"]
        TT.side_path(book).write_text(json.dumps(side, ensure_ascii=False), encoding="utf-8")
        code6, text6 = run("check_flow.py", w, "--stage", "matrix")
        out.append(("業務マトリクス一覧: 中身が無くて入れなかったシートに、いまは中身がある（書き出し直していない）は、点検が FAIL", code6 != 0 and "書き出したときは中身が無く" in text6))

    def bad(name, change, expect):
        cw = tmp / "narrow_case"
        if cw.exists():
            shutil.rmtree(cw)
        shutil.copytree(w, cw)
        edit(cw / "structure/systems.json", change)
        code, text = run("check_flow.py", cw, "--stage", "systems")
        hit = next((ln for ln in text.splitlines() if ln.startswith("FAIL") and expect in ln), None)
        out.append((f"{name}: {(hit or '')[:110]}" if hit and code != 0 else f"{name}: 「{expect}」が出ない（終了コード {code}）\n{text[-600:]}", bool(hit) and code != 0))
    fold = lambda proc, reason="探す・取り出す": lambda j: j["folded_procedures"].append({"procedure": proc, "reason": reason})
    bad("図にしない手続き: シートに書く手続きを書く", fold("M01.締め.印を付ける"), "補助の候補でない")
    bad("図にしない手続き: 決まっていない理由を書く", lambda j: j["folded_procedures"][0].update(reason="細かいから"), "folded_procedures[0].reason")
    bad("図にしない手続き: 起動点の手続きを書く", fold("M01.締め.月末の締め"), "起動点の手続き")
    bad("図にしない手続き: 無い手続きを書く", fold("M01.締め.控えを消す"), "という手続きが無い")
    bad("図にしない手続き: 同じ手続きを 2 回書く", fold("M01.締め.帳票の名前", "保存先・名前の組み立て"), "2回書かれている")
    return out


def book_dump(path):
    """xlsx を、比べられる形に（セルの文字・字・塗り・折り返し／列の幅・固定する枠・フィルター・入力規則）"""
    from openpyxl import load_workbook
    out = []
    for ws in load_workbook(path).worksheets:
        out.append((ws.title, ws.freeze_panes, ws.auto_filter.ref, sorted((k, round(v.width or 0, 2)) for k, v in ws.column_dimensions.items()),
                    sorted((str(dv.sqref), dv.formula1) for dv in ws.data_validations.dataValidation)))
        for row in ws.iter_rows():
            for x in row:
                if x.value is not None or x.has_style:
                    out.append((ws.title, x.coordinate, x.value, x.font.b, x.font.sz, str(x.font.color.rgb) if x.font.color is not None else None,
                                str(x.fill.fgColor.rgb) if x.fill.fill_type else None, x.alignment.wrap_text, x.border.left.style))
    return out


def template_tests(tmp, good):
    """納品物の様式（テンプレート）: 表（xlsx）は、差し替えたテンプレートのとおりに出て、点検が通る。決まりを外れたテンプレートは FAIL。アクティビティ図は、そのテンプレートの形・色・凡例で描く"""
    import copy
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    import build_matrix as BM
    import build_reports as BR
    import harvest_template as HT
    import table_template as TT
    assets, out, d, w = SKILL / "assets", [], tmp / "tpl", tmp / "tpl_work"
    d.mkdir()
    shutil.copytree(good, w)
    same = True
    for _kind, (spec, look, name) in TT.books().items():
        TT.make_default(spec, look, d / name)
        same = same and book_dump(d / name) == book_dump(assets / name)
    r = subprocess.run([sys.executable, str(S / "table_template.py"), "check"], capture_output=True, text=True)
    out.append(("テンプレート（xlsx）: assets/ の 2 つが、既定の作り方（table_template.py default）と同じ。確かめ（check）が FAIL 0" + ("" if r.returncode == 0 else r.stdout[-400:]),
                same and r.returncode == 0 and "0 FAIL" in r.stdout))
    tok = lambda ws, t: next(x for row in ws.iter_rows() for x in row if x.value == t)

    def swap_cols(ws, a, b, rows):
        for i in rows:
            x, y = ws.cell(i, a), ws.cell(i, b)
            x.value, y.value = y.value, x.value
            x._style, y._style = copy.copy(y._style), copy.copy(x._style)
    # 業務マトリクス一覧: シートの名前・見出しの文字・列の順・判断の塗りを変える。任意の欄を、人が埋める列に替える
    wb = load_workbook(assets / BM.TEMPLATE)
    ws = wb["業務マトリクス"]
    a, b = tok(ws, "{分類名の根拠}").column, tok(ws, "{原典（資料とページ）}").column
    swap_cols(ws, a, b, (1, 2))
    ws["A1"].value, ws.title = "業務領域", "業務一覧"
    ws = wb["業務フローのノード一覧"]
    memo = tok(ws, "{参考}")
    mcol = memo.column
    ws.cell(1, mcol).value, memo.value = "メモ（人が書く）", None
    for row in wb[TT.STYLE_SHEET].iter_rows():
        if row[0].value == "判断":
            row[1].fill = PatternFill("solid", fgColor="FFD966")
    wb.save(d / "matrix_mod.xlsx")
    code, text = run("build_matrix.py", w, "--template", str(d / "matrix_mod.xlsx"))
    got, ref = load_workbook(w / "out" / BM.BOOK), load_workbook(good / "out" / BM.BOOK)["業務マトリクス"]
    ws, nl = got["業務一覧"] if "業務一覧" in got.sheetnames else got.worksheets[0], got["業務フローのノード一覧"]
    back = TT.read_back(w / "out" / BM.BOOK, BM.SPEC) if code == 0 else {"absent": {}}
    code2, text2 = run("check_flow.py", w, "--stage", "matrix")
    out.append(("テンプレート（xlsx）: シートの名前・見出しの文字・列の順・判断の塗りを変え、任意の欄を人が埋める列に替えたテンプレートのとおりに出て、点検（--stage matrix）が通る"
                + ("" if code == 0 and code2 == 0 else text[-300:] + text2[-500:]),
                code == 0 and "業務マトリクス" not in got.sheetnames and ws["A1"].value == "業務領域" and ws.cell(1, a).value == ref.cell(1, b).value
                and ws.cell(2, a).value == ref.cell(2, b).value and ws.cell(2, b).value == ref.cell(2, a).value and ws.max_row == ref.max_row
                and nl.cell(1, mcol).value == "メモ（人が書く）" and all(nl.cell(i, mcol).value is None for i in range(2, nl.max_row + 1))
                and any(str(x.fill.fgColor.rgb).endswith("FFD966") for row in ws.iter_rows(min_row=2) for x in row if x.fill.fill_type)
                and "参考" in back["absent"].get("業務フローのノード一覧", ()) and code2 == 0 and "0 FAIL" in text2))
    # 確認事項一覧: 表題の行を足す・列を入れ替える・人が埋める列（選択式）を足す
    wb = load_workbook(assets / BR.Q_TEMPLATE)
    ws = wb[BR.Q_SHEET]
    a, b, last = tok(ws, "{質問}").column, tok(ws, "{仮置きした内容}").column, ws.max_column
    swap_cols(ws, a, b, (1, 2))
    ws.insert_rows(1, 2)
    ws["A1"].value, ws.freeze_panes = "確認事項一覧（お客様への質問）", "A4"
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A3:{get_column_letter(last + 1)}3"
    for i, v in ((3, "ステータス"), (4, "未回答")):
        ws.cell(i, last + 1).value, ws.cell(i, last + 1)._style = v, copy.copy(ws.cell(i, last)._style)
    col = get_column_letter(last + 1)
    dv = DataValidation(type="list", formula1='"未回答,回答済,保留"', allow_blank=True)
    dv.add(f"{col}4")
    ws.add_data_validation(dv)
    wb.save(d / "questions_mod.xlsx")
    base = TT.read_back(good / "out" / BR.Q_BOOK, BR.Q_SPEC)["rows"]
    code, text = run("build_reports.py", w, "--template", str(d / "questions_mod.xlsx"))
    ws = load_workbook(w / "out" / BR.Q_BOOK)[BR.Q_SHEET]
    rows = TT.read_back(w / "out" / BR.Q_BOOK, BR.Q_SPEC)["rows"] if code == 0 else None
    out.append(("テンプレート（xlsx）: 確認事項一覧に、表題の行・列の入れ替え・人が埋める列（選択式）を足しても、中身は同じ。見本の行の文字と入力規則が、すべての行に写る" + ("" if code == 0 else text[-400:]),
                code == 0 and rows == base and len(base[BR.Q_SHEET]) > 2 and ws["A1"].value == "確認事項一覧（お客様への質問）" and ws.cell(3, a).value == "仮置きした内容"
                and ws.max_row == len(base[BR.Q_SHEET]) + 2 and all(ws.cell(i, last + 1).value == "未回答" for i in range(4, ws.max_row + 1))
                and [str(v.sqref) for v in ws.data_validations.dataValidation if v.formula1 == '"未回答,回答済,保留"'] == [f"{col}4:{col}{ws.max_row}"]))

    def broken(name, book, change, expect):
        path = d / f"bad{len(out)}.xlsx"
        wb = load_workbook(assets / (BM.TEMPLATE if book == "matrix" else BR.Q_TEMPLATE))
        change(wb)
        wb.save(path)
        r = subprocess.run([sys.executable, str(S / "table_template.py"), "check", str(path), "--book", book], capture_output=True, text=True)
        hit = next((ln for ln in r.stdout.splitlines() if ln.startswith("FAIL") and expect in ln), None)
        out.append((f"{name}: {(hit or '')[:110]}" if hit and r.returncode != 0 else f"{name}: 「{expect}」が出ない（終了コード {r.returncode}）\n{r.stdout[-600:]}", bool(hit) and r.returncode != 0))
        return path
    setv = lambda x, v: setattr(x, "value", v)
    bad_book = broken("テンプレート（xlsx）: 知らない印を書く", "matrix", lambda wb: setv(tok(wb["業務フローのノード一覧"], "{参考}"), "{無い欄}"), "知らない印")
    broken("テンプレート（xlsx）: 点検が使う欄を消す", "questions", lambda wb: setv(tok(wb[BR.Q_SHEET], "{質問}"), None), "点検が使う欄が無い")
    broken("テンプレート（xlsx）: 見本の行の下に文字を書く", "matrix", lambda wb: setv(wb["業務フローのノード一覧"].cell(5, 1), "メモ"), "より下に、文字がある")
    def two_repeats(wb):
        ws = wb["業務マトリクス"]
        ws.cell(1, ws.max_column + 1).value, ws.cell(2, ws.max_column).value = "①", "{作業}"
    broken("テンプレート（xlsx）: 繰り返す列（作業）を、2 つ置く", "matrix", two_repeats, "1つだけ置く")
    wb = load_workbook(assets / BM.TEMPLATE)                          # 繰り返す列の右に、人が埋める列を足す。作業の列が増えたぶん、右へ送られる
    ws = wb["業務マトリクス"]
    last = ws.max_column
    ws.cell(1, last + 1).value, ws.cell(1, last + 1)._style = "備考（人が書く）", copy.copy(ws.cell(1, last)._style)
    wb.save(d / "matrix_right.xlsx")
    w2 = tmp / "tpl_right"
    shutil.copytree(good, w2)
    code, text = run("build_matrix.py", w2, "--template", str(d / "matrix_right.xlsx"))
    code2, text2 = run("check_flow.py", w2, "--stage", "matrix")
    ws = load_workbook(w2 / "out" / BM.BOOK)["業務マトリクス"] if code == 0 else None
    heads = [c.value for c in ws[1]] if ws is not None else []
    out.append(("テンプレート（xlsx）: 繰り返す列（作業）の右に置いた欄（人が埋める列）は、作業の列が増えたぶん、右へ送られる。点検が通る" + ("" if code == 0 and code2 == 0 else text[-300:] + text2[-500:]),
                code == 0 and code2 == 0 and "0 FAIL" in text2 and heads[-1] == "備考（人が書く）" and heads[-2] in TT.CIRCLED[1:] and heads.count("①") == 1))
    code, text = run("build_matrix.py", w, "--template", str(bad_book))
    shutil.copy(assets / BM.TEMPLATE, d / "macro.xlsm")
    r = subprocess.run([sys.executable, str(S / "table_template.py"), "check", str(d / "macro.xlsm"), "--book", "matrix"], capture_output=True, text=True)
    out.append(("テンプレート（xlsx）: 決まりを外れたテンプレートでは、書き出さない。マクロ付き（.xlsm）は、受け取らない",
                code != 0 and "決まりどおりでない" in text and r.returncode != 0 and "テンプレートは .xlsx だけ" in r.stdout))
    # アクティビティ図のテンプレート（draw.io）
    in_sync = all(HT.harvest(str(assets / t), act) == json.loads((assets / o).read_text(encoding="utf-8")) for act, (t, o) in HT.TEMPLATES.items())
    out.append(("テンプレート（draw.io）: assets/ の template_styles.json・activity_styles.json が、テンプレートから抜き直したものと同じ", in_sync))
    cw, name = tmp / "tpl_code", HT.TEMPLATES[True][0]
    shutil.copytree(tmp / "code_good", cw)
    old = json.loads((assets / HT.TEMPLATES[True][1]).read_text(encoding="utf-8"))["palettes"][0]["fill"]
    body = (assets / name).read_text(encoding="utf-8")
    (d / name).write_text(body.replace(f"fillColor={old}", "fillColor=#E2F0D9").replace("凡例（アクティビティ図）", "凡例（処理の流れの図）"), encoding="utf-8")
    harvest = lambda path: subprocess.run([sys.executable, str(S / "harvest_template.py"), str(path), "-o", str(d / "act.json"), "--activity"], capture_output=True, text=True)
    r = harvest(d / name)
    code, text = run("build_activity.py", cw, "--styles", str(d / "act.json"))
    xml = "".join(f.read_text(encoding="utf-8") for f in (cw / "out").glob("アクティビティ図*.drawio"))
    code2, text2 = run("check_flow.py", cw, "--stage", "drawio")
    out.append(("テンプレート（draw.io）: アクティビティ図のテンプレートの、箱の色と凡例の文を変えると、図がそのとおりに変わる" + ("" if r.returncode == 0 and code == 0 else (r.stdout + r.stderr + text)[-400:]),
                r.returncode == 0 and code == 0 and "fillColor=#E2F0D9" in xml and f"fillColor={old}" not in xml and "凡例（処理の流れの図）" in xml and old in body))
    out.append(("テンプレート（draw.io）: assets/ のテンプレートに無い色のアクティビティ図は、点検で FAIL（差し替えるときは、assets/ のテンプレートを置き換えて、抜き直す）",
                code2 != 0 and any(ln.startswith("FAIL") and "はテンプレートに無い色" in ln for ln in text2.splitlines())))
    (d / "nostart.drawio").write_text(body.replace("shape=startState", "shape=ellipse"), encoding="utf-8")
    r = harvest(d / "nostart.drawio")
    out.append(("テンプレート（draw.io）: 開始の丸の無いアクティビティ図のテンプレートは、読み取らない", r.returncode != 0 and "開始（shape=startState）" in r.stdout + r.stderr))
    return out


def graph(pairs, lanes=None):
    """列の決め方を確かめるための、小さい図（ノードの種類は、番号の頭の字から）"""
    ids = list(dict.fromkeys(x for p in pairs for x in p))
    kind = lambda i: "start" if i == "start" else "end" if i.startswith("end") else "decision" if i[0] == "D" else "process"
    return [{"id": i, "lane": (lanes or {}).get(i, "p"), "type": kind(i)} for i in ids], [{"from": a, "to": b} for a, b in pairs]


def column_tests():
    """分岐の先を横に並べる配置（build_drawio.assign_cols）"""
    import build_drawio as B
    import terms as T
    out = []
    at = lambda pairs, cap=4: (lambda c, r: {i: (c[i], r[i]) for i in c})(*B.assign_cols(*graph(pairs), cap))
    two = [("start", "D1"), ("D1", "P1"), ("D1", "P2"), ("P1", "P3"), ("P2", "P3"), ("P3", "end")]
    p = at(two)
    out.append(("列: 2 つに分かれて合流する道は、同じ高さに横に並び、合流する箱は、もとの列", p["P1"] == (0, 2) and p["P2"] == (1, 2) and p["P3"] == (0, 3)))
    p = at([("start", "D1"), ("D1", "P1"), ("D1", "P2"), ("P1", "P2"), ("P2", "end")])
    out.append(("列: 条件に当たったときだけ行う短い処理は右の列に出し、当たらない側は、まっすぐ下", p["P1"] == (1, 2) and p["P2"] == (0, 3) and p["D1"] == (0, 1)))
    p = at([("start", "D1"), ("D1", "end2"), ("D1", "P1"), ("P1", "end")])
    out.append(("列: 途中で終わる側は右の列、ふつうの終了に着く道は、まっすぐ下", p["end2"] == (1, 2) and p["P1"] == (0, 2) and p["end"] == (0, 3)))
    p = at([("start", "D1"), ("D1", "P1"), ("P1", "D1"), ("D1", "P2"), ("P2", "end")])
    out.append(("列: 繰り返しは、中身も、抜けた先も、同じ列（抜けた先は、中身より下）", {c for c, r in p.values()} == {0} and p["P2"][1] > p["P1"][1]))
    p = at([("start", "D1"), ("D1", "P1"), ("D1", "P2"), ("D1", "P3"), ("P1", "P4"), ("P2", "P4"), ("P3", "P4"), ("P4", "end")])
    out.append(("列: 3 つに分かれる振り分けは、右へ順に並ぶ", [p[i] for i in ("P1", "P2", "P3", "P4")] == [(0, 2), (1, 2), (2, 2), (0, 3)]))
    p = at(two, 1)
    out.append(("列: 列を 1 つにすると、いまの縦の並び（1 箱に 1 行）", {c for c, r in p.values()} == {0} and sorted(r for c, r in p.values()) == list(range(6))))
    wide = [("start", "D1")] + [("D1", f"P{i}") for i in range(1, 7)] + [(f"P{i}", "P9") for i in range(1, 7)] + [("P9", "end")]
    p = at(wide, 3)
    out.append(("列: 上限を超える分は、いちばん右の列に縦に並ぶ（同じ所に 2 つ置かない）", max(c for c, r in p.values()) == 2 and len(set(p.values())) == len(p)))
    import activity as AC
    proc = lambda pid, to: {"id": pid, "body": [{"kind": "call", "to": to}] if to else [{"kind": "seq"}]}
    fake = type("Led", (), {"units": [{"kind": "shared", "project": "M01", "procedure": f"m.{x}", "activity": f"A-01-{k}"} for k, x in enumerate("abc", 1)],
                            "structs": {"M01": {"procedures": [proc("m.a", "m.c"), proc("m.b", None), proc("m.c", "m.b")]}},
                            "activity": {("M01", f"m.{x}"): f"A-01-{k}" for k, x in enumerate("abc", 1)}})()
    out.append(("図を作る順: 共通の処理どうしでも、呼ばれる側が先（呼び出しの箱に、行き先の図のいまの名前が入る）",
                [u["activity"] for u in AC.ordered(fake)] == ["A-01-2", "A-01-3", "A-01-1"]))
    # 名前の見分け方
    code = 'Set wb = Workbooks.Open(savePath & "uriage_" & TargetMonth)  \' open_file を呼ぶ'
    names = T.code_names(code, "vba")
    out.append(("名前の見分け: 文字列・コメントの中の文字は、プログラムの中の名前にしない", {"savePath", "TargetMonth", "Workbooks"} <= names and "uriage_" not in names and "open_file" not in names))
    out.append(("名前の見分け: 記号でつないだ形・大文字小文字まじり・大文字 4 字以上だけを、名前らしい形と見る",
                [T.shaped(x) for x in ("savePath", "ACCT-ID", "tax_rate", "TOTAL", "PDF", "Excel", "Outlook")] == [True, True, True, True, False, False, False]))
    out.append(("名前の見分け: 箱の名前の、かぎかっこの外に出ている名前だけを拾う",
                T.inner_names("savePath の「TargetMonth」を PDF にする", names, set()) == ["savePath"] and T.inner_names("Workbooks を開く", names, {"Workbooks"}) == ["Workbooks"]))
    out.append(("名前の見分け: かぎかっこの中が、文字列の中の文字・入出力の名前なら通し、プログラムの中だけの名前なら拾う（不明として載せた名前は通す）",
                T.quoted_inner("「uriage_」を「TargetMonth」で開き「請求一覧」に写す", code, "vba", "請求一覧", set()) == ["TargetMonth"]
                and T.quoted_inner("「TargetMonth」で開く", code, "vba", "", {"TargetMonth"}) == []))
    out.append(("名前の見分け: VBA の文字列は、\\ で終わっても、そこで閉じる（\\ は、ただの文字）",
                T.quoted_inner("場所に「.dat\\」を付ける", 'saveTo = basePath & ".dat\\"', "vba", "", set()) == [] and "basePath" in T.code_names('saveTo = basePath & ".dat\\"  \' x', "vba")))
    out.append(("名前の見分け: COBOL の、数字で始まる段落の名前を 1 つに取り、言語の語は名前にしない",
                "9000-CLOSE-FILE" in T.code_names("           PERFORM 9000-CLOSE-FILE THRU 9000-EXIT", "cobol") and "PERFORM" not in T.code_names("           PERFORM 9000-CLOSE-FILE", "cobol")))
    return out


def form_tests(tmp, good):
    """出典は枠の外に、切らずに出す／箱の名前は動詞で終え、補足はかっこ書き／確認事項は、答えで図が変わるものだけ"""
    import xml.etree.ElementTree as ET
    import activity as A
    import build_drawio as B
    from common import collect_questions, ends_as_question, ends_with_verb, load, load_flows, manifest
    from openpyxl import load_workbook
    out = []
    xml = next(p for p in (good / "out").glob("*.drawio") if "アクティビティ図" not in p.name).read_text(encoding="utf-8")
    root = ET.fromstring(xml)
    inside = [o.get("label", "") for o in root.iter("object") if o.get("node_id")]
    caps = [x.get("value", "") for x in root.iter("mxCell") if str(x.get("id", "")).startswith("src-")]
    f3 = next(pg for lp in (good / "out").glob("*.layout.json") for pg in json.loads(lp.read_text(encoding="utf-8"))["pages"] if pg["id"] == "flow-F03")
    box, cap = f3["nodes"], f3["captions"]
    out.append(("出典: 箱・ひし形の中に、出典を書かない。枠の外（形のすぐ下・真ん中の線の右）に、切らずに出る",
                not any("出典" in v or "M02 p." in v or "シート C" in v for v in inside) and any("経費精算チェックシート「手順」シート C2" in v for v in caps) and not any("…" in v for v in caps)
                and all(b[1] >= box[i][1] + box[i][3] and b[0] > box[i][0] + box[i][2] / 2 for i, b in cap.items())
                and sorted(cap) == sorted(i for i in box if not i.startswith(("F-", "start", "end")))))
    out.append(("出典: 枠の外に出したので、ひし形は、番号と条件だけの大きさ（テンプレートの 150×80）", box["D-001"][2:] == [150, 80] and box["D-002"][2:] == [150, 80]))
    ctx = {"styles": load(SKILL / "assets" / "template_styles.json"), "manifest": manifest(good), "taxonomy": load(good / "taxonomy.json", {}) or {},
           "hierarchy": load(good / "structure" / "hierarchy.json", {}) or {}, "questions": collect_questions(good, write=False), "today": "2026-01-01", "page_ids": set()}
    keep = B.src_lines
    try:                                                             # 長い出典（架空の名前）: 切らずに、その列を広げる
        B.src_lines = lambda sources, px=B.CAP_MAX: ["M99「とても長い名前の、架空の資料の題」シート AB123"] if sources else []
        d = B.FlowDrawing(next(f for f in load_flows(good) if f["flow_id"] == "F03"), ctx)
        d.measure()
        d.place()
    finally:
        B.src_lines = keep
    out.append(("出典: 長い出典でも切らず、その列を、出典が収まる幅に広げる（テンプレートの幅より広くなる）",
                max(d.vw) > d.L["width"] and all(b[0] + b[2] <= d.lx(d.vcol(i)) + d.vw[d.vcol(i)] for i, b in d.cap_box.items())))
    out.append(("出典: 幅を超える行は、空白の所で折る。1語が長ければ字で折る（どちらも、切らない）",
                B.wrap("M99「長い 名前の 資料」 12〜34 行", 90) == ["M99「長い 名前の", "資料」 12〜34 行"] and "".join(B.wrap("ABCDEFGHIJKLMNOPQRSTUVWXYZ", 50)) == "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
                and len(B.wrap("ABCDEFGHIJKLMNOPQRSTUVWXYZ", 50)) > 1 and abs(B.text_px("M40「SAMPLE01.cbl」 319 行", 9) - 125.0) < 0.5))
    def diamond(label):
        """(高さ, 行, どの行も、その高さでのひし形の幅に収まっているか)"""
        h, rows = B.diamond_lines(label, 150, 80, B.text_px("D-001", 9), 10)
        n = 1 + len(rows)
        return h, rows, all(B.text_px(r, 10) <= 150 * (1 - 2 * (abs(i + 1 - (n - 1) / 2) * 12 + 5) / h) for i, r in enumerate(rows))

    long = "申請書の担当者の番号が、申請書の宛先にある担当者の番号と違うか"
    (h1, r1, in1), (h2, r2, in2), (h3, r3, in3) = diamond("不備があるか"), diamond(long), diamond("見つかった取引先が 0 件か")
    out.append(("ひし形: 条件の文は、形に合わせて機械が折る（上下の細い所ほど、行が短い。文は欠けない。行の頭に読点・閉じかっこを置かない）。長い条件でも、低いまま",
                (h1, r1) == (80, ["不備があるか"]) and in1 and in2 and in3 and h2 <= 100 and len(r2) >= 3 and "".join(r2).replace("\u00a0", "") == long
                and not any(r[:1] in B.NO_HEAD or r[-1:] in B.NO_TAIL for r in r2) and (h3, r3) == (80, ["見つかった取引先が 0 件か", "\u00a0"])
                and B.tokens("「ABC」が、空か") == ["「ABC」", "が、", "空", "か"] and B.diamond_lines("あ" * 80, 150, 80, 30, 10) == (170, ["あ" * 80])))
    out.append(("名前: 動詞で終わるかの見分け（末尾のかっこ書きは除いて見る。体言止め・「〜へ」は違う）。判断は「〜か」",
                all(ends_with_verb(x) for x in ("注文書を受領する", "メールを送る（Outlook）", "空の台帳を用意する（番号があれば読む（画面用））", "「取込が終わりました」と、メッセージを出す"))
                and not any(ends_with_verb(x) for x in ("承認へ", "振込結果を会計システムに連携", "申請書の確認", "メール（Outlook）"))
                and ends_as_question("不備があるか") and ends_as_question("不備があるか？") and not ends_as_question("不備あり？")))
    out.append(("名前: 機械が付ける名前（やり取りの箱・外部の箱）の形は、どれも動詞で終わる",
                all(ends_with_verb(a) and ends_with_verb(b.format("文字")) for a, b in list(A.INTERACT_LABEL.values()) + [A.INTERACT_OTHER])
                and all(ends_with_verb(v + "（相手の名前）") for v in A.EXTERNAL_LABEL.values())))
    out.append(("補足: 図には、かっこ書きで出る（すでに全体がかっこなら、重ねない）",
                "（毎月末）" in xml and B.paren("（自動）") == "（自動）" and B.paren("（a）と（b）") == "（（a）と（b））"))
    qmd = (good / "out" / "確認事項一覧.md").read_text(encoding="utf-8")
    head = [x.value for x in load_workbook(good / "out" / "確認事項一覧.xlsx").worksheets[0][1]]
    out.append(("確認事項: 一覧に「答えで変わる所」の列が出る（md・xlsx）。値だけの食い違い（支払日）は、確認事項にせず、箱の補足に両方の値",
                "| 答えで変わる所 |" in qmd and "答えで変わる所" in head and "規程では、毎月25日" in xml and "支払日は" not in qmd))
    for name, text, expect in (("値だけを聞く問い", "承認が要る金額の基準は、いくらですか。", "値（金額・日数・件数・期限）だけを聞く問い"),
                               ("書かれていることが、いまも同じかの確かめ", "締め日は、いまも毎月25日ですか。", "いまも同じかを確かめるだけ")):
        w = tmp / "form_q"
        shutil.rmtree(w, ignore_errors=True)
        shutil.copytree(good, w)
        edit(w / "flows/F01.flow.json", lambda f: f["questions"][0].update(question=text))
        _code, shown = run("check_flow.py", w, "--stage", "flow")
        out.append((f"確認事項: {name}は、言い回しから拾って WARN にする", any(ln.startswith("WARN") and expect in ln for ln in shown.splitlines())))
    return out


def shown_tests(tmp, good):
    """式のセル（仮の値・合計・ノード一覧の時間と根拠）には、計算の結果も保存する。計算をしない表示（プレビュー・保護ビュー）でも、最初の出力から値が出る。
    判断の仮の値は 2 分（決めの値）。システム処理・別のフローへは 0 分。基準の表の値を人が直したかは、書き出したときの表の値と比べる"""
    import zipfile
    from openpyxl import load_workbook
    import build_matrix as BM
    import office_to_pdf
    import table_template as TT
    out, w, M = [], tmp / "shown", "業務マトリクス"
    shutil.copytree(good, w)
    book = w / "out" / BM.BOOK
    side = TT.side_path(book)

    def read(path=None):
        """(式の本, 保存してある値の本, {(フロー, ノード): 時間のセルの番地})"""
        f, v = load_workbook(path or book), load_workbook(path or book, data_only=True)
        return f, v, {(m.group(1), m.group(2)): c.coordinate for row in f[M].iter_rows() for c in row if c.comment is not None for m in [BM.TAG_TIME.match(c.comment.text)] if m}

    def formulas(f):
        return [(ws.title, c.coordinate) for ws in f for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]

    def listed(v):
        nl = v[BM.NODES]
        col = {c.value: k for k, c in enumerate(nl[1])}
        return {(r[col["フローID"]].value, r[col["ノード番号"]].value): (r[col["処理時間（分）"]].value, r[col["時間の根拠"]].value) for r in nl.iter_rows(min_row=2)}

    def base(wb, action, volume):
        tb = TT.time_table(wb, BM.SPEC["time"])[0]
        return wb[tb["title"]].cell(tb["rows"][action], tb["col"] + tb["cols"][volume] - 1)

    def table():
        d = json.loads((w / "times.json").read_text(encoding="utf-8")) if (w / "times.json").exists() else {}
        return {(x["action"], x["volume"]): x["value"] for x in d.get("table") or []}
    f, v, t = read()
    at = lambda key, wb=None: (wb or v)[M][t[key]].value
    every = formulas(f)
    out.append(("計算の結果: 書き出したブックの式のセル（仮の値・合計・ノード一覧）のすべてに、値も保存してある（計算をしない表示でも、出る）",
                len(every) == 28 + 3 + 2 * 28 and all(v[s][c].value is not None for s, c in every)))
    out.append(("計算の結果: 仮の値は、基準の表のとおり（入力する×小 2・入力する×中 4・確認する×大 15）。判断は 2 分。システム処理・別のフローへは 0 分。合計は、その行の足し算",
                (at(("F01", "P-001")), at(("F01", "P-002")), at(("F03", "P-003")), at(("F03", "D-001")), at(("F03", "S-001")), at(("F01", "SUB-001"))) == (2, 4, 15, 2, 0, 0)
                and [v[M].cell(i, 9).value for i in (3, 5, 7)] == [7, 24, 53] and listed(v)[("F01", "P-002")] == (4, "仮") and listed(v)[("F03", "D-001")] == (2, "仮")))
    office = office_to_pdf.find_soffice()
    if office:
        d = tmp / "shown_calc"
        d.mkdir()
        subprocess.run([office, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(d), str(book)], capture_output=True, timeout=300)
        calc = load_workbook(d / book.name, data_only=True)
        out.append(("計算の結果: 保存してある値が、LibreOffice で計算し直した値と、すべての式のセルで同じ", all(v[s][c].value == calc[s][c].value for s, c in every)))
    else:
        out.append(("計算の結果: 計算し直しての確かめは、飛ばした（LibreOffice が無い）", True))
    # 人が値を入れて、書き出し直す: 保存してある値も、人の値・直した表のとおり
    wb = load_workbook(book)
    wb[M][t[("F01", "P-002")]].value, base(wb, "確認する", "中").value = 20, 6
    wb.save(book)
    c1, _o = run("build_matrix.py", w)
    f, v, t = read()
    soft6 = [k for k in t if re.fullmatch(r'=VLOOKUP\("確認する",.+,3,FALSE\)', str(f[M][t[k]].value))]
    out.append(("計算の結果: 人が時間・基準の表を直して書き出し直すと、保存してある値も、そのとおり（合計・ノード一覧の「人」・直した表の値）",
                c1 == 0 and v[M].cell(3, 9).value == 23 and listed(v)[("F01", "P-002")] == (20, "人") and soft6 and all(at(k, v) == 6 for k in soft6)
                and table() == {("確認する", "中"): 6} and all(v[s][c].value is not None for s, c in formulas(f))))
    # 前の版のブック（判断が 0 分。書き出したときの表の値を、残していない）: 0 を、人の直しと見ない
    wb = load_workbook(book)
    base(wb, "判断する", "中").value = 0
    wb.save(book)
    d = json.loads(side.read_text(encoding="utf-8"))
    kept_table = d.pop("table")
    side.write_text(json.dumps(d, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    c2, _o = run("build_matrix.py", w)
    f, v, t = read()
    out.append(("基準の表: 前の版のブックに残る、判断の 0 分（前の既定の値）は、人が直した値と見ない。書き出し直すと、2 分になる（人が直した表の値は、残る）",
                c2 == 0 and table() == {("確認する", "中"): 6} and base(f, "判断する", "中").value == 2 and at(("F03", "D-001"), v) == 2
                and ["判断する", "中", 2] in json.loads(side.read_text(encoding="utf-8"))["table"] and len(kept_table) == 15 * 4))
    # 書き出したあとで、テンプレートの表が変わった（ブックの値は、書き出したときのまま）: 人の直しと見ない。人が 0 と打てば、その値が残る
    wb = load_workbook(book)
    base(wb, "入力する", "小").value = 7
    wb.save(book)
    d = json.loads(side.read_text(encoding="utf-8"))
    d["table"] = [[a, vol, 7 if (a, vol) == ("入力する", "小") else x] for a, vol, x in d["table"]]
    side.write_text(json.dumps(d, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    c3, _o = run("build_matrix.py", w)
    got3 = (table(), base(load_workbook(book), "入力する", "小").value)
    wb = load_workbook(book)
    base(wb, "判断する", "中").value = 0
    wb.save(book)
    c4, _o = run("build_matrix.py", w)
    c5, _o = run("build_matrix.py", w)
    f, v, t = read()
    out.append(("基準の表: ブックの値が、書き出したときのままなら、テンプレートの値と違っていても、人の直しと見ない（テンプレートの値になる）。人が判断を 0 と打てば、書き出し直しても 0 のまま",
                c3 == 0 and got3 == ({("確認する", "中"): 6}, 2) and c4 == 0 and c5 == 0 and table() == {("確認する", "中"): 6, ("判断する", "中"): 0}
                and base(f, "判断する", "中").value == 0 and at(("F03", "D-001"), v) == 0 and v[M].cell(7, 9).value == sum(at(k, v) for k in t if k[0] == "F03")))
    # 部品: 値の無い式のセルだけに、書き足す（<v/> の書き方・文字を返す式・無いシート）
    z = tmp / "kr.xlsx"
    sheet = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1">'
             '<c r="A1" s="2"><f>SUM(B1:C1)</f><v/></c><c r="B1"><f>1+1</f><v></v></c><c r="C1" t="str"><f>"x"</f><v>x</v></c><c r="D1"><f>IF(A1&gt;0,"仮","人")</f><v></v></c>'
             '<c r="E1"><v>5</v></c></row></sheetData></worksheet>')
    with zipfile.ZipFile(z, "w") as zz:
        zz.writestr("xl/workbook.xml", '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                                       '<sheets><sheet name="表" sheetId="1" r:id="rId1"/></sheets></workbook>')
        zz.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="t" Target="/xl/worksheets/sheet1.xml"/></Relationships>')
        zz.writestr("xl/worksheets/sheet1.xml", sheet)
    left = TT.keep_results(z, {"表": {"A1": 3, "B1": 2.5, "C1": "y", "D1": "仮", "E1": 9, "Z9": 1}, "無いシート": {"A1": 1}})
    with zipfile.ZipFile(z) as zz:
        got = zz.read("xl/worksheets/sheet1.xml").decode("utf-8")
    out.append(("計算の結果: 値の無い式のセルだけに書き足す（数・文字。値のあるセル・式でないセルは、そのまま）。書き足せなかった数を返す",
                '<c r="A1" s="2"><f>SUM(B1:C1)</f><v>3</v></c>' in got and '<c r="B1"><f>1+1</f><v>2.5</v></c>' in got and '<c r="C1" t="str"><f>"x"</f><v>x</v></c>' in got
                and '<c r="D1" t="str"><f>IF(A1&gt;0,"仮","人")</f><v>仮</v></c>' in got and '<c r="E1"><v>5</v></c>' in got and left == 4 and TT.keep_results(z, {}) == 0))
    return out


def title_tests(tmp, good):
    """資料は、納品物では題名で書く（資料ID「M01」は、機械が振った番号なので出さない）。題名は docs[].title → 表紙のカード → ファイル名。点検の記録は ID のまま"""
    import html
    from common import Labels
    from openpyxl import load_workbook
    out, pat, hits = [], re.compile(r"(?<![A-Za-z0-9/])M0\d(?![0-9])"), []
    for p in sorted((good / "out").iterdir()):
        if p.suffix == ".xlsx":
            hits += [f"{p.name}「{ws.title}」{c.coordinate}" for ws in load_workbook(p) for row in ws.iter_rows() for c in row
                     if isinstance(c.value, str) and pat.search(c.value)]
        elif p.suffix in (".drawio", ".md", ".csv"):
            text = p.read_text(encoding="utf-8-sig")
            if p.suffix == ".drawio":                                  # 見える文字（値・ラベル・マウスを乗せると出る文）だけ。データの属性（doc_id など）は除く
                text = " ".join(html.unescape(m.group(1)) for m in re.finditer(r'(?:value|label|tooltip)="([^"]*)"', text))
            hits += [p.name] if pat.search(text) else []
    out.append(("資料は題名で書く: 納品物（業務マトリクス一覧・確認事項一覧・図・出典対応表・階層と業務分類）に、資料ID（M01）が出ない", not hits))
    if hits:
        print("   ", hits[:6])
    t = tmp / "titles"
    shutil.copytree(good, t, ignore=shutil.ignore_patterns("out"))
    m = json.loads((t / "manifest.json").read_text(encoding="utf-8"))
    lab = Labels(t, m)
    before = {d["doc_id"]: lab.title(d["doc_id"]) for d in m["docs"]}

    def put(h):
        for d in h["docs"]:
            d.update({"M02": {"title": "支払の手順"}, "M04": {"title": "経費精算業務マニュアル"}}.get(d["doc_id"], {}))   # M04 は、M01 の表紙の題名と重ねる
    edit(t / "structure/hierarchy.json", put)
    lab2 = Labels(t, m)
    out.append(("資料の題名: 表紙のカードの見出し → ファイル名（拡張子なし）。docs[].title に書けば、それ。重なる題名は、ファイル名を添える",
                before == {"M01": "経費精算業務マニュアル", "M02": "支払処理手順", "M03": "経費精算チェックシート", "M04": "経費精算規程"}
                and lab2.title("M02") == "支払の手順" and lab2.title("M01") == "経費精算業務マニュアル（経費精算業務マニュアル.pptx）"
                and lab2.title("M04") == "経費精算業務マニュアル（経費精算規程.pdf）"))
    lab2.titles["M03"] = "月次集計"
    out.append(("資料の題名: プログラムのモジュールの名前が題名と同じなら、1回だけ書く（シートの名前は、題名に続けて書く）",
                lab2.join("M03", "「月次集計」") == "「月次集計」" and lab2.join("M03", "マクロ「月次集計」") == "マクロ「月次集計」"
                and lab2.join("M03", "「月次集計」シート") == "月次集計「月次集計」シート" and lab2.join("M03", "p.3") == "月次集計 p.3"))
    out.append(("資料ID は、点検の記録（ref）にだけ出る。出典の書き方（src）は題名",
                lab.ref("M04", 2) == "M04 p.1〔M04/p0002〕" and lab.src({"doc_id": "M04", "page": 2}) == "経費精算規程 p.1" and lab.src({"doc_id": "M04"}) == "経費精算規程"))
    w = tmp / "docref"
    shutil.copytree(good, w)
    edit(w / "flows/F03.flow.json", lambda f: f["questions"][0].update(question="M02 の一覧の出力と、領収書原本の照合は、どちらが先か。"))
    code, text = run("check_flow.py", w, "--stage", "flow")
    out.append(("資料を ID で呼ぶ文（確認事項）は、点検が WARN（担当者には通じない。題名で書く）", code == 0 and "資料を ID で呼んでいる（M02）" in text))
    return out


def review_tests(tmp, good):
    """業務マトリクス一覧のレビュー（2026-10-14）への対応: 件数は合計の左。前の形の様式（前の名前のシート・件数が作業の列の右・起動点一覧のシート）でも書き出せ、
    中身の無いプログラム用のシートは入れない。前に書き出した本（件数が右端・列の対応が前の一覧の名前）からも、人が入れた時間・件数が残る"""
    import copy
    from openpyxl import load_workbook
    import build_matrix as BM
    import table_template as TT
    out, M = [], BM.M
    old, w = tmp / "review_old.xlsx", tmp / "review"
    shutil.copytree(good, w)
    book = w / "out" / BM.BOOK
    wb = load_workbook(BM.ASSETS / BM.TEMPLATE)                        # 前の形の様式
    for was, now in BM.WAS_SHEET.items():
        wb[now].title = was
    ws = wb[M]
    k, last = next(c.column for c in ws[2] if c.value == "{件数}"), ws.max_column + 1
    for r in (1, 2):                                                  # 件数を、作業の列の右（行の最後）へ
        ws.cell(r, last).value, ws.cell(r, last)._style = ws.cell(r, k).value, copy.copy(ws.cell(r, k)._style)
    ws.delete_cols(k)
    ent = wb.create_sheet("起動点一覧", 3)
    for j, f in enumerate(BM.ENTRY_HEAD, 1):
        ent.cell(1, j).value, ent.cell(2, j).value = f, "{" + f + "}"
    wb.save(old)
    c1, o1 = run("build_matrix.py", w, "--template", str(old))
    c2, o2 = run("check_flow.py", w, "--stage", "matrix")
    names = load_workbook(book).sheetnames if c1 == 0 else []
    out.append(("前の形の様式（シート名 ノード一覧 など・件数が作業の列の右・起動点一覧のシート）でも書き出せて、点検が通る。中身の無い起動点一覧・プログラム用のシートは、入れない",
                c1 == 0 and "ノード一覧" in names and "起動点一覧" not in names and "処理ノード一覧" not in names and c2 == 0 and "0 FAIL" in o2))
    wb = load_workbook(book)                                          # 人が値を入れる（前の並びの本）
    for row in wb[M].iter_rows():
        for c in row:
            text = c.comment.text if c.comment is not None else ""
            if text.startswith("処理時間 F01 P-002｜"):
                c.value = 20
            elif text.startswith("件数 F01｜"):
                c.value = 120
    wb.save(book)
    side = json.loads(TT.side_path(book).read_text(encoding="utf-8"))   # 前に書き出した本の、列の対応（前の一覧の名前）
    back = {v: k for k, v in BM.WAS_SHEET.items()}
    for s in side["sheets"].values():
        s["canon"] = back.get(s["canon"], s["canon"])
    TT.side_path(book).write_text(json.dumps(side, ensure_ascii=False), encoding="utf-8")
    c3, o3 = run("build_matrix.py", w, "--template", str(BM.ASSETS / BM.TEMPLATE))   # 既定の様式に戻す（付けなければ、前の書き出しの様式を使う）
    c4, o4 = run("check_flow.py", w, "--stage", "matrix")
    f, v = load_workbook(book), load_workbook(book, data_only=True)
    head = [c.value for c in f[M][1]]
    tags = {c.comment.text.split("｜")[0]: c for row in f[M].iter_rows() for c in row if c.comment is not None}
    ci, ti = head.index("件数（人が入れる）"), head.index("処理時間の合計（分・全作業）")
    out.append(("前の並びの本（件数が右端・列の対応が前の一覧の名前）から書き出し直しても、人が入れた時間・件数が残り、件数は合計の左に移る。点検が通る",
                c3 == 0 and "件数 1 行" in o3 and ti == ci + 1 and head[ti + 1] == "①" and v[M][tags["処理時間 F01 P-002"].coordinate].value == 20
                and tags["件数 F01"].value == 120 and tags["件数 F01"].column == ci + 1 and f.sheetnames == [s for s in BM.LOOK["order"] if s not in BM.PROGRAM]
                and c4 == 0 and "0 FAIL" in o4))
    return out


def supply_tests(tmp, good):
    """基準の表（名前「処理時間の表」）の無いテンプレート（前の形）でも、最初の出力から、処理時間が出る: 既定のテンプレートから、基準の表・処理時間の行の書式・処理時間の欄を補う。
    テンプレートのファイルは変えない。table_template.py supply で、補ったテンプレートを保存できる"""
    from openpyxl import load_workbook
    import build_matrix as BM
    import table_template as TT
    out, w, M = [], tmp / "supply", "業務マトリクス"
    shutil.copytree(good, w)
    book, old, new = w / "out" / BM.BOOK, tmp / "supply_old.xlsx", tmp / "supply_new.xlsx"
    wb = load_workbook(BM.ASSETS / BM.TEMPLATE)                        # 前の形のテンプレート: 基準の表・処理時間の書式・処理時間の欄が、どれも無い
    del wb.defined_names[BM.TIME_NAME]
    del wb[BM.TIME_SHEET]
    for title, gone in ((M, ("{処理時間の合計}", "{件数}")), (BM.NODES, ("{動作の種類}", "{量}", "{処理時間（分）}", "{時間の根拠}"))):
        for col in sorted((c.column for c in wb[title][2] if c.value in gone), reverse=True):
            wb[title].delete_cols(col)
    for r in [c.row for c in wb[TT.STYLE_SHEET]["A"] if str(c.value or "").startswith("処理時間")][::-1]:
        wb[TT.STYLE_SHEET].delete_rows(r)
    wb.save(old)
    before = old.read_bytes()

    def read(path=None):
        f, v = load_workbook(path or book), load_workbook(path or book, data_only=True)
        tags = {(m.group(1), m.group(2)): c.coordinate for row in f[M].iter_rows() for c in row if c.comment is not None for m in [BM.TAG_TIME.match(c.comment.text)] if m}
        nl = v[BM.NODES]
        col = {c.value: k for k, c in enumerate(nl[1])}
        listed = {(r[col["フローID"]].value, r[col["ノード番号"]].value): (r[col["動作の種類"]].value, r[col["処理時間（分）"]].value, r[col["時間の根拠"]].value) for r in nl.iter_rows(min_row=2)} if "時間の根拠" in col else {}
        head = [c.value for c in f[M][1]]
        total = head.index("処理時間の合計（分・全作業）") + 1 if "処理時間の合計（分・全作業）" in head else None
        return f, v, tags, listed, head, total
    c1, o1 = run("build_matrix.py", w, "--template", str(old))
    c2, o2 = run("check_flow.py", w, "--stage", "matrix")
    f, v, tags, listed, head, total = read()
    names = f.sheetnames
    out.append(("基準の表の無いテンプレート: 既定のテンプレートから、基準の表（シートと名前）を補って、処理時間の行を出す（1 フロー＝2 行）。補ったことを知らせる。テンプレートのファイルは変えない",
                c1 == 0 and "補って、書き出す" in o1 and "処理時間は扱わない" not in o1 and f[M].max_row == 7 and len(tags) == 28 and BM.TIME_NAME in f.defined_names
                and BM.TIME_SHEET in names and names.index(BM.TIME_SHEET) == names.index("読み方") - 1 and old.read_bytes() == before))
    out.append(("基準の表の無いテンプレート: 処理時間の欄（合計・件数・ノード一覧の 4 つ）も、右端に補う。合計・仮の値が、表のとおり（計算の結果も保存）。点検が通る",
                total is not None and head[-2:] == ["件数（人が入れる）", "処理時間の合計（分・全作業）"] and head.index("①") < total - 2
                and [v[M].cell(i, total).value for i in (3, 5, 7)] == [7, 24, 53] and v[M][tags[("F01", "P-002")]].value == 4
                and listed.get(("F01", "P-002")) == ("入力する", 4, "仮") and c2 == 0 and "0 FAIL / 0 WARN" in o2))
    rules = [r for rng in f[M].conditional_formatting for r in rng.rules]
    t = f[M][tags[("F01", "P-002")]]
    out.append(("基準の表の無いテンプレート: 処理時間の行の書式（塗り・仮の値の条件付き書式）も、既定のテンプレートのものが付く",
                any("ISFORMULA" in x for r in rules for x in r.formula) and t.fill.fill_type == "solid" and t.alignment.horizontal == "center"))
    wb = load_workbook(book)                                          # 人が値を入れても、同じテンプレートで書き出し直せば、残る
    wb[M][tags[("F01", "P-002")]].value = 20
    wb[M].cell(3, total - 1).value = 120
    wb.save(book)
    c3, _o = run("build_matrix.py", w, "--template", str(old))
    f, v, tags, listed, head, total = read()
    d = json.loads((w / "times.json").read_text(encoding="utf-8")) if (w / "times.json").exists() else {}
    out.append(("基準の表の無いテンプレート: 人が入れた時間・件数は、書き出し直しても残る（補った欄でも、同じ）",
                c3 == 0 and v[M][tags[("F01", "P-002")]].value == 20 and v[M].cell(3, total).value == 23 and v[M].cell(3, total - 1).value == 120
                and listed.get(("F01", "P-002")) == ("入力する", 20, "人") and [(x["flow"], x["value"]) for x in d.get("counts") or []] == [("F01", 120)]))
    r1 = subprocess.run([sys.executable, str(S / "table_template.py"), "supply", str(old), "-o", str(new)], capture_output=True, text=True)
    r2 = subprocess.run([sys.executable, str(S / "table_template.py"), "check", str(new), "--book", "matrix"], capture_output=True, text=True)
    r3 = subprocess.run([sys.executable, str(S / "table_template.py"), "supply", str(BM.ASSETS / BM.TEMPLATE), "-o", str(tmp / "supply_none.xlsx")], capture_output=True, text=True)
    w2 = tmp / "supply2"
    shutil.copytree(good, w2)
    c4, o4 = run("build_matrix.py", w2, "--template", str(new))
    c5, o5 = run("check_flow.py", w2, "--stage", "matrix")
    g = load_workbook(w2 / "out" / BM.BOOK, data_only=True)
    tw = TT.Template(new, BM.SPEC)
    out.append(("table_template.py supply: 補ったテンプレートを、別の名前で保存する。それは、基準の表のあるテンプレートとして点検に通り、もう補わない。書き出した合計も同じ。基準の表のあるテンプレートには、何もしない",
                r1.returncode == 0 and new.exists() and old.read_bytes() == before and r2.returncode == 0 and "0 FAIL" in r2.stdout and "補って" not in r2.stdout
                and not tw.supplied and tw.time is not None and not tw.problems and c4 == 0 and "補って" not in o4 and c5 == 0 and "0 FAIL / 0 WARN" in o5
                and [g[M].cell(i, total).value for i in (3, 5, 7)] == [7, 24, 53] and r3.returncode == 0 and "補うものは無い" in r3.stdout and not (tmp / "supply_none.xlsx").exists()))
    part = tmp / "supply_part.xlsx"                                   # 基準の表はあって、件数の欄を消したテンプレート: 消したまま（補うのは、基準の表が無いときだけ）
    wb = load_workbook(BM.ASSETS / BM.TEMPLATE)
    wb[M].delete_cols(next(c.column for c in wb[M][2] if c.value == "{件数}"))
    wb.save(part)
    w3 = tmp / "supply3"
    shutil.copytree(good, w3)
    c6, o6 = run("build_matrix.py", w3, "--template", str(part))
    head6 = [c.value for c in load_workbook(w3 / "out" / BM.BOOK)[M][1]]
    out.append(("基準の表のあるテンプレートで、欄を消してあれば、消したまま出す（補わない）", c6 == 0 and "補って" not in o6 and "件数（人が入れる）" not in head6 and "処理時間の合計（分・全作業）" in head6))
    return out


def big_tests(tmp, good):
    """量「特大」（時間・日の単位でかかる作業）: LLM が、理由つきで付ける。基準の表の 4 列目を引き、まだ仮の値のセルに色が付く。
    前の形のテンプレート（表が 小・中・大 の 3 列）でも止まらない。
    同じ作業名の箱が 2 つあるフロー（F02 の「承認する」）で番号が入れ替わっても、人が入れた値は、担当・直前の箱で、もとの作業に付いていく"""
    from openpyxl import load_workbook
    import build_matrix as BM
    import office_to_pdf
    import table_template as TT
    from common import Times, locate, node_clues, time_basis
    out, w, M = [], tmp / "big", "業務マトリクス"
    shutil.copytree(good, w)
    book, fp = w / "out" / BM.BOOK, w / "flows" / "F03.flow.json"

    def cells(path):
        wb = load_workbook(path)
        return wb, {(m.group(1), m.group(2)): c for row in wb[M].iter_rows() for c in row if c.comment is not None for m in [BM.TAG_TIME.match(c.comment.text)] if m}

    def edit(path, fn):
        d = json.loads(path.read_text(encoding="utf-8"))
        fn(d)
        path.write_text(json.dumps(d, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    box = lambda d: next(n for n in d["nodes"] if n["id"] == "P-005")       # F03 P-005 は「入力する × 中」（システムの画面で作る。決まり 3.22）
    edit(fp, lambda d: box(d).update(volume="特大", volume_reason="月に 1 回、全社の分をまとめて入力する"))
    c1, o1 = run("build_matrix.py", w)
    c2, o2 = run("check_flow.py", w, "--stage", "matrix")
    wb, t = cells(book)
    at = t[("F03", "P-005")].coordinate
    rules = [(str(rng.sqref), r) for rng in wb[M].conditional_formatting for r in rng.rules]
    own = [r for where, r in rules if where == at]
    wide = [r for where, r in rules if ":" in where]
    out.append(("処理時間の量「特大」: 基準の表の 4 列目（特大）を引く式になる。ほかの箱は変わらない。理由が、時間のセルのコメントに出る",
                c1 == 0 and re.fullmatch(r'=VLOOKUP\("入力する",.+,5,FALSE\)', t[("F03", "P-005")].value) is not None
                and re.fullmatch(r'=VLOOKUP\("入力する",.+,3,FALSE\)', t[("F01", "P-002")].value) is not None and "月に 1 回、全社の分をまとめて入力する" in t[("F03", "P-005")].comment.text))
    out.append(("処理時間の量「特大」: まだ仮の値のセルに、そのセルだけの条件付き書式（塗りつき。先に効く）が付く。点検が、先に入れる所として知らせる",
                len(own) == 1 and bool(own[0].stopIfTrue) and own[0].dxf.fill is not None and "ISFORMULA" in own[0].formula[0] and wide and own[0].priority < min(r.priority for r in wide)
                and c2 == 0 and "量が「特大」" in o2 and "1 個" in o2 and "F03 P-005" in o2))
    office = office_to_pdf.find_soffice()
    if office:
        d = tmp / "big_calc"
        d.mkdir()
        subprocess.run([office, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(d), str(book)], capture_output=True, timeout=300)
        ws = load_workbook(d / book.name, data_only=True)[M]
        out.append(("処理時間の量「特大」: 計算し直すと、表の特大の値（入力する 60 分）。合計にも入る", (ws[at].value, ws.cell(7, 9).value) == (60, 109)))
    else:
        out.append(("処理時間の量「特大」: 計算し直しての確かめは、飛ばした（LibreOffice が無い）", True))
    old = tmp / "old3.xlsx"                                           # 前の形のテンプレート: 表が 小・中・大 の 3 列。書式の見本に、特大の行が無い
    tw = load_workbook(BM.ASSETS / BM.TEMPLATE)
    tw[BM.TIME_SHEET].delete_cols(5)
    tw.defined_names[BM.TIME_NAME].attr_text = tw.defined_names[BM.TIME_NAME].attr_text.replace("$E$", "$D$")
    for r in [c.row for c in tw[TT.STYLE_SHEET]["A"] if c.value in ("処理時間（長い作業）", "処理時間（仮・長い作業）")][::-1]:
        tw[TT.STYLE_SHEET].delete_rows(r)
    tw.save(old)
    c3, o3 = run("build_matrix.py", w, "--template", str(old))
    c4, o4 = run("check_flow.py", w, "--stage", "matrix")
    wb3, t3 = cells(book)
    out.append(("処理時間の量「特大」: 前の形のテンプレート（表が 3 列・特大の書式が無い）でも止まらない。特大の作業は「大」の列を引き、そう知らせる",
                c3 == 0 and "「特大」の列が無い" in o3 and re.fullmatch(r'=VLOOKUP\("入力する",.+,4,FALSE\)', t3[("F03", "P-005")].value) is not None and c4 == 0 and "FAIL  " not in o4))
    edit(fp, lambda d: box(d).pop("volume_reason"))
    c5, o5 = run("check_flow.py", w, "--stage", "flow")
    same = [{"type": "process", "sources": [{"doc_id": "M01", "page": 3, "item": f"s{k}", "quote": "x"} for k in range(1, 9)], "action": "作成する"}]
    out.append(("処理時間の量「特大」: 理由の無い特大は FAIL。機械は、特大を付けない（同じページの手順を 8 つまとめた箱も「大」）",
                c5 != 0 and "理由（volume_reason）が無い" in o5 and time_basis(same[0])[1] == "大" and time_basis(dict(same[0], volume="特大", volume_reason="決算の書類一式"))[1:] == ("特大", "決算の書類一式")))

    # 同じ作業名の箱（F02 P-004 課長「承認する」・P-006 部長「承認する」）: 番号が入れ替わっても、値が、もとの作業に付いていく
    def swap(work):
        f = work / "flows" / "F02.flow.json"
        s = f.read_text(encoding="utf-8")
        f.write_text(s.replace('"P-004"', '"@@"').replace('"P-006"', '"P-004"').replace('"@@"', '"P-006"'), encoding="utf-8")

    def stored(work):
        d = json.loads((work / "times.json").read_text(encoding="utf-8"))
        return {(x["flow"], x["node"]): (x["value"], x.get("lane")) for x in d["cells"]}

    def put(path, values):
        wb, t = cells(path)
        for k, v in values.items():
            t[k].value = v
        wb.save(path)
    a = tmp / "twins"
    shutil.copytree(good, a)
    put(a / "out" / BM.BOOK, {("F02", "P-004"): 4, ("F02", "P-006"): 11})
    run("build_matrix.py", a)
    before, kept = stored(a), tmp / "twins_old.xlsx"
    shutil.copy(a / "out" / BM.BOOK, kept)                           # 入れ替える前に書き出したブック（お客様の手元に残っている、という見立て）
    swap(a)
    c6, o6 = run("build_matrix.py", a)
    c7, o7 = run("check_flow.py", a, "--stage", "matrix")
    out.append(("処理時間: 同じ作業名の箱が 2 つあるフローで、番号が入れ替わっても、人が入れた値は、担当の同じ作業に付いていく（times.json に、担当・直前の箱を書いてある）",
                before == {("F02", "P-004"): (4, "課長"), ("F02", "P-006"): (11, "部長")} and c6 == 0 and o6.count("付け替えた") == 2
                and stored(a) == {("F02", "P-006"): (4, "課長"), ("F02", "P-004"): (11, "部長")} and c7 == 0 and "FAIL  " not in o7))
    b = tmp / "twins_keep"
    shutil.copytree(good, b)
    swap(b)
    run("build_matrix.py", b)
    c8, o8 = run("build_matrix.py", b, "--keep", str(kept))
    out.append(("処理時間: 番号が入れ替わる前に書き出したブックからの取り込み（--keep）でも、時間のセルのコメントの担当・直前の箱で、同じ作業に付く",
                c8 == 0 and stored(b) == {("F02", "P-006"): (4, "課長"), ("F02", "P-004"): (11, "部長")}))
    # 絞り方そのもの: 担当 → 直前の箱。手がかりの無い値・決められない値は、動かさない
    node = lambda i, label, lane: {"id": i, "type": "process", "label": label, "lane": lane}
    flow = lambda nodes, order: {"flow_id": "F09", "lanes": [{"id": "L1", "name": "係"}, {"id": "L2", "name": "課長"}],
                                 "nodes": [{"id": "start", "type": "start"}] + nodes + [{"id": "end", "type": "end"}],
                                 "edges": [{"from": x, "to": y} for x, y in zip(["start"] + order, order + ["end"])]}
    was = flow([node("P-001", "起案する", "L1"), node("P-002", "押印する", "L1"), node("P-003", "回付する", "L1"), node("P-004", "押印する", "L1"), node("P-005", "押印する", "L2")],
               ["P-001", "P-002", "P-003", "P-004", "P-005"])
    now = flow([node("P-001", "起案する", "L1"), node("P-004", "押印する", "L1"), node("P-003", "回付する", "L1"), node("P-005", "押印する", "L1"), node("P-002", "押印する", "L2")],
               ["P-001", "P-004", "P-003", "P-005", "P-002"])
    ts = Times(tmp / "no_such_work")
    for i, v in (("P-002", 1), ("P-004", 2), ("P-005", 3)):
        ts.cells[("F09", i)] = {"flow": "F09", "node": i, "label": "押印する", "value": v}
    bare = Times(tmp / "no_such_work")
    bare.cells = {k: dict(x) for k, x in ts.cells.items()}
    ts.note([was])
    moved, still = ts.reconcile([now]), bare.reconcile([now])
    tie = dict(node_clues([now])[("F09", "P-004")][0:0] or {}, label="押印する", lane="係", after=["どこにも無い箱"])
    out.append(("処理時間: 同じ作業名の箱が 3 つ（うち 2 つは同じ担当）でも、担当 → 直前の箱の順に絞って、付け替える。手がかりの無い値・絞り切れない値は、動かさない",
                {k[1]: x["value"] for k, x in ts.cells.items()} == {"P-004": 1, "P-005": 2, "P-002": 3} and len(moved) == 3 and still == []
                and {k[1]: x["value"] for k, x in bare.cells.items()} == {"P-002": 1, "P-004": 2, "P-005": 3}
                and locate(node_clues([now]), "F09", "P-009", tie) is None and locate(node_clues([now]), "F09", "P-004", tie) == "P-004"
                and locate(node_clues([now]), "F09", "P-009", {"label": "回付する"}) == "P-003"))
    # Excel で保存したブック: 式のシート名の、要らない引用符が外れる（'処理時間の基準'!$A$2 → 処理時間の基準!$A$2）。書き換えた、と見ない
    x = tmp / "resaved"
    shutil.copytree(good, x)
    wbx, n = load_workbook(x / "out" / BM.BOOK), 0
    for sh in wbx.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and "'" in c.value:
                    c.value, n = c.value.replace("'", ""), n + 1
    wbx.save(x / "out" / BM.BOOK)
    c9, o9 = run("check_flow.py", x, "--stage", "matrix")
    c10, _o10 = run("build_matrix.py", x)
    out.append(("処理時間: Excel で保存したブック（式のシート名の引用符が外れる）を、点検が FAIL にしない。仮の値の式を、人が入れた値と取り違えない",
                n >= 28 * 3 and c9 == 0 and "FAIL  " not in o9 and c10 == 0 and not (x / "times.json").exists()))
    return out


def volume_tests():
    """処理時間の量の決め方: まとめた手順は、同じページの手順（s1・s2…）の数。出典の数ではない"""
    from common import time_basis
    src = lambda *xs: [{"doc_id": d, "page": p, "item": i, "quote": "x"} for d, p, i in xs]
    box = lambda sources, **more: dict({"type": "process", "action": "確認する", "sources": sources}, **more)
    same = time_basis(box(src(("M01", 3, "s1"), ("M01", 3, "s2"), ("M01", 3, "s3"))))
    spread = time_basis(box(src(("M01", 2, "s1"), ("M01", 3, "s4"), ("M02", 1, "s1")) + [{"doc_id": "M01", "page": 3, "quote": "注記"}]))
    mixed = time_basis(box(src(("M01", 3, "s1"), ("M01", 3, "s2"), ("M01", 3, "g1"), ("M01", 3, "d1"), ("M01", 2, "s9"))))
    return [("処理時間の量: 同じページの手順を 3 つまとめた箱は「大」。同じ作業が、概要・図・詳しい手順の 3 か所に書かれているだけの箱は「中」（図の箱・分岐・注記の出典は、数えない）",
             same == ("確認する", "大", "手順 3 つ") and spread == ("確認する", "中", "") and mixed == ("確認する", "中", "")),
            ("処理時間の量: 見る帳票が 2 つ以上の箱は「大」。理由つきの指定が、機械の決め方より先。判断・システム処理は、機械が決める種類の「中」",
             time_basis(box([]), 2) == ("確認する", "大", "見る帳票 2 つ") and time_basis(box([]), 1) == ("確認する", "中", "") and time_basis(box(src(("M01", 3, "s1"), ("M01", 3, "s2"), ("M01", 3, "s3")), volume="小", volume_reason="押すだけ")) == ("確認する", "小", "押すだけ")
             and time_basis({"type": "decision", "sources": src(("M01", 3, "s1"), ("M01", 3, "s2"), ("M01", 3, "s3"))}) == ("判断する", "中", "")
             and time_basis({"type": "process"}) == ("その他", "中", ""))]


def unit_tests():
    """部品が単体で正しく動くか"""
    from common import Report, bare_number, printed_numbers, question_ids, quote_match
    from validate import validate
    import pages
    out = column_tests()
    tree = [{"id": "H1", "title": "経費精算規程", "level": 1}, {"id": "H1.1", "title": "予算", "level": 2}, {"id": "H1.2", "title": "決算", "level": 2},
            {"id": "H1.2.1", "title": "予算", "level": 3}]
    toc = [{"title": "予算", "level": 1}, {"title": "決算", "level": 1}, {"title": "予算", "level": 2}]
    got = pages.pair_up(tree, toc)
    from common import action_like
    acts = ["申請書の内容を確かめる。", "書類を総務へ送る。", "領収書を封筒に入れる。", "月末に締めの処理を行う。"]
    nons = ["金額は十万円までとする。", "承認が必要である。", "入力欄が表示されている。", "担当者が提出できる。", "内容に誤りがある。", "手続きは規程による。"]
    out.append(("作業らしい文: 「る」で終わる作業（確かめる・送る・入れる）を拾い、状態・可能・受け身・決まりの言い回しは拾わない",
                all(action_like(s) for s in acts) and not any(action_like(s) for s in nons)))
    out.append(("目次と業務の木: 木の根を 1 段目に置いた木（目次は章が 1 段目）でも、段の差をそろえて突き合わせる",
                got.get("H1.1") is toc[0] and got.get("H1.2") is toc[1] and got.get("H1.2.1") is toc[2] and "H1" not in got))
    out.append(("刷られた番号の飾りを外す（- 13 -・P.13・13/20。3-2 はそのまま）",
                [bare_number(x) for x in ("- 13 -", "P.13", "13/20", "13 ページ", "3-2", "ii")] == ["13", "13", "13", "13", "3-2", "ii"]))
    deco = printed_numbers([{"page": i, "cands": [], "card": f"- {i + 1} -"} for i in range(1, 8)])
    out.append(("飾りつきで読んだ番号（- 3 -）も番号と見る", deco.get(2) == ("3", "card")))
    far = printed_numbers([{"page": i, "cands": [{"n": str(i - 2), "x": 0.5, "y": 0.05}] if i in (3, 7, 11) else [], "card": None} for i in range(1, 13)])
    out.append(("離れたページの数字を、番号の並びにしない", far == {}))
    src = "上長は、申請内容と領収書を照合する。申請金額が5万円未満の場合は課長が承認する。5万円以上の場合は部長が承認する。不備がある場合は、申請者に差し戻す。"
    out.append(("原典の連続した一部は一致 1、離れた断片をつないだ文は 1 にしない",
                quote_match("5万円以上の場合は部長が承認する。", src) == 1.0 and quote_match("上長は、申請者に差し戻す。", src) <= 0.5
                and quote_match("5万円未満の場合は部長が承認する。", src) < 1.0))
    pg = [{"page": i, "cands": ([{"n": str(i - 1), "x": 0.5, "y": 0.96}] if i not in (1, 7) else []) + [{"n": "2026", "x": 0.9, "y": 0.04}], "card": None}
          for i in range(1, 13)]
    got = printed_numbers(pg)
    out.append(("表紙を数えない資料の番号（ファイルの番号 − 1）", got.get(14 - 2) == ("11", "machine") and 1 not in got and got.get(7) == ("6", "inferred")))
    out.append(("番号でない数字（年）を番号にしない", all(v[0] != "2026" for v in got.values())))
    scan = printed_numbers([{"page": i, "cands": [], "card": str(i + 2) if i != 4 else None} for i in range(1, 8)])
    out.append(("スキャン資料はカードの読みから決める", scan.get(1) == ("3", "card") and scan.get(4) == ("6", "inferred")))
    out.append(("番号の無い資料は空", printed_numbers([{"page": i, "cands": [], "card": None} for i in range(1, 5)]) == {}))
    chap = printed_numbers([{"page": i, "cands": [{"n": f"3-{i}", "x": 0.5, "y": 0.96}], "card": None} for i in range(1, 7)])
    out.append(("「3-2」の形の番号", chap.get(2) == ("3-2", "machine")))
    rp = Report("t", fold=3)
    for i in range(9):
        rp.fail(f"M01 p.{i}〔M01/p{i:04d}〕 手順 s{i}「あ{i}」: 引用が機械抽出の文字に見当たらない（一致 {i}0%）: 「い{i}」")
    out.append(("同じ種類の行を畳む（9 件 → 3 行＋件数）", len(rp.lines(3)) == 6 and len(rp.lines()) == 11))
    out.append(("余分な項目を違反にする", any("スキーマに無い項目" in e for e in validate({"terms": [], "questions": [], "memo": 1}, "glossary"))))
    # 章と節が同じ名前の目次（第5章 出張手続／1 出張手続）と、名前が含み合う分類（出張手続・海外出張手続）
    from check_flow import toc_category
    from pages import named_refs
    cats = [{"id": "F05", "number": "第5章", "name": "出張手続", "pages": [], "make_flow": False, "hierarchy_id": "H5", "parent": None},
            {"id": "F05-1", "number": "第5章 1", "name": "出張手続", "pages": [{"doc_id": "M01", "pages": [1]}], "make_flow": True, "hierarchy_id": "H5.1", "parent": "F05"},
            {"id": "F06", "number": "第6章", "name": "備品の購入", "pages": [], "make_flow": False, "hierarchy_id": "H6", "parent": None},
            {"id": "F06-1", "number": "第6章 1", "name": "備品の購入", "pages": [{"doc_id": "M01", "pages": [2]}], "make_flow": True, "hierarchy_id": "H6.1", "parent": "F06"},
            {"id": "F07-1", "number": "第7章 1", "name": "海外出張手続", "pages": [{"doc_id": "M01", "pages": [3]}], "make_flow": True, "hierarchy_id": "H7.1", "parent": None}]
    pick = lambda **e: (toc_category(e, cats) or {}).get("id")
    out.append(("目次の項目を、同じ名前の章でなく、ページを持つ節に当てる", pick(number="第5章 1", title="出張手続") == "F05-1" and pick(title="出張手続") == "F05-1"
                and pick(number="第7章 1", title="海外出張手続") == "F07-1" and pick(title="無い業務") is None))
    with tempfile.TemporaryDirectory() as t:
        w = pathlib.Path(t)
        (w / "cards" / "M01").mkdir(parents=True)
        (w / "cards" / "M01" / "p0001.json").write_text(json.dumps({"cross_refs": [
            {"text": "（備品を買う場合 → 第6 章「備品の購入」）", "target": "第6章 備品の購入"},
            {"text": "海外出張手続と同じ", "target": ""}]}, ensure_ascii=False), encoding="utf-8")
        got = named_refs(w, {"docs": [{"doc_id": "M01", "page_list": [{"page": 1}]}]}, {"categories": cats}, {("M01", 1): [cats[1]]}, set())
    out.append(("番号と名前で名指しした参照だけを「続く」の候補にする（章を指せば、その下のフローへ）", [(x["from"], x["to"]) for x in got] == [("H5.1", "H6.1")]))
    # 1件の確認事項が複数のノードを指す（nodes）
    many = {"id": "Q-001", "flow": "F01", "node": "P-001", "nodes": ["P-002", "P-001"]}
    out.append(("確認事項の対象（node と nodes）を、重なりなく並べる", q_targets(many) == ["P-001", "P-002"] and q_targets({"edge": ["a", "b"]}) == []
                and question_ids([many], "F01")[0] == {"P-001": ["Q-001"], "P-002": ["Q-001"]}))
    from renumber import renumber
    fl = {"nodes": [{"id": "start", "type": "start"}, {"id": "a", "type": "process"}, {"id": "b", "type": "process"}, {"id": "end", "type": "end"}],
          "edges": [{"from": "start", "to": "a"}, {"from": "a", "to": "b"}, {"from": "b", "to": "end"}], "questions": [{"id": "q1", "nodes": ["b", "a"]}]}
    renumber(fl)
    out.append(("番号の振り直しが、確認事項の nodes にも及ぶ", fl["questions"][0]["nodes"] == ["P-002", "P-001"]))
    from common import step_order_by_branch
    ns = lambda *ids: [{"id": i, "type": "start" if i == "start" else "end" if i.startswith("end") else "process"} for i in ids]
    es = lambda *pairs: [{"from": a_, "to": b_} for a_, b_ in pairs]
    fork = {"nodes": ns("start", "d", "a1", "b1", "a2", "b2", "m", "end"),
            "edges": es(("start", "d"), ("d", "a1"), ("d", "b1"), ("a1", "a2"), ("b1", "b2"), ("a2", "m"), ("b2", "m"), ("m", "end"))}
    loop = {"nodes": ns("start", "a", "b", "c", "end"), "edges": es(("start", "a"), ("a", "b"), ("b", "c"), ("c", "a"), ("c", "end"))}
    lost = {"nodes": ns("start", "a", "x", "end"), "edges": es(("start", "a"), ("a", "end"))}
    out.append(("表の列の順: 枝を続けて並べ、合流はどの枝よりもあと。戻りの線は数えず、開始から届かないノードは末尾",
                step_order_by_branch(fork) == ["start", "d", "a1", "a2", "b1", "b2", "m", "end"]
                and step_order_by_branch(loop) == ["start", "a", "b", "c", "end"] and step_order_by_branch(lost) == ["start", "a", "end", "x"]))
    from common import infer_nodes
    fl = {"nodes": [{"id": "P-001", "type": "process", "sources": [{"item": "s1"}], "inferred": True},
                    {"id": "D-001", "type": "decision", "sources": [{"quote": "x"}]},
                    {"id": "SUB-001", "type": "subprocess", "sources": [{"quote": "x"}]}]}
    out.append(("推定の印: カードの項目を引く作業のノードは外し、引かない作業のノードに付ける（リンクには触らない）",
                infer_nodes(fl) == [("P-001", False), ("D-001", True)] and "inferred" not in fl["nodes"][0] and fl["nodes"][1]["inferred"] is True
                and "inferred" not in fl["nodes"][2] and infer_nodes(fl) == []))
    from common import sentences, span_text
    sys.path.insert(0, str(S))
    from quotes import find_span
    pj = {"text": ["手順を次に示す", "\uf06e", "\uf09e 台帳を確認します。", "長い文の途中で折れた行がここにあって", "\uf06c次の項目を確認する。", "概要"]}
    out.append(("文の単位: 外字（フォントごとの箇条書きの印）は落とし、行の頭の外字で切る",
                sentences(pj) == ["手順を次に示す", "台帳を確認します。", "長い文の途中で折れた行がここにあって", "次の項目を確認する。", "概要"]))
    units = ["見出し", "経理部は、申請を確認し、", "承認する。", "なお、手数料は会社が負担する。"]
    out.append(("行の番号: 範囲の文字をつなぐ。範囲の外・逆順・数でないものは None",
                span_text(units, [2, 3]) == "経理部は、申請を確認し、承認する。" and span_text(units, [4, 4]) == units[3]
                and all(span_text(units, x) is None for x in ([0, 1], [3, 2], [4, 5], [1], "1-2", [True, 1]))))
    out.append(("行の番号: 引用を含む、いちばん短い範囲を引く（無ければ None）",
                find_span(units, "申請を確認し、承認する") == [2, 3] and find_span(units, "承認する。") == [3, 3] and find_span(units, "作り話") is None))
    from check_flow import placeholder
    out.append(("原典の語「不明」は仮置きの語と見ない（引用に無ければ見る）",
                placeholder({"label": "調査しても原因が不明か", "sources": [{"quote": "原因を調べても不明なときは、次の処理を行います。"}]}) is None
                and placeholder({"label": "担当不明の処理", "sources": [{"quote": "経理部は、支払データを作成する。"}]}) == "不明"))
    from check_layout import stale_pngs
    with tempfile.TemporaryDirectory() as td:
        td = pathlib.Path(td)
        (td / "x.layout.json").write_text(json.dumps({"file": "図.drawio", "pages": [{"index": 1, "id": "overview"}, {"index": 2, "id": "flow-F01"}]}), encoding="utf-8")
        for name in ("図_p01_overview.png", "図_p02_flow-F01.png", "図_p02_flow-F02.png", "図_p03_flow-F01.png", "別の図_p02_flow-F02.png", "メモ.png"):
            (td / name).write_bytes(b"")
        out.append(("ページが増減したあとに残る古い PNG だけを拾う",
                    [q.name for q in stale_pngs(td, [td / "x.layout.json"])] == ["図_p02_flow-F02.png", "図_p03_flow-F01.png"]))
    from build_drawio import node_attrs
    at = node_attrs({"id": "P-001", "type": "process", "label": "見積を徴取する", "reference": "10万円未満は1社以上"}, "経理部", [], {}, None)
    out.append(("箱に収まらない参考情報（reference）は、マウスを乗せたときの説明に出る", "参考: 10万円未満は1社以上" in at.get("tooltip", "")))
    from check_flow import lineage, tied
    cats = {"F08": {"parent": None}, "F08-3": {"parent": "F08"}, "F08-3-3": {"parent": "F08-3"}}
    out.append(("下位フローの上位は、フローを作る分類だけをたどる",
                lineage(cats, {"F08-3", "F08-3-3"}, "F08-3-3") == ["F08-3-3", "F08-3"] and lineage(cats, {"F08-3"}, "F09") == []))
    fl = {"F08-3", "F08-3-3", "F09"}
    cats["F09"] = {"parent": None}
    out.append(("上位フローの「続く」は、下位フローに置いたリンクで足りる（逆向き・無関係は足りない）",
                tied(cats, fl, "F08-3-3", "F09", {("F08-3", "F09")}) and not tied(cats, fl, "F09", "F08-3-3", {("F08-3", "F09")})
                and not tied(cats, fl, "F08-3", "F09", {("F08-3-3", "F09")})))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-w", "--work", required=True, help="fixture を描画済みの作業フォルダ")
    ap.add_argument("--keep", help="正しいデータで作った一式（図・報告書）を、このフォルダに残す")
    a = ap.parse_args()
    base = pathlib.Path(a.work).resolve()
    if not (base / "manifest.json").exists() or not (base / "pages").exists():
        sys.exit("manifest.json / pages が無い。先に inventory.py と render_pages.py を fixture に対して実行する")
    tmp = pathlib.Path(tempfile.mkdtemp(prefix="asisflow_selftest_"))
    ok, skipped = True, 0
    try:
        good = tmp / "good"
        shutil.copytree(base, good, ignore=shutil.ignore_patterns("out", "cards", "flows", "structure", "reports", "review", "taxonomy.json",
                                                                 "glossary.json", "questions.json", "batches.json", "*.pdf"))
        shutil.copytree(SKILL / "evals" / "fixture_expected", good, dirs_exist_ok=True)
        steps = [("validate.py", []), ("renumber.py", ["--dry-run"]), ("build_drawio.py", []), ("check_extract.py", []), ("check_flow.py", []),
                 ("check_layout.py", ["--no-png"]), ("build_reports.py", []), ("build_matrix.py", []), ("pages.py", ["pagemap"]), ("status.py", [])]
        for script, args in steps:
            code, out = run(script, good, *args)
            passed = code == 0 and "0 FAIL" in out if script.startswith("check") else code == 0
            if script == "status.py":                                # 何を表示しても終了コードは 0。次の一手が合っているかを文字で見る
                passed = passed and "make_compare.py" in out and "要直し" not in out
            if script == "renumber.py":
                passed = passed and "フロー 0 本" in out                # 期待データの番号・推定の印は、機械が付けたものと同じ
            print(f"{'OK  ' if passed else 'NG  '} 正しいデータ: {script}")
            if not passed:
                ok = False
                print(out[-1500:])
        code, out = run("validate.py", good, "--check-doc")
        print(f"{'OK  ' if code == 0 else 'NG  '} schemas.md の項目表がスキーマと合っている")
        ok = ok and code == 0
        # 納品物の表記
        src = (good / "out" / "出典対応表.md").read_text(encoding="utf-8")
        drawio = "".join(p.read_text(encoding="utf-8") for p in (good / "out").glob("*.drawio"))
        looks = [
            ("表紙を数えない PDF は、刷られた番号で出る（M04 のファイル2枚目 → p.1）", "経費精算規程 p.1" in src and "経費精算規程 p.1" in drawio),
            ("ファイルの何枚目かは納品物に出ない（M04 p.3 が無い）", "経費精算規程 p.3" not in src and "経費精算規程 p.3" not in drawio and "p0002" not in src),
            ("PowerPoint はスライド番号", "経費精算業務マニュアル スライド 3" in src and "経費精算業務マニュアル p.3" not in src),
            ("Excel はシート名とセル", "経費精算チェックシート「手順」シート C2" in src and "経費精算チェックシート「手順」シート C2" in drawio),
            ("Excel の図形は、置かれたセル", "経費精算チェックシート「フロー」シート B3" in src),
            ("補足由来の手順に印が付く", "| 補足 |" in src),
            ("入れなかった手順と理由が載る", "支払日が手順書" in src),
            ("階層と業務分類に、結び付きと根拠が載る", "詳しくしている" in (good / "out" / "階層と業務分類.md").read_text(encoding="utf-8")),
            ("まとめた確認事項は、対象のノードを並べて載る", "| P-001・P-002 |" in (good / "out" / "確認事項一覧.md").read_text(encoding="utf-8")),
        ]
        for name, passed in looks + matrix_looks(good) + title_tests(tmp, good) + axis_tests(tmp) + heading_tests(tmp) + pointer_tests(tmp, good) + draft_tests(tmp) + form_tests(tmp, good) + unit_tests() + code_tests() + vba_stmt_tests() + io_fix_tests() + lang_tests() + code_pipeline_tests(tmp) + call_run_tests(tmp) + narrow_tests(tmp) + template_tests(tmp, good) + time_tests(tmp, good) + shown_tests(tmp, good) + supply_tests(tmp, good) + review_tests(tmp, good) + big_tests(tmp, good) + volume_tests() + access_tests(tmp) + direct_tests(tmp) + sample_fix_tests(tmp) + tree_tests(tmp) + tree_rule_tests(tmp) + review_fix_tests(tmp) + schema_tests(tmp, good) + review_gate_tests(tmp, good) + output_check_tests(tmp, good) + answer_tests(tmp, good) + gate_tests(tmp, good) + mech_tests(tmp, good):
            print(f"{'OK  ' if passed else 'SKIP' if passed is None else 'NG  '} {name}")   # None は、道具が無くて飛ばした（合格に数えない）
            ok = ok and passed is not False
            skipped += passed is None
        # 揺れの測定（diff_runs.py）: 同じものどうしは 100%。手順を1つ落とし、種別を変えた写しとは差が出る
        other = tmp / "other"
        shutil.copytree(good, other)
        r = subprocess.run([sys.executable, str(S / "diff_runs.py"), "cards", str(good), str(other), "--min", "1.0"], capture_output=True, text=True)
        same_ok = r.returncode == 0
        edit(other / "cards/M01/p0003.json", lambda c: c.update(page_type="overview", steps=c["steps"][:-1]))
        edit(other / "flows/F01.flow.json", lambda f: f["lanes"][0].update(name="申請部門"))
        r1 = subprocess.run([sys.executable, str(S / "diff_runs.py"), "cards", str(good), str(other), "--min", "1.0"], capture_output=True, text=True)
        r2 = subprocess.run([sys.executable, str(S / "diff_runs.py"), "flows", str(good), str(other), "--min", "1.0"], capture_output=True, text=True)
        diff_ok = same_ok and r1.returncode == 1 and "ページの種別が違う" in r1.stdout and "片方だけにある手順" in r1.stdout \
            and r2.returncode == 1 and "レーンが違う" in r2.stdout
        print(f"{'OK  ' if diff_ok else 'NG  '} 揺れの測定: 同じものは 100%、種別・手順・レーンの違いを数える")
        ok = ok and diff_ok
        # 進み具合: 棚卸しだけのフォルダでは、ページ化を案内する。点検が動かない（壊れたカード）ときは「済み」にしない
        inv = tmp / "inv"
        inv.mkdir()
        m = json.loads((good / "manifest.json").read_text(encoding="utf-8"))
        (inv / "manifest.json").write_text(json.dumps(dict(m, docs=[dict(d, rendered=False, page_list=[]) for d in m["docs"]]), ensure_ascii=False),
                                           encoding="utf-8")
        _c, out_inv = run("status.py", inv)
        brk = tmp / "brk"
        shutil.copytree(good, brk)
        (brk / "cards/M02/p0001.json").write_text('{"doc_id": "M02",', encoding="utf-8")
        _c, out_brk = run("status.py", brk)
        st_ok = "render_pages.py" in out_inv and "ページ化     途中" in out_inv and "4 読取       要直し" in out_brk and "build_drawio.py" not in out_brk
        print(f"{'OK  ' if st_ok else 'NG  '} 進み具合: 棚卸しの直後はページ化を案内する。壊れたカードがあれば「要直し」で、書き出しを案内しない")
        if not st_ok:
            print(out_inv[-600:], out_brk[-900:])
        ok = ok and st_ok
        # 目次の番号をそのまま割当に使った冊子: 節が2ページ以上あると、目次の指すページは割当に入ったまま。見出しがそろってずれることで見つける
        make_book(tmp / "book")
        make_book(tmp / "book_shift", shift=-1)
        code0, out0 = run("check_flow.py", tmp / "book", "--stage", "structure")
        code1, out1 = run("check_flow.py", tmp / "book_shift", "--stage", "structure")
        book_ok = code0 == 0 and "0 FAIL / 0 WARN" in out0 and code1 == 1 and "系統的なずれ — 目次に準拠した分類 3 件" in out1 \
            and not any(ln.startswith("FAIL") and "が指すのは" in ln for ln in out1.splitlines())
        print(f"{'OK  ' if book_ok else 'NG  '} 目次の番号をそのまま割当に使う（そろって1ページずれる）: 正しい割当は FAIL 0、ずらすと「系統的なずれ」")
        if not book_ok:
            print(out0[-800:], out1[-1200:])
        ok = ok and book_ok
        # 用語の正本の下書き（pages.py terms --draft）: 寄せずに全部を載せる。もう一度実行しても増えない。すでにある用語には触らない
        dr = tmp / "draft"
        shutil.copytree(good, dr)
        (dr / "glossary.json").unlink()
        run("pages.py", dr, "terms", "--draft")
        g1 = json.loads((dr / "glossary.json").read_text(encoding="utf-8"))
        code, out = run("check_flow.py", dr, "--stage", "structure")
        edit(dr / "glossary.json", lambda g: g["terms"][0].update(aliases=["別の呼び方"], evidence="「A（以下、別の呼び方）」M01 p.1"))
        run("pages.py", dr, "terms", "--draft")
        g2 = json.loads((dr / "glossary.json").read_text(encoding="utf-8"))
        draft_ok = bool(g1["terms"]) and all(t["aliases"] == [] for t in g1["terms"]) and code == 0 and len(g2["terms"]) == len(g1["terms"]) \
            and g2["terms"][0]["aliases"] == ["別の呼び方"]
        print(f"{'OK  ' if draft_ok else 'NG  '} 用語の正本の下書き: 寄せずに {len(g1['terms'])} 件を載せ、点検が通る。やり直しても、ある用語には触らない")
        ok = ok and draft_ok
        # 進行役（run.py）: 機械の工程を回し、LLM の出番（原典照合）で止まる。推定の印を消したフローは、番号の振り直しで付け直してから書き出す
        auto = tmp / "auto"
        shutil.copytree(good, auto)
        edit(auto / "flows/F03.flow.json", lambda f: node(f, "D-002").pop("inferred"))
        r = subprocess.run([sys.executable, str(S / "run.py"), "-w", str(auto)], capture_output=True, text=True)
        fixed = node(json.loads((auto / "flows/F03.flow.json").read_text(encoding="utf-8")), "D-002").get("inferred") is True
        run_ok = r.returncode == 0 and fixed and "次は LLM・人の仕事" in r.stdout and "要直し" not in r.stdout \
            and list((auto / "out").glob("照合シート_*.pdf")) != [] and (auto / "out" / "業務マトリクス一覧.xlsx").exists()
        print(f"{'OK  ' if run_ok else 'NG  '} 進行役（run.py）: 推定の印を付け直し、図・業務マトリクス一覧・照合シートまで作って、原典照合（LLM の出番）で止まる")
        if not run_ok:
            print(r.stdout[-1500:], r.stderr[-500:])
        ok = ok and run_ok
        if a.keep:
            keep = pathlib.Path(a.keep).expanduser().resolve()
            shutil.copytree(good, keep, dirs_exist_ok=True)
            print(f"正しいデータの一式 → {keep}")
        for name, script, args, mutate, expect in CASES:
            w = tmp / "case"
            if w.exists():
                shutil.rmtree(w)
            shutil.copytree(good, w)
            mutate(w)
            code, out = run(script, w, "--all-lines", *args)
            hit = [ln for ln in out.splitlines() if ln.startswith("FAIL") and expect in ln]
            passed = code == 1 and bool(hit)
            print(f"{'OK  ' if passed else 'NG  '} {name}: " + (hit[0][:110] if hit else f"FAIL「{expect}」が出ない（終了コード {code}）"))
            if not passed:
                ok = False
                print("\n".join(ln for ln in out.splitlines() if ln.startswith(("FAIL", "WARN")))[:1200])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("\nselftest " + ("OK" if ok else "NG") + (f"（飛ばした項目 {skipped}。合格には数えていない）" if skipped else ""))
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
