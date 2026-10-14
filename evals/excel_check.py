#!/usr/bin/env python3
"""excel_check.py — 業務マトリクス一覧の処理時間が、Microsoft Excel でも、決めたとおりに動くかを確かめる。
macOS と Excel が要る（AppleScript で動かす）。自己点検（selftest.py）には入っていない。LibreOffice での確かめは、自己点検の側にある。

  python3 evals/excel_check.py -w <作業フォルダ>                # 業務マトリクス一覧まで書き出してある作業フォルダ（自己点検の -w のフォルダでよい）
  python3 evals/excel_check.py -w <作業フォルダ> --keep <フォルダ>   # 試したブック（Excel で保存したもの）を、そのフォルダに残す

作業フォルダは変えない（写しで試す）。Excel は、試し用の、名前の重ならないブックだけを開いて、閉じる。開いているほかのブックには触らない。

確かめること:
  1. Excel で開いて計算した値が、基準の表のとおり（VLOOKUP）。合計（SUM）。ノード一覧の「仮」（ISFORMULA）
  2. Excel で、時間・件数・基準の表に数を入れると、その場で、合計・同じ種類の仮の値・ノード一覧の「人」が変わる
  3. Excel で保存したブックから、値を取り込める（build_matrix.py --keep。保存でコメント・式の書き方が変わっても）
  4. Excel で保存したブックを out に置いても、点検（check_flow.py --stage matrix）が FAIL 0。書き出し直すと、値が残る
  5. 見た目（Excel が、そのセルに表示している書式を読む）: 仮の値は、条件付き書式の字（灰色の斜体）。特大は、色つき。数を打つと、ふつうの字に戻る
  6. 式のセルに保存してある計算の結果: 計算をしない表示（Finder のプレビュー。qlmanage）でも、合計・仮の値が出る。Excel が計算した値と、同じ
"""
import argparse
import html
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

S = pathlib.Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(S))
M, NODES = "業務マトリクス", "業務フローのノード一覧"
HUMAN, COUNT, STEP = 12, 120, 7                                      # 人が入れる時間・件数、基準の表の値に足す数


def run(script, work, *args):
    r = subprocess.run([sys.executable, str(S / script), "-w", str(work), *args], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr


def number(s):
    try:
        return float(str(s).replace(",", "."))
    except ValueError:
        return None


def preview_totals(book, d):
    """Finder のプレビュー（クイックルック）が出す、業務マトリクスの時間の行の、合計の字。計算はしないので、ブックに保存してある値が出る。qlmanage が無ければ None"""
    if shutil.which("qlmanage") is None:
        return None
    d.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.run(["qlmanage", "-p", "-o", str(d), str(book)], capture_output=True, timeout=120)
    except subprocess.TimeoutExpired:
        return None
    out = []
    for p in sorted(d.rglob("*.html")):
        at = None                                                     # 合計の列（見出しの「処理時間の合計」で探す）
        for r in re.findall(r"<tr.*?</tr>", p.read_text(encoding="utf-8", errors="replace"), re.S):
            cells = [re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", "", c))).strip() for c in re.findall(r"<td.*?</td>", r, re.S)]
            at = next((j for j, c in enumerate(cells) if c.startswith("処理時間の合計")), at)
            if at is not None and len(cells) > at and any("｜処理時間（分）" in c for c in cells[:2]):     # 時間の行（行の名前のセルに、印がある）
                out.append(cells[at])
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-w", "--work", required=True)
    ap.add_argument("--keep", help="試したブックを残すフォルダ")
    a = ap.parse_args()
    import build_matrix as BM
    import office_to_pdf as OP
    import table_template as TT
    from openpyxl import load_workbook
    ok, why = OP.office_ready("excel")
    if not ok:
        sys.exit(f"Excel を動かせない（{why}）")
    tmp = pathlib.Path(a.keep).resolve() if a.keep else pathlib.Path(tempfile.mkdtemp(prefix="excel_check_"))
    tmp.mkdir(parents=True, exist_ok=True)
    for d in ("work", "work_out"):
        if (tmp / d).exists():
            shutil.rmtree(tmp / d)
    w = tmp / "work"
    shutil.copytree(a.work, w, ignore=shutil.ignore_patterns("pages", "png", "*.pdf", "~$*"))
    code, said = run("build_matrix.py", w)
    if code:
        sys.exit(f"業務マトリクス一覧を書き出せない:\n{said[-800:]}")
    shutil.copytree(w, tmp / "work_out")                              # 4 で使う（書き出したままの作業フォルダ）
    src = w / "out" / BM.BOOK
    wb, kept = load_workbook(src), load_workbook(src, data_only=True)   # 式／式のセルに保存してある、計算の結果
    ws, table = wb[M], TT.time_table(wb, BM.SPEC["time"])[0]
    every = [(s.title, c.coordinate) for s in wb for r in s.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith("=")]
    some = every[::max(1, len(every) // 300)]                         # Excel に聞く式のセル（多ければ、間を空けて 300 個ほど）
    sums = [kept[M][c.coordinate].value for r in ws.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith("=SUM(")]
    tb = wb[table["title"]]
    times, counts = {}, {}
    for row in ws.iter_rows():
        for c in row:
            text = c.comment.text if c.comment is not None else ""
            m, k = BM.TAG_TIME.match(text), BM.TAG_COUNT.match(text)
            if m:
                times[(m.group(1), m.group(2))] = c
            elif k:
                counts[k.group(1)] = c
    looked = {}                                                       # 仮の値のセル → (動作の種類, 表の何列目)
    for key, c in times.items():
        f = re.fullmatch(r'=VLOOKUP\("([^"]*)",.+,(\d+),FALSE\)', c.value) if isinstance(c.value, str) else None
        if f:
            looked[key] = (f.group(1), int(f.group(2)))
    base = lambda act, col: tb.cell(table["rows"][act], table["col"] + col - 1)
    value = lambda key, edits={}: edits.get(key, edits.get(looked.get(key), base(*looked[key]).value if key in looked else times[key].value))
    groups = {}
    for key, k in looked.items():
        groups.setdefault(k, []).append(key)
    shared = next(k for k, v in sorted(groups.items(), key=lambda x: -len(x[1])) if len(v) >= 2 and base(*k).value)   # 2 つ以上のセルが引いている、表の値
    h = next(key for key in looked if looked[key] != shared and base(*looked[key]).value)                               # 人が値を入れるセル
    F = h[0]
    row = [key for key in times if key[0] == F]
    total = next(c for c in ws[times[h].row] if isinstance(c.value, str) and c.value.startswith("=SUM("))
    nl = wb[NODES]
    col = {str(c.value): c.column for c in nl[1] if c.value}
    line = next(r for r in nl.iter_rows(min_row=2) if (r[col["フローID"] - 1].value, r[col["ノード番号"] - 1].value) == h)
    basis, listed = line[col["時間の根拠"] - 1].coordinate, line[col["処理時間（分）"] - 1].coordinate
    peer = groups[shared][0]
    name = f"excel_check_{os.getpid()}.xlsx"
    book = tmp / name
    shutil.copy(src, book)
    ask = [(M, times[k].coordinate) for k in row] + [(M, total.coordinate), (M, times[peer].coordinate), (NODES, listed), (NODES, basis)]

    def read():
        lines = ["set TB to character id 9", 'set res to ""']
        for sh, at in ask:
            lines += [f'set v to value of range "{at}" of worksheet "{sh}" of workbook "{name}"', "try", "set res to res & (v as string) & TB",
                      "on error", 'set res to res & "" & TB', "end try"]
        got = OP.excel(lines + ["return res"]).split("\t")
        nums = [number(x) for x in got[:len(row) + 3]]
        return dict(zip(row, nums)), nums[len(row)], nums[len(row) + 1], nums[len(row) + 2], got[len(row) + 3]

    def look(at):
        """Excel が、そのセルに表示している書式（条件付き書式を掛けたあと）: (斜体か, 字の色, 塗りの色)"""
        got = OP.excel(['set AppleScript\'s text item delimiters to ","', f'set r to range "{at}" of worksheet "{M}" of workbook "{name}"',
                        "set a to (italic of font object of display format of r) as string", "set b to (color of font object of display format of r) as string",
                        "set c to (color of interior object of display format of r) as string", 'return a & ";" & b & ";" & c']).split(";")
        return got[0] == "true", tuple(int(x) for x in got[1].split(",")), tuple(int(x) for x in got[2].split(","))

    rgb = lambda color: tuple(int(str(color.rgb)[-6:][k:k + 2], 16) for k in (0, 2, 4))
    rules = [(str(rng.sqref), r) for rng in ws.conditional_formatting for r in rng.rules]
    wide = next(r for where, r in rules if ":" in where)             # 仮の値の字（囲む四角に掛けた決まり）
    big = next(((where, r) for where, r in rules if ":" not in where and r.stopIfTrue), None)   # 特大のセル（あれば）

    def put(sheet, at, v):
        OP.excel([f'set value of range "{at}" of worksheet "{sheet}" of workbook "{name}" to {v}'])

    def read_all(pairs):
        lines = ["set TB to character id 9", 'set res to ""']
        for sh, at in pairs:
            lines += [f'set v to value of range "{at}" of worksheet "{sh}" of workbook "{name}"', "try", "set res to res & (v as string) & TB",
                      "on error", 'set res to res & "" & TB', "end try"]
        return OP.excel(lines + ["return res"], timeout=300).split("\t")[:len(pairs)]

    results, new = [], base(*shared).value + STEP
    seen = preview_totals(src, tmp / "preview")
    results.append(("計算をしない表示（Finder のプレビュー）でも、合計が出る（式のセルに、計算の結果を保存してある）" if seen is not None else "プレビューでの確かめは、飛ばした（qlmanage が無い）",
                    seen is None or (all(x is not None for x in sums) and [number(x) for x in seen] == [float(x) for x in sums] and any(sums)),
                    f"プレビューの合計 {seen}・保存してある合計 {sums}"))
    try:
        OP.excel([f'open POSIX file "{book}"'], timeout=120)
        got = read_all(some)
        same = lambda mine, x: mine is not None and (x == mine if isinstance(mine, str) else number(x) == float(mine))
        off = [(s, at, kept[s][at].value, x) for (s, at), x in zip(some, got) if not same(kept[s][at].value, x)]
        results.append((f"Excel が、ブックを、そのまま開く。Excel が計算した値が、式のセルに保存してある値と同じ（式のセル {len(every)} 個のうち {len(some)} 個）",
                        len(got) == len(some) and not off and all(kept[s][at].value is not None for s, at in every), f"違うセル {off[:6]}"))
        v, tot, pv, lv, bs = read()
        results.append(("Excel で開いて計算した値が、基準の表のとおり。合計は、その行の足し算。ノード一覧も同じ値で「仮」",
                        all(v[k] == value(k) for k in row) and tot == sum(value(k) for k in row) and pv == value(peer) and lv == value(h) and bs == "仮",
                        f"{F} の行 {[v[k] for k in row]}・合計 {tot}・ノード一覧 {lv}「{bs}」"))
        soft = look(times[h].coordinate)
        far = look(big[0]) if big else None
        results.append(("見た目: 仮の値のセルは、条件付き書式の字（斜体・色）で表示される" + ("。特大のセルは、その字と塗り" if big else "（特大のセルは、このブックに無い）"),
                        soft[:2] == (bool(wide.dxf.font.i), rgb(wide.dxf.font.color)) and (far is None or far[1:] == (rgb(big[1].dxf.font.color), rgb(big[1].dxf.fill.end_color))),
                        f"仮 {soft}・特大 {far}"))
        put(M, times[h].coordinate, HUMAN)
        if counts.get(F) is not None:
            put(M, counts[F].coordinate, COUNT)
        put(table["title"], base(*shared).coordinate, new)
        edits = {h: HUMAN, shared: new}
        v, tot, pv, lv, bs = read()
        results.append((f"Excel で数を入れると、その場で変わる（時間 {HUMAN} 分・基準の表「{shared[0]}」の {shared[1]} 列目を {new} 分に）: 合計・同じ種類の仮の値・ノード一覧の「人」",
                        v[h] == HUMAN and all(v[k] == value(k, edits) for k in row) and tot == sum(value(k, edits) for k in row) and pv == new and lv == HUMAN and bs == "人",
                        f"{F} の行 {[v[k] for k in row]}・合計 {tot}・同じ種類のセル {pv}・ノード一覧 {lv}「{bs}」"))
        typed = look(times[h].coordinate)
        results.append(("見た目: 数を打ったセルは、ふつうの字に戻る（斜体でない・仮の値の色でない）", not typed[0] and typed[1] != soft[1], f"{typed}"))
        OP.excel([f'save workbook "{name}"'], timeout=120)
        OP.wait_file(book)
    finally:
        try:
            OP.excel(["try", f'close workbook "{name}" saving no', "end try", "set display alerts to true"])
        except OP.ConvertError:
            pass
    saved = load_workbook(book)
    tags = sum(1 for r in saved[M].iter_rows() for c in r if c.comment is not None and (BM.TAG_TIME.match(c.comment.text) or BM.TAG_COUNT.match(c.comment.text)))
    ours = sum(1 for k, c in times.items() if isinstance(saved[M][c.coordinate].value, str) and BM.OURS.match(saved[M][c.coordinate].value))
    code, said = run("build_matrix.py", w, "--keep", str(book))
    d = json.loads((w / "times.json").read_text(encoding="utf-8")) if (w / "times.json").exists() else {}
    got = ({(x["flow"], x["node"]): x["value"] for x in d.get("cells") or []}, {x["flow"]: x["value"] for x in d.get("counts") or []},
           {x["action"]: x["value"] for x in d.get("table") or []})
    results.append(("Excel で保存したブックから、値を取り込める（コメントの札が残り、仮の値の式を見分けられる）",
                    code == 0 and tags == len(times) + len(counts) and ours == len(looked) - 1 and got[0] == {h: HUMAN}
                    and got[1] == ({F: COUNT} if counts.get(F) is not None else {}) and got[2] == {shared[0]: new},
                    f"札 {tags}/{len(times) + len(counts)}・仮の式 {ours}/{len(looked) - 1}・取り込んだ値 {got}"))
    w2 = tmp / "work_out"
    shutil.copy(book, w2 / "out" / BM.BOOK)
    c1, o1 = run("check_flow.py", w2, "--stage", "matrix")
    c2, o2 = run("build_matrix.py", w2)
    c3, o3 = run("check_flow.py", w2, "--stage", "matrix")
    kept = json.loads((w2 / "times.json").read_text(encoding="utf-8")) if (w2 / "times.json").exists() else {}
    results.append(("Excel で保存したブックを out に置いても、点検が FAIL 0。書き出し直すと、値が残る",
                    c1 == 0 and "FAIL  " not in o1 and c2 == 0 and c3 == 0 and "FAIL  " not in o3 and {(x["flow"], x["node"]): x["value"] for x in kept.get("cells") or []} == {h: HUMAN},
                    (o1.strip().splitlines() or [""])[-1] + "／" + (o3.strip().splitlines() or [""])[-1]))
    for text, passed, detail in results:
        print(("OK   " if passed else "NG   ") + text + ("" if passed else f"\n       {detail}"))
    if not a.keep:
        shutil.rmtree(tmp, ignore_errors=True)
    print("excel_check OK" if all(p for _t, p, _d in results) else "excel_check NG")
    sys.exit(0 if all(p for _t, p, _d in results) else 1)


if __name__ == "__main__":
    main()
