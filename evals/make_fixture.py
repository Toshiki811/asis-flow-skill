#!/usr/bin/env python3
"""make_fixture.py — スクリプトの動作確認用の小さな入力一式を作る（架空の「経費精算」マニュアル）。

  python3 evals/make_fixture.py <出力フォルダ>
  python3 evals/make_fixture.py <出力フォルダ> --atypical     # 別の入力: 原典の順が、一般的な流れと逆の手順書1つ（仮払金。現金を渡したあとで申請・承認）
  python3 evals/make_fixture.py <出力フォルダ> --code         # 別の入力: 手順書1つ（月次の請求集計）と、そのマクロの書き出し（.bas・.cls）

  00_全体/経費精算業務マニュアル.pptx   表紙・目次・手順（文字）・フロー図（図形とコネクタ）・画面の画像・非表示のスライド・改訂履歴
  01_手順書/支払処理手順.pdf           文字のページと、文字層のない画像だけのページ
  01_手順書/経費精算チェックシート.xlsx  手順の表のシートと、図形・コネクタ・画像のあるシート
  02_規程/経費精算規程.pdf             表紙（番号なし）＋本文2ページ（下端に「- 1 -」「- 2 -」）

点検の仕組みを確かめるための仕込み:
  階層        規程（M04）› 全体マニュアル（M01）› 手順書（M02）› 別紙（M03）。M01「3. 支払処理」は1行、M02 が詳細（上位と下位で粒度が違う）
  受け渡し    承認（M01）→ 支払処理（M02）。資料をまたいで業務が続く
  食い違い    支払日が、規程（毎月25日）と手順書（翌月10日）で違う
  補足の作業  M01 スライド5の「※ …確認すること。」は作業。M02 の「なお、…負担します。」は作業でない
  ページ番号  M04 は表紙を数えないので、刷られた番号がファイルの番号より 1 小さい。M01 はスライド7が非表示（番号が飛ぶ）
"""
import argparse
import io
import pathlib
import zipfile

from PIL import Image, ImageDraw, ImageFont

JP_FONTS = ["/System/Library/Fonts/ヒラギノ角ゴシック W3.ttc", "/System/Library/Fonts/Hiragino Sans GB.ttc",
            "/Library/Fonts/Arial Unicode.ttf", "C:/Windows/Fonts/meiryo.ttc",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]


def font(size):
    for f in JP_FONTS:
        if pathlib.Path(f).exists():
            return ImageFont.truetype(f, size)
    return ImageFont.load_default()


def screen_png(lines, title, size=(900, 520)):
    """画面キャプチャ風の画像（文字は画像の中にしかない）"""
    im = Image.new("RGB", size, "#f4f6f8")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, size[0], 46], fill="#2f5d8a")
    d.text((16, 10), title, font=font(22), fill="white")
    y = 70
    for label, value in lines:
        d.text((30, y + 6), label, font=font(20), fill="#222222")
        d.rectangle([260, y, 700, y + 36], outline="#888888", fill="white")
        d.text((270, y + 6), value, font=font(20), fill="#222222")
        y += 54
    d.rectangle([560, y + 10, 700, y + 54], fill="#2f8a4b")
    d.text((596, y + 20), "登録", font=font(20), fill="white")
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def make_pptx(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)

    def text_slide(title, lines):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        tb = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.5), Inches(5)).text_frame
        tb.word_wrap = True
        for i, line in enumerate(lines):
            p = tb.paragraphs[0] if i == 0 else tb.add_paragraph()
            p.text = line
            p.font.size = Pt(20)
        return s

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "経費精算業務マニュアル"
    s.placeholders[1].text = "管理部 第3版"
    text_slide("目次", ["1. 経費申請 ………… 3", "2. 承認 ………… 5", "3. 支払処理 ………… 6"])
    text_slide("1. 経費申請", [
        "(1) 申請者は、領収書を受け取ったら経費精算システムにログインする。",
        "(2) 申請者は、経費申請書を作成し、領収書の画像を添付する。",
        "(3) 申請者は、申請ボタンを押して上長に提出する。",
        "※ 申請の締め日は毎月25日。"])

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "1. 経費申請（全体の流れ）"

    def box(shape, text, x, y, w=2.3, h=0.9):
        b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        b.text_frame.text = text
        return b

    a = box(MSO_SHAPE.ROUNDED_RECTANGLE, "申請者: 経費申請書を作成", 0.8, 2.0)
    b = box(MSO_SHAPE.DIAMOND, "5万円以上？", 4.2, 1.8, 2.3, 1.3)
    c = box(MSO_SHAPE.ROUNDED_RECTANGLE, "課長: 承認", 7.6, 1.2)
    d = box(MSO_SHAPE.ROUNDED_RECTANGLE, "部長: 承認", 7.6, 3.2)
    e = box(MSO_SHAPE.ROUNDED_RECTANGLE, "経理部: 支払処理", 10.6, 2.2)

    def link(p, q, i=3, j=1):
        cn = s.shapes.add_connector(MSO_CONNECTOR.ELBOW, 0, 0, 0, 0)
        cn.begin_connect(p, i)
        cn.end_connect(q, j)

    link(a, b)
    link(b, c)
    link(b, d)
    link(c, e)
    link(d, e)
    n = s.shapes.add_textbox(Inches(6.4), Inches(1.4), Inches(1.2), Inches(0.4)).text_frame
    n.text = "いいえ"
    n = s.shapes.add_textbox(Inches(6.4), Inches(3.3), Inches(1.2), Inches(0.4)).text_frame
    n.text = "はい"

    text_slide("2. 承認", [
        "(1) 上長は、申請内容と領収書を照合する。",
        "(2) 申請金額が5万円未満の場合は課長が承認する。5万円以上の場合は部長が承認する。",
        "(3) 不備がある場合は、申請者に差し戻す。",
        "※ 5万円以上の申請は、承認の前に見積書の添付を確認すること。"])

    s = text_slide("3. 支払処理", ["経理部は、承認済みの申請を下の画面で確認し、支払データを登録する。"])
    png = screen_png([("申請番号", "EX-2026-00123"), ("申請者", "山田 太郎"), ("金額", "48,600 円"),
                      ("支払予定日", "翌月10日"), ("振込先", "給与口座")], "経費精算システム ｜ 支払登録")
    s.shapes.add_picture(io.BytesIO(png), Inches(2.2), Inches(2.6), width=Inches(7.5))
    s = text_slide("（旧）3. 支払処理", ["経理部は、承認済みの申請を紙の一覧で確認し、支払依頼書を作成する。"])
    s._element.set("show", "0")                                      # 非表示のスライド（PDF に出ないので、ページの番号が1つ飛ぶ）
    text_slide("改訂履歴", ["第1版 2024年4月 新規作成", "第2版 2025年4月 承認の金額基準を改定", "第3版 2026年4月 経費精算システムの画面を更新"])
    prs.save(path)


def make_pdf(path):
    import fitz
    doc = fitz.open()
    p = doc.new_page(width=595, height=842)
    lines = ["支払処理手順", "", "1. 経理部は、毎月末に承認済みの経費申請を一覧で出力する。",
             "2. 経理部は、支払データを作成し、経理課長の確認を受ける。",
             "3. 経理部は、翌月10日に銀行のファームバンキングで振込を実行する。",
             "4. 振込結果は会計システムに自動で連携される。", "なお、振込手数料は会社が負担します。", "",
             "詳細は別紙「経費精算チェックシート」を参照。"]
    y = 80
    for i, t in enumerate(lines):
        p.insert_text((60, y), t, fontname="japan", fontsize=18 if i == 0 else 12)
        y += 34 if i == 0 else 24
    # 2ページ目は画像だけ（スキャンした紙を模す。文字層なし）
    im = Image.new("RGB", (1190, 1684), "white")
    d = ImageDraw.Draw(im)
    d.text((120, 140), "振込エラー時の対応", font=font(44), fill="black")
    for i, t in enumerate(["① 経理部は、エラー一覧を印刷する。", "② 口座情報の誤りは、申請者に電話で確認する。",
                           "③ 修正後、翌営業日に再振込を行う。", "④ 再振込の記録を支払台帳に記入する。"]):
        d.text((140, 300 + i * 90), t, font=font(34), fill="black")
    b = io.BytesIO()
    im.save(b, "PNG")
    p2 = doc.new_page(width=595, height=842)
    p2.insert_image(p2.rect, stream=b.getvalue())
    doc.save(path)


def make_rule_pdf(path):
    """規程。表紙には番号が無く、本文の下端に「- 1 -」「- 2 -」と刷ってある（ファイルの番号と 1 ずれる）"""
    import fitz
    doc = fitz.open()
    pages = [["経費精算規程", "", "管理部"],
             ["第1章 申請と承認", "", "第5条 申請者は、経費を使った日から1か月以内に申請しなければならない。",
              "第6条 5万円以上の経費は、部長の承認を得るものとする。"],
             ["第2章 支払", "", "第7条 経理部は、承認済みの経費を毎月25日に支払う。",
              "第8条 この規程の細則は、経費精算業務マニュアルに定める。"]]
    for i, lines in enumerate(pages):
        p = doc.new_page(width=595, height=842)
        y = 300 if i == 0 else 90
        for j, t in enumerate(lines):
            p.insert_text((60, y), t, fontname="japan", fontsize=22 if (i == 0 and j == 0) else 16 if j == 0 else 12)
            y += 40 if j == 0 else 26
        if i:
            p.insert_text((283, 812), f"- {i} -", fontname="helv", fontsize=10)
    doc.save(path)


DRAWING_SHAPES = """
<xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="101" name="処理 1"/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr><a:prstGeom prst="flowChartProcess"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="DDEBF7"/></a:solidFill><a:ln><a:solidFill><a:srgbClr val="1F4E79"/></a:solidFill></a:ln></xdr:spPr><xdr:txBody><a:bodyPr/><a:p><a:r><a:rPr lang="ja-JP"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:rPr><a:t>領収書を確認</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>
<xdr:twoCellAnchor><xdr:from><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>6</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="102" name="判断 2"/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr><a:prstGeom prst="flowChartDecision"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="FFF2CC"/></a:solidFill><a:ln><a:solidFill><a:srgbClr val="7F6000"/></a:solidFill></a:ln></xdr:spPr><xdr:txBody><a:bodyPr/><a:p><a:r><a:rPr lang="ja-JP"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:rPr><a:t>不備あり？</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>
<xdr:twoCellAnchor><xdr:from><xdr:col>7</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>9</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="103" name="処理 3"/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr><a:prstGeom prst="flowChartProcess"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="DDEBF7"/></a:solidFill><a:ln><a:solidFill><a:srgbClr val="1F4E79"/></a:solidFill></a:ln></xdr:spPr><xdr:txBody><a:bodyPr/><a:p><a:r><a:rPr lang="ja-JP"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:rPr><a:t>申請者に差し戻す</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>
<xdr:twoCellAnchor><xdr:from><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>95250</xdr:rowOff></xdr:from><xdr:to><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>95250</xdr:rowOff></xdr:to>
<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="104" name="矢印 4"/><xdr:cNvCxnSpPr><a:stCxn id="101" idx="3"/><a:endCxn id="102" idx="1"/></xdr:cNvCxnSpPr></xdr:nvCxnSpPr><xdr:spPr><a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom><a:ln w="19050"><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:tailEnd type="triangle"/></a:ln></xdr:spPr></xdr:cxnSp><xdr:clientData/></xdr:twoCellAnchor>
<xdr:twoCellAnchor><xdr:from><xdr:col>6</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>95250</xdr:rowOff></xdr:from><xdr:to><xdr:col>7</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>95250</xdr:rowOff></xdr:to>
<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="105" name="矢印 5"/><xdr:cNvCxnSpPr><a:stCxn id="102" idx="3"/><a:endCxn id="103" idx="1"/></xdr:cNvCxnSpPr></xdr:nvCxnSpPr><xdr:spPr><a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom><a:ln w="19050"><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:tailEnd type="triangle"/></a:ln></xdr:spPr></xdr:cxnSp><xdr:clientData/></xdr:twoCellAnchor>
"""


def make_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    wb = Workbook()
    ws = wb.active
    ws.title = "手順"
    rows = [["No", "担当", "作業", "使う帳票"], [1, "経理部", "承認済み申請の領収書原本を回収する", "領収書"],
            [2, "経理部", "領収書の金額と申請金額を照合する", "経費申請書"],
            [3, "経理部", "不備があれば申請者に差し戻す", ""], [4, "経理課長", "支払データを確認する", "支払一覧表"]]
    for r in rows:
        ws.append(r)
    ws.column_dimensions["C"].width = 44
    ws2 = wb.create_sheet("フロー")
    ws2["B1"] = "領収書チェックの流れ"
    png = screen_png([("領収書番号", "R-0456"), ("照合結果", "一致")], "チェック結果の入力", size=(700, 300))
    tmp = path.with_suffix(".tmp.png")
    tmp.write_bytes(png)
    ws2.add_image(XLImage(str(tmp)), "B9")
    wb.save(path)
    tmp.unlink()
    # openpyxl は図形を書けないので、できた drawing に図形とコネクタを差し込む
    src = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/drawings/drawing1.xml":
                xml = data.decode("utf-8")
                ns = ('xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" ')
                for root in ("wsDr", "xdr:wsDr"):                    # openpyxl は接頭辞なしで書く
                    if f"</{root}>" in xml:
                        xml = xml.replace(f"<{root} ", f"<{root} {ns}", 1).replace(f"</{root}>", DRAWING_SHAPES + f"</{root}>")
                        break
                data = xml.encode("utf-8")
            out.writestr(item, data)
    src.close()
    path.write_bytes(buf.getvalue())


def make_atypical_pdf(path):
    """原典の順が、一般的な流れ（申請 → 承認 → 支払）と逆の手順書。読取・フロー化が、よくある流れに並べ直さないかを見る"""
    import fitz
    doc = fitz.open()
    p = doc.new_page(width=595, height=842)
    lines = ["仮払金の取扱手順", "", "1. 経理部は、出張者に仮払金を現金で渡す。", "2. 出張者は、仮払金を受け取ったあとで、仮払申請書を作成して課長に提出する。",
             "3. 課長は、仮払申請書を承認し、経理部に回付する。", "4. 出張者は、出張から戻ったら、領収書と残金を経理部に提出する。",
             "5. 経理部は、領収書と残金を受け取ってから、仮払精算書を作成する。", "",
             "※ 急な出張に間に合わせるため、当社では現金を先に渡し、申請と承認はあとから行う。"]
    y = 80
    for i, t in enumerate(lines):
        p.insert_text((60, y), t, fontname="japan", fontsize=18 if i == 0 else 12)
        y += 34 if i == 0 else 24
    doc.save(path)


CODE_BAS = r"""Attribute VB_Name = "月次集計"
Option Explicit

Private Const 売上フォルダ As String = "\\fileserver\keiri\売上\"

' 「集計」ボタンから動く
Public Sub Main()
    Dim wb As Workbook, ws As Worksheet, 対象月 As String
    Dim i As Long, 最終行 As Long, 件数 As Long

    On Error GoTo ErrHandler
    Application.ScreenUpdating = False

    ' 対象月を確かめる
    対象月 = Worksheets("メニュー").Range("B2").Value
    If 対象月 = "" Then
        MsgBox "対象月を入力してください"
        Exit Sub
    End If
    If MsgBox(対象月 & " の集計を始めますか", vbYesNo) = vbNo Then Exit Sub

    ' 売上データを読み込む
    Set wb = Workbooks.Open(売上フォルダ & "売上_" & 対象月 & ".csv")
    Set ws = ThisWorkbook.Worksheets("請求一覧")
    ws.Range("A2:F5000").ClearContents
    wb.Sheets(1).UsedRange.Copy Destination:=ws.Range("A1")
    wb.Close SaveChanges:=False

    ' 取引先ごとに税額を入れる
    最終行 = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row
    For i = 2 To 最終行
        If ws.Cells(i, 3).Value >= 1000000 Then
            ws.Cells(i, 7).Value = "要確認"
            件数 = 件数 + 1
        End If
        ws.Cells(i, 6).Value = 消費税計算(ws.Cells(i, 3).Value, ws.Cells(i, 4).Value)
    Next i

    Call 集計表を保存する(対象月)

    If 件数 > 0 Then
        メールを送る 件数
    End If

    Application.ScreenUpdating = True
    MsgBox "集計が終わりました"
    Exit Sub

ErrHandler:
    Application.ScreenUpdating = True
    MsgBox "エラーが起きました: " & Err.Description
End Sub

Private Sub 集計表を保存する(対象月 As String)
    ' PDF にして共有フォルダに置く
    Worksheets("集計表").ExportAsFixedFormat Type:=xlTypePDF, Filename:="\\fileserver\keiri\月次\集計表.pdf"
    ThisWorkbook.Save
End Sub

Private Sub メールを送る(件数 As Long)
    Dim app As Object, mail As Object
    Set app = CreateObject("Outlook.Application")
    Set mail = app.CreateItem(0)
    mail.To = "kacho@example.com"
    mail.Subject = "月次集計: 要確認 " & 件数 & " 件"
    mail.Send
End Sub

' 入金の消込（ボタンは無い）
Public Sub 入金消込()
    Dim cn As Object, rs As Object, r As Long
    Set cn = CreateObject("ADODB.Connection")
    cn.Open "Provider=SQLOLEDB;Data Source=db01;Initial Catalog=hanbai;User ID=keiri;Password=abc12345;"
    Set rs = cn.Execute("SELECT 請求番号, 入金額 FROM 入金明細 WHERE 消込済 = 0")
    r = 2
    Do Until rs.EOF
        Select Case rs("入金額")
            Case Is < 0
                Worksheets("エラー").Cells(r, 1).Value = rs("請求番号")
            Case 0
                ' 何もしない
            Case Else
                Worksheets("請求一覧").Cells(r, 8).Value = rs("入金額") + 消費税計算(rs("入金額"), "標準")
        End Select
        r = r + 1
        rs.MoveNext
    Loop
    cn.Execute "UPDATE 入金明細 SET 消込済 = 1 WHERE 消込済 = 0"
    cn.Close
End Sub

Public Function 消費税計算(金額 As Currency, 区分 As String) As Currency
    If 区分 = "非課税" Then
        消費税計算 = 0
        Exit Function
    End If
    If 区分 = "軽減" Then
        消費税計算 = Int(金額 * 0.08)
    Else
        消費税計算 = Int(金額 * 0.1)
    End If
End Function

Sub 古い集計()
    Worksheets("旧集計").Range("A1").Value = Now
End Sub
"""

CODE_CLS = r"""VERSION 1.0 CLASS
BEGIN
  MultiUse = -1  'True
END
Attribute VB_Name = "ThisWorkbook"
Private Sub Workbook_Open()
    Worksheets("メニュー").Activate
    Worksheets("メニュー").Range("B2").Value = Format(DateAdd("m", -1, Date), "yyyymm")
End Sub
"""


def make_code_input(out):
    """プログラムのソースを含む入力: 手順書1つ（月次の請求集計）と、そのマクロの書き出し（標準モジュールとブックのモジュール）。
    仕込み: マニュアルが名指しするマクロ（Main）・マニュアルに出てこない起動点（ブックを開いたとき・古い集計）・2か所から呼ばれる手続き（消費税計算）・
    呼ぶ所が1か所だけの小さな手続き（集計表を保存する・メールを送る）・エラー処理・操作者への確認・接続文字列のパスワード"""
    import fitz
    (out / "手順書").mkdir(parents=True, exist_ok=True)
    (out / "マクロ").mkdir(parents=True, exist_ok=True)
    (out / "マクロ" / "月次集計.bas").write_text(CODE_BAS, encoding="utf-8")
    (out / "マクロ" / "ThisWorkbook.cls").write_text(CODE_CLS, encoding="utf-8")
    doc = fitz.open()
    p = doc.new_page(width=595, height=842)
    lines = ["月次の請求集計の手順", "", "1. 営業担当は、月末までに売上データ（CSV）を共有フォルダに置く。",
             "2. 経理担当は、「月次集計」のブックを開き、メニューのシートで対象月を確かめる。", "3. 経理担当は、マクロ「Main」を実行する（メニューの「集計」ボタン）。",
             "4. 経理担当は、請求一覧の「要確認」の行を、販売管理システムの売上と照合する。", "5. 経理課長は、集計表を確認し、承認する。", "6. 経理担当は、入金の消込を行う。"]
    y = 80
    for i, t in enumerate(lines):
        p.insert_text((50, y), t, fontname="japan", fontsize=18 if i == 0 else 11)
        y += 34 if i == 0 else 24
    doc.save(out / "手順書" / "月次の請求集計の手順.pdf")


# ---------- Access を書き出したテキスト（架空。形は SaveAsText の出力と、assets/access_export/ の道具の出力に合わせてある） ----------

def acc_ctl(kind, name, caption="", events=(), child="", extra=""):
    ev = "".join(f'                    {k} ="{v}"\n' for k, v in events)
    cap = f'                    Caption ="{caption}"\n' if caption else ""
    return (f"                Begin {kind}\n                    OverlapFlags =85\n                    Left =1800\n                    Top =960\n"
            f'                    Name ="{name}"\n{cap}{ev}{extra}'
            "                    GUID = Begin\n                        0x3822e4b79de09b4b8acf9150c95debd3\n                    End\n"
            f"{child}                    LayoutCachedLeft =1800\n                End\n")


def acc_form(obj, caption, source, top_events, controls, code):
    ev = "".join(f'    {k} ="{v}"\n' for k, v in top_events)
    src = f'    RecordSource ="{source}"\n' if source else ""
    return ("Version =21\nVersionRequired =20\nPublishOption =1\nChecksum =-540535037\n"
            f"Begin {obj}\n    DividingLines = NotDefault\n    AllowDesignChanges = NotDefault\n    DefaultView =0\n    Width =8884\n"
            "    GUID = Begin\n        0x215548c7105bb74f9c5a62a34852d43e\n    End\n"
            "    NameMap = Begin\n        0x0acc0e5500000000000000000000000000000000000000000c00000004000000 ,\n        0x0000000000000000000000000000\n    End\n"
            f'{src}    Caption ="{caption}"\n{ev}    DatasheetFontName ="Yu Gothic UI"\n'
            "    Begin\n        Begin Label\n            BackStyle =0\n            FontSize =11\n        End\n"
            "        Begin CommandButton\n            FontSize =11\n            FontWeight =400\n        End\n"
            "        Begin TextBox\n            AddColon = NotDefault\n            BorderLineStyle =0\n        End\n"
            '        Begin Section\n            Height =7560\n            Name ="詳細"\n            AutoHeight =1\n            Begin\n'
            f"{controls}            End\n        End\n    End\nEnd\n" + (f"CodeBehindForm\nAttribute VB_GlobalNameSpace = False\nAttribute VB_Creatable = True\n"
                                                                      f"Attribute VB_PredeclaredId = True\nAttribute VB_Exposed = False\n{code}" if code else ""))


ACC_LABEL = lambda name, caption: ("                    Begin\n                        Begin Label\n                            OverlapFlags =85\n"
                                   f'                            Name ="{name}"\n                            Caption ="{caption}"\n                        End\n                    End\n')
ACC_EMBEDDED = ("                    OnClickEmMacro = Begin\n                        Version =196611\n                        ColumnsShown =0\n                        Begin\n"
                '                            Action ="Quit"\n                            Argument ="1"\n                        End\n                        Begin\n'
                '                            Comment ="_AXL:<?xml version=\\"1.0\\" encoding=\\"UTF-16\\" standalone=\\"no\\"?>\\015\\012<UserI"\n'
                '                                "nterfaceMacro><Statements><Action Name=\\"QuitAccess\\"/></Statements></UserInterfaceMacro>"\n'
                "                        End\n                    End\n")

ACC_MENU_CODE = r"""Option Compare Database
Option Explicit

Private Sub Form_Load()
    Me.lbl日付.Caption = Format(Date, "yyyy/mm/dd")
End Sub

Private Sub cmd取込_Click()
    Call 受注データを取り込む
End Sub

Private Sub cmd請求書_Click()
    Dim 対象月 As String
    対象月 = InputBox("請求する月を入れてください（例: 202609）")
    If 対象月 = "" Then
        Exit Sub
    End If
    DoCmd.OpenReport "請求書", acViewPreview, , "請求月='" & 対象月 & "'"
End Sub
"""

ACC_ORDER_CODE = r"""Option Compare Database
Option Explicit

Private Sub txt数量_AfterUpdate()
    If Me.txt数量 <= 0 Then
        MsgBox "数量は 1 以上を入れてください"
        Me.txt数量 = Null
    End If
End Sub

Private Sub cmd登録_Click()
    If IsNull(Me.cbo顧客) Then
        MsgBox "顧客を選んでください"
        Exit Sub
    End If
    DoCmd.RunCommand acCmdSaveRecord
    ' 在庫を引き当てる
    DoCmd.SetWarnings False
    DoCmd.OpenQuery "Q_在庫_引当"
    DoCmd.SetWarnings True
    MsgBox "登録しました"
End Sub

Private Sub cmd削除_Click()
    DoCmd.RunCommand acCmdDeleteRecord
End Sub
"""

ACC_REPORT_CODE = r"""Option Compare Database
Option Explicit

Private Sub Report_NoData(Cancel As Integer)
    MsgBox "請求するデータがありません"
    Cancel = True
End Sub
"""

ACC_MODULE = r"""Option Compare Database
Option Explicit

Private Const 取込フォルダ As String = "\\fs01\juchu\取込\"

' 受注データ（CSV）を取り込む
Public Sub 受注データを取り込む()
    DoCmd.SetWarnings False
    DoCmd.RunSQL "DELETE FROM T_受注_取込"
    DoCmd.TransferText acImportDelim, "受注取込定義", "T_受注_取込", 取込フォルダ & "受注データ.csv", True
    If DCount("*", "T_受注_取込") = 0 Then
        DoCmd.SetWarnings True
        MsgBox "取り込むデータがありません"
        Exit Sub
    End If
    DoCmd.OpenQuery "Q_受注_追加"
    DoCmd.SetWarnings True
    MsgBox "受注データを取り込みました"
End Sub

Public Function 月次ログを書く()
    CurrentDb.Execute "INSERT INTO T_処理ログ ( 処理名, 処理日時 ) VALUES ( '月次集計', Now() )"
End Function
"""


def acc_macro(rows):
    out = "Version =196611\nColumnsShown =8\n"
    for r in rows:
        out += "Begin\n" + "".join(f'    {k} ="{v}"\n' for k, v in r) + "End\n"
    return out


ACC_MACRO_MONTHLY = [[("Action", "SetWarnings"), ("Argument", "0")],
                     [("Action", "OpenQuery"), ("Argument", "Q_月次売上_削除"), ("Argument", "0"), ("Argument", "1")],
                     [("Action", "OpenQuery"), ("Argument", "Q_月次売上_追加"), ("Argument", "0"), ("Argument", "1")],
                     [("Condition", "DCount(\\\"*\\\",\\\"T_月次売上\\\")=0"), ("Action", "MsgBox"), ("Argument", "対象のデータがありません"), ("Argument", "-1"), ("Argument", "0")],
                     [("Condition", "..."), ("Action", "StopMacro")],
                     [("Action", "TransferSpreadsheet"), ("Argument", "1"), ("Argument", "10"), ("Argument", "T_月次売上"), ("Argument", "\\\\\\\\fs01\\\\juchu\\\\月次\\\\月次売上.xlsx"),
                      ("Argument", "-1")],
                     [("Action", "RunCode"), ("Argument", "月次ログを書く()")],
                     [("Action", "MsgBox"), ("Argument", "月次集計が終わりました"), ("Argument", "-1"), ("Argument", "0")],
                     [("Action", "SetWarnings"), ("Argument", "-1")],
                     [("Comment", "_AXL:<?xml version=\\\"1.0\\\" encoding=\\\"UTF-16\\\" standalone=\\\"no\\\"?>\\015\\012<UserInterfaceMacro/>")]]
ACC_MACRO_AUTO = [[("Action", "OpenForm"), ("Argument", "メニュー"), ("Argument", "0"), ("Argument", ""), ("Argument", ""), ("Argument", "-1"), ("Argument", "0")]]

ACC_QUERIES = """-- access_queries 1

-- query: Q_月次売上_削除\ttype=delete
DELETE T_月次売上.*
FROM T_月次売上;

-- query: Q_月次売上_追加\ttype=append
INSERT INTO T_月次売上 ( 顧客コード, 売上月, 金額 )
SELECT T_受注.顧客コード, Format([受注日],"yyyymm") AS 売上月, Sum(T_受注.金額) AS 金額
FROM T_受注
GROUP BY T_受注.顧客コード, Format([受注日],"yyyymm");

-- query: Q_受注_追加\ttype=append
INSERT INTO T_受注 ( 受注日, 顧客コード, 商品コード, 数量, 金額 )
SELECT T_受注_取込.受注日, T_受注_取込.顧客コード, T_受注_取込.商品コード, T_受注_取込.数量, T_受注_取込.金額
FROM T_受注_取込;

-- query: Q_在庫_引当\ttype=update
UPDATE T_在庫 INNER JOIN T_受注 ON T_在庫.商品コード = T_受注.商品コード SET T_在庫.引当数 = [T_在庫].[引当数]+[T_受注].[数量]
WHERE (((T_受注.引当済)=False));

-- query: Q_請求書明細\ttype=select
SELECT T_受注.受注番号, T_顧客.顧客名, T_受注.金額, Format([受注日],"yyyymm") AS 請求月
FROM T_受注 INNER JOIN T_顧客 ON T_受注.顧客コード = T_顧客.顧客コード;

-- query: Q_売掛_照会\ttype=passthrough
-- connect: ODBC;DSN=KEIRI;UID=euc;PWD=***;DATABASE=keiri
SELECT * FROM dbo.売掛残高
"""

ACC_TABLES = """# access_tables 1
# Table definitions only. No rows are written.
[startup]\tform=メニュー
[table]\tT_受注\trows=1284
  受注番号\tautonumber\tpk
  受注日\tdatetime
  顧客コード\ttext(10)
  商品コード\ttext(10)
  数量\tlong
  金額\tcurrency
  引当済\tyesno
[table]\tT_受注_取込\trows=0
  受注日\tdatetime
  顧客コード\ttext(10)
  商品コード\ttext(10)
  数量\tlong
  金額\tcurrency
[table]\tT_月次売上\trows=36
  顧客コード\ttext(10)\tpk
  売上月\ttext(6)\tpk
  金額\tcurrency
[table]\tT_在庫\trows=120
  商品コード\ttext(10)\tpk
  在庫数\tlong
  引当数\tlong
[table]\tT_処理ログ\trows=52
  処理名\ttext(50)
  処理日時\tdatetime
[linked]\tT_顧客\tsource=顧客\tconnect=;DATABASE=\\\\fs01\\master\\顧客マスタ.accdb
  顧客コード\ttext(10)
  顧客名\ttext(50)
[relation]\tT_顧客.顧客コード\tT_受注.顧客コード
"""

ACC_LOG = """# access_export 1
tool\tasis-flow-skill access_export 1
source\t受注管理.accdb
date\t2026/10/01 10:00:00
access\t16.0
tables\t6\taccess_tables.txt\tok
queries\t6\taccess_queries.sql\tok
form\tメニュー\tF_メニュー.form.txt\tok
form\t受注入力\tF_受注入力.form.txt\tok
report\t請求書\tR_請求書.report.txt\tok
report\t旧請求書\tR_旧請求書.report.txt\tFAILED 2501 The action was canceled.
macro\tAutoExec\tX_AutoExec.macro.txt\tok
macro\t月次集計\tX_月次集計.macro.txt\tok
module\tmodCommon\tM_modCommon.bas\tok
done\t6\tfailed\t1
"""


def make_access_input(out):
    """Access を書き出したテキスト（架空の「受注管理」）。フォーム2・レポート1・マクロ2・モジュール1・クエリ6・テーブル6。
    仕込み: ボタンに割り当てたコード・マクロ・埋め込みマクロ／項目の見出し（付いたラベル）／割り当てていない手続き（cmd削除_Click）／
    条件つきのマクロの行と「...」／マクロから VBA の関数（RunCode）／クエリの定義から引くテーブル／リンク テーブル／書き出せなかったレポート／
    文字コード（フォーム・レポート・マクロと道具の出力は UTF-16、モジュールは Shift_JIS）"""
    d = out / "受注管理_export"
    d.mkdir(parents=True, exist_ok=True)
    u16 = lambda name, text: (d / name).write_bytes(b"\xff\xfe" + text.replace("\n", "\r\n").encode("utf-16-le"))
    menu = (acc_ctl("Label", "lbl日付", "日付") + acc_ctl("CommandButton", "cmd取込", "受注データ取込", [("OnClick", "[Event Procedure]")])
            + acc_ctl("CommandButton", "cmd月次", "月次集計", [("OnClick", "月次集計")]) + acc_ctl("CommandButton", "cmd請求書", "請求書印刷", [("OnClick", "[Event Procedure]")])
            + acc_ctl("CommandButton", "cmd終了", "終了", [("OnClick", "[Embedded Macro]")], extra=ACC_EMBEDDED))
    order = (acc_ctl("ComboBox", "cbo顧客", child=ACC_LABEL("lbl顧客", "顧客")) + acc_ctl("TextBox", "txt数量", events=[("AfterUpdate", "[Event Procedure]")], child=ACC_LABEL("lbl数量", "数量"))
             + acc_ctl("CommandButton", "cmd登録", "登録", [("OnClick", "[Event Procedure]")]) + acc_ctl("CommandButton", "cmd削除", "削除"))
    u16("F_メニュー.form.txt", acc_form("Form", "受注管理メニュー", "", [("OnLoad", "[Event Procedure]")], menu, ACC_MENU_CODE))
    u16("F_受注入力.form.txt", acc_form("Form", "受注入力", "T_受注", [], order, ACC_ORDER_CODE))
    u16("R_請求書.report.txt", acc_form("Report", "請求書", "Q_請求書明細", [("OnNoData", "[Event Procedure]")], acc_ctl("TextBox", "txt金額"), ACC_REPORT_CODE))
    u16("X_月次集計.macro.txt", acc_macro(ACC_MACRO_MONTHLY))
    u16("X_AutoExec.macro.txt", acc_macro(ACC_MACRO_AUTO))
    u16("access_queries.sql", ACC_QUERIES)
    u16("access_tables.txt", ACC_TABLES)
    u16("access_export.txt", ACC_LOG)
    (d / "M_modCommon.bas").write_bytes(ACC_MODULE.replace("\n", "\r\n").encode("cp932"))

# ---------- Access を、版管理アドイン（msaccess-vcs-addin）で書き出したフォルダ（架空。並びと形は、実物の書き出しに合わせてある） ----------

VCS_FORM = """Version =21
VersionRequired =20
Begin Form
    Caption ="受注管理メニュー"
    OnLoad ="[Event Procedure]"
    Begin
        Begin Label
            FontName ="Calibri"
        End
        Begin Section
            Name ="詳細"
            Begin
                Begin CommandButton
                    Name ="cmd月次"
                    Caption ="月次集計"
                    OnClick ="[Event Procedure]"
                End
                Begin CommandButton
                    Name ="cmd取込"
                    Caption ="受注データ"
                        "取込"
                    OnClick ="月次取込"
                End
                Begin TextBox
                    Name ="txt対象月"
                    ControlSource ="=今月の初日()"
                    Begin
                        Begin Label
                            Name ="lbl対象月"
                            Caption ="対象月"
                        End
                    End
                End
            End
        End
    End
End
CodeBehindForm
' See "メニュー.cls"
"""
VCS_FORM_CODE = """Attribute VB_GlobalNameSpace = False
Attribute VB_Creatable = True
Attribute VB_PredeclaredId = True
Attribute VB_Exposed = False
Option Compare Database
Option Explicit

Private Sub Form_Load()
    Me.txt対象月.Requery
End Sub

Private Sub cmd月次_Click()
    Dim 明細 As cls明細
    Set 明細 = New cls明細
    If 明細.件数 = 0 Then
        MsgBox "対象のデータがありません"
        Exit Sub
    End If
    DoCmd.OpenQuery "Q_月次/売上_追加"
End Sub

Private Sub cmd削除_Click()
    DoCmd.RunSQL "DELETE FROM T_受注"
End Sub
"""
VCS_CLASS = """VERSION 1.0 CLASS
BEGIN
  MultiUse = -1  'True
END
Attribute VB_Name = "cls明細"
Attribute VB_GlobalNameSpace = False
Attribute VB_Creatable = False
Attribute VB_PredeclaredId = False
Attribute VB_Exposed = False
Option Compare Database
Option Explicit

Private m_件数 As Long

Private Sub Class_Initialize()
    m_件数 = DCount("*", "T_受注")
End Sub

Public Property Get 件数() As Long
    件数 = m_件数
End Property

Public Property Let 件数(ByVal v As Long)
    m_件数 = v
End Property
"""
VCS_MODULE = """Attribute VB_Name = "modCommon"
Option Compare Database
Option Explicit

Public Function 今月の初日() As Date
    今月の初日 = DateSerial(Year(Date), Month(Date), 1)
End Function

Public Function 起動する()
#If VBA7 Then
    DoCmd.OpenForm "メニュー"
#Else
    MsgBox "この版では動きません": Exit Function
#End If
End Function
"""
VCS_AXL = ('<?xml version="1.0" encoding="UTF-16" standalone="no"?>\r\n<UserInterfaceMacro MinimumClientDesignVersion="14.0.0000.0000" '
           'xmlns="http://example.invalid/2009/11/application"><Statements><ConditionalBlock><If><Condition>[CurrentProject].[IsTrusted]</Condition><Statements>'
           '<Action Name="RunCode"><Argument Name="FunctionName">起動する()</Argument></Action></Statements></If><Else><Statements><Action Name="MessageBox">'
           '<Argument Name="Message">コンテンツを有効にしてください</Argument><Argument Name="Title">受注管理</Argument></Action></Statements></Else>'
           '</ConditionalBlock></Statements></UserInterfaceMacro>')


def vcs_macro(rows, axl=""):
    """マクロの書き出し（アクションの行と、あれば XML の定義を 70 字ずつに切った _AXL: の行）"""
    esc = lambda v: v.replace("\\", "\\\\").replace('"', '\\"').replace("\r\n", "\\015\\012")
    out = "Version =196611\nColumnsShown =0\n"
    for r in rows:
        out += "Begin\n" + "".join(f'    {k} ="{esc(v)}"\n' for k, v in r) + "End\n"
    for k in range(0, len(axl), 250):
        chunk = esc("_AXL:" + axl[k:k + 250])
        out += "Begin\n    Comment =" + "\n        ".join(f'"{chunk[j:j + 70]}"' for j in range(0, len(chunk), 70)) + "\nEnd\n"
    return out


def make_vcs_input(out):
    """Access を、版管理アドインで書き出したフォルダ（架空の「受注管理.accdb.src」）。
    仕込み: フォームのコードが隣の .cls／下のフォルダのモジュール／名前に記号のあるクエリ（%2F）／リンク テーブル（JSON）／起動時に開くフォーム／
    条件のブロックのあるマクロ（アクションの行＋Access が足した行＋XML の定義）／読まないファイル（画像・テーマ・テーブルの行の中身・設定）。文字コードは BOM つき UTF-8"""
    d = out / "受注管理.accdb.src"

    def put(rel, text):
        f = d / rel
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_bytes(b"\xef\xbb\xbf" + text.replace("\r\n", "\n").replace("\n", "\r\n").encode("utf-8"))
    put("vcs-options.json", '{\n  "Info": {"AddinVersion": "0.0.0"},\n  "Options": {}\n}\n')
    put("vbe-project.json", '{\n  "Info": {"Class": "clsDbVbeProject"},\n  "Items": {"Name": "受注管理", "FileName": "受注管理.accdb"}\n}\n')
    put("dbs-properties.json", '{\n  "Info": {"Class": "clsDbProperty"},\n  "Items": {"StartupForm": {"Value": "メニュー", "Type": 10}, "AppTitle": {"Value": "受注管理", "Type": 10}}\n}\n')
    put("forms/メニュー.form", VCS_FORM)
    put("forms/メニュー.cls", VCS_FORM_CODE)
    put("modules/modCommon.bas", VCS_MODULE)
    put("modules/明細/cls明細.cls", VCS_CLASS)
    cond = "[CurrentProject].[IsTrusted]"
    put("macros/AutoExec.macro", vcs_macro([
        [("Action", "SetLocalVar"), ("Argument", "__*L0_"), ("Argument", cond)],
        [("Action", "SetLocalVar"), ("Argument", "__*L0C_"), ("Argument", "[LocalVars]![__*L0_]")],
        [("Condition", "[LocalVars]![__*L0C_]"), ("Action", "RunCode"), ("Argument", "起動する()")],
        [("Action", "SetLocalVar"), ("Argument", "__*L0C_"), ("Argument", "(Not [LocalVars]![__*L0_])")],
        [("Condition", "[LocalVars]![__*L0C_]"), ("Action", "MsgBox"), ("Argument", "コンテンツを有効にしてください"), ("Argument", "-1"), ("Argument", "2"), ("Argument", "受注管理")],
        [("Action", "SetLocalVar"), ("Argument", "__*L0C_"), ("Argument", "False")]], VCS_AXL))
    put("macros/月次取込.macro", vcs_macro([[("Action", "OpenQuery"), ("Argument", "Q_月次/売上_追加"), ("Argument", "0"), ("Argument", "1")]]))
    put("queries/Q_月次%2F売上_追加.sql", "INSERT INTO T_月次売上 ( 品番, 数量 )\nSELECT T_受注.品番, Sum(T_受注.数量) AS 数量\nFROM T_受注 INNER JOIN L_得意先 ON T_受注.得意先ID = L_得意先.ID\nGROUP BY T_受注.品番;\n")
    put("queries/Q_月次%2F売上_追加.json", '{"Info": {"Class": "clsDbQuery"}, "Items": {}}\n')
    put("tbldefs/T_受注.sql", "CREATE TABLE [T_受注] (\n  [ID] AUTOINCREMENT CONSTRAINT [PrimaryKey] PRIMARY KEY UNIQUE NOT NULL,\n  [品番] VARCHAR (20),\n  [数量] LONG,\n  [受注日] DATETIME\n)\n")
    put("tbldefs/T_月次売上.sql", "CREATE TABLE [T_月次売上] (\n  [品番] VARCHAR (20),\n  [数量] LONG\n)\n")
    put("tbldefs/L_得意先.json", '{\n  "Info": {"Class": "clsDbTableDef"},\n  "Items": {"Name": "L_得意先", "Connect": "ODBC;DSN=基幹;PWD=himitsu99", "SourceTableName": "dbo.得意先"}\n}\n')
    put("tables/T_受注.txt", "ID\t品番\t数量\n1\tA-100\t3\n")
    put("themes/Office.json", "{}\n")
    put("AGENTS.md", "# 開発の手引き\n")
    (d / "images").mkdir(exist_ok=True)
    (d / "images" / "logo.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
    (out / "受注管理.accdb").write_bytes(b"\x00\x01\x00\x00Standard ACE DB\x00" + b"\x00" * 64)   # Access のファイルそのもの（書き出しがあるので、読まない）


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("out", help="入力一式を書き出すフォルダ")
    ap.add_argument("--code", action="store_true", help="プログラムのソースを含む入力（手順書1つと、マクロの書き出し）だけを作る")
    ap.add_argument("--atypical", action="store_true", help="原典の順が一般的な流れと逆の手順書1つだけを作る")
    ap.add_argument("--access", action="store_true", help="Access を書き出したテキスト（架空）だけを作る")
    ap.add_argument("--vcs", action="store_true", help="Access を、版管理アドインで書き出したフォルダ（架空）だけを作る")
    a = ap.parse_args()
    out = pathlib.Path(a.out).expanduser().resolve()
    if a.vcs:
        make_vcs_input(out)
        print(f"入力（Access を、版管理アドインで書き出したフォルダ） → {out}")
        return
    if a.access:
        make_access_input(out)
        print(f"入力（Access を書き出したテキスト） → {out}")
        return
    if a.code:
        make_code_input(out)
        print(f"入力（手順書とマクロの書き出し） → {out}")
        return
    if a.atypical:
        out.mkdir(parents=True, exist_ok=True)
        make_atypical_pdf(out / "仮払金の取扱手順.pdf")
        print(f"入力（逆順の手順書） → {out}")
        return
    (out / "00_全体").mkdir(parents=True, exist_ok=True)
    (out / "01_手順書").mkdir(parents=True, exist_ok=True)
    (out / "02_規程").mkdir(parents=True, exist_ok=True)
    make_pptx(out / "00_全体" / "経費精算業務マニュアル.pptx")
    make_pdf(out / "01_手順書" / "支払処理手順.pdf")
    make_xlsx(out / "01_手順書" / "経費精算チェックシート.xlsx")
    make_rule_pdf(out / "02_規程" / "経費精算規程.pdf")
    print(f"入力一式 → {out}")


if __name__ == "__main__":
    main()
