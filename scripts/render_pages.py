#!/usr/bin/env python3
"""render_pages.py — すべての入力を「ページ画像＋機械抽出」にそろえる。

  python3 scripts/render_pages.py -w <作業フォルダ>                 # まだ描画していない資料だけ
  python3 scripts/render_pages.py -w <作業フォルダ> --only M03,M07  # 資料を選ぶ
  python3 scripts/render_pages.py -w <作業フォルダ> --force         # 描画し直す

資料ごとに pages/<資料ID>/ へ:
  pNNNN.png   ページ画像（長辺 1568px。LLM はこれを見て読む）
  pNNNN.json  機械抽出。text（PDF の文字層）・shapes / connectors（図形と、矢印がどの図形をつないでいるか）・
              tables・cells（Excel のセル）・images（画像の数）・notes・small_text（字が小さく、拡大して読むべきページ）・
              vector_figure（線で描いた図がありそうなページ。PDF では矢印のつながりは取れないので画像で読む）・
              printed_candidates（ページの上端・下端にある数字と位置。刷られたページ番号の候補）・
              cell_at（Excel。cells の各行の、値のあるセルの番地）・align / misaligned（ページ画像の文字と、同じ番号の
              スライド・シートから抜いた文字の重なり。低ければ画像と機械抽出が別のページ）
ページの単位: PowerPoint＝スライド、PDF・Word＝ページ、Excel＝シート（印刷で複数ページになれば分ける）、
画像＝1枚、マクロ＝モジュール（画像なし。ソースを読む）、文字のファイル＝300行ごと。
機械抽出は LLM の読取の突き合わせ用。画像の中の文字は入らない（そこは LLM が画像を見て読む）。
"""
import argparse
import math
import pathlib
import re
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET

import code_struct
import extract_access
import extract_vba
import office_to_pdf as otp
from common import load, manifest, norm, page_stem, page_text, save, work_dir
from validate import validate

LONG_SIDE = 1568          # 画像の長辺（px）。これより大きくしても LLM 側で縮むだけ
TEXT_LINES = 300          # 文字のファイル・マクロを1ページに入れる行数
BAND = 0.10               # ページの上端・下端のこの割合の帯にある数字を、刷られたページ番号の候補にする
NUM_TOKEN = re.compile(r"^[\-‐－—–(（]*((?:[0-9A-Za-z]{1,2}[\-‐－—–])?\d{1,4})[\-‐－—–)）]*$")
NS = extract_vba.NS


class RenderError(Exception):
    pass


def local(tag):
    return tag.rsplit("}", 1)[-1]


def col_letter(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


# ---------- PDF のページを画像と文字にする ----------

def render_pdf_page(pdf, i, png):
    import fitz
    page = pdf[i]
    r = page.rect
    zoom = min(3.0, LONG_SIDE / max(r.width, r.height))
    page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False).save(str(png))
    lines = [ln.strip() for ln in page.get_text("text").splitlines() if ln.strip()]
    sizes = []
    for b in page.get_text("dict").get("blocks", []):
        for ln in b.get("lines", []):
            for sp in ln.get("spans", []):
                t = sp.get("text", "").strip()
                if t:
                    sizes += [sp["size"]] * len(t)
    info = {"text": lines, "images": len(page.get_images(full=True)),
            "links": sorted({l["uri"] for l in page.get_links() if l.get("uri")}),
            "size_pt": [round(r.width), round(r.height)], "zoom": round(zoom, 3)}
    cands = []                                                       # 刷られたページ番号の候補（番号かどうかは、資料の中での並びで決める）
    for x0, y0, x1, y1, w, *_ in page.get_text("words"):
        yc = (y0 + y1) / 2 / r.height if r.height else 0
        mt = NUM_TOKEN.match(w.strip())
        if mt and (yc <= BAND or yc >= 1 - BAND):
            cands.append({"n": mt.group(1), "x": round((x0 + x1) / 2 / r.width, 2), "y": round(yc, 2)})
    info["printed_candidates"] = cands[:12]
    if not lines and info["images"]:
        info["no_text_layer"] = True                                 # スキャン。文字は画像からしか読めない
    if sizes:
        sizes.sort()
        small = [x for x in sizes if x * zoom < 10]                  # 画像の上で 10px に満たない字
        # 本文が大きくても、図・帳票の中だけ字が小さいページがある。小さい字が 20 字以上あれば、分割して拡大して読む
        if len(small) >= 20 or sizes[len(sizes) // 2] * zoom < 10:
            need = 12 / (small[len(small) // 2] if small else sizes[len(sizes) // 2])
            info["small_text"] = True
            info["small_chars"] = len(small)
            info["suggest_grid"] = [max(1, math.ceil(r.width * need / LONG_SIDE)),
                                    max(1, math.ceil(r.height * need / LONG_SIDE))]
    kinds = {"l": 0, "c": 0, "re": 0}                                # 線で描かれた図（箱・矢印）。PDF では図形や矢印のつながりは取れず、画像でしか読めない
    try:
        for dr in page.get_drawings():
            for it in dr.get("items", []):
                if it[0] in kinds:
                    kinds[it[0]] += 1
    except Exception:
        pass
    info["drawings"] = {"line": kinds["l"], "curve": kinds["c"], "rect": kinds["re"]}
    if kinds["l"] + kinds["c"] >= 10:
        info["vector_figure"] = True                                 # 罫線の表だけのページは矩形ばかりで、線・曲線は少ない
    return info


def grams3(t):
    return {t[i:i + 3] for i in range(max(0, len(t) - 2))}


def alignment(pdf_lines, extra):
    """ページ画像（PDF）の文字が、同じ番号のスライド・シートから抜いた文字にどれだけ入っているか（0〜1）。比べる文字が少なければ None。
    非表示のスライド・1シートが複数ページ・変換で改ページが変わった、で対応が1つずれると、ここが下がる"""
    a, b = norm("".join(pdf_lines)).replace("|", ""), norm(page_text(extra)).replace("|", "")
    if len(a) < 20 or len(b) < 20:
        return None
    ga = grams3(a)
    return round(len(ga & grams3(b)) / len(ga), 2)


def pdf_pages(pdf_path, doc_dir, entries, base):
    """PDF の各ページを描画する。entries は [(ページ番号, ラベル, 種別, 追加の機械抽出)]（PDF のページ順）"""
    import fitz
    pdf = fitz.open(str(pdf_path))
    out = []
    for i, (no, label, kind, extra) in enumerate(entries):
        if i >= pdf.page_count:
            break
        png = doc_dir / f"{page_stem(no)}.png"
        pj = dict(base, page=no, label=label, kind=kind)
        pj.update(render_pdf_page(pdf, i, png))
        if extra and kind in ("slide", "sheet"):
            al = alignment(pj.get("text") or [], extra)
            if al is not None:
                pj["align"] = al
                if al < 0.5:
                    pj["misaligned"] = True
        pj.update(extra or {})
        if extra and "pictures" in extra:                            # Office の文書は、PDF 上の画像数（影なども数える）ではなく貼られた画像・グラフの数
            pj["images"] = extra["pictures"] + extra.get("charts", 0)
        pj["render"] = {"pdf": str(pdf_path), "pdf_page": i}
        save(doc_dir / f"{page_stem(no)}.json", pj)
        out.append({"page": no, "label": label, "kind": kind, "png": f"pages/{base['doc_id']}/{png.name}",
                    "json": f"pages/{base['doc_id']}/{page_stem(no)}.json",
                    **({"sheet": extra["sheet"]} if kind == "sheet" and (extra or {}).get("sheet") else {})})
    pdf.close()
    return out


def text_pages(lines, doc_dir, base, first_no, label, kind, cell_at=None):
    """画像のないページ（文字のファイル・マクロのソース・非表示シート）"""
    out = []
    chunks = [lines[i:i + TEXT_LINES] for i in range(0, max(len(lines), 1), TEXT_LINES)]
    for k, chunk in enumerate(chunks):
        no = first_no + k
        lab = label if len(chunks) == 1 else f"{label}（{k + 1}/{len(chunks)}）"
        pj = dict(base, page=no, label=lab, kind=kind, text=chunk, images=0, line_from=k * TEXT_LINES + 1)
        if cell_at:
            pj["cell_at"] = cell_at[k * TEXT_LINES:(k + 1) * TEXT_LINES]
        save(doc_dir / f"{page_stem(no)}.json", pj)
        out.append({"page": no, "label": lab, "kind": kind, "png": None,
                    "json": f"pages/{base['doc_id']}/{page_stem(no)}.json", **({"sheet": base["sheet"]} if base.get("sheet") else {})})
    return out


def code_pages(lines, doc_dir, base, first_no, label, kind, meta, offset=0):
    """プログラムのソースのページ（画像なし。1行が1単位）。手続きの切れ目で切る（1つの手続きを途中で切らない。長い手続きだけ TEXT_LINES 行で続きにする）。
    認証情報（接続文字列のパスワード・キー・トークン）は、ページを書く前に伏せる。戻り値は (ページの一覧, 伏せた行の数)。
    offset: モジュールの中での、このファイルの最初の行のずれ（フォームの定義のあとにつなぐコードのファイル）"""
    lines, hit = code_struct.mask_secrets(lines)
    cuts = code_struct.page_cuts(meta["language"], lines, TEXT_LINES)
    out = []
    for k, at in enumerate(cuts):
        end = cuts[k + 1] - 1 if k + 1 < len(cuts) else len(lines)
        no = first_no + k
        lab = label if len(cuts) == 1 else f"{label}（{k + 1}/{len(cuts)}）"
        code = dict(meta, line_from=at + offset, secrets=[x - at + 1 for x in hit if at <= x <= end])
        save(doc_dir / f"{page_stem(no)}.json", dict(base, page=no, label=lab, kind=kind, text=lines[at - 1:end], images=0, line_from=at + offset, code=code))
        out.append({"page": no, "label": lab, "kind": kind, "png": None, "json": f"pages/{base['doc_id']}/{page_stem(no)}.json",
                    "code": {"language": meta["language"], "module": meta["module"], "line_from": at + offset}})
    return out, len(hit)


def vba_pages(d, path, doc_dir, base, first_no, read_locked=False):
    res = extract_vba.extract(path)
    if res["error"]:
        d["render_note"] = (d.get("render_note", "") + " " + res["error"]).strip()
    out, no, masked, total = [], first_no, 0, 0
    d["vba_modules"] = [{"name": m["name"], "type": m["type"], "procedures": m["procedures"], "touches": m["touches"]}
                        for m in res["modules"]]
    d["vba_buttons"], d["vba_sheets"], d["vba_shortcuts"], d["vba_locked"] = res["buttons"], res["sheets"], res["shortcuts"], bool(res["locked"])
    for key, val in (("vba_ribbon", res["ribbon"]), ("vba_book", res["book"]), ("vba_unread", res["unread"])):
        if val:
            d[key] = val
        else:
            d.pop(key, None)
    if res["locked"] and not read_locked:                            # 保護の意思表示があるものは、利用者に伺ってから読む
        d["render_note"] = (d.get("render_note", "") + " VBA プロジェクトに閲覧のロックが掛かっているので、マクロは読んでいない"
                            "（読んでよいかを利用者に確かめ、よければ render_pages.py --read-locked --force --only <資料ID> で読む）").strip()
        return out
    for m in res["modules"]:
        kind = code_struct.module_kind(m["name"], m["ext"], res["sheets"], res["book"])
        meta = {"language": "vba", "module": m["name"], "module_kind": kind}
        if kind == "sheet":
            meta["sheet"] = next(v for k, v in res["sheets"].items() if k.lower() == m["name"].lower())
        lines = m["code"].split("\n")
        pages, hit = code_pages(lines, doc_dir, base, no, f"VBA:{m['name']}", "vba", meta)
        out += pages
        no, masked, total = no + len(pages), masked + hit, total + len(lines)
    if res["modules"]:
        d["code"] = {"language": "vba", "modules": len(res["modules"]), "lines": total, "secrets": masked}
    return out


# ---------- PowerPoint ----------

def pptx_extract(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.shapes.connector import Connector
    prs = Presentation(str(path))
    W, H = prs.slide_width or 1, prs.slide_height or 1
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        shapes, conns, tables, images = [], [], [], [0]

        def walk(shs):
            for sh in shs:
                try:
                    st = sh.shape_type
                except Exception:
                    st = None
                if st == MSO_SHAPE_TYPE.GROUP:
                    walk(sh.shapes)
                    continue
                if isinstance(sh, Connector):
                    el = sh._element
                    s, e = el.find(f".//{A}stCxn"), el.find(f".//{A}endCxn")
                    conns.append({"id": str(sh.shape_id), "from": s.get("id") if s is not None else None,
                                  "to": e.get("id") if e is not None else None})
                    continue
                if st == MSO_SHAPE_TYPE.PICTURE:
                    images[0] += 1
                    continue
                if getattr(sh, "has_table", False) and sh.has_table:
                    tables.append([[c.text.strip() for c in row.cells] for row in sh.table.rows])
                    continue
                text = sh.text_frame.text.strip() if getattr(sh, "has_text_frame", False) and sh.has_text_frame else ""
                kind = ""
                try:
                    kind = str(sh.auto_shape_type).split(".")[-1].split(" ")[0]
                except Exception:
                    kind = "placeholder" if getattr(sh, "is_placeholder", False) else (str(st).split(".")[-1].split(" ")[0] if st else "")
                if not text and kind in ("", "placeholder"):
                    continue
                bbox = [round((sh.left or 0) / W, 3), round((sh.top or 0) / H, 3),
                        round((sh.width or 0) / W, 3), round((sh.height or 0) / H, 3)]
                shapes.append({"id": str(sh.shape_id), "name": sh.name, "type": kind, "text": text.replace("\x0b", "\n"), "bbox": bbox})   # 行内改行

        walk(slide.shapes)
        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text.strip()
        except Exception:
            pass
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"slide": idx, "hidden": slide._element.get("show") == "0", "title": title, "shapes": shapes,
                       "connectors": conns, "tables": tables, "pictures": images[0], "notes": notes})
    sections = []
    ids = [s.get("id") for s in prs.slides._sldIdLst]
    for sec in prs.part._element.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}section"):
        nums = [ids.index(x.get("id")) + 1 for x in sec.iter("{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId")
                if x.get("id") in ids]
        sections.append({"name": sec.get("name"), "slides": nums})
    return slides, sections


def slide_extra(s):
    """スライド1枚の機械抽出（pptx から抜いた文字・図形・表）"""
    return {"title": s["title"], "shapes": s["shapes"], "connectors": s["connectors"], "tables": s["tables"],
            "pictures": s["pictures"], "notes": s["notes"], **({"hidden": True} if s["hidden"] else {})}


def do_pptx(work, d, a, doc_dir, base):
    tmp = otp.stage_dir()
    safe = otp.defang(d["abs_path"], tmp)
    try:
        pdf = otp.to_pdf(safe, tmp, a.backend)
        src_pdf = doc_dir / "_src.pdf"
        shutil.move(str(pdf), str(src_pdf))
        slides, sections = pptx_extract(safe)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    import fitz
    n_pdf = fitz.open(str(src_pdf)).page_count
    visible = [s for s in slides if not s["hidden"]]
    use = slides if n_pdf == len(slides) else visible if n_pdf == len(visible) else None
    d["sections"] = sections
    d["hidden_slides"] = [s["slide"] for s in slides if s["hidden"]]
    if use is None:
        d["render_note"] = f"スライド {len(slides)} 枚に対して PDF が {n_pdf} ページ。スライド番号との対応が取れていない"
        entries = [(i + 1, f"p.{i + 1}", "slide", None) for i in range(n_pdf)]
    else:
        entries = [(s["slide"], f"スライド{s['slide']}" + ("（非表示）" if s["hidden"] else ""), "slide", slide_extra(s)) for s in use]
    pages = pdf_pages(src_pdf, doc_dir, entries, base)
    if use is visible and use is not slides:                         # 非表示のスライドは PDF に出ない。文字だけのページとして残す（Excel の非表示シートと同じ）
        for s in slides:
            if s["hidden"]:
                pj = dict(base, page=s["slide"], label=f"スライド{s['slide']}（非表示）", kind="slide", text=[], images=0, **slide_extra(s))
                save(doc_dir / f"{page_stem(s['slide'])}.json", pj)
                pages.append({"page": s["slide"], "label": pj["label"], "kind": "slide", "png": None,
                              "json": f"pages/{base['doc_id']}/{page_stem(s['slide'])}.json"})
        pages.sort(key=lambda p: p["page"])
    if d.get("has_macro"):
        pages += vba_pages(d, d["abs_path"], doc_dir, base, max([p["page"] for p in pages] + [len(slides)]) + 1)
    return pages


# ---------- Excel ----------

def drawing_items(z, part):
    shapes, conns, counts = [], [], {"images": 0, "charts": 0}

    def walk(el, cell):
        for ch in el:
            t = local(ch.tag)
            if t == "sp":
                nv = ch.find(".//xdr:cNvPr", NS)
                geom = ch.find(".//a:prstGeom", NS)
                paras = ["".join(x.text or "" for x in p.iter(f"{{{NS['a']}}}t")) for p in ch.iter(f"{{{NS['a']}}}p")]
                text = "\n".join(p for p in paras if p.strip())
                shapes.append({"id": nv.get("id") if nv is not None else "", "name": nv.get("name") if nv is not None else "",
                               "type": geom.get("prst") if geom is not None else "", "text": text, "cell": cell})
            elif t == "cxnSp":
                nv = ch.find(".//xdr:cNvPr", NS)
                s, e = ch.find(".//a:stCxn", NS), ch.find(".//a:endCxn", NS)
                conns.append({"id": nv.get("id") if nv is not None else "", "from": s.get("id") if s is not None else None,
                              "to": e.get("id") if e is not None else None, "cell": cell})
            elif t == "pic":
                counts["images"] += 1
            elif t == "graphicFrame":
                counts["charts"] += 1
            elif t == "grpSp":
                walk(ch, cell)

    for anchor in ET.fromstring(z.read(part)):
        if local(anchor.tag) not in ("twoCellAnchor", "oneCellAnchor", "absoluteAnchor"):
            continue
        frm = anchor.find("xdr:from", NS)
        cell = ""
        if frm is not None:
            cell = f"{col_letter(int(frm.findtext('xdr:col', '0', NS)))}{int(frm.findtext('xdr:row', '0', NS)) + 1}"
        walk(anchor, cell)
    return shapes, conns, counts


def xlsx_extract(path, max_rows=3000, max_chars=30000):
    """シートごとの機械抽出: セル（行ごとに「 | 」でつなぐ。cell_at に各行のセルの番地）・図形・コネクタ・画像とグラフの数"""
    import openpyxl
    sheets = []
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        for name, part, state in extract_vba.sheet_parts(z):
            s = {"name": name, "state": state, "cells": [], "cell_at": [], "shapes": [], "connectors": [], "pictures": 0,
                 "charts": 0, "chartsheet": "/chartsheets/" in part}
            for dr in extract_vba.rel_targets(z, part, "/drawing"):
                if dr in names:
                    sh, cn, cnt = drawing_items(z, dr)
                    s["shapes"] += sh
                    s["connectors"] += cn
                    s["pictures"] += cnt["images"]
                    s["charts"] += cnt["charts"]
            sheets.append(s)
    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        for s in sheets:
            if s["chartsheet"] or s["name"] not in wb.sheetnames:
                continue
            chars = 0
            for r, row in enumerate(wb[s["name"]].iter_rows(values_only=True), start=1):   # 読み取り専用では A1 から順に返る
                if r > max_rows or chars >= max_chars:
                    s["cells_truncated"] = True
                    break
                vals = [(i, str(v).strip()) for i, v in enumerate(row) if v is not None and str(v).strip()]
                if vals:
                    line = " | ".join(v for _i, v in vals)
                    s["cells"].append(line)
                    s["cell_at"].append([f"{col_letter(i)}{r}" for i, _v in vals])
                    chars += len(line)
        wb.close()
    except Exception as e:
        for s in sheets:
            s["cells_error"] = f"{type(e).__name__}: {e}"
    return sheets


def sheet_extra(s):
    return {"sheet": s["name"], "cells": s["cells"], "cell_at": s["cell_at"], "shapes": s["shapes"],
            "connectors": s["connectors"], "pictures": s["pictures"], "charts": s["charts"]}


def do_xlsx(work, d, a, doc_dir, base):
    import fitz
    legacy = d["ext"] == ".xls"
    sheets = [] if legacy else xlsx_extract(d["abs_path"])
    how = otp.pick_backend("excel", a.backend)
    if legacy and how == "office" and otp.legacy_has_macro(d["abs_path"]):
        raise RenderError("旧形式（xls）のマクロ付きは Microsoft Office では開かない。LibreOffice を入れるか、xlsm に保存し直す")
    tmp = otp.stage_dir()
    safe = otp.defang(d["abs_path"], tmp)
    pages, no = [], 1
    try:
        by_name = {s["name"]: s for s in sheets}
        if how == "office":
            for i, name, pdf in otp.excel_sheet_pdfs(safe, tmp):
                if pdf is None:
                    continue
                keep = doc_dir / f"_sheet_{i}.pdf"
                shutil.move(str(pdf), str(keep))
                n = fitz.open(str(keep)).page_count
                shown = min(n, a.max_sheet_pages)
                entries = []
                for k in range(shown):
                    extra = sheet_extra(by_name[name]) if (k == 0 and name in by_name) else {"sheet": name}
                    if k == 0 and n > shown:
                        extra["truncated_pages"] = n - shown
                    entries.append((no + k, name if n == 1 else f"{name}（{k + 1}/{n}）", "sheet", extra))
                pages += pdf_pages(keep, doc_dir, entries, base)
                no += shown
                by_name.pop(name, None)
        else:
            pdf = otp.soffice_pdf(safe, tmp, single_page_sheets=True)
            keep = doc_dir / "_src.pdf"
            shutil.move(str(pdf), str(keep))
            doc = fitz.open(str(keep))
            texts = [norm(doc[i].get_text()) for i in range(doc.page_count)]
            doc.close()
            cand = [s for s in sheets if s["state"] == "visible"
                    and (s["cells"] or s["shapes"] or s["pictures"] or s["charts"] or s["chartsheet"])]
            if len(cand) != len(texts):                              # 数が合わなければ、文字の重なりで順に当てる
                picked, j = [], 0
                for t in texts:
                    best = None
                    for k in range(j, len(cand)):
                        head = norm("".join(cand[k]["cells"][:5]))[:40]
                        if head and head[:12] in t:
                            best = k
                            break
                    picked.append(cand[best] if best is not None else None)
                    j = best + 1 if best is not None else j
                cand = picked
                if any(c is None for c in cand):
                    d["render_note"] = "PDF のページとシートの対応が取れないものがある（「シート不明」）"
            entries = []
            for k in range(len(texts)):
                s = cand[k] if k < len(cand) else None
                entries.append((no + k, s["name"] if s else f"シート不明({k + 1})", "sheet", sheet_extra(s) if s else {}))
                if s:
                    by_name.pop(s["name"], None)
            pages += pdf_pages(keep, doc_dir, entries, base)
            no += len(entries)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    # 画像にならなかったシート（非表示・空）。中身があれば文字だけのページとして残す
    lost = [s["name"] for s in sheets if s["name"] in by_name and s["state"] == "visible"
            and (s["cells"] or s["shapes"] or s["pictures"] or s["charts"])]
    if lost:
        d["render_note"] = (d.get("render_note", "") + f" 画像にできなかったシート: {'、'.join(lost)}（セルの文字と図形の文字だけ読める。画像は無い。Excel の画面に確認のダイアログが出ていないか確かめて再実行する）").strip()
    hidden = []
    for s in sheets:
        if s["name"] in by_name and (s["cells"] or s["shapes"]):
            lab = f"{s['name']}（{'非表示' if s['state'] != 'visible' else '画像なし'}）"
            ps = text_pages(s["cells"], doc_dir, dict(base, sheet=s["name"], hidden=s["state"] != "visible", shapes=s["shapes"],
                                                      connectors=s["connectors"]), no, lab, "sheet", cell_at=s["cell_at"])
            pages += ps
            no += len(ps)
        if s["state"] != "visible":
            hidden.append(s["name"])
    d["hidden_sheets"] = hidden
    d["sheets"] = [s["name"] for s in sheets]
    if d.get("has_macro") or legacy:
        pages += vba_pages(d, d["abs_path"], doc_dir, base, no, a.read_locked)
    return pages


# ---------- Word・旧形式・画像・文字 ----------

def do_doc(work, d, a, doc_dir, base):
    """Word と、旧形式の PowerPoint。PDF にしてページ単位で読む（図形の機械抽出はない）"""
    app = otp.APP_OF_EXT.get(d["ext"])
    legacy = d["ext"] in (".doc", ".ppt", ".pps")
    if legacy and otp.pick_backend(app, a.backend) == "office" and otp.legacy_has_macro(d["abs_path"]):
        raise RenderError("旧形式のマクロ付きは Microsoft Office では開かない。LibreOffice を入れるか、新形式で保存し直す")
    tmp = otp.stage_dir()
    safe = otp.defang(d["abs_path"], tmp)
    try:
        pdf = otp.to_pdf(safe, tmp, a.backend)
        src_pdf = doc_dir / "_src.pdf"
        shutil.move(str(pdf), str(src_pdf))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    import fitz
    n = fitz.open(str(src_pdf)).page_count
    unit = "スライド" if app == "powerpoint" else "p."
    pages = pdf_pages(src_pdf, doc_dir, [(i + 1, f"{unit}{i + 1}", "slide" if app == "powerpoint" else "page", None)
                                         for i in range(n)], base)
    if d.get("has_macro") or legacy:
        pages += vba_pages(d, d["abs_path"], doc_dir, base, n + 1, a.read_locked)
    return pages


def do_pdf(work, d, a, doc_dir, base):
    import fitz
    pdf = fitz.open(d["abs_path"])
    if pdf.needs_pass:
        raise RenderError("パスワード付きの PDF は開けない")
    d["outline"] = [{"level": lv, "title": t, "page": p} for lv, t, p in pdf.get_toc()]
    n = pdf.page_count
    pdf.close()
    return pdf_pages(d["abs_path"], doc_dir, [(i + 1, f"p.{i + 1}", "page", None) for i in range(n)], base)


def do_image(work, d, a, doc_dir, base):
    from PIL import Image
    im = Image.open(d["abs_path"])
    im = im.convert("RGB")
    scale = min(1.0, LONG_SIDE / max(im.size))
    if scale < 1:
        im = im.resize((round(im.width * scale), round(im.height * scale)))
    png = doc_dir / f"{page_stem(1)}.png"
    im.save(png)
    save(doc_dir / f"{page_stem(1)}.json", dict(base, page=1, label="画像", kind="image", text=[], images=1,
                                               no_text_layer=True, render={"image": d["abs_path"]}))
    return [{"page": 1, "label": "画像", "kind": "image", "png": f"pages/{d['doc_id']}/{png.name}",
             "json": f"pages/{d['doc_id']}/{page_stem(1)}.json"}]


def read_text(path):
    raw = pathlib.Path(path).read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff") or raw[:400].count(b"\x00") > 40:   # UTF-16（Access の書き出しなど）。cp932 として読むと、誤りにならずに化ける
        try:
            return raw.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for enc in ("utf-8-sig", "cp932", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def do_text(work, d, a, doc_dir, base):
    return text_pages(read_text(d["abs_path"]).splitlines(), doc_dir, base, 1, "本文", "text")


def do_code(work, d, a, doc_dir, base):
    """プログラムのソース（文字として読む。実行も取り込みもしない）。Access を書き出したテキストは、フォーム・レポート・マクロ・定義として読む"""
    lines = read_text(d["abs_path"]).splitlines()
    acc = extract_access.describe(d["abs_path"], lines)
    offset = 0
    if acc:
        lang, name = acc["language"], acc["module"]
        meta = {"language": lang, "module": name, "module_kind": acc["module_kind"]}
        d["access"] = {"object": acc["access"], "name": acc["name"]}
        if acc.get("root"):                                          # 版管理アドインの書き出し: 根のフォルダ（入力からの相対）で、1つのまとまりにする
            depth = len(pathlib.Path(d["abs_path"]).relative_to(acc["root"]).parts)
            d["access"]["root"] = str(pathlib.PurePosixPath(*pathlib.PurePath(d["rel_path"]).parts[:-depth])) if depth < len(pathlib.PurePath(d["rel_path"]).parts) else "."
        if acc.get("part") == "code":                                # コードだけのファイル: 同じ名前の定義のあとにつなぐ
            meta["part"] = d["access"]["part"] = "code"
            lay = extract_access.layout_file(d["abs_path"])
            offset = len(read_text(str(lay)).splitlines()) if lay else 0
    else:
        lang = code_struct.LANG_OF_EXT.get(d["ext"], d["ext"].lstrip("."))
        name = code_struct.module_name(lines, pathlib.Path(d["name"]).stem) if lang == "vba" else d["name"]
        meta = {"language": lang, "module": name, "module_kind": code_struct.module_kind(name, d["ext"]) if lang == "vba" else "file"}
        d.pop("access", None)
    pages, hit = code_pages(lines, doc_dir, base, 1, name, "code", meta, offset)
    d["code"] = {"language": lang, "modules": 1, "lines": len(lines), "secrets": hit}
    return pages


def do_macrobook(work, d, a, doc_dir, base):
    """シートを開かずに、マクロ（VBA）だけを読む（バイナリ ブック・アドイン）"""
    pages = vba_pages(d, d["abs_path"], doc_dir, base, 1, a.read_locked)
    if not pages and not d.get("vba_locked"):
        raise RenderError("マクロ（VBA）を読めない・入っていない。シートも読むなら、xlsm で保存し直したものを入力にする")
    d["render_note"] = (d.get("render_note", "") + " マクロだけを読んだ（シートは読んでいない。シートも読むなら、xlsm で保存し直したものを入力にする）").strip()
    return pages


def do_access(work, d, a, doc_dir, base):
    """Access のファイルそのものは開かない。書き出したテキスト（access_export.txt の source が、このファイルの名前）が入力にあれば、そちらを読む"""
    d.pop("read_via", None)
    d.pop("read_direct", None)
    for x in manifest(work)["docs"]:
        if x["name"].lower() in ("access_export.txt", "vbe-project.json"):   # 書き出しの道具の一覧／版管理アドインの書き出し
            folder = pathlib.Path(x["abs_path"]).parent
            idx = extract_access.log_index(str(folder)) or extract_access.parse_log(extract_access.read_head(folder / x["name"], 200_000).splitlines())
            if idx and idx["source"].lower() == d["name"].lower():
                d["read_via"] = str(pathlib.PurePosixPath(x["rel_path"]).parent)
                if idx.get("via") == "direct":                       # 棚卸し（inventory.py）が、ファイルを直接読んで作ったテキスト
                    d["read_direct"] = True
                    d["render_note"] = f"ファイルを直接読んで作ったテキスト（{d['read_via']}）を読む。Access では開いていない。読めた範囲だけ"
                else:
                    d["render_note"] = f"書き出したテキスト（{d['read_via']}）を読む。このファイルは開いていない"
                return []
    import access_direct
    try:
        access_direct.JetFile(d["abs_path"])
        why = "直接読んだテキストが、まだ無い。inventory.py を（--no-direct を付けずに）もう一度実行する。それでも読めなければ"
    except access_direct.Unreadable as e:
        why = f"直接は読めない（{e}）。"
    raise RenderError(f"Access のファイルを読めていない: {why}Access のあるパソコンで、書き出しの道具（assets/access_export/）を使ってテキストに書き出し、"
                      "書き出したフォルダ（<ファイル名>_export）を入力に足す（手順は references/access-export.md）")


HANDLERS = {"pdf": do_pdf, "pptx": do_pptx, "xlsx": do_xlsx, "docx": do_doc, "image": do_image, "text": do_text, "code": do_code,
            "macrobook": do_macrobook, "access": do_access}
LEGACY = {".xls": do_xlsx, ".doc": do_doc, ".ppt": do_doc, ".pps": do_doc}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-w", "--work", required=True)
    ap.add_argument("--only", help="資料IDをカンマ区切りで")
    ap.add_argument("--force", action="store_true")
    ap.add_argument("--backend", choices=["auto", "office", "soffice"], default="auto")
    ap.add_argument("--max-sheet-pages", type=int, default=15, help="1シートが印刷で何ページにもなるとき、先頭の何ページまで読むか")
    ap.add_argument("--read-locked", action="store_true", help="閲覧のロックが掛かった VBA プロジェクトのマクロも読む（利用者に確かめてから）")
    a = ap.parse_args()
    work = work_dir(a.work)
    m = manifest(work)
    only = set(a.only.split(",")) if a.only else None
    done = failed = skipped = 0
    stuck = {}                                                       # 止まった Office アプリ → 理由（残りの資料で待ち続けないため）
    for d in m["docs"]:
        if only and d["doc_id"] not in only:
            continue
        if d.get("duplicate_of"):
            continue
        fn = HANDLERS.get(d["kind"]) or (LEGACY.get(d["ext"]) if d["kind"] == "legacy" else None)
        if fn is None:
            d["render_note"] = "未対応の形式（読んでいない）"
            continue
        if d.get("rendered") and not a.force:
            skipped += 1
            continue
        app = otp.APP_OF_EXT.get(otp.PLAIN_EXT.get(d["ext"], d["ext"]))
        if app in stuck:
            d["rendered"], d["page_list"] = False, []
            d["render_note"] = f"先の資料で {otp.APP_NAME[app]} が止まったので読んでいない（{stuck[app]}）"
            failed += 1
            print(f"{d['doc_id']} {d['rel_path']}: 飛ばした — {otp.APP_NAME[app]} が止まっている")
            continue
        doc_dir = work / "pages" / d["doc_id"]
        if doc_dir.exists():                                         # 描画し直しても、拡大して読んだ跡（crops/）は残す
            for x in doc_dir.iterdir():
                if x.name == "crops":
                    continue
                shutil.rmtree(x) if x.is_dir() else x.unlink()
        doc_dir.mkdir(parents=True, exist_ok=True)
        d.pop("render_note", None)
        base = {"doc_id": d["doc_id"], "doc_name": d["name"]}
        try:
            if d["kind"] in ("xlsx", "pptx", "docx", "macrobook", "legacy") and extract_vba.is_encrypted(d["abs_path"]):
                raise RenderError("ファイルが暗号化されている（パスワードで保護して保存してある）。中は読んでいない。保護を外した写しをもらう")
            d["page_list"] = fn(work, d, a, doc_dir, base)
            d["pages"] = len(d["page_list"])
            d["rendered"] = True
            done += 1
            print(f"{d['doc_id']} {d['rel_path']}: {d['pages']} ページ" + (f"  ※ {d['render_note']}" if d.get("render_note") else ""))
        except (RenderError, otp.ConvertError) as e:
            d["rendered"], d["page_list"], d["render_note"] = False, [], str(e)
            failed += 1
            print(f"{d['doc_id']} {d['rel_path']}: 描画できない — {e}")
            if app and ("応えない" in str(e) or otp.is_timeout(e)):
                stuck[app] = "画面のダイアログを閉じて再実行するか、LibreOffice を入れる"
        except Exception as e:                                       # 1件の失敗で全体を止めない
            d["rendered"], d["page_list"], d["render_note"] = False, [], f"{type(e).__name__}: {e}"
            failed += 1
            print(f"{d['doc_id']} {d['rel_path']}: 描画できない — {type(e).__name__}: {e}")
        save(work / "manifest.json", m)                              # 1資料ごとに保存（途中で止まっても続きから）
    total = sum(len(d.get("page_list") or []) for d in m["docs"])
    small = 0
    for d in m["docs"]:
        for p in d.get("page_list") or []:
            pj = load(work / p["json"], {})
            small += bool(pj.get("small_text"))
    codes = [d for d in m["docs"] if d.get("code")]
    if codes:
        print(f"プログラムのソース {len(codes)} 件・{sum(d['code']['lines'] for d in codes)} 行"
              + (f"（認証情報を伏せた行 {sum(d['code']['secrets'] for d in codes)}）" if any(d["code"]["secrets"] for d in codes) else "")
              + "。構造の取り出しは python3 scripts/code_struct.py -w <作業フォルダ>")
    for d in m["docs"]:
        if d.get("vba_locked") and not a.read_locked:
            print(f"※ {d['doc_id']} {d['rel_path']}: VBA プロジェクトに閲覧のロックが掛かっている。マクロは読んでいない（読んでよいかを利用者に確かめる）")
    print(f"\n描画 {done} 件 / 済み {skipped} 件 / 失敗 {failed} 件。全 {total} ページ"
          + (f"（うち字が小さく拡大して読むページ {small}。crop_page.py を使う）" if small else ""))
    if failed:
        print("失敗した資料は manifest.json の render_note に理由がある。直せないものは確認事項に載る")
    bad = off = unknown = 0
    for d in m["docs"]:
        for p in d.get("page_list") or []:
            pj = load(work / p["json"], {}) or {}
            bad += bool(validate(pj, "page"))
            off += bool(pj.get("misaligned"))
            unknown += str(p.get("label", "")).startswith("シート不明")
    if bad:
        print(f"※ 機械抽出の形が崩れているページ {bad}（スクリプトの不具合の疑い。python3 scripts/validate.py --file page <ファイル> で見る）")
    merrs = validate(m, "manifest")
    for e in merrs[:10]:
        print(f"※ manifest.json の形が崩れている（スクリプトの不具合の疑い）: {e}")
    if off:
        print(f"※ ページ画像と機械抽出が合わないページ {off}（別のスライド・シートを指している。--force でやり直す。check_extract.py が FAIL にする）")
    if unknown:
        print(f"※ どのシートの画像か対応が取れないページ {unknown}（出典にシート名を書けない。--backend office でやり直すか、利用者に知らせる）")
    sys.exit(1 if (failed and not done and not skipped) or bad or merrs else 0)  # 形の崩れた抽出のまま先へ進まない


if __name__ == "__main__":
    main()
