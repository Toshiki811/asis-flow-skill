#!/usr/bin/env python3
"""check_flow.py — 階層・業務分類・フロー・書き出した draw.io の機械チェック。FAIL 0 にしてから納品する。

  python3 scripts/check_flow.py -w <作業フォルダ>                       # ある成果物をすべて点検
  python3 scripts/check_flow.py -w <作業フォルダ> --stage hierarchy     # 階層の判定だけ（業務分類を決める前）
  python3 scripts/check_flow.py -w <作業フォルダ> --stage structure     # 階層・用語・業務分類だけ
  python3 scripts/check_flow.py -w <作業フォルダ> --stage flow          # flow.json だけ（draw.io を書き出す前）
  python3 scripts/check_flow.py -w <作業フォルダ> --stage drawio        # 書き出した draw.io だけ
  python3 scripts/check_flow.py -w <作業フォルダ> --stage matrix        # 書き出した業務マトリクス一覧（xlsx）だけ
  python3 scripts/check_flow.py -w <作業フォルダ> --stage review        # 原典照合の採否表だけ（工程 8）
  python3 scripts/check_flow.py -w <作業フォルダ> --stage reports       # 書き出した報告書（確認事項一覧・出典対応表・階層と業務分類・採否表）だけ
  python3 scripts/check_flow.py -w <作業フォルダ> --stage deliverables  # 納品物の文字（内部の番号・場所・項目名・【下書き】）だけ
  python3 scripts/check_flow.py -w <作業フォルダ> --stage flow --only F03   # 自分のフローの行だけを出す
  python3 scripts/check_flow.py -w <作業フォルダ> --report <作業フォルダ>/out/チェック_フロー.txt

正典: references/flow-rules.md（§2 階層と分類・§3 フロー・§4 作図）
[structure]
  FAIL  階層の判定（hierarchy.json）・用語の正本（glossary.json）・業務分類（taxonomy.json）が無い・スキーマと違う・存在しない資料やページを指す
  FAIL  役割の無い資料／どの業務にも結び付かない資料（対象外の理由も無い）／根拠の無い紐付け／上位・下位の関係が業務の木に落ちていない
  FAIL  「続く」が業務の木に無い業務を指す／分類の親子が循環している
  FAIL  目次に業務フローの一覧があるのに、分類がそれに準拠していない／目次の項目に対応する分類が無い／名称が目次と違う
  FAIL  目次が指すページが割当に入っていない（N ページずれ）／分類の見出しが、割当の外の隣のページにある
  FAIL  目次に準拠した分類の見出しが、そろって同じ向きにずれている（系統的なずれ。目次の番号をファイルの番号として使った疑い）
  FAIL  LLM が分類したのに、根拠と「分類」の確認事項が無い
  FAIL  フローを描かず、下にフローも無い分類に、対象外の理由（out_of_scope）が無い／対象外の分類にページを割り当てている
  FAIL  目次に無く足した分類を、目次の項目の間に置いている（足した分類ごとの確認事項は、機械が作る）
  FAIL  名前の根拠（toc_ref・name_ref）の文字が、その場所に無い／原典の手がかりの無い上の段を作っている
  WARN  結び付き・関係・「続く」の根拠の文が、出典のページに見当たらない（原文でなく説明を書いている）
  FAIL  どの分類にも割り当てていない内容ページ（除外理由の記録が無いもの）／分類の pages と unassigned_pages の両方にあるページ
  FAIL  確認事項の id がファイルの中で重複している／目次に無い分類（basis: llm）があるのに全体の basis が toc
  FAIL  1つの分類の材料（カードの文字数）が上限を超える（1エージェントで読み切れない。下位の分類に分ける）
  FAIL  同じものとした用語に根拠が無い／同じ名前が2つの用語に入っている
[flow]
  FAIL  スキーマと違う・id の重複・存在しないレーンやノードを指す線／開始・終了が無い
  FAIL  出典の無いノード・存在しないページへの出典・引用が原典と合わない（前後のページにあれば、どのページかを名指しする）
  FAIL  出典が引くカードの項目（item）が無い・引用がその項目と合わない
  FAIL  カードの手順・分岐（補足の中の作業を含む）が、どのノードの出典にも引かれず、対象外の理由（excluded_items）も無い
  FAIL  開始から届かないノード・終了に届かないノード／出口が1本以下の分岐・ラベルの無い分岐・同じラベルの分岐
  FAIL  推定（inferred）なのに確認事項が無い／確認事項が存在しないノード・線を指す／ラベルに仮置きの語が残っている
  FAIL  確認事項の重要度が、決まり 5.4（答えで変わる所と、指す箱の数）と違う（renumber.py が付け直す）
  FAIL  カードの手順・分岐を出典に引いているノードが、推定（inferred）になっている（作業は原典にある）／引いていない作業のノードに、推定の印が無い
  FAIL  もとの手順に担当が書かれていない（カードの actor が空）のに、ノードを指す確認事項が無い（レーンを黙って決めている）
  FAIL  どのノードにも使われていない内容ページ（unused_pages に理由が無いもの）
  FAIL  レーン名・帳票名が用語の正本に無い／「続く」とした業務のフローに、続き先へのリンクが無い
  WARN  同じ手順が2つのフローに描かれている／帳票を受け渡しているのに「続く」が無い／リンクがあるのに「続く」が無い
  WARN  読めなかった・曖昧だったページなのに確認事項が無い／出典が、この分類に割り当てたページでない
  WARN  分岐でないノードから線が2本以上／ラベルが長い／1フロー 25 ステップ超／レーン 8 本超
  WARN  改善案らしい語（as-is に to-be を混ぜない）／採番の飛び（renumber.py で振り直す）
  WARN  引用が原典と完全には一致しない／文章の中にページ番号・ノードの番号がある／判断の出口が多い・線のラベルが長い
  WARN  直しの一覧の確認事項が、機械の確認事項（m-）と同じ所を指し、effect も同じ
  WARN  JSON をまるごと書いたフローで、システムのレーンが右端に無い（決まり 4.2）／使うシステムの札が付いた作業で、動作の種類が「作成する」（決まり 3.22 では入力する）
[code]
  FAIL  認証情報らしい文字が伏せられていない／構造（code/*.struct.json）が無い・いまのページと合わない／木の形の誤り（範囲・入れ子・実行文の抜けと重なり）
[systems]
  FAIL  システム台帳が無い（プログラムがあるのに）・形が違う／番号・名前の重なり／用語の正本に無いシステム・台帳に無いシステム／
        プログラムの資料が、1つのシステムに入っていない／紐付けの根拠の文字が、手順の文とプログラムの側の両方に無い／無い起動点・手順を指す／
        図にしない起動点の理由が中身と合わない（中身が空なのに実行文がある・表示だけなのに入出力がある・呼ぶだけなのにほかの文がある）
[activity]
  FAIL  図にする単位の図が無い／いまの構造・直しの一覧から作り直したものと違う／名前の付いていない箱／
        行の網羅（実行文が、どの箱にも対象外にも入っていない・2つに入っている）／分岐の網羅（判断にも畳みにもなっていない分岐・繰り返し）／
        かぎかっこの中の名前が、指したソースの行に無い／呼び出しの行き先の図が無い／確認事項の重要度が決まり 5.4 と違う
  WARN  箱が 25 個を超える／改善案らしい語／技術の用語（配列・状態コード など）／用語の対応で、定義の行にコメントがあるのに根拠が「つづり」
[drawio]
  FAIL  XML の破綻・id の重複・両端の無い線／テンプレートに無い色・形／flow.json と図の不一致（ノード・線・レーン）
  FAIL  出典の属性が無いノード／行き先の無いページリンク／フローのページが足りない
  FAIL  図・配置が、いまのデータから作り直したものと違う（作成日を除いて完全一致。手で直した・データを直したあと書き出していない）
[matrix]
  FAIL  業務マトリクス一覧が無い・読めない／いまのデータから組み立て直した中身と違うセルがある（フローを直したあと書き出し直していない・手で直した）
  FAIL  式のセルに保存した計算の結果が、式を計算し直した値と違う／フロー図の PNG へのリンクの先が無い・PNG があるのにリンクが無い
  WARN  計算の結果が保存されていない式のセル（計算をしない表示では 0・空に見える）／ブックに、まだ取り込んでいない人の値（書き出し直すと取り込む）
  FAIL  人が入れた時間の控え（times.json）を読めない／基準の表の版の記録が無い（前の版の道具で書き出したブック）
  FAIL  行がフローを描く分類の順でない／作業のノードがセルに無い・重なっている／行き先から組み直した線がフローと合わない／確認事項の番号が行に出ていない
[reports]
  FAIL  確認事項一覧・出典対応表・階層と業務分類・採否表が無い・いまのデータから作り直したものと違う／
        確認事項一覧.xlsx の機械が埋める欄（番号の並び・重要度・質問・仮置き・出典 ほか）が、いまのデータから作った行と違う
[deliverables]
  FAIL  納品物（図・業務マトリクス一覧・確認事項一覧・出典対応表・階層と業務分類）に、資料 ID（ページ・「」と並ぶ形）・内部の番号（業務の木の id・
        システムの番号・カードの項目のキー・ファイルの番号）・ファイルの場所（資料の表の「ファイル」の欄は除く）・スキーマの項目名・【下書き】が出ている
[review]
  FAIL  採否表の無い図（業務フロー・アクティビティ図）・図の無い採否表／形が違う・採否か対応の書かれていない指摘／判定・指摘が照合の報告と違う／
        照合したあとに変わった・足した箱（報告が実際に判定した版（review/sheets.json の控え）のどれとも、いまの図が違う。差分照合で照合し直す）／
        報告の sheet が無い・記録に無い・照合に出していない箱を判定した・名前が決まりと違う・形が崩れている・読めない（その図だけ止めて、次の図へ）／
        採用した指摘があるのに、図のデータが、その指摘を受けたシートのときのまま／
        支持率（支持 ÷ 開始・終了を除く全ノード）が 95% 未満（業務フロー・アクティビティ図それぞれ）
        sheet の無い前の版の報告だけの採否表は、前のとおり: 判定したノードが図に無い（採用した指摘が無いのに）／採用した指摘があるのに、図のデータが採否表より古い（ファイルの日時）
  WARN  保留の指摘／照合に出した版の記録が無い（前の版の採否表）
最後に 網羅率（ページ・項目）・補足由来の手順の反映数・引用照合率・推定の数・確認事項の件数を出す。
終了コード: FAIL があれば 1。
"""
import argparse
import collections
import html
import json
import pathlib
import re
import xml.etree.ElementTree as ET

from common import order_of  # noqa: E402
from common import (ASSETS, BadJson, MATERIAL_CAP, MAX_LANES, MAX_STEPS, NODE_TYPES, QUOTE_OK, SHIFT_MIN, SHIFT_SHARE, NON_CONTENT_TYPES, SOURCED_TYPES, STEP_TYPES, WORK_TYPES, CardItems, Report,
                    all_pages, card_items, card_json, flow_key, guarded, is_code_page, labels, load_card, load_flows, load_page, manifest, node_systems, norm, page_json,
                    page_text, page_units, q_targets, quote_match, same_text, severity_of, span_text, style_dict, system_names, tree_nodes, try_load, work_dir, flow_heads, ends_as_question, ends_with_verb, Times, seen_files, time_basis, actions_of, soft_formula, ACTIONS, DELIVERABLES, REVIEW_MIN, VOLUME_SPARE, Answers, answer_gaps, save, FLOW_FIG, ACT_FIG, MATRIX_BOOK, split_figs, ENDS_ASKED, LABEL_LONG)
from validate import validate

TECH_WORD = re.compile(r"配列|可変長|状態コード|ステータスコード|パラメータ|パラメーター|引数|戻り値|変数|フラグ|インデックス|レコードセット|オブジェクト|URL|例外")   # 箱に書かない技術の用語
PLACEHOLDER = re.compile(r"(TBD|TODO|XXX|？？|\?\?|◯◯|○○|〇〇|不明|未定|要確認)")
TOBE = re.compile(r"(すべき|するべき|べきである|望ましい|推奨|改善|効率化|自動化を検討|見直す|廃止する予定|導入予定)")
COLOR_KEYS = ("fillColor", "strokeColor", "fontColor", "labelBackgroundColor")
NEAR = 2                                                             # 引用を探す前後のページ数
HEADING_NEAR = 3                                                     # 分類の見出しを探す前後のページ数
SHAPE_ERRORS = 10                                                    # フロー1本につき、形の違反を何件まで出すか
STRUCT_ERRORS = 12                                                   # 階層・用語・業務分類のファイル1つにつき、同じく
MIN_QUOTE = 2                                                        # 引用は、記号と空白を除いてこの字数以上（「。」だけの引用は、どのページにも当たる）
ITEM_WORD = {"step": "手順", "decision": "分岐", "box": "図の箱"}


def shape_sig(st):
    """形の種類（スタイルの先頭の語と shape）"""
    lead = [k for k, v in st.items() if v == ""]
    return (lead[0] if lead else ("edge" if "edgeStyle" in st or "endArrow" in st else "rect"), st.get("shape", ""))


def card_text(card):
    """カードの中で、画像から読んだ文字として照合に使えるもの"""
    out = []
    for f in card.get("figures") or []:
        out += [str(t) for t in f.get("texts") or []] + [str(f.get("description", ""))]
    for s in (card.get("steps") or []) + (card.get("decisions") or []):
        out.append(str(s.get("quote", "")))
    for n in (card.get("diagram") or {}).get("nodes") or []:
        out.append(str(n.get("label", "")))
    out.append(str(card.get("title", "")))
    return "\n".join(out)


def file_ref(k):
    """ファイルの場所での書き方（M01/p0014）。存在しないページを指しているときにも使える（番号が数でなければ、書かれたまま出す）"""
    try:
        return f"{k[0]}/p{int(k[1] or 0):04d}"
    except (TypeError, ValueError):
        return f"{k[0]}/「{k[1]}」"


def shape_fails(rp, name, errs, limit, hint):
    """形の違反を limit 件まで出し、残りは件数と、全部を見る方法を1行で出す"""
    for e in errs[:limit]:
        rp.fail(f"{name}: 形が違う — {e}")
    if len(errs) > limit:
        rp.fail(f"{name}: 形が違う — ほか {len(errs) - limit} 件（{hint} で全部出る）")


def read_checked(rp, path, name, missing):
    """点検する JSON を読む。無い・JSON として読めない・全体がオブジェクトでないときは、FAIL にして None を返す"""
    d, err = try_load(path)
    if err is not None:
        rp.fail(f"{name}: JSON として読めない（{err}）。書きかけ・閉じ忘れを直す")
    elif d is None:
        rp.fail(missing)
    elif not isinstance(d, dict):
        rp.fail(f"{name}: 全体がオブジェクト（{{ … }}）になっていない")
        d = None
    return d


def dupes(xs):
    xs = [x for x in xs if isinstance(x, (str, int))]
    return sorted({x for x in xs if xs.count(x) > 1}, key=str)


STALE_REF = re.compile(r"M\d+\s*(?:/\s*)?p\.?\s?\d+|\bp\d{4}\b|\b(?:P|D|S|F|SUB)-\d{3}\b|(?<![A-Za-z])q\d+\b")
DOC_REF = re.compile(r"(?<![A-Za-z0-9])M\d{2,}(?![0-9])")             # 資料ID（M01）。機械が読み込んだ順に振った番号で、納品物には出さない（資料は題名で書く）
NUMBERS = {}                                                         # 納品の数（--numbers のファイルに書く。deliver.py が、納品の数.md・.json にする）
DETAIL_MAX = 120                                                     # ノードの補足（detail）の字数。超えると箱が縦に伸び、図が読めなくなる
LABEL_MAX, EXITS_MAX = 12, 4                                         # 判断から出る線のラベルの長さ・出口の数の目安


def lineage(cats, flow_ids, x):
    """フロー x と、その上位のフロー（フローを作る分類だけ）。近い順"""
    out, seen = [], set()
    while x in cats and x not in seen:
        seen.add(x)
        if x in flow_ids:
            out.append(x)
        x = cats[x].get("parent")
    return out


def tied(cats, flow_ids, a, b, pairs):
    """フロー a・b か、その上位のフローどうしが pairs にあるか（上位フローの「続く」・リンクは、下位フローの分も表す）"""
    return any((x, y) in pairs for x in lineage(cats, flow_ids, a) for y in lineage(cats, flow_ids, b))


def stale_refs(rp, name, texts, more=()):
    """文章で書く項目の中の、ページ番号・ノードの番号・確認事項の id。機械が引き直し・振り直すので、文の中の番号は古くなる。
    あわせて、資料を ID（M01）で呼んでいる所（more の文も見る。箱の名前・開始条件など。どれも、納品物にそのまま出る）"""
    hit = sorted({m.group(0) for t in texts if isinstance(t, str) for m in STALE_REF.finditer(t)})
    if hit:
        rp.warn(f"{name}: 文章の中に番号がある（{'・'.join(hit[:6])}）。ページは sources、対象は node・edge に書く。"
                f"文の中の番号は、機械が引き直し・振り直したあとも古いまま納品物に出る（flow-rules.md 5.7）")
    ids = sorted({m.group(0) for t in list(texts) + list(more) if isinstance(t, str) for m in DOC_REF.finditer(STALE_REF.sub(" ", t))})
    if ids:
        rp.warn(f"{name}: 文章の中で、資料を ID で呼んでいる（{'・'.join(ids[:6])}）。資料ID は機械が振った番号で、担当者には通じない。"
                f"資料は題名で書く（hierarchy.json の docs[].title。flow-rules.md 5.7）")


def evidence_pages(cx, rp, where, x):
    """根拠の文（evidence）があるなら、出典のページ（sources）が要り、そのページが実在すること（無いと、根拠を確かめようがない）"""
    if blank(x.get("evidence")):
        return
    srcs = expand(x.get("sources"))
    if not srcs:
        rp.fail(f"{where}: 根拠の文（evidence）があるのに、出典のページ（sources）が無い（その文が書いてある資料とページを書く）")
    for k in sorted(srcs - set(cx.pages), key=str):
        rp.fail(f"{where}: 根拠の出典が、存在しないページを指す: {file_ref(k)}")


def question_sources(cx, rp, name, qs):
    """確認事項の出典のページが実在するか（フロー以外のファイルの確認事項も）"""
    for q in qs or []:
        if isinstance(q, dict):
            for k in sorted(expand(q.get("sources")) - set(cx.pages), key=str):
                rp.fail(f"{name}: 確認事項 {q.get('id')} の出典が、存在しないページを指す: {file_ref(k)}")


def evidence_found(cx, text, sources):
    """根拠の文が、出典のページ（機械抽出＋カード）にあるか。出典が無ければ確かめない"""
    pages = [k for k in expand(sources) if k in cx.pages]
    return not pages or any(quote_match(text, cx.heading_text(k)) >= QUOTE_OK for k in pages)


def safe_items(card):
    """カードの項目（card_items）。カードの形が崩れていれば空"""
    try:
        return card_items(card)
    except (AttributeError, TypeError):
        return {}


def blank(v):
    return not str(v or "").strip()


def listed(keys, ref, limit=30):
    return ", ".join(ref(k) for k in sorted(keys)[:limit]) + (" …" if len(keys) > limit else "")


def expand(page_sets):
    """[{doc_id, pages: [..]}]（または {doc_id, page}）を (資料, page) の集合に。形が崩れた要素は飛ばす（形の違反は validate が言う）"""
    out = set()
    for ps in page_sets if isinstance(page_sets, list) else []:
        if not isinstance(ps, dict) or not isinstance(ps.get("doc_id"), str):
            continue
        pages = ps.get("pages") if isinstance(ps.get("pages"), list) else [ps["page"]] if "page" in ps else []
        out |= {(ps["doc_id"], p) for p in pages if isinstance(p, (int, str))}
    return out


class Ctx:
    def __init__(self, work):
        self.work = work
        self.m = manifest(work)
        self.lab = labels(self.m)
        self.pages = {(d["doc_id"], p["page"]): p for d, p in all_pages(self.m, code=True)}   # 出典に引けるページ（プログラムのソースを含む）
        self.docs = {d["doc_id"] for d in self.m["docs"]}
        self.active = {d for d, _p in self.pages}                    # 読む対象の資料（重複でなく、ページ化できたもの）
        tool = {d["doc_id"] for d in self.m["docs"] if d.get("tool")}   # 道具のブック（シートは参考。手順書として読まない）
        self.code_only = {d for d in self.active if d in tool or all(is_code_page(p) for (x, _n), p in self.pages.items() if x == d)}   # ソースだけの資料（業務の木でなく、システム台帳に結び付く）
        self.seq = {}                                                # 資料 → ページの並び
        for d, p in self.pages:
            self.seq.setdefault(d, []).append(p)
        self._pj, self._card = {}, {}

    def ref(self, k):
        """人に見せる書き方＋ファイルの場所（M01 p.13〔M01/p0014〕）"""
        return self.lab.ref(k[0], k[1]) if k in self.pages else file_ref(k)

    def pj(self, doc, page):
        k = (doc, page)
        if k not in self._pj:
            self._pj[k] = load_page(page_json(self.work, doc, page))
        return self._pj[k]

    def card(self, doc, page):
        k = (doc, page)
        if k not in self._card:                                      # 無ければ None。壊れていれば {}（check_extract.py が FAIL にする）
            self._card[k] = load_card(self.work, doc, page) if card_json(self.work, doc, page).exists() else None
        return self._card[k]

    def card_chars(self, doc, page):
        p = card_json(self.work, doc, page)
        return len(p.read_text(encoding="utf-8")) if p.exists() else 0

    def heading_text(self, k):
        """そのページの文字（機械抽出＋カードの見出し・画像から読んだ文字）。見出しがあるかを見るのに使う"""
        c = self.card(*k) or {}
        return norm(page_text(self.pj(*k))) + norm(card_text(c)) + norm(" ".join(map(str, c.get("section_path") or [])))

    def content_pages(self):
        """網羅チェックの対象（表紙・目次・白紙以外）。カードが無いページは対象に入れる（check_extract が FAIL にする）。
        種別が表紙・目次・白紙でも、手順・分岐・図の箱が書かれていれば対象に入れる（種別より、書かれた中身を信じる）"""
        out = []
        for k in self.pages:
            if is_code_page(self.pages[k]):                          # プログラムのソースは、カード・業務分類の対象でない（アクティビティ図の行の網羅で見る）
                continue
            c = self.card(*k)
            if c is None or c.get("page_type") not in NON_CONTENT_TYPES or safe_items(c):
                out.append(k)
        return out

    def neighbors(self, k, n):
        """同じ資料の前後 n ページを、近い順に [(ページ, 何ページ離れているか)]"""
        seq = self.seq.get(k[0], [])
        if k[1] not in seq:
            return []
        at = seq.index(k[1])
        near = sorted(range(max(0, at - n), min(len(seq), at + n + 1)), key=lambda j: abs(j - at))
        return [((k[0], seq[j]), j - at) for j in near if j != at]


# ---------- 階層・用語・業務分類 ----------

def glossary_names(cx):
    """用語の正本: {正規化した名前: (種類, 正本の名前)}。別の表記（aliases）は alias に入れる"""
    g = try_load(cx.work / "glossary.json")[0]
    names, alias = {}, {}
    for t in (g.get("terms") if isinstance(g, dict) else None) or []:
        if not isinstance(t, dict):
            continue
        names[norm(t.get("name", ""))] = (t.get("kind"), t.get("name"))
        for x in t.get("aliases") if isinstance(t.get("aliases"), list) else []:
            alias[norm(x)] = (t.get("kind"), t.get("name"))
    return (g if isinstance(g, dict) else None), names, alias


def check_hierarchy(cx, rp, hier, final):
    """階層の判定の中身。読取の前には決めきれないものは、final でなければ（階層だけの段階では）WARN"""
    seen, tied, detailed = set(), set(), set()
    toc = hier.get("toc_flow_list") if isinstance(hier.get("toc_flow_list"), dict) else {}
    axis = {x.get("doc_id") for x in toc.get("source") or [] if isinstance(x, dict)} if toc.get("found") else set()   # 縦軸にする目次のある資料
    if toc.get("found") and (not axis or not toc.get("entries")):  # 目次との突き合わせ（分類の割当・名前・順）が、黙って飛ぶ
        rp.fail("hierarchy.json: 目次に業務の一覧がある（toc_flow_list.found）としたのに、目次の資料（source）か項目（entries）が空。目次のページと項目を書く")
    for n, _parent in tree_nodes(hier.get("tree")):
        if n.get("id") in seen:
            rp.fail(f"hierarchy.json: 階層ノードの id が重複: {n.get('id')}")
        seen.add(n.get("id"))
        for r in n.get("refs") or []:
            if r.get("doc_id") not in cx.docs:
                rp.fail(f"hierarchy.json: {n.get('id')} が存在しない資料を指す: {r.get('doc_id')}")
                continue
            tied.add(r["doc_id"])
            for pg in r.get("pages") or []:
                if (r["doc_id"], pg) not in cx.pages:
                    rp.fail(f"hierarchy.json: {n.get('id')} が存在しないページを指す: {file_ref((r['doc_id'], pg))}")
            if r.get("rel") == "details":
                detailed.add(r["doc_id"])
                if blank(r.get("evidence")):
                    rp.fail(f"hierarchy.json: {n.get('id')}「{n.get('title')}」に {r['doc_id']} を「詳しくしている」と結び付けたのに、"
                            f"根拠（evidence）が無い（根拠の文とページを書く。無ければ結び付けない）")
            elif axis and r["doc_id"] not in axis and n.get("source") == "toc" and blank(r.get("evidence")):
                rp.fail_if(final, f"hierarchy.json: {n.get('id')}「{n.get('title')}」に、目次の資料でない {r['doc_id']} を結び付けたのに、根拠（evidence）が無い"
                           f"（根拠の文とページを書く。書けなければ結び付けず、目次に無い業務として分類に足す。flow-rules.md 2.20）")
            evidence_pages(cx, rp, f"hierarchy.json: {n.get('id')} と {r['doc_id']} の結び付き", r)
            if not blank(r.get("evidence")) and not evidence_found(cx, r["evidence"], r.get("sources")):
                rp.warn(f"hierarchy.json: {n.get('id')} と {r['doc_id']} の結び付きの根拠「{str(r['evidence'])[:30]}」が、出典のページに見当たらない"
                        f"（根拠は原文のまま写す。説明の文を書かない）")
    if hier.get("is_hierarchical") in ("yes", "partial") and not hier.get("tree"):
        rp.fail("hierarchy.json: 階層ありと判定したのに tree が空")
    roles = {x.get("doc_id"): x for x in hier.get("docs") or [] if isinstance(x, dict)}
    for d in sorted(cx.active - cx.code_only):
        if d not in roles:
            rp.fail(f"hierarchy.json: 資料 {d} の役割（docs[].role）が書かれていない")
        elif d not in tied and blank(roles[d].get("out_of_scope")):
            rp.fail_if(final, f"hierarchy.json: 資料 {d} がどの業務にも結び付いていない（tree の refs に入れるか、docs の out_of_scope に対象外の理由を書く）")
    for d in roles:
        if d not in cx.docs:
            rp.fail(f"hierarchy.json: docs が存在しない資料を指す: {d}")
    for r in hier.get("doc_relations") or []:
        if r.get("from") not in cx.docs or r.get("to") not in cx.docs:
            rp.fail(f"hierarchy.json: 資料間の関係が存在しない資料を指す: {r.get('from')}→{r.get('to')}")
        elif r.get("type") == "parent_of" and r["to"] not in detailed:
            rp.fail_if(final, f"hierarchy.json: {r['from']} が {r['to']} の上位、としたのに、{r['to']} が業務の木のどの節を詳しくしているか"
                       f"（tree の refs に rel: details）が書かれていない。資料どうしの親子だけでは、フローの親子に落とせない")
        evidence_pages(cx, rp, f"hierarchy.json: {r.get('from')}→{r.get('to')} の関係", r)
        if not blank(r.get("evidence")) and not evidence_found(cx, r["evidence"], r.get("sources")):
            rp.warn(f"hierarchy.json: {r.get('from')}→{r.get('to')} の関係の根拠「{str(r['evidence'])[:30]}」が、出典のページに見当たらない"
                    f"（根拠は原文のまま写す。説明の文を書かない）")
    for q in hier.get("sequence") or []:
        for end in ("from", "to"):
            if q.get(end) not in seen:
                rp.fail(f"hierarchy.json: 「続く」が業務の木に無い id を指す: {q.get(end)}")
        evidence_pages(cx, rp, f"hierarchy.json: 「続く」{q.get('from')}→{q.get('to')}", q)
        if not blank(q.get("evidence")) and not evidence_found(cx, q["evidence"], q.get("sources")):
            rp.warn(f"hierarchy.json: 「続く」{q.get('from')}→{q.get('to')} の根拠「{str(q['evidence'])[:30]}」が、出典のページに見当たらない"
                    f"（根拠は原文のまま写す）")
    asked = [x.get(k) for x in hier.get("questions") or [] if isinstance(x, dict) for k in ("question", "assumption")]
    stale_refs(rp, "hierarchy.json", [hier.get("rationale")] + asked)


def check_glossary(rp, g):
    shape_fails(rp, "glossary.json", validate(g, "glossary"), STRUCT_ERRORS, "python3 scripts/validate.py -w <作業フォルダ> --kind glossary")
    used = {}                                                        # 正規化した名前 → (何件目の用語か, その用語の名前)
    for i, t in enumerate(g.get("terms") or []):
        if not isinstance(t, dict):
            continue
        aliases = [x for x in t.get("aliases") or [] if isinstance(x, str)] if isinstance(t.get("aliases"), list) else []
        if aliases and blank(t.get("evidence")):
            rp.fail(f"glossary.json: 「{t.get('name')}」に別の表記（{'・'.join(aliases)}）を寄せたのに、同じものとした根拠（evidence）が無い。"
                    f"確かめられなければ寄せずに別の用語にし、確認事項にする")
        for x in [t.get("name")] + aliases:
            k = norm(x)
            if k and k in used and used[k][0] != i:                  # 同じ名前の用語を2件並べたものも、ここで止める
                rp.fail(f"glossary.json: 「{x}」が2つの用語（{used[k][1]}・{t.get('name')}）に入っている")
            used.setdefault(k, (i, t.get("name")))
    check_question_ids(rp, "glossary.json", g.get("questions"))


# 答えで図が変わらない問いの言い回し（言い回しでの見分けなので、WARN）
Q_STILL = re.compile(r"(いま|今|現在|現状)(で)?も[^。？?]{0,40}(同じ|変わ|まま|ですか|でしょうか|よいか)")
Q_VALUE = re.compile(r"(いくら|何円|何万円|何日|何時|何件|何回|何％|何パーセント|何か月|何ヶ月|何年|いつまで)[^。？?、]{0,12}(ですか|ますか|でしょうか|か[。？?]?$)")


def question_wording(rp, name, qs):
    """確認事項は、答えで図（箱・線・分岐・担当・システムや帳票・フローの範囲）が変わるものだけ。値だけを聞く問い・いまも同じかの確かめを、言い回しから拾う"""
    for q in qs or []:
        text = str(q.get("question", "")) if isinstance(q, dict) else ""
        if Q_STILL.search(text):
            rp.warn(f"{name}: 確認事項 {q.get('id')}「{text[:30]}」は、書かれていることが、いまも同じかを確かめるだけの問いに見える。"
                    "答えで図が変わらないなら、確認事項にしない（資料どうしで、箱・順・担当が食い違うなら、その食い違いを聞く）")
        elif Q_VALUE.search(text):
            rp.warn(f"{name}: 確認事項 {q.get('id')}「{text[:30]}」は、値（金額・日数・件数・期限）だけを聞く問いに見える。"
                    "答えで図が変わらないなら、確認事項にしない（値は、箱の補足に書く。資料どうしで値が違えば、両方を補足に書く）")


def check_question_ids(rp, name, qs):
    """確認事項の id は、そのファイルの中で重ならないこと（通し番号 Q-001… の引き継ぎに使う。重なると、別の2件に同じ番号が付く）"""
    for i in dupes([q.get("id") for q in qs or [] if isinstance(q, dict) and isinstance(q.get("id"), str)]):
        rp.fail(f"{name}: 確認事項の id が重複: {i}")
    question_wording(rp, name, qs)


def check_label_form(rp, where, n):
    """箱の名前は、動詞で終える（「何を・どうする」）。判断は、問いの形（「〜か」。繰り返しの判断は、動詞で終えてもよい）。
    最後に目的語・補足を置くなら、かっこ書きにする（かっこ書きは、除いて見る）。開始・終了と帳票（物の名前）は、対象にしない"""
    t, lab = n.get("type"), str(n.get("label", "")).strip()
    if not lab:
        return
    if t in ("process", "system", "subprocess") and not ends_with_verb(lab):
        rp.fail(f"{where}「{lab[:24]}」: 箱の名前が、動詞で終わっていない（「何を・どうする」で終える。最後に置く目的語・補足は、かっこ書きにするか、detail に書く）")
    elif t == "decision" and not (ends_as_question(lab) or ends_with_verb(lab)):
        rp.fail(f"{where}「{lab[:24]}」: 判断の名前が、「〜か」で終わっていない（問いの形にする。繰り返しの判断は、動詞で終えてもよい）")


def check_category(cx, rp, c, by_id, hier_ids, cap):
    """分類1つ: 親子・業務の木との対応・ページの割当・目次の名称・材料の大きさ・見出しの位置。
    (割り当てたページ, 見出しのずれ) を返す"""
    cid, name = c.get("id"), c.get("name")
    if c.get("parent") and c["parent"] not in by_id:
        rp.fail(f"taxonomy.json: {cid} の parent「{c['parent']}」が無い")
    up, hops = c.get("parent"), 0
    while up in by_id and hops <= len(by_id):
        up, hops = by_id[up].get("parent"), hops + 1
    if hops > len(by_id):
        rp.fail(f"taxonomy.json: {cid} の親子（parent）が循環している")
    want = (by_id[c["parent"]].get("level") + 1) if c.get("parent") in by_id and isinstance(by_id[c["parent"]].get("level"), int) else None if c.get("parent") else 1
    if want is not None and c.get("level") != want:
        rp.fail(f"taxonomy.json: {cid} の段（level {c.get('level')}）が、親から数えた段（{want}）と違う（根は 1、子は親の段＋1）")
    if c.get("hierarchy_id") and hier_ids and c["hierarchy_id"] not in hier_ids:
        rp.fail(f"taxonomy.json: {cid} の hierarchy_id「{c['hierarchy_id']}」が業務の木に無い")
    pages = expand(c.get("pages"))
    for k in sorted(pages, key=str):
        if k not in cx.pages:
            rp.fail(f"taxonomy.json: {cid} が存在しないページを指す: {file_ref(k)}")
    real = sorted(k for k in pages if k in cx.pages)
    if c.get("basis") == "toc":
        ref = c.get("toc_ref") or {}
        k = (ref.get("doc_id"), ref.get("page"))
        if k not in cx.pages:
            rp.fail(f"taxonomy.json: {cid}「{name}」は目次準拠なのに、目次の場所（toc_ref）が無い・存在しない")
        else:
            toc_card = cx.card(*k) or {}
            src = norm(page_text(cx.pj(*k))) + norm(card_text(toc_card)) + norm(
                " ".join(str(e.get("title", "")) for e in toc_card.get("toc_entries") or []))
            if norm(c.get("name", "")) not in src:
                rp.fail(f"taxonomy.json: {cid}「{name}」が目次（{cx.ref(k)}）の文字に見当たらない。"
                        f"目次準拠の分類は名称を目次のままにする")
            elif not blank(ref.get("text")) and norm(ref["text"]) not in src and quote_match(ref["text"], src) < QUOTE_OK:
                rp.fail(f"taxonomy.json: {cid}「{name}」の toc_ref の文字「{str(ref['text'])[:30]}」が、目次（{cx.ref(k)}）に見当たらない"
                        f"（目次の文字のまま写す）")
    check_name_ref(cx, rp, c)
    if c.get("make_flow"):                                           # 材料の大きさ（1エージェントが読み切れるか）
        size = sum(cx.card_chars(*k) for k in real)
        if size > cap:
            rp.fail(f"taxonomy.json: {cid}「{name}」の材料が大きすぎる（カード {size:,} 字。上限 {cap:,} 字）。"
                    f"1つのエージェントでは読み切れない。原典の節・項の切れ目で下位の分類（{cid}-1…）に分け、"
                    f"親のフローには下位フローへのリンクを置く（pages.py plan-flows で大きさを見る）")
    shift = None
    if c.get("basis") == "toc" and c.get("pages"):
        shift = check_heading_page(cx, rp, c, [k for k in real if k[0] == c["pages"][0].get("doc_id")])
    return pages, shift


def check_name_ref(cx, rp, c):
    """LLM が分類した分類の、名前の根拠（name_ref）: 書いた文字が、その場所（見出しのページ・資料名）に本当にあるか"""
    ref, cid, name = c.get("name_ref"), c.get("id"), c.get("name")
    if not isinstance(ref, dict):
        return
    if c.get("basis") == "toc":
        rp.fail(f"taxonomy.json: {cid}「{name}」は目次準拠なのに name_ref がある（目次の場所は toc_ref に書く）")
        return
    doc, text = ref.get("doc_id"), str(ref.get("text") or "")
    if doc not in cx.docs:
        rp.fail(f"taxonomy.json: {cid} の name_ref が存在しない資料を指す: {doc}")
        return
    k = (doc, ref.get("page"))
    if ref.get("page") is not None and k not in cx.pages:
        rp.fail(f"taxonomy.json: {cid} の name_ref が存在しないページを指す: {file_ref(k)}")
        return
    if ref.get("kind") == "heading":
        if ref.get("page") is None:
            rp.fail(f"taxonomy.json: {cid}「{name}」の name_ref（本文の見出し）にページが無い")
            return
        found = norm(text) in cx.heading_text(k)
        where = cx.ref(k)
    else:                                                            # 資料の題名・ファイル名・フォルダ名
        d = next((x for x in cx.m["docs"] if x["doc_id"] == doc), {})
        names = norm(" ".join([str(d.get("rel_path", "")), str(d.get("name", ""))] + [str(x) for x in d.get("folder") or []]))
        found = norm(text) in names or (ref.get("page") is not None and norm(text) in cx.heading_text(k))
        where = f"{doc} の資料名・フォルダ名" + (f"・{cx.ref(k)}" if ref.get("page") is not None else "")
    if not found:
        rp.fail(f"taxonomy.json: {cid}「{name}」の名前の根拠「{text[:30]}」が、{where} に見当たらない（原典の文字のまま写す。"
                f"原典に名前が無ければ name_ref を書かず、その分類を指す確認事項にする）")
    elif norm(name or "") not in norm(text):
        rp.fail(f"taxonomy.json: {cid}「{name}」が、名前の根拠「{text[:30]}」に含まれていない（原典の見出し・資料名の言葉のまま名付ける）")


def check_axis(rp, tax, cats, by_id):
    """縦軸（業務マトリクスの行の並び）の決まり: フローを描かない分類の扱い、目次に無く足した分類の置き場所と確認事項"""
    kids = {}
    for c in cats:
        kids.setdefault(c.get("parent"), []).append(c)

    def draws_below(c, seen=()):
        return any(x.get("make_flow") or (x.get("id") not in seen and draws_below(x, seen + (c.get("id"),))) for x in kids.get(c.get("id"), []))
    asked = {}                                                       # 分類 id → その分類を指す、種別「分類」の確認事項の数
    for q in tax.get("questions") or []:
        if isinstance(q, dict) and q.get("flow") is not None:
            if q["flow"] not in by_id:
                rp.fail(f"taxonomy.json: 確認事項 {q.get('id')} の flow「{q['flow']}」が、分類に無い")
            elif "分類" in str(q.get("category", "")):
                asked[q["flow"]] = asked.get(q["flow"], 0) + 1
    with_toc = tax.get("basis") in ("toc", "mixed")
    for c in cats:
        cid, name, out = c.get("id"), c.get("name"), not blank(c.get("out_of_scope"))
        if c.get("make_flow"):
            if out:
                rp.fail(f"taxonomy.json: {cid}「{name}」はフローを描く分類なのに、対象外の理由（out_of_scope）がある（どちらかにする）")
        elif draws_below(c):
            if out:
                rp.fail(f"taxonomy.json: {cid}「{name}」は下にフローを描く分類がある（束ねる見出し）のに、対象外の理由（out_of_scope）がある")
            if c.get("basis") == "llm" and not isinstance(c.get("name_ref"), dict):
                rp.fail(f"taxonomy.json: {cid}「{name}」— 束ねる段（上の分類）を作ったのに、原典の手がかり（name_ref: フォルダ名・資料名・章の見出し）が無い。"
                        f"原典に無い上の段は作らない（1段のままにする）")
        else:
            if not out:
                rp.fail(f"taxonomy.json: {cid}「{name}」はフローを描かず、下にフローを描く分類も無い。業務でない項目なら対象外の理由（out_of_scope）を書く。"
                        f"業務なら make_flow を true にする")
            if expand(c.get("pages")):
                rp.fail(f"taxonomy.json: {cid}「{name}」は対象外なのに、ページが割り当てられている（pages を空にし、unassigned_pages に理由を書く）")
        if c.get("basis") != "llm":
            continue
        named = isinstance(c.get("name_ref"), dict)
        if with_toc:                                                 # 目次に無く足した分類: 兄弟の末尾に置き、1件ずつ確かめる
            later = [x for x in kids.get(c.get("parent"), []) if x.get("basis") == "toc" and cats.index(x) > cats.index(c)]
            if later:
                rp.fail(f"taxonomy.json: {cid}「{name}」（目次に無く足した分類）が、目次の項目 {later[0].get('id')}「{later[0].get('name')}」より前にある。"
                        f"足した分類は、同じ段の目次の項目の末尾に置く（目次の並びを崩さない）")
        elif c.get("make_flow") and not named and not asked.get(cid) and not any(
                isinstance(q, dict) and "分類" in str(q.get("category", "")) and q.get("flow") is None for q in tax.get("questions") or []):
            rp.fail(f"taxonomy.json: {cid}「{name}」は原典に名前が無い分類（name_ref なし）なのに、分類を確かめる確認事項が無い")


def check_tree_items(rp, hier, cats):
    """目次から採った業務の木の項目が、どれも扱いが決まっているか（フローを描く・束ねる見出し・上位のフローに含む・対象外）"""
    of_h = {}
    for c in cats:
        if c.get("hierarchy_id"):
            of_h.setdefault(c["hierarchy_id"], c)
    state = {}
    for n, parent in tree_nodes(hier.get("tree")):
        c, up = of_h.get(n.get("id")), (state.get(id(parent)) if parent is not None else None)
        if c is not None:
            state[id(n)] = "flow" if c.get("make_flow") else "out" if not blank(c.get("out_of_scope")) else "group"
        elif up in ("flow", "out", "unset"):
            state[id(n)] = up
        else:
            state[id(n)] = "unset"
            if n.get("source") == "toc":
                rp.fail(f"taxonomy.json: 目次の項目 {n.get('id')}「{n.get('title')}」の扱いが決まっていない（対応する分類が無く、上の項目もフローを描く分類・対象外の分類でない）。"
                        f"分類にして、フローを描くか（make_flow）、業務でなければ対象外の理由（out_of_scope）を書く")


def check_heading_page(cx, rp, c, mine):
    """見出しのあるページから割り当てているか（目次の番号をそのまま使うと、表紙を数えない資料では1ページずれる）。
    見出しが別のページにあれば、(資料, 何ページ先か) を返す"""
    name = norm(c.get("name", ""))
    if not mine or len(name) < 2 or name in cx.heading_text(mine[0]):
        return None
    near = cx.neighbors(mine[0], HEADING_NEAR)
    for k, off in [x for x in near if x[0] in mine] + [x for x in near if x[0] not in mine]:   # 割当の中を先に（本文の中の語に先に当たらないように）
        if (cx.card(*k) or {}).get("page_type") in ("toc", "cover"):
            continue
        if name in cx.heading_text(k):
            msg = (f"taxonomy.json: {c.get('id')}「{c.get('name')}」の見出しが、割当の先頭（{cx.ref(mine[0])}）に無く、"
                   f"{abs(off)} ページ{'後ろ' if off > 0 else '前'}の {cx.ref(k)} にある（ページの割当が {abs(off)} ページずれている疑い）")
            rp.fail_if(k not in mine, msg)
            return mine[0][0], off
    return None


def toc_page(cx, doc, printed):
    """目次に刷られた番号 → ファイルの番号。目次の番号は「刷られた番号」で、データの鍵とは違うことがある。機械が引いて確かめる"""
    if cx.lab.docs.get(doc, {}).get("basis") == "printed":
        return cx.lab.find(doc, printed)
    return int(printed) if str(printed).isdigit() and (doc, int(printed)) in cx.pages else None


def toc_category(e, cats):
    """目次の項目に当たる分類。番号と名前が合うもの → 名前が同じもの → 名前が含み合うもの、の順に探す。
    章と節が同じ名前のとき（第5章 出張手続／1 出張手続）は、ページを持つほうを採る"""
    t, num = norm(e.get("title", "")), norm(str(e.get("number") or ""))
    if not t:
        return None
    named = [(c, norm(c.get("name", ""))) for c in cats]
    for same in (lambda c, n: bool(num) and n == t and norm(str(c.get("number") or "")) == num,
                 lambda c, n: n == t,
                 lambda c, n: bool(n) and (t in n or n in t)):
        hit = [c for c, n in named if same(c, n)]
        if hit:
            return min(hit, key=lambda c: (not c.get("pages"), abs(len(norm(c.get("name", ""))) - len(t))))
    return None


def check_toc_list(cx, rp, toc, tax, cats):
    """目次に業務フローの一覧があるとき: 分類が目次の項目・ページ・順に合っているか"""
    if tax.get("basis") == "llm":
        rp.fail("目次に業務フローの一覧があるのに、業務分類が目次に準拠していない（taxonomy.basis が llm）")
    entries = [e for e in toc.get("entries") or [] if isinstance(e, dict)]
    tdoc = next((s.get("doc_id") for s in toc.get("source") or [] if isinstance(s, dict)), None)
    bare = lambda t: re.sub(r"^(?:第?[\d.\-()]+[章節項編部]?)+", "", norm(t))        # 先頭の番号（「3.」「第5章」「(1)」）は比べない
    titles = {bare(e.get("title", "")) for e in entries} - {""}
    for c in cats:                                                   # 名称は目次の項目のまま。縮めた名前・言い換えは、目次の文字の一部でも通さない
        n = bare(c.get("name", ""))
        if c.get("basis") == "toc" and c.get("make_flow") and n and n not in titles and any(n in t or t in n for t in titles):
            near = next(e.get("title") for e in entries if bare(e.get("title", "")) and (n in bare(e.get("title", "")) or bare(e.get("title", "")) in n))
            rp.fail(f"taxonomy.json: {c.get('id')}「{c.get('name')}」が目次の項目と一致しない（目次は「{near}」）。目次準拠の分類は名称を目次のままにする")
    seq = []
    for e in entries:
        t = norm(e.get("title", ""))
        c = toc_category(e, cats)
        if t and c is None:
            rp.fail(f"目次の項目「{e.get('title')}」に対応する業務分類が無い")
            continue
        if c is not None:
            seq.append(cats.index(c))
        if not t or e.get("page") in (None, "") or tdoc is None:
            continue
        page = toc_page(cx, tdoc, e["page"])
        if page is None:
            rp.warn(f"目次の項目「{e.get('title')}」の番号 {e['page']} に当たるページが {tdoc} に見つからない（目次の番号を確かめる）")
            continue
        mine = sorted(k[1] for k in expand(c.get("pages")) if k[0] == tdoc and k in cx.pages)
        if mine and page not in mine:
            rp.fail(f"taxonomy.json: {c.get('id')}「{c.get('name')}」— 目次の番号 {e['page']} が指すのは {cx.ref((tdoc, page))} なのに、"
                    f"割当は {cx.ref((tdoc, mine[0]))} から始まっていて、そのページが入っていない（{abs(mine[0] - page)} ページずれ。"
                    f"目次の番号を、ファイルの番号としてそのまま使っていないか。pages.py pagemap で対応を引く）")
    if seq != sorted(seq):
        rp.warn("業務分類の並びが目次の順と違う")


def check_taxonomy(cx, rp, tax, hier, cap):
    """業務分類とページの割当。(分類の数, 割り当てたページの数, 除外したページの数) を返す"""
    shape_fails(rp, "taxonomy.json", validate(tax, "taxonomy"), STRUCT_ERRORS, "python3 scripts/validate.py -w <作業フォルダ> --kind taxonomy")
    cats = [c for c in (tax.get("categories") if isinstance(tax.get("categories"), list) else []) if isinstance(c, dict)]
    ids = [c.get("id") for c in cats]
    for i in dupes(ids):
        rp.fail(f"taxonomy.json: 分類の id が重複: {i}")
    by_id = {c.get("id"): c for c in cats}
    hier_ids = {n.get("id") for n, _p in tree_nodes(hier.get("tree"))}
    assigned, shifts = set(), collections.Counter()
    for c in cats:
        pages, shift = check_category(cx, rp, c, by_id, hier_ids, cap)
        assigned |= pages
        if shift:
            shifts[shift] += 1
    based = sum(1 for c in cats if c.get("basis") == "toc" and c.get("pages"))
    for (doc, off), n in shifts.items():                             # 1件ずつなら扉のページを含めただけかもしれない。そろってずれていれば番号の取り違え
        if n >= SHIFT_MIN and n >= SHIFT_SHARE * based:
            rp.fail(f"系統的なずれ — 目次に準拠した分類 {n} 件で、見出しが割当の先頭の {abs(off)} ページ{'後ろ' if off > 0 else '前'}にある（{doc}）。"
                    f"目次の番号（刷られた番号）を、ファイルの番号としてそのまま割当に使っていないか。pages.py pagemap で対応を引く", first=True)
    toc = hier.get("toc_flow_list") or {}
    if toc.get("found"):
        check_toc_list(cx, rp, toc, tax, cats)
    check_axis(rp, tax, cats, by_id)
    check_tree_items(rp, hier, cats)
    kinds = {c.get("basis") for c in cats}
    if tax.get("basis") == "toc" and "llm" in kinds:                 # 目次に無い分類を足したのに、全体を toc のままにした
        who = "・".join(str(c.get("id")) for c in cats if c.get("basis") == "llm")
        rp.fail(f"taxonomy.json: {who} の basis が llm なのに、全体の basis が toc（mixed にして、分類の考え方を書く。足した分類ごとの確認事項は、機械が作る）")
    if tax.get("basis") in ("llm", "mixed") or "llm" in kinds:
        if blank(tax.get("rationale")):
            rp.fail("taxonomy.json: LLM が分類したのに、分類の考え方（rationale）が空")
        if tax.get("basis") == "llm" and not any("分類" in str(q.get("category", "")) for q in tax.get("questions") or [] if isinstance(q, dict)):
            rp.fail("taxonomy.json: LLM が分類したのに、種別「分類」の確認事項（分類案と別案を利用者に確かめるもの）が無い"
                    "（目次に無い分類を足しただけ（mixed）なら、足した分類ごとの問いは機械が作る）")
    check_question_ids(rp, "taxonomy.json", tax.get("questions"))
    question_sources(cx, rp, "taxonomy.json", tax.get("questions"))
    asked = [x.get(k) for x in tax.get("questions") or [] if isinstance(x, dict) for k in ("question", "assumption")]
    stale_refs(rp, "taxonomy.json", [tax.get("rationale")] + asked, [c.get("name") for c in tax.get("categories") or [] if isinstance(c, dict)])
    skipped = {(u.get("doc_id"), u.get("page")): u.get("reason") for u in tax.get("unassigned_pages") or [] if isinstance(u, dict)}
    for k, why in sorted(skipped.items(), key=str):
        if blank(why):
            rp.fail(f"taxonomy.json: {file_ref(k)} を割り当てない理由が空")
        if k not in cx.pages:
            rp.fail(f"taxonomy.json: unassigned_pages が存在しないページを指す: {file_ref(k)}")
        elif k in assigned:
            rp.fail(f"taxonomy.json: {cx.ref(k)} が、分類の pages と unassigned_pages の両方にある（どちらかにする）")
    lost = [k for k in cx.content_pages() if k not in assigned and k not in skipped]
    if lost:
        rp.fail(f"どの業務分類にも割り当てていない内容ページ {len(lost)} 件: {listed(lost, cx.ref)}（分類に入れるか、unassigned_pages に理由を書く）")
    return len(cats), len(assigned), len(skipped)


def check_structure(cx, rp, hierarchy_only=False, cap=MATERIAL_CAP):
    """階層・用語・業務分類。ファイルごとに、形の違反を出してから中身を見る。形が崩れていて中身を見られないときも、ほかのファイルは見る"""
    hier = read_checked(rp, cx.work / "structure" / "hierarchy.json", "structure/hierarchy.json",
                        "structure/hierarchy.json が無い（階層の事前チェックの結果を書く。references/structure-guide.md）") or {}
    if hier:
        shape_fails(rp, "hierarchy.json", validate(hier, "hierarchy"), STRUCT_ERRORS, "python3 scripts/validate.py -w <作業フォルダ> --kind hierarchy")
        hier = {k: v for k, v in hier.items() if isinstance(v, (list, dict, str))}      # 形が違っても、読める所までは点検する
        with guarded(rp, "hierarchy.json"):
            check_hierarchy(cx, rp, hier, not hierarchy_only)
            check_question_ids(rp, "hierarchy.json", hier.get("questions"))
            question_sources(cx, rp, "hierarchy.json", hier.get("questions"))
    toc = hier.get("toc_flow_list") if isinstance(hier.get("toc_flow_list"), dict) else {}
    state = f"階層: {hier.get('is_hierarchical', '未判定')}"
    if hierarchy_only:
        rp.info(f"{state}／目次の業務フロー一覧: {'あり' if toc.get('found') else 'なし'}／"
                f"資料どうしの関係 {len(hier.get('doc_relations') or [])} 件")
        return
    g = read_checked(rp, cx.work / "glossary.json", "glossary.json",
                     "glossary.json が無い（担当・システム・帳票の名前の正本。pages.py terms で表記の一覧を出して書く。references/structure-guide.md）")
    if g is not None:
        with guarded(rp, "glossary.json"):
            check_glossary(rp, g)
            question_sources(cx, rp, "glossary.json", g.get("questions"))
    tax = read_checked(rp, cx.work / "taxonomy.json", "taxonomy.json",
                       "taxonomy.json が無い（業務分類とページの割当を書く。references/structure-guide.md）")
    if tax is None:
        return
    counts = ("?", "?", "?")
    with guarded(rp, "taxonomy.json"):
        counts = check_taxonomy(cx, rp, tax, dict(hier, toc_flow_list=toc), cap)
    rp.info(f"{state}／業務分類 {counts[0]}（根拠 {tax.get('basis')}）／割当 {counts[1]} ページ・除外 {counts[2]} ページ／"
            f"用語 {len((g or {}).get('terms') or [])}／業務どうしの「続く」 {len(hier.get('sequence') or [])} 件")


# ---------- フロー ----------

def placeholder(n):
    """ノードのラベルに残った仮置きの語。原典がその語で書いている（引用に「原因を調べても不明なとき」とある）なら、仮置きでない"""
    held = PLACEHOLDER.search(str(n.get("label", "")))
    quoted = norm(" ".join(str(s.get("quote", "")) for s in n.get("sources") or [] if isinstance(s, dict)))
    return held.group(0) if held and norm(held.group(0)) not in quoted else None


class FlowView:
    """フロー1本を、点検で引きやすい形にしたもの"""

    def __init__(self, fl):
        rows = lambda key: [x for x in (fl.get(key) if isinstance(fl.get(key), list) else []) if isinstance(x, dict)]
        named = lambda x: isinstance(x, str) and bool(x)
        self.raw = fl
        self.F = flow_key(fl)
        every = rows("nodes")
        self.nodes = [n for n in every if named(n.get("id"))]         # id の無いノードは点検できない（形の違反として出す）
        self.unnamed = len(every) - len(self.nodes)
        self.ids = [n["id"] for n in self.nodes]
        self.by_id = {n["id"]: n for n in self.nodes}
        self.main = [n for n in self.nodes if n.get("type") != "file"]          # 帳票以外（線でつながるノード）
        self.lanes = {l["id"]: l for l in rows("lanes") if named(l.get("id"))}
        self.edges = [e for e in rows("edges") if named(e.get("from")) and named(e.get("to"))]
        self.qs = rows("questions")
        pair = lambda q: tuple(q["edge"]) if isinstance(q.get("edge"), list) and len(q["edge"]) == 2 and all(map(named, q["edge"])) else None
        self.q_nodes = {x for q in self.qs for x in q_targets(q)}
        self.q_edges = {pair(q) for q in self.qs} - {None}
        self.q_pages = expand([s for q in self.qs for s in (q.get("sources") if isinstance(q.get("sources"), list) else [])])
        mids = {n["id"] for n in self.main}
        self.out_e, self.in_e = {i: [] for i in mids}, {i: [] for i in mids}    # 線でつながるノードどうしの線（edges の点検で埋める）


IF_WORDS = re.compile(r"場合|とき|時は|あれば|なければ|なら|に応じ")      # 分かれ目の条件（「翌月10日」「承認後」のような時期は、条件に数えない）


def reach(seeds, nxt):
    seen, todo = set(seeds), list(seeds)
    while todo:
        u = todo.pop()
        for v in nxt(u):
            if v not in seen:
                seen.add(v)
                todo.append(v)
    return seen


class FlowChecks:
    """全フローの点検。フローをまたいで数えるもの（網羅・ページのずれ・業務どうしの受け渡し）は、ここに貯めて最後に見る"""

    def __init__(self, cx, rp, flows):
        self.cx, self.rp, self.flows = cx, rp, flows
        asdict = lambda d: d if isinstance(d, dict) else {}
        tax = asdict(try_load(cx.work / "taxonomy.json")[0])
        self.hier = asdict(try_load(cx.work / "structure" / "hierarchy.json")[0])
        self.has_tax = bool(tax)
        self.cats = {c["id"]: c for c in (tax.get("categories") if isinstance(tax.get("categories"), list) else [])
                     if isinstance(c, dict) and isinstance(c.get("id"), str)}
        self.cat_pages = {i: expand(c.get("pages")) for i, c in self.cats.items()}
        self.flow_ids = {flow_key(f) for f in flows}
        self.g, self.gnames, self.galias = glossary_names(cx)
        self.sys_names, self.items_of, self.led = system_names(cx.work), CardItems(cx.work), None
        self.cited = set()                                           # 出典に引かれたページ
        self.excused = {k: "分類の対象外" for k in expand(tax.get("unassigned_pages"))}   # ページごと使わない理由のあるページ
        self.item_excused = set()                                    # 項目ごとの対象外（excluded_items）があるページ。残りの項目は、項目の網羅で確かめる
        self.cited_items, self.excluded_items, self.item_users = {}, {}, {}     # (資料, page, item) → 引いたノード／対象外の理由／作業として描いたノード
        self.actors = {}                                             # (フロー, ノード) → 引いた手順の担当（カードの actor。空は「書かれていない」）
        self.n = collections.Counter()                               # quotes・ok・img・nodes・inferred・checked
        self.shifted = {}                                            # (資料, 何ページずれているか) → ノード
        self.q_pages = set()                                         # 確認事項の出典になったページ
        self.files_out, self.files_in, self.links = {}, {}, set()    # 帳票名 → 作るフロー／使うフロー、(フロー, リンク先)

    def run(self, only=None):
        rp = self.rp
        for fl in self.flows:
            with guarded(rp, flow_key(fl)):
                self.check_one(FlowView(fl))
        content, items = self.cx.content_pages(), collections.Counter()
        for name, step in (("ページのずれ", self.systematic_shift), ("確認事項の無い曖昧なページ", self.unasked_pages),
                           ("分類とフローの対応", self.unmatched_flows), ("ページの網羅", lambda: self.page_coverage(content)),
                           ("項目の網羅", lambda: items.update(self.item_coverage(content))), ("業務どうしの受け渡し", self.handoffs)):
            with guarded(rp, f"フローをまたぐ点検（{name}）"):
                step()
        self.summary(content, items)
        if only:                                                     # 自分のフローの行だけを残す。行の頭がフローの id でない行（共通の行）は残す
            mine = lambda ln: any(re.match(rf"{re.escape(x)}[ :]", ln) for x in only)
            other = lambda ln: any(re.match(rf"{re.escape(x)}[ :]", ln) for x in self.flow_ids | set(self.cats))
            rp.fails[:] = [x for x in rp.fails if mine(x) or not other(x)]
            rp.warns[:] = [x for x in rp.warns if mine(x) or not other(x)]

    # ---- フロー1本 ----

    def check_one(self, v):
        self.shape(v)
        self.lanes(v)
        for n in v.nodes:
            self.node(v, n)
        self.edges(v)
        self.ends(v)
        self.card_order(v)
        self.entry_boxes(v)
        self.branches(v)
        self.questions(v)
        self.exclusions(v)
        self.size(v)
        self.timing(v)
        self.q_pages |= v.q_pages

    def shape(self, v):
        shape_fails(self.rp, v.F, validate({k: x for k, x in v.raw.items() if k != "_file"}, "flow"), SHAPE_ERRORS,
                    f"python3 scripts/validate.py -w <作業フォルダ> --kind flow --only {v.F}")
        if v.unnamed:
            self.rp.fail(f"{v.F}: id の無い（または文字でない）ノードが {v.unnamed} 個ある。点検できないので、先に id を付ける")
        if v.raw.get("_file") and v.raw["_file"].split(".")[0] != v.F:
            self.rp.fail(f"{v.F}: ファイル名（flows/{v.raw['_file']}）と flow_id が違う。ファイル名は <flow_id>.flow.json にする")
        for i in dupes(v.ids):
            self.rp.fail(f"{v.F}: ノードの id が重複: {i}")
        check_question_ids(self.rp, v.F, v.qs)

    def lanes(self, v):
        rp = self.rp
        if len(v.lanes) != len(v.raw.get("lanes") or []):
            rp.fail(f"{v.F}: レーンの id が重複")
        for l in v.lanes.values():                                   # kind の決まった語は、形の点検（スキーマ）が見る
            nm = norm(l.get("name", ""))
            if self.g is None or not nm or nm in self.gnames:
                continue
            if nm in self.galias:
                rp.fail(f"{v.F}: レーン名「{l.get('name')}」は、用語の正本では「{self.galias[nm][1]}」の別の表記。正本の名前にそろえる")
            else:
                rp.fail(f"{v.F}: レーン名「{l.get('name')}」が用語の正本（glossary.json）に無い。原典の呼び方から選ぶ。"
                        f"原典にある呼び方なら、親に知らせて glossary.json に足してもらう")
        if (self.cx.work / "flows" / f"{v.F}.edits.json").exists():  # 直しの一覧のフローは、機械が並べる（lane_order で決めた並びは、そのまま）
            return
        kinds = [v.lanes[x].get("kind") == "system" for x in v.lanes]
        if True in kinds and False in kinds[kinds.index(True):]:
            rp.warn(f"{v.F}: システムのレーンが右端に無い（決まり 4.2）。lanes の順を入れ替えるか、python3 scripts/renumber.py -w <作業フォルダ> --only {v.F} --lanes で並べ直す")

    def node(self, v, n):
        rp, nid, t = self.rp, f"{v.F} {n.get('id')}", n.get("type")
        if t not in NODE_TYPES:                                      # 決まった語でない種別は、形の点検（スキーマ）が FAIL にする。中身は見られない
            return
        if n.get("lane") not in v.lanes:
            rp.fail(f"{nid}: レーン「{n.get('lane')}」が無い")
        lab = str(n.get("label", "")).strip()
        if not lab and t not in ("start", "end"):
            rp.fail(f"{nid}: ラベルが空")
        if placeholder(n) and not n.get("inferred"):
            rp.fail(f"{nid}: ラベルに仮置きの語が残っている「{lab}」。分からないことは推定（inferred）にして確認事項に出す")
        if TOBE.search(lab + str(n.get("detail", "")) + str(n.get("reference", ""))):
            rp.warn(f"{nid}: 改善案らしい語がある「{lab}」。as-is のフローには現状だけを書く")
        if t in ("process", "system", "subprocess") and len(lab) > LABEL_LONG:
            rp.warn(f"{nid}: ラベルが長い（{len(lab)} 字）。「何をどうする」に絞り、条件や補足は detail に回す")
        if not placeholder(n):
            check_label_form(rp, nid, n)
        if len(str(n.get("detail") or "")) > DETAIL_MAX:
            rp.warn(f"{nid}: 補足（detail）が長い（{len(str(n['detail']))} 字。{DETAIL_MAX} 字まで）。図の箱には条件・時期の要点だけを書き、一覧や表の中身は reference に書く（原文は出典の引用に残る）")
        want = node_systems(n, str((v.lanes.get(n.get("lane")) or {}).get("name", "")), self.items_of, self.sys_names)
        if (n.get("systems") or []) != want:
            rp.fail(f"{nid}「{lab[:20]}」: 札（systems）が、出典に引いた手順の system と合わない（図 {'・'.join(map(str, n.get('systems') or [])) or 'なし'}／"
                    f"カード {'・'.join(want) or 'なし'}）。札は機械が付ける。python3 scripts/renumber.py -w <作業フォルダ> を実行する")
        for x in want:
            self.n["tags"] += 1
            if self.g is not None and norm(x) not in self.sys_names:
                rp.warn(f"{nid}: 札のシステム「{x}」が、用語の正本（glossary.json）にシステムとして無い。カードの system の書き方を確かめ、原典にある呼び方なら親に知らせる")
        if t == "file":
            self.file_node(v, n, nid, lab)
        if t == "subprocess":
            to = n.get("link_flow")
            if isinstance(to, str) and to in self.flow_ids:
                self.links.add((v.F, to))
            elif isinstance(to, str) and self.cats.get(to, {}).get("make_flow"):   # 行き先の分類はある。フローがまだ書かれていないだけ（行の頭は行き先の id）
                rp.fail(f"{to}: フロー（flows/{to}.flow.json）がまだ無い（{nid} がリンクしている。そろってから親が見る）")
            else:
                rp.fail(f"{nid}: 下位フロー link_flow「{to}」が無い（業務分類の id を書く）")
        if n.get("inferred"):
            self.n["inferred"] += 1
            if n["id"] not in v.q_nodes:
                rp.fail(f"{nid}: 推定（inferred）で置いたのに、このノードを指す確認事項が無い")
            cited = [s["item"] for s in n.get("sources") or [] if isinstance(s, dict) and isinstance(s.get("item"), str) and s["item"]]
            if cited:
                rp.fail(f"{nid}「{lab[:20]}」: カードの手順・分岐（{'・'.join(cited[:3])}）を出典に引いているのに、推定（inferred）になっている"
                        f"（作業は原典にある。推定の印は、原典に無いノードと線だけに付ける。担当・時期・現状行っているかが決まらないときは、印を外して確認事項で聞く）")
        elif t in WORK_TYPES and not any(isinstance(s, dict) and s.get("item") for s in n.get("sources") or []):
            rp.fail(f"{nid}「{lab[:20]}」: カードの手順・分岐を出典に引いていない（item が無い）のに、推定（inferred）の印が無い。原典に無いノードなら、"
                    f"python3 scripts/renumber.py で印を付けて確認事項を書く。原典にある作業なら、出典に item を書く（カードに無ければ、カードに手順を足す）")
        if t in SOURCED_TYPES:
            self.n["nodes"] += 1
            srcs = [s for s in n.get("sources") or [] if isinstance(s, dict)]
            if not srcs:
                rp.fail(f"{nid}「{lab[:20]}」: 出典が無い（どの資料の何ページに書いてあるか）")
            for s in srcs:
                self.source(v, n, nid, s)
            named = self.actors.get((v.F, n["id"]))
            if t in ("process", "decision") and named and not any(named) and n["id"] not in v.q_nodes and not self.lane_quoted(v, n, srcs):
                rp.fail(f"{nid}「{lab[:20]}」: もとの手順に担当が書かれていない（カードの actor が空）のに、このノードを指す確認事項が無い"
                        f"（レーンは仮に決めたもの。同じ趣旨の問いは1件にまとめ、対象を nodes に並べる）")

    def entry_boxes(self, v):
        """プログラムの起動点の箱（activity つきの system ノード）が、いまのシステム台帳の紐付け・アクティビティ図の名前のとおりか"""
        import copy

        from systems import Ledger, has_ledger, place_entries, placed
        boxes = placed(v.raw)
        for n in boxes:
            if n.get("type") != "system":
                self.rp.fail(f"{v.F} {n.get('id')}: activity（アクティビティ図の番号）は、システム処理のノードだけに付く")
        if not has_ledger(self.cx.work):
            if boxes:
                self.rp.fail(f"{v.F}: プログラムの起動点の箱（{'・'.join(str(n.get('id')) for n in boxes)}）があるのに、システム台帳（structure/systems.json）が無い")
            return
        if self.led is None:
            self.led = Ledger(self.cx.work, self.cx.m)
        self.n["boxes"] += len(boxes)
        want = copy.deepcopy({k: x for k, x in v.raw.items() if k != "_file"})
        if place_entries(want, self.led):
            now, new = {n.get("activity") for n in boxes}, {n.get("activity") for n in placed(want)}
            self.rp.fail(f"{v.F}: プログラムの起動点の箱が、いまのシステム台帳の紐付け・アクティビティ図の名前と合わない"
                         f"（足りない {'・'.join(sorted(new - now)) or 'なし'}／余分 {'・'.join(sorted(x for x in now - new if x)) or 'なし'}／ほかは名前・線・番号の違い）。"
                         f"箱は機械が置く。python3 scripts/renumber.py -w <作業フォルダ> を実行する")

    def lane_quoted(self, v, n, srcs):
        """ノードを置いたレーンの名前（別の表記を含む）が、出典の引用に出てくるか（「課長の承認を得てから発注します」→ 課長のレーンの「承認する」）"""
        lane = str((v.lanes.get(n.get("lane")) or {}).get("name", ""))
        names = {norm(lane)} | {k for k, x in self.galias.items() if x[1] == lane}
        quoted = norm(" ".join(str(s.get("quote", "")) for s in srcs))
        return any(k and k in quoted for k in names)

    def file_node(self, v, n, nid, lab):
        """帳票: どのプロセスで使うか・向き・名前が用語の正本にあるか"""
        rp = self.rp
        tgt = v.by_id.get(n.get("attach_to"))
        if tgt is None or tgt.get("type") == "file":
            rp.fail(f"{nid}: ファイルの attach_to（どのプロセスで使うか）が無い・不正")
        if n.get("direction") not in ("in", "out"):
            rp.fail(f"{nid}: ファイルの direction は in（参照する）/ out（作る）のどちらか")
        docs_known = [k for k, x in self.gnames.items() if x[0] == "document" and k]
        if self.g is not None and lab and not any(k in norm(lab) for k in docs_known):
            hit = next((x[1] for k, x in self.galias.items() if x[0] == "document" and k and k in norm(lab)), None)
            rp.fail(f"{nid}: 帳票名「{lab}」が用語の正本（glossary.json）に無い"
                    + (f"（「{hit}」の別の表記。正本の名前にそろえる）" if hit else "（原典の呼び方から選ぶ。原典にあるなら親に知らせて足してもらう）"))
        (self.files_out if n.get("direction") == "out" else self.files_in).setdefault(norm(lab), set()).add(v.F)

    def source(self, v, n, nid, s):
        """出典1件: ページがあるか・この分類のページか・引いた項目があるか・引用がそのページにあるか"""
        cx, rp = self.cx, self.rp
        k = (s.get("doc_id"), s.get("page"))
        if k not in cx.pages:
            rp.fail(f"{nid}: 存在しないページを出典にしている: {file_ref(k)}（page はファイル名 pNNNN の番号。刷られたページ番号ではない）")
            return
        self.cited.add(k)
        linked = self.cat_pages.get(n.get("link_flow")) or ()         # 別フローへのリンクは、リンク先のページを根拠にしてよい
        upper = {x for a in lineage(self.cats, self.flow_ids, v.F)[1:] for x in self.cat_pages.get(a) or ()}   # 下位フローは、上位の分類のページ（節の概要など）も根拠にしてよい
        code = is_code_page(cx.pages[k])                              # プログラムのソース（業務分類には割り当てない）
        if self.cat_pages.get(v.F) and k not in self.cat_pages[v.F] and k not in linked and k not in upper and not code:
            rp.warn(f"{nid}: 出典 {cx.ref(k)} は、この業務分類に割り当てたページでない（taxonomy.json を直すか、出典を見直す）")
        if code and not n.get("activity"):
            rp.fail(f"{nid}: プログラムのソース（{cx.ref(k)}）を出典にしている。業務フローでソースを出典にできるのは、機械が置く起動点の箱（activity つき）だけ。"
                    f"プログラムの中身は、アクティビティ図に描く")
        q = s["quote"].strip() if isinstance(s.get("quote"), str) else ""
        if not q:
            rp.fail(f"{nid}: 出典 {cx.ref(k)} に引用（quote）が無い")
            return
        if len(re.sub(r"[\W_]", "", norm(q))) < MIN_QUOTE:
            rp.fail(f"{nid}: 出典 {cx.ref(k)} の引用「{q}」が短すぎる（どの箇所かが分かるよう、原文の連続した一部を写す）")
            return
        self.n["quotes"] += 1
        card = cx.card(*k) or {}
        if isinstance(s.get("item"), str) and s["item"] and not code:
            self.cited_item(v, n, nid, k, s["item"], q, card)
        self.quote(nid, k, s, q, card)

    def cited_item(self, v, n, nid, k, item, q, card):
        it = safe_items(card).get(item)
        if it is None:
            self.rp.fail(f"{nid}: 出典が引くカードの項目 {item} が {self.cx.ref(k)} のカードに無い"
                         f"（項目の id はカード cards/{file_ref(k)}.json の steps・decisions・diagram にある）")
            return
        self.cited_items.setdefault(k + (item,), []).append(nid)
        if it["kind"] == "step":
            self.actors.setdefault((v.F, n["id"]), []).append(it.get("actor", ""))
        if n.get("type") in WORK_TYPES:
            self.item_users.setdefault(k + (item,), set()).add((v.F, n["id"]))
        if not same_text(q, it["quote"]) and not same_text(q, it["text"]):
            self.rp.fail(f"{nid}: 引用「{q[:30]}」が、引いた項目 {item}（{self.cx.ref(k)}）の引用「{it['quote'][:30]}」と合わない"
                         f"（項目の取り違えか、ページの取り違え）")

    def quote(self, nid, k, s, q, card):
        """引用がそのページの文字にあるか。無くて前後のページにあれば、どのページかを名指しする"""
        cx, rp, img = self.cx, self.rp, s.get("from_image")
        if "lines" in s and not img:                                 # 行を番号で指した出典: 引用は、その行の文字そのままのはず（機械が入れる）
            text = span_text(page_units(cx.pj(*k)), s["lines"])
            self.n["checked"] += 1
            if text is None:
                rp.fail(f"{nid}: 出典（{cx.ref(k)}）の行の番号 {s['lines']} が範囲の外（python3 scripts/pages.py lines で番号を見る）")
            elif norm(q) != norm(text):
                rp.fail(f"{nid}: 引用が、指した行 {s['lines']}（{cx.ref(k)}）の文字と違う。引用は手で書かず、python3 scripts/quotes.py fill で入れる: 「{q[:30]}」")
            else:
                self.n["ok"] += 1
                self.n["exact"] += 1
            return
        item = next((x for place in ("steps", "decisions") for x in (card or {}).get(place) or []
                     if isinstance(x, dict) and x.get("id") == s.get("item")), None) if s.get("item") else None
        if item is not None and "lines" in item and not item.get("from_image"):   # カードが行で指している項目: 出典の引用は、その項目の引用を機械が写したもののはず
            self.n["checked"] += 1
            if norm(q) == norm(str(item.get("quote", ""))):
                self.n["ok"] += 1                                    # 行の文字そのままかは、カードの点検が確かめている
                self.n["exact"] += 1
            else:
                rp.fail(f"{nid}: 引用が、引いた項目 {s['item']}（{cx.ref(k)}）の引用と違う。カードの行を直したあと、フローに写していない。"
                        f"python3 scripts/quotes.py fill で写す: 「{q[:30]}」")
            return
        r = quote_match(q, page_text(cx.pj(*k)))
        if r >= QUOTE_OK:
            self.n["ok"] += 1
            self.n["exact"] += r >= 1.0
            self.n["checked"] += 1
            if r < 1.0:                                              # 連続した一部そのままではない。主語・条件の違いが紛れていないかを人が見る
                rp.warn(f"{nid}: 引用が原典（{cx.ref(k)}）の文字と完全には一致しない（一致 {r:.0%}）。原文の連続した一部をそのまま写す: 「{q[:40]}」")
            return
        if img and quote_match(q, card_text(card)) >= QUOTE_OK:
            self.n["ok"] += 1
            self.n["img"] += 1
            return
        self.n["checked"] += 1
        other = next((o for o, _off in cx.neighbors(k, NEAR) if quote_match(q, page_text(cx.pj(*o))) >= 0.9
                      or (img and quote_match(q, card_text(cx.card(*o) or {})) >= 0.9)), None)
        if other:
            self.shifted.setdefault((k[0], other[1] - k[1]), []).append(nid)
            rp.fail(f"{nid}: 引用が出典のページ（{cx.ref(k)}）に無く、{cx.ref(other)} にある（出典のページが "
                    f"{abs(other[1] - k[1])} ページずれている。page を {other[1]} に直す）: 「{q[:30]}」")
        elif img:
            rp.fail(f"{nid}: 画像由来の引用が、ページカード（{cx.ref(k)}）に読み取った文字と合わない: 「{q[:40]}」")
        else:
            rp.fail(f"{nid}: 引用が原典（{cx.ref(k)}）の文字と合わない（一致 {r:.0%}）: 「{q[:40]}」。"
                    f"原文どおりに直す。画像の中の文字なら from_image を true にする")

    def edges(self, v):
        rp = self.rp
        for e in v.edges:
            a, b = e.get("from"), e.get("to")
            if a not in v.by_id or b not in v.by_id:
                rp.fail(f"{v.F}: 線 {a}→{b} が存在しないノードを指す")
                continue
            if a in v.out_e and b in v.in_e:
                v.out_e[a].append(e)
                v.in_e[b].append(e)
            else:
                rp.fail(f"{v.F}: 線 {a}→{b} の端が帳票（file）。帳票は線でつながず、attach_to で作業のノードに付ける")
            if e.get("inferred") and (a, b) not in v.q_edges and a not in v.q_nodes and b not in v.q_nodes:
                rp.fail(f"{v.F}: 線 {a}→{b} は推定（inferred）なのに、この線か両端のノードを指す確認事項が無い")
            if e.get("inferred"):
                self.n["inferred"] += 1

    def ends(self, v):
        """開始・終了があるか。全ノードが開始から届き、終了に届くか"""
        rp = self.rp
        starts = [n["id"] for n in v.main if n.get("type") == "start"]
        ends = [n["id"] for n in v.main if n.get("type") == "end"]
        if not starts:
            rp.fail(f"{v.F}: 開始ノードが無い")
        if not ends:
            rp.fail(f"{v.F}: 終了ノードが無い")
        for s in starts:
            if v.in_e[s]:
                rp.fail(f"{v.F}: 開始ノード {s} に入る線がある")
        for s in ends:
            if v.out_e[s]:
                rp.fail(f"{v.F}: 終了ノード {s} から出る線がある")
        fwd = reach(starts, lambda u: [e["to"] for e in v.out_e.get(u, [])])
        bwd = reach(ends, lambda u: [e["from"] for e in v.in_e.get(u, [])])
        for i in sorted(set(v.out_e) - fwd):
            rp.fail(f"{v.F} {i}: 開始からたどり着けない（線が抜けている）")
        for i in sorted(set(v.out_e) - bwd):
            rp.fail(f"{v.F} {i}: 終了にたどり着けない（行き止まり）")

    def card_order(self, v):
        """同じカードの続く手順（s2 → s3）を、フローが逆向きにつないでいないか（原典の順を、一般的な流れに合わせて入れ替えていないか）"""
        try:
            depth = order_of(v.raw)
        except Exception:                                            # 形の崩れたフロー。形の点検が FAIL にする
            return
        at = {}                                                      # (資料, page) → {項目: その項目を引く最初の作業のノード}
        for n in v.nodes:
            for s in (n.get("sources") or []) if n.get("type") in WORK_TYPES else []:
                if isinstance(s, dict) and isinstance(s.get("item"), str) and isinstance(s.get("page"), int):
                    at.setdefault((s.get("doc_id"), s["page"]), {}).setdefault(s["item"], n["id"])
        fwd = lambda u: [e["to"] for e in v.out_e.get(u, []) if depth.get(e["to"], 0) > depth.get(u, 0)]      # 戻りの線は見ない
        for k, items in at.items():
            steps = [x for x in (self.cx.card(*k) or {}).get("steps") or [] if isinstance(x, dict) and x.get("origin", "body") == "body" and x.get("id") in items]
            for x, y in zip(steps, steps[1:]):                       # 条件つきの手順（「不備がある場合は差し戻す」）は、どこで分かれるかの読み方なので見ない
                a, b = items[x["id"]], items[y["id"]]
                if a != b and x.get("section", "") == y.get("section", "") and not IF_WORDS.search(str(x.get("condition") or "") + str(y.get("condition") or "")) \
                        and v.by_id[a].get("type") != "decision" and v.by_id[b].get("type") != "decision" and a in reach([b], fwd) and b not in reach([a], fwd):
                    self.rp.warn(f"{v.F}: カードの手順の順（{self.cx.ref(k)} {x['id']} → {y['id']}）と逆につないでいる（{b}「{str(v.by_id[b].get('label', ''))[:16]}」が、"
                                 f"{a}「{str(v.by_id[a].get('label', ''))[:16]}」より前）。原典の順が一般的な流れと違っても、原典のままにする（flow-rules.md 3.3）。あとの項目が、前の作業の説明・前提を書いた文なら、このままでよい")

    def branches(self, v):
        rp = self.rp
        for n in v.main:
            oe = v.out_e.get(n.get("id"), [])
            if n.get("type") == "decision":
                head = f"{v.F} {n['id']}「{n.get('label')}」"
                if len(oe) < 2:
                    rp.fail(f"{head}: 分岐の出口が {len(oe)} 本（2本以上要る。片方が原典に無ければ推定で置いて確認事項に出す）")
                labs = [str(e.get("label", "")).strip() for e in oe]
                if any(not x for x in labs):
                    rp.fail(f"{head}: ラベルの無い分岐の線がある")
                if len(set(labs)) < len(labs):
                    rp.fail(f"{head}: 同じラベルの分岐がある（{labs}）")
                if len(oe) > EXITS_MAX:
                    rp.warn(f"{head}: 出口が {len(oe)} 本ある（{EXITS_MAX} 本まで。図は下・左・右にしか出せず、重なって読みにくい。判断を2段に分ける）")
                for x in labs:
                    if len(x) > LABEL_MAX:
                        rp.warn(f"{head}: 分岐の線のラベル「{x[:20]}」が長い（{len(x)} 字。{LABEL_MAX} 字まで。条件の中身は判断ノードの文言に書く）")
            elif len(oe) > 1:
                rp.warn(f"{v.F} {n['id']}: 分岐でないノードから線が {len(oe)} 本出ている（条件で分かれるなら判断ノードを置く。並行なら detail に書く）")

    def questions(self, v):
        rp = self.rp
        for end, (qid, key, _text, _assume) in ENDS_ASKED.items():
            if blank(v.raw.get(key)) and end in v.by_id and not any(end in q_targets(q) for q in v.qs if isinstance(q, dict)):
                rp.fail(f"{v.F}: {'開始' if end == 'start' else '終了'}条件が空なのに、それを聞く確認事項が無い（決まり 3.6）。"
                        f"python3 scripts/renumber.py -w <作業フォルダ> --only {v.F} で、機械の問い {qid} を付ける")
        for q in v.qs:
            for x in q_targets(q):
                if x not in v.by_id:
                    rp.fail(f"{v.F}: 確認事項 {q.get('id')} が存在しないノード {x} を指す")
            if q.get("edge") and not any([e["from"], e["to"]] == q["edge"] for e in v.edges):
                rp.fail(f"{v.F}: 確認事項 {q.get('id')} が存在しない線 {q['edge']} を指す（[\"出どころの id\", \"行き先の id\"] で書く）")
            if blank(q.get("question")):
                rp.fail(f"{v.F}: 確認事項 {q.get('id')} の質問文が空")
            where = ("edge", tuple(q["edge"])) if isinstance(q.get("edge"), list) else ("nodes", frozenset(q_targets(q)))
            twin = next((x.get("id") for x in v.qs if str(x.get("id", "")).startswith("m-") and x is not q and x.get("effect") == q.get("effect")
                         and (("edge", tuple(x["edge"])) if isinstance(x.get("edge"), list) else ("nodes", frozenset(q_targets(x)))) == where), None)
            if twin and not str(q.get("id", "")).startswith("m-") and where[1]:
                rp.warn(f"{v.F}: 確認事項 {q.get('id')} は、機械の確認事項 {twin} と同じ所を指し、答えで変わる所（effect）も同じ。同じことを聞いていれば、直しの一覧の問いを消す（決まり 5.6）")
            if q.get("effect") and q.get("severity") != severity_of(q):
                rp.fail(f"{v.F}: 確認事項 {q.get('id')} の重要度（{q.get('severity')}）が、決まり 5.4 の {severity_of(q)} と違う"
                        f"（flow.json を手で直したか、前の版の道具で作った）。python3 scripts/renumber.py -w <作業フォルダ> --only {v.F} で付け直す")
            for k in sorted(expand(q.get("sources")) - set(self.cx.pages), key=str):
                rp.fail(f"{v.F}: 確認事項 {q.get('id')} の出典が、存在しないページを指す: {file_ref(k)}")
        rows = lambda key: [x for x in (v.raw.get(key) if isinstance(v.raw.get(key), list) else []) if isinstance(x, dict)]
        shown = [q.get(k) for q in v.qs for k in ("question", "assumption")] + [n.get(k) for n in v.nodes for k in ("detail", "reference")]
        stale_refs(rp, v.F, shown + [x.get("reason") for key in ("unused_pages", "excluded_items") for x in rows(key)],
                   [n.get("label") for n in v.nodes] + [v.raw.get(k) for k in ("name", "start_condition", "end_condition")])

    def exclusions(self, v):
        """使わなかったページ（unused_pages）と、フローに入れなかった項目（excluded_items）"""
        cx, rp = self.cx, self.rp
        rows = lambda key: [x for x in (v.raw.get(key) if isinstance(v.raw.get(key), list) else []) if isinstance(x, dict)]
        mine = self.cat_pages.get(v.F) or set()
        for u in rows("unused_pages"):
            k = (u.get("doc_id"), u.get("page"))
            if blank(u.get("reason")):
                rp.fail(f"{v.F}: unused_pages の {file_ref(k)} に理由が無い")
            if k not in cx.pages:
                rp.fail(f"{v.F}: unused_pages が存在しないページを指す: {file_ref(k)}")
            elif mine and k not in mine:                             # ほかの分類のページを「使わない」と書いても、そのページの点検は外さない
                rp.fail(f"{v.F}: unused_pages の {cx.ref(k)} は、この分類に割り当てたページでない。ほかの分類のページは、その分類の担当が見る（書かない）")
            else:
                self.excused[k] = str(u.get("reason") or "")
        for x in rows("excluded_items"):
            k = (x.get("doc_id"), x.get("page"))
            if k not in cx.pages or x.get("item") not in safe_items(cx.card(*k) or {}):
                rp.fail(f"{v.F}: excluded_items が、存在しない項目を指す: {file_ref(k)} {x.get('item')}")
            else:
                self.excluded_items[k + (x["item"],)] = x["reason"].strip() if isinstance(x.get("reason"), str) else ""
                self.item_excused.add(k)

    def size(self, v):
        """採番の飛びと、フローの大きさ"""
        rp = self.rp
        frozen = (self.cx.work / "review" / f"{v.F}.review.json").exists()    # 原典照合のあとは番号を振り直さない（採否表の番号とずれる）。足し引きで飛んでもよい
        for pre in ("P", "S", "F", "D"):
            nums = [int(m.group(1)) for m in (re.fullmatch(rf"{pre}-(\d+)", str(i)) for i in v.ids) if m]
            if nums and sorted(nums) != list(range(1, len(nums) + 1)) and not frozen:
                rp.warn(f"{v.F}: {pre}- の採番に飛び・重なりがある（{sorted(nums)[:12]}…）。"
                        f"python3 scripts/renumber.py -w <作業フォルダ> --only {v.F} で振り直す")
        steps = [n for n in v.main if n.get("type") in STEP_TYPES]
        if len(steps) > MAX_STEPS:
            rp.warn(f"{v.F}: ステップが {len(steps)} 個。1ページ {MAX_STEPS} 個を目安に下位フロー（subprocess）へ分ける")
        if len(v.lanes) > MAX_LANES:
            rp.warn(f"{v.F}: レーンが {len(v.lanes)} 本。関係の薄い担当は別フローに分ける")

    def timing(self, v):
        """処理時間の仮の値のもと（動作の種類）が、人の作業の箱に付いているか。無ければ「その他」の行を引くので、止めない"""
        bare = [n["id"] for n in v.nodes if n.get("type") == "process" and not n.get("action")]
        if bare:
            self.rp.warn(f"{v.F}: 動作の種類（action）の無い作業が {len(bare)} 個（{'・'.join(bare[:6])}{' …' if len(bare) > 6 else ''}）。"
                         "業務マトリクス一覧の処理時間の仮の値が、「その他」の行になる。直しの一覧の op time で、動作の種類を選ぶ（flow-edits-guide.md）")
        made = [n["id"] for n in v.nodes if n.get("type") == "process" and "作成する" in actions_of(n.get("action")) and n.get("systems")]
        if made:
            self.rp.warn(f"{v.F}: 使うシステムの札が付いた作業で、動作の種類が「作成する」の箱が {len(made)} 個（{'・'.join(made[:6])}{' …' if len(made) > 6 else ''}）。"
                         "システムの画面で作る（入力して登録する）作業は「入力する」。紙・ファイルを一から作る作業なら、そのままでよい（flow-rules.md 3.22）")
        for n in v.nodes:
            if n.get("type") == "process" and n.get("volume") in VOLUME_SPARE and not str(n.get("volume_reason") or "").strip():
                self.rp.fail(f"{v.F} {n.get('id')}: 量が「{n['volume']}」なのに、理由（volume_reason）が無い。資料から、時間・日の単位でかかると読めるときだけ、そう読める所を書いて付ける（flow-rules.md 3.22）")

    # ---- 全フローを見終えてから ----

    def systematic_shift(self):
        """同じ資料の引用が、同じ向きに何件もずれているとき（出典を刷られた番号や目次の番号で書いた疑い）。先頭に出す"""
        for (doc, delta), who in self.shifted.items():
            if len(who) >= SHIFT_MIN and len(who) >= SHIFT_SHARE * max(1, self.n["checked"]):
                self.rp.fail(f"系統的なずれ — {doc} を出典にした引用 {len(who)} 件が、{abs(delta)} ページ{'後ろ' if delta > 0 else '前'}のページにある。"
                             f"出典の page を、刷られたページ番号や目次の番号で書いていないか（page はファイル名 pNNNN の番号）。{', '.join(who[:6])} …", first=True)

    def unasked_pages(self):
        for k in sorted(self.cited):
            c = self.cx.card(*k)
            if c and (c.get("unreadable") or c.get("uncertain")) and k not in self.q_pages:
                self.rp.warn(f"{self.cx.ref(k)} は読めなかった・曖昧な箇所があるのに、このページを出典にした確認事項がどのフローにも無い")

    def owners(self):
        """ページ → そのページを割り当てた、フローを作る分類（分類の並び順）"""
        out = {}
        for F, ps in self.cat_pages.items():
            if self.cats[F].get("make_flow"):
                for k in ps:
                    out.setdefault(k, []).append(F)
        return out

    def unmatched_flows(self):
        """フローと業務分類の対応。フローを書く担当は --stage flow しか回さないので、ここでも見る"""
        if not self.has_tax:
            return
        for F in sorted(self.flow_ids - set(self.cats)):
            self.rp.fail(f"{F}: 対応する業務分類が taxonomy.json に無い（flow_id を確かめる）")
        for F, c in self.cats.items():
            if c.get("make_flow") and F not in self.flow_ids:
                self.rp.fail(f"{F}: フロー（flows/{F}.flow.json）がまだ無い（「{c.get('name')}」）")
        for fl in self.flows:                                        # フローの名前・番号・親・段・準拠は、業務分類の写し（図の見出し・一覧の行に出る）
            c = self.cats.get(flow_key(fl))
            if c is None:
                continue
            want = flow_heads(c)
            diff = [f"{k}「{fl.get(k)}」→「{x}」" for k, x in want.items() if fl.get(k) != x]
            if diff:
                self.rp.fail(f"{flow_key(fl)}: 業務分類と違う（{'・'.join(diff)}）。python3 scripts/renumber.py -w <作業フォルダ> が分類から写す")
        flows_dir = self.cx.work / "flows"
        for f in sorted(flows_dir.glob("*.edits.json")) if flows_dir.exists() else []:
            F = f.name.split(".")[0]
            if F in self.flow_ids and F in self.cats:                # 直しの一覧のあるフローは、当て直したものと同じか（flow.json を手で直すと、次の apply_edits.py で消える）
                self.remade(F, f)
            if F not in self.cats:
                self.rp.fail(f"flows/{f.name}: 業務分類に無いフローの直しの一覧（分類の id を変えたなら、ファイル名も変える。要らなければ消す）")
            elif not self.cats[F].get("make_flow"):
                self.rp.fail(f"flows/{f.name}: フローを描かない分類（make_flow が false）の直しの一覧")

    def remade(self, F, ef):
        import apply_edits
        edits, err = try_load(ef)
        if err or not isinstance(edits, dict) or validate(edits, "edits"):
            self.rp.fail(f"{F}: 直しの一覧（flows/{ef.name}）が読めない・形が違う（python3 scripts/apply_edits.py -w <作業フォルダ> --flow {F} で全部出る）")
            return
        want, errs, _d = apply_edits.make(self.cx.work, F, edits)
        have = next(fl for fl in self.flows if flow_key(fl) == F)
        if errs:
            self.rp.fail(f"{F}: 直しの一覧に誤り {len(errs)} 件（例: {errs[0][:80]}）。python3 scripts/apply_edits.py -w <作業フォルダ> --flow {F} で全部出る")
        elif json.dumps(want, ensure_ascii=False, sort_keys=True) != json.dumps({k: x for k, x in have.items() if k != "_file"}, ensure_ascii=False, sort_keys=True):
            self.rp.fail(f"{F}: flow.json が、いまの直しの一覧・カード・業務分類から当て直したものと違う（直しの一覧を直したあと当てていない・flow.json を手で直した）。"
                         f"python3 scripts/apply_edits.py -w <作業フォルダ> --flow {F} を実行する（原典照合のあとも、番号は引き継ぐ）")

    def page_coverage(self, content):
        """内容ページが、どれかのノードの出典になっているか。行は、そのページを割り当てた分類ごとに出す（自分の分だけを直せるように）"""
        own, lost = self.owners(), {}
        for k in content:
            if k not in self.cited and k not in self.excused and k not in self.item_excused:
                for F in own.get(k) or [None]:
                    lost.setdefault(F, []).append(k)
        for F, ks in sorted(lost.items(), key=lambda x: str(x[0])):
            if F is None:
                self.rp.fail(f"どのノードにも使われていない内容ページ {len(ks)} 件: {listed(ks, self.cx.ref)}（どの分類にも割り当てていない。taxonomy.json を見直す）")
            elif F in self.flow_ids:
                self.rp.fail(f"{F}: この分類に割り当てたのに、どのノードにも使われていない内容ページ {len(ks)} 件: {listed(ks, self.cx.ref)}"
                             f"（フローに反映するか、flow.json の unused_pages に理由を書く）")

    def item_coverage(self, content):
        """カードの手順・分岐（補足の中の作業を含む）が、ノードの出典に引かれているか、対象外の理由があるか。数えた結果を返す"""
        cx, n = self.cx, collections.Counter()
        owner = {k: [F for F in fs if F in self.flow_ids] for k, fs in self.owners().items()}    # ページ → そのページを受け持つフロー
        for k in sorted(content, key=str):
            c = cx.card(*k)
            if (k in self.excused and k not in self.cited) or not c or not owner.get(k):
                continue                                             # ページごと使わない理由が書いてある・受け持つフローが無い
            items = safe_items(c)
            pure_diagram = not any(x["kind"] in ("step", "decision") for x in items.values())
            shapes = {b.get("id"): b.get("shape") for b in (c.get("diagram") or {}).get("nodes") or []}
            for i, it in items.items():
                if it["kind"] == "box" and not (pure_diagram and shapes.get(i) in WORK_TYPES):
                    continue                                         # 図の箱は、手順が1つも無い図だけのページのときに数える
                note = it["origin"] == "note"
                key = k + (i,)
                where = "in" if key in self.cited_items else "out" if self.excluded_items.get(key, "").strip() else "lost"
                n["all"] += 1
                n[where] += 1
                n["note"] += note
                n["note_" + where] += note
                if where == "lost":
                    what = ITEM_WORD[it["kind"]] + ("（補足の中の作業）" if note else "")
                    shared = f"このページは {'・'.join(owner[k])} にまたがる。どちらかで引く。" if len(owner[k]) > 1 else ""
                    for F in owner[k]:                               # またがるページは、どのフローの担当にも見えるように出す
                        self.rp.fail(f"{F}: {cx.ref(k)} の{what} {i}「{it['text'][:24]}」がフローに無い"
                                     f"（{shared}ノードの出典で item: \"{i}\" を引く。入れないなら excluded_items に理由を書く）")
        for key, users in self.item_users.items():
            fs = sorted({f for f, _n in users})
            if len(fs) > 1:
                self.rp.warn(f"同じ手順（{cx.ref(key[:2])} {key[2]}）が {len(fs)} つのフローに描かれている（"
                             + "・".join(f"{f} {x}" for f, x in sorted(users))
                             + "）。上位と下位に重ねて描いていないか。上位フローには下位フローへのリンクを置き、詳細は下位フローだけに描く")
        return n

    def handoffs(self):
        """業務どうしの「続く」（hierarchy.json の sequence）と、フローのリンク・帳票の受け渡しが合っているか"""
        rp, cats = self.rp, self.cats
        of_h = {}                                                    # 業務の木の id → フロー
        for i, c in cats.items():
            if c.get("hierarchy_id") and c.get("make_flow"):
                of_h.setdefault(c["hierarchy_id"], []).append(i)
        seqs = set()
        for q in self.hier.get("sequence") or []:
            for a in of_h.get(q.get("from"), []):
                for b in of_h.get(q.get("to"), []):
                    seqs.add((a, b))
                    if a in self.flow_ids and b in self.flow_ids and not any(tied(cats, self.flow_ids, x, y, {(a, b)}) for x, y in self.links):
                        rp.fail(f"{a}: 「{cats[a].get('name')}」は「{cats[b].get('name')}」（{b}）に続く、と hierarchy.json にあるのに、"
                                f"{a} のフロー（下位フローを含む）に {b} へのリンク（subprocess・link_flow）が無い")

        def family(a, b):
            return cats.get(a, {}).get("parent") == b or cats.get(b, {}).get("parent") == a
        for a, b in sorted(self.links - seqs):
            if not family(a, b) and not tied(cats, self.flow_ids, a, b, seqs):
                rp.warn(f"{a}: {b} へのリンクがあるのに、業務どうしの「続く」（hierarchy.json の sequence）に無い（根拠の文とページを添えて足す）")
        known = seqs | self.links
        for name, makers in self.files_out.items():
            for a in makers:
                for b in self.files_in.get(name, set()) - {a}:
                    kin = set(lineage(cats, self.flow_ids, a)) & set(lineage(cats, self.flow_ids, b))     # 同じ上位フローの下にある
                    if not kin and not tied(cats, self.flow_ids, a, b, known) and not tied(cats, self.flow_ids, b, a, known):
                        rp.warn(f"{a} が作る帳票「{name}」を {b} が使っているのに、{a} から {b} への「続く」もリンクも無い（受け渡しの抜けの疑い）")

    def summary(self, content, items):
        rp, n = self.rp, self.n
        used = sum(1 for k in content if k in self.cited)
        rp.info(f"フロー {len(self.flows)} 本・出典つきノード {n['nodes']} 個・推定 {n['inferred']} か所")
        rp.info(f"網羅率（ページ）{used}/{len(content)}（{(used / len(content) if content else 1):.0%}）。"
                f"使っていない {len(content) - used} ページのうち理由つき "
                f"{sum(1 for k in content if k not in self.cited and (k in self.excused or k in self.item_excused))}")
        rp.info(f"網羅率（手順・分岐の項目）{items['in'] + items['out']}/{items['all']}: フローに入れた {items['in']}・"
                f"対象外（理由つき）{items['out']}・どちらでもない {items['lost']}")
        rp.info(f"補足由来の手順 {items['note']} 件: 反映 {items['note_in']}・対象外（理由つき）{items['note_out']}")
        rp.info(f"引用照合率 {n['ok']}/{n['quotes']}（{(n['ok'] / n['quotes'] if n['quotes'] else 1):.0%}）。うち画像由来（カードとの照合）{n['img']} 件")
        sev = collections.Counter(q.get("severity") if q.get("severity") in ("high", "mid", "low") else "mid"
                                  for f in self.flows for q in (f.get("questions") if isinstance(f.get("questions"), list) else [])
                                  if isinstance(q, dict))
        rp.info(f"フローの確認事項 高 {sev['high']}・中 {sev['mid']}・低 {sev['low']}（階層・用語・分類の分は、確認事項一覧に合わせて載る）")
        NUMBERS.update(flows=len(self.flows), nodes=n["nodes"], inferred=n["inferred"], pages_content=len(content), pages_used=used,
                       items_all=items["all"], items_in=items["in"], items_out=items["out"], items_lost=items["lost"],
                       quotes=n["quotes"], quotes_ok=n["ok"], quotes_exact=n["exact"], quotes_image=n["img"])


def check_flows(cx, rp, only=None):
    bad = []
    flows = load_flows(cx.work, bad)
    for name, why in bad:
        rp.fail(f"flows/{name}: {why}（このフローは点検していない）")
    if not flows:
        rp.fail("flows/*.flow.json が無い")
        return
    known = {flow_key(f) for f in flows}
    if only and not only & known:
        rp.fail(f"--only の {'・'.join(sorted(only))} に当たるフローが無い（あるフロー: {'・'.join(sorted(known))}）。何も点検していない")
    FlowChecks(cx, rp, flows).run(only)


# ---------- 書き出した draw.io ----------

def template_vocab(styles, extra=()):
    """テンプレートが使っている色と形。図に、ここに無い色・形があれば FAIL にする。extra は、アクティビティ図で足すスタイル"""
    colors, sigs = {"none", "default", "#ffffff"}, set()

    def learn(style):
        st = style_dict(style)
        for k in COLOR_KEYS:
            if st.get(k):
                colors.add(st[k].lower())
        sigs.add(shape_sig(st))
    for p in styles["palettes"] + [styles["system_palette"]]:
        for k in ("header_style", "container_style", "process_style", "system_style"):
            if p.get(k):
                learn(p[k])
    for s in list(styles["node"].values()) + list(styles["edge"].values()) + [styles["info_style"], styles["text_style"]]:
        learn(s)
    for it in styles["legend"]:
        learn(it["style"])
    for k, s in styles["extensions"].items():
        if not k.endswith("_font"):                                  # 字の指定（CSS）は、形のスタイルでない
            learn(s)
    for s in extra:
        learn(s)
    sigs.add(("rounded", "process"))                                 # 下位フロー（プロセスの形に shape=process を足したもの）
    colors.add("#0066cc")
    return colors, sigs


def read_cells(rp, where, rt):
    """ページの中身を [(id, mxCell, 包んでいる object か None)] に"""
    cells = []
    for el in rt:
        if el.tag == "mxCell":
            cells.append((el.get("id"), el, None))
        elif el.tag in ("object", "UserObject"):
            c = el.find("mxCell")
            if c is None:
                rp.fail(f"{where}: 中身の無い object: {el.get('id')}")
                continue
            cells.append((el.get("id"), c, el))
    return cells


def check_cells(rp, where, cells, page_ids, colors, sigs):
    """図として壊れていないか（id・線の両端・位置）、テンプレートに無い色・形が無いか、ページリンクの行き先があるか"""
    ids = [i for i, _c, _o in cells]
    if "0" not in ids or "1" not in ids:
        rp.fail(f"{where}: 根のセル（id 0 と 1）が無い")
    for i in dupes(ids):
        rp.fail(f"{where}: セルの id が重複: {i}")
    idset = set(ids)
    for i, c, o in cells:
        if i in ("0", "1"):
            continue
        st = style_dict(c.get("style"))
        if c.get("edge") == "1":
            if c.get("source") not in idset or c.get("target") not in idset:
                rp.fail(f"{where}: 線 {i} の両端が無い（{c.get('source')}→{c.get('target')}）")
            if c.find("mxGeometry") is None:
                rp.fail(f"{where}: 線 {i} に mxGeometry が無い")
        elif c.get("vertex") == "1" and c.find("mxGeometry") is None:
            rp.fail(f"{where}: 図形 {i} に mxGeometry が無い")
        for k in COLOR_KEYS:
            if st.get(k) and st[k].lower() not in colors:
                rp.fail(f"{where}: {i} の {k}={st[k]} はテンプレートに無い色")
        if shape_sig(st) not in sigs:
            rp.fail(f"{where}: {i} の形（{shape_sig(st)}）はテンプレートに無い")
        link = o.get("link") if o is not None else None
        if link and link.startswith("data:page/id,") and link.split(",", 1)[1] not in page_ids:
            rp.fail(f"{where}: {i} のリンク先のページ（{link.split(',', 1)[1]}）がこのファイルに無い")


def check_page_matches_flow(rp, where, cells, fl, builder="build_drawio.py"):
    """図のページが flow.json（アクティビティ図は activity.json）と同じか（ノード・ラベル・出典の属性・線・レーン）。違えば、図を手で直した疑い"""
    v = FlowView(fl)
    want = v.by_id
    got = {o.get("node_id"): o for _i, _c, o in cells if o is not None and o.get("node_id")}
    for i in sorted(set(want) - set(got)):
        rp.fail(f"{where}: flow.json のノード {i} が図に無い（{builder} を実行し直す）")
    for i in sorted(set(got) - set(want)):
        rp.fail(f"{where}: 図のノード {i} が flow.json に無い（図を手で直した疑い。flow.json か直しの一覧を直して書き出し直す）")
    for i, o in got.items():
        n = want.get(i)
        if n is None:
            continue
        lab = html.unescape(re.sub(r"<[^>]+>", "", o.get("label", "")))   # 箱の名前の & や < は、図の中では文字参照になっている
        if norm(n.get("label", "")) not in norm(lab) and n.get("type") not in ("start", "end"):
            rp.fail(f"{where}: ノード {i} のラベルが flow.json と違う（図「{lab[:24]}」／JSON「{str(n.get('label', ''))[:24]}」）")
        if n.get("type") in SOURCED_TYPES and not o.get("src"):
            rp.fail(f"{where}: ノード {i} に出典の属性（src）が無い")
        for x in n.get("systems") or []:
            if norm(f"［{x}］") not in norm(lab):
                rp.fail(f"{where}: ノード {i} の札「［{x}］」が図に無い（build_drawio.py を実行し直す）")
    main_ids = {n["id"] for n in v.main}
    want_e = sorted((e["from"], e["to"]) for e in v.edges if e["from"] in main_ids and e["to"] in main_ids)
    got_e = sorted((c.get("source"), c.get("target")) for _i, c, _o in cells
                   if c.get("edge") == "1" and c.get("source") in main_ids and c.get("target") in main_ids)
    if want_e != got_e:
        miss = [e for e in want_e if e not in got_e]
        extra = [e for e in got_e if e not in want_e]
        rp.fail(f"{where}: 線が flow.json と違う（図に無い {miss[:5]}／JSON に無い {extra[:5]}）")
    heads = [c.get("value") for i, c, _o in cells if str(i).startswith("lane-") and str(i).endswith("-h")]
    if [norm(x) for x in heads] != [norm(l.get("name")) for l in v.lanes.values()]:
        rp.fail(f"{where}: レーンが flow.json と違う（図 {heads}）")


DATED = re.compile(r"作成: \d{4}-\d{2}-\d{2}")


def check_regenerated(cx, rp, files, styles):
    """書き出した図（と配置）が、いまのデータから作り直したものと同じか（作成日を除いて完全一致）。
    違えば、図を手で直したか、データを直したあと書き出していない。ラベルの部分一致・線の両端だけでは見えない違い（補足・札・出典・色・位置）も拾う"""
    import build_drawio as BD
    want = {}
    if any(p.name.startswith(FLOW_FIG) for p in files):
        groups, _ctx = BD.figures(cx.work, styles, split=split_figs(cx.work / "out"), write_questions=False)
        want.update({f"{FLOW_FIG}{suffix}.drawio": BD.render(f"{FLOW_FIG}{suffix}.drawio", pages) for suffix, pages, _n in groups})
    if any(p.name == ACT_FIG for p in files):
        import build_activity as BA
        acts = BA.load_activities(cx.work)
        sty = try_load(ASSETS / BA.STYLES)[0]
        if acts and isinstance(sty, dict):
            want[ACT_FIG] = BD.render(ACT_FIG, BA.figure(cx.work, acts, sty, write_questions=False))
    for p in files:
        if p.name not in want:
            rp.fail(f"{p.name}: いまのデータから書き出す図の名前に無い（手で写した図・前の分け方の残り。納品物に写るので、out/ から外す）")
            continue
        text, layout = want[p.name]
        if DATED.sub("作成: ", p.read_text(encoding="utf-8")) != DATED.sub("作成: ", text):
            rp.fail(f"{p.name}: いまのデータから作り直した図と違う（図を手で直した・データを直したあと書き出していない）。"
                    f"python3 scripts/{'build_activity' if p.name.startswith('アクティビティ図') else 'build_drawio'}.py -w <作業フォルダ> で書き出し直す")
        elif try_load(p.with_suffix(".layout.json"))[0] != json.loads(json.dumps(layout)):
            rp.fail(f"{p.with_suffix('.layout.json').name}: いまのデータから作り直した配置と違う（書き出し直す）")
    for name in sorted(set(want) - {p.name for p in files}):
        rp.fail(f"{name} が無い（書き出していない）")


def check_drawio(cx, rp, files, styles):
    flows = {flow_key(f): f for f in load_flows(cx.work, [])}        # 読めないフローは、flow の段階が FAIL にする
    colors, sigs = template_vocab(styles)
    ast = try_load(ASSETS / "activity_styles.json")[0]                # アクティビティ図は、そのテンプレートから抜いたスタイルで確かめる
    if not isinstance(ast, dict) or "legend_lines" not in ast:
        ast = dict(styles, start="", end="", legend_style="", process_extra="")
        if (cx.work / "activities").exists():
            rp.fail("assets/activity_styles.json が無い・古い（python3 scripts/harvest_template.py で、アクティビティ図のテンプレートから作り直す）")
    a_colors, a_sigs = template_vocab(ast, [ast.get(k, "") for k in ("start", "end", "legend_style")] + [
        (p.get("process_style") or "") + ast.get("process_extra", "") for p in ast["palettes"] + [ast["system_palette"]]])
    acts = {}
    for f in sorted((cx.work / "activities").glob("*.activity.json")) if (cx.work / "activities").exists() else []:
        a = try_load(f)[0]
        if isinstance(a, dict) and a.get("flow_id"):
            acts[a["flow_id"]] = a
    from systems import Ledger, has_ledger
    if acts and has_ledger(cx.work):                                 # 図にしない手続き（補助）の、前に作った図は、ページにならない
        led = Ledger(cx.work, cx.m)
        for u in led.units:
            if u["kind"] == "shared" and led.off(u) is not None:
                acts.pop(u["activity"], None)
    seen_pages, seen_acts, act_file = set(), set(), False
    for path in files:
        try:
            root = ET.parse(path).getroot()
        except ET.ParseError as e:
            rp.fail(f"{path.name}: XML として読めない（{e}）")
            continue
        diagrams = root.findall("diagram")
        page_ids = {d.get("id") for d in diagrams}
        is_act = "systems" in page_ids or any(str(i).startswith("act-") for i in page_ids)    # アクティビティ図のファイル
        act_file = act_file or is_act
        if is_act and "systems" not in page_ids:
            rp.fail(f"{path.name}: システムと起動点の一覧のページが無い")
        elif not is_act and "overview" not in page_ids:
            rp.fail(f"{path.name}: 全体構成のページが無い")
        for d in diagrams:
            where = f"{path.name}［{d.get('name')}］"
            rt = d.find("mxGraphModel/root")
            if rt is None:
                rp.fail(f"{where}: mxGraphModel/root が無い")
                continue
            cells = read_cells(rp, where, rt)
            check_cells(rp, where, cells, page_ids, *((a_colors, a_sigs) if is_act else (colors, sigs)))
            did = d.get("id") or ""
            if did.startswith("act-"):
                seen_acts.add(did[4:])
                if did[4:] not in acts:
                    rp.fail(f"{where}: 対応する activity.json（{did[4:]}）が無い")
                else:
                    with guarded(rp, f"{where}（activity.py apply で作り直してから、build_activity.py で書き出し直す）"):
                        check_page_matches_flow(rp, where, cells, acts[did[4:]], "build_activity.py")
                continue
            if not did.startswith("flow-"):
                continue
            F = did[5:]
            seen_pages.add(F)
            if F not in flows:
                rp.fail(f"{where}: 対応する flow.json（{F}）が無い")
                continue
            with guarded(rp, f"{where}（flow.json の形を直してから、build_drawio.py で書き出し直す）"):
                check_page_matches_flow(rp, where, cells, flows[F])
    for F in sorted(set(flows) - seen_pages):
        rp.fail(f"フロー {F} のページが draw.io に無い")
    for A in sorted(set(acts) - seen_acts) if act_file or acts else []:
        rp.fail(f"アクティビティ図 {A} のページが draw.io に無い（python3 scripts/build_activity.py -w <作業フォルダ> で書き出す）")
    rp.info(f"draw.io {len(files)} ファイル・フローのページ {len(seen_pages)}" + (f"・アクティビティ図のページ {len(seen_acts)}" if seen_acts else ""))


# ---------- プログラムのソース ----------

def check_code(cx, rp):
    """プログラムのソース: 認証情報が伏せてある／構造（code/*.struct.json）があり、いまのページから取り出し直したものと同じ／木の形が正しい"""
    import json

    import code_struct as cs
    projs = cs.projects(cx.m)
    locked = [d["doc_id"] for d in cx.m["docs"] if d.get("vba_locked") and not d.get("code")]
    if locked:
        rp.warn(f"VBA プロジェクトに閲覧のロックが掛かっていて、マクロを読んでいない資料: {'・'.join(locked)}（読んでよいかを利用者に確かめる）")
    if not projs:
        rp.info("プログラムのソースは無い")
        return
    have, n = cs.load_structs(cx.work), collections.Counter()
    again = "python3 scripts/code_struct.py -w <作業フォルダ> を実行する"
    for proj in projs:
        for d in proj["docs"]:
            for p in d.get("page_list") or []:
                if is_code_page(p):
                    hit = cs.mask_secrets((cx.pj(d["doc_id"], p["page"]) or {}).get("text") or [])[1]
                    if hit:
                        rp.fail(f"{cx.ref((d['doc_id'], p['page']))}: 認証情報らしい文字が伏せられていない（{len(hit)} 行）。"
                                f"python3 scripts/render_pages.py -w <作業フォルダ> --force --only {d['doc_id']} でページ化し直す")
        name, st = f"code/{proj['id']}.struct.json", have.get(proj["id"])
        if st is None:
            rp.fail(f"{name} が無い（{again}）")
            continue
        now = cs.struct_of(cx.work, proj)
        if json.dumps(now, ensure_ascii=False, sort_keys=True) != json.dumps(st, ensure_ascii=False, sort_keys=True):
            rp.fail(f"{name}: いまのページ" + ("・LLM が書いた木" if st.get("tier") == "llm" else "") + f"から取り出し直した構造と違う（ページ化し直した・木を直した・手で直した）。{again}")
            continue
        rd = st.get("read")
        if rd is not None:                                           # 機械で構文を読めない言語: LLM が書いた木（code/<まとまり>.tree.json）
            tree_name = f"code/{proj['id']}.tree.json"
            if not rd["tree"]:
                n[st.get("tier")] += 1
                rp.fail(f"{tree_name} が無い（機械で構文を読めない言語。LLM がソースを読んで、ブロックの木を書く。読まないと決めたなら、木に skip と理由を書く。references/code-guide.md）")
                continue
            tree = read_checked(rp, cx.work / tree_name, tree_name, "") or {}   # 呼び元を 1 行に 1 つ（被覆の控えが、どちらの部品かを見分けられるように）
            shape_fails(rp, tree_name, validate(tree, "code_tree"), STRUCT_ERRORS, "python3 scripts/validate.py -w <作業フォルダ> --kind code_tree")
            for e in (rd.get("errors") or [])[:STRUCT_ERRORS]:
                rp.fail(f"{tree_name}: {e}")
            if not rd.get("errors") and not st["procedures"] and not rd.get("skip"):
                rp.fail(f"{tree_name}: 手続きが1つも無い（読まないと決めたなら、skip に理由を書く）")
            if not rd.get("errors") and st["procedures"] and not st["entries"]:
                rp.fail(f"{tree_name}: 起動点（entries）が1つも無い（外から始まる所を、根拠の行つきで書く）")
        for e in cs.tree_errors(st)[:STRUCT_ERRORS]:
            rp.fail(f"{name}: 取り出した構造の誤り — {e}")
        n.update(procs=len(st["procedures"]), entries=len(st["entries"]), units=len(st["units"]), unreached=len(st["unreached"]),
                 dynamic=sum(len(p.get("dynamic") or []) for p in st["procedures"]))
        n[st.get("tier")] += 1
    rp.info(f"プログラム {len(projs)} まとまり（構造の取り出し: 構文から {n['parsed']}・行の形から {n['heuristic']}・LLM の読取 {n['llm']}）／手続き {n['procs']}・"
            f"起動点 {n['entries']}・図にする単位 {n['units']}・たどれない手続き {n['unreached']}・実行時に決まる名前 {n['dynamic']}")


# ---------- システム台帳 ----------

def check_systems(cx, rp):
    """システム台帳: 形／番号・名前の重なり／用語の正本との対応／プログラムの資料の割り当て／紐付けの根拠（両方に文字がある）／図にしない起動点"""
    import code_struct as cs
    from systems import SYS_ID, Ledger, has_ledger
    projs = cs.projects(cx.m)
    name = "structure/systems.json"
    if not has_ledger(cx.work):
        if projs:
            rp.fail(f"{name} が無い（プログラムのソースがある。python3 scripts/systems.py draft -w <作業フォルダ> で下書きを書く）")
        else:
            rp.info("システム台帳は無い（プログラムのソースが無いので、無くてよい。システム一覧は、用語の正本から作る）")
        return
    data = read_checked(rp, cx.work / "structure" / "systems.json", name, "")
    if data is None:
        return
    shape_fails(rp, name, validate(data, "systems"), STRUCT_ERRORS, "python3 scripts/validate.py -w <作業フォルダ> --kind systems")
    led = Ledger(cx.work, cx.m)
    for i in dupes([s.get("id") for s in led.systems]):
        rp.fail(f"{name}: システムの番号が重複: {i}")
    for i in dupes([norm(s.get("name")) for s in led.systems]):
        rp.fail(f"{name}: 同じ名前のシステムが2つある: {i}")
    g = try_load(cx.work / "glossary.json")[0]
    terms = {norm(t.get("name")): t.get("name") for t in (g.get("terms") if isinstance(g, dict) else None) or [] if isinstance(t, dict) and t.get("kind") == "system"}
    if isinstance(g, dict):
        for s in led.systems:
            if norm(s["name"]) not in terms:
                rp.fail(f"{name}: システム「{s['name']}」が、用語の正本（glossary.json）に kind: system として無い。原典での呼び方にそろえるか、"
                        f"プログラムのファイル名から付けた名前なら、glossary.json に足す（レーンの名前に使う）")
        lost = [t for k, t in terms.items() if k not in {norm(s["name"]) for s in led.systems}]
        if lost:
            rp.fail(f"{name}: 用語の正本のシステムが、台帳に無い: {'・'.join(lost[:8])}（python3 scripts/systems.py draft で足す）")
    code_docs = {d["doc_id"] for p in projs for d in p["docs"]}
    owner = {}
    for s in led.systems:
        for d in s.get("docs") or []:
            if d not in code_docs:
                rp.fail(f"{name}: {s.get('id')}「{s['name']}」の docs の {d} は、プログラムのある資料でない（マクロ付きのブック・ソースのファイルの資料 ID を書く）")
            if d in owner:
                rp.fail(f"{name}: 資料 {d} が、2つのシステム（{owner[d]}・{s.get('id')}）に入っている")
            owner[d] = s.get("id")
    for p in projs:
        ids = [d["doc_id"] for d in p["docs"]]
        if p["id"] not in led.structs:
            rp.fail(f"code/{p['id']}.struct.json が無い（python3 scripts/code_struct.py -w <作業フォルダ> を実行する）")
        if len({owner.get(d) for d in ids}) > 1 or None in {owner.get(d) for d in ids}:
            rp.fail(f"{name}: プログラムの資料 {'・'.join(ids)}（1つのまとまり）が、同じ1つのシステムの docs に入っていない"
                    f"（入っていない {'・'.join(d for d in ids if d not in owner) or 'なし'}）")
    entries = led.entries()
    keys = {e["key"] for e in entries}
    seen = set()
    for i, x in enumerate(led.links):
        for err in led.link_errors(x):
            rp.fail(f"{name}: links[{i}]（{x.get('entry')} ⇔ {x.get('step')}）: {err}")
        if (x.get("entry"), x.get("step")) in seen:
            rp.fail(f"{name}: links[{i}]（{x.get('entry')} ⇔ {x.get('step')}）が、2回書かれている")
        seen.add((x.get("entry"), x.get("step")))
    for k in led.excluded:
        if k not in keys:
            rp.fail(f"{name}: excluded_entries の「{k}」という起動点が無い")
    for err in led.excluded_errors():                                # 図にしない起動点の理由が、中身と合っているか
        rp.fail(f"{name}: excluded_entries の{err}")
    for err in led.folded_errors():                                  # 図にしない手続き（補助）は、機械が出す候補（入出力・やり取り・終了が無い、共通の処理）だけ
        rp.fail(f"{name}: folded_procedures の{err}")
    helpers = led.helpers()
    if helpers:
        off = [h for h in helpers if h["reason"] is not None]
        rp.info(f"図にしない手続き（補助）{len(off)}（実行文 {sum(h['exec'] for h in off)} 行）／補助の候補のうち、図のままにしたもの {sum(h['reason'] is None for h in helpers)}")
    check_question_ids(rp, name, [q for q in data.get("questions") or [] if isinstance(q, dict)])
    for s in led.systems:
        if s.get("docs") and not SYS_ID.match(str(s.get("id") or "")):
            rp.fail(f"{name}: システム「{s['name']}」の id が SYS-01 の形でない（アクティビティ図の番号を振れない）")
    linked = {x.get("entry") for x in led.good_links()}
    rp.info(f"システム {len(led.systems)}（プログラムのあるもの {sum(bool(s.get('docs')) for s in led.systems)}）／起動点 {len(entries)}"
            f"（業務フローに紐付けたもの {sum(e['key'] in linked for e in entries)}・紐付けの無いもの {sum(e['key'] not in linked and e['excluded'] is None for e in entries)}・"
            f"図にしないもの {sum(e['excluded'] is not None for e in entries)}）／紐付け {len(led.links)}"
            f"（名指し {sum(x.get('basis') == 'named' for x in led.links)}・入出力の一致 {sum(x.get('basis') == 'io' for x in led.links)}・推定 {sum(x.get('basis') == 'inferred' for x in led.links)}）")


# ---------- アクティビティ図 ----------

def check_activity(cx, rp):
    """アクティビティ図のデータ: 図にする単位ごとに図がある／直しの一覧から作り直したものと同じ／名前が付いている／
    行の網羅（実行文が、どれかの箱か対象外に1回だけ入る）／分岐の網羅（分岐・繰り返しが、判断か畳みのどちらかにある）／
    かぎかっこの中の名前が、指したソースの行にある／呼び出しの箱の行き先がある"""
    import json

    from activity import MAX_BOXES, Draft, make, ordered
    from systems import Ledger, has_ledger
    if not has_ledger(cx.work):
        rp.info("アクティビティ図は無い（システム台帳が無い）")
        return
    led = Ledger(cx.work, cx.m)
    adir = cx.work / "activities"
    units = [u for u in ordered(led) if led.off(u) is None]
    for u in led.units:
        if u["kind"] == "shared" and led.off(u) is not None and any((adir / f"{u['activity']}.{x}.json").exists() for x in ("activity", "edits")):
            rp.warn(f"{u['activity']}: 図にしない手続き（{u['procedure']}）なのに、前に作った図のデータが残っている（activities/{u['activity']}.*.json。図には出ない。消してよい）")
    n = collections.Counter()
    import terms as T
    try:
        tm = T.Terms(cx.work, led.structs)
    except BadJson:
        rp.fail(f"{T.FILE}: JSON として読めない。書きかけ・閉じ忘れを直す（名付けの点検ができないので、アクティビティ図の点検は、ここで止めた）")
        return
    if tm.exists():                                                  # 用語の対応（プログラムの中の名前 → 業務の言葉）
        with guarded(rp, T.FILE):
            raw, terr = try_load(cx.work / T.FILE)
            if terr or not isinstance(raw, dict):
                rp.fail(f"{T.FILE}: 読めない（{terr or '形が違う'}）")
            else:
                shape_fails(rp, T.FILE, validate(raw, "code_terms"), SHAPE_ERRORS, "references/schemas.md の「用語の対応」")
                for msg in tm.problems():
                    rp.fail(f"{T.FILE}: {msg}")
                for msg in tm.warnings():
                    rp.warn(f"{T.FILE}: {msg}")
    for u in units:
        aid = u["activity"]
        name = f"activities/{aid}.activity.json"
        with guarded(rp, name):
            a, err = try_load(adir / f"{aid}.activity.json")
            if err or not isinstance(a, dict):
                rp.fail(f"{aid}: アクティビティ図がまだ無い・読めない（{u['procedure']}。python3 scripts/activity.py draft -w <作業フォルダ> --id {aid} で下書きを見て、"
                        f"直しの一覧 activities/{aid}.edits.json を書き、activity.py apply で作る）")
                continue
            shape_fails(rp, name, validate(a, "activity"), SHAPE_ERRORS, f"python3 scripts/activity.py apply -w <作業フォルダ> --id {aid}")
            edits, eerr = try_load(adir / f"{aid}.edits.json")
            if eerr or not isinstance(edits, dict):
                rp.fail(f"{aid}: 直しの一覧（activities/{aid}.edits.json）が無い・読めない。下書きのままの図（--no-edits）は、確かめ用")
                edits = {}
            else:
                shape_fails(rp, f"activities/{aid}.edits.json", validate(edits, "activity_edits"), SHAPE_ERRORS, "references/schemas.md の「アクティビティ図の直しの一覧」")
            want, errs, _d = make(led, u, edits)
            if errs:
                rp.fail(f"{aid}: 直しの一覧に誤り {len(errs)} 件（例: {errs[0][:80]}）。python3 scripts/activity.py apply -w <作業フォルダ> --id {aid} で全部出る")
            elif json.dumps(want, ensure_ascii=False, sort_keys=True) != json.dumps(a, ensure_ascii=False, sort_keys=True):
                rp.fail(f"{aid}: 図のデータが、いまの構造・直しの一覧・台帳から作り直したものと違う（直したあと作り直していない・手で直した）。"
                        f"python3 scripts/activity.py apply -w <作業フォルダ> --all を実行する")
            ids = {x.get("id") for x in a.get("nodes") or [] if isinstance(x, dict)}
            draft = Draft(led, u)
            need = draft.exec_lines()                                # 構造から数え直した、この図が受け持つ実行文の行
            got = collections.Counter((o.get("module"), ln) for o in a.get("lines") or [] for ln in o.get("lines") or [])
            lost, twice, extra = sorted(need - set(got)), sorted(k for k, v in got.items() if v > 1), sorted(set(got) - need)
            show = lambda ks: "、".join(f"{m} {ln} 行" for m, ln in ks[:6]) + (f" ほか {len(ks) - 6} 行" if len(ks) > 6 else "")
            if lost:
                rp.fail(f"{aid}: 行の網羅: 実行文 {len(lost)} 行が、どの箱にも対象外にも入っていない（{show(lost)}）")
            if twice:
                rp.fail(f"{aid}: 行の網羅: 同じ行が、2つの箱・対象外に入っている（{show(twice)}）")
            if extra:
                rp.fail(f"{aid}: 行の網羅: この図の手続きに無い行が入っている（{show(extra)}）")
            for o in a.get("lines") or []:
                if o.get("state") in ("box", "folded") and o.get("node") not in ids:
                    rp.fail(f"{aid}: 行の持ち主の箱 {o.get('node')}（{o.get('key')}）が、図に無い")
            forks = {x["key"] for x in draft.walk() if x["kind"] in ("if", "select", "loop")}
            logged = {b.get("key"): b for b in a.get("branches") or []}
            miss = sorted(forks - set(logged))
            if miss:
                rp.fail(f"{aid}: 分岐の網羅: 分岐・繰り返し {len(miss)} か所が、判断にも畳みにもなっていない（{'・'.join(miss[:6])}）")
            for b in logged.values():
                if b.get("state") in ("decision", "loop_box", "folded") and b.get("node") not in ids:
                    rp.fail(f"{aid}: 分岐 {b.get('key')} の箱 {b.get('node')} が、図に無い")
            boxes = [x for x in a.get("nodes") or [] if x.get("type") in STEP_TYPES]
            question_wording(rp, aid, a.get("questions"))
            for q in a.get("questions") or []:
                if isinstance(q, dict) and q.get("effect") and q.get("severity") != severity_of(q):
                    rp.fail(f"{aid}: 確認事項 {q.get('id')} の重要度（{q.get('severity')}）が、決まり 5.4 の {severity_of(q)} と違う"
                            f"（前の版の道具で作った）。python3 scripts/activity.py apply -w <作業フォルダ> --id {aid} をやり直す")
            proj, own = u["project"], {}
            lang, known, unknown = tm.lang(proj), tm.known(proj), tm.unknown(proj)
            for o in a.get("lines") or []:
                if o.get("state") != "excluded":
                    own.setdefault(o.get("node"), []).append(tm.text(proj, o.get("module"), o.get("lines") or []))
            whole = "\n".join(t for ts in own.values() for t in ts)
            titled = f"{a.get('name', '')} {a.get('summary', '')}"
            for tok in T.inner_names(titled, T.code_names(whole, lang) | known, known):
                rp.fail(f"{aid}: 図の名前・説明に、プログラムの中の名前「{tok}」が、そのまま出ている（読むのは、業務部門の担当。直しの一覧の name・summary を、業務の言葉で書く）")
            for q in T.quoted_inner(titled, whole, lang, norm(" ".join(f.get("label", "") for f in a["nodes"] if f.get("type") == "file")), unknown):
                rp.fail(f"{aid}: 図の名前・説明の、かぎかっこの中の「{q}」は、プログラムの中だけの名前（変数・項目・手続き）。業務の言葉に直す")
            for x in boxes:
                if not x.get("unnamed"):                             # 名前の中に、プログラムの中だけの名前が出ていないか（判断は、出口のラベルも）
                    code = "\n".join(own.get(x["id"], []))
                    said = " ".join([str(x.get("label", "")), str(x.get("detail", ""))] + [str(e.get("label", "")) for e in a.get("edges") or [] if e.get("from") == x["id"]])
                    mine = norm(" ".join(f.get("label", "") for f in a["nodes"] if f.get("attach_to") == x["id"]))
                    for tok in T.inner_names(said, T.code_names(code, lang) | (known if x.get("type") == "subprocess" else set()), known):
                        rp.fail(f"{aid} {x['id']}「{x.get('label', '')[:24]}」: プログラムの中の名前「{tok}」が、そのまま出ている（読むのは、業務部門の担当。業務の言葉に直し、"
                                f"言い換えを {T.FILE} に書く。担当の目に見える名前＝シート・ファイル・メッセージなら「」で囲む。呼び出しの箱は、行き先の図に名前を付ける）")
                    for q in T.quoted_inner(said, code, lang, mine, unknown):
                        rp.fail(f"{aid} {x['id']}「{x.get('label', '')[:24]}」: かぎかっこの中の「{q}」は、プログラムの中だけの名前（変数・項目・手続き）。業務の言葉に直す"
                                f"（意味が読めなければ、{T.FILE} に basis: unknown で載せて、そのまま書く）")
                if x.get("unnamed"):
                    rp.fail(f"{aid} {x['id']}「{x.get('label', '')[:24]}」: 名前が付いていない（直しの一覧の name で、業務の言葉の名前を付ける）")
                else:
                    check_label_form(rp, f"{aid} {x['id']}", x)
                if TOBE.search(str(x.get("label", "")) + str(x.get("detail", ""))):
                    rp.warn(f"{aid} {x['id']}: 改善案らしい語がある「{x.get('label', '')[:24]}」。アクティビティ図には、いまの処理だけを書く")
                jargon = TECH_WORD.search(re.sub(r"「[^」]*」", "", str(x.get("label", "")) + " " + str(x.get("detail", ""))))   # 「」の中は、担当の目に見える名前
                if jargon:
                    rp.warn(f"{aid} {x['id']}「{x.get('label', '')[:24]}」: 技術の用語「{jargon.group(0)}」がある。何のためのものかを、業務の言葉で書く"
                            "（「読み書きの結果のコード」。code-guide の名付け）")
                src = norm("".join(span_text(page_units(cx.pj(s["doc_id"], s["page"])), s.get("lines")) or "" for s in x.get("sources") or [] if (s.get("doc_id"), s.get("page")) in cx.pages))
                held = norm(" ".join(f.get("label", "") for f in a["nodes"] if f.get("attach_to") == x["id"]))
                for q in re.findall(r"「([^」]+)」", str(x.get("label", "")) + str(x.get("detail", ""))):
                    if norm(q) not in src and norm(q) not in held:
                        rp.fail(f"{aid} {x['id']}「{x.get('label', '')[:24]}」: かぎかっこの中の名前「{q}」が、指したソースの行にも、付いている帳票の名前にも無い"
                                f"（シート・ファイル・テーブル・メッセージの名前は、ソースの文字のまま書く）")
                if x.get("type") == "subprocess" and not (adir / f"{x.get('link_flow')}.activity.json").exists():
                    rp.fail(f"{aid} {x['id']}: 呼び出しの行き先のアクティビティ図（{x.get('link_flow')}）が無い")
                for c in x.get("calls") or []:
                    if not (adir / f"{c.get('to')}.activity.json").exists():
                        rp.fail(f"{aid} {x['id']}: まとめた呼び出しの行き先のアクティビティ図（{c.get('to')}）が無い")
            if len(boxes) > MAX_BOXES:
                rp.warn(f"{aid}: 箱が {len(boxes)} 個ある（{MAX_BOXES} 個まで）。技術的な分岐を畳む（fold）・隣り合うまとまりをまとめる（merge）")
            n.update(acts=1, boxes=len(boxes), lines=len(need), off=sum(len(o["lines"]) for o in a.get("lines") or [] if o.get("state") == "excluded"),
                     folded=sum(b.get("state") == "folded" for b in logged.values()), forks=len(forks))
    known = {u["activity"] for u in units} | {u["activity"] for u in led.units if u["kind"] == "shared" and led.off(u) is not None}   # 図にしない手続きの残りは、上で知らせた
    for f in sorted(adir.glob("*.activity.json")) if adir.exists() else []:
        if f.name.split(".")[0] not in known:
            rp.fail(f"activities/{f.name}: 図にする単位に無い番号の図が残っている（台帳・構造が変わった。消すか、python3 scripts/activity.py list で番号を確かめる）")
    if tm.rows:
        rp.info(f"用語の対応 {len(tm.rows)} 件（コメント・定義から {sum(t.get('basis') in ('comment', 'definition') for t in tm.rows)}／"
                f"つづりからの推定 {sum(t.get('basis') == 'spelling' for t in tm.rows)}／不明 {sum(t.get('basis') == 'unknown' for t in tm.rows)}）")
    rp.info(f"アクティビティ図 {n['acts']}/{len(units)} 枚・箱 {n['boxes']}／行の網羅: 実行文 {n['lines']} 行（うち対象外 {n['off']} 行）／"
            f"分岐・繰り返し {n['forks']} か所（うち畳んだもの {n['folded']}）")


# ---------- 業務マトリクス一覧 ----------

def check_review(cx, rp):
    """原典照合（工程 8）: 全部の図に採否表がある／形・全部の指摘の採否／照合したあとに変わった箱が無い（照合に出した版の記録と、いまの図）／
    照合の報告と同じ／採用した指摘の反映／支持率"""
    import review as RV
    figs = RV.figures(cx.work)
    try:
        rec = RV.read_record(cx.work)
    except ValueError as e:
        rp.fail(f"{e}。消さずに、ファイルを前の版に戻す（照合に出した版の記録は、作り直せない）")
        rec = []
    rdir = cx.work / "review"
    have = {f.name.split(".")[0]: f for f in sorted(rdir.glob("*.review.json"))} if rdir.exists() else {}
    for F in sorted(set(figs) - set(have)):
        rp.fail(f"{F}: 原典照合の採否表（review/{F}.review.json）が無い（照合シートで照合させ、python3 scripts/review.py draft -w <作業フォルダ> で下書きを作る。"
                f"references/{'code-review-prompt' if F.startswith('A-') else 'source-review-prompt'}.md）")
    reviews = {}
    for F, f in have.items():
        name = f"review/{f.name}"
        if F not in figs:
            rp.fail(f"{name}: 業務フローにもアクティビティ図にも無い番号の採否表（図の番号を変えたなら、照合し直す。要らなければ消す）")
            continue
        rv = read_checked(rp, f, name, "")
        if rv is None:
            continue
        reviews[F] = rv
        open_ = [x for x in rv.get("findings") or [] if isinstance(x, dict) and (x.get("decision") not in ("採用", "不採用", "保留") or blank(x.get("action")))]
        shape_fails(rp, name, [e for e in validate(rv, "review") if not (open_ and re.search(r"\.(decision|action)", e))], STRUCT_ERRORS,
                    f"python3 scripts/validate.py --file review {name}")
        if open_:
            rp.fail(f"{F}: 採否・対応の書かれていない指摘 {len(open_)} 件（{'・'.join(str(x.get('id')) for x in open_[:8])}）。"
                    "指摘ごとに decision（採用／不採用／保留）と action（何を直したか・採らない理由）を書く")
        if rv.get("flow_id") not in (None, F):
            rp.fail(f"{name}: flow_id が {rv.get('flow_id')}（ファイル名は {F}）")
        held = [x for x in rv.get("findings") or [] if isinstance(x, dict) and x.get("decision") == "保留"]
        if held:
            rp.warn(f"{F}: 保留の指摘 {len(held)} 件（原典の読みが割れるなら、確認事項に足したか）")
        adopted = any(isinstance(x, dict) and x.get("decision") == "採用" for x in rv.get("findings") or [])
        act = F.startswith("A-")
        data = cx.work / ("activities" if act else "flows") / f"{F}.{'activity' if act else 'flow'}.json"
        try:                                                         # 報告と、照合に出した版（報告の sheet）。崩れていれば、この図だけ止めて次の図へ
            reps = RV.reports(cx.work, F)
            maps = RV.mapped(F, reps, rec) if reps else None
            want = RV.from_reports(cx.work, F, figs[F], rec, reps, maps) if reps else None
        except ValueError as e:
            rp.fail(f"{F}: {e}")
            continue
        if want is None:
            rp.fail(f"{F}: 照合の報告（reports/）が無い。採否表の判定と指摘は、作り方を知らない担当の報告から機械が写す（親が書いた採否表は、照合の代わりにならない）。"
                    f"照合させてから python3 scripts/review.py draft -w <作業フォルダ> --only {F}")
            continue
        got_v = {(v.get("node"), v.get("verdict"), v.get("evidence")) for v in rv.get("verdicts") or [] if isinstance(v, dict)}
        got_f = collections.Counter((RV.key(x), x.get("sheet")) for x in rv.get("findings") or [] if isinstance(x, dict))
        if got_v != {(v["node"], v["verdict"], v["evidence"]) for v in want[0]} or got_f != collections.Counter((RV.key(x), x.get("sheet")) for x in want[1]) \
                or (maps is not None and rv.get("sheets") != {f.name: e["id"] for f, _d, e in maps}):
            rp.fail(f"{F}: 採否表の判定・指摘が、照合の報告（reports/）と違う（写し違い・言い換え・報告が増えた・変わった）。"
                    f"python3 scripts/review.py draft -w <作業フォルダ> --only {F} で写し直す（採否は引き継ぐ）")
        if maps is None and not RV.legacy(rv, F, reps, rec):
            rp.fail(f"{F}: 照合の報告に、照合したシートの番号（sheet）が無い。照合シートの見出しの［S-<n>］を、報告の sheet に写させる"
                    "（アクティビティ図は、python3 scripts/review.py record で控えた番号。照合に出す前からあった前の版の報告には、書き足さない）")
        elif maps is None:                                           # 前の版の採否表: 前のとおり、番号と、ファイルの日時で見る
            rp.warn(f"{F}: 照合に出した版の記録が無い（前の版の採否表。報告に sheet が無い）。照合のあとに図を直していても、気づけない。"
                    f"前の版の報告のまま、いまの図を照合済みと認めるなら、"
                    f"python3 scripts/review.py snapshot -w <作業フォルダ> --only {F} --assume-current（利用者が決める）。認めないなら、全体を照合し直す")
            ids = {n.get("id") for n in figs[F].get("nodes") or [] if isinstance(n, dict)}
            gone = [str(v.get("node")) for v in rv.get("verdicts") or [] if isinstance(v, dict) and v.get("node") not in ids]
            if gone and not adopted:
                rp.fail(f"{F}: 判定したノード（{'・'.join(gone[:8])}）が図に無い。採用した指摘が無いのに番号が変わっている（照合のあとに番号を振り直した・別の図の判定を写した疑い）")
            if adopted and data.exists() and data.stat().st_mtime < f.stat().st_mtime:
                rp.fail(f"{F}: 採用した指摘があるのに、図のデータ（{data.parent.name}/{data.name}）が採否表より古い（採用した指摘を、まだ図に反映していない）")
        else:
            gaps = RV.unreviewed(F, figs[F], maps)
            if gaps:
                ids_ = ",".join(i for i, _w in gaps)
                rp.fail(f"{F}: 照合したあとに変わった・足した箱 {len(gaps)} 個（{'／'.join(f'{i} {w}' for i, w in gaps[:8])}{' …' if len(gaps) > 8 else ''}）。"
                        "採用した指摘の直しも含めて、その箱だけを照合し直す（差分照合）: "
                        + (f"python3 scripts/review.py record -w <作業フォルダ> --only {F} --nodes {ids_} → references/code-review-prompt.md の差分照合で、"
                           f"報告を reports/code_review_{F}_差分<N>.json に書かせ" if act else
                           f"python3 scripts/make_compare.py -w <作業フォルダ> --only {F} --nodes {ids_} → references/source-review-prompt.md の差分照合で、"
                           f"報告を reports/review_{F}_差分<N>.json に書かせ")
                        + f" → python3 scripts/review.py draft -w <作業フォルダ> --only {F}")
            at, now, cur = {e["id"]: e for e in rec if e["figure"] == F}, RV.figure_mark(figs[F]), RV.summaries(figs[F])

            def kept(x):
                """採用した指摘を、まだ反映していないか。箱を指す指摘は、その箱がそのシートのときのまま（消していない）。ページ・線・図全体を指す指摘は、図のデータがそのときのまま"""
                e = at.get(x.get("sheet"))
                if e is None or e["kind"] == RV.ASSUMED:
                    return False
                t = str(x.get("target") or "")
                return cur.get(t) == e["nodes"][t] if t in e["nodes"] else e["mark"] == now
            same = [str(x.get("id")) for x in rv.get("findings") or [] if isinstance(x, dict) and x.get("decision") == "採用" and kept(x)]
            if same:
                rp.fail(f"{F}: 採用した指摘（{'・'.join(same[:8])}）があるのに、指摘の箱（箱でない指摘は、図のデータ {data.parent.name}/{data.name}）が、"
                        "その指摘を受けたシートのときのまま（採用した指摘を、まだ図に反映していない）")
            said = sorted({f"{e['id']}（{e['date']}）" for _f, _d, e in maps if e["kind"] == RV.ASSUMED})
            if said:
                rp.info(f"{F}: 利用者が照合済みと認めた版 {'・'.join(said)} の上で点検した（review.py snapshot --assume-current）")
    for kind, s in RV.tally(figs, reviews).items():
        if not s["figs"] or not s["judged"]:
            continue
        rate, name = s["tot"]["支持"] / s["nodes"] if s["nodes"] else 1, {"flow": "業務フロー", "activity": "アクティビティ図"}[kind]
        line = f"{name}の支持率 {rate:.0%}（支持 {s['tot']['支持']} ÷ 開始・終了を除く全ノード {s['nodes']}。判定 {s['judged']}）"
        NUMBERS.setdefault("review", {})[kind] = {"figures": s["figs"], "reviewed": s["figs"] - len(s["rest"]), "nodes": s["nodes"], "judged": s["judged"],
                                                  "support": s["tot"]["支持"], "partial": s["tot"]["一部支持"], "against": s["tot"]["不支持"]}
        if rate < REVIEW_MIN:
            rp.fail(f"{line} が {REVIEW_MIN:.0%} 未満。不支持・一部支持の多い図を直し、照合し直す")
        else:
            rp.info(line)


def check_reports(cx, rp):
    """書き出した報告書（確認事項一覧・出典対応表・階層と業務分類・採否表）が、いまのデータから作り直したものと同じか。
    確認事項一覧の xlsx は、機械が埋める欄（様式にある欄）の全部が、いまのデータから作った行と同じか（人が埋める欄は別に引き継ぐ）"""
    import build_reports as BR
    r = BR.Reports(cx.work, write=False)
    r.all()
    again = "python3 scripts/build_reports.py -w <作業フォルダ> で書き出し直す"
    out = cx.work / "out"
    for name, text in r.made.items():
        p = out / name
        if not p.exists():
            rp.fail(f"out/{name} が無い。{again}")
            continue
        have = p.read_text(encoding="utf-8-sig" if name.endswith(".csv") else "utf-8")
        if have.replace("\r\n", "\n") != text.replace("\r\n", "\n"):
            rp.fail(f"out/{name} が、いまのデータから作り直したものと違う（データを直したあと書き出していない・手で直した）。{again}")
    if (out / "原典照合レビュー_採否表.md").exists() and "原典照合レビュー_採否表.md" not in r.made:
        rp.fail(f"out/原典照合レビュー_採否表.md があるのに、採否表（review/）が無い（古い報告書。消すか、照合をやり直す）")
    book = out / BR.Q_BOOK
    try:
        import table_template as tt
        back = tt.read_back(book, BR.Q_SPEC)
    except ImportError:
        rp.warn("openpyxl が無いので、確認事項一覧.xlsx を点検していない")
        return
    except Exception as e:
        rp.fail(f"out/{BR.Q_BOOK} を読めない（{type(e).__name__}: {str(e)[:60]}）。{again}")
        return
    rows = back["rows"][BR.Q_SHEET]
    cols = [k for k, f in enumerate(BR.Q_HEAD) if f not in back["absent"].get(BR.Q_SHEET, set())]
    want = [r.question_row(q) for q in sorted(r.qs, key=lambda q: (BR.SEV_ORDER.get(q["severity"], 1), q["id"]))]
    got_rows, want_rows = [[x[k] for k in cols] for x in rows[1:]], [[tt.clean(str(x[k])) for k in cols] for x in want]
    if [x[0] for x in got_rows] != [x[0] for x in want_rows]:
        rp.fail(f"out/{BR.Q_BOOK}: 番号の並びが、いまの確認事項（questions.json）と違う（{len(rows) - 1} 行／いま {len(want)} 件）。{again}")
    elif got_rows != want_rows:
        bad = [(g[0], BR.Q_HEAD[cols[k]]) for g, w in zip(got_rows, want_rows) for k in range(len(cols)) if g[k] != w[k]]
        rp.fail(f"out/{BR.Q_BOOK}: いまのデータと違うセル {len(bad)} 個（{'・'.join(f'{q}「{f}」' for q, f in bad[:6])}）。手で直した・書き出し直していない。{again}")
    rp.info(f"報告書 {len(r.made)} ファイル・確認事項一覧 {len(want)} 件（高 {sum(q['severity'] == 'high' for q in r.qs)}・中 {sum(q['severity'] == 'mid' for q in r.qs)}・低 {sum(q['severity'] == 'low' for q in r.qs)}）")
    try:                                                             # 回答（人が埋める欄）の引き継ぎ
        ans = Answers(cx.work)
    except BadJson as e:
        rp.fail(f"人が埋める欄の控え（answers.json）を読めない — {e}")
        return
    now = {q["id"]: q["question"] for q in r.qs}
    gone, changed = answer_gaps(ans, BR.Q_BOOK, BR.Q_SHEET, now)
    some = lambda xs: "・".join(xs[:8]) + (f" ほか {len(xs) - 8}" if len(xs) > 8 else "")
    if gone:
        rp.warn(f"確認事項に無くなった番号の回答 {len(gone)} 件（{some(gone)}）。answers.json に残してある（図に反映して確認事項が消えたのなら、そのままでよい）")
    if changed:
        rp.warn(f"回答を取り込んだあとで、質問の文が変わった確認事項 {len(changed)} 件（{some(changed)}）。回答が、いまの問いに合っているか見直す")
    answered = sorted(k for k in ans.rows(BR.Q_BOOK, BR.Q_SHEET) if k in now)
    NUMBERS["questions"] = {"all": len(r.qs), **{k: sum(q["severity"] == k for q in r.qs) for k in ("high", "mid", "low")}, "answered_open": len(answered)}
    if answered:
        rp.warn(f"回答済みなのに、まだ確認事項に残っている {len(answered)} 件（{some(answered)}）。回答を図に反映して、確認事項を消す（反映しないまま渡すなら、そのことを添える）")
    try:
        got, _notes = tt.manual_back(book, BR.Q_SPEC, ASSETS / BR.Q_TEMPLATE)
    except Exception:                                                # 読めない・列を動かした本は、書き出しが止まって知らせる
        got = {}
    have = ans.kept(BR.Q_BOOK).get(BR.Q_SHEET, {})
    new = [k for k, cells in (got.get(BR.Q_SHEET) or {}).items() if cells != have.get(k, {})]
    if new:
        rp.info(f"{BR.Q_BOOK} に、まだ取り込んでいない人の値 {len(new)} 行（{some(new)}）。書き出し直すと、answers.json に取り込む")


INNER_ID = re.compile(r"(?<![A-Za-z0-9])(?:H\d+(?:\.\d+)+|SYS-\d{2,}|M\d+\.\d+\.[sdgn]\d+|p\d{4})(?![A-Za-z0-9])")   # 業務の木の id・システムの番号・カードの項目のキー・ファイルの番号
WORK_PATH = re.compile(r"(?:^|[\s（(「])(?:flows|cards|pages|structure|reports|review|activities|code)/[^\s）)」]+")


def deliverable_texts(path):
    """納品物 1 つの文字: [(どこ, 文字)]。draw.io はセルの値と属性、xlsx は全セルとコメント、md・csv は行"""
    out = []
    if path.suffix == ".drawio":
        for d in ET.parse(path).getroot().iter("diagram"):
            for c in d.iter():
                for k, v in c.attrib.items():
                    if k in ("value", "label", "tooltip", "src", "quote") and v:
                        out.append((f"［{d.get('name')}］{c.get('id', '')}", html.unescape(re.sub(r"<[^>]+>", " ", v))))
    elif path.suffix == ".xlsx":
        from openpyxl import load_workbook
        for ws in load_workbook(path):
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        out.append((f"「{ws.title}」{c.coordinate}", str(c.value)))
                    if c.comment is not None:
                        out.append((f"「{ws.title}」{c.coordinate} のコメント", str(c.comment.text)))
    else:
        section = ""
        for i, ln in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
            section = ln if ln.startswith("#") else section
            out.append((f"{i} 行目" + ("（資料の表）" if section == "## 資料" else ""), ln))
    return out


def check_deliverables(cx, rp):
    """お客様に渡す納品物の文字に、内部の番号・ファイルの場所・スキーマの項目名・下書きの印が出ていないか（出してよい番号は F・P・Q・A・S と、ノードの番号だけ）"""
    out = cx.work / "out"
    files = sorted({p for pat in DELIVERABLES for p in out.glob(pat)}) if out.exists() else []
    if not files:
        rp.fail("納品物（out/ の図・業務マトリクス一覧・確認事項一覧・出典対応表・階層と業務分類）が無い")
        return
    ids = {d["doc_id"] for d in cx.m["docs"]}
    paths = sorted({d.get("abs_path") for d in cx.m["docs"] if d.get("abs_path")} | set(cx.m.get("inputs") or []), key=len, reverse=True)
    rels = sorted({d["rel_path"] for d in cx.m["docs"] if "/" in str(d.get("rel_path", ""))}, key=len, reverse=True)
    titles = {cx.lab.title(d) for d in ids}
    fields = sorted({k for f in (ASSETS / "schemas").glob("*.schema.json") for k in re.findall(r'"([a-z]+_[a-z_]+)"\s*:', f.read_text(encoding="utf-8"))}, key=len, reverse=True)
    field_re = re.compile(r"(?<![A-Za-z0-9_])(" + "|".join(map(re.escape, fields)) + r")(?![A-Za-z0-9_])") if fields else None
    n = 0
    for p in files:
        try:
            texts = deliverable_texts(p)
        except Exception as e:
            rp.fail(f"{p.name}: 読めない（{type(e).__name__}: {str(e)[:60]}）")
            continue
        n += len(texts)
        hits = collections.OrderedDict()
        for where, t in texts:
            exempt = "（資料の表）" in where or "「資料一覧」" in where
            found = [] if exempt else [m.group(0) for m in DOC_REF.finditer(t) if m.group(0) in ids]   # 資料 ID は、続く文字を問わない（「M01 の」も）
            found += [f"内部の番号 {m.group(0)}" for m in INNER_ID.finditer(t)]
            found += [f"作業フォルダの場所 {m.group(0).strip()}" for m in WORK_PATH.finditer(t)]
            found += [f"ファイルの場所 {x}" for x in paths if x and x in t]
            found += [f"ファイルの場所 {x}" for x in rels if not exempt and x in t and not any(x in y for y in titles)]   # 題名が重なる資料は、題名に場所を添えて見分ける
            found += [f"項目名 {m.group(1)}" for m in field_re.finditer(t)] if field_re else []
            found += ["【下書き】"] if "【下書き】" in t else []
            for x in found:
                hits.setdefault(x, []).append(where)
        for x, at in hits.items():
            rp.fail(f"{p.name}: 納品物に「{x[:60]}」が出ている（{'・'.join(at[:3])}{' ほか ' + str(len(at) - 3) if len(at) > 3 else ''}）。"
                    + ("資料は題名で書く" if re.fullmatch(r"M\d+", x) else "図にしない（作り直すか、名前の付いていない箱に名前を付ける）" if x == "【下書き】"
                       else "内部の番号・場所・項目名は、納品物に出さない"))
    rp.info(f"納品物 {len(files)} ファイル・文字 {n} か所を走査した")


def matrix_times(rp, where, F, fl, pairs, row, times, table, seen, again):
    """処理時間の行（作業の行のすぐ下）を、ブックだけを読んで確かめる。pairs は [(作業のセル, その真下のセル)]。返すのは [(ノード番号, 時間のセルの番地)]。
    真下のセルが、その作業の時間のセル（コメントの 1 行目の札）か／中身が、数（人が入れた値）か、基準の表を引く式（動作の種類・量が、いまのフローのとおり）か／
    合計の式が、その行の時間のセルを足しているか。ブックと times.json の値が違う所は、数えて WARN（人がブックを直して、まだ書き出し直していない）"""
    from build_matrix import MATRIX_HEAD, OURS, TAG_COUNT, TAG_TIME, number
    node = {x.get("id"): x for x in fl.get("nodes") or [] if isinstance(x, dict)}
    held = seen_files(fl)
    note = lambda c: str(c.comment.text) if c is not None and c.comment is not None else ""
    differs = lambda x, mine: mine is None or abs(float(x) - float(mine)) > 0.005

    def human(c, what, mine):
        """人が埋めるセルの中身を数える。数として読めない字は WARN"""
        v = c.value
        if v is None or (isinstance(v, str) and not v.strip()):
            seen["pending"] += mine is not None
            return
        if isinstance(v, str) and v.startswith("="):                 # 人が書いた式（計算の結果を、書き出し直しで取り込む）
            seen[what] += 1
            seen["pending"] += mine is None
            return
        x = number(v)
        if x is None:
            rp.warn(f"{where}: {c.coordinate} に、数として読めない値がある「{str(v)[:20]}」（取り込めない。数だけを入れる）")
            return
        seen[what] += 1
        seen["pending"] += differs(x, mine)

    out, bare = [], []
    for above, c in pairs:
        i = note(above).split("（", 1)[0]
        m = TAG_TIME.match(note(c))
        if not note(c):
            bare.append(c.coordinate)
        elif m is None or (m.group(1), m.group(2)) != (F, i):
            rp.fail(f"{where}: 作業 {i} の真下のセル（{c.coordinate}）が、その作業の時間のセルでない（コメントの札が「処理時間 {F} {i}｜」で始まらない。行・列を動かした）。{again}")
            continue
        out.append((i, c.coordinate))
        v, mine = c.value, times.cells.get((F, i), {}).get("value")
        if isinstance(v, str) and OURS.match(v):                     # 仮の値（基準の表を引く式）
            act, vol, _why = time_basis(node.get(i) or {}, held[i])
            plain = lambda f: f.replace("'", "")                        # Excel は、保存のとき、要らない引用符を外す（'処理時間の基準'!$A$2 → 処理時間の基準!$A$2）
            col = table is not None and (table["cols"].get(vol) or table["cols"][VOLUME_SPARE[vol]])
            if table is not None and plain(v) != plain("=" + soft_formula(act, table["ref"], col)):
                rp.fail(f"{where}: 作業 {i} の仮の値の式が、基準の表の「{act} × {vol}」を引いていない（{v[:70]}）。{again}")
            seen["soft"] += 1
            seen["pending"] += mine is not None
        elif v is None or (isinstance(v, str) and not v.strip()):    # 空（書き出し直すと、仮の値に戻る）
            seen["pending"] += 1
        else:
            human(c, "human", mine)
    total, count = row[MATRIX_HEAD.index("処理時間の合計")], row[MATRIX_HEAD.index("件数")]
    if total is not None and out:
        want = f"=SUM({out[0][1]}:{out[-1][1]})"
        if str(total.value or "").replace("$", "") != want:
            rp.fail(f"{where}: 合計の式が、この行の時間のセルを足していない（ブック「{str(total.value)[:40]}」／あるはずの式「{want}」）。{again}")
    if count is not None:
        m = TAG_COUNT.match(note(count))
        if not note(count):
            bare.append(count.coordinate)
        elif m is None or m.group(1) != F:
            rp.fail(f"{where}: 件数のセル（{count.coordinate}）のコメントの札が、「件数 {F}｜」で始まらない（行を動かした）。{again}")
        human(count, "count", times.counts.get(F, {}).get("value"))
    if bare:
        rp.warn(f"{where}: 時間・件数のセル {len(bare)} 個のコメントが消えている（{'・'.join(bare[:6])}{' …' if len(bare) > 6 else ''}。だれの時間かの札）。"
                f"{again}（書き出したときの番地で突き合わせて、付け直す。行・列を動かしていなければ、値は残る）")
    return out


CELL_REF = re.compile(r"^(?:'((?:[^']|'')+)'!)?\$?([A-Z]{1,3})\$?(\d+)$")


def matrix_results(rp, book, again):
    """式のセルに保存した計算の結果（計算をしない表示・プレビューが見る値）が、ブックの中の値で式を計算し直したものと同じか。
    書き出す式は 5 つの形（基準の表を引く VLOOKUP・2 つを引いて大きいほうの MAX・合計の SUM・別のシートのセル・仮か人かの IF）。人が書き足した式は、計算せずに数える"""
    from openpyxl import load_workbook
    from openpyxl.utils import range_boundaries
    wf, wv = load_workbook(book), load_workbook(book, data_only=True)
    memo = {}

    def where(ref, here):
        m = CELL_REF.match(ref.strip())
        return ((m.group(1).replace("''", "'") if m.group(1) else here), f"{m.group(2)}{m.group(3)}") if m else None

    def area(ref, here):
        title, _s, rng = ref.rpartition("!")
        title = title.strip("'").replace("''", "'") if title else here
        c1, r1, c2, r2 = range_boundaries(rng.replace("$", ""))
        return [[(title, wf[title].cell(r, c).coordinate) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]

    def value(title, at):
        if (title, at) not in memo:
            v = wf[title][at].value
            memo[(title, at)] = calc(title, v) if isinstance(v, str) and v.startswith("=") else v
        return memo[(title, at)]

    def calc(here, f):
        m = re.fullmatch(r"=MAX\((.+)\)", f)
        if m:
            one = re.findall(r'VLOOKUP\("[^"]*",.+?,\d+,FALSE\)', m.group(1))
            if ",".join(one) != m.group(1):
                raise KeyError(f)
            got = [calc(here, "=" + x) for x in one]
            return max(got) if all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in got) else None
        m = re.fullmatch(r'=VLOOKUP\("([^"]*)",(.+),(\d+),FALSE\)', f)
        if m:
            row = next((r for r in area(m.group(2), here) if value(*r[0]) == m.group(1)), None)
            return value(*row[int(m.group(3)) - 1]) if row else None
        m = re.fullmatch(r"=SUM\(([^)]+)\)", f)
        if m:
            return sum(v for r in area(m.group(1), here) for v in [value(*x) for x in r] if isinstance(v, (int, float)) and not isinstance(v, bool))
        m = re.fullmatch(r'=IF\(_xlfn\.ISFORMULA\((.+)\),"仮","人"\)', f)
        if m and where(m.group(1), here):
            v = wf[where(m.group(1), here)[0]][where(m.group(1), here)[1]].value
            return "仮" if isinstance(v, str) and v.startswith("=") else "人"
        if where(f[1:], here):
            return value(*where(f[1:], here))
        raise KeyError(f)
    bad, missing, other = [], 0, 0
    for ws in wf:
        for row in ws.iter_rows():
            for c in row:
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    continue
                try:
                    want = calc(ws.title, c.value)
                except KeyError:
                    other += 1
                    continue
                kept = wv[ws.title][c.coordinate].value
                if kept is None and want not in (None, ""):
                    missing += 1
                elif kept != want and not (isinstance(kept, (int, float)) and isinstance(want, (int, float)) and abs(kept - want) < 1e-6):
                    bad.append(f"「{ws.title}」{c.coordinate}（保存 {kept!r}／計算 {want!r}）")
    for x in bad[:5]:
        rp.fail(f"{book.name}: 式のセルに保存した計算の結果が、式を計算し直した値と違う: {x}。{again}")
    if len(bad) > 5:
        rp.fail(f"{book.name}: 保存した計算の結果が違う式のセルが、ほかに {len(bad) - 5} 個。{again}")
    if missing:
        rp.warn(f"{book.name}: 計算の結果が保存されていない式のセル {missing} 個（計算をしない表示・プレビューでは 0・空に見える。{again}すか、Excel で開いて保存する）")
    if other:
        rp.info(f"{book.name}: 人が書き足した式 {other} 個（計算の結果は確かめていない）")


def matrix_links(rp, book, flow_ids, again):
    """業務マトリクスの行の名前から、フロー図の PNG へのリンク: リンク先があるか。PNG があるのにリンクの無いフローが無いか（PNG を作る前に書き出した）"""
    from openpyxl import load_workbook
    targets = [c.hyperlink.target for ws in load_workbook(book) for row in ws.iter_rows() for c in row if c.hyperlink is not None and c.hyperlink.target]
    for t in targets:
        if t.startswith("png/") and not (book.parent / t).exists():
            rp.fail(f"{book.name}: フロー図へのリンクの先（{t}）が無い。python3 scripts/check_layout.py -w <作業フォルダ> で PNG を作ってから、{again}")
    png = book.parent / "png"
    have = {re.sub(r".*_flow-", "", p.stem) for p in png.glob("*_flow-*.png")} if png.exists() else set()
    linked = {re.sub(r".*_flow-", "", pathlib.PurePosixPath(t).stem) for t in targets if t.startswith("png/")}
    lost = sorted((have & set(flow_ids)) - linked)
    if lost:
        rp.fail(f"{book.name}: PNG があるのに、フロー図へのリンクの無い行 {len(lost)}（{'・'.join(lost[:8])}）。PNG を作る前に書き出した。{again}")


def check_matrix(cx, rp):
    """書き出した業務マトリクス一覧（xlsx）が、いまのデータのとおりか。
    (1) データから組み立て直した中身と、ブックの読み戻しを、全セルで比べる（フローを直したあと書き出し直していない・手で直した）
    (2) 組み立てとは別に、ブックだけを読んで確かめる（書き出しの誤りを、同じ部品で見逃さないように）:
        行がフローを描く分類の順／作業のノードが、どれも1回だけセルにある／業務フローのノード一覧の「行き先」から組み直した線が、フローの線と同じ／
        そのフローの確認事項の番号が、すべて行に出ている／処理時間の行（作業の行のすぐ下）の、札・式・合計（matrix_times）
    人が埋めるセル（時間・件数）は、(1) で比べない。人が入れた値は、FAIL にしない（書き出し直すと、times.json に取り込む）"""
    try:
        import openpyxl  # noqa: F401
        import table_template as tt
        from build_matrix import BOOK, EXIT_RE, MATRIX_HEAD, NODE_HEAD, NODES, Q_RE, SPEC, Matrix, calibration
    except ImportError:
        rp.warn("openpyxl が無いので、業務マトリクス一覧を点検していない（pip3 install openpyxl）")
        return
    book = cx.work / "out" / BOOK
    if not book.exists():
        rp.fail(f"out/{BOOK} が無い（python3 scripts/build_matrix.py -w <作業フォルダ> を実行する）")
        return
    again = "python3 scripts/build_matrix.py -w <作業フォルダ> で書き出し直す"
    try:                                                             # 人が入れた時間の控え。読めなければ、組み立て直しもできないので、ここで止める
        times = Times(cx.work)
    except BadJson as e:
        rp.fail(f"人が入れた時間の控え（times.json）を読めない — {e}")
        return
    try:
        back = tt.read_back(book, SPEC)                                # テンプレートが決めた列の順を、欄の決まった順に戻して読む
    except tt.TemplateError as e:
        rp.fail(f"out/{BOOK}: {e}。{again}")
        return
    except Exception as e:                                           # 壊れた xlsx は、どの例外で落ちるか決まっていない
        rp.fail(f"out/{BOOK} を読めない（{type(e).__name__}: {str(e)[:60]}）。build_matrix.py で書き出し直す"
                "（人が入れた処理時間・件数は、読めないブックからは取り込めない。取り込まずに作り直すなら --fresh）")
        return
    got = back["rows"]
    mx = Matrix(cx.work, False)
    mx.timed = timed = bool(back["side"].get("timed"))               # 処理時間の行のあるブックか（書き出したときの、テンプレートで決まる）
    want, diffs = mx.texts(), 0
    lost = [n for n in SPEC["required"] if n not in got]
    if lost:
        rp.fail(f"{BOOK}: 一覧「{'・'.join(lost)}」のシートが無い（テンプレートに要る）。{again}")
        return
    for name in got:
        a, b = got[name], want[name]
        head, skip = b[0] if b else [], {j for j, f in enumerate(SPEC["sheets"][name]) if f in back["absent"][name]}
        if len(a) != len(b):
            rp.fail(f"{BOOK}「{back['titles'][name]}」: 行の数が、いまのデータと違う（ブック {len(a) - 1} 行・データ {len(b) - 1} 行）。{again}")
            continue
        for i, (ra, rb) in enumerate(zip(a, b)):
            for j in range(max(len(ra), len(rb))):
                x, y = (ra[j] if j < len(ra) else ""), (rb[j] if j < len(rb) else "")
                if y == tt.EDIT or (y == tt.FORMULA and x.startswith("=")):      # 人が埋めるセル（時間・件数）・式のセル。中身は、下で、ブックだけを読んで確かめる
                    continue
                if x != y and j not in skip:
                    diffs += 1
                    if diffs <= 5:
                        col = head[j] if j < len(head) else str(j + 1)
                        at = next((n for n, (p_, q_) in enumerate(zip(x, y)) if p_ != q_), min(len(x), len(y)))      # 初めて食い違う所の前後を見せる
                        cut = lambda t: ("…" if at > 12 else "") + t[max(0, at - 12):at + 28].replace("\n", "⏎")
                        rp.fail(f"{BOOK}「{back['titles'][name]}」のデータの {i} 行目「{col}」が、いまのデータと違う: ブック「{cut(x)}」／データ「{cut(y)}」。{again}")
    values = mx.values()
    for title, cells in back["static"].items():                      # 文のシート（読み方）: 数を入れた所だけを比べる（文は、テンプレートのもの）
        for at, (x, src) in cells.items():
            y = tt.TOKEN.sub(lambda m: str(values.get(m.group(1), m.group(0))), src)
            if x != y:
                diffs += 1
                if diffs <= 5:
                    rp.fail(f"{BOOK}「{title}」{at} の数が、いまのデータと違う: ブック「{x[:40]}」／データ「{y[:40]}」。{again}")
    if diffs > 5:
        rp.fail(f"{BOOK}: いまのデータと違うセルが、ほかに {diffs - 5} 個ある。{again}")
    matrix_results(rp, book, again)
    matrix_links(rp, book, [flow_key(fl) for fl in load_flows(cx.work, [])], again)
    for name in back["side"].get("empty") or []:                     # 中身が無くて、本に入れなかった一覧（プログラム用のシート）: いまは中身があるなら、書き出し直していない
        if len(want.get(name) or []) > 1:
            rp.fail(f"{BOOK}: 一覧「{name}」は、書き出したときは中身が無く、本に入れていないが、いまは {len(want[name]) - 1} 行ある。{again}")
    # ここから、ブックだけを読んで確かめる
    flows = {flow_key(fl): fl for fl in load_flows(cx.work, [])}
    tax = try_load(cx.work / "taxonomy.json")[0]
    cats = [c for c in (tax.get("categories") if isinstance(tax, dict) else None) or [] if isinstance(c, dict)]
    kids, order, todo = {}, [], []
    for c in cats:
        kids.setdefault(c.get("parent") if any(x.get("id") == c.get("parent") for x in cats) else None, []).append(c.get("id"))
    todo = kids.get(None, [])[::-1]
    while todo:                                                      # 分類の木を上から（下位フローは親の直後）
        i = todo.pop()
        if i not in order:
            order.append(i)
            todo += kids.get(i, [])[::-1]
    expect = [i for i in order if i in flows] + [i for i in flows if i not in order]
    col, nt = {h: k for k, h in enumerate(NODE_HEAD)}, back["titles"][NODES]
    listed = {}
    for r in got[NODES][1:]:
        listed.setdefault(r[col["フローID"]], []).append(r)
    if list(listed) != expect:
        rp.fail(f"{BOOK}「{nt}」のフローの並びが、業務分類の順でない（ブック {'・'.join(list(listed)[:8])}／分類 {'・'.join(expect[:8])}）。{again}")
        return
    rows, sheet = got["業務マトリクス"][1:], back["cells"]["業務マトリクス"]
    per = 2 if timed else 1                                          # 処理時間の行があれば、1 フロー＝2 行（作業の行・時間の行）
    if len(rows) != per * len(expect):
        return                                                       # 行の数の違いは、上で FAIL にしている
    loose = [f"{F} {i}" for (F, i), c in sorted(times.cells.items()) if not c.get("book")] + [f"{F} の件数" for F, c in sorted(times.counts.items()) if not c.get("book")]
    if loose:                                                        # ブックから取り込んだ値には、取り込み元（book）が必ず付く
        rp.fail(f"times.json に、ブックから取り込んでいない値がある（{'・'.join(loose[:8])}{' …' if len(loose) > 8 else ''}）。処理時間・件数は、人が業務マトリクス一覧に入れた値だけ"
                "（LLM・手で times.json に書かない。flow-rules 3.22）。その行を消して、書き出し直す")
    seen, table = collections.Counter(), None
    if timed:
        table, bad = tt.time_table(back["book"], SPEC["time"])
        if table is None:
            rp.fail(f"{BOOK}: 処理時間の基準の表を読めない（{bad or '名前「' + SPEC['time']['name'] + '」の範囲が無い'}）。{again}")
        elif not back["side"].get("table_version"):                 # 版の記録の前の道具で書き出したブック（納品の門が、現状の処理時間の写しに版を書けない）
            rp.fail(f"{BOOK}: 基準の表の版の記録が無い（前の版の道具で書き出したブック）。{again}")
        else:                                                        # 人がブックで直した基準の表（書き出し直すと取り込み、版が変わる）
            put = {(a, v): x for a, v, x in back["side"].get("table") or []}
            seen["pending"] += sum(1 for k, x in table["values"].items() if k in put and abs(float(x) - float(put[k])) > 0.005)
    questions = (try_load(cx.work / "questions.json")[0] or {}).get("questions") or []
    for k, F in enumerate(expect):
        fl, r = flows[F], rows[per * k]
        where = f"{BOOK}「業務マトリクス」{per * k + 2} 行目（{F} {fl.get('name', '')}）"
        if (r[0], r[1]) != (listed[F][0][0], listed[F][0][1]):
            rp.fail(f"{where}: 大分類・中分類が「{nt}」と違う。{again}")
        steps = [n.get("id") for n in fl.get("nodes") or [] if isinstance(n, dict) and n.get("type") in STEP_TYPES]
        cells = [c for c in sheet[per * k][len(MATRIX_HEAD):] if c is not None and c.value is not None and str(c.value) != ""]
        ids = [str(c.comment.text).split("（", 1)[0] if c.comment is not None else "" for c in cells]
        if sorted(ids) != sorted(steps):
            lost, extra = [i for i in steps if i not in ids], [i for i in ids if i not in steps or ids.count(i) > 1]
            rp.fail(f"{where}: 作業のセルが、フローの作業のノードと合わない（セルに無い {'・'.join(lost[:6]) or 'なし'}／"
                    f"余分・重なり {'・'.join(dict.fromkeys(extra)) or 'なし'}）。{again}")
        lines = [(v[col["ノード番号"]], ln) for v in listed[F] for ln in v[col["行き先"]].split("\n") if ln]
        parsed = [(i, EXIT_RE.match(ln)) for i, ln in lines]
        drawn = sorted((i, mt["to"], mt["label"] or "", "（推定）" in mt["marks"]) for i, mt in parsed if mt)
        real = sorted((str(e.get("from")), str(e.get("to")), e.get("label") or "", bool(e.get("inferred"))) for e in fl.get("edges") or [] if isinstance(e, dict))
        if drawn != real or len(drawn) != len(lines):
            gone = [f"{a_}→{b_}" for a_, b_, _l, _i in real if not any(x[:2] == (a_, b_) for x in drawn)]
            rp.fail(f"{where}: 「{nt}」の行き先から組み直した線が、フローの線と合わない"
                    + (f"（行き先に無い線 {'・'.join(gone[:6])}）" if gone else "（ラベル・推定の印・本数の違い）") + f"。{again}")
        asked = {q.get("id") for q in questions if isinstance(q, dict) and q.get("flow") == F}
        shown = set(Q_RE.findall(" ".join(r)))
        if questions and asked != shown:
            rp.fail(f"{where}: 行に出ている確認事項の番号が、このフローの確認事項と合わない（行に無い {'・'.join(sorted(asked - shown)[:6]) or 'なし'}／"
                    f"余分 {'・'.join(sorted(shown - asked)[:6]) or 'なし'}）。{again}")
        if timed:
            below = sheet[per * k + 1]
            pairs = [(a_, b_) for a_, b_ in zip(sheet[per * k][len(MATRIX_HEAD):], below[len(MATRIX_HEAD):]) if a_ is not None and a_.value not in (None, "") and b_ is not None]
            at = dict(matrix_times(rp, f"{BOOK}「業務マトリクス」{per * k + 3} 行目（{F} の処理時間）", F, fl, pairs, below, times, table, seen, again))
            ref = tt.quoted(back["titles"]["業務マトリクス"])
            for v in listed[F]:                                      # 業務フローのノード一覧は、業務マトリクスの時間のセルを指す（数を打つ場所は、1 か所）
                i = v[col["ノード番号"]]
                for name, want_f in (("処理時間（分）", f"={ref}!{at.get(i)}"), ("時間の根拠", f'=IF(_xlfn.ISFORMULA({ref}!{at.get(i)}),"仮","人")')):
                    if i in at and name not in back["absent"][NODES] and v[col[name]].replace("'", "") != want_f.replace("'", ""):   # Excel は、保存のとき、要らない引用符を外す
                        rp.fail(f"{BOOK}「{nt}」{F} {i} の「{name}」が、業務マトリクスの時間のセル（{at[i]}）を指していない（{v[col[name]][:50]}）。{again}")
    flow_rows = rows[::per]
    rp.info(f"業務マトリクス一覧 {len(flow_rows)} 行・作業 {sum(len([c for c in r[len(MATRIX_HEAD):] if c]) for r in flow_rows)} セル")
    try:                                                             # 様式で足した、人が埋める欄の引き継ぎ
        ans = Answers(cx.work)
    except BadJson as e:
        rp.fail(f"人が埋める欄の控え（answers.json）を読めない — {e}")
        ans = None
    if ans is not None:
        from build_matrix import M, answer_rows
        now = answer_rows(list(flows.values()))
        for canon in (M, NODES):
            gone, changed = answer_gaps(ans, BOOK, canon, now[canon])
            if gone:
                rp.warn(f"「{canon}」の人が埋める欄: 行が無くなった値 {len(gone)} 行（{'・'.join(gone[:6])}）。answers.json に残してある")
            if changed:
                rp.warn(f"「{canon}」の人が埋める欄: 取り込んだあとで、行の中身（フロー名・作業名）が変わった {len(changed)} 行（{'・'.join(changed[:6])}）。値を見直す")
    if not timed:
        return
    gone, renamed = times.gaps(list(flows.values()))
    for F, i, v, was, _now in gone:
        rp.warn(f"{F} {i}: 人が入れた処理時間（{v} 分）の箱「{was}」が、もう無い。値は times.json に残してある（要らなければ build_matrix.py --forget）")
    for F, i, v, was, now in renamed:
        rp.warn(f"{F} {i}: 人が入れた処理時間（{v} 分）の作業名が変わった（「{was}」→「{now}」）。値を見直す（そのままでよければ build_matrix.py --accept）")
    rp.info(f"処理時間: 人が入れた値 {seen['human']} セル・仮の値（基準の表を引く式）{seen['soft']} セル／件数 {seen['count']} 行")
    NUMBERS["times"] = {"human": seen["human"], "soft": seen["soft"], "counts": seen["count"]}
    steps = [(r["id"], s) for r in mx.rows for s in r["steps"]]
    some = lambda xs: "・".join(xs[:12]) + (f" ほか {len(xs) - 12}" if len(xs) > 12 else "")
    long = [f"{F} {s['id']}" for F, s in steps if s["time"][1] in VOLUME_SPARE and (F, s["id"]) not in times.cells]
    if long:
        rp.info(f"量が「特大」（時間・日の単位でかかる作業）で、人の値がまだ無いセルが {len(long)} 個（表の値が、とくに粗い。先に、実際の時間を入れる所）: {some(long)}")
    misc = [f"{F} {s['id']}" for F, s in steps if s["node"].get("type") == "process" and s["time"][0] == ACTIONS[-1] and (F, s["id"]) not in times.cells]
    if misc:
        rp.info(f"動作の種類が「{ACTIONS[-1]}」で、人の値がまだ無い作業が {len(misc)} 箱（仮の値は、「{ACTIONS[-1]}」の行の決めの値）: {some(misc)}")
    if seen["pending"]:                                              # 納品の門は、書き出し直して取り込んでから、現状の処理時間の写しを作る
        rp.warn(f"ブックに、まだ取り込んでいない人の値が {seen['pending']} 個（処理時間・件数・基準の表。人がブックを直したあと、まだ書き出し直していない。{again}と、取り込む）")
    for key, cnt, mid, cur in calibration(mx, (table or {}).get("values") or {}):
        if mid != cur:
            rp.info(f"基準の表「{key[0]} × {key[1]}」は {cur} 分・人が入れた {cnt} セルの中央値は {mid} 分（表を引き直すなら build_matrix.py --calibrate）")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-w", "--work", required=True)
    ap.add_argument("--stage", choices=["code", "systems", "activity", "hierarchy", "structure", "flow", "drawio", "matrix", "review", "reports", "deliverables", "all"], default="all")
    ap.add_argument("--only", help="フロー ID（カンマ区切り）。flow の点検で、そのフローの行だけを出す")
    ap.add_argument("--max-material", type=int, default=MATERIAL_CAP, help=f"1つの分類の材料（カードの文字数）の上限（既定 {MATERIAL_CAP}）")
    ap.add_argument("--report")
    ap.add_argument("--all-lines", action="store_true", help="同じ種類の行を畳まずに全部出す")
    ap.add_argument("--styles", default=str(ASSETS / "template_styles.json"))
    ap.add_argument("--numbers", help="数（網羅率・引用照合率・確認事項・支持率・処理時間）を JSON で書くファイル（deliver.py が使う）")
    a = ap.parse_args()
    work = work_dir(a.work)
    cx = Ctx(work)
    rp = Report("階層・業務分類・フロー・draw.io の点検（check_flow）", fold=None if a.all_lines else 10)
    only = {x.strip() for x in a.only.split(",") if x.strip()} if a.only else None
    has_code = any(is_code_page(p) for _d, p in all_pages(cx.m, code=True))
    if a.stage == "code" or (a.stage == "all" and has_code):
        with guarded(rp, "プログラムのソース"):
            check_code(cx, rp)
    if a.stage == "systems" or (a.stage == "all" and (has_code or (work / "structure" / "systems.json").exists())):
        with guarded(rp, "システム台帳"):
            check_systems(cx, rp)
    if a.stage == "activity" or (a.stage == "all" and ((work / "activities").exists() or (work / "structure" / "systems.json").exists())):   # 台帳があれば、図がまだ無いことも見る
        with guarded(rp, "アクティビティ図"):
            check_activity(cx, rp)
    manual_files = [x for x in ("structure/hierarchy.json", "glossary.json", "taxonomy.json") if (work / x).exists()]
    if a.stage == "all" and has_code and not (cx.active - cx.code_only) and not manual_files:
        rp.info("読む資料が、プログラムのソースだけ（マニュアルが無い）。階層・用語・業務分類・業務フローは作らない（アクティビティ図と、システム・起動点の一覧を出す）")
    elif a.stage in ("hierarchy", "structure", "all"):
        with guarded(rp, "階層・用語・業務分類"):
            check_structure(cx, rp, hierarchy_only=a.stage == "hierarchy", cap=a.max_material)
    has_flows = (work / "flows").exists() and any((work / "flows").glob("*.flow.json"))
    if a.stage == "flow" or (a.stage == "all" and has_flows):
        with guarded(rp, "フロー"):
            check_flows(cx, rp, only)
    files = sorted((work / "out").glob("*.drawio")) if (work / "out").exists() else []
    if a.stage == "drawio" or (a.stage == "all" and files):
        styles = try_load(a.styles)[0]
        if not files:
            rp.fail("out/*.drawio が無い（build_drawio.py を実行する）")
        elif not isinstance(styles, dict):
            rp.fail(f"--styles のファイルが無い・読めない: {a.styles}")
        else:
            with guarded(rp, "draw.io"):
                check_drawio(cx, rp, files, styles)
            with guarded(rp, "draw.io の作り直し"):
                check_regenerated(cx, rp, files, styles)
    if a.stage == "matrix" or (a.stage == "all" and (work / "out" / MATRIX_BOOK).exists()):
        with guarded(rp, "業務マトリクス一覧"):
            check_matrix(cx, rp)
    if a.stage == "reports" or (a.stage == "all" and (work / "out" / "確認事項一覧.md").exists()):
        with guarded(rp, "報告書"):
            check_reports(cx, rp)
    if a.stage == "deliverables" or (a.stage == "all" and files):
        with guarded(rp, "納品物の文字"):
            check_deliverables(cx, rp)
    if a.stage == "review" or (a.stage == "all" and any((work / "review").glob("*.review.json"))):   # 採否表ができてから（照合シートを作ると、review/ に記録だけができる）
        with guarded(rp, "原典照合"):
            check_review(cx, rp)
    if a.numbers:
        NUMBERS["activities"] = len(list((work / "activities").glob("*.activity.json"))) if (work / "activities").exists() else 0
        save(pathlib.Path(a.numbers), dict(NUMBERS, fails=len(rp.fails), warns=len(rp.warns)))
    rp.finish(a.report)


if __name__ == "__main__":
    main()
