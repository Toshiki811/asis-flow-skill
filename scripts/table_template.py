#!/usr/bin/env python3
"""table_template.py — 表の納品物（xlsx）を、テンプレートから作る部品（build_matrix.py・build_reports.py が使う）。

  python3 scripts/table_template.py check                                    # assets/ のテンプレート 2 つを確かめる
  python3 scripts/table_template.py check <テンプレート.xlsx> --book matrix    # 差し替えるテンプレートを確かめる（matrix＝業務マトリクス一覧／questions＝確認事項一覧）
  python3 scripts/table_template.py default                                  # 既定のテンプレートを assets/ に作り直す
  python3 scripts/table_template.py supply <テンプレート.xlsx> [-o <保存先>]     # 基準の表の無いテンプレートに、処理時間に要るものを補って、別の名前で保存する

テンプレートは、欄と見た目だけを決める。中身は、機械が同じデータから流し込む（LLM は、テンプレートに書かない）。
  表のシート    {…} の印がある行が「見本の行」。印（{質問}・{大分類} …）が、機械が埋める欄。その上の行は、見出し・表題として、そのまま出る。
                印の無い列は、人が埋める欄（見本の行に書いた文字・書式・入力規則を、すべての行に写す）。見出しの文字・列の順は自由。列は消してよい（点検が使う欄は除く）。
                列の幅・字・塗り・罫線・折り返し・固定する枠・フィルター・印刷の設定は、テンプレートのもの。見本の行より下は、空にしておく。
  繰り返す列    業務マトリクスの作業の列（印は {作業}、見出しは ①）。1つだけ置く。機械が、要る数だけ ②③… と繰り返す。右に置いた欄は、そのぶん右へ送る。
  基準の表      名前（定義された名前）「処理時間の表」の範囲。1 行目が見出し（小・中・大・特大。特大の列は、無くてもよい — 無ければ、大の列を引く）、1 列目が動作の種類。仮の処理時間は、この表を引く式で入る。
                名前が無いテンプレート（前の形）には、既定のテンプレートから、基準の表・処理時間の行の書式・処理時間の欄を補って、書き出す（ファイルは変えない。補ったことを知らせる）。
  文のシート    見本の行が無いシート（読み方・表紙 など）。そのまま出る。文の中の {フローの数} のような印にだけ、機械が数を入れる。
  シートの名前  自由（どの一覧かは、印で決まる）。一覧の前の名前（spec の was）のシートも、その一覧として受ける。
  空の一覧      spec の drop_empty に挙げた一覧（業務マトリクス一覧の、プログラム用のシート）は、中身が1行も無ければ、本に入れない。
  書式の見本    シート「書式の見本」の A 列＝名前、B 列＝見本のセル。判断・推定などの塗りと、セルの中の字の種類を、ここから読む（出力には入れない）。
書き出すと、列の対応を、本のとなりに残す（<本の名前>.columns.json）。点検（check_flow.py --stage matrix）が、それを通して読み戻す。
マクロ付き（.xlsm）のテンプレートは、受け取らない。図形・グラフは、写らないことがある（文字・書式・入力規則・印刷の設定を置く）。
"""
import argparse
import copy
import datetime
import json
import pathlib
import re
import sys

from common import ASSETS, save

STYLE_SHEET = "書式の見本"
TOKEN = re.compile(r"\{([^{}\n]+)\}")
WHOLE = re.compile(r"^\{([^{}\n]+)\}$")
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳㉑㉒㉓㉔㉕㉖㉗㉘㉙㉚㉛㉜㉝㉞㉟㊱㊲㊳㊴㊵㊶㊷㊸㊹㊺㊻㊼㊽㊾㊿"
FORMULA, EDIT = "＝式", "＝人が埋める"                                  # 中身を文字で比べるときの、式のセル・人が書き換えてよいセルの書き方（点検が、別に確かめる）
FIRST = CIRCLED[0]                                                    # 繰り返す列の見出しに書く字。機械が ②③… に置き換える
SPREAD = 200                                                          # まとめて幅を決めた列（<col min max>）を、1列ずつに開く範囲（使っている列の右、ここまで）


class TemplateError(Exception):
    """テンプレート・列の対応が、決まりどおりでない"""


class Cell:
    """表の1セル。runs は (文字, 字の種類) の並びで、つないだものがセルの文字（text）。fill は塗りの種類（「書式の見本」の名前に当たる鍵）。
    formula は、式を返す関数（番地は、書くときにテンプレートが決めるので、At を受け取って式の文字を返す）。font は、式・数のセルの字の種類。
    cond は、条件付き書式の鍵（式のセルだけ、その字にする）。tag は、人が書き換えてよいセルの札（読み戻して引き継ぐときの、突き合わせの鍵）。
    shown は、式のセルの、計算の結果（数・文字。または、At を受け取って返す関数）。式と一緒に保存する（計算をしない表示でも、値が出る。keep_results）"""

    def __init__(self, runs="", fill=None, note=None, link=None, formula=None, font=None, cond=None, tag=None, shown=None):
        self.runs = [(runs, "text")] if isinstance(runs, str) else [r for r in runs if r[0] != ""]
        self.fill, self.note, self.link = fill, note, link
        self.formula, self.font, self.cond, self.tag, self.shown = formula, font, cond, tag, shown

    @property
    def text(self):
        if self.tag is not None:
            return EDIT
        return FORMULA if self.formula is not None else "".join(str(t) for t, _s in self.runs)


class At:
    """式に入れる番地を返す（列と行は、テンプレートが決める。書くときに分かる）"""

    def __init__(self, plans, tpl):
        self.plans, self.tpl, self.here = plans, tpl, None
        self.table, self.results = dict((tpl.time or {}).get("values") or {}), {}   # 基準の表の値（直した値を入れたあと）／式のセルの、計算の結果 {シート名: {番地: 値}}

    def cell(self, canon, i, key):
        """一覧 canon の、データの i 行目（0 から）の、欄 key（欄の名前。繰り返す列は 1 からの番号）の番地。別のシートなら、シート名つき。欄が無ければ None"""
        from openpyxl.utils import get_column_letter
        plan = self.plans.get(canon)
        col = plan and plan["col"].get(key)
        if not col:
            return None
        at = f"{get_column_letter(col)}{plan['sample'] + i}"
        return at if canon == self.here else f"{quoted(plan['title'])}!{at}"

    def lookup(self, action, volume):
        """基準の表から、動作の種類 × 量の分を引く式"""
        t = self.tpl.time
        col = t["cols"].get(volume) or t["cols"][(self.tpl.spec["time"].get("spare") or {})[volume]]     # 表に、その量の列が無い（前の形のテンプレート）: 代わりの列
        return f'VLOOKUP("{action}",{t["ref"]},{col},FALSE)'

    def minutes(self, action, volume):
        """基準の表の、動作の種類 × 量の分（lookup の式の、計算の結果）"""
        spare = (self.tpl.spec["time"].get("spare") or {}).get(volume)
        return self.table.get((action, volume), self.table.get((action, spare)))


def quoted(title):
    """式の中の、シートの名前"""
    return "'" + str(title).replace("'", "''") + "'"


def circled(k):
    """列の番号（①〜㊿。それより先は (51)）"""
    return CIRCLED[k - 1] if 1 <= k <= len(CIRCLED) else f"({k})"


def clean(text):
    """xlsx に書けない文字（制御文字）を落とす"""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub("", str(text))


def head_of(ws, c, sample):
    """表の c 列の見出し（見本の行・データの 1 行目 sample より上で、いちばん近い字）"""
    return str(next((ws.cell(r, c).value for r in range(sample - 1, 0, -1) if ws.cell(r, c).value not in (None, "")), "") or "").replace("\n", " ")


def lock_of(path):
    """表計算ソフトで開いている印のファイル（Excel の ~$・LibreOffice の .~lock.…#）。無ければ None"""
    path = pathlib.Path(path)
    return next((x for x in (path.with_name("~$" + path.name), path.with_name(f".~lock.{path.name}#")) if x.exists()), None)


def stored(c):
    """人が埋める欄のセル 1 つの値を、answers.json に書ける形に（空なら None）: {"value": …, "type": 日時の種類, "format": 表示の形}"""
    v = None if c is None else c.value
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        out = {"value": v.isoformat(), "type": type(v).__name__}
    elif isinstance(v, (bool, int, float)):
        out = {"value": v}
    else:
        out = {"value": str(v)}
    if c.number_format and c.number_format != "General":
        out["format"] = c.number_format
    return out


def restored(s):
    """stored の形から、セルに書く値"""
    t = s.get("type")
    return getattr(datetime, t).fromisoformat(s["value"]) if t in ("datetime", "date", "time") else s["value"]


def plain(v):
    """見本の行の値（人が埋める欄の、既定の値）を、columns.json に書ける形に"""
    return v.isoformat() if isinstance(v, (datetime.datetime, datetime.date, datetime.time)) else v if isinstance(v, (str, int, float, bool)) or v is None else str(v)


def side_path(book):
    """列の対応のファイル（本のとなり）"""
    book = pathlib.Path(book)
    return book.with_name(book.stem + ".columns.json")


def remembered(book, default):
    """前の書き出しが使ったテンプレート（列の対応の template_path）。記録が無い・既定と同じなら None。
    記録した様式のファイルが無くなっていれば止める（黙って既定の様式に戻すと、様式で足した列と人の値が、納品物から消える）"""
    try:
        p = json.loads(side_path(book).read_text(encoding="utf-8")).get("template_path")
    except (OSError, ValueError, AttributeError):
        return None
    if not p or pathlib.Path(p).resolve() == pathlib.Path(default).resolve():
        return None
    if not pathlib.Path(p).exists():
        raise SystemExit(f"前の書き出しで使ったテンプレート（{p}）が無い。--template で様式を指定する（既定の様式に戻すなら --template {default}）")
    return pathlib.Path(p)


def note_lines(text, per_line=58):
    """コメントの枠の高さを決めるための、折り返したあとの行数（全角を 2 と数えて per_line で折る）"""
    import math
    import unicodedata
    return sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in "WFA" else 1 for ch in ln) / per_line)) for ln in text.split("\n"))


def blocks(runs):
    """リッチテキストの区切り [(文字, 字の種類)]。空白・改行だけの区切りは、前の区切りにつなぐ
    （空白だけの区切りがあると、Excel はブックを壊れていると見て開かない。LibreOffice は開くので、気づきにくい）"""
    out = []
    for t, s in runs:
        if out and not str(t).strip():
            out[-1][0] += str(t)
        else:
            out.append([str(t), s])
    if len(out) > 1 and not out[0][0].strip():
        out[1][0] = out[0][0] + out[1][0]
        out = out[1:]
    return [(t, s) for t, s in out]


def ungroup(ws):
    """まとめて幅を決めた列を、1列ずつに開く（1列だけ幅を変える・列を足すときに、範囲が重ならないように）"""
    from openpyxl.utils import get_column_letter
    limit = ws.max_column + SPREAD
    for d in list(ws.column_dimensions.values()):
        lo, hi = d.min, d.max
        if not (lo and hi and hi > lo):
            continue
        d.min = d.max = lo
        for c in range(lo + 1, min(hi, limit) + 1):
            nd = ws.column_dimensions[get_column_letter(c)]
            nd.width, nd.hidden, nd.min, nd.max = d.width, d.hidden, c, c
            if d.customWidth is not None:
                nd.customWidth = d.customWidth


def copy_cell(src, dst):
    """ほかの本のセルを写す（中身と書式。書式の番号は本ごとなので、1 つずつ写す）"""
    dst.value = src.value
    if src.has_style:
        dst.font, dst.fill, dst.border, dst.alignment = copy.copy(src.font), copy.copy(src.fill), copy.copy(src.border), copy.copy(src.alignment)
        dst.number_format, dst.protection = src.number_format, copy.copy(src.protection)


def col_width(ws, c):
    from openpyxl.utils import get_column_letter
    d = ws.column_dimensions.get(get_column_letter(c))
    return d.width if d is not None else None


def inline_font(font):
    """セルの字 → セルの中の区切りの字"""
    from openpyxl.cell.text import InlineFont
    return InlineFont(rFont=font.name, sz=font.sz, b=True if font.b else None, i=True if font.i else None, u=font.u, color=copy.copy(font.color) if font.color is not None else None)


class Template:
    """テンプレートを読んだもの。problems が空でなければ、使えない"""

    def __init__(self, path, spec):
        from openpyxl import load_workbook
        self.path, self.spec = pathlib.Path(path), spec
        self.problems, self.notes = [], []
        self.tables, self.static, self.fills, self.fonts = {}, {}, {}, {}   # シート名 → 表の読み取り／シート名 → {セル: 印のある文}／塗り／字
        if self.path.suffix.lower() != ".xlsx":
            raise TemplateError(f"{self.path.name}: テンプレートは .xlsx だけ（マクロ付き・古い形式は、受け取らない）")
        if not self.path.exists():
            raise TemplateError(f"{self.path}: テンプレートが無い（python3 scripts/table_template.py default で、既定のテンプレートを作り直せる）")
        try:
            self.wb = load_workbook(self.path)
        except Exception as e:                                       # 壊れた xlsx は、どの例外で落ちるか決まっていない
            raise TemplateError(f"{self.path.name}: 読めない（{type(e).__name__}: {str(e)[:80]}）")
        self.repeat = spec.get("repeat") or {}
        self.aligns, self.time = {}, None                            # 塗りと一緒に写す、文字の位置／基準の表（無ければ None）
        self.late = set((spec.get("time") or {}).get("styles") or ())  # 基準の表があるときだけ要る書式（処理時間の行。古いテンプレートには無い）
        self.soft = set((spec.get("time") or {}).get("soft") or ())    # 無くてもよい書式（特大のセル。無ければ、ふつうの仮の値と同じ見た目）。条件付き書式は、セルごとに掛ける
        self.supplied = self.supply_time()                           # 基準の表の無いテンプレート: 既定のテンプレートから補う（補ったら、その知らせの文。補わなければ、空）
        self.read_styles()
        for ws in self.wb.worksheets:
            if ws.title != STYLE_SHEET:
                ungroup(ws)
                self.read_sheet(ws)
        self.whole_book()
        self.read_time()

    def bad(self, text):
        self.problems.append(f"{self.path.name}: {text}")

    def canon_of(self, title, have):
        """シートの印（have）から、どの一覧かを決める: (一覧の名前か None, 印をすべて持つ一覧の並び)。
        シート名が、その一覧の名前（前の名前も）なら、それ。違えば、印をすべて持つ一覧が 1 つだけのとき、それ"""
        fits = [name for name in self.spec["sheets"] if set(have) <= set(self.fields(name))]
        was = (self.spec.get("was") or {}).get(title)
        return (title if title in fits else was if was in fits else fits[0] if len(fits) == 1 else None), fits

    def supply_time(self):
        """基準の表（名前の付いた範囲）の無いテンプレートに、既定のテンプレートから、処理時間に要るものを補う（どのテンプレートでも、最初の出力から、処理時間が出る）。
        補うのは、基準の表のシートと名前／「書式の見本」の、処理時間の行／処理時間の欄（見本の行の、使っている列の右に足す。繰り返す列の右なら、作業の列のあとに並ぶ）。
        開いた本（self.wb）を直すだけで、テンプレートのファイルは変えない（table_template.py supply で、補ったテンプレートを保存できる）"""
        from openpyxl.utils import get_column_letter
        from openpyxl.workbook.defined_name import DefinedName
        want = self.spec.get("time") or {}
        src = want.get("default")
        if not src or self.wb.defined_names.get(want["name"]) is not None or pathlib.Path(src).resolve() == self.path.resolve():
            return ""
        try:
            base = Template(src, self.spec)
        except TemplateError:
            base = None
        if base is None or base.time is None or base.problems:
            self.notes.append(f"既定のテンプレート（{pathlib.Path(src).name}）を読めないので、処理時間の基準の表を補えない")
            return ""
        title = base.time["title"]                                    # 1. 基準の表のシート（「読み方」の前。無ければ、書式の見本の前・最後）
        while title in self.wb.sheetnames:
            title += "（既定）"
        names = self.wb.sheetnames
        at = next((names.index(t) for t in (want.get("before"), STYLE_SHEET) if t in names), len(names))
        ws, from_ws = self.wb.create_sheet(title, at), base.wb[base.time["title"]]
        for row in from_ws.iter_rows():
            for c in row:
                copy_cell(c, ws.cell(c.row, c.column))
        for k, d in from_ws.column_dimensions.items():
            if d.width is not None:
                ws.column_dimensions[k].width = d.width
        ws.freeze_panes = from_ws.freeze_panes
        ws.page_setup.orientation, ws.page_setup.paperSize = from_ws.page_setup.orientation, from_ws.page_setup.paperSize
        area = base.wb.defined_names[want["name"]].attr_text.rsplit("!", 1)[1]
        self.wb.defined_names[want["name"]] = DefinedName(want["name"], attr_text=f"{quoted(title)}!{area}")
        keys = set(want.get("styles") or ()) | set(want.get("soft") or ())   # 2. 書式の見本の、処理時間の行
        labels = {n for n, k in list((self.spec.get("fills") or {}).items()) + list((self.spec.get("fonts") or {}).items()) if k in keys}
        if STYLE_SHEET in self.wb.sheetnames and STYLE_SHEET in base.wb.sheetnames:
            mine = self.wb[STYLE_SHEET]
            have = {str(r[0].value or "").strip() for r in mine.iter_rows(min_col=1, max_col=1)}
            for row in base.wb[STYLE_SHEET].iter_rows():
                if str(row[0].value or "").strip() in labels - have:
                    r = mine.max_row + 1
                    for c in row:
                        copy_cell(c, mine.cell(r, c.column))
        every, added = {f for name in self.spec["sheets"] for f in self.fields(name)}, []   # 3. 処理時間の欄
        for ws in self.wb.worksheets:
            marks = [(c.row, c.column, m.group(1)) for row in ws.iter_rows() for c in row if isinstance(c.value, str)
                     for m in [WHOLE.match(c.value.strip())] if m and m.group(1) in every] if ws.title != STYLE_SHEET else []
            if not marks:
                continue
            sample = min(r for r, _c, _t in marks)
            have = {t for r, _c, t in marks if r == sample}
            canon, _fits = self.canon_of(ws.title, have)
            lost = [f for f in (want.get("fields") or {}).get(canon, []) if f not in have]
            if not lost:
                continue
            ungroup(ws)
            used = max(c.column for row in ws.iter_rows(max_row=sample) for c in row if c.value not in (None, ""))
            for k, f in enumerate(lost, 1):
                for r in range(max(1, sample - 1), sample + 1):       # 見出しの行と、見本の行（書式は、左隣の列のもの）
                    x = ws.cell(r, used + k)
                    x._style = copy.copy(ws.cell(r, used)._style)
                    x.value = "{" + f + "}" if r == sample else (want.get("heads") or {}).get(f, f)
                ws.column_dimensions[get_column_letter(used + k)].width = 14
            if ws.auto_filter.ref:
                ws.auto_filter.ref = f"{ws.auto_filter.ref.split(':')[0]}:{get_column_letter(used + len(lost))}{sample}"
            added.append(f"「{ws.title}」に {'・'.join('{' + f + '}' for f in lost)}")
        self.notes.append(f"名前「{want['name']}」の範囲（処理時間の基準の表）が無いテンプレート。既定のテンプレートから、基準の表のシート「{title}」・処理時間の行の書式"
                          + (f"・処理時間の欄（{'／'.join(added)}。右端に足した）" if added else "") + "を補って、書き出す（テンプレートのファイルは、変えない。"
                          "補ったテンプレートを残して、欄の場所・表の数を直すなら: python3 scripts/table_template.py supply <テンプレート.xlsx>）")
        return self.notes[-1]

    def read_styles(self):
        want_fill, want_font = self.spec.get("fills") or {}, self.spec.get("fonts") or {}
        if STYLE_SHEET not in self.wb.sheetnames:
            if want_fill or want_font:
                self.bad(f"シート「{STYLE_SHEET}」が無い（判断・推定などの塗りと、セルの中の字の種類を、ここから読む）")
            return
        for row in self.wb[STYLE_SHEET].iter_rows(min_col=1, max_col=2):
            name = str(row[0].value or "").strip()
            if name in want_fill:
                self.fills[want_fill[name]] = copy.copy(row[1].fill)
                if want_fill[name] in (self.spec.get("aligns") or ()):
                    self.aligns[want_fill[name]] = copy.copy(row[1].alignment)
            if name in want_font:
                self.fonts[want_font[name]] = copy.copy(row[1].font)
        lost = self.lost_styles(lambda k: k not in self.late)
        if lost:
            self.bad(f"シート「{STYLE_SHEET}」の A 列に、次の名前の行が無い: {'・'.join(lost)}")

    def lost_styles(self, mine):
        """書式の見本に無い、塗り・字の名前（mine が真の鍵だけ）"""
        fills, fonts = self.spec.get("fills") or {}, self.spec.get("fonts") or {}
        need = lambda k, have: mine(k) and k not in self.soft and k not in have
        return list(dict.fromkeys([n for n, k in fills.items() if need(k, self.fills)] + [n for n, k in fonts.items() if need(k, self.fonts)]))

    def read_time(self):
        """基準の表（名前の付いた範囲。1 行目＝見出し、1 列目＝動作の種類）。名前の無いテンプレートには、先に補ってある（supply_time）。補えなかったときだけ、処理時間の行を出さない"""
        want = self.spec.get("time")
        if not want:
            return
        self.time, problem = time_table(self.wb, want)
        lost = self.lost_styles(lambda k: k in self.late) if self.time is not None else []
        if lost:
            self.bad(f"処理時間の基準の表があるのに、シート「{STYLE_SHEET}」の A 列に、処理時間の行の書式の、次の名前の行が無い: {'・'.join(lost)}")
        if problem:
            self.bad(problem)
        elif self.time is None:
            self.notes.append(f"名前「{want['name']}」の範囲（処理時間の基準の表）が無い。処理時間の行は、出さない")
        else:
            spare = want.get("spare") or {}
            for v in self.time["absent"]:
                self.notes.append(f"処理時間の基準の表に、「{v}」の列が無い（前の形）。量が「{v}」の作業は、「{spare[v]}」の列を引く")
            fonts = self.spec.get("fonts") or {}
            for n in [n for n, k in fonts.items() if k in self.soft and k not in self.fonts]:
                self.notes.append(f"シート「{STYLE_SHEET}」に、「{n}」の行が無い（前の形）。そのセルは、ふつうの仮の値と同じ見た目になる")

    def fields(self, canon):
        return list(self.spec["sheets"][canon]) + ([self.repeat[canon]] if canon in self.repeat else [])

    def read_sheet(self, ws):
        every = {f for name in self.spec["sheets"] for f in self.fields(name)}
        marks = [(c.row, c.column, m.group(1)) for row in ws.iter_rows() for c in row if isinstance(c.value, str)
                 for m in [WHOLE.match(c.value.strip())] if m and m.group(1) in every]
        if not marks:                                                 # 文のシート
            inline = {c.coordinate: c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and TOKEN.search(c.value)}
            unknown = sorted({t for v in inline.values() for t in TOKEN.findall(v) if t not in (self.spec.get("values") or [])})
            if unknown:
                self.bad(f"シート「{ws.title}」: 知らない印 {'・'.join('{' + t + '}' for t in unknown)}"
                         f"（表の欄の印は、1つのセルに1つだけ書く。文の中に書ける印: {'・'.join(self.spec.get('values') or []) or 'なし'}）")
            self.static[ws.title] = inline
            return
        sample = min(r for r, _c, _t in marks)
        if any(r != sample for r, _c, _t in marks):
            self.bad(f"シート「{ws.title}」: 印が、{sample} 行目のほかの行にもある（見本の行は、1行だけ）")
        cols = {c: t for r, c, t in marks if r == sample}
        canon, fits = self.canon_of(ws.title, cols.values())
        if canon is None:
            self.bad(f"シート「{ws.title}」: 印が、どの一覧のものか決まらない（{'・'.join(fits) if fits else '全部の印を持つ一覧が無い'}）。シートの名前を、既定の名前にする")
            return
        if any(t["canon"] == canon for t in self.tables.values()):
            self.bad(f"シート「{ws.title}」: 一覧「{canon}」のシートが、2つある")
            return
        for c in ws[sample]:
            if isinstance(c.value, str) and TOKEN.search(c.value) and c.column not in cols:
                self.bad(f"シート「{ws.title}」{c.coordinate}: 知らない印「{c.value.strip()[:30]}」（「{canon}」の欄: {'・'.join(self.fields(canon))}。1つのセルに、印を1つだけ書く）")
        below = [c.coordinate for row in ws.iter_rows(min_row=sample + 1) for c in row if c.value not in (None, "")]
        if below:
            self.bad(f"シート「{ws.title}」: 見本の行（{sample} 行目）より下に、文字がある（{'・'.join(below[:5])}）。下は、空にする")
        rep = next((c for c, t in cols.items() if t == self.repeat.get(canon)), None)
        used = max([c.column for row in ws.iter_rows(max_row=sample) for c in row if c.value not in (None, "")] or [1])
        if rep is not None and list(cols.values()).count(self.repeat[canon]) > 1:
            self.bad(f"シート「{ws.title}」: 繰り返す列（{{{self.repeat[canon]}}}）は、1つだけ置く")
        lost = [f for f in (self.spec.get("required") or {}).get(canon, []) if f not in cols.values()]
        if lost:
            self.bad(f"シート「{ws.title}」: 点検が使う欄が無い: {'・'.join('{' + f + '}' for f in lost)}")
        head = lambda c: str(next((ws.cell(r, c).value for r in range(sample - 1, 0, -1) if ws.cell(r, c).value not in (None, "")), "") or "").replace("\n", " ")
        self.tables[ws.title] = {"canon": canon, "sample": sample, "cols": cols, "repeat": rep, "used": used}
        manual = [head(c) or f"{c} 列目" for c in range(1, used + 1) if c not in cols]
        gone = [f for f in self.fields(canon) if f not in cols.values() and f not in ((self.spec.get("quiet") or {}).get("fields") or {}).get(canon, ())]
        if manual:
            self.notes.append(f"シート「{ws.title}」: 機械が埋めない列（人が埋める欄）: {'・'.join(manual)}")
        if gone:
            self.notes.append(f"シート「{ws.title}」: テンプレートに無い欄（出さない）: {'・'.join(gone)}")

    def whole_book(self):
        have = {t["canon"] for t in self.tables.values()}
        for name in (self.spec.get("required") or {}):
            if name not in have:
                self.bad(f"一覧「{name}」のシートが無い（点検が使う）")
        gone = [n for n in self.spec["sheets"] if n not in have and n not in ((self.spec.get("quiet") or {}).get("lists") or ())]
        if gone:
            self.notes.append(f"テンプレートに無い一覧（出さない）: {'・'.join(gone)}")


def time_table(wb, want):
    """本（テンプレート・書き出した本）の、基準の表。(表, 困ること)。名前が無ければ (None, None)。
    表は {"title": シート名, "col": 1 列目, "cols": {量: 何列目か}, "rows": {動作の種類: 行}, "values": {(動作の種類, 量): 値}, "ref": 式に書く範囲}"""
    from openpyxl.utils import get_column_letter, range_boundaries
    dn = wb.defined_names.get(want["name"])
    if dn is None:
        return None, None
    try:
        (title, area), = list(dn.destinations)
        c0, r0, c1, r1 = range_boundaries(area.replace("$", ""))
        ws = wb[title]
    except Exception:                                                 # 範囲が 2 つ・シートが無い・番地でない
        return None, f"名前「{want['name']}」が、シートの中の 1 つの範囲を指していない"
    head = [str(ws.cell(r0, c).value or "").strip() for c in range(c0, c1 + 1)]
    cols = {v: head.index(v) + 1 for v in want["volumes"] if v in head}
    rows = {str(ws.cell(r, c0).value or "").strip(): r for r in range(r0 + 1, r1 + 1)}
    spare = want.get("spare") or {}                                    # 無くてもよい量（列が無ければ、代わりの列を引く）
    lost = [v for v in want["volumes"] if v not in cols and v not in spare] + [a for a in want["actions"] if a not in rows]
    if lost:
        return None, (f"「{title}」の基準の表（名前「{want['name']}」の範囲）に、次の見出し・動作の種類が無い: {'・'.join(lost)}"
                      "（1 行目に 小・中・大、1 列目に動作の種類）")
    values = {(a, v): ws.cell(rows[a], c0 + cols[v] - 1).value for a in want["actions"] for v in want["volumes"] if v in cols}
    odd = [f"{a}×{v}" for (a, v), x in values.items() if isinstance(x, bool) or not isinstance(x, (int, float)) or x < 0]
    if odd:
        return None, f"「{title}」の基準の表に、数でない・負の値がある: {'・'.join(odd[:8])}（分を、0 以上の数で書く）"
    return {"title": title, "col": c0, "cols": cols, "rows": rows, "values": values, "absent": [v for v in want["volumes"] if v not in cols],
            "ref": f"{quoted(title)}!${get_column_letter(c0)}${r0 + 1}:${get_column_letter(c1)}${r1}"}, None


def put(x, cell, tpl, base_font, at=None):
    """セル1つに、中身を書く（書式は、先に見本の行のものを写してある）"""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.comments import Comment
    x.value = None
    if cell.fill and cell.fill in tpl.fills:
        x.fill = copy.copy(tpl.fills[cell.fill])
        if cell.fill in tpl.aligns:
            x.alignment = copy.copy(tpl.aligns[cell.fill])
    runs = cell.runs
    if cell.formula is not None:                                     # 式（番地は、いま決まる）
        x.value = cell.formula(at)
    elif len(runs) == 1 and not isinstance(runs[0][0], str):         # 数（順番・段・人が入れた時間）は数のまま書く
        x.value = runs[0][0]
    elif len(runs) == 1:
        x.value = clean(runs[0][0])
        x.data_type = "s"                                            # 「=」で始まる原文を、式として書かない
        if runs[0][1] in tpl.fonts:
            x.font = copy.copy(tpl.fonts[runs[0][1]])
    elif runs:
        x.value = CellRichText(*[TextBlock(inline_font(tpl.fonts.get(s, base_font)), clean(t)) for t, s in blocks(runs)])
    if cell.font in tpl.fonts:
        x.font = copy.copy(tpl.fonts[cell.font])
    if cell.link:
        x.hyperlink = cell.link
    if cell.note:
        x.comment = Comment(clean(cell.note)[:3000], "asis-flow-skill", width=420, height=min(600, 30 + 16 * note_lines(cell.note)))


def shift_columns(ws, sample, first, last, by):
    """見本の行までの、first〜last 列を、by 列だけ横へ送る（見出し・書式・幅・入力規則）。繰り返す列の右に置いた欄のため"""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
    order = range(last, first - 1, -1) if by > 0 else range(first, last + 1)
    widths = {c: col_width(ws, c) for c in range(first, last + 1)}
    for c in order:
        for r in range(1, sample + 1):
            src = ws._cells.pop((r, c), None)
            if src is not None:
                dst = ws.cell(r, c + by)
                dst.value, dst._style = src.value, copy.copy(src._style)
            else:
                ws._cells.pop((r, c + by), None)
    for c in range(first, last + 1):
        ws.column_dimensions.pop(get_column_letter(c), None)
    for c, w in widths.items():
        if w is not None:
            ws.column_dimensions[get_column_letter(c + by)].width = w
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange([CellRange(min_col=g.min_col + by, min_row=g.min_row, max_col=g.max_col + by, max_row=g.max_row)
                                   if first <= g.min_col and g.max_col <= last else g for g in dv.sqref.ranges])


def arrange(ws, tb, rows, tpl):
    """表のシート1つの、列の割り付け。繰り返す列を、要る数だけ増やし（右の欄は、そのぶん右へ送る）、見出しを書く。
    返すのは、書くときに使うもの（列 → 行の中の位置、書く列、見本の書式）と、欄 → 列の対応"""
    from openpyxl.utils import get_column_letter
    canon, sample, cols, rep = tb["canon"], tb["sample"], dict(tb["cols"]), tb["repeat"]
    fixed = list(tpl.spec["sheets"][canon])
    head = [c.text for c in rows[0]]
    if head[:len(fixed)] != fixed:
        raise TemplateError(f"一覧「{canon}」の欄が、決まった並びでない（{head[:len(fixed)]}）")
    index, steps, used = {f: j for j, f in enumerate(fixed)}, head[len(fixed):], tb["used"]
    by = len(steps) - 1 if rep is not None else 0
    if by and used > rep:                                             # 繰り返す列の右の欄
        if by < 0:                                                    # 繰り返す欄が1つも無い: 見本の列を消して、左へ詰める
            for r in range(1, sample + 1):
                ws._cells.pop((r, rep), None)
        shift_columns(ws, sample, rep + 1, used, by)
        cols = {(c + by if c > rep else c): f for c, f in cols.items()}
    width = [c for c in range(1, used + by + 1)] if rep is not None else list(range(1, used + 1))
    reps = list(range(rep, rep + len(steps))) if rep is not None else []
    if rep is not None and steps:
        heads = [(r, ws.cell(r, rep).value, copy.copy(ws.cell(r, rep)._style)) for r in range(1, sample)]
        mark, w = copy.copy(ws.cell(sample, rep)._style), col_width(ws, rep)
        for k, name in enumerate(steps):
            c = rep + k
            if k:
                ws.cell(sample, c)._style = copy.copy(mark)
                if w is not None:
                    ws.column_dimensions[get_column_letter(c)].width = w
            for r, v, st in heads:
                x = ws.cell(r, c)
                x._style = copy.copy(st)
                x.value = v.replace(FIRST, name) if isinstance(v, str) else v
    elif rep is not None and used == rep:                             # 繰り返す欄が1つも無く、右の欄も無い: 見本の列ごと消す
        for r in range(1, sample + 1):
            ws._cells.pop((r, rep), None)
        ws.column_dimensions.pop(get_column_letter(rep), None)
        width = list(range(1, rep))
    where = {c: index[f] for c, f in cols.items() if f in index}     # 列 → 行の中の位置（欄の順）
    where.update({c: len(fixed) + k for k, c in enumerate(reps)})
    col = {f: c for c, f in cols.items() if f in index}
    col.update({k + 1: c for k, c in enumerate(reps)})
    return {"title": ws.title, "canon": canon, "sample": sample, "width": width, "where": where, "col": col, "fixed": fixed,
            "style": {c: copy.copy(ws.cell(sample, c)._style) for c in width}, "font": {c: copy.copy(ws.cell(sample, c).font) for c in width},
            "const": {c: ws.cell(sample, c).value for c in width if c not in where},
            "manual": {c: head_of(ws, c, sample) or f"{c} 列目" for c in width if c not in where}}       # 人が埋める欄（印の無い列）の見出し


def row_keys(spec, canon, rows, keys=None):
    """行ごとの鍵（人が埋める欄の値を、書き出し直したあとも同じ行に戻すため）: [(鍵, 行の名前)]。
    spec の keys に欄があれば、その欄の文字（確認事項は番号・業務フローのノード一覧はフローIDとノード番号。並べ替えても決まる）。
    欄で決まらない一覧（業務マトリクス）は、作る側が渡す鍵と、key_check の欄の文字（読み戻すとき、行を動かしていないかを、これで確かめる）"""
    fixed = list(spec["sheets"][canon])
    text = lambda r, fields: "｜".join(clean(r[fixed.index(f)].text) for f in fields if f in fixed)
    fields = (spec.get("keys") or {}).get(canon)
    if fields:
        return [(text(r, fields), "") for r in rows[1:]]
    if keys is not None:
        return [(k, text(r, (spec.get("key_check") or {}).get(canon) or [])) for k, r in zip(keys, rows[1:])]
    return None


def write_table(ws, plan, rows, tpl, at, keys=None, kept=None):
    """表のシート1つに、行を流し込む（列の割り付けは、arrange が済ませている）。返すのは、列の対応（columns.json の1シート分）。
    keys: 行ごとの (鍵, 行の名前)。kept: 鍵 → {人が埋める欄の見出し: 値}（answers.json から。書き出し直す前に、人が入れた値）"""
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
    canon, sample, width, where, fixed, data = plan["canon"], plan["sample"], plan["width"], plan["where"], plan["fixed"], rows[1:]
    at.here, cond, tags = canon, {}, {}
    for i, row in enumerate(data):
        for c in width:
            x = ws.cell(sample + i, c)
            x._style = copy.copy(plan["style"][c])
            if c in where:
                cell = row[where[c]]
                put(x, cell, tpl, plan["font"][c], at)
                if cell.formula is not None and cell.shown is not None:   # 式のセルの、計算の結果（保存したあとで、式と一緒に書く）
                    at.results.setdefault(ws.title, {})[x.coordinate] = cell.shown(at) if callable(cell.shown) else cell.shown
                key = next((k for k in ([cell.cond] if isinstance(cell.cond, str) else cell.cond or ()) if k in tpl.fonts), None)   # 並びなら、テンプレートにある最初の書式
                if key:
                    cond.setdefault(key, []).append((sample + i, c))
                if cell.tag is not None:
                    tags[x.coordinate] = list(cell.tag)
            else:
                s = (kept or {}).get(keys[i][0], {}).get(plan["manual"][c]) if keys else None
                x.value = plan["const"].get(c) if s is None else restored(s)
                if s is not None and s.get("format"):
                    x.number_format = s["format"]
    if not data:                                                      # 行が無い: 見本の行を消す（印を残さない）
        for c in width:
            ws._cells.pop((sample, c), None)
    last_row, last_col = max(sample - 1, sample + len(data) - 1), max(width)
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"{ws.auto_filter.ref.split(':')[0]}:{get_column_letter(last_col)}{last_row}"
    for dv in ws.data_validations.dataValidation:                    # 見本の行に掛かる入力規則を、すべての行に広げる
        dv.sqref = MultiCellRange([CellRange(min_col=g.min_col, min_row=g.min_row, max_col=g.max_col, max_row=max(g.max_row, last_row))
                                   if g.min_row <= sample <= g.max_row else g for g in dv.sqref.ranges])
    def rule(key, first, **more):
        f = tpl.fonts[key]
        return FormulaRule(formula=[f"_xlfn.ISFORMULA({first})"], font=Font(b=True if f.b else None, i=True if f.i else None, color=copy.copy(f.color), u=f.u), **more)

    for key in sorted(cond, key=lambda k: k not in tpl.soft):         # 式のセルだけ、字を変える（人が数を打つと、ふつうの字に戻る）。セルごとの決まり（特大）を先に書く（先に効く）
        cells = cond[key]
        if key in tpl.soft:                                           # 飛び飛びのセル: 1 つずつ掛ける。塗りも付け、あとの決まり（囲む四角）は効かせない
            shade = tpl.fills.get(key)
            paint = {"fill": PatternFill(fill_type="solid", start_color=copy.copy(shade.fgColor), end_color=copy.copy(shade.fgColor))} if shade is not None and shade.fill_type else {}
            for r, c in cells:
                ws.conditional_formatting.add(f"{get_column_letter(c)}{r}", rule(key, f"{get_column_letter(c)}{r}", stopIfTrue=True, **paint))
            continue
        r0, r1, c0, c1 = min(r for r, _ in cells), max(r for r, _ in cells), min(c for _, c in cells), max(c for _, c in cells)   # 囲む四角に、1つの決まりで掛ける
        ws.conditional_formatting.add(f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1}", rule(key, f"{get_column_letter(c0)}{r0}"))
    repeat = tpl.repeat.get(canon)
    side = {"canon": canon, "header_rows": sample - 1, "tags": tags,
            "columns": [(repeat if c in where and where[c] >= len(fixed) else fixed[where[c]] if c in where else None) for c in range(1, last_col + 1)],
            "manual": [[c, h, plain(plan["const"].get(c))] for c, h in sorted(plan["manual"].items())]}
    if keys and not (tpl.spec.get("keys") or {}).get(canon):       # 欄で決まらない鍵は、行の順に残す
        side["keys"] = [list(k) for k in keys]
    return side


RESULT = re.compile(r'<c r="([A-Z]+[0-9]+)"((?: [A-Za-z:]+="[^"]*")*)>(<f(?: [^>]*)?>[^<]*</f>)<v(?:/>|></v>)</c>')   # 式だけで、値の無いセル（openpyxl の書き方）


def keep_results(path, results):
    """保存した xlsx の、式のセルに、計算の結果も書き足す（results: {シートの名前: {番地: 数 か 文字}}）。返すのは、書き足せなかったセルの数。
    openpyxl は、式だけを書く。計算をしない表示（Finder・メールのプレビュー、Excel の保護ビュー）では、そのセルが 0・空に見える。
    Excel・LibreOffice は、開くときに計算し直す（fullCalcOnLoad）。ここの値は、計算をしない表示のためのもの"""
    import zipfile
    from xml.etree import ElementTree as ET
    from xml.sax.saxutils import escape
    path, left = pathlib.Path(path), sum(len(v) for v in results.values())
    if not left:
        return 0
    with zipfile.ZipFile(path) as z:
        names = [(i.filename, i.date_time) for i in z.infolist()]
        data = {name: z.read(name) for name, _d in names}
    rel = {x.get("Id"): x.get("Target", "") for x in ET.fromstring(data.get("xl/_rels/workbook.xml.rels", b"<r/>"))}
    part = {}                                                         # シートの名前 → 中のファイル
    for s in ET.fromstring(data["xl/workbook.xml"]).iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"):
        target = rel.get(s.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"), "")
        part[s.get("name")] = target.lstrip("/") if target.startswith("/") else "xl/" + target
    for title, cells in results.items():
        if part.get(title) not in data:
            continue
        done = set()

        def one(m):
            at, attrs, v = m.group(1), m.group(2), cells.get(m.group(1))
            if v is None or isinstance(v, bool) or ' t="' in attrs:
                return m.group(0)
            done.add(at)
            if isinstance(v, str):                                    # 文字を返す式
                return f'<c r="{at}"{attrs} t="str">{m.group(3)}<v>{escape(v)}</v></c>'
            return f'<c r="{at}"{attrs}>{m.group(3)}<v>{int(v) if float(v).is_integer() else float(v)}</v></c>'
        data[part[title]] = RESULT.sub(one, data[part[title]].decode("utf-8")).encode("utf-8")
        left -= len(done)
    tmp = path.with_name(path.name + ".tmp")                          # 書き終えてから、置き換える（途中で止まっても、もとの本は残る）
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, date in names:
            z.writestr(zipfile.ZipInfo(name, date), data[name], zipfile.ZIP_DEFLATED)
    tmp.replace(path)
    return left


def fill(template, spec, sheets, values, out_path, table=None, info=None, keys=None, kept=None):
    """テンプレートに流し込んで、保存する。
    sheets: 一覧の名前 → 行（1行目は欄の名前。セルは .runs・.text・.fill・.note・.link を持つもの）。values: 文のシートの印 → 値。
    table: 基準の表の、テンプレートから直す値 {(動作の種類, 量): (分, わけ)}。info: 列の対応（columns.json）に一緒に残すもの。
    keys: 一覧 → 行ごとの鍵（spec の keys の欄で決まらない一覧だけ）。kept: 一覧 → 鍵 → {人が埋める欄の見出し: 値}（answers.json）"""
    from openpyxl.comments import Comment
    tpl = template if isinstance(template, Template) else Template(template, spec)
    if tpl.problems:
        raise TemplateError("\n".join(tpl.problems))
    wb, side = tpl.wb, dict({"template": tpl.path.name, "template_path": str(tpl.path.resolve()), "sheets": {}, "static": {}}, **(info or {}))
    empty = {title: tb["canon"] for title, tb in tpl.tables.items() if tb["canon"] in (spec.get("drop_empty") or ()) and len(sheets[tb["canon"]]) <= 1}
    tables = {title: tb for title, tb in tpl.tables.items() if title not in empty}  # 中身が1行も無い一覧（プログラム用のシート）は、本に入れない
    if empty:
        side["empty"] = sorted(empty.values())                        # 点検が、あとで中身が増えたのに書き出し直していないかを見る
    plans = {tb["canon"]: arrange(wb[title], tb, sheets[tb["canon"]], tpl) for title, tb in tables.items()}
    at = At(plans, tpl)
    at.table.update({k: v for k, (v, _why) in (table or {}).items() if k in at.table})
    if tpl.time:                                                      # 書き出したときの、基準の表の値（次に取り込むとき、人が直したかを、これと比べる）
        side["table"] = [[a, v, x] for (a, v), x in sorted(at.table.items())]
        side["table_version"] = table_version(at.table)
    for title, tb in tables.items():
        canon = tb["canon"]
        side["sheets"][title] = write_table(wb[title], plans[canon], sheets[canon], tpl, at, row_keys(spec, canon, sheets[canon], (keys or {}).get(canon)),
                                            (kept or {}).get(canon))
    for title in empty:
        del wb[title]
    for title, inline in tpl.static.items():
        for where, text in inline.items():
            wb[title][where].value = TOKEN.sub(lambda m: str(values[m.group(1)]), text)
        if inline:
            side["static"][title] = inline
    if tpl.time:
        t = tpl.time
        for (action, volume), (v, why) in (table or {}).items():     # 基準の表の、直した値（テンプレートの値は、コメントに残す）
            if (action, volume) in t["values"] and v != t["values"][(action, volume)]:
                x = wb[t["title"]].cell(t["rows"][action], t["col"] + t["cols"][volume] - 1)
                x.value, x.comment = v, Comment(f"テンプレートの値は {t['values'][(action, volume)]}（{why}）", "asis-flow-skill", width=320, height=60)
        head = wb[t["title"]].cell(min(t["rows"].values()) - 1, t["col"])   # 表の左上の見出し: 表の版（見直しの前と後は、同じ版の表で比べる）
        was = [p for p in (head.comment.text.split("\n\n") if head.comment is not None else []) if p.strip() and not p.startswith("基準の表の版")]
        head.comment = Comment("\n\n".join(was + [""]) + f"基準の表の版 {side['table_version']}（表の値から作る印。値が 1 つでも変われば、版が変わる）。"
                               "見直しの前と後の処理時間は、同じ版の表で比べる", "asis-flow-skill", width=320, height=90)
    if STYLE_SHEET in wb.sheetnames:
        del wb[STYLE_SHEET]
    wb.active = 0
    out_path = pathlib.Path(out_path)
    wb.save(out_path)
    tpl.unkept = keep_results(out_path, at.results)                   # 式のセルに、計算の結果も書く（書けなかった数）
    save(side_path(out_path), side)
    return tpl


def table_version(table):
    """基準の表の版: 表の値（動作の種類 × 量 → 分）から作る短い指紋（T-3f9a1c）。値が 1 つでも違えば変わる。
    見直しの前と後の処理時間は、同じ版の表で比べる（決まり 3.22）"""
    import hashlib
    rows = sorted([a, v, x] for (a, v), x in table.items())
    return "T-" + hashlib.sha1(json.dumps(rows, ensure_ascii=False).encode("utf-8")).hexdigest()[:6]


def read_side(book):
    """列の対応（columns.json）を読んで確かめる。無い・崩れていれば TemplateError"""
    try:
        side = json.loads(side_path(book).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        raise TemplateError(f"{side_path(book).name} が無い・読めない（列の対応。本と一緒に書き出される）")
    from validate import validate
    errs = validate(side, "columns")
    if errs:                                                         # 崩れた対応で読むと、人が入れた値を別の欄に取り込む
        raise TemplateError(f"{side_path(book).name} の形が崩れている（{'／'.join(errs[:3])}）。本を書き出し直す")
    return side


def guess_side(wb, tpl):
    """列の対応（columns.json）の無い本（前の版で書き出した本・戻ってきた本だけを受け取ったとき）の列の対応を、テンプレートの見出しの文字から推し量る。
    分かるのは、欄の列と、人が埋める欄の列だけ（行の鍵は、欄で決まる一覧だけ）"""
    sheets = {}
    for title, tb in tpl.tables.items():
        if title not in wb.sheetnames:
            continue
        tws, ws, sample = tpl.wb[title], wb[title], tb["sample"]
        field = {head_of(tws, c, sample): f for c, f in tb["cols"].items()}
        manual = {head_of(tws, c, sample) or f"{c} 列目": tws.cell(sample, c).value for c in range(1, tb["used"] + 1) if c not in tb["cols"]}
        cols, hands = [], []
        for c in range(1, ws.max_column + 1):
            h = head_of(ws, c, sample)
            cols.append(None if h in manual else field.get(h, ""))
            if h in manual:
                hands.append([c, h, plain(manual[h])])
        sheets[title] = {"canon": tb["canon"], "header_rows": sample - 1, "columns": cols, "manual": hands}
    return {"template": tpl.path.name, "sheets": sheets}


def manual_back(book, spec, template):
    """書き出した本の、人が埋める欄（テンプレートの、印の無い列）を、行の鍵で読む。
    → ({一覧: {鍵: {欄の見出し: 値}}}（本にある行はすべて。値の無い行は {}。見本の行と同じ値は、入れていないものと見る）, [知らせ])。
    鍵は row_keys のとおり。欄で決まらない一覧は、行の名前が書き出したときと同じかを確かめる。
    列の対応（columns.json）が本のとなりに無ければ、テンプレート（パスか Template）の見出しの文字で列を探す（guess_side）。
    列・行を動かした（見出し・行の名前が書き出したときと違う）・どの行のものか決められない値があるときは TemplateError（別の行・欄に取り込まない）"""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    book = pathlib.Path(book)
    wb = load_workbook(book)
    if side_path(book).exists():
        side = read_side(book)
    else:
        side = guess_side(wb, template if isinstance(template, Template) else Template(template, spec))
    text = lambda c: clean(c.value) if c is not None and c.value is not None else ""
    got, notes = {}, []
    for title, s in side["sheets"].items():
        if title not in wb.sheetnames:
            raise TemplateError(f"シート「{title}」が、本に無い")
        ws, canon, columns, top = wb[title], (spec.get("was") or {}).get(s["canon"], s["canon"]), s["columns"], s["header_rows"] + 1
        manual = s.get("manual")
        if manual is None:                                            # 前の版の columns.json: 印の無い列と、本の見出し
            manual = [[k + 1, head_of(ws, k + 1, top) or f"{k + 1} 列目", None] for k, x in enumerate(columns) if x is None]
        if not manual:
            continue
        for c, h, _const in manual:
            now = head_of(ws, c, top)
            if now and now != h and not h.endswith("列目"):
                raise TemplateError(f"「{title}」の {get_column_letter(c)} 列の見出しが「{now}」（書き出したときは「{h}」）。列を足した・動かしたので、人が埋める欄を読めない。"
                                    "列を戻すか、値を書き出し直したブックに入れ直す")
        fields = (spec.get("keys") or {}).get(canon) or []
        at = {f: columns.index(f) + 1 for f in fields if f in columns}
        check = [columns.index(f) + 1 for f in (spec.get("key_check") or {}).get(canon) or [] if f in columns]
        keyed, mine, lost = s.get("keys"), {}, 0
        for i, row in enumerate(ws.iter_rows(min_row=top)):
            cell = lambda c: row[c - 1] if c - 1 < len(row) else None
            vals = {h: x for c, h, const in manual for x in [stored(cell(c))] if x is not None and x["value"] != const}
            if fields and len(at) == len(fields):
                key = "｜".join(text(cell(at[f])) for f in fields)
                if not key.strip("｜"):
                    lost += bool(vals)
                    continue
            elif keyed is not None and i < len(keyed):
                key, name = keyed[i]
                if "｜".join(text(cell(c)) for c in check) != name:
                    if vals or any(text(cell(c)) for c in check):
                        raise TemplateError(f"「{title}」の {top + i} 行目が、書き出したときの行と違う（行を足した・動かした・並べ替えた）。人が埋める欄を、どの行のものか決められない。"
                                            "行を戻すか、値を書き出し直したブックに入れ直す")
                    continue
            else:
                lost += bool(vals)
                continue
            mine[key] = vals
        if lost and not fields and (spec.get("key_check") or {}).get(canon):   # 行の鍵が、列の対応にしか無い一覧（業務マトリクス）
            raise TemplateError(f"「{title}」の人が埋める欄に値のある行が {lost} あるが、どの行のものか決められない（列の対応 {side_path(book).name} が本のとなりに無い）。"
                                "列の対応を本のとなりに置くか、値を書き出し直したブックに入れ直す")
        if lost:
            notes.append(f"WARN  {book.name}「{title}」: 人が埋める欄に値のある {lost} 行は、どの行のものか決まらないので、引き継がない"
                         + ("（番号の欄が空）" if fields else "（この一覧は、書き出し直すと人が埋める欄が空に戻る）"))
        got[canon] = mine
    return got, notes


def read_back(book, spec):
    """書き出した本を、列の対応を通して読み戻す。
    → {"rows": {一覧: [[文字]]}（1行目は欄の名前。欄の決まった順。テンプレートに無い欄は ""）, "cells": {一覧: [[セルか None]]}（データの行だけ）,
       "absent": {一覧: {欄}}, "static": {シート: {セル: (本の文字, 印のある文)}}, "titles": {一覧: シート名}, "side": 列の対応（columns.json）, "book": 開いた本}"""
    from openpyxl import load_workbook
    book = pathlib.Path(book)
    side = read_side(book)
    wb = load_workbook(book, rich_text=True)
    text = lambda v: "" if v is None else str(v).replace("\r", "")
    out = {"rows": {}, "cells": {}, "absent": {}, "static": {}, "titles": {}, "side": side, "book": wb}
    for title, s in side["sheets"].items():
        if title not in wb.sheetnames:
            raise TemplateError(f"シート「{title}」が、本に無い")
        canon, columns = (spec.get("was") or {}).get(s["canon"], s["canon"]), s["columns"]     # 前の名前で書いた一覧も、いまの名前で読む
        fixed, rep = list(spec["sheets"][canon]), (spec.get("repeat") or {}).get(canon)
        at = {f: next((k for k, x in enumerate(columns) if x == f), None) for f in fixed}
        more = [k for k, x in enumerate(columns) if rep is not None and x == rep]
        order = [at[f] for f in fixed] + more
        cells = [[(row[k] if k is not None and k < len(row) else None) for k in order] for row in wb[title].iter_rows(min_row=s["header_rows"] + 1)]
        while cells and all(c is None or c.value in (None, "") for c in cells[-1]):      # 人が埋める列だけに字がある行・書式だけの行
            cells.pop()
        out["cells"][canon] = cells
        out["rows"][canon] = [fixed + [circled(k + 1) for k in range(len(more))]] + [[text(c.value) if c is not None else "" for c in r] for r in cells]
        out["absent"][canon] = {f for f in fixed if at[f] is None}
        out["titles"][canon] = title
    for title, inline in (side.get("static") or {}).items():
        if title in wb.sheetnames:
            out["static"][title] = {at: (text(wb[title][at].value), src) for at, src in inline.items()}
    return out


# ---------- 既定のテンプレート ----------

def make_default(spec, look, path):
    """既定のテンプレートを作る（look は、いままでプログラムの中で決めていた体裁）"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin) if look.get("border") else None
    align = Alignment(wrap_text=True, vertical="top")
    name, size = look.get("font"), look.get("size")
    mk = lambda f: Font(name=name, size=f.get("sz", size), bold=f.get("b"), italic=f.get("i"), color=f.get("color"), underline=f.get("u"))

    def dress(x, font=None, fill=None, head=False):
        if font is not None:
            x.font = font
        if fill:
            x.fill = PatternFill("solid", fgColor=fill)
        if border is not None:
            x.border = border
        if not head or look.get("head_align", True):
            x.alignment = align

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True) if name is None else mk(look["head_font"])
    body_font = None if name is None else Font(name=name, size=size)
    for title in look["order"]:
        ws = wb.create_sheet(title)
        if title in spec["sheets"]:
            fields, wide, heads = spec["sheets"][title], look["widths"][title], look.get("heads") or {}
            omit = (look.get("omit") or {}).get(title, [])           # 既定の様式に、入れない欄（様式に印を置けば出る）
            col = lambda f: (heads.get(f, f), "{" + f + "}", wide[fields.index(f)])
            cols = [col(f) for f in fields if f not in omit]
            cols += [(m, None, wide[len(fields) + j] if len(fields) + j < len(wide) else None) for j, m in enumerate((look.get("manual") or {}).get(title, []))]
            if title in (spec.get("repeat") or {}):
                cols.append((FIRST, "{" + spec["repeat"][title] + "}", None))
            for k, (h, mark, w) in enumerate(cols, 1):
                dress(ws.cell(1, k, h), head_font, look["head_fill"], head=True)
                dress(ws.cell(2, k, mark), body_font)
                ws.column_dimensions[get_column_letter(k)].width = w or look.get("step_width")
            if title in look.get("filter", []):
                ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}2"
        else:                                                         # 文のシート
            st = look["static"][title]
            for k, (h, w) in enumerate(zip(st["head"], look["widths"][title]), 1):
                dress(ws.cell(1, k, h), head_font, look["head_fill"], head=True)
                ws.column_dimensions[get_column_letter(k)].width = w
            for i, row in enumerate(list(st["rows"]) + [[]] * bool(st.get("tail")) + list(st.get("tail") or []), 2):
                for k, (v, f) in enumerate(zip(row, st["fonts"]), 1):
                    dress(ws.cell(i, k, v), mk(look["fonts"][f][1]) if f else body_font)
            if st.get("name"):                                        # 名前の付いた範囲（見出しの行から、表の終わりまで）
                from openpyxl.workbook.defined_name import DefinedName
                nm, width = st["name"]
                wb.defined_names[nm] = DefinedName(nm, attr_text=f"{quoted(title)}!$A$1:${get_column_letter(width)}${1 + len(st['rows'])}")
        ws.freeze_panes = (look.get("freeze") or {}).get(title, "A2")
        page = look.get("page")
        if page:
            ws.page_setup.orientation, ws.page_setup.paperSize = page["orientation"], page["paper"]
            ws.print_title_rows = page["title_rows"]
            if title in page.get("title_cols", {}):
                ws.print_title_cols = page["title_cols"][title]
            if title in page.get("scale", {}):
                ws.page_setup.scale = page["scale"][title]
            else:
                ws.sheet_properties.pageSetUpPr.fitToPage = True
                ws.page_setup.fitToHeight = 0
    if look.get("fills") or look.get("fonts"):
        ws = wb.create_sheet(STYLE_SHEET)
        for k, (h, w) in enumerate((("名前（変えない）", 22), ("見本（ここの塗り・字を変える）", 34), ("使われる所", 80)), 1):
            dress(ws.cell(1, k, h), head_font, look["head_fill"], head=True)
            ws.column_dimensions[get_column_letter(k)].width = w
        r = 2
        for label, (color, where) in look.get("fills", {}).items():
            dress(ws.cell(r, 1, label), body_font)
            dress(ws.cell(r, 2, "塗りの見本"), body_font, color)
            if label in (look.get("aligns") or {}):                   # 字の位置も、この見本から写す塗り
                ws.cell(r, 2).alignment = Alignment(**look["aligns"][label])
            dress(ws.cell(r, 3, where), body_font)
            r += 1
        for label, (_key, f, where) in look.get("fonts", {}).items():
            dress(ws.cell(r, 1, label), body_font)
            dress(ws.cell(r, 2, "字の見本 Aa 123"), mk(f))
            dress(ws.cell(r, 3, where), body_font)
            r += 1
    wb.save(path)


def books():
    """本の種類 → (欄の決まり, 既定の体裁, テンプレートのファイル名)"""
    import build_matrix
    import build_reports
    return {"matrix": (build_matrix.SPEC, build_matrix.LOOK, build_matrix.TEMPLATE), "questions": (build_reports.Q_SPEC, build_reports.Q_LOOK, build_reports.Q_TEMPLATE)}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=["check", "default", "supply"])
    ap.add_argument("template", nargs="?", help="check: 確かめるテンプレート（省くと、assets/ の 2 つ）。supply: 補うテンプレート（業務マトリクス一覧）")
    ap.add_argument("-o", "--out", help="supply: 補ったテンプレートの保存先（省くと、<もとの名前>_処理時間つき.xlsx）")
    ap.add_argument("--book", choices=["matrix", "questions"], help="check でテンプレートを指定したとき: どの本のテンプレートか。default: 作り直す本（省くと、両方）")
    a = ap.parse_args()
    bk = books()
    if a.command == "default":
        for kind, (spec, look, name) in bk.items():
            if not a.book or kind == a.book:
                make_default(spec, look, ASSETS / name)
                print(f"既定のテンプレート → {ASSETS / name}")
        return
    if a.command == "supply":                                         # 基準の表の無いテンプレートに、処理時間に要るものを補って、別の名前で保存する（もとのファイルは変えない）
        if not a.template:
            sys.exit("補うテンプレート（業務マトリクス一覧の xlsx）を書く")
        try:
            tpl = Template(a.template, bk["matrix"][0])
        except TemplateError as e:
            sys.exit(f"FAIL  {e}")
        if tpl.problems:
            sys.exit("\n".join(f"FAIL  {p}" for p in tpl.problems))
        if not tpl.supplied:
            print("補うものは無い（処理時間の基準の表がある）" if tpl.time else "補えなかった: " + "／".join(tpl.notes[-1:]))
            return
        out = pathlib.Path(a.out) if a.out else tpl.path.with_name(tpl.path.stem + "_処理時間つき.xlsx")
        if out.resolve() == tpl.path.resolve() or out.suffix.lower() != ".xlsx":
            sys.exit("保存先は、もとのテンプレートと別の名前の .xlsx にする")
        tpl.wb.save(out)
        print(f"補ったテンプレート → {out}\n{tpl.supplied}")
        return
    if a.template and not a.book:
        sys.exit("テンプレートを指定するときは、--book matrix か --book questions も書く")
    bad = 0
    for kind, (spec, _look, name) in bk.items():
        if a.book and kind != a.book:
            continue
        try:
            tpl = Template(a.template or ASSETS / name, spec)
        except TemplateError as e:
            print(f"FAIL  {e}")
            bad += 1
            continue
        for p in tpl.problems:
            print(f"FAIL  {p}")
        for title, tb in tpl.tables.items():
            print(f"INFO  {tpl.path.name}「{title}」＝一覧「{tb['canon']}」・見出し {tb['sample'] - 1} 行・機械が埋める列 {len(tb['cols'])}")
        for title in tpl.static:
            print(f"INFO  {tpl.path.name}「{title}」＝文のシート（印 {len(tpl.static[title])} か所）")
        if tpl.time:
            print(f"INFO  {tpl.path.name}「{tpl.time['title']}」＝処理時間の基準の表（動作の種類 {len(tpl.time['rows'])} 行。仮の処理時間は、ここを引く式で入る）")
        for n in tpl.notes:
            print(f"INFO  {tpl.path.name}: {n}")
        bad += len(tpl.problems)
    print(f"{bad} FAIL")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
