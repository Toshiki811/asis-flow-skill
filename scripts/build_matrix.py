#!/usr/bin/env python3
"""build_matrix.py — 業務マトリクス一覧（out/業務マトリクス一覧.xlsx）を書き出す。

  python3 scripts/build_matrix.py -w <作業フォルダ>

フロー図と同じデータ（flows/*.flow.json・taxonomy.json・structure/hierarchy.json）から機械で作る、もう1つの見え方。
このブックのために LLM が書くものは無い。直すときは flow.json（直しの一覧で作ったフローは、直しの一覧）・taxonomy.json を直して、図と一緒に書き出し直す。

  業務マトリクス  1行＝1フロー。左に大分類・中分類（分類の段を上から詰める）・分類名の根拠・原典・開始条件・終了条件・フロー全体の確認事項・件数・処理時間の合計、
                  右に作業を ①②③… と並べる。作業の行のすぐ下に、処理時間（分）の行（人が数を打つ所。打っていない所は、仮の値＝「処理時間の基準」を引く式。件数も、この行）
  業務フローのノード一覧  1行＝業務フロー図の1ノード（絞り込み用。開始・終了・帳票も1行）。線は「行き先」の列に、1本を1行で書く
  システム一覧    1行＝1システム（用語の正本のシステムすべて）。そのシステムを使う作業（札）と、システムが行う処理（箱）の場所
  起動点一覧      1行＝プログラムの起動点1つ（ボタン・イベント・バッチ）。既定の様式には入れない（同じ表が、アクティビティ図の1ページ目にある）。様式にシートを置けば出る
  プログラムのノード一覧  1行＝アクティビティ図の箱1つと、そのソースの場所。図にしなかった文（対象外）も、理由つきで1行
  プログラムの図にしない手続き  1行＝図にしなかった補助の手続き1つ（理由・呼ぶ側の図・ソースの場所）
  プログラムの用語対応  1行＝プログラムの中だけの名前1つと、アクティビティ図に書いた業務の言葉
  縦軸の根拠      業務の木（目次）の項目すべてと、目次に無く足した分類の、扱いと理由
  資料一覧        資料の題名・ファイル・役割・ページ番号の数え方・資料どうしの関係（資料は、どのシートでも題名で書く。資料ID は出さない）
  読み方          印・色の意味と、数（文は、テンプレートのもの）
プログラム用の一覧（起動点一覧・プログラムの…）は、中身が1行も無ければ（プログラムを読んでいない）、ブックに入れない。
「開始条件」「終了条件」の欄の印は {きっかけ}・{成果}（前の様式が、そのまま使える）。前の名前のシート（ノード一覧・処理ノード一覧・図にしない手続き・用語の対応）も、そのまま受ける。

シートの並び・列の順・見出し・幅・色・字は、テンプレート（assets/業務マトリクス一覧_テンプレート.xlsx）が決める（table_template.py）。
ここが組み立てるのは、中身（一覧 → 行 → セル）だけ。--template で、別のテンプレートを使える。

処理時間（分）は、人が入れた値が正。書き出す前に、いま out にあるブックから、人が入れた値（時間・件数・基準の表の直し）を times.json に取り込み、書き出しで戻す。
  --keep <ブック.xlsx>   お客様から戻ってきたブックから、取り込む（入っている値を足す・書き換える）
  --calibrate            人が入れた時間が 3 つ以上ある「動作の種類 × 量」は、その中央値を、基準の表に入れる（残りの仮の値が、実態に寄る）
  --accept               作業名が変わった箱の、人が入れた時間を、そのまま認める
  --forget               無くなった箱の、人が入れた時間を、times.json から消す
  --fresh                いま out にあるブックから、取り込まない（読めないブックを、作り直すとき）
仮の値は、動作の種類（flow.json の action）× 量（volume。無ければ、同じページからまとめた手順の数・見る帳票の数から）で、基準の表を引く式。分の数字は、表にしか無い。
基準の表（名前「処理時間の表」）の無いテンプレートには、既定のテンプレートから、基準の表・処理時間の行の書式・処理時間の欄を補って、書き出す（table_template.py の supply_time）。
最初の出力から、すべての箱に時間が出る（人の作業・判断は、仮の値。システムの自動処理・別のフローへ続く箱は 0）。式のセルには、計算の結果も保存する（計算をしない表示でも、値が出る）。

列の順は、枝ごとに続けて並べる（common.step_order_by_branch）。左隣が前の作業でないセルには、入口（← ④「No」）を書く。
行き先は省かない（判断のすべての枝・次の列でない行き先・戻り・終了）。次の列へ進むだけの線は書かない。
確認事項の番号は、確認事項一覧・図のバッジと同じ（questions.json）。そのフローの確認事項は、どれも行のどこかに出る。
"""
import argparse
import datetime
import json
import pathlib
import re
import statistics
import unicodedata

from build_reports import KIND, Reports
from common import (ACTIONS, REL_LABEL, STEP_TYPES, SYSTEM_KIND, TYPE_ACTION, VOLUME_SPARE, VOLUMES, Answers, BadJson, Times, answer_gaps, category_path, flow_key, locate, node_clues, node_details, order_of, own_number,
                    question_ids, step_order_by_branch, system_ledger, time_basis, tree_nodes, whole, work_dir, unlinked_note, ASSETS)
from table_template import (Cell, Template, TemplateError, circled, clean, fill as fill_template, lock_of, manual_back, remembered,
                            side_path, time_table)
from validate import blocking

BOOK = "業務マトリクス一覧.xlsx"
SHEETS = ["業務マトリクス", "業務フローのノード一覧", "システム一覧", "起動点一覧", "プログラムのノード一覧", "プログラムの図にしない手続き", "プログラムの用語対応", "縦軸の根拠",
          "資料一覧", "読み方"]
WAS_SHEET = {"ノード一覧": SHEETS[1], "処理ノード一覧": SHEETS[4], "図にしない手続き": SHEETS[5], "用語の対応": SHEETS[6]}   # 前の名前（前の様式のシート名・前に書き出したブックの列の対応）
PROGRAM = SHEETS[3:7]                                                 # プログラム用の一覧（中身が1行も無ければ、ブックに入れない）
MATRIX_HEAD = ["大分類", "中分類", "分類名の根拠", "原典（資料とページ）", "きっかけ", "成果", "フロー全体の確認事項", "件数", "処理時間の合計"]   # きっかけ・成果の見出しは、開始条件・終了条件（LOOK）
NODE_HEAD = ["大分類", "中分類", "フローID", "順番", "ノード番号", "種別", "担当", "作業", "補足", "参考", "システム", "入口", "行き先", "帳票",
             "出典", "原文", "推定", "補足由来", "確認事項", "動作の種類", "量", "処理時間（分）", "時間の根拠"]
SYSTEM_HEAD = ["番号", "システム", "別の呼び方", "種類", "コード", "起動点", "使う作業（札）", "システムが行う処理（箱）", "備考"]
CODE_HEAD = ["アクティビティ図", "図の名前", "ノード番号", "種別", "レーン", "処理", "補足", "行き先", "読む・書く", "ソースの場所", "ソースの先頭の行", "確認事項"]
HELPER_HEAD = ["システム", "手続き", "扱い", "理由", "呼ぶ側の図", "呼ぶ所", "ソースの場所", "実行文の行数"]
TERM_HEAD = ["システム", "もとの名前", "業務の言葉", "根拠", "定義の場所", "使っている図", "備考"]
ENTRY_HEAD = ["起動点", "システム", "きっかけ", "手続き", "アクティビティ図", "使われる業務の箱", "紐付けの根拠", "確認事項", "備考"]
AXIS_HEAD = ["段", "項目", "出どころ", "扱い", "フロー（行）", "理由", "ページ"]
DOC_HEAD = ["略称", "ファイル", "役割", "ページ数", "ページ番号の数え方", "備考", "資料どうしの関係"]   # 印は {略称} のまま（前の様式が使える）。中身は資料の題名（既定の見出しは「資料の題名」）
INFERRED, EDGE_KIND = "（推定）", {"system": "（システム連携）", "data": "（データ）"}
MARK = {"decision": "◇ ", "subprocess": "⇒ "}                       # 作業名の前に付ける印
TREE_SOURCE = {"toc": "目次", "folder": "フォルダ", "heading": "見出し", "inferred": "推定"}
NAME_REF = {"heading": "本文の見出し", "title": "資料名"}
# 縦軸の扱い（閉じた6種）
DRAW, GROUP, INSIDE, SPLIT, OUT, ADDED = "フローを描く", "束ねる見出し", "上位のフローに含む", "下位フローに分けた", "対象外（業務でない）", "目次に無く足した"
UNSET = "未設定"
Q_RE = re.compile(r"Q-\d+")
# 「行き先」の1行（ノード一覧）: 「ラベル」→（推定）行き先のノード番号 …。点検（check_flow.py --stage matrix）が、ここから線を組み直す
EXIT_RE = re.compile(r"^(?:「(?P<label>.*?)」)?→(?P<marks>(?:（推定）|（システム連携）|（データ）)*) ?(?P<to>\S+)")
# 塗りの種類（色は、テンプレートの「書式の見本」が決める）
FILL = {k: k for k in ("head", "decision", "subprocess", "system", "inferred", "file", "time")}
# 処理時間: 基準の表（テンプレートの、名前の付いた範囲）と、人が書き換えてよいセルの札（コメントの 1 行目。読み戻しは、これで突き合わせる）
TIME_NAME, TIME_ACTIONS = "処理時間の表", ACTIONS + list(TYPE_ACTION.values())
TIME_WAS = {"判断する": (0,)}                                         # 前の版の、既定の値。書き出したときの表の値を残していないブック（前の版）に、この値があっても、人が直した値と見ない
TAG_TIME, TAG_COUNT, TAG_LABEL = re.compile(r"^処理時間 (\S+) (\S+)｜"), re.compile(r"^件数 (\S+)｜"), re.compile(r"^作業: (.*)$", re.M)
TAG_CLUE = re.compile(r"^担当: (.*?)｜前: (.*)$", re.M)                 # 時間のセルのコメントの、付け替えの手がかり（担当・直前の箱。同じ作業名の箱を見分ける）
OURS = re.compile(r'^=VLOOKUP\("')                                    # 仮の値の式（基準の表を引く）
TEMPLATE = "業務マトリクス一覧_テンプレート.xlsx"
M, NODES = SHEETS[0], SHEETS[1]                                      # 業務マトリクスのシート（式の番地は、ここを指す）・業務フローのノード一覧
STEP = "作業"                                                         # 業務マトリクスの、繰り返す列（①②③…）の印
VALUES = ["フローの数", "作業の数", "帳票の数", "推定のノードの数", "推定の線の数", "確認事項の数", "フローに掛かる確認事項の数", "それ以外の確認事項の数"]


def lines(parts, style="text"):
    """[文字, …] を改行でつないだ runs（空の行は落とす）"""
    out = []
    for p in parts:
        runs = [(p, style)] if isinstance(p, str) else list(p)
        if not "".join(t for t, _s in runs):
            continue
        out += ([("\n", "text")] if out else []) + runs
    return out


# ---------- 欄の決まり（テンプレートの印）と、既定のテンプレートの体裁 ----------

TABLES = SHEETS[:-1]                                                  # 表の一覧（読み方は、文のシート）
HEADS = dict(zip(TABLES, (MATRIX_HEAD, NODE_HEAD, SYSTEM_HEAD, ENTRY_HEAD, CODE_HEAD, HELPER_HEAD, TERM_HEAD, AXIS_HEAD, DOC_HEAD)))
# 「書式の見本」の行: (名前, 鍵, 既定の色, 使われる所)／(名前, 鍵, 既定の字, 使われる所)
FILLS = [("判断", "decision", "FFF2CC", "判断の作業。「縦軸の根拠」の、目次に無く足した分類"),
         ("別のフローへ", "subprocess", "E2EFDA", "別のフローへ続く作業"),
         ("システム処理", "system", "DDEBF7", "システムが自動で行う処理"),
         ("推定", "inferred", "FCE4D6", "推定のノード。つづりから読んだ言葉・読めない名前。台帳に無いシステム。未設定の項目"),
         ("帳票", "file", "F2F2F2", "「業務フローのノード一覧」の帳票の行。対象外・図にしないものの行"),
         ("処理時間の行", "time", "F7F7F7", "業務マトリクスの、処理時間（分）の行。字の位置（中央）も、この見本のもの"),
         ("処理時間（長い作業）", "time_long", "E4DFEC", "量が「特大」（時間・日の単位でかかる作業）で、まだ仮の値のセル。先に、実際の時間を入れる所。条件付き書式で付く（数を打つと、消える）")]
FONTS = [("担当", "actor", dict(sz=8, color="1F4E79"), "作業のセルの1行目（フロー図のレーン）"),
         ("入口", "entry", dict(sz=8, color="7F7F7F"), "作業のセルの「← ④」"),
         ("作業の名前", "label", dict(sz=10, b=True), "作業のセルの、作業の名前"),
         ("行き先", "exit", dict(sz=9), "作業のセルの「→ ⑤」"),
         ("出典", "src", dict(sz=8, color="7F7F7F"), "作業のセルの末尾のページ。分類名の根拠の場所"),
         ("確認事項の番号", "q", dict(sz=8, b=True, color="C55A11"), "Q-012 の形の番号"),
         ("推定の印", "inferred", dict(sz=9, b=True, color="C00000"), "（推定）の字"),
         ("薄い字", "pale", dict(sz=9, color="A6A6A6"), "上の行と同じ大分類。対象外の理由"),
         ("太字", "bold", dict(sz=10, b=True), "大分類・中分類・システムの名前・番号"),
         ("リンク", "link", dict(sz=10, b=True, color="0563C1", u="single"), "フロー図（PNG）を開ける、中分類の名前"),
         ("種別", "kind", dict(sz=9, b=True), "扱い・種別。「読み方」の項目"),
         ("札", "sys", dict(sz=8, color="0066CC"), "作業のセルの［◯◯システム］"),
         ("処理時間", "time", dict(sz=10), "人が入れた処理時間・件数（数を打ったセル）"),
         ("処理時間（仮）", "time_soft", dict(sz=10, i=True, color="808080"), "仮の処理時間（基準の表を引く式のセル）。条件付き書式で付くので、効くのは、太字・斜体・色だけ"),
         ("処理時間（仮・長い作業）", "time_long", dict(sz=10, b=True, i=True, color="7030A0"), "量が「特大」で、まだ仮の値のセルの字。塗りは、「処理時間（長い作業）」の行。効くのは、太字・斜体・色だけ")]
SPEC = {"sheets": HEADS, "repeat": {SHEETS[0]: STEP}, "values": VALUES,
        "required": {"業務マトリクス": ["大分類", "中分類", "きっかけ", "フロー全体の確認事項", STEP], NODES: ["大分類", "中分類", "フローID", "ノード番号", "行き先"]},
        "was": WAS_SHEET, "drop_empty": PROGRAM,
        "keys": {NODES: ["フローID", "ノード番号"]}, "key_check": {"業務マトリクス": ["大分類", "中分類"]},   # 人が埋める欄を戻す行の鍵（業務マトリクスは、作る側が渡すフロー）
        "quiet": {"lists": [SHEETS[3]], "fields": {"システム一覧": ["起動点"]}},          # 既定の様式に、わざと入れていないもの（無くても、知らせない）
        "fills": {label: key for label, key, _c, _w in FILLS}, "fonts": {label: key for label, key, _f, _w in FONTS},
        "aligns": ["time"], "time": {"name": TIME_NAME, "actions": TIME_ACTIONS, "volumes": VOLUMES, "styles": ["time", "time_soft"],
                                "spare": VOLUME_SPARE, "soft": ["time_long"],
                                "default": ASSETS / TEMPLATE, "before": "読み方",         # 基準の表の無いテンプレートに、補うもとのテンプレート・表のシートを置く場所（このシートの前）
                                "fields": {"業務マトリクス": ["件数", "処理時間の合計"], NODES: ["動作の種類", "量", "処理時間（分）", "時間の根拠"]},
                                "heads": {"処理時間の合計": "処理時間の合計（分・全作業）", "件数": "件数（人が入れる）"}}}                # 無くてもよい列（無ければ、代わりの列を引く）・無くてもよい書式（前の形のテンプレートを、止めない）
# 基準の表の、既定の値（分）。裏付けのある数ではない（目安）。使う側が、実際の時間に直して使う
TIME_ROWS = [("入力する", 2, 4, 10, 60, "システム・台帳に、入力する・登録する"), ("作成する", 5, 15, 40, 240, "書類・資料を作る"), ("確認する", 2, 5, 15, 120, "内容を確かめる・照合する・点検する"),
             ("承認する", 1, 3, 10, 30, "承認する・決裁する・押印する"), ("提出・回付する", 1, 2, 5, 15, "書類を出す・回す・送る"), ("受け取る", 1, 2, 5, 15, "書類・連絡を受け取る"),
             ("出力する", 1, 2, 5, 15, "印刷する・出力する"), ("連絡する", 2, 5, 15, 60, "知らせる・依頼する・報告する"), ("計算する", 3, 10, 30, 120, "計算する・集計する"),
             ("調べる", 5, 15, 45, 240, "調査する・検討する・選定する"), ("保管する", 1, 2, 5, 30, "整理する・保存する・廃棄する"),
             ("その他", 2, 5, 15, 120, "上のどれにも当たらない作業。種類を決めていない作業"),
             ("判断する", 2, 2, 2, 2, "判断（ひし形）。読んで、どちらへ進むかを決める時間（決めの値。量は分けないので、どの列も同じ数）。前の作業に含めて数えるなら、0 にする"), ("システムが処理する", 0, 0, 0, 0, "システムの自動処理。人の時間でない"),
             ("別のフローへ進む", 0, 0, 0, 0, "別のフローへ続く箱。時間は、行き先のフローで数える")]
TIME_TAIL = [("使い方", "", "", "", "", "業務マトリクスの仮の処理時間（灰色の斜体）は、この表を「動作の種類 × 量」で引いた値。ここの数を直すと、その種類の仮の値が、すべて変わる（人が数を打ったセルは、変わらない）。"),
             ("数について", "", "", "", "", "既定の数は、目安（裏付けのある数ではない）。実際の時間に合わせて、直して使う。直した数は、書き出し直しても残る。"),
             ("量", "", "", "", "", "同じページの手順を 3 つ以上まとめた箱・見る帳票が 2 つ以上の箱は「大」、ほかは「中」。資料から量が言えるときだけ、小・大・特大に変えてある（わけは、時間のセルのコメント）。"),
             ("特大", "", "", "", "", "時間・日の単位でかかる作業（資料から、そう読めるものだけ）。数は、とくに粗い目安。業務マトリクスで色の付いた時間のセルは、まだ仮の値の「特大」。"
                                      "先に、実際の時間を入れる（分で入れる。1 時間は 60 分。1 日を何分と数えるかは、使う側で決める）。"),
             ("変えないもの", "", "", "", "", "1 列目の語と、見出しの 小・中・大・特大（式が、この語で引く）。行の順・説明・見た目は、変えてよい。")]
LEGEND = [
    ("このブック", "as-is フロー図（業務フロー_as-is.drawio）と同じデータから機械で作った一覧。出典・確認事項・推定の印は、図と同じものが出る。"
                   "直すときは、もとのデータ（flow.json・taxonomy.json）を直して、図と一緒に書き出し直す。このブックに書き込んだ内容は、書き出し直すと消える"
                   "（処理時間・件数・「処理時間の基準」の数と、様式で足した、人が埋める欄（業務マトリクス・業務フローのノード一覧）は、引き継ぐ）。"),
    ("数", "フロー {フローの数} 本／作業（列になるノード）{作業の数}／帳票 {帳票の数}／"
           "推定のノード {推定のノードの数}・推定の線 {推定の線の数}／"
           "確認事項 {確認事項の数} 件（フローに掛かるもの {フローに掛かる確認事項の数}・それ以外 {それ以外の確認事項の数}。それ以外は「確認事項一覧」を参照）"),
    ("行", "1フロー＝2行（作業の行と、そのすぐ下の、処理時間（分）の行）。業務分類の順（下位フローは、親のフローの直後）。"),
    ("大分類・中分類", "業務分類の段を、上から詰めて入れる。大分類＝いちばん上の段、中分類＝その下の段（3段以上あるときは、下の段を「›」でつなぐ。「›」は「〜の中の」）。"
                       "1段しか無いときは、大分類に名前、中分類は「—」。行の名前（いちばん下の段）から、フロー図（PNG）を開ける。"),
    ("分類名の根拠", "目次＝資料の目次・業務一覧表にある名前／本文の見出し＝目次に無いが、見出しにある名前／資料名＝資料の題名・ファイル名から／"
                     "LLM が命名＝原典に名前が無い（確認事項で確かめる）。下の行は、その名前が書かれている場所。"),
    ("原典（資料とページ）", "そのフローの作業が出典にしている資料とページ。資料が複数のときは、引いているノードの多い資料から並べ、〔 〕にその数を書く。"
                             "資料どうしの関係（上位・下位、別紙、根拠）があれば添える。資料は、題名で書く（表紙の題名。表紙の無い資料は、ファイル名）。ファイルの場所と役割は「資料一覧」。"),
    ("開始条件・終了条件", "その業務が、何を受けて始まるか（開始条件）と、何ができたら終わるか（終了条件）。原典に書かれていなければ「（記載なし）」にして、確認事項で聞く。"),
    ("①②③…", "作業を、枝ごとに続けて並べた順。判断で分かれたら、片方の枝を終わりまで並べてから、もう片方を並べる。番号はこの表の中だけのもの"
               "（フロー図のノード番号 P-001 などは、セルのコメントと「業務フローのノード一覧」にある）。"),
    ("セルの1行目", "担当（フロー図のレーン）。"),
    ("← ④「No」", "入口。左隣の列が前の作業でないとき・前の作業が複数あるときに、どの列から来るかを書く。書いていないセルは、左隣の列から続く。"
                  "「（戻り）」は、あとの列から戻ってくる線。"),
    ("◇", "判断。行き先は、すべての枝を「ラベル」→ 行き先 で書く。"),
    ("⇒", "別のフローへ続く作業。"),
    ("→ ⑤／→ 終了", "行き先。次の列へ進むだけのときは書かない。「③へ戻る」は差し戻しなどの戻り。"),
    (INFERRED, "原典に書かれておらず、推定で置いたノード・線。必ず確認事項がある。フロー図では赤い破線。"),
    ("Q-012", "確認事項の番号（「確認事項一覧」と同じ）。作業に掛かるものはセルの末尾、線に掛かるものは行き先の後ろ、"
              "フロー全体に掛かるものは「フロー全体の確認事項」の列に出る。"),
    ("セルの末尾", "出典の資料とページ（資料の題名と、資料に刷られたページ番号。PowerPoint はスライド番号、Excel はシート名）。原文は、セルのコメントと「業務フローのノード一覧」。"),
    ("［会計システム］", "札。その作業で使うシステム（原典の手順に書かれているもの）。フロー図では、箱の中の青い字。どのシステムがどの作業で使われるかは「システム一覧」。"),
    ("色", "業務マトリクス: 青＝システムが自動で行う処理（プログラムが動く箱）／黄＝判断／緑＝別のフローへ続く作業／赤みのある色＝推定のノード／"
           "薄い灰の行＝処理時間（分）の行／紫＝量が「特大」で、まだ仮の値の時間。"
           "業務フローのノード一覧: 黄＝判断／灰＝帳票／赤みのある色＝推定。システム一覧: 赤みのある色＝用語の正本に無い名前。"
           "プログラムのノード一覧: 黄＝判断。プログラムの用語対応: 赤みのある色＝つづりから読んだ言葉・読めない名前（推定）。"
           "縦軸の根拠: 黄＝目次に無く足した分類／灰＝対象外／赤みのある色＝扱いが未設定。"),
    ("処理時間（分）の行", "作業の行のすぐ下の行。その作業を 1 回行うのに、手を動かす時間（分）。待ち時間は含めない。"
                           "灰色の斜体＝仮の値（「処理時間の基準」の表を、動作の種類 × 量で引いた式）。数を打つと、その値が正になり、字が黒くなる。"
                           "色の付いたセル＝量が「特大」（時間・日の単位でかかる作業）で、まだ仮の値。表の値が、とくに粗いので、先に、実際の時間を入れる（分で入れる）。"
                           "書き出し直しても、打った数は残る（消すと、仮の値に戻る）。動作の種類と量は、セルのコメントと「業務フローのノード一覧」。"),
    ("処理時間の合計", "その行の、すべての作業の時間の合計（分）。判断で分かれた両側の作業を、どちらも足している（1 件あたりの時間より、大きくなることがある）。"),
    ("件数", "処理時間の行の、合計の左。人が入れる欄（数え方は、決めて使う。月の件数など）。仮の値は入れない。書き出し直しても残る。"),
    ("処理時間の基準", "仮の処理時間のもとになる表。動作の種類（行）× 量（小・中・大・特大）の分。数を直すと、その種類の仮の値が、すべて変わる。"
                       "既定の数は、目安（裏付けのある数ではない）。判断（ひし形）は、決めの値（既定は 2 分）。"
                       "システムの自動処理・別のフローへ続く箱は、0 分（人の時間でない／行き先のフローで数える）。"),
    ("業務フローのノード一覧", "1行＝業務フロー図の1ノード（人の作業・判断・開始・終了・帳票も1行）。絞り込み・並べ替え用。「行き先」は線1本を1行で書き、行き先のノード番号を添える。"
                   "「処理時間（分）」は、業務マトリクスのセルを指す式（数を打つのは、業務マトリクスの側）。「時間の根拠」は、仮＝基準の表の値／人＝人が入れた値。"),
    ("システム一覧", "1行＝1システム（原典に出てくるシステムすべて）。「使う作業（札）」は、人がそのシステムを使う作業のノード。"
                     "「システムが行う処理（箱）」は、システムのレーンに置いた、システムが自動で行う処理のノード。"),
    ("▶ A-02-1", "アクティビティ図の番号。プログラム（マクロ・スクリプト）が動く所に置いたシステム処理の箱に付く。中身は「アクティビティ図.drawio」の、その番号のページ。"
                 "番号は、「システム一覧」の番号（02）と、そのシステムの中での連番。"),
    ("プログラムのシート", "「プログラムのノード一覧」「プログラムの図にしない手続き」「プログラムの用語対応」は、プログラム（マクロ・ソース）を読んだときだけ出る（中身の無いシートは、入れない）。"
                           "業務フローの箱とは、「▶ A-01-2」の番号でつながる。プログラムの起動点（ボタン・イベント・バッチ）の一覧は、アクティビティ図の1ページ目。"),
    ("プログラムのノード一覧", "1行＝アクティビティ図（プログラムの中身）の箱1つ（開始・終了・帳票を除く）と、そのソースの場所（資料・モジュール・行）。"
                       "種別が「対象外」の行は、図にしなかった文（画面の更新の制御・変数の準備・後片付け など。理由が「処理」の列にある）。"
                       "プログラムの実行文は、どれも、この表のどこかの行に入っている。"),
    ("プログラムの図にしない手続き", "1行＝図にしなかった補助の手続き1つ。ほかの処理から呼ばれる共通の処理のうち、シート・ファイル・テーブル・画面のメッセージに触れず、"
                                   "値を探す・寸法を計算する・文字を整える、のような補助だけをするもの。呼び出しは、呼ぶ側の図の箱（何のために呼ぶかを、名前に書いた）に入っている。"
                                   "プログラムを読む段で、機械が候補を出し、決まった理由に当たるものだけを図にしない（書き出すたびに、台帳から作り直す。人が書き込む表ではない）。"),
    ("プログラムの用語対応", "1行＝プログラムの中だけの名前（変数・項目・手続き・定数）1つと、アクティビティ図に書いた業務の言葉。図の箱には、業務の言葉だけを書いている。"
                   "「根拠」は、コメント（ソースのコメントに書いてある）／定義（画面のラベル・定数の値などから読める）／つづり（推定。名前のつづりから読んだ。色つき。"
                   "確認事項で確かめる）／不明（読めない。図には、もとの名前のまま書いた）。業務での正しい呼び方が違えば、確認事項と同じように知らせる（この表は、書き出すたびに作り直す）。"),
    ("縦軸の根拠", "業務の木（目次）の項目すべてについて、フローを描く／束ねる見出し／上位のフローに含む／下位フローに分けた／対象外（業務でない）／"
                   "目次に無く足した、のどれとして扱ったかと、その理由。"),
]
TIME_SHEET = "処理時間の基準"
LOOK = {"font": "游ゴシック", "size": 9, "border": True, "head_fill": "D9E1F2", "head_font": dict(sz=10, b=True), "order": [s for s in SHEETS[:-1] if s != SHEETS[3]] + [TIME_SHEET, SHEETS[-1]],   # 起動点一覧は入れない（アクティビティ図の1ページ目にある。様式に置けば出る）
        "step_width": 30, "omit": {"システム一覧": ["起動点"]}, "aligns": {"処理時間の行": dict(horizontal="center", vertical="center")},
        "heads": {"きっかけ": "開始条件", "成果": "終了条件", "処理時間の合計": "処理時間の合計（分・全作業）", "件数": "件数（人が入れる）", "略称": "資料の題名"},
        "widths": {"業務マトリクス": [20, 34, 24, 34, 26, 26, 12, 10, 12], TIME_SHEET: [20, 7, 7, 7, 7, 80],
                   NODES: [18, 30, 9, 6, 9, 10, 16, 34, 40, 30, 18, 22, 34, 22, 22, 50, 6, 8, 14, 14, 14, 10, 8],
                   "システム一覧": [9, 28, 28, 16, 22, 8, 50, 50, 40], SHEETS[3]: [10, 24, 44, 28, 12, 24, 44, 16, 36],
                   SHEETS[4]: [12, 28, 9, 12, 16, 40, 30, 34, 34, 34, 50, 14], SHEETS[5]: [26, 40, 16, 28, 24, 44, 34, 10],
                   SHEETS[6]: [26, 30, 34, 16, 34, 26, 40],
                   "縦軸の根拠": [5, 44, 9, 20, 40, 50, 22], "資料一覧": [32, 44, 14, 9, 30, 30, 44], "読み方": [24, 96]},
        "freeze": {"業務マトリクス": "C2", NODES: "F2"},
        "filter": [s for s in SHEETS[:-2] if s != SHEETS[3]],
        "page": {"orientation": "landscape", "paper": 8, "title_rows": "1:1", "title_cols": {"業務マトリクス": "A:B"}, "scale": {"業務マトリクス": 60}},   # A3 横
        "fills": {label: (color, where) for label, _k, color, where in FILLS}, "fonts": {label: (key, font, where) for label, key, font, where in FONTS},
        "static": {"読み方": {"head": ["項目", "説明"], "rows": LEGEND, "fonts": ["種別", None]},
                   TIME_SHEET: {"head": ["動作の種類"] + VOLUMES + ["説明"], "rows": TIME_ROWS, "tail": TIME_TAIL, "fonts": ["種別"] + [None] * (len(VOLUMES) + 1),
                                "name": (TIME_NAME, 1 + len(VOLUMES))}}}


def number(v):
    """人がセルに入れたものを、数として読む（7・「７」・「7分」・「1.5」）。読めなければ None"""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return v if v >= 0 else None
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(?:分|件|回)?", unicodedata.normalize("NFKC", str(v)).strip().replace(",", ""))
    return float(m.group(1)) if m else None


def pull(times, book, flows, replace, base):
    """ブックに人が入れた値（処理時間・件数・基準の表の直し）を、times に取り込む。返すのは、知らせること（文の並び）。
    replace: いま out にあるブック（最後に書き出したもの）なら True — 札のあるセルが、式・空に戻っていれば、その値を消す。
    お客様から戻ってきた別のブック（--keep）なら False — 入っている値を足す・書き換えるだけ。base は、テンプレートの基準の表の値。
    だれのセルかは、セルのコメント（「処理時間 F03 P-003｜」「件数 F03｜」）で突き合わせる（行・列を入れ替えられても合う）。
    番号を振り直したあとの古いブックは、コメントの作業名で、いまのノードを探す"""
    from openpyxl import load_workbook
    book = pathlib.Path(book)
    wb, shown = load_workbook(book), load_workbook(book, data_only=True)
    today, said, unread, n = datetime.date.today().isoformat(), [], [], {"時間": 0, "件数": 0, "基準の表": 0, "消した値": 0}
    label = {(flow_key(fl), x["id"]): x.get("label", "") for fl in flows for x in fl["nodes"] if x.get("type") in STEP_TYPES}
    clues = node_clues(flows)

    def value(ws, c):
        """(人が入れた数, 読めたか)。式・空は (None, True)"""
        v = c.value
        if v is None or (isinstance(v, str) and (not v.strip() or OURS.match(v))):
            return None, True
        x = number(shown[ws.title][c.coordinate].value if isinstance(v, str) and v.startswith("=") else v)   # 人が書いた式は、保存されている計算の結果を読む
        if x is None:
            unread.append(f"{ws.title}!{c.coordinate}「{str(v)[:20]}」")
        return x, x is not None

    def take(store, key, entry, x):
        if x is None:
            if replace and key in store:
                del store[key]
                n["消した値"] += 1
            return False
        if store.get(key, {}).get("value") != whole(x):
            store[key] = dict(entry, value=whole(x), book=book.name, date=today)
            return True
        return False

    def one(ws, c, kind, F, i="", was="", clue=None):
        x, ok = value(ws, c)
        if not ok:
            return
        if kind == "count":
            n["件数"] += take(times.counts, F, {"flow": F}, x)
            return
        now = label.get((F, i))
        to = locate(clues, F, i, dict(clue or {}, label=was)) if was else None
        if was and now == was:                                        # 作業名は同じ。同じ名前の箱が 2 つ以上あって、番号が入れ替わっていれば、手がかり（担当・直前の箱）の合うほうへ
            i = to or i
        elif was:                                                     # ブックを書き出したあとで、フローが変わっている
            if to is not None:                                        # 番号がずれた: 同じ作業名のノードへ（2 つ以上あれば、担当 → 直前の箱の順に絞る）
                i = to
            elif now is not None and claimed.get((F, now), set()) - {i}:   # その番号は、いまは、ブックの別のセルの作業のもの。だれの値か、決められない
                if x is not None:
                    lost.append(f"{c.coordinate}（{F} {i}「{was}」{whole(x)} 分）")
                return
        n["時間"] += take(times.cells, (F, i), {"flow": F, "node": i, "label": was or now or ""}, x)

    tags, used, bare, lost = [], set(), 0, []                         # 札の付いたセル: (シート, セル, 種類, フロー, ノード, 作業名)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.comment is None:
                    continue
                text = str(c.comment.text or "")
                m, k = TAG_TIME.match(text), TAG_COUNT.match(text)
                if m:
                    lab, cl = TAG_LABEL.search(text), TAG_CLUE.search(text)
                    clue = {"lane": cl.group(1).strip(), "after": [a for a in cl.group(2).strip().split("｜") if a]} if cl else None
                    tags.append((ws, c, "time", m.group(1), m.group(2), lab.group(1).strip() if lab else "", clue))
                elif k:
                    tags.append((ws, c, "count", k.group(1), "", ""))
                if m or k:
                    used.add((ws.title, c.coordinate))
    side = json.loads(side_path(book).read_text(encoding="utf-8")) if side_path(book).exists() else {}
    if side:                                                          # コメントの消えたセル: 書き出したときの番地で突き合わせる
        found = {t[2:5] for t in tags}
        for title, sh in (side.get("sheets") or {}).items():
            for where, tag in (sh.get("tags") or {}).items():
                tag = (list(tag) + ["", ""])[:4]
                if title in wb.sheetnames and tuple(tag[:3]) not in found and (title, where) not in used:
                    tags.append((wb[title], wb[title][where], *tag))
                    bare += 1
        if bare:
            said.append(f"{book.name}: コメント（だれの時間かの札）の消えた時間・件数のセル {bare} 個は、書き出したときの番地で突き合わせた（行・列を動かしていると、ずれる）")
    titles = {canon: next((t for t, s in (side.get("sheets") or {}).items() if s.get("canon") == canon), canon) for canon in (M, NODES)}
    stray = []                                                        # 処理時間の行の、札で突き合わせられない数（だれの時間か決められない）
    known = {(t, w) for t, s in (side.get("sheets") or {}).items() for w in (s.get("tags") or {})}
    if titles[M] in wb.sheetnames:
        for row in wb[titles[M]].iter_rows():
            if any(isinstance(c.value, str) and c.value.endswith("｜処理時間（分）") for c in row):
                stray += [c.coordinate for c in row if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
                          and (titles[M], c.coordinate) not in used and (titles[M], c.coordinate) not in known]
    if stray:
        raise SystemExit(f"{book.name}: 処理時間の行に、だれの時間か決められない数がある（{'・'.join(stray[:8])}{' …' if len(stray) > 8 else ''}。"
                         "セルのコメント（札）が消えていて、書き出したときの列の対応（" + side_path(book).name + "）も無い）。"
                         "人が入れた値を落とさないよう止める。列の対応をブックのとなりに置くか、書き出し直したブックに入れ直す")
    if titles[NODES] in wb.sheetnames:                                # 業務フローのノード一覧の時間の列は、業務マトリクスの値を写す式。数を打つと、書き出し直したときに消える
        ws = wb[titles[NODES]]
        cols = {c.column for row in ws.iter_rows() for c in row if isinstance(c.value, str) and re.match(rf"^='?{re.escape(titles[M])}'?!", c.value)}
        typed = [c.coordinate for row in ws.iter_rows(min_row=2) for c in row
                 if c.column in cols and isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
        if typed:
            raise SystemExit(f"{book.name}「{titles[NODES]}」の {'・'.join(typed[:8])}{' …' if len(typed) > 8 else ''} に、数が打ってある。"
                             f"ここは「{titles[M]}」の処理時間を写す式で、書き出し直すと消える。値は「{titles[M]}」の処理時間の行（その作業のセルの真下）に入れ直してから、もう一度")
    claimed = {}                                                      # ブックの中で、その作業名を持っている番号
    for _ws, _c, kind, F, i, was, *_clue in tags:
        if kind == "time" and was:
            claimed.setdefault((F, was), set()).add(i)
    for t in tags:
        one(*t)
    if lost:
        said.append(f"WARN  {book.name}: 次の値は、取り込んでいない（ブックを書き出したあとで、番号がずれ、作業名も変わっていて、どの作業の値か決められない。書き出し直したブックに、入れ直す）: "
                    + "、".join(lost[:8]) + (f" ほか {len(lost) - 8}" if len(lost) > 8 else ""))
    table, bad = time_table(wb, SPEC["time"])
    if table is None and side.get("timed"):                           # 処理時間を出したブックなのに、基準の表を読めない。人が直した表の値を落とさない
        raise SystemExit(f"{book.name}: 処理時間の基準の表を読めない（{bad or '名前「' + SPEC['time']['name'] + '」の範囲が無い'}）。人が直した表の値を取り込めないので止める。"
                         "表を戻すか、取り込まずに作り直すなら --fresh")
    wrote = {(a, v): x for a, v, x in side.get("table") or []}        # 書き出したときの、表の値（テンプレートの表が変わっても、ブックに残る古い値を、人の直しと見ない）
    for key, x in ((table or {}).get("values") or {}).items():        # 基準の表: 書き出したときの値と違う所が、人が直した値
        if key not in base or (key in wrote and x == wrote[key]):     # テンプレートの表に無い列（前の形のテンプレート）: 比べる相手が無い／書き出したときのまま: 人は、触っていない
            continue
        if x != base[key] and (key in wrote or x not in TIME_WAS.get(key[0], ())):
            if times.table.get(key, {}).get("value") != whole(x):
                times.table[key] = {"action": key[0], "volume": key[1], "value": whole(x), "from": "book"}
                n["基準の表"] += 1
        elif replace and key in times.table:
            del times.table[key]
            n["消した値"] += 1
    if any(n.values()):
        said.append(f"{book.name}: 人が入れた値を取り込んだ（処理時間 {n['時間']} セル・件数 {n['件数']} 行・基準の表 {n['基準の表']} か所"
                    + (f"。式・空に戻っていた値 {n['消した値']} 個は、消した" if n["消した値"] else "") + "）")
    if unread:                                                       # 取り込まずに書き出すと、人が入れた値が消える
        raise SystemExit(f"{book.name}: 数として読めない値が、時間・件数のセルにある: {'、'.join(unread[:8])}" + (f" ほか {len(unread) - 8}" if len(unread) > 8 else "")
                         + "。分の数だけ（「60」「1.5」）にしてから、もう一度（上書きしていない。取り込まずに作り直すなら --fresh。そのセルの値は消える）")
    return said


def calibration(mx, base):
    """人が入れた時間から見た、基準の表: [((動作の種類, 量), 人の値の数, その中央値, いまの表の値)]。人の作業の箱で、3 つ以上あるものだけ"""
    groups = {}
    for r in mx.rows:
        for s in r["steps"]:
            x = mx.times.cells.get((r["id"], s["id"]))
            if x and s["node"]["type"] == "process":
                groups.setdefault(s["time"][:2], []).append(x["value"])
    return [(k, len(v), whole(statistics.median(v)), mx.times.table.get(k, {}).get("value", base.get(k))) for k, v in sorted(groups.items()) if len(v) >= 3]


class Matrix:
    """業務マトリクス一覧の中身（sheets: シート名 → 行 → Cell）を、作業フォルダのデータから組み立てる"""

    def __init__(self, work, write=True, tax=None):
        self.r = r = Reports(work, write)                            # 報告書と同じ材料（資料・ページの書き方・フロー・確認事項・階層・分類）
        self.timed, self.times = False, Times(r.work)                # 処理時間の行を出すか（テンプレートに基準の表があれば。write が決める）／人が入れた値
        if tax is not None:                                          # まだ書いていない分類の下書きで、縦軸だけを見る（pages.py rows）
            r.tax = tax
        self.work, self.out, self.lab = r.work, r.out, r.lab
        self.cats = [c for c in r.tax.get("categories") or [] if isinstance(c, dict)]
        self.by_cat = {c.get("id"): c for c in self.cats}
        self.flows = self.ordered(r.flows)
        self.clues = node_clues(self.flows)                          # 時間のセルのコメントに書く、付け替えの手がかり（担当・直前の箱）
        self.rows = [self.flow_row(fl) for fl in self.flows]
        self.ledger = system_ledger(self.work, self.flows)
        from systems import Ledger, has_ledger
        self.led = Ledger(self.work, r.m) if has_ledger(self.work) else None

    # ---- 行の並びと、左の列 ----

    def ordered(self, flows):
        """分類の木を上からたどった順（下位フローは親の直後。兄弟は taxonomy に書いた順）。分類に無いフローは末尾"""
        kids = {}
        for c in self.cats:
            kids.setdefault(c.get("parent") if c.get("parent") in self.by_cat else None, []).append(c.get("id"))
        order, todo = [], kids.get(None, [])[::-1]
        while todo:
            i = todo.pop()
            if i not in order:
                order.append(i)
                todo += kids.get(i, [])[::-1]
        by = {flow_key(fl): fl for fl in flows}
        return [by[i] for i in order if i in by] + [fl for fl in flows if flow_key(fl) not in order]

    def names(self, fl):
        """(大分類, 中分類, 行の名前の列)。分類の親をたどった並び c0…ck を、上から詰める: 大分類＝c0、中分類＝c1…ck を「›」でつないだもの。
        1段しか無ければ（分類に無いフローも）、大分類に名前、中分類は「—」。行の名前の列は、0＝大分類／1＝中分類"""
        path = category_path(self.cats, flow_key(fl))
        seg = [f"{own_number(c, path[i - 1] if i else None)} {c.get('name', '')}".strip() for i, c in enumerate(path)] or [str(fl.get("name", ""))]
        return (seg[0], " › ".join(seg[1:]), 1) if len(seg) > 1 else (seg[0], "—", 0)

    def name_basis(self, fl):
        """分類名の根拠（目次／本文の見出し／資料名／LLM が命名）と、その場所"""
        c = self.by_cat.get(flow_key(fl), {})
        ref = c.get("toc_ref") if c.get("basis", fl.get("basis")) == "toc" else c.get("name_ref")
        kind = "目次" if c.get("basis", fl.get("basis")) == "toc" else NAME_REF.get((ref or {}).get("kind"), "LLM が命名")
        where = ""
        if isinstance(ref, dict) and ref.get("doc_id"):
            at = self.lab.join(ref["doc_id"], self.lab.short(ref["doc_id"], ref["page"])) if ref.get("page") is not None else self.lab.title(ref["doc_id"])
            where = at + (f"「{ref['text']}」" if ref.get("text") else "")
        return lines([[(kind, "kind")], [(where, "src")]])

    def page_ranges(self, doc, pages):
        """1資料のページの並びを、刷られた番号で短く書く（p.41〜43、p.45）。番号で書けないページ（シート・表紙）は1つずつ"""
        pages, out, i = sorted({p for p in pages if isinstance(p, int)}), [], 0
        num = lambda p: re.fullmatch(r"(p\.|スライド )(\d+)", self.lab.short(doc, p))
        while i < len(pages):
            j = i
            while num(pages[i]) and j + 1 < len(pages) and pages[j + 1] == pages[j] + 1 and num(pages[j + 1]) \
                    and int(num(pages[j + 1]).group(2)) == int(num(pages[j]).group(2)) + 1:
                j += 1
            out.append(self.lab.short(doc, pages[i]) + (f"〜{num(pages[j]).group(2)}" if j > i else ""))
            i = j + 1
        return "、".join(dict.fromkeys(out))

    def cited(self, nodes):
        """ノードの並びが出典にしているページを、資料ごとに: [(資料, [page], 引いているノードの数)]（初めて出てきた順）"""
        out = {}
        for n in nodes:
            for d in dict.fromkeys(s.get("doc_id") for s in n.get("sources") or []):
                e = out.setdefault(d, [[], 0])
                e[0] += [s.get("page") for s in n["sources"] if s.get("doc_id") == d]
                e[1] += 1
        return [(d, ps, k) for d, (ps, k) in out.items()]

    def origin(self, fl):
        """原典（資料とページ）: ノードの出典を資料ごとに集めたもの（資料は題名で書く。役割は「資料一覧」）。引いているノードの多い資料から。資料どうしの関係を添える"""
        rank = {d["doc_id"]: i for i, d in enumerate(self.r.m["docs"])}
        docs = sorted(self.cited(n for n in fl["nodes"] if n["type"] not in ("start", "end")), key=lambda x: (-x[2], rank.get(x[0], 10 ** 6)))
        parts = []
        for d, pages, k in docs:
            parts.append([(self.lab.join(d, self.page_ranges(d, pages)), "text"), (f"〔{k} ノード〕" if len(docs) > 1 else "", "src")])
        used = {d for d, _p, _k in docs}
        for rel in self.r.hier.get("doc_relations") or []:
            a, b, t = rel.get("from"), rel.get("to"), rel.get("type")
            if a in used and b in used:
                A, B = (f"「{self.lab.title(x)}」" for x in (a, b))
                text = {"parent_of": f"{B}は{A}の下位（詳細）", "appendix_of": f"{B}は{A}の別紙", "basis_of": f"{A}は{B}の根拠"}.get(
                    t, f"{A}と{B}: {REL_LABEL.get(t, t)}")
                parts.append([(text, "src")])
        return lines(parts)

    # ---- 作業のセル ----

    def flow_row(self, fl):
        """フロー1本の行の材料: 列の順に並べた作業（steps）、開始・終了・帳票、確認事項の置き場所"""
        F, qs = flow_key(fl), self.r.qs
        q_node, q_edge = question_ids(qs, F)
        node = {n["id"]: n for n in fl["nodes"]}
        det = {d["id"]: d for d in node_details(fl, self.r.m, qs, self.r.items_of)}
        order = [i for i in step_order_by_branch(fl) if node[i]["type"] in STEP_TYPES]
        no, depth = {i: k for k, i in enumerate(order, 1)}, order_of(fl)
        lane = {l["id"]: l["name"] for l in fl.get("lanes") or []}
        files = {}
        for n in fl["nodes"]:
            if n["type"] == "file":
                files.setdefault(n.get("attach_to") if n.get("attach_to") in no else None, []).append(n)

        def name(v, back=False):
            """行き先・入口のノードの呼び方: 作業は列の番号、終了はそのラベル"""
            n = node.get(v)
            if n is None:
                return str(v)
            if v in no:
                return circled(no[v]) + ("へ戻る" if back else "")
            return {"start": "開始", "end": n.get("label") or "終了"}.get(n["type"], f"帳票「{n.get('label', '')}」")

        edges = []
        for e in fl.get("edges") or []:
            u, v = e.get("from"), e.get("to")
            marks = (INFERRED if e.get("inferred") else "") + EDGE_KIND.get(e.get("kind"), "")
            head = (f"「{e['label']}」" if e.get("label") else "") + "→" + (marks or " ")
            qe = list(q_edge.get((u, v), []))
            if node.get(v, {}).get("type") == "end":                  # 終了に掛かる確認事項は、そこへ行く線の所に出す
                qe += [q for q in q_node.get(v, []) if q not in qe]
            back = u in depth and v in depth and depth[v] <= depth[u]   # 戻りの線（差し戻しなど。前の列へ戻る）
            edges.append({"from": u, "to": v, "label": e.get("label") or "", "plain": not e.get("label") and not marks and not qe, "qs": qe,
                          "inferred": bool(e.get("inferred")), "from_start": node.get(u, {}).get("type") == "start",
                          "exit": head + name(v, back),                                                             # マトリクスのセルの書き方
                          "exit_id": head + f"{v} {name(v)}",                                                       # ノード一覧の書き方（ノード番号つき）
                          "entry": name(u) + (f"「{e['label']}」" if e.get("label") else "") + ("（戻り）" if back else "")})
        steps = []
        for k, i in enumerate(order, 1):
            n, d = node[i], det[i]
            ins, outs = [e for e in edges if e["to"] == i], [e for e in edges if e["from"] == i]
            left, right = (order[k - 2] if k > 1 else None), (order[k] if k < len(order) else None)
            # 入口を書かないのは、左隣から来る線1本だけのとき（1列目は、開始から来る線1本だけのとき）
            usual = len(ins) == 1 and (ins[0]["from"] == left if left else ins[0]["from_start"] and ins[0]["plain"])
            # 行き先を書かないのは、判断でなく、次の列へ進むだけの線1本のとき
            quiet = n["type"] != "decision" and len(outs) == 1 and outs[0]["to"] == right and outs[0]["plain"]
            fq = [q for f in files.get(i, []) for q in q_node.get(f["id"], [])]        # 帳票に掛かる確認事項は、付け先の作業のセルに出す
            steps.append({"no": k, "id": i, "node": n, "detail": d, "lane": lane.get(n.get("lane"), ""), "ins": ins, "outs": outs,
                          "show_in": bool(ins) and not usual, "show_out": bool(outs) and not quiet, "files": files.get(i, []),
                          "qs": list(dict.fromkeys(d["questions"] + fq)), "pages": self.cited([n]), "time": time_basis(n, sum(f.get("direction") == "in" for f in files.get(i, [])))})
        starts = [n for n in fl["nodes"] if n["type"] == "start"]
        row = {"flow": fl, "id": F, "steps": steps, "edges": edges, "node": node, "det": det, "no": no, "lane": lane, "files": files,
               "start_qs": list(dict.fromkeys(q for s in starts for q in q_node.get(s["id"], []))), "q_node": q_node}
        shown = set(row["start_qs"])
        for s in steps:
            shown |= set(s["qs"]) | {q for e in s["outs"] for q in e["qs"]} | {q for e in s["ins"] if e["from_start"] for q in e["qs"]}
        # ノードにも線にも掛からない確認事項と、置き場所が無かったもの（指す先のノードが無い など）は、「フロー全体の確認事項」に出す
        row["whole_qs"] = [q["id"] for q in qs if q.get("flow") == F and q["id"] not in shown]
        return row

    def step_cell(self, s):
        """作業1つのセル: 担当／入口／作業名／行き先／出典と確認事項。ノード番号・補足・帳票・原文はコメントに"""
        n, d = s["node"], s["detail"]
        parts = [[(s["lane"], "actor")]]
        if s["show_in"]:
            entry = []
            for e in s["ins"]:                                       # 開始から来る線の推定・確認事項は、ここにしか出せない（開始は列にならない）
                entry.append(e["entry"] + ((INFERRED if e["inferred"] else "") + (" " + ", ".join(e["qs"]) if e["qs"] else "") if e["from_start"] else ""))
            parts.append([("← " + "、".join(entry), "entry")])
        parts.append([(MARK.get(n["type"], "") + d["label"], "label")] + ([(INFERRED, "inferred")] if d["inferred"] else []))
        if d["systems"]:                                             # 札: その作業で使うシステム
            parts.append([("".join(f"［{x}］" for x in d["systems"]), "sys")])
        if n.get("activity"):                                        # プログラムの起動点の箱: 中身は、この番号のアクティビティ図
            parts.append([(f"▶ {n['activity']}", "sys")])
        if s["show_out"]:
            for e in s["outs"]:
                parts.append([(e["exit"], "exit")] + ([(" " + ", ".join(e["qs"]), "q")] if e["qs"] else []))
        src = "／".join(self.lab.join(doc, self.page_ranges(doc, pages)) for doc, pages, _k in s["pages"])
        parts.append([(src, "src")] + ([(("　" if src else "") + ", ".join(s["qs"]), "q")] if s["qs"] else []))
        fill = FILL["inferred"] if d["inferred"] else FILL.get(n["type"])
        return Cell(lines(parts), fill=fill, note=self.note(s))

    def files_text(self, files):
        return "／".join(f"{'見る' if f.get('direction') == 'in' else '作る'}「{f.get('label', '')}」（{f['id']}）" for f in files)

    def note(self, s):
        d = s["detail"]
        return "\n".join(x for x in [f"{d['id']}（{KIND.get(d['type'], d['type'])}）", d["detail"] and f"補足: {d['detail']}",
                                     d["reference"] and f"参考: {d['reference']}", d["systems"] and f"使うシステム: {'、'.join(d['systems'])}",
                                     s["node"].get("activity") and f"プログラムの中身: アクティビティ図 {s['node']['activity']}",
                                     s["files"] and f"帳票: {self.files_text(s['files'])}",
                                     d["sources"] and f"出典: {'、'.join(d['sources'])}", d["quote"] and f"原文: {d['quote']}"] if x)

    # ---- シート ----

    def matrix(self):
        width = max((len(r["steps"]) for r in self.rows), default=0)
        out, prev = [[Cell([(h, "head")], fill=FILL["head"]) for h in MATRIX_HEAD + [circled(k) for k in range(1, width + 1)]]], None
        self.row_keys = []                                            # 行ごとの鍵（人が埋める欄を、書き出し直したあとも同じ行に戻す）
        pngs = {re.sub(r".*_flow-", "", p.stem): p.name for p in sorted((self.out / "png").glob("*_flow-*.png"))} if (self.out / "png").exists() else {}
        for r in self.rows:
            fl = r["flow"]
            top, mid, k = self.names(fl)
            link = f"png/{pngs[r['id']]}" if r["id"] in pngs else None
            start, end = str(fl.get("start_condition") or ""), str(fl.get("end_condition") or "")
            # 行の名前のセル（1段なら大分類、2段以上なら中分類）: 太字と、フロー図へのリンク。リンクのセルは、字の種類を1つにする（途中で変えると、表計算ソフトによっては字の指定が落ちる）
            named = Cell([((top, mid)[k], "link" if link else "bold")], link=link)
            out.append([named if k == 0 else Cell([(top, "pale" if top == prev else "bold")]), named if k == 1 else Cell([(mid, "pale")]),
                        Cell(self.name_basis(fl)), Cell(self.origin(fl)),
                        Cell([(start or "（記載なし）", "text")] + ([(" " + ", ".join(r["start_qs"]), "q")] if r["start_qs"] else [])),
                        Cell(end or "（記載なし）"), Cell([(", ".join(r["whole_qs"]), "q")]), Cell(), Cell()]
                       + [self.step_cell(s) for s in r["steps"]] + [Cell() for _ in range(width - len(r["steps"]))])
            self.row_keys.append(r["id"])
            if self.timed:
                out.append(self.time_row(r, len(out) - 1, width, (top, mid), k))
                self.row_keys.append(f"{r['id']}｜処理時間")
            prev = top
        return out

    def time_row(self, r, i, width, axis, k):
        """処理時間（分）の行（作業の行のすぐ下。データの i 行目）: 作業のセルの真下に、分。人が入れた値があれば、その数。無ければ、基準の表を引く式（仮の値）。
        どのセルも、コメントの 1 行目が札（だれの時間か）。書き出し直すときに、これで読み戻す。axis は (大分類, 中分類)、k は行の名前の列"""
        F, n, T, name = r["id"], len(r["steps"]), FILL["time"], axis[k]
        total = Cell(formula=lambda at: f"=SUM({at.cell(M, i, 1)}:{at.cell(M, i, n)})", font="bold", fill=T,
                     shown=lambda at: whole(sum(self.minutes(at, F, s) for s in r["steps"]))) if n else Cell(fill=T)
        cnt = self.times.counts.get(F)
        count = Cell([(cnt["value"], "text")] if cnt else "", fill=T, font="time", tag=("count", F),
                     note=f"件数 {F}｜人が入れる欄（数）\n{name}\n書き出し直しても残る。仮の値は、入れない")
        cells = [Cell([(x + ("｜処理時間（分）" if j == k else ""), "pale")], fill=T) for j, x in enumerate(axis)] + [Cell(fill=T) for _ in range(5)] + [count, total]
        for s in r["steps"]:
            (act, vol, why), got = s["time"], self.times.cells.get((F, s["id"]))
            _name, lane, after = self.clues.get((F, s["id"]), ("", "", []))
            note = "\n".join([f"処理時間 {F} {s['id']}｜{'人が入れた値' if got else '仮の値'}（{act} × {vol}{'：' + why if why else ''}）", f"作業: {s['detail']['label']}",
                              "担当: {}｜前: {}".format(lane, "｜".join(after)),
                              "数を打つと、その値が正になる（書き出し直しても残る）。消すと、仮の値に戻る"])
            if got:
                cells.append(Cell([(got["value"], "text")], fill=T, font="time", note=note, tag=("time", F, s["id"], s["detail"]["label"])))
            else:
                cells.append(Cell(formula=lambda at, a=act, v=vol: "=" + at.lookup(a, v), shown=lambda at, a=act, v=vol: at.minutes(a, v), fill=T, font="time",
                                  cond=("time_long", "time_soft") if vol in VOLUME_SPARE else "time_soft", note=note, tag=("time", F, s["id"], s["detail"]["label"])))
        return cells + [Cell(fill=T) for _ in range(width - n)]

    def minutes(self, at, F, s):
        """その作業の時間（分）: 人が入れた値。無ければ、基準の表の値（仮の値の式の、計算の結果）"""
        got = self.times.cells.get((F, s["id"]))
        return got["value"] if got else at.minutes(*s["time"][:2])

    def node_list(self):
        out = [[Cell([(h, "head")], fill=FILL["head"]) for h in NODE_HEAD]]
        for at_row, r in enumerate(self.rows):
            fl, q_node = r["flow"], r["q_node"]
            top, mid, _k = self.names(fl)
            head = [top, mid, r["id"]]

            def timing(s, i=2 * at_row + 1, F=r["id"]):
                """動作の種類・量・処理時間（業務マトリクスの、時間のセルを指す式。数を打つ場所は、1 か所）・時間の根拠"""
                act, vol, why = s["time"]
                k, basis = s["no"], vol + (f"（{why}）" if why else "")
                if not self.timed:
                    return [act, basis, "", ""]
                return [act, basis, Cell(formula=lambda at: "=" + at.cell(M, i, k), shown=lambda at: self.minutes(at, F, s)),
                        Cell(formula=lambda at: f'=IF(_xlfn.ISFORMULA({at.cell(M, i, k)}),"仮","人")', shown="人" if (F, s["id"]) in self.times.cells else "仮")]
            flow_of = lambda i, key: "\n".join(e[key] + (" " + ", ".join(e["qs"]) if e["qs"] and key == "exit_id" else "") for e in r["edges"]
                                               if e["from" if key == "exit_id" else "to"] == i)
            ins = lambda i: "\n".join(f"← {e['from']} {e['entry']}" for e in r["edges"] if e["to"] == i)

            def plain(n, order=""):
                """開始・終了・帳票の行"""
                d = r["det"].get(n["id"], {})
                who = f"{n.get('attach_to', '')} が{'見る' if n.get('direction') == 'in' else '作る'}" if n["type"] == "file" else ""
                return head + [order, n["id"], KIND.get(n["type"], n["type"]), r["lane"].get(n.get("lane"), ""), n.get("label", ""),
                               d.get("detail", ""), d.get("reference", ""), "", ins(n["id"]), flow_of(n["id"], "exit_id"), who,
                               "、".join(d.get("sources", [])), d.get("quote", ""), "推定" if n.get("inferred") else "", "補足" if d.get("note") else "",
                               ", ".join(q_node.get(n["id"], []))] + ["", "", "", ""]
            rows = [plain(n) for n in fl["nodes"] if n["type"] == "start"]
            for s in r["steps"]:
                d = s["detail"]
                rows.append(head + [s["no"], d["id"], KIND.get(d["type"], d["type"]), s["lane"], d["label"], d["detail"], d["reference"],
                                    "、".join(d["systems"]) or (f"▶ {s['node']['activity']}" if s["node"].get("activity") else ""), ins(d["id"]),
                                    flow_of(d["id"], "exit_id"), self.files_text(s["files"]), "、".join(d["sources"]), d["quote"],
                                    "推定" if d["inferred"] else "", "補足" if d["note"] else "", ", ".join(d["questions"])] + timing(s))
                rows += [plain(f, s["no"]) for f in s["files"]]
            rows += [plain(n) for n in fl["nodes"] if n["type"] == "end"] + [plain(f) for f in r["files"].get(None, [])]
            fills, soft, kind = {KIND["decision"]: FILL["decision"], KIND["file"]: FILL["file"]}, NODE_HEAD.index("推定"), NODE_HEAD.index("種別")
            for x in rows:
                fill = FILL["inferred"] if x[soft] else fills.get(x[kind])
                for v in x:
                    if isinstance(v, Cell):
                        v.fill = fill
                out.append([v if isinstance(v, Cell) else Cell(v if isinstance(v, str) else [(v, "text")], fill=fill) for v in x])
        return out

    def systems(self):
        """システム一覧: 1行＝1システム。使う作業（札）と、システムが行う処理（箱）を、フローごとに並べる"""
        def places(items):
            by = {}
            for it in items:
                by.setdefault(it[0], []).append(it[1] + (f"（▶ {it[2]}）" if len(it) > 2 and it[2] else ""))
            return "\n".join(f"{F} {'・'.join(ids)}" for F, ids in by.items())
        out = [[Cell([(h, "head")], fill=FILL["head"]) for h in SYSTEM_HEAD]]
        count = {}
        for e in self.led.entries() if self.led else []:
            if e["system"] is not None:
                count[e["system"]["name"]] = count.get(e["system"]["name"], 0) + 1
        for x in self.ledger:
            note = "" if x["known"] else "用語の正本（glossary.json）に無い名前"
            if x["known"] and not x["tags"] and not x["boxes"]:
                note = "フローに出てこない"
            # 番号は数だけ（アクティビティ図の番号 A-02-1 の 02。SYS- は内部の番号なので出さない）
            out.append([Cell(str(x["id"] or "—").replace("SYS-", "")), Cell([(x["name"], "bold")]), Cell("、".join(x["aliases"])), Cell(SYSTEM_KIND.get(x["kind"], x["kind"] or "—")),
                        Cell("あり（" + "・".join(self.lab.title(d) for d in x["docs"]) + "）" if x["docs"] else "なし"),
                        Cell([(count[x["name"]], "text")] if x["name"] in count else "—"), Cell(places(x["tags"])), Cell(places(x["boxes"])),
                        Cell(note, fill=None if x["known"] else FILL["inferred"])])
        return out

    def code_nodes(self):
        """プログラムのノード一覧: アクティビティ図の箱（開始・終了・帳票を除く）を1行ずつ。続けて、その図で図にしなかった文（対象外）を、理由つきで"""
        from build_activity import load_activities
        out = [[Cell([(h, "head")], fill=FILL["head"]) for h in CODE_HEAD]]
        for a in load_activities(self.work):
            q_node, _q_edge = question_ids(self.r.qs, a["flow_id"])
            lane = {l["id"]: l["name"] for l in a["lanes"]}
            for n in a["nodes"]:
                if n["type"] in ("start", "end", "file"):
                    continue
                outs = "\n".join((f"「{e['label']}」" if e.get("label") else "") + "→ " + (e["to"] if not e["to"].startswith("end") else "終了") for e in a["edges"] if e["from"] == n["id"])
                files = "／".join(f"{'読む' if f.get('direction') == 'in' else '書く'}: {f['label']}" for f in a["nodes"] if f.get("attach_to") == n["id"])
                note = "／".join(x for x in (n.get("detail", ""), f"繰り返し（{n['loop']}）" if n.get("loop") else "", f"中の分岐を畳んだ（{n['folded']}）" if n.get("folded") else "",
                                           f"別の図 {n['link_flow']}" if n.get("link_flow") else "",
                                           "・".join(f"別の図 {c['to']} を {c['times']} 回" for c in n.get("calls") or [])) if x)
                srcs = n.get("sources") or []
                out.append([Cell(a["flow_id"]), Cell(a["name"]), Cell(n["id"]), Cell(KIND.get(n["type"], n["type"]) if n["type"] != "subprocess" else "呼び出し"),
                            Cell(lane.get(n["lane"], "")), Cell([(n["label"], "bold")], fill=FILL.get(n["type"]) if n["type"] == "decision" else None), Cell(note), Cell(outs),
                            Cell(files), Cell("、".join(self.lab.src(s) for s in srcs)), Cell(srcs[0].get("quote", "") if srcs else ""),
                            Cell([(", ".join(q_node.get(n["id"], [])), "q")])])
            st = (self.led.structs.get(a.get("project")) or {}) if self.led else {}
            for o in a.get("lines") or []:
                if o.get("state") == "excluded" and o.get("lines"):
                    from code_struct import cite
                    at = "、".join(self.lab.src(c) for c in cite(st, o["module"], min(o["lines"]), max(o["lines"]))) if st else f"{o['module']} {min(o['lines'])}〜{max(o['lines'])} 行"
                    out.append([Cell(a["flow_id"]), Cell(a["name"]), Cell("—"), Cell("対象外"), Cell(""), Cell([(o.get("reason", ""), "pale")]), Cell(f"実行文 {len(o['lines'])} 行"),
                                Cell(""), Cell(""), Cell(at), Cell(""), Cell("")])
        return out

    def terms(self):
        """用語の対応: 1行＝プログラムの中の名前1つと、図に書いた業務の言葉（structure/code_terms.json）"""
        out = [[Cell([(h, "head")], fill=FILL["head"]) for h in TERM_HEAD]]
        if self.led is None:
            return out
        import terms as T
        from build_activity import load_activities
        from code_struct import cite
        tm, used = T.Terms(self.work, self.led.structs), {}
        for a in load_activities(self.work):                         # その名前が、箱の受け持つ行に出てくる図
            proj = a.get("project")
            text = "\n".join(tm.text(proj, o.get("module"), o.get("lines") or []) for o in a.get("lines") or [] if o.get("state") != "excluded")
            for name in T.code_names(text, tm.lang(proj)):
                used.setdefault((proj, name), []).append(a["flow_id"])
        for t in sorted(tm.rows, key=lambda t: (str(t.get("project")), str(t.get("name")))):
            proj, s, src = t.get("project"), self.led.sys_of_proj.get(t.get("project")) or {}, t.get("source") or {}
            st = self.led.structs.get(proj) or {}
            at = "、".join(self.lab.src(c) for c in cite(st, src["module"], src["line"], src["line"])) if src.get("line") and st else ""
            guess = t.get("basis") in ("spelling", "unknown")
            out.append([Cell(s.get("name", "")), Cell([(str(t.get("name", "")), "bold")]), Cell(str(t.get("word", "")) or "（読めない。もとの名前のまま書いた）", fill=FILL["inferred"] if guess else None),
                        Cell(T.BASIS.get(t.get("basis"), str(t.get("basis", ""))), fill=FILL["inferred"] if guess else None), Cell(at or (f"{src.get('module')} {src.get('line')} 行" if src.get("line") else "")),
                        Cell("、".join(used.get((proj, t.get("name")), []))), Cell(str(t.get("note", "")))])
        return out

    def entries(self):
        """起動点一覧: 1行＝プログラムの起動点1つ。続けて、ほかの処理から呼ばれる共通の処理（別ページのアクティビティ図）"""
        from code_struct import TRIGGER_LABEL
        from systems import BASIS, placed
        out = [[Cell([(h, "head")], fill=FILL["head"]) for h in ENTRY_HEAD]]
        if self.led is None:
            return out
        boxes = {}
        for fl in self.flows:
            for n in placed(fl):
                boxes.setdefault(n["activity"], []).append(f"{flow_key(fl)} {n['id']}")
        qs = {}
        for q in self.r.qs:
            mt = re.match(r"^(?:unlinked|unplaced|manual):(.+)$", str(q.get("local", ""))) if q.get("origin") == "systems" else \
                re.match(r"^m-sys-(.+)$", str(q.get("local", "")))
            if mt:
                qs.setdefault(re.sub(r"-(e[0-9]+)$", r".\1", mt.group(1)), []).append(q["id"])
        def at(step):
            """手順の指し方（資料.ページ.項目「M03.1.s3」。台帳の書き方）→ 人に見せる場所（資料の題名とページ）"""
            m = re.fullmatch(r"(.+)\.(\d+)\.(\w+)", str(step or ""))
            return self.lab.join(m.group(1), self.lab.short(m.group(1), int(m.group(2)))) if m else str(step or "")

        for i, e in enumerate(self.led.entries(), 1):                 # 起動点は、表の中の連番で書く（台帳の番号「M01.e1」は、資料ID を含むので出さない）
            links = "\n".join(f"{BASIS.get(x.get('basis'), x.get('basis'))}" + (f"「{x['evidence']}」" if x.get("evidence") else "") + f"（{at(x.get('step'))}）" for x in e["links"])
            note = f"図にしない: {e['excluded']}" if e["excluded"] is not None else "" if e["links"] else unlinked_note(self.led.m)
            out.append([Cell(str(i)), Cell(e["system"]["name"] if e["system"] else "（システムが未設定）"), Cell(f"{e['label']}〔{TRIGGER_LABEL.get(e['trigger'], e['trigger'])}〕"),
                        Cell(e["procedure"]), Cell([(e["activity"] if e["excluded"] is None else "—", "bold")]), Cell("\n".join(boxes.get(e["activity"], []))),
                        Cell(links), Cell([(", ".join(qs.get(e["key"], [])), "q")]), Cell(note, fill=FILL["file"] if e["excluded"] is not None else None)])
        for u in self.led.units:
            if u["kind"] == "shared":
                off = self.led.off(u)                                 # 図にしない手続き（補助）は、番号を出さない
                out.append([Cell("（共通）"), Cell(u["system"]["name"]), Cell("ほかの処理から呼ばれる"), Cell(u["procedure"]), Cell([(u["activity"] if off is None else "—", "bold")]),
                            Cell(""), Cell(""), Cell(""), Cell("" if off is None else f"図にしない: {off}（呼ぶ側の図の箱に入れた）", fill=None if off is None else FILL["file"])])
        return out

    def helpers(self):
        """プログラムの図にしない手続き: 台帳で図にしないとした補助の手続き（理由つき）。機械が候補に出したが、図のままにしたものは、出さない（点検の INFO に、数が出る）"""
        from build_activity import load_activities
        from code_struct import cite
        out = [[Cell([(h, "head")], fill=FILL["head"]) for h in HELPER_HEAD]]
        if self.led is None:
            return out
        scope = {}                                                   # (まとまり, 手続き) → その手続きが入っている図
        for a in load_activities(self.work):
            for pid in a.get("scope") or []:
                scope.setdefault((a.get("project"), pid), []).append(a["flow_id"])
        for h in [x for x in self.led.helpers() if x["reason"] is not None]:
            st = self.led.structs[h["project"]]
            p = next(x for x in st["procedures"] if x["id"] == h["procedure"])
            callers = list(dict.fromkeys(a for pid, _ln in h["sites"] for a in scope.get((h["project"], pid), [])))
            out.append([Cell(h["system"]["name"]), Cell([(h["procedure"], "bold")]), Cell([("図にしない", "kind")], fill=FILL["file"]),
                        Cell(h["reason"]), Cell("、".join(callers)),
                        Cell("\n".join(f"{pid} {ln} 行" for pid, ln in h["sites"][:12]) + (f"\nほか {len(h['sites']) - 12} か所" if len(h["sites"]) > 12 else "")),
                        Cell("、".join(self.lab.src(s) for s in cite(st, p["module"], *p["lines"]))), Cell([(h["exec"], "text")])])
        return out

    def axis(self):
        """縦軸の根拠: 業務の木の項目すべてと、木に無い分類の、扱いと理由"""
        tax, lab = self.r.tax, self.lab
        drawn, no_toc = {r["id"] for r in self.rows}, tax.get("basis") == "llm"     # 目次が無い（業務の木は、見出し・資料名から）
        of_h = {}
        for c in self.cats:
            if c.get("hierarchy_id"):
                of_h.setdefault(c["hierarchy_id"], c)
        below = lambda c: any(x.get("make_flow") for x in self.cats if c in category_path(self.cats, x.get("id"))[:-1])
        why = {}
        for u in tax.get("unassigned_pages") or []:
            if isinstance(u, dict):
                why[(u.get("doc_id"), u.get("page"))] = str(u.get("reason") or "")

        def treat(c):
            if c.get("make_flow"):
                return SPLIT if self.by_cat.get(c.get("parent"), {}).get("make_flow") else DRAW
            return GROUP if below(c) else OUT

        def row_name(c):
            top, mid, _k = self.names({"flow_id": c.get("id"), "name": c.get("name")})
            return " › ".join([top] + ([mid] if mid != "—" else [])) + ("" if c.get("id") in drawn or not c.get("make_flow") or not drawn else "（フローが無い）")

        out, state = [[Cell([(h, "head")], fill=FILL["head"]) for h in AXIS_HEAD]], {}
        for h, parent in tree_nodes(self.r.hier.get("tree")):
            c = of_h.get(h.get("id"))
            up = state.get(id(parent)) if parent is not None else None      # 上の項目の (扱い, 分類)
            pages = [(r_.get("doc_id"), p) for r_ in h.get("refs") or [] for p in r_.get("pages") or []]
            own = "、".join(dict.fromkeys(why[k] for k in pages if why.get(k)))
            if c is not None:
                kind = treat(c)
                reason = {SPLIT: f"上位の「{self.by_cat.get(c.get('parent'), {}).get('name', '')}」から分けた",
                          OUT: str(c.get("out_of_scope") or own)}.get(kind, "")
                if not reason and no_toc and c.get("basis") == "llm":       # 目次の無い資料: 名前の根拠は、ここに書く（足した分類の行を、重ねて出さない）
                    ref = c.get("name_ref") or {}
                    reason = "名前: " + NAME_REF.get(ref.get("kind"), "LLM が命名") + (f"「{ref['text']}」" if ref.get("text") else "")
                state[id(h)] = (kind, c)
            elif up and up[0] in (DRAW, SPLIT, INSIDE):
                kind, reason = INSIDE, f"「{up[1].get('name', '')}」のフローに含む"
                state[id(h)] = (INSIDE, up[1])
            elif up and up[0] == OUT:
                kind, reason = OUT, own or "上の項目と同じ"
                state[id(h)] = up
            else:
                kind, reason = UNSET, ""
                state[id(h)] = (UNSET, None)
            title = "　" * (int(h.get("level") or 1) - 1) + f"{h.get('number', '')} {h.get('title', '')}".strip()
            where = "／".join(lab.join(r_["doc_id"], self.page_ranges(r_["doc_id"], r_.get("pages") or [])) for r_ in h.get("refs") or [] if r_.get("doc_id"))
            target = state[id(h)][1]
            out.append([Cell([(h.get("level") or "", "text")]), Cell(title), Cell(TREE_SOURCE.get(h.get("source"), "")),
                        Cell([(kind, "kind")], fill=FILL["file"] if kind == OUT else FILL["inferred"] if kind == UNSET else None),
                        Cell(row_name(target) if target and kind in (DRAW, SPLIT, INSIDE) else ""), Cell(reason), Cell(where)])
        ids = {h.get("id") for h, _p in tree_nodes(self.r.hier.get("tree"))}
        for c in self.cats:
            if c.get("hierarchy_id") in ids and (c.get("basis") != "llm" or no_toc):    # 業務の木の項目として、上に出ている
                continue
            ref = c.get("name_ref") or {}
            reason = f"{NAME_REF.get(ref.get('kind'), 'LLM が命名')}" + (f"「{ref['text']}」" if ref.get("text") else "")
            where = "／".join(lab.join(ps["doc_id"], self.page_ranges(ps["doc_id"], ps.get("pages") or [])) for ps in c.get("pages") or [] if ps.get("doc_id"))
            out.append([Cell([(c.get("level") or "", "text")]), Cell("　" * (int(c.get("level") or 1) - 1) + str(c.get("name", ""))), Cell("（目次に無い）"),
                        Cell([(ADDED, "kind")], fill=FILL["decision"]), Cell(row_name(c) if c.get("make_flow") else ""), Cell(reason), Cell(where)])
        return out

    def docs(self):
        rel = {}
        for x in self.r.hier.get("doc_relations") or []:
            a, b, t = x.get("from"), x.get("to"), x.get("type")
            A, B = (f"「{self.lab.title(d)}」" for d in (a, b))
            said = {"parent_of": (f"{B}の上位", f"{A}の下位（詳細）"), "appendix_of": (f"別紙に{B}", f"{A}の別紙"), "basis_of": (f"{B}の根拠", f"根拠は{A}"),
                    "refers_to": (f"{B}を参照", f"{A}から参照される")}.get(t, (f"{B}と{REL_LABEL.get(t, t)}", f"{A}と{REL_LABEL.get(t, t)}"))
            rel.setdefault(a, []).append(said[0])
            rel.setdefault(b, []).append(said[1])
        return [[Cell([(h, "head")], fill=FILL["head"]) for h in DOC_HEAD]] \
            + [[Cell(v if isinstance(v, str) else [(v, "text")]) for v in row + ["、".join(rel.get(d["doc_id"], []))]]
               for d, row in zip(self.r.m["docs"], self.r.doc_rows())]

    def values(self):
        """文のシート（読み方）の印に入れる数"""
        rows, qs = self.rows, self.r.qs
        steps = [s for r in rows for s in r["steps"]]
        on_rows = {q["id"] for q in qs if q.get("flow") in {r["id"] for r in rows}}
        return dict(zip(VALUES, [len(rows), len(steps), sum(len(v) for r in rows for v in r["files"].values()), sum(s["detail"]["inferred"] for s in steps),
                                 sum(e["inferred"] for r in rows for e in r["edges"]), len(qs), len(on_rows), len(qs) - len(on_rows)]))

    def sheets(self):
        """表の一覧（読み方は、テンプレートの文のシート）: 一覧の名前 → 行 → セル。1行目は、欄の名前"""
        return dict(zip(TABLES, (self.matrix(), self.node_list(), self.systems(), self.entries(), self.code_nodes(), self.helpers(), self.terms(), self.axis(), self.docs())))

    def texts(self):
        """表の一覧の中身を、文字だけで: {一覧の名前: [[セルの文字, …], …]}。点検（check_flow.py --stage matrix）が、書き出した xlsx の読み戻しと比べる"""
        return {name: [[clean(c.text) for c in row] for row in rows] for name, rows in self.sheets().items()}


def model(work, write=False):
    """表の一覧の中身を、文字だけで（Matrix.texts。自己点検が使う）"""
    return Matrix(work, write).texts()


def answer_rows(flows):
    """人が埋める欄の行の鍵 → いまの行の中身（取り込んだあとで変わったかを見る）: {一覧: {鍵: 文字}}"""
    out = {M: {}, NODES: {}}
    for fl in flows:
        F = flow_key(fl)
        out[M][F] = out[M][f"{F}｜処理時間"] = str(fl.get("name") or "")
        out[NODES].update({f"{F}｜{n.get('id')}": str(n.get("label") or "") for n in fl.get("nodes") or [] if isinstance(n, dict)})
    return out


def write(mx, path, template=None, kept=None):
    """テンプレートに流し込んで、本を書き出す（列の対応 <本の名前>.columns.json も、となりに書く）。
    処理時間の行を出すかは、テンプレートが決める（基準の表があれば出す）。kept: 人が埋める欄に戻す値（answers.json）"""
    tpl = template if isinstance(template, Template) else Template(template or ASSETS / TEMPLATE, SPEC)
    mx.timed = tpl.time is not None
    table = {k: (x["value"], "人が入れた時間の中央値から、引き直した" if x.get("from") == "calibrate" else "ブックで直した") for k, x in mx.times.table.items()} if mx.timed else None
    sheets = mx.sheets()
    return fill_template(tpl, SPEC, sheets, mx.values(), path, table=table, info={"timed": mx.timed}, keys={M: mx.row_keys}, kept=kept)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-w", "--work", required=True)
    ap.add_argument("--template", help=f"テンプレートの xlsx（既定: 前の書き出しで使ったもの。はじめは assets/{TEMPLATE}。決まりは references/template-spec.md）")
    ap.add_argument("--keep", help="人が値を入れたブック（お客様から戻ってきた xlsx）。処理時間・件数・基準の表の直しを取り込んでから、書き出す")
    ap.add_argument("--calibrate", action="store_true", help="人が入れた時間が 3 つ以上ある「動作の種類 × 量」は、その中央値を、基準の表に入れる")
    ap.add_argument("--accept", action="store_true", help="作業名が変わった箱の、人が入れた時間を、そのまま認める")
    ap.add_argument("--forget", action="store_true", help="無くなった箱の、人が入れた時間を、times.json から消す")
    ap.add_argument("--fresh", action="store_true", help="いま out にあるブックから、人が入れた値を取り込まない（読めないブックを、作り直すとき）")
    a = ap.parse_args()
    work = work_dir(a.work)
    blocking(work)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise SystemExit("openpyxl が無いので、業務マトリクス一覧.xlsx を作れない（pip3 install openpyxl）")
    again = "テンプレートが、決まりどおりでない（python3 scripts/table_template.py check で確かめる）"
    try:
        chosen = a.template or remembered(work / "out" / BOOK, ASSETS / TEMPLATE)
        if chosen and not a.template:
            print(f"テンプレート: 前の書き出しと同じ {pathlib.Path(chosen).name} を使う（既定の様式に戻すなら --template {ASSETS / TEMPLATE}）")
        tpl = Template(chosen or ASSETS / TEMPLATE, SPEC)
    except TemplateError as e:
        raise SystemExit(f"{again}:\n{e}")
    mx = Matrix(work)
    mx.timed, book, said = tpl.time is not None, mx.out / BOOK, []
    for p in [book] + ([pathlib.Path(a.keep)] if a.keep else []):
        lock = lock_of(p)
        if lock:                                                      # 開いたまま書き換えると、人が入れた値を取り込めず、上書きもできない
            raise SystemExit(f"{p.name} を表計算ソフトで開いている（{lock.name} がある）。閉じてから、もう一度")
    try:
        ans = Answers(work)
    except BadJson as e:
        raise SystemExit(str(e))
    now = answer_rows(mx.flows)
    for p in ([book] if book.exists() and not a.fresh else []) + ([pathlib.Path(a.keep)] if a.keep else []):   # 様式で足した、人が埋める欄
        try:
            got, notes = manual_back(p, SPEC, tpl)
        except TemplateError as e:
            raise SystemExit(f"{p.name}: {e}。人が入れた値を取り込めないので、上書きしない（取り込まずに作り直すなら --fresh）")
        except Exception as e:                                        # 壊れた xlsx は、どの例外で落ちるか決まっていない
            raise SystemExit(f"{p.name} を読めない（{type(e).__name__}: {str(e)[:60]}）。人が入れた値を取り込めないので、上書きしない。取り込まずに作り直すなら、--fresh を付ける")
        n = ans.take(BOOK, got, now)
        said += notes + ([f"人が埋める欄の値を取り込んだ: {p.name} から {n} 行（answers.json）"] if n else [])
    for canon in (M, NODES):
        gone, changed = answer_gaps(ans, BOOK, canon, now[canon])
        said += ([f"WARN  「{canon}」の人が埋める欄: 行が無くなった値 {len(gone)} 行（{'・'.join(gone[:6])}）。answers.json に残してある"] if gone else []) \
            + ([f"WARN  「{canon}」の人が埋める欄: 取り込んだあとで、行の中身（フロー名・作業名）が変わった {len(changed)} 行（{'・'.join(changed[:6])}）。値を見直す"] if changed else [])
    ans.save()                                                        # ブックより先に書く
    if mx.timed and not tpl.problems:
        base, t = tpl.time["values"], mx.times
        for F, was, now, lab, v in t.reconcile(mx.flows):               # フローを作り直して番号がずれていれば、作業名で付け替える（ブックを読む前に）
            said.append(f"人が入れた時間（{v} 分）を、番号の変わった箱に付け替えた: {F} {was} → {now}「{lab}」" if now else
                        f"WARN  人が入れた時間（{v} 分）を、times.json から外した: {F} {was}「{lab}」（同じ作業に、あとから入れた値がある／作業名が変わり、その番号を、いまは、ほかの作業が使っている。"
                        "要るなら、書き出し直したブックに、入れ直す）")
        if book.exists() and not a.fresh:                             # 人が入れた値を、消さない: いまあるブックから取り込んでから、書き出す
            try:
                said += pull(t, book, mx.flows, True, base)
            except Exception as e:                                    # 壊れた xlsx は、どの例外で落ちるか決まっていない
                raise SystemExit(f"out/{BOOK} を読めない（{type(e).__name__}: {str(e)[:60]}）。人が入れた値を取り込めないので、上書きしない。"
                                 "取り込まずに作り直すなら、--fresh を付ける")
        if a.keep:
            try:
                said += pull(t, a.keep, mx.flows, False, base)
            except Exception as e:
                raise SystemExit(f"{a.keep} を読めない（{type(e).__name__}: {str(e)[:60]}）")
        gone, renamed = t.gaps(mx.flows)
        if a.accept:
            for F, i, _v, _was, now in renamed:
                t.cells[(F, i)]["label"] = now
            said.append(f"作業名が変わった箱 {len(renamed)} 個の、人が入れた時間を、そのまま認めた")
            renamed = []
        if a.forget:
            for F, i, _v, _was, _now in gone:
                del t.cells[(F, i)]
            said.append(f"無くなった箱 {len(gone)} 個の、人が入れた時間を、times.json から消した")
            gone = []
        if a.calibrate:
            for k, cnt, mid, cur in calibration(mx, base):
                if mid != cur:
                    t.table[k] = {"action": k[0], "volume": k[1], "value": mid, "from": "calibrate"}
                    said.append(f"基準の表「{k[0]} × {k[1]}」を {cur} → {mid} 分に引き直した（人が入れた {cnt} セルの中央値）")
        said += [f"WARN  人が入れた時間（{v} 分）の箱が、もう無い: {F} {i}「{was}」。値は times.json に残してある（要らなければ --forget）" for F, i, v, was, _n in gone]
        said += [f"WARN  人が入れた時間（{v} 分）の作業名が変わった: {F} {i}「{was}」→「{now}」。値を見直す（そのままでよければ --accept）" for F, i, v, was, now in renamed]
        t.note(mx.flows)
        t.save()
    elif a.keep or a.calibrate or a.accept or a.forget:
        said.append(f"テンプレートに、基準の表（名前「{TIME_NAME}」の範囲）が無いので、処理時間は扱わない")
    try:
        write(mx, book, tpl, kept=ans.kept(BOOK))
    except TemplateError as e:
        raise SystemExit(f"{again}:\n{e}")
    for n in tpl.notes:
        print(f"テンプレート: {n}")
    if getattr(tpl, "unkept", 0):
        print(f"WARN  式のセル {tpl.unkept} 個に、計算の結果を書き足せなかった（Excel で開けば、計算される。計算をしない表示（プレビュー）では、0・空に見える）")
    for x in said:
        print(x)
    steps = sum(len(r["steps"]) for r in mx.rows)
    print(f"業務マトリクス一覧 → {book}（フロー {len(mx.rows)} 行・作業 {steps} セル・いちばん長い行 {max((len(r['steps']) for r in mx.rows), default=0)} 列）")
    if mx.timed:
        human = sum((r["id"], s["id"]) in mx.times.cells for r in mx.rows for s in r["steps"])
        print(f"処理時間: 人が入れた値 {human} セル・仮の値（基準の表を引く式）{steps - human} セル／件数 {sum(r['id'] in mx.times.counts for r in mx.rows)} 行／"
              f"基準の表の直し {len(mx.times.table)} か所")


if __name__ == "__main__":
    main()
